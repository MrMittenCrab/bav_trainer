"""Step 9M.2.4.1.1.1.24 — source-supported share-repurchase cash-use diagnostics."""

from __future__ import annotations

import copy
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    SHARE_REPURCHASE_COMPONENT_CATALOG,
    expand_share_repurchase_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.historical_expected import share_repurchase_expected_series
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.share_repurchase import (
    cash_after_ppe_capex_acquisitions_and_repurchases_applicable,
    compute_share_repurchase_series,
    resolve_share_repurchase_source,
    share_repurchase_applicable,
    share_repurchase_availability,
)
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_acquisition_cash import _add_acq
from core.tests.test_capex import P1, P2, _anchor, _dupont_row_by_label, _tiny
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "benchmark" / "fast_retailing" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _add_rp(
    fin: StandardizedFinancials,
    *,
    values=(-30.0, -40.0),
    label="Repurchase of common stock",
    concept="repurchase_of_common_stock",
    duplicate: bool = False,
    missing_period: bool = False,
    none_period: bool = False,
    on_balance_sheet: bool = False,
) -> LineItem:
    if missing_period:
        pay_values = {P1: values[0]}
    elif none_period:
        pay_values = {P1: values[0], P2: None}
    else:
        pay_values = {P1: values[0], P2: values[1]}
    item = _li(label, pay_values, concept=concept)
    target = fin.balance_sheet if on_balance_sheet else fin.cash_flow
    target.append(item)
    if duplicate and not on_balance_sheet:
        fin.cash_flow.append(
            _li("Repurchase duplicate", pay_values, concept=concept)
        )
    return item


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(SHARE_REPURCHASE_COMPONENT_CATALOG) == 3
    assert [f.order for f in SHARE_REPURCHASE_COMPONENT_CATALOG] == [132, 133, 134]
    specs = expand_share_repurchase_specs(periods, start_order=1000)
    assert len(specs) == 10
    assert {s.family_id for s in specs} == {
        "share_repurchase_outflow",
        "share_repurchase_to_revenue",
    }
    assert all(s.semantic_key.startswith("share_repurchase.") for s in specs)
    residual = expand_share_repurchase_specs(
        periods, start_order=1000, include_residual=True
    )
    assert len(residual) == 15
    assert {s.family_id for s in residual} == {
        "share_repurchase_outflow",
        "share_repurchase_to_revenue",
        "cash_after_ppe_capex_acquisitions_and_repurchases",
    }
    with pytest.raises(ValueError, match="duplicate"):
        expand_share_repurchase_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_share_repurchase_specs(list(reversed(periods)), start_order=1)


def test_unique_concept_resolution_and_renamed_label():
    fin = _tiny()
    _add_acq(fin)
    stored = _add_rp(fin)
    item = resolve_share_repurchase_source(fin)
    assert item is stored
    assert item.concept == "repurchase_of_common_stock"
    assert item.label == "Repurchase of common stock"
    assert share_repurchase_applicable(fin)
    avail = share_repurchase_availability(fin)
    assert avail.repurchase_of_common_stock is True
    assert avail.ambiguous is False
    assert cash_after_ppe_capex_acquisitions_and_repurchases_applicable(fin)

    renamed = _tiny()
    _add_acq(renamed)
    _add_rp(renamed, label="Common stock repurchase")
    assert resolve_share_repurchase_source(renamed).label == "Common stock repurchase"


def test_absent_label_only_duplicate_wrong_statement():
    absent = _tiny()
    assert not share_repurchase_applicable(absent)
    assert resolve_share_repurchase_source(absent) is None
    assert share_repurchase_availability(absent).repurchase_of_common_stock is False
    with pytest.raises(MissingLineError):
        compute_share_repurchase_series(
            absent, list(canonical_fiscal_periods(absent)), _anchor(absent)
        )
    assert ReferenceModelBuilder(absent).share_repurchase_specs == ()

    label_only = _tiny()
    _add_rp(label_only, concept="")
    assert not share_repurchase_applicable(label_only)
    assert resolve_share_repurchase_source(label_only) is None
    assert (
        resolve_line(
            label_only.cash_flow, "repurchase_of_common_stock", required=False
        ).item
        is None
    )
    assert ReferenceModelBuilder(label_only).share_repurchase_specs == ()

    dup = _tiny()
    _add_rp(dup, duplicate=True)
    avail = share_repurchase_availability(dup)
    assert avail.ambiguous is True
    assert avail.repurchase_of_common_stock is False
    assert not share_repurchase_applicable(dup)
    assert resolve_share_repurchase_source(dup) is None
    with pytest.raises(AmbiguousLineError):
        resolve_line(dup.cash_flow, "repurchase_of_common_stock", required=False)
    assert ReferenceModelBuilder(dup).share_repurchase_specs == ()

    wrong = _tiny()
    _add_rp(wrong, on_balance_sheet=True)
    assert not share_repurchase_applicable(wrong)
    assert resolve_share_repurchase_source(wrong) is None
    assert (
        resolve_line(
            wrong.balance_sheet, "repurchase_of_common_stock", required=False
        ).item
        is not None
    )


def test_no_substitution_from_sbc_treasury_share_count_or_withholding():
    fin = _tiny()
    fin.balance_sheet.append(
        _li("Treasury stock", {P1: -80.0, P2: -90.0}, concept="treasury_stock")
    )
    fin.cash_flow.append(
        _li(
            "Stock-based compensation expense",
            {P1: 5.0, P2: 6.0},
            concept="stock_based_compensation",
        )
    )
    fin.cash_flow.append(
        _li(
            "Proceeds from settlement of stock-based compensation",
            {P1: 7.0, P2: 8.0},
            concept="proceeds_from_stock_based_compensation",
        )
    )
    fin.cash_flow.append(
        _li(
            "Shares withheld related to net share settlement",
            {P1: -9.0, P2: -10.0},
            concept="shares_withheld_for_stock_based_compensation",
        )
    )
    fin.cash_flow.append(
        _li(
            "Payments for repurchase of common stock",
            {P1: -11.0, P2: -12.0},
            concept="payments_for_repurchase_of_common_stock",
        )
    )
    fin.cash_flow.append(
        _li("Dividends paid", {P1: -13.0, P2: -14.0}, concept="dividends_paid")
    )
    assert resolve_share_repurchase_source(fin) is None
    assert not share_repurchase_applicable(fin)
    assert share_repurchase_availability(fin).repurchase_of_common_stock is False


def test_sign_conversion_ratios_positive_negative_mixed_zero():
    fin = _tiny()
    _add_acq(fin, values=(-10.0, -20.0))
    _add_rp(fin, values=(-30.0, -40.0))
    series = compute_share_repurchase_series(
        fin, list(canonical_fiscal_periods(fin)), _anchor(fin)
    )
    assert series.payments_reported == (-30.0, -40.0)
    assert series.share_repurchase_outflow == (30.0, 40.0)
    assert series.share_repurchase_to_revenue[0] == pytest.approx(30.0 / 1000.0)
    assert series.cash_after_ppe_capex_acquisitions_and_repurchases == (
        80.0 - 100.0 - 10.0 - 30.0,
        90.0 - 120.0 - 20.0 - 40.0,
    )

    inflow = _tiny()
    _add_acq(inflow)
    _add_rp(inflow, values=(30.0, 40.0))
    series_in = compute_share_repurchase_series(
        inflow, list(canonical_fiscal_periods(inflow)), _anchor(inflow)
    )
    assert series_in.payments_reported == (30.0, 40.0)
    assert series_in.share_repurchase_outflow == (-30.0, -40.0)

    zero = _tiny()
    _add_acq(zero)
    _add_rp(zero, values=(0.0, -10.0))
    series_z = compute_share_repurchase_series(
        zero, list(canonical_fiscal_periods(zero)), _anchor(zero)
    )
    assert series_z.payments_reported == (0.0, -10.0)
    assert series_z.share_repurchase_outflow == (0.0, 10.0)
    assert series_z.share_repurchase_to_revenue[0] == pytest.approx(0.0)

    mapped = share_repurchase_expected_series(series)
    assert set(mapped) == {f.id for f in SHARE_REPURCHASE_COMPONENT_CATALOG}


def test_zero_revenue_undefined_ratio():
    fin = _tiny(revenue=(0.0, 2000.0))
    _add_acq(fin)
    _add_rp(fin, values=(-50.0, -80.0))
    series = compute_share_repurchase_series(
        fin, list(canonical_fiscal_periods(fin)), _anchor(fin)
    )
    assert series.share_repurchase_outflow == (50.0, 80.0)
    assert series.share_repurchase_to_revenue[0] == UNDEFINED_RATIO
    assert series.share_repurchase_to_revenue[1] == pytest.approx(80.0 / 2000.0)


def test_missing_none_versus_reported_zero():
    missing = _tiny()
    _add_rp(missing, missing_period=True)
    periods = list(canonical_fiscal_periods(missing))
    with pytest.raises(MissingHistoricalValueError, match="repurchase_of_common_stock"):
        compute_share_repurchase_series(missing, periods, _anchor(missing))

    none_period = _tiny()
    _add_rp(none_period, none_period=True)
    periods_n = list(canonical_fiscal_periods(none_period))
    with pytest.raises(MissingHistoricalValueError, match="repurchase_of_common_stock"):
        compute_share_repurchase_series(none_period, periods_n, _anchor(none_period))

    zero = _tiny()
    stored = _add_rp(zero, values=(0.0, -10.0))
    series = compute_share_repurchase_series(
        zero, list(canonical_fiscal_periods(zero)), _anchor(zero)
    )
    assert series.payments_reported[0] == 0.0
    assert stored.values[P1] == 0.0


def test_absent_ambiguous_cfo_capex_and_acquisition_omit_only_residual():
    no_cfo = _tiny(with_cfo=False)
    _add_acq(no_cfo)
    _add_rp(no_cfo)
    assert share_repurchase_applicable(no_cfo)
    assert not cash_after_ppe_capex_acquisitions_and_repurchases_applicable(no_cfo)
    series = compute_share_repurchase_series(
        no_cfo, list(canonical_fiscal_periods(no_cfo)), _anchor(no_cfo)
    )
    assert series.share_repurchase_outflow == (30.0, 40.0)
    assert series.cash_after_ppe_capex_acquisitions_and_repurchases is None
    mapped = share_repurchase_expected_series(series)
    assert set(mapped) == {"share_repurchase_outflow", "share_repurchase_to_revenue"}
    builder = ReferenceModelBuilder(no_cfo)
    assert {s.family_id for s in builder.share_repurchase_specs} == {
        "share_repurchase_outflow",
        "share_repurchase_to_revenue",
    }
    assert len(builder.share_repurchase_specs) == 4

    no_capex = _tiny(with_payments=False)
    _add_acq(no_capex)
    _add_rp(no_capex)
    assert share_repurchase_applicable(no_capex)
    assert not cash_after_ppe_capex_acquisitions_and_repurchases_applicable(no_capex)
    series_nc = compute_share_repurchase_series(
        no_capex, list(canonical_fiscal_periods(no_capex)), _anchor(no_capex)
    )
    assert series_nc.cash_after_ppe_capex_acquisitions_and_repurchases is None
    assert len(ReferenceModelBuilder(no_capex).share_repurchase_specs) == 4

    no_acq = _tiny()
    _add_rp(no_acq)
    assert share_repurchase_applicable(no_acq)
    assert not cash_after_ppe_capex_acquisitions_and_repurchases_applicable(no_acq)
    series_na = compute_share_repurchase_series(
        no_acq, list(canonical_fiscal_periods(no_acq)), _anchor(no_acq)
    )
    assert series_na.share_repurchase_outflow == (30.0, 40.0)
    assert series_na.cash_after_ppe_capex_acquisitions_and_repurchases is None
    assert len(ReferenceModelBuilder(no_acq).share_repurchase_specs) == 4

    amb_cfo = _tiny(duplicate_cfo=True, cfo_concept="operating_cash_flow")
    _add_acq(amb_cfo)
    _add_rp(amb_cfo)
    avail = share_repurchase_availability(amb_cfo)
    assert avail.operating_cash_flow_ambiguous is True
    assert share_repurchase_applicable(amb_cfo)
    assert not cash_after_ppe_capex_acquisitions_and_repurchases_applicable(amb_cfo)
    series_a = compute_share_repurchase_series(
        amb_cfo, list(canonical_fiscal_periods(amb_cfo)), _anchor(amb_cfo)
    )
    assert series_a.cash_after_ppe_capex_acquisitions_and_repurchases is None

    amb_capex = _tiny(duplicate=True)
    _add_acq(amb_capex)
    _add_rp(amb_capex)
    avail_c = share_repurchase_availability(amb_capex)
    assert avail_c.ppe_capex_ambiguous is True
    assert share_repurchase_applicable(amb_capex)
    assert not cash_after_ppe_capex_acquisitions_and_repurchases_applicable(amb_capex)

    amb_acq = _tiny()
    _add_acq(amb_acq, duplicate=True)
    _add_rp(amb_acq)
    avail_acq = share_repurchase_availability(amb_acq)
    assert avail_acq.acquisition_ambiguous is True
    assert share_repurchase_applicable(amb_acq)
    assert not cash_after_ppe_capex_acquisitions_and_repurchases_applicable(amb_acq)
    series_aa = compute_share_repurchase_series(
        amb_acq, list(canonical_fiscal_periods(amb_acq)), _anchor(amb_acq)
    )
    assert series_aa.cash_after_ppe_capex_acquisitions_and_repurchases is None


def test_independent_gating_without_share_history_or_sbc():
    fin = _tiny()
    _add_acq(fin)
    _add_rp(fin)
    assert fin.historical_shares is None
    assert "stock_based_compensation" not in {item.concept for item in fin.cash_flow}
    assert share_repurchase_applicable(fin)
    builder = ReferenceModelBuilder(fin)
    assert len(builder.share_repurchase_specs) == 6
    assert builder.per_share_specs == ()


def test_period_ordering_and_unchanged_input():
    fin = _tiny()
    _add_acq(fin)
    stored = _add_rp(fin, values=(-30.0, -40.0))
    reverse = [P2, P1]
    from core.model.financial_math import compute_anchor

    series = compute_share_repurchase_series(
        fin, reverse, compute_anchor(fin, reverse)
    )
    assert series.payments_reported == (-40.0, -30.0)
    assert series.share_repurchase_outflow == (40.0, 30.0)

    before = copy.deepcopy(stored.values)
    compute_share_repurchase_series(fin, [P1, P2], _anchor(fin))
    assert stored.values == before
    assert stored.concept == "repurchase_of_common_stock"


def test_workbook_gating_structure_formulas_notes_and_check(tmp_path):
    data = _tiny()
    _add_acq(data)
    _add_rp(data)
    trainer, answer = build_training_workbook(data, tmp_path / "RP_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.share_repurchase_series is not None
    assert len(builder.share_repurchase_specs) == 6
    assert {s.family_id for s in builder.share_repurchase_specs} == {
        "share_repurchase_outflow",
        "share_repurchase_to_revenue",
        "cash_after_ppe_capex_acquisitions_and_repurchases",
    }

    smap = load_semantic_map(answer)
    rp_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in SHARE_REPURCHASE_COMPONENT_CATALOG}
    ]
    assert len(rp_comps) == 6

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "SHARE REPURCHASE CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        reported_row = _dupont_row_by_label(
            ws, "Repurchase of common stock (reported)"
        )
        outflow_row = _dupont_row_by_label(
            ws, "Share-repurchase outflow (−reported)"
        )
        ratio_row = _dupont_row_by_label(ws, "Share-repurchase / Revenue")
        residual_row = _dupont_row_by_label(
            ws,
            "Operating cash after PP&E capex, acquisitions and repurchases",
        )
        assert isinstance(ws.cell(reported_row, 2).value, str)
        assert ws.cell(reported_row, 2).value.startswith("=")
        for col in (2, 3):
            outflow_cell = ws.cell(outflow_row, column=col)
            ratio_cell = ws.cell(ratio_row, column=col)
            residual_cell = ws.cell(residual_row, column=col)
            if path == answer:
                assert isinstance(outflow_cell.value, str) and outflow_cell.value.startswith(
                    "="
                )
                assert outflow_cell.value.replace(" ", "").startswith("=-")
                assert "IF(" in str(ratio_cell.value) and "NA()" in str(ratio_cell.value)
                assert residual_cell.value.replace(" ", "").count("-") >= 3
                note = (outflow_cell.comment.text or "") if outflow_cell.comment else ""
                assert "distribution" in note.lower()
                assert "dilution" in note.lower()
                assert "free cash flow" in note.lower()
                assert residual_cell.comment is not None
                residual_note = residual_cell.comment.text or ""
                assert "exceed" in residual_note.lower()
                assert "debt" in residual_note.lower()
                assert "free cash flow" in residual_note.lower()
                assert "mechanical" in residual_note.lower()
            else:
                assert outflow_cell.value is None
                assert ratio_cell.value is None
                assert residual_cell.value is None
                assert outflow_cell.comment is None
                assert ratio_cell.comment is None
                assert residual_cell.comment is None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in rp_comps if c.family_id == "share_repurchase_outflow")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Share-repurchase outflow" not in dumped


def test_source_row_fidelity_round_trip_and_immutability(tmp_path):
    fin = _tiny()
    _add_acq(fin)
    stored = _add_rp(fin, values=(-30.0, -40.0))
    original_index = fin.cash_flow.index(stored)
    resolved = resolve_line(
        fin.cash_flow, "repurchase_of_common_stock", required=True
    )
    assert resolved.index == original_index
    assert resolved.item is stored

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt = resolve_share_repurchase_source(restored)
    assert rt is not None
    assert rt.concept == "repurchase_of_common_stock"
    assert rt.values == stored.values

    trainer, answer = build_training_workbook(fin, tmp_path / "RP_ROW.xlsx")
    wb = load_workbook(answer, data_only=False)
    ws = wb["ALT DuPont"]
    reported_row = _dupont_row_by_label(
        ws, "Repurchase of common stock (reported)"
    )
    level_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
    expected_row = 7 + original_index
    assert f"'CashFlowStatement'!B{expected_row}" in level_f or (
        f"'Cash Flow Statement'!B{expected_row}" in level_f
    )
    wb.close()
    assert stored.concept == "repurchase_of_common_stock"
    assert stored.values == {P1: -30.0, P2: -40.0}
    assert trainer.exists() and answer.exists()


def test_demo_and_fast_retailing_omit_share_repurchase(tmp_path):
    demo = HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )
    assert not share_repurchase_applicable(demo)
    builder = ReferenceModelBuilder(demo)
    assert builder.share_repurchase_series is None
    assert builder.share_repurchase_specs == ()
    trainer, answer = build_training_workbook(demo, tmp_path / "RP_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312

    payload = json.loads(FR_JSON.read_text(encoding="utf-8"))
    fr = standardized_from_payload(payload)
    assert not share_repurchase_applicable(fr)
    fr_builder = ReferenceModelBuilder(fr)
    assert fr_builder.share_repurchase_specs == ()
    assert len(fr_builder.expected_specs) == 521
    assert "share_repurchase_outflow" not in {
        s.family_id for s in fr_builder.expected_specs
    }

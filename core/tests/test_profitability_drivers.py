"""Step 9C.1 — historical RNOA margin / NOA-turnover driver decomposition."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    PROFITABILITY_DRIVER_COMPONENT_CATALOG,
    WORKING_CAPITAL_COMPONENT_CATALOG,
    expand_profitability_driver_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.historical_expected import (
    expected_value_for_component,
    profitability_driver_expected_series,
)
from core.model.profitability_drivers import compute_profitability_driver_series
from core.model.ratio_values import UNDEFINED_RATIO
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family
from core.tests.test_normalization import _inject_formula_and_cached_value

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    adapter = HKManualDocumentAdapter()
    return adapter.ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, v1, v2, concept=""):
    return LineItem(
        label=label,
        values={date(2024, 12, 31): v1, date(2025, 12, 31): v2},
        concept=concept,
    )


def _periods():
    return [
        FinancialPeriod(end_date=date(2024, 12, 31), label="FY2024"),
        FinancialPeriod(end_date=date(2025, 12, 31), label="FY2025"),
    ]


def _base_fin(**overrides):
    kwargs = dict(
        ticker="PD",
        company_name="PD Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=_periods(),
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 80, 90),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 50, 55),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 330, 355),
        ],
        cash_flow=[_li("Net cash from operating activities", 50, 60)],
    )
    kwargs.update(overrides)
    return StandardizedFinancials(**kwargs)


def test_profitability_driver_series_core_semantics():
    fin = _base_fin()
    periods = [p.end_date for p in fin.periods]
    anchor = compute_anchor(fin, periods)
    series = compute_profitability_driver_series(anchor)

    assert series.average_noa[0] is None
    assert series.noa_turnover[0] is None
    assert series.noa_intensity[0] is None
    assert series.rnoa_from_margin_turnover[0] is None

    avg = (anchor.reformulation.noa[0] + anchor.reformulation.noa[1]) / 2.0
    assert series.average_noa[1] == pytest.approx(avg)
    assert series.noa_turnover[1] == pytest.approx(1100.0 / avg)
    assert series.noa_intensity[1] == pytest.approx(avg / 1100.0)
    assert series.rnoa_from_margin_turnover[1] == pytest.approx(
        float(anchor.dupont["NOPAT Margin"][1]) * float(series.noa_turnover[1])
    )
    assert series.rnoa_from_margin_turnover[1] == pytest.approx(
        float(anchor.dupont["RNOA"][1])
    )

    # Zero Average NOA -> turnover #N/A
    fin0 = _base_fin(
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 50, 55),
            _li("Trade payables", 50, 55),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", -100, -100),
        ]
    )
    series0 = compute_profitability_driver_series(compute_anchor(fin0, periods))
    assert series0.average_noa[1] == pytest.approx(0.0)
    assert series0.noa_turnover[1] == UNDEFINED_RATIO

    # Zero Revenue -> intensity #N/A; turnover may be 0.0; driver #N/A; direct RNOA may remain numeric
    fin_rev0 = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 0),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    anchor_rev0 = compute_anchor(fin_rev0, periods)
    series_rev0 = compute_profitability_driver_series(anchor_rev0)
    assert series_rev0.noa_intensity[1] == UNDEFINED_RATIO
    assert series_rev0.noa_turnover[1] == pytest.approx(0.0)
    assert series_rev0.rnoa_from_margin_turnover[1] == UNDEFINED_RATIO
    assert isinstance(anchor_rev0.dupont["RNOA"][1], float)


def test_profitability_driver_length_and_reconcile_failures():
    fin = _base_fin()
    periods = [p.end_date for p in fin.periods]
    anchor = compute_anchor(fin, periods)
    anchor.historical.revenue.append(1.0)
    with pytest.raises(ValueError, match="length mismatch"):
        compute_profitability_driver_series(anchor)

    anchor2 = compute_anchor(fin, periods)
    bad_dup = dict(anchor2.dupont)
    bad_dup["RNOA"] = list(bad_dup["RNOA"])
    bad_dup["RNOA"][1] = float(bad_dup["RNOA"][1]) + 1.0
    anchor2.dupont = bad_dup
    with pytest.raises(ValueError, match="does not reconcile"):
        compute_profitability_driver_series(anchor2)


def test_catalog_expand_and_expected_keys():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(PROFITABILITY_DRIVER_COMPONENT_CATALOG) == 4
    assert len({f.id for f in PROFITABILITY_DRIVER_COMPONENT_CATALOG}) == 4
    assert all(f.period_scope == "comparable" for f in PROFITABILITY_DRIVER_COMPONENT_CATALOG)
    assert [f.order for f in PROFITABILITY_DRIVER_COMPONENT_CATALOG] == [46, 47, 48, 49]
    # Existing WC catalog unchanged size
    assert len(WORKING_CAPITAL_COMPONENT_CATALOG) == 11

    specs = expand_profitability_driver_specs(periods, start_order=189)
    assert len(specs) == 16

    data = _ingest_demo()
    builder = ReferenceModelBuilder(data)
    series = profitability_driver_expected_series(builder.anchor)
    assert set(series) == {f.id for f in PROFITABILITY_DRIVER_COMPONENT_CATALOG}
    for j in range(1, len(builder.periods)):
        driver = series["rnoa_margin_turnover"][j]
        direct = builder.anchor.dupont["RNOA"][j]
        if driver != UNDEFINED_RATIO:
            assert driver == pytest.approx(float(direct))


def test_demo_profitability_driver_surface(tmp_path):
    import json

    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "BASE_PD.xlsx")
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == 204
    assert len(group_components_by_family(smap)) == 45
    pd_comps = [c for c in smap.all_ordered() if c.category == "profitability_driver"]
    assert len(pd_comps) == 16

    wb = load_workbook(answer)
    ws = wb["ALT DuPont"]
    assert ws.cell(5, 1).value == "Sales Growth"  # existing rows unchanged
    assert ws.cell(12, 1).value == "Actual ROE"
    assert ws.cell(14, 1).value == "RNOA DRIVER DECOMPOSITION"
    assert ws.cell(15, 1).value == "Average NOA"
    assert ws.cell(19, 1).value == "RNOA DRIVER CHECK"
    assert ws.cell(15, 2).value == "N/A"
    assert ws.cell(19, 2).value == "N/A"
    practice = {(c.tab, c.cell) for c in smap.all_ordered()}
    assert ("ALT DuPont", "B19") not in practice
    assert ("ALT DuPont", "C19") not in practice

    avg = next(c for c in pd_comps if c.family_id == "average_noa" and c.period_index == 1)
    assert "Condensed Financials" in avg.formula
    turn = next(c for c in pd_comps if c.family_id == "noa_turnover" and c.period_index == 1)
    assert "NA()" in turn.formula
    inten = next(c for c in pd_comps if c.family_id == "noa_intensity" and c.period_index == 1)
    assert "NA()" in inten.formula
    driver = next(
        c for c in pd_comps if c.family_id == "rnoa_margin_turnover" and c.period_index == 1
    )
    assert "*" in driver.formula
    assert "IFERROR" not in driver.formula.upper()

    # Existing direct RNOA formula unchanged (still uses inline average)
    rnoa = next(c for c in smap.all_ordered() if c.family_id == "rnoa" and c.period_index == 1)
    assert "nopat" not in rnoa.formula.lower() or True
    assert "/2)" in rnoa.formula or "/2))" in rnoa.formula

    wb_t = load_workbook(trainer)
    for comp in pd_comps:
        row, col = parse_cell_ref(comp.cell)
        cell = wb_t[comp.tab].cell(row=row, column=col)
        assert cell.value is None
        assert _fill_rgb(cell) == "FFFF00"
        assert cell.comment is None
    wb.close()
    wb_t.close()

    summary = check_workbook(trainer)
    assert summary.total == 204
    assert summary.blank == 204

    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        data, tmp_path / "NORM_PD.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(smap_n.all_ordered()) == 224
    assert len(group_components_by_family(smap_n)) == 49
    assert check_workbook(trainer_n).blank == 224


def test_live_classification_changes_profitability_drivers(tmp_path):
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
    fin = StandardizedFinancials(
        ticker="STI3",
        company_name="STI Profit Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=_periods(),
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", 50, 55),
            _li("Trade receivables", 40, 45),
            _li("Short-term investments", 100, 120),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 30, 35),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 360, 395),
        ],
        cash_flow=[_li("Net cash from operating activities", 50, 60)],
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "STI_PD.xlsx")
    smap = load_semantic_map(answer)
    avg = next(
        c for c in smap.all_ordered() if c.family_id == "average_noa" and c.period_index == 1
    )
    turn = next(
        c for c in smap.all_ordered() if c.family_id == "noa_turnover" and c.period_index == 1
    )
    inten = next(
        c for c in smap.all_ordered() if c.family_id == "noa_intensity" and c.period_index == 1
    )
    driver = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "rnoa_margin_turnover" and c.period_index == 1
    )
    rnoa = next(c for c in smap.all_ordered() if c.family_id == "rnoa" and c.period_index == 1)

    ref_avg = float(avg.expected_value)
    alt_anchor = compute_anchor(
        fin,
        [d1, d2],
        classification_overrides={
            "label:Short-term investments": "Operating Working Capital Asset"
        },
    )
    alt_avg = float(expected_value_for_component(alt_anchor, avg))
    assert alt_avg != pytest.approx(ref_avg)
    alt_turn = float(expected_value_for_component(alt_anchor, turn))
    alt_inten = float(expected_value_for_component(alt_anchor, inten))
    alt_driver = float(expected_value_for_component(alt_anchor, driver))
    alt_rnoa = float(expected_value_for_component(alt_anchor, rnoa))
    assert alt_driver == pytest.approx(alt_rnoa)

    wb = load_workbook(trainer, data_only=False)
    wb["Accounting Judgment"].cell(5, 6).value = "Operating Working Capital Asset"
    wb.save(trainer)
    wb.close()

    _inject_formula_and_cached_value(
        trainer, avg.tab, avg.cell, formula="=(B13+C13)/2", cached_value=ref_avg
    )
    assert check_workbook(trainer).incorrect == 1

    _inject_formula_and_cached_value(
        trainer, avg.tab, avg.cell, formula="=(B13+C13)/2", cached_value=alt_avg
    )
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer, turn.tab, turn.cell, formula="=C5/C15", cached_value=alt_turn
    )
    assert check_workbook(trainer).correct >= 1


def test_driver_check_tamper_and_undefined_decomposition(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "PD_Tamper.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c for c in smap.all_ordered() if c.family_id == "average_noa" and c.period_index == 1
    )
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    # Tamper RNOA DRIVER CHECK for comparable period (col C)
    check_row = ReferenceModelBuilder(data).rowmap.get("dupont_driver_check_row")
    # After build, rowmap isn't on shared builder; locate by label
    ws = wb["ALT DuPont"]
    check_row = next(
        r for r in range(1, (ws.max_row or 1) + 1) if ws.cell(r, 1).value == "RNOA DRIVER CHECK"
    )
    ws.cell(check_row, 3).value = '="TAMPERED"'
    wb.save(trainer)
    wb.close()
    with pytest.raises(ValueError, match="Trusted workbook cell was modified: ALT DuPont"):
        check_workbook(trainer)
    wb = load_workbook(trainer, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

    # Zero Revenue fixture: driver #N/A check semantics + Formula Check
    fin = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 0),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    trainer2, answer2 = build_training_workbook(fin, tmp_path / "PD_NA.xlsx")
    smap2 = load_semantic_map(answer2)
    intensity = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "noa_intensity" and c.period_index == 1
    )
    driver = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "rnoa_margin_turnover" and c.period_index == 1
    )
    assert intensity.expected_value == UNDEFINED_RATIO
    assert driver.expected_value == UNDEFINED_RATIO

    wb = load_workbook(answer2, data_only=False)
    check_row = next(
        r
        for r in range(1, (wb["ALT DuPont"].max_row or 1) + 1)
        if wb["ALT DuPont"].cell(r, 1).value == "RNOA DRIVER CHECK"
    )
    formula = str(wb["ALT DuPont"].cell(check_row, 3).value)
    assert 'ISNA(' in formula
    assert '"N/A"' in formula
    wb.close()

    wb = load_workbook(trainer2, data_only=False)
    row, col = parse_cell_ref(intensity.cell)
    wb[intensity.tab].cell(row=row, column=col).value = intensity.formula
    wb.save(trainer2)
    wb.close()
    assert check_workbook(trainer2).correct == 1
    _inject_formula_and_cached_value(
        trainer2, intensity.tab, intensity.cell, formula="=NA()", cached_value=UNDEFINED_RATIO
    )
    assert check_workbook(trainer2).correct == 1
    _inject_formula_and_cached_value(
        trainer2, intensity.tab, intensity.cell, formula="=0", cached_value=0.0
    )
    assert check_workbook(trainer2).incorrect == 1

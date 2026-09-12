"""Step 9B.1 — historical working-capital diagnostics."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    WORKING_CAPITAL_COMPONENT_CATALOG,
    expand_working_capital_specs,
)
from core.engine.reference_model import WORKING_CAPITAL_SHEET, ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.historical_expected import (
    expected_value_for_component,
    working_capital_expected_series,
)
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.working_capital import (
    compute_working_capital_series,
    working_capital_applicable,
)
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family
from core.tests.test_normalization import _inject_formula_and_cached_value

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    adapter = HKManualDocumentAdapter()
    return adapter.ingest([DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)])


def _li(label, v1, v2, concept=""):
    return LineItem(
        label=label,
        values={date(2024, 12, 31): v1, date(2025, 12, 31): v2},
        concept=concept,
    )


def _periods():
    return [
        FinancialPeriod(end_date=date(2024, 12, 31), label="FY2024"),
        FinancialPeriod(end_date=date(2025, 12, 31), label="FY2025"),
    ]


def _base_wc_fin(**overrides):
    kwargs = dict(
        ticker="WC",
        company_name="WC Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=_periods(),
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 80, 90),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 50, 55),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 330, 355),
        ],
        cash_flow=[_li("Net cash from operating activities", 50, 60)],
    )
    kwargs.update(overrides)
    return StandardizedFinancials(**kwargs)


def test_working_capital_series_core_semantics():
    fin = _base_wc_fin()
    periods = [p.end_date for p in fin.periods]
    anchor = compute_anchor(fin, periods)
    series = compute_working_capital_series(anchor)

    assert working_capital_applicable(anchor) is True
    assert series.owca_to_revenue[0] == pytest.approx(80.0 / 1000.0)
    assert series.owcl_to_revenue[0] == pytest.approx(50.0 / 1000.0)
    assert series.nowc_to_revenue[0] == pytest.approx(30.0 / 1000.0)
    assert series.revenue_change[0] is None
    assert series.nowc_change[0] is None
    assert series.incremental_nowc_to_revenue_change[0] is None
    assert series.revenue_change[1] == pytest.approx(100.0)
    assert series.nowc_change[1] == pytest.approx(5.0)
    assert series.incremental_nowc_to_revenue_change[1] == pytest.approx(0.05)
    # Driver bridge
    assert series.owca_change[0] is None
    assert series.owcl_change[0] is None
    assert series.nowc_change_from_components[0] is None
    assert series.incremental_owca_to_revenue_change[0] is None
    assert series.incremental_owcl_to_revenue_change[0] is None
    assert series.owca_change[1] == pytest.approx(10.0)  # 80 -> 90
    assert series.owcl_change[1] == pytest.approx(5.0)  # 50 -> 55
    assert series.nowc_change_from_components[1] == pytest.approx(5.0)
    assert series.nowc_change_from_components[1] == pytest.approx(series.nowc_change[1])
    assert series.incremental_owca_to_revenue_change[1] == pytest.approx(0.10)
    assert series.incremental_owcl_to_revenue_change[1] == pytest.approx(0.05)

    # Zero revenue -> #N/A intensity
    fin0 = _base_wc_fin(
        income_statement=[
            _li("Revenue", 0, 1100),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    series0 = compute_working_capital_series(compute_anchor(fin0, periods))
    assert series0.owca_to_revenue[0] == UNDEFINED_RATIO
    assert series0.owcl_to_revenue[0] == UNDEFINED_RATIO
    assert series0.nowc_to_revenue[0] == UNDEFINED_RATIO

    # Zero numerator with nonzero revenue -> 0.0
    fin_z = _base_wc_fin(
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 0, 0),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 0, 0),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 300, 320),
        ]
    )
    # Still applicable? all OWCA/OWCL zero -> not applicable, but compute still works
    assert working_capital_applicable(compute_anchor(fin_z, periods)) is False
    series_z = compute_working_capital_series(compute_anchor(fin_z, periods))
    assert series_z.owca_to_revenue[0] == pytest.approx(0.0)
    assert series_z.owcl_to_revenue[0] == pytest.approx(0.0)
    assert series_z.nowc_to_revenue[0] == pytest.approx(0.0)

    # Zero ΔRevenue -> incremental #N/A; negative ΔNOWC retained
    fin_flat = _base_wc_fin(
        income_statement=[
            _li("Revenue", 1000, 1000),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 90, 80),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 50, 55),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 340, 345),
        ],
    )
    series_flat = compute_working_capital_series(compute_anchor(fin_flat, periods))
    assert series_flat.revenue_change[1] == pytest.approx(0.0)
    assert series_flat.nowc_change[1] == pytest.approx(-15.0)
    assert series_flat.incremental_nowc_to_revenue_change[1] == UNDEFINED_RATIO
    assert series_flat.incremental_owca_to_revenue_change[1] == UNDEFINED_RATIO
    assert series_flat.incremental_owcl_to_revenue_change[1] == UNDEFINED_RATIO

    # Negative ΔRevenue retains signed ratio (not abs)
    fin_neg = _base_wc_fin(
        income_statement=[
            _li("Revenue", 1100, 1000),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    series_neg = compute_working_capital_series(compute_anchor(fin_neg, periods))
    assert series_neg.revenue_change[1] == pytest.approx(-100.0)
    assert series_neg.incremental_nowc_to_revenue_change[1] == pytest.approx(
        series_neg.nowc_change[1] / -100.0
    )


def test_working_capital_length_mismatch_fails():
    fin = _base_wc_fin()
    periods = [p.end_date for p in fin.periods]
    anchor = compute_anchor(fin, periods)
    # Corrupt reformulation length via a shallow copy of series is hard; mutate revenue.
    anchor.historical.revenue.append(1.0)  # type: ignore[attr-defined]
    with pytest.raises(ValueError, match="length mismatch"):
        compute_working_capital_series(anchor)


def test_catalog_expand_and_expected_keys():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(WORKING_CAPITAL_COMPONENT_CATALOG) == 11
    assert len({f.id for f in WORKING_CAPITAL_COMPONENT_CATALOG}) == 11
    assert [f.order for f in WORKING_CAPITAL_COMPONENT_CATALOG] == list(range(35, 46))
    # First six Step 9B.1 IDs unchanged
    assert [f.id for f in WORKING_CAPITAL_COMPONENT_CATALOG[:6]] == [
        "owca_to_revenue",
        "owcl_to_revenue",
        "nowc_to_revenue",
        "revenue_change",
        "nowc_change",
        "incremental_nowc_to_revenue_change",
    ]
    specs = expand_working_capital_specs(periods, start_order=142)
    assert len(specs) == 47
    assert len([s for s in specs if s.family_id in {
        "owca_to_revenue", "owcl_to_revenue", "nowc_to_revenue"
    }]) == 15
    assert len([s for s in specs if s.family_id in {
        "revenue_change", "nowc_change", "incremental_nowc_to_revenue_change"
    }]) == 12
    assert len([s for s in specs if s.family_id in {
        "owca_change",
        "owcl_change",
        "nowc_change_from_components",
        "incremental_owca_to_revenue_change",
        "incremental_owcl_to_revenue_change",
    }]) == 20

    data = _ingest_demo()
    builder = ReferenceModelBuilder(data)
    series = working_capital_expected_series(builder.anchor)
    assert set(series) == {f.id for f in WORKING_CAPITAL_COMPONENT_CATALOG}
    for j in range(1, len(builder.periods)):
        assert series["nowc_change_from_components"][j] == pytest.approx(
            series["nowc_change"][j]
        )


def test_demo_working_capital_surface(tmp_path):
    import json

    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "BASE_WC.xlsx")
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == 188
    assert len(group_components_by_family(smap)) == 41
    wc = [c for c in smap.all_ordered() if c.category == "working_capital"]
    assert len(wc) == 47
    wb = load_workbook(answer)
    assert WORKING_CAPITAL_SHEET in wb.sheetnames
    ws = wb[WORKING_CAPITAL_SHEET]
    assert ws["A1"].value == "Working Capital Analysis"
    assert "diagnostics, not automatic judgments" in str(ws["A3"].value).lower()
    assert "seasonality" in str(ws["A3"].value).lower()
    # Source links to Condensed
    assert str(ws.cell(6, 2).value).startswith("='Condensed Financials'!")
    assert str(ws.cell(7, 2).value).startswith("='Condensed Financials'!")
    # First-period change cells are N/A text, not practice
    assert ws.cell(14, 2).value == "N/A"
    assert ws.cell(15, 2).value == "N/A"
    assert ws.cell(16, 2).value == "N/A"
    assert ws.cell(18, 1).value == "WORKING-CAPITAL DRIVER DECOMPOSITION"
    assert ws.cell(19, 2).value == "N/A"
    assert ws.cell(24, 2).value == "N/A"
    assert ws.cell(24, 1).value == "DRIVER DECOMPOSITION CHECK"
    practice = {(c.tab, c.cell) for c in smap.all_ordered()}
    assert (WORKING_CAPITAL_SHEET, "B14") not in practice
    assert (WORKING_CAPITAL_SHEET, "B24") not in practice
    assert (WORKING_CAPITAL_SHEET, "C24") not in practice
    # Intensity uses NA(); incremental uses NA(); no IFERROR
    for comp in wc:
        assert "IFERROR" not in comp.formula.upper()
        if comp.family_id in {"owca_to_revenue", "owcl_to_revenue", "nowc_to_revenue"}:
            assert "NA()" in comp.formula
        if comp.family_id in {
            "incremental_nowc_to_revenue_change",
            "incremental_owca_to_revenue_change",
            "incremental_owcl_to_revenue_change",
        }:
            assert "NA()" in comp.formula
            assert "ABS(" not in comp.formula.upper()
    bridge = next(
        c for c in wc if c.family_id == "nowc_change_from_components" and c.period_index == 1
    )
    assert "-" in bridge.formula
    wb_t = load_workbook(trainer)
    assert WORKING_CAPITAL_SHEET in wb_t.sheetnames
    for comp in wc:
        row, col = parse_cell_ref(comp.cell)
        cell = wb_t[comp.tab].cell(row=row, column=col)
        assert cell.value is None
        assert _fill_rgb(cell) == "FFFF00"
    wb.close()
    wb_t.close()

    summary = check_workbook(trainer)
    assert summary.total == 188
    assert summary.blank == 188

    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        data, tmp_path / "NORM_WC.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(smap_n.all_ordered()) == 208
    assert len(group_components_by_family(smap_n)) == 45
    assert check_workbook(trainer_n).blank == 208


def test_working_capital_absent_when_owca_owcl_zero(tmp_path):
    fin = _base_wc_fin(
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Property, plant and equipment", 400, 420),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 300, 320),
        ]
    )
    builder = ReferenceModelBuilder(fin)
    assert builder.working_capital_specs == ()
    assert builder.working_capital_series is None
    trainer, answer = build_training_workbook(fin, tmp_path / "NoWC.xlsx")
    wb = load_workbook(answer)
    assert WORKING_CAPITAL_SHEET not in wb.sheetnames
    wb.close()
    # Historical + quality still present (CFO present)
    smap = load_semantic_map(answer)
    assert any(c.family_id == "nopat_fy" for c in smap.all_ordered())
    assert any(c.family_id == "operating_cash_flow_link" for c in smap.all_ordered())
    assert not any(c.category == "working_capital" for c in smap.all_ordered())


def test_trusted_source_link_tamper_rejects_before_recolor(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "WC_Tamper.xlsx")
    smap = load_semantic_map(answer)
    comp = next(c for c in smap.all_ordered() if c.family_id == "owca_to_revenue")
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    # Tamper trusted Revenue source link
    wb[WORKING_CAPITAL_SHEET].cell(6, 2).value = 1
    wb.save(trainer)
    wb.close()
    with pytest.raises(
        ValueError, match="Trusted workbook cell was modified: Working Capital Analysis"
    ):
        check_workbook(trainer)
    wb = load_workbook(trainer, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

    # Equivalent formula on practice cell is allowed
    trainer2, answer2 = build_training_workbook(data, tmp_path / "WC_Eq.xlsx")
    smap2 = load_semantic_map(answer2)
    comp2 = next(c for c in smap2.all_ordered() if c.family_id == "owca_to_revenue")
    _inject_formula_and_cached_value(
        trainer2,
        comp2.tab,
        comp2.cell,
        formula="=B7/B6",
        cached_value=float(comp2.expected_value),
    )
    assert check_workbook(trainer2).correct == 1


def test_live_classification_changes_working_capital_check(tmp_path):
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
    fin = StandardizedFinancials(
        ticker="STI",
        company_name="STI Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=_periods(),
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", 50, 55),
            _li("Trade receivables", 40, 45),
            _li("Short-term investments", 100, 120),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 30, 35),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 360, 395),
        ],
        cash_flow=[_li("Net cash from operating activities", 50, 60)],
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "STI_WC.xlsx")
    builder = ReferenceModelBuilder(fin)
    assert any(
        c.topic.startswith("Short-term investment") for c in builder.judgment_cases
    )
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "owca_to_revenue" and c.period_index == 1
    )
    ref_expected = float(comp.expected_value)

    alt_anchor = compute_anchor(
        fin,
        [d1, d2],
        classification_overrides={"label:Short-term investments": "Operating Working Capital Asset"},
    )
    assert alt_anchor.reformulation.category_totals["Operating Working Capital Asset"][1] != (
        builder.anchor.reformulation.category_totals["Operating Working Capital Asset"][1]
    )
    alt_expected = float(
        expected_value_for_component(alt_anchor, comp)
    )
    assert alt_expected != pytest.approx(ref_expected)

    wb = load_workbook(trainer, data_only=False)
    wb["Accounting Judgment"].cell(5, 6).value = "Operating Working Capital Asset"
    wb.save(trainer)
    wb.close()

    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula="=B7/B6",
        cached_value=alt_expected,
    )
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula="=B7/B6",
        cached_value=ref_expected,
    )
    assert check_workbook(trainer).incorrect == 1


def test_working_capital_check_na_and_numeric(tmp_path):
    # Zero revenue intensity
    fin = _base_wc_fin(
        income_statement=[
            _li("Revenue", 0, 1100),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "WC_NA.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "owca_to_revenue" and c.period_index == 0
    )
    assert comp.expected_value == UNDEFINED_RATIO
    assert "NA()" in comp.formula

    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula="=NA()", cached_value=UNDEFINED_RATIO
    )
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula="=0", cached_value=0.0
    )
    assert check_workbook(trainer).incorrect == 1

    # Flat revenue -> incremental #N/A
    fin2 = _base_wc_fin(
        income_statement=[
            _li("Revenue", 1000, 1000),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    trainer2, answer2 = build_training_workbook(fin2, tmp_path / "WC_IncrNA.xlsx")
    smap2 = load_semantic_map(answer2)
    incr = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "incremental_nowc_to_revenue_change" and c.period_index == 1
    )
    assert incr.expected_value == UNDEFINED_RATIO
    wb = load_workbook(trainer2, data_only=False)
    row, col = parse_cell_ref(incr.cell)
    wb[incr.tab].cell(row=row, column=col).value = incr.formula
    wb.save(trainer2)
    wb.close()
    assert check_workbook(trainer2).correct == 1
    _inject_formula_and_cached_value(
        trainer2, incr.tab, incr.cell, formula="=NA()", cached_value=UNDEFINED_RATIO
    )
    assert check_workbook(trainer2).correct == 1
    _inject_formula_and_cached_value(
        trainer2, incr.tab, incr.cell, formula="=0", cached_value=0.0
    )
    assert check_workbook(trainer2).incorrect == 1

    # Ordinary numeric exact/equivalent/wrong
    data = _ingest_demo()
    trainer3, answer3 = build_training_workbook(data, tmp_path / "WC_Num.xlsx")
    smap3 = load_semantic_map(answer3)
    num = next(c for c in smap3.all_ordered() if c.family_id == "nowc_to_revenue")
    wb = load_workbook(trainer3, data_only=False)
    row, col = parse_cell_ref(num.cell)
    wb[num.tab].cell(row=row, column=col).value = num.formula
    wb.save(trainer3)
    wb.close()
    assert check_workbook(trainer3).correct == 1
    _inject_formula_and_cached_value(
        trainer3,
        num.tab,
        num.cell,
        formula="=B9/B6",
        cached_value=float(num.expected_value),
    )
    assert check_workbook(trainer3).correct == 1
    _inject_formula_and_cached_value(
        trainer3,
        num.tab,
        num.cell,
        formula="=B9/B6",
        cached_value=float(num.expected_value) + 10.0,
    )
    assert check_workbook(trainer3).incorrect == 1


def test_driver_bridge_semantics_and_reconciliation_failure():
    from dataclasses import replace

    periods = [p.end_date for p in _periods()]
    # Positive OWCA/OWCL bridge: 100->130 OWCA, 40->50 OWCL => bridge +20
    fin = _base_wc_fin(
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 100, 130),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 40, 50),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 360, 400),
        ]
    )
    series = compute_working_capital_series(compute_anchor(fin, periods))
    assert series.owca_change[1] == pytest.approx(30.0)
    assert series.owcl_change[1] == pytest.approx(10.0)
    assert series.nowc_change[1] == pytest.approx(20.0)
    assert series.nowc_change_from_components[1] == pytest.approx(20.0)

    # Negative side changes retained
    fin_neg = _base_wc_fin(
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 130, 100),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 50, 40),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 380, 380),
        ]
    )
    series_neg = compute_working_capital_series(compute_anchor(fin_neg, periods))
    assert series_neg.owca_change[1] == pytest.approx(-30.0)
    assert series_neg.owcl_change[1] == pytest.approx(-10.0)
    assert series_neg.nowc_change_from_components[1] == pytest.approx(-20.0)

    # Zero side numerator with nonzero ΔRevenue -> 0.0
    fin_z = _base_wc_fin(
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 80, 80),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 50, 50),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 330, 350),
        ]
    )
    series_z = compute_working_capital_series(compute_anchor(fin_z, periods))
    assert series_z.incremental_owca_to_revenue_change[1] == pytest.approx(0.0)
    assert series_z.incremental_owcl_to_revenue_change[1] == pytest.approx(0.0)

    # OWCL-side unit: change OWCL holding OWCA constant
    fin_owcl = _base_wc_fin(
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 80, 80),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 40, 55),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 340, 345),
        ]
    )
    series_owcl = compute_working_capital_series(compute_anchor(fin_owcl, periods))
    assert series_owcl.owca_change[1] == pytest.approx(0.0)
    assert series_owcl.owcl_change[1] == pytest.approx(15.0)
    assert series_owcl.nowc_change_from_components[1] == pytest.approx(-15.0)
    assert series_owcl.incremental_owcl_to_revenue_change[1] == pytest.approx(0.15)

    # Deliberately inconsistent NOWC identity fails
    anchor = compute_anchor(fin, periods)
    bad_reform = replace(anchor.reformulation, nowc=(0.0, 999.0))
    anchor.reformulation = bad_reform
    with pytest.raises(ValueError, match="does not reconcile"):
        compute_working_capital_series(anchor)


def test_live_classification_changes_owca_driver_check(tmp_path):
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
    fin = StandardizedFinancials(
        ticker="STI2",
        company_name="STI Driver Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=_periods(),
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", 50, 55),
            _li("Trade receivables", 40, 45),
            _li("Short-term investments", 100, 120),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 30, 35),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 360, 395),
        ],
        cash_flow=[_li("Net cash from operating activities", 50, 60)],
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "STI_Driver.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "owca_change" and c.period_index == 1
    )
    ref_expected = float(comp.expected_value)

    alt_anchor = compute_anchor(
        fin,
        [d1, d2],
        classification_overrides={
            "label:Short-term investments": "Operating Working Capital Asset"
        },
    )
    alt_expected = float(expected_value_for_component(alt_anchor, comp))
    assert alt_expected != pytest.approx(ref_expected)

    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    wb = load_workbook(trainer, data_only=False)
    wb["Accounting Judgment"].cell(5, 6).value = "Operating Working Capital Asset"
    wb.save(trainer)
    wb.close()
    # Stale reference cached value under alternative treatment -> red
    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula="=B7-A7",
        cached_value=ref_expected,
    )
    assert check_workbook(trainer).incorrect == 1

    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula="=B7-A7",
        cached_value=alt_expected,
    )
    assert check_workbook(trainer).correct == 1


def test_driver_check_tamper_and_incremental_side_na(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "WC_DriverTamper.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c for c in smap.all_ordered() if c.family_id == "owca_change" and c.period_index == 1
    )
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    # Tamper trusted DRIVER DECOMPOSITION CHECK (col C = period index 1)
    wb[WORKING_CAPITAL_SHEET].cell(24, 3).value = '="TAMPERED"'
    wb.save(trainer)
    wb.close()
    with pytest.raises(
        ValueError, match="Trusted workbook cell was modified: Working Capital Analysis"
    ):
        check_workbook(trainer)
    wb = load_workbook(trainer, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

    # Equivalent learner formula on driver practice cell is allowed
    trainer2, answer2 = build_training_workbook(data, tmp_path / "WC_DriverEq.xlsx")
    smap2 = load_semantic_map(answer2)
    comp2 = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "nowc_change_from_components" and c.period_index == 1
    )
    _inject_formula_and_cached_value(
        trainer2,
        comp2.tab,
        comp2.cell,
        formula="=C19-C20",
        cached_value=float(comp2.expected_value),
    )
    assert check_workbook(trainer2).correct == 1

    # Incremental OWCA #N/A when ΔRevenue = 0
    fin = _base_wc_fin(
        income_statement=[
            _li("Revenue", 1000, 1000),
            _li("Finance costs", -10, -11),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    trainer3, answer3 = build_training_workbook(fin, tmp_path / "WC_SideNA.xlsx")
    smap3 = load_semantic_map(answer3)
    side = next(
        c
        for c in smap3.all_ordered()
        if c.family_id == "incremental_owca_to_revenue_change" and c.period_index == 1
    )
    assert side.expected_value == UNDEFINED_RATIO
    assert "NA()" in side.formula
    assert "ABS(" not in side.formula.upper()
    wb = load_workbook(trainer3, data_only=False)
    row, col = parse_cell_ref(side.cell)
    wb[side.tab].cell(row=row, column=col).value = side.formula
    wb.save(trainer3)
    wb.close()
    assert check_workbook(trainer3).correct == 1
    _inject_formula_and_cached_value(
        trainer3, side.tab, side.cell, formula="=NA()", cached_value=UNDEFINED_RATIO
    )
    assert check_workbook(trainer3).correct == 1
    _inject_formula_and_cached_value(
        trainer3, side.tab, side.cell, formula="=0", cached_value=0.0
    )
    assert check_workbook(trainer3).incorrect == 1

"""Step 9M.7 — deferred-tax balance diagnostics."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    DEFERRED_TAX_COMPONENT_CATALOG,
    expand_deferred_tax_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.deferred_tax import (
    compute_deferred_tax_series,
    deferred_tax_applicable,
    deferred_tax_availability,
    resolve_deferred_tax_sources,
)
from core.model.historical_expected import deferred_tax_expected_series
from core.model.line_resolver import AmbiguousLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"


def _ingest_demo():
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _tiny(
    *,
    dta=(100.0, 120.0),
    dtl=(40.0, 50.0),
    with_dta: bool = True,
    with_dtl: bool = True,
    label_only: bool = False,
    duplicate_dta: bool = False,
    duplicate_dtl: bool = False,
    missing_period: bool = False,
    none_period: bool = False,
    dta_first: bool = False,
    single_period: bool = False,
    dta_concept: str = "deferred_tax_assets",
    dtl_concept: str = "deferred_tax_liabilities",
    dta_label: str = "Deferred tax assets",
    dtl_label: str = "Deferred tax liabilities",
):
    if single_period:
        d1 = date(2025, 12, 31)
        periods = [FinancialPeriod(end_date=d1, label="FY2025")]

        def vals(a, b=None):
            return {d1: a}

        dta_vals = (dta[0],)
        dtl_vals = (dtl[0],)
    else:
        d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
        periods = [
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ]

        def vals(a, b=None):
            return {d1: a, d2: b if b is not None else a}

        dta_vals = dta
        dtl_vals = dtl

    cash, ar, ap, bank = 50.0, 40.0, 30.0, 20.0
    dta0 = dta_vals[0] if with_dta else 0.0
    dta1 = dta_vals[-1] if with_dta else 0.0
    dtl0 = dtl_vals[0] if with_dtl else 0.0
    dtl1 = dtl_vals[-1] if with_dtl else 0.0
    eq0 = cash + ar + dta0 - ap - bank - dtl0
    eq1 = cash + ar + dta1 - ap - bank - dtl1

    dta_item = None
    if with_dta:
        concept = "" if label_only else dta_concept
        if missing_period and not single_period:
            dta_values = {date(2024, 12, 31): dta_vals[0]}
        elif none_period and not single_period:
            dta_values = {date(2024, 12, 31): dta_vals[0], date(2025, 12, 31): None}
        elif single_period:
            dta_values = vals(dta_vals[0])
        else:
            dta_values = vals(*dta_vals)
        dta_item = _li(dta_label, dta_values, concept=concept)

    dtl_item = None
    if with_dtl:
        concept = "" if label_only else dtl_concept
        if missing_period and not single_period:
            dtl_values = {date(2024, 12, 31): dtl_vals[0]}
        elif none_period and not single_period:
            dtl_values = {date(2024, 12, 31): dtl_vals[0], date(2025, 12, 31): None}
        elif single_period:
            dtl_values = vals(dtl_vals[0])
        else:
            dtl_values = vals(*dtl_vals)
        dtl_item = _li(dtl_label, dtl_values, concept=concept)

    bs = [
        _li("Cash and cash equivalents", vals(cash, cash)),
        _li("Trade receivables", vals(ar, ar)),
        _li("Trade payables", vals(ap, ap)),
        _li("Bank borrowings", vals(bank, bank)),
    ]
    if dta_item is not None:
        if dta_first:
            bs.insert(0, dta_item)
        else:
            bs.insert(2, dta_item)
        if duplicate_dta:
            bs.insert(
                3,
                _li(
                    "Deferred tax assets duplicate",
                    vals(*dta_vals) if not single_period else vals(dta_vals[0]),
                    concept=dta_concept,
                ),
            )
    if dtl_item is not None:
        bs.append(dtl_item)
        if duplicate_dtl:
            bs.append(
                _li(
                    "Deferred tax liabilities duplicate",
                    vals(*dtl_vals) if not single_period else vals(dtl_vals[0]),
                    concept=dtl_concept,
                )
            )
    bs.append(_li("Total equity", vals(eq0, eq1)))

    if single_period:
        is_rows = [
            _li("Revenue", vals(1000)),
            _li("Finance costs", vals(-10)),
            _li("Finance income", vals(0)),
            _li("Profit before tax", vals(200)),
            _li("Income tax expense", vals(-30)),
            _li("Profit for the year", vals(170)),
        ]
        cf = [_li("Net cash from operating activities", vals(80))]
    else:
        is_rows = [
            _li("Revenue", vals(1000, 1100)),
            _li("Finance costs", vals(-10, -11)),
            _li("Finance income", vals(0, 0)),
            _li("Profit before tax", vals(200, 220)),
            _li("Income tax expense", vals(-30, -33)),
            _li("Profit for the year", vals(170, 187)),
        ]
        cf = [_li("Net cash from operating activities", vals(80, 90))]

    return StandardizedFinancials(
        ticker="DTAX",
        company_name="Deferred Tax Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=is_rows,
        balance_sheet=bs,
        cash_flow=cf,
    )


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(DEFERRED_TAX_COMPONENT_CATALOG) == 4
    assert [f.order for f in DEFERRED_TAX_COMPONENT_CATALOG] == list(range(116, 120))
    specs = expand_deferred_tax_specs(periods, start_order=1000)
    assert len(specs) == 17
    assert {s.family_order for s in specs} == set(range(116, 120))
    with pytest.raises(ValueError, match="duplicate"):
        expand_deferred_tax_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_deferred_tax_specs(list(reversed(periods)), start_order=1)


def test_ordinary_two_period_math():
    fin = _tiny()
    periods = list(canonical_fiscal_periods(fin))
    assert deferred_tax_applicable(fin)
    series = compute_deferred_tax_series(fin, periods)
    assert series.deferred_tax_assets == (100.0, 120.0)
    assert series.deferred_tax_liabilities == (40.0, 50.0)
    assert series.net_deferred_tax_position == (60.0, 70.0)
    assert series.deferred_tax_assets_change == (None, 20.0)
    assert series.deferred_tax_liabilities_change == (None, 10.0)
    assert series.net_deferred_tax_position_change == (None, 10.0)
    mapped = deferred_tax_expected_series(series)
    assert set(mapped) == {f.id for f in DEFERRED_TAX_COMPONENT_CATALOG}


def test_signed_net_liability_and_explicit_zero():
    fin = _tiny(dta=(10.0, 20.0), dtl=(40.0, 55.0))
    periods = list(canonical_fiscal_periods(fin))
    series = compute_deferred_tax_series(fin, periods)
    assert series.net_deferred_tax_position == (-30.0, -35.0)
    assert series.net_deferred_tax_position_change == (None, -5.0)

    zero = _tiny(dta=(0.0, 0.0), dtl=(0.0, 5.0))
    series_z = compute_deferred_tax_series(zero, list(canonical_fiscal_periods(zero)))
    assert series_z.deferred_tax_assets == (0.0, 0.0)
    assert series_z.net_deferred_tax_position == (0.0, -5.0)
    assert series_z.deferred_tax_assets_change == (None, 0.0)


def test_absent_partial_label_only_duplicate_incomplete():
    absent = _tiny(with_dta=False, with_dtl=False)
    assert not deferred_tax_applicable(absent)
    assert resolve_deferred_tax_sources(absent) is None
    assert ReferenceModelBuilder(absent).deferred_tax_specs == ()

    partial_dta = _tiny(with_dtl=False)
    assert not deferred_tax_applicable(partial_dta)
    assert deferred_tax_availability(partial_dta).deferred_tax_assets is True
    assert deferred_tax_availability(partial_dta).deferred_tax_liabilities is False
    assert ReferenceModelBuilder(partial_dta).deferred_tax_specs == ()

    partial_dtl = _tiny(with_dta=False)
    assert not deferred_tax_applicable(partial_dtl)
    assert ReferenceModelBuilder(partial_dtl).deferred_tax_specs == ()

    label_only = _tiny(label_only=True)
    assert not deferred_tax_applicable(label_only)
    assert resolve_deferred_tax_sources(label_only) is None
    assert ReferenceModelBuilder(label_only).deferred_tax_specs == ()

    dup = _tiny(duplicate_dta=True)
    avail = deferred_tax_availability(dup)
    assert avail.ambiguous is True
    assert not deferred_tax_applicable(dup)
    assert resolve_deferred_tax_sources(dup) is None

    missing = _tiny(missing_period=True)
    periods = list(canonical_fiscal_periods(missing))
    with pytest.raises(MissingHistoricalValueError):
        compute_deferred_tax_series(missing, periods)


def test_one_period_and_source_reordering(tmp_path):
    one = _tiny(single_period=True)
    periods = list(canonical_fiscal_periods(one))
    series = compute_deferred_tax_series(one, periods)
    assert series.net_deferred_tax_position == (60.0,)
    assert series.deferred_tax_assets_change == (None,)
    builder = ReferenceModelBuilder(one)
    assert len(builder.deferred_tax_specs) == 1
    assert builder.deferred_tax_specs[0].family_id == "net_deferred_tax_position"

    for dta_first in (False, True):
        fin = _tiny(dta_first=dta_first)
        trainer, answer = build_training_workbook(
            fin, tmp_path / f"DTAX_REORDER_{dta_first}.xlsx"
        )
        idx = next(
            i
            for i, item in enumerate(fin.balance_sheet)
            if (item.concept or "") == "deferred_tax_assets"
        )
        expected_row = 7 + idx
        wb = load_workbook(answer, data_only=False)
        ws = wb["ALT DuPont"]
        level_row = next(
            r
            for r in range(1, (ws.max_row or 1) + 1)
            if ws.cell(r, 1).value == "Deferred Tax Assets"
        )
        level_f = str(ws.cell(level_row, 2).value).replace(" ", "")
        assert f"'BalanceSheet'!B{expected_row}" in level_f or (
            f"'Balance Sheet'!B{expected_row}" in level_f
        )
        wb.close()


def _dupont_row_by_label(ws, label: str) -> int:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"ALT DuPont label not found: {label!r}")


def test_workbook_first_period_blank_and_check(tmp_path):
    data = _tiny()
    trainer, answer = build_training_workbook(data, tmp_path / "DTAX_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.deferred_tax_series is not None
    assert len(builder.deferred_tax_specs) == 5
    assert {s.family_id for s in builder.deferred_tax_specs} == {
        f.id for f in DEFERRED_TAX_COMPONENT_CATALOG
    }

    smap = load_semantic_map(answer)
    practice_cells = {(c.tab, c.cell) for c in smap.all_ordered()}
    dt_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in DEFERRED_TAX_COMPONENT_CATALOG}
    ]
    assert len(dt_comps) == 5

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "DEFERRED-TAX BALANCE CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        dta_row = _dupont_row_by_label(ws, "Deferred Tax Assets")
        dtl_row = _dupont_row_by_label(ws, "Deferred Tax Liabilities")
        net_row = _dupont_row_by_label(
            ws, "Net Deferred-Tax Asset Position (DTA − DTL)"
        )
        change_rows = (
            _dupont_row_by_label(ws, "Change in Deferred Tax Assets"),
            _dupont_row_by_label(ws, "Change in Deferred Tax Liabilities"),
            _dupont_row_by_label(ws, "Change in Net Deferred-Tax Asset Position"),
        )
        assert ws.cell(row=dta_row, column=2).value is not None
        assert ws.cell(row=dtl_row, column=2).value is not None
        # Net position is practice for all periods, including first.
        first_net = ws.cell(row=net_row, column=2)
        if path == answer:
            assert isinstance(first_net.value, str) and first_net.value.startswith("=")
            assert first_net.comment is not None
            assert (first_net.comment.text or "").strip()
        else:
            assert first_net.value is None
            assert first_net.comment is None
        for row in change_rows:
            cell = ws.cell(row=row, column=2)
            assert cell.value is None
            assert ("ALT DuPont", cell.coordinate) not in practice_cells
            later = ws.cell(row=row, column=3)
            if path == answer:
                assert isinstance(later.value, str) and later.value.startswith("=")
                assert later.comment is not None
                assert (later.comment.text or "").strip()
            else:
                assert later.value is None
                assert later.comment is None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(
        c for c in dt_comps if c.family_id == "net_deferred_tax_position_change"
    )
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Change in Net Deferred-Tax Asset Position" not in dumped


def test_demo_omits_and_practice_counts_unchanged(tmp_path):
    data = _ingest_demo()
    assert not deferred_tax_applicable(data)
    builder = ReferenceModelBuilder(data)
    assert builder.deferred_tax_series is None
    assert builder.deferred_tax_specs == ()

    trainer, answer = build_training_workbook(data, tmp_path / "DTAX_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312
    assert check_workbook(trainer).blank == 312


def _alias_dta_dtl(**extra):
    kwargs = {
        "dta_concept": "deferred_tax_asset",
        "dtl_concept": "deferred_tax_liability",
        "dta_label": "Deferred income tax assets",
        "dtl_label": "Deferred income tax liabilities",
    }
    kwargs.update(extra)
    return kwargs


def test_deferred_tax_alias_resolution_mixed_zero_missing_and_none():
    alias = _tiny(**_alias_dta_dtl())
    sources = resolve_deferred_tax_sources(alias)
    assert sources is not None
    assert sources.deferred_tax_assets.concept == "deferred_tax_asset"
    assert sources.deferred_tax_liabilities.concept == "deferred_tax_liability"
    assert deferred_tax_applicable(alias)
    avail = deferred_tax_availability(alias)
    assert avail.deferred_tax_assets is True
    assert avail.deferred_tax_liabilities is True
    assert avail.ambiguous is False

    series = compute_deferred_tax_series(alias, list(canonical_fiscal_periods(alias)))
    assert series.deferred_tax_assets == (100.0, 120.0)
    assert series.deferred_tax_liabilities == (40.0, 50.0)
    assert series.net_deferred_tax_position == (60.0, 70.0)

    mixed = _tiny(
        dta_concept="deferred_tax_asset",
        dtl_concept="deferred_tax_liabilities",
        dta_label="Deferred income tax assets",
    )
    mixed_series = compute_deferred_tax_series(
        mixed, list(canonical_fiscal_periods(mixed))
    )
    assert mixed_series.net_deferred_tax_position == series.net_deferred_tax_position

    mixed_rev = _tiny(
        dta_concept="deferred_tax_assets",
        dtl_concept="deferred_tax_liability",
        dtl_label="Deferred income tax liabilities",
    )
    assert deferred_tax_applicable(mixed_rev)

    canonical = _tiny()
    canon_series = compute_deferred_tax_series(
        canonical, list(canonical_fiscal_periods(canonical))
    )
    assert series.net_deferred_tax_position == canon_series.net_deferred_tax_position

    zero = _tiny(dta=(0.0, 0.0), dtl=(0.0, 5.0), **_alias_dta_dtl())
    series_z = compute_deferred_tax_series(zero, list(canonical_fiscal_periods(zero)))
    assert series_z.deferred_tax_assets == (0.0, 0.0)
    assert series_z.net_deferred_tax_position == (0.0, -5.0)

    missing = _tiny(missing_period=True, **_alias_dta_dtl())
    with pytest.raises(MissingHistoricalValueError):
        compute_deferred_tax_series(missing, list(canonical_fiscal_periods(missing)))

    none_period = _tiny(none_period=True, **_alias_dta_dtl())
    with pytest.raises(MissingHistoricalValueError):
        compute_deferred_tax_series(
            none_period, list(canonical_fiscal_periods(none_period))
        )


def test_deferred_tax_alias_ambiguity_and_label_only():
    both = _tiny()
    extra_vals = {date(2024, 12, 31): 100.0, date(2025, 12, 31): 120.0}
    both.balance_sheet.insert(
        3,
        _li(
            "Deferred income tax assets",
            extra_vals,
            concept="deferred_tax_asset",
        ),
    )
    equity = next(item for item in both.balance_sheet if item.label == "Total equity")
    for period, amount in extra_vals.items():
        equity.values[period] = float(equity.values[period]) + amount
    avail = deferred_tax_availability(both)
    assert avail.ambiguous is True
    assert not deferred_tax_applicable(both)
    assert resolve_deferred_tax_sources(both) is None
    with pytest.raises(AmbiguousLineError):
        resolve_line(both.balance_sheet, "deferred_tax_assets", required=False)
    assert ReferenceModelBuilder(both).deferred_tax_specs == ()

    dup_alias = _tiny(duplicate_dta=True, **_alias_dta_dtl())
    assert deferred_tax_availability(dup_alias).ambiguous is True
    assert not deferred_tax_applicable(dup_alias)

    label_only = _tiny(
        label_only=True,
        dta_label="Deferred income tax assets",
        dtl_label="Deferred income tax liabilities",
    )
    assert not deferred_tax_applicable(label_only)
    assert resolve_deferred_tax_sources(label_only) is None

    absent = _tiny(with_dta=False, with_dtl=False)
    assert not deferred_tax_applicable(absent)


def test_deferred_tax_alias_round_trip_preserves_stored_identity():
    fin = _tiny(**_alias_dta_dtl())
    original = resolve_deferred_tax_sources(fin)
    assert original is not None
    restored = standardized_from_payload(standardized_to_payload(fin))
    sources = resolve_deferred_tax_sources(restored)
    assert sources is not None
    assert sources.deferred_tax_assets.concept == "deferred_tax_asset"
    assert sources.deferred_tax_liabilities.concept == "deferred_tax_liability"
    assert sources.deferred_tax_assets.values == original.deferred_tax_assets.values
    assert sources.deferred_tax_liabilities.values == (
        original.deferred_tax_liabilities.values
    )
    assert deferred_tax_applicable(restored)
    series = compute_deferred_tax_series(
        restored, list(canonical_fiscal_periods(restored))
    )
    assert series.net_deferred_tax_position == (60.0, 70.0)


def test_deferred_tax_alias_python_excel_source_identity(tmp_path):
    for dta_first in (False, True):
        fin = _tiny(dta_first=dta_first, **_alias_dta_dtl())
        dta_resolved = resolve_line(
            fin.balance_sheet, "deferred_tax_assets", required=True
        )
        dtl_resolved = resolve_line(
            fin.balance_sheet, "deferred_tax_liabilities", required=True
        )
        assert dta_resolved.item is not None and dtl_resolved.item is not None
        assert dta_resolved.item.concept == "deferred_tax_asset"
        assert dtl_resolved.item.concept == "deferred_tax_liability"
        dta_row = 7 + dta_resolved.index
        dtl_row = 7 + dtl_resolved.index
        python_net = compute_deferred_tax_series(
            fin, list(canonical_fiscal_periods(fin))
        ).net_deferred_tax_position
        assert python_net == (60.0, 70.0)

        trainer, answer = build_training_workbook(
            fin, tmp_path / f"DTAX_ALIAS_{dta_first}.xlsx"
        )
        builder = ReferenceModelBuilder(fin)
        assert builder.deferred_tax_series is not None
        assert builder.deferred_tax_series.net_deferred_tax_position == python_net
        assert len(builder.deferred_tax_specs) == 5

        wb = load_workbook(answer, data_only=False)
        ws = wb["ALT DuPont"]
        dta_level = _dupont_row_by_label(ws, "Deferred Tax Assets")
        dtl_level = _dupont_row_by_label(ws, "Deferred Tax Liabilities")
        dta_f = str(ws.cell(dta_level, 2).value).replace(" ", "")
        dtl_f = str(ws.cell(dtl_level, 2).value).replace(" ", "")
        assert f"'BalanceSheet'!B{dta_row}" in dta_f or (
            f"'Balance Sheet'!B{dta_row}" in dta_f
        )
        assert f"'BalanceSheet'!B{dtl_row}" in dtl_f or (
            f"'Balance Sheet'!B{dtl_row}" in dtl_f
        )
        wb.close()

        blank = check_workbook(trainer)
        assert blank.incorrect == 0
        assert blank.blank == blank.total
        assert blank.correct == 0

        smap = load_semantic_map(answer)
        wb = load_workbook(trainer, data_only=False)
        for comp in smap.all_ordered():
            row, col = parse_cell_ref(comp.cell)
            wb[comp.tab].cell(row=row, column=col).value = comp.formula
        wb.save(trainer)
        wb.close()
        filled = check_workbook(trainer)
        assert filled.correct == filled.total
        assert filled.incorrect == 0
        assert filled.blank == 0

"""Step 9M.2.4.1.1.1.26 — source-supported reported operating-margin bridge."""

from __future__ import annotations

import copy
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.current_build import prepare_company_input, resolve_company
from core.data.interface import (
    DocumentManifest,
    DocumentType,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    REPORTED_MARGIN_COMPONENT_CATALOG,
    expand_reported_margin_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.historical_expected import reported_margin_expected_series
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.reported_margin import (
    GROSS_PROFIT_CONCEPT,
    OPERATING_PROFIT_CONCEPT,
    REVENUE_CONCEPT,
    compute_reported_margin_series,
    gross_margin_applicable,
    net_operating_expense_burden_applicable,
    reported_margin_applicable,
    reported_margin_availability,
    reported_operating_margin_applicable,
    resolve_reported_margin_sources,
)
from core.model.source_values import MissingHistoricalValueError
from core.research.drivers import assemble_drivers_view, render_drivers_markdown
from core.tests.test_capex import P1, P2, _dupont_row_by_label, _tiny
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "build" / "input" / "fast_retailing" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _vals(a, b, *, missing_period=False, none_period=False):
    if missing_period:
        return {P1: a}
    if none_period:
        return {P1: a, P2: None}
    return {P1: a, P2: b}


def _add_margins(
    fin: StandardizedFinancials,
    *,
    revenue=(1000.0, 1100.0),
    gross=(560.0, 600.0),
    operating=(200.0, 180.0),
    revenue_concept=REVENUE_CONCEPT,
    gross_concept=GROSS_PROFIT_CONCEPT,
    operating_concept=OPERATING_PROFIT_CONCEPT,
    include_revenue=True,
    include_gross=True,
    include_operating=True,
    duplicate=None,
    missing_period=None,
    none_period=None,
    on_balance_sheet=None,
):
    rows = (
        ("revenue", include_revenue, "Revenue", revenue_concept, revenue),
        ("gross", include_gross, "Gross profit", gross_concept, gross),
        (
            "operating",
            include_operating,
            "Operating profit",
            operating_concept,
            operating,
        ),
    )
    stored = {}
    for key, include, label, concept, values in rows:
        if not include:
            continue
        item = _li(
            label,
            _vals(
                values[0],
                values[1],
                missing_period=missing_period == key,
                none_period=none_period == key,
            ),
            concept=concept,
        )
        if key == "revenue":
            existing = next(
                (
                    row
                    for row in fin.income_statement
                    if row.concept == "revenue" or row.label.lower() == "revenue"
                ),
                None,
            )
            if existing is not None:
                fin.income_statement.remove(existing)
        target = (
            fin.balance_sheet if on_balance_sheet == key else fin.income_statement
        )
        target.append(item)
        stored[key] = item
        if duplicate == key and on_balance_sheet != key:
            target.append(
                _li(
                    f"{label} duplicate",
                    _vals(values[0], values[1]),
                    concept=concept,
                )
            )
    return stored


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(REPORTED_MARGIN_COMPONENT_CATALOG) == 6
    assert [f.order for f in REPORTED_MARGIN_COMPONENT_CATALOG] == list(
        range(139, 145)
    )
    levels = expand_reported_margin_specs(
        periods,
        start_order=1000,
        include_gross_margin=True,
        include_operating_margin=True,
        include_burden=True,
    )
    assert len(levels) == 15
    assert {s.family_id for s in levels} == {
        "gross_margin",
        "reported_operating_margin",
        "net_operating_expense_burden",
    }
    assert all(s.semantic_key.startswith("reported_margin.") for s in levels)
    full = expand_reported_margin_specs(
        periods,
        start_order=1000,
        include_gross_margin=True,
        include_operating_margin=True,
        include_burden=True,
        include_gross_margin_change=True,
        include_burden_change=True,
        include_reconstructed=True,
    )
    assert len(full) == 27
    assert {s.family_id for s in full} == {
        "gross_margin",
        "reported_operating_margin",
        "net_operating_expense_burden",
        "gross_margin_change",
        "net_operating_expense_burden_change",
        "reconstructed_operating_margin_change",
    }
    assert not any(
        s.family_id.endswith("_change") and s.period_index == 0 for s in full
    )
    with pytest.raises(ValueError, match="gross-margin change requires"):
        expand_reported_margin_specs(
            periods, start_order=1, include_gross_margin_change=True
        )
    with pytest.raises(ValueError, match="reconstructed change requires"):
        expand_reported_margin_specs(
            periods,
            start_order=1,
            include_gross_margin=True,
            include_gross_margin_change=True,
            include_reconstructed=True,
        )
    with pytest.raises(ValueError, match="duplicate"):
        expand_reported_margin_specs(
            [periods[0], periods[0]], start_order=1, include_gross_margin=True
        )
    with pytest.raises(ValueError, match="chronological"):
        expand_reported_margin_specs(
            list(reversed(periods)), start_order=1, include_gross_margin=True
        )


def test_unique_concept_resolution_and_renamed_label():
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_margins(fin)
    sources = resolve_reported_margin_sources(fin)
    assert sources.gross_profit is stored["gross"]
    assert sources.operating_profit is stored["operating"]
    assert sources.revenue is stored["revenue"]
    assert reported_margin_applicable(fin)
    assert gross_margin_applicable(fin)
    assert reported_operating_margin_applicable(fin)
    assert net_operating_expense_burden_applicable(fin)
    avail = reported_margin_availability(fin)
    assert avail.gross_profit is True
    assert avail.operating_profit_ambiguous is False

    renamed = _tiny(with_cfo=False, with_payments=False)
    _add_margins(renamed, operating_concept="operating_income")
    assert (
        resolve_reported_margin_sources(renamed).operating_profit.concept
        == "operating_income"
    )


def test_absent_label_only_duplicate_wrong_statement():
    absent = _tiny(with_cfo=False, with_payments=False)
    assert not reported_margin_applicable(absent)
    assert resolve_reported_margin_sources(absent).gross_profit is None
    with pytest.raises(MissingLineError):
        compute_reported_margin_series(
            absent, list(canonical_fiscal_periods(absent))
        )
    assert ReferenceModelBuilder(absent).reported_margin_specs == ()

    label_only = _tiny(with_cfo=False, with_payments=False)
    _add_margins(label_only, gross_concept="", operating_concept="")
    assert not reported_margin_applicable(label_only)
    assert resolve_reported_margin_sources(label_only).gross_profit is None
    assert (
        resolve_line(
            label_only.income_statement, GROSS_PROFIT_CONCEPT, required=False
        ).item
        is None
    )
    assert ReferenceModelBuilder(label_only).reported_margin_specs == ()

    dup = _tiny(with_cfo=False, with_payments=False)
    _add_margins(dup, duplicate="operating")
    avail = reported_margin_availability(dup)
    assert avail.operating_profit_ambiguous is True
    assert avail.operating_profit is False
    assert not reported_operating_margin_applicable(dup)
    assert gross_margin_applicable(dup)
    with pytest.raises(AmbiguousLineError):
        resolve_line(dup.income_statement, OPERATING_PROFIT_CONCEPT, required=False)
    builder = ReferenceModelBuilder(dup)
    assert {s.family_id for s in builder.reported_margin_specs} == {
        "gross_margin",
        "gross_margin_change",
    }

    wrong = _tiny(with_cfo=False, with_payments=False)
    _add_margins(wrong, on_balance_sheet="gross")
    assert not gross_margin_applicable(wrong)
    assert reported_operating_margin_applicable(wrong)
    assert (
        resolve_line(
            wrong.balance_sheet, GROSS_PROFIT_CONCEPT, required=False
        ).item
        is not None
    )


def test_independent_family_omissions():
    no_gross = _tiny(with_cfo=False, with_payments=False)
    _add_margins(no_gross, include_gross=False)
    assert reported_operating_margin_applicable(no_gross)
    assert not gross_margin_applicable(no_gross)
    assert not net_operating_expense_burden_applicable(no_gross)
    series = compute_reported_margin_series(
        no_gross, list(canonical_fiscal_periods(no_gross))
    )
    assert series.reported_operating_margin is not None
    assert series.gross_margin is None
    assert series.net_operating_expense_burden is None
    assert {s.family_id for s in ReferenceModelBuilder(no_gross).reported_margin_specs} == {
        "reported_operating_margin",
    }

    no_op = _tiny(with_cfo=False, with_payments=False)
    _add_margins(no_op, include_operating=False)
    assert gross_margin_applicable(no_op)
    assert not reported_operating_margin_applicable(no_op)
    builder = ReferenceModelBuilder(no_op)
    assert {s.family_id for s in builder.reported_margin_specs} == {
        "gross_margin",
        "gross_margin_change",
    }

    no_rev = _tiny(with_cfo=False, with_payments=False)
    _add_margins(no_rev, include_revenue=False)
    no_rev.income_statement = [
        row
        for row in no_rev.income_statement
        if row.concept != "revenue" and row.label.lower() != "revenue"
    ]
    assert not reported_margin_applicable(no_rev)
    with pytest.raises(MissingLineError):
        compute_reported_margin_series(
            no_rev, list(canonical_fiscal_periods(no_rev))
        )


def test_zeros_undefined_negative_and_missing_values():
    fin = _tiny(with_cfo=False, with_payments=False)
    _add_margins(fin, gross=(0.0, 50.0), operating=(0.0, -10.0), revenue=(1000.0, 2000.0))
    series = compute_reported_margin_series(fin, [P1, P2])
    assert series.gross_margin == (0.0, pytest.approx(50.0 / 2000.0))
    assert series.reported_operating_margin[0] == pytest.approx(0.0)
    assert series.reported_operating_margin[1] == pytest.approx(-10.0 / 2000.0)
    assert series.net_operating_expense_burden[0] == pytest.approx(0.0)
    assert series.net_operating_expense_burden[1] == pytest.approx(60.0 / 2000.0)

    zero_rev = _tiny(with_cfo=False, with_payments=False)
    _add_margins(zero_rev, revenue=(0.0, 2000.0))
    series_z = compute_reported_margin_series(zero_rev, [P1, P2])
    assert series_z.gross_margin[0] == UNDEFINED_RATIO
    assert series_z.reported_operating_margin[0] == UNDEFINED_RATIO
    assert series_z.net_operating_expense_burden[0] == UNDEFINED_RATIO
    assert series_z.gross_margin_change[1] == UNDEFINED_RATIO
    assert series_z.net_operating_expense_burden_change[1] == UNDEFINED_RATIO
    assert series_z.reconstructed_operating_margin_change[1] == UNDEFINED_RATIO
    assert series_z.gross_margin[1] == pytest.approx(600.0 / 2000.0)

    mapped = reported_margin_expected_series(series)
    assert set(mapped) == {f.id for f in REPORTED_MARGIN_COMPONENT_CATALOG}

    missing = _tiny(with_cfo=False, with_payments=False)
    _add_margins(missing, missing_period="gross")
    with pytest.raises(MissingHistoricalValueError):
        compute_reported_margin_series(missing, [P1, P2])
    none_period = _tiny(with_cfo=False, with_payments=False)
    _add_margins(none_period, none_period="operating")
    with pytest.raises(MissingHistoricalValueError):
        compute_reported_margin_series(none_period, [P1, P2])


def test_reconstructed_change_reconciles_and_immutability():
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_margins(fin)
    series = compute_reported_margin_series(fin, [P1, P2])
    gm0 = 560.0 / 1000.0
    gm1 = 600.0 / 1100.0
    om0 = 200.0 / 1000.0
    om1 = 180.0 / 1100.0
    b0 = (560.0 - 200.0) / 1000.0
    b1 = (600.0 - 180.0) / 1100.0
    assert series.gross_margin[0] == pytest.approx(gm0)
    assert series.reported_operating_margin[1] == pytest.approx(om1)
    assert series.net_operating_expense_burden[0] == pytest.approx(b0)
    assert series.gross_margin_change[1] == pytest.approx(gm1 - gm0)
    assert series.net_operating_expense_burden_change[1] == pytest.approx(b1 - b0)
    reconstructed = series.reconstructed_operating_margin_change[1]
    assert reconstructed == pytest.approx((gm1 - gm0) - (b1 - b0))
    assert reconstructed == pytest.approx(om1 - om0)
    assert series.gross_margin_change[0] is None
    assert series.reconstructed_operating_margin_change[0] is None

    reverse = [P2, P1]
    series_rev = compute_reported_margin_series(fin, reverse)
    assert series_rev.gross_margin[0] == pytest.approx(gm1)

    before = copy.deepcopy(stored["gross"].values)
    compute_reported_margin_series(fin, [P1, P2])
    assert stored["gross"].values == before
    assert stored["gross"].concept == GROSS_PROFIT_CONCEPT
    assert stored["operating"].concept == OPERATING_PROFIT_CONCEPT


def test_workbook_gating_structure_formulas_notes_and_check(tmp_path):
    data = _tiny(with_cfo=False, with_payments=False)
    _add_margins(data)
    trainer, answer = build_training_workbook(data, tmp_path / "RM_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.reported_margin_series is not None
    assert len(builder.reported_margin_specs) == 9
    assert {s.family_id for s in builder.reported_margin_specs} == {
        "gross_margin",
        "reported_operating_margin",
        "net_operating_expense_burden",
        "gross_margin_change",
        "net_operating_expense_burden_change",
        "reconstructed_operating_margin_change",
    }

    smap = load_semantic_map(answer)
    rm_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in REPORTED_MARGIN_COMPONENT_CATALOG}
    ]
    assert len(rm_comps) == 9
    assert not any(c.period_index == 0 and "change" in c.family_id for c in rm_comps)

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "REPORTED MARGIN BRIDGE CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        gp_row = _dupont_row_by_label(ws, "Gross profit (reported)")
        gm_row = _dupont_row_by_label(ws, "Gross margin")
        om_row = _dupont_row_by_label(ws, "Reported operating margin")
        burden_row = _dupont_row_by_label(ws, "Net operating expense burden")
        gm_change_row = _dupont_row_by_label(ws, "Change in gross margin")
        recon_row = _dupont_row_by_label(
            ws, "Reconstructed change in reported operating margin"
        )
        assert isinstance(ws.cell(gp_row, 2).value, str)
        assert ws.cell(gp_row, 2).value.startswith("=")
        assert "Income Statement" in str(ws.cell(gp_row, 2).value)
        assert ws.cell(gm_change_row, 2).value in (None, "")
        for col in (2, 3):
            gm_cell = ws.cell(gm_row, column=col)
            om_cell = ws.cell(om_row, column=col)
            burden_cell = ws.cell(burden_row, column=col)
            recon_cell = ws.cell(recon_row, column=col)
            if path == answer:
                assert isinstance(gm_cell.value, str) and gm_cell.value.startswith("=")
                assert "NA()" in str(gm_cell.value)
                assert "NA()" in str(burden_cell.value)
                note = (gm_cell.comment.text or "") if gm_cell.comment else ""
                assert "gross profit" in note.lower()
                om_note = (om_cell.comment.text or "") if om_cell.comment else ""
                assert "nopat" in om_note.lower()
                burden_note = (
                    (burden_cell.comment.text or "") if burden_cell.comment else ""
                )
                assert "sg&a" in burden_note.lower() or "sga" in burden_note.lower()
                if col == 3:
                    recon_note = (
                        (recon_cell.comment.text or "") if recon_cell.comment else ""
                    )
                    assert "percentage" in recon_note.lower()
                    assert "nopat" in recon_note.lower()
            else:
                assert gm_cell.value is None
                assert om_cell.value is None
                assert burden_cell.value is None
                assert gm_cell.comment is None
                assert om_cell.comment is None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in rm_comps if c.family_id == "gross_margin")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Gross margin" not in dumped


def test_source_row_fidelity_round_trip_and_immutability(tmp_path):
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_margins(fin)
    original_index = fin.income_statement.index(stored["gross"])
    trainer, answer = build_training_workbook(fin, tmp_path / "RM_FID.xlsx")
    smap = load_semantic_map(answer)
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    gp_row = _dupont_row_by_label(ws, "Gross profit (reported)")
    src_f = str(ws.cell(gp_row, 2).value).replace(" ", "")
    assert src_f.startswith("='IncomeStatement'!") or src_f.startswith(
        "='Income Statement'!"
    )
    assert str(original_index + 7) in src_f
    awb.close()
    restored = standardized_from_payload(standardized_to_payload(fin))
    rt = resolve_reported_margin_sources(restored)
    assert rt.gross_profit is not None
    assert rt.gross_profit.concept == GROSS_PROFIT_CONCEPT
    assert rt.operating_profit.concept == OPERATING_PROFIT_CONCEPT
    assert stored["gross"].values == {P1: 560.0, P2: 600.0}
    assert len(smap.all_ordered()) == len(ReferenceModelBuilder(fin).expected_specs)


def test_demo_omits_reported_margin_fast_retailing_activates(tmp_path):
    demo = HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )
    assert not reported_margin_applicable(demo)
    builder = ReferenceModelBuilder(demo)
    assert builder.reported_margin_series is None
    assert builder.reported_margin_specs == ()
    trainer, answer = build_training_workbook(demo, tmp_path / "RM_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312

    payload = json.loads(FR_JSON.read_text(encoding="utf-8"))
    fr = standardized_from_payload(payload)
    assert reported_margin_applicable(fr)
    fr_builder = ReferenceModelBuilder(fr)
    assert len(fr_builder.reported_margin_specs) == 27
    assert len(fr_builder.expected_specs) == 577
    assert "reconstructed_operating_margin_change" in {
        s.family_id for s in fr_builder.expected_specs
    }
    fr_series = compute_reported_margin_series(fr, canonical_fiscal_periods(fr))
    assert fr_series.sga is not None
    assert fr_series.impairment is None
    assert fr_series.operating_profit_residual is not None
    assert any(
        value is not None and abs(value) > 1.0
        for value in fr_series.operating_profit_residual
    )
    assert any(
        item.kind == "unestablished_inference" and not item.established
        for item in fr_series.assessments
    )


LULU_RECONCILED = (
    ROOT / "build" / "input" / "lululemon" / "reconciled" / "standardized.json"
)


def test_lululemon_component_margin_bridge_reconciles_to_filings():
    payload = json.loads(LULU_RECONCILED.read_text(encoding="utf-8"))
    fin = standardized_from_payload(payload)
    sources = resolve_reported_margin_sources(fin)
    assert sources.sga is not None
    assert sources.impairment is not None
    assert sources.amortization is not None
    periods = canonical_fiscal_periods(fin)
    series = compute_reported_margin_series(fin, periods)
    independent = {
        date(2022, 1, 30): (3608565.0, 2225034.0, 0.0, 50176.0, 1333355.0),
        date(2023, 1, 29): (4492340.0, 2757447.0, 407913.0, -1428.0, 1328408.0),
        date(2024, 1, 28): (5609405.0, 3397218.0, 74501.0, 5010.0, 2132676.0),
        date(2025, 2, 2): (6270811.0, 3762379.0, 0.0, 2735.0, 2505697.0),
        date(2026, 2, 1): (6284132.0, 4066556.0, 0.0, 6961.0, 2210615.0),
    }
    assert series.sga[-1] == 4066556.0
    assert series.impairment[1] == 407913.0
    assert series.impairment[-1] == 0.0
    assert series.other_operating_items[-1] == 6961.0
    for index, period in enumerate(periods):
        gp, sga, imp, other, op = independent[period]
        assert series.gross_profit[index] == gp
        assert series.sga[index] == sga
        assert series.impairment[index] == imp
        assert series.other_operating_items[index] == other
        assert series.reconstructed_operating_profit[index] == op
        assert series.operating_profit_residual[index] == 0.0
        assert series.operating_margin_residual[index] == 0.0
        gm = gp / series.revenue[index]
        sga_r = sga / series.revenue[index]
        imp_r = imp / series.revenue[index]
        other_r = other / series.revenue[index]
        assert series.reconstructed_component_operating_margin[index] == pytest.approx(
            gm - sga_r - imp_r - other_r
        )
    assert series.sga_change[0] is None
    assert series.gross_profit_interaction[0] is None
    for index in range(1, len(periods)):
        assert abs(series.gross_profit_change_residual[index]) < 1e-6
        assert series.operating_profit_change_residual[index] == 0.0
        gm = series.gross_profit[index] / series.revenue[index]
        prior_gm = series.gross_profit[index - 1] / series.revenue[index - 1]
        sga_r = series.sga[index] / series.revenue[index]
        prior_sga = series.sga[index - 1] / series.revenue[index - 1]
        imp_r = series.impairment[index] / series.revenue[index]
        prior_imp = series.impairment[index - 1] / series.revenue[index - 1]
        other_r = series.other_operating_items[index] / series.revenue[index]
        prior_other = (
            series.other_operating_items[index - 1] / series.revenue[index - 1]
        )
        reported = (
            series.operating_profit[index] / series.revenue[index]
            - series.operating_profit[index - 1] / series.revenue[index - 1]
        )
        expected_gm = gm - prior_gm
        expected_sga = -(sga_r - prior_sga)
        expected_imp = -(imp_r - prior_imp)
        expected_other = -(other_r - prior_other)
        expected_sum = expected_gm + expected_sga + expected_imp + expected_other
        assert series.gross_margin_contribution[index] == pytest.approx(expected_gm)
        assert series.sga_ratio_contribution[index] == pytest.approx(expected_sga)
        assert series.impairment_ratio_contribution[index] == pytest.approx(
            expected_imp
        )
        assert series.other_operating_ratio_contribution[index] == pytest.approx(
            expected_other
        )
        assert series.reconstructed_contribution_sum[index] == pytest.approx(
            expected_sum
        )
        assert series.reported_operating_margin_change[index] == pytest.approx(
            reported
        )
        assert series.contribution_residual[index] == pytest.approx(
            reported - expected_sum
        )
        assert abs(series.contribution_residual[index]) < 1e-12
    assert series.gross_margin_contribution[0] is None
    assert series.sga_ratio_contribution[0] is None
    assert any(
        item.name == "component operating-margin contributions" and item.established
        for item in series.assessments
    )
    assert sources.acquisition_related.values[date(2026, 2, 1)] is None
    assert sources.gain_on_disposal.values[date(2026, 2, 1)] is None
    kinds = {item.kind for item in series.assessments}
    assert "identity" in kinds
    assert "unestablished_inference" in kinds
    assert any(
        item.name.startswith("mix") and not item.established
        for item in series.assessments
    )
    restored = standardized_from_payload(standardized_to_payload(fin))
    assert resolve_reported_margin_sources(restored).impairment.concept == (
        "impairment_and_restructuring"
    )


def _add_component_lines(
    fin: StandardizedFinancials,
    *,
    sga=(300.0, 300.0),
    impairment=(0.0, 50.0),
    other=(10.0, 0.0),
    omit_impairment=False,
    impairment_none_current=False,
):
    fin.income_statement.append(
        _li(
            "SG&A",
            _vals(sga[0], sga[1]),
            concept="selling_general_and_administrative_expenses",
        )
    )
    if not omit_impairment:
        fin.income_statement.append(
            _li(
                "Impairment",
                {P1: impairment[0], P2: None}
                if impairment_none_current
                else _vals(impairment[0], impairment[1]),
                concept="impairment_and_restructuring",
            )
        )
    fin.income_statement.append(
        _li(
            "Amortization",
            _vals(other[0], other[1]),
            concept="amortization_of_intangible_assets",
        )
    )


def test_component_contributions_sign_missing_zero_and_residual():
    falling_burden = _tiny(with_cfo=False, with_payments=False)
    _add_margins(falling_burden, gross=(560.0, 600.0), operating=(250.0, 250.0))
    _add_component_lines(falling_burden)
    series = compute_reported_margin_series(falling_burden, [P1, P2])
    gm = 600.0 / 1100.0 - 560.0 / 1000.0
    sga = -(300.0 / 1100.0 - 300.0 / 1000.0)
    imp = -(50.0 / 1100.0 - 0.0 / 1000.0)
    other = -(0.0 / 1100.0 - 10.0 / 1000.0)
    reported = 250.0 / 1100.0 - 250.0 / 1000.0
    assert series.gross_margin_contribution[1] == pytest.approx(gm)
    assert series.sga_ratio_contribution[1] == pytest.approx(sga)
    assert sga > 0
    assert series.impairment_ratio_contribution[1] == pytest.approx(imp)
    assert imp < 0
    assert series.other_operating_ratio_contribution[1] == pytest.approx(other)
    assert other > 0
    assert series.reconstructed_contribution_sum[1] == pytest.approx(
        gm + sga + imp + other
    )
    assert series.contribution_residual[1] == pytest.approx(
        reported - (gm + sga + imp + other)
    )
    assert abs(series.contribution_residual[1]) < 1e-12
    assert series.gross_margin_contribution[0] is None
    assert round(series.sga_ratio_contribution[1] * 100, 2) == round(sga * 100, 2)
    assert round(series.sga_ratio_contribution[1] * 10000) == round(sga * 10000)

    zeros = _tiny(with_cfo=False, with_payments=False)
    _add_margins(zeros, gross=(560.0, 616.0), operating=(250.0, 276.0))
    _add_component_lines(zeros, sga=(300.0, 330.0), impairment=(0.0, 0.0), other=(10.0, 10.0))
    zero_series = compute_reported_margin_series(zeros, [P1, P2])
    assert zero_series.impairment_ratio_contribution[1] == pytest.approx(0.0)
    assert zero_series.impairment[0] == 0.0
    assert zero_series.impairment[1] == 0.0

    missing_line = _tiny(with_cfo=False, with_payments=False)
    _add_margins(missing_line, gross=(560.0, 600.0), operating=(250.0, 300.0))
    _add_component_lines(missing_line, omit_impairment=True)
    missing_series = compute_reported_margin_series(missing_line, [P1, P2])
    assert missing_series.impairment is None
    assert missing_series.impairment_ratio_contribution is None
    assert missing_series.reconstructed_contribution_sum[1] is not None

    sparse = _tiny(with_cfo=False, with_payments=False)
    _add_margins(sparse, gross=(560.0, 600.0), operating=(250.0, 250.0))
    _add_component_lines(sparse, impairment_none_current=True)
    sparse_series = compute_reported_margin_series(sparse, [P1, P2])
    assert sparse_series.impairment[1] is None
    assert sparse_series.impairment_ratio_contribution[1] is None
    assert sparse_series.reconstructed_contribution_sum[1] is None
    assert sparse_series.contribution_residual[1] is None

    fr = standardized_from_payload(json.loads(FR_JSON.read_text(encoding="utf-8")))
    fr_series = compute_reported_margin_series(fr, canonical_fiscal_periods(fr))
    assert fr_series.impairment_ratio_contribution is None
    assert any(
        value is not None and abs(value) > 1e-8
        for value in (fr_series.contribution_residual or ())
    )


def test_rendered_component_contribution_schedule(tmp_path):
    company = resolve_company("Lululemon")
    fin = prepare_company_input(company, tmp_path / "input")
    trainer, answer = build_training_workbook(fin, tmp_path / "LULU_CONTRIB.xlsx")
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    assert any(
        ws.cell(row, 1).value == "COMPONENT OPERATING-MARGIN CONTRIBUTIONS"
        for row in range(1, (ws.max_row or 1) + 1)
    )
    gm_row = _dupont_row_by_label(ws, "Δ gross margin contribution")
    sga_row = _dupont_row_by_label(ws, "−Δ SG&A/revenue contribution")
    imp_row = _dupont_row_by_label(ws, "−Δ impairment/revenue contribution")
    other_row = _dupont_row_by_label(
        ws, "−Δ other operating items/revenue contribution"
    )
    sum_row = _dupont_row_by_label(ws, "Reconstructed contribution sum")
    resid_row = _dupont_row_by_label(
        ws, "Contribution residual (reported − reconstructed)"
    )
    bps_row = _dupont_row_by_label(ws, "Δ gross margin contribution (bps)")
    assert ws.cell(gm_row, 2).value in (None, "")
    gm_f = str(ws.cell(gm_row, 3).value)
    sga_f = str(ws.cell(sga_row, 3).value)
    assert gm_f.startswith("=")
    assert sga_f.startswith("=")
    assert "-(" in sga_f or sga_f.startswith("=-")
    assert "*10000" in str(ws.cell(bps_row, 3).value)
    assert '""' in str(ws.cell(sum_row, 3).value)
    assert "reported" not in str(ws.cell(resid_row, 3).value).lower() or "-" in str(
        ws.cell(resid_row, 3).value
    )
    awb.close()

    text = render_drivers_markdown(assemble_drivers_view(fin, company.name))
    assert "Δgross margin" in text
    assert "−Δ(SG&A/revenue)" in text
    assert "−Δ(impairment/revenue)" in text
    assert "Reconstructed sum" in text
    assert "unrounded" in text
    assert "10,000" in text
    names = [
        line.split("|")[1].strip()
        for line in text.splitlines()
        if line.startswith("| ") and " | " in line
    ]
    assert names.count("component operating-margin identity") == 1
    assert names.count("latest adjacent operating-margin movement") == 1
    assert "component operating-margin contributions" in text

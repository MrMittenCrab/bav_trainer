"""Step 9M.2.4.1.1.1.25 — source-supported historical cash roll-forward practice."""

from __future__ import annotations

import copy
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    CASH_ROLLFORWARD_COMPONENT_CATALOG,
    expand_cash_rollforward_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.cash_rollforward import (
    BEGINNING_CONCEPT,
    CHANGE_CONCEPT,
    ENDING_CONCEPT,
    FINANCING_CONCEPT,
    FX_CONCEPT,
    INVESTING_CONCEPT,
    OPERATING_CONCEPT,
    cash_ending_difference_applicable,
    cash_ending_from_flows_applicable,
    cash_movement_difference_applicable,
    cash_movement_from_flows_applicable,
    cash_rollforward_applicable,
    cash_rollforward_availability,
    compute_cash_rollforward_series,
    resolve_cash_rollforward_sources,
)
from core.model.historical_expected import cash_rollforward_expected_series
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_capex import P1, P2, _dupont_row_by_label, _tiny
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "benchmark" / "fast_retailing" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _vals(a, b, *, missing_period=False, none_period=False):
    if missing_period:
        return {P1: a}
    if none_period:
        return {P1: a, P2: None}
    return {P1: a, P2: b}


def _add_cash(
    fin: StandardizedFinancials,
    *,
    operating=(80.0, 90.0),
    investing=(-20.0, -25.0),
    financing=(-30.0, -35.0),
    fx=(-5.0, 2.0),
    change=(25.0, 32.0),
    beginning=(100.0, 125.0),
    ending=(125.0, 157.0),
    operating_concept=OPERATING_CONCEPT,
    investing_concept=INVESTING_CONCEPT,
    financing_concept=FINANCING_CONCEPT,
    fx_concept=FX_CONCEPT,
    change_concept=CHANGE_CONCEPT,
    beginning_concept=BEGINNING_CONCEPT,
    ending_concept=ENDING_CONCEPT,
    include_operating=True,
    include_investing=True,
    include_financing=True,
    include_fx=True,
    include_change=True,
    include_beginning=True,
    include_ending=True,
    duplicate=None,
    missing_period=None,
    none_period=None,
    on_balance_sheet=None,
):
    rows = (
        (
            "operating",
            include_operating,
            "Net cash from operating activities",
            operating_concept,
            operating,
        ),
        (
            "investing",
            include_investing,
            "Net cash from investing activities",
            investing_concept,
            investing,
        ),
        (
            "financing",
            include_financing,
            "Net cash from financing activities",
            financing_concept,
            financing,
        ),
        ("fx", include_fx, "Effect of FX on cash", fx_concept, fx),
        ("change", include_change, "Increase (decrease) in cash", change_concept, change),
        (
            "beginning",
            include_beginning,
            "Cash beginning of period",
            beginning_concept,
            beginning,
        ),
        ("ending", include_ending, "Cash end of period", ending_concept, ending),
    )
    stored = {}
    for key, include, label, concept, values in rows:
        if not include:
            continue
        item = _li(
            label,
            _vals(
                values[0],
                values[1],
                missing_period=missing_period == key,
                none_period=none_period == key,
            ),
            concept=concept,
        )
        target = (
            fin.balance_sheet if on_balance_sheet == key else fin.cash_flow
        )
        target.append(item)
        stored[key] = item
        if duplicate == key and on_balance_sheet != key:
            fin.cash_flow.append(
                _li(
                    f"{label} duplicate",
                    _vals(values[0], values[1]),
                    concept=concept,
                )
            )
    return stored


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(CASH_ROLLFORWARD_COMPONENT_CATALOG) == 4
    assert [f.order for f in CASH_ROLLFORWARD_COMPONENT_CATALOG] == [135, 136, 137, 138]
    specs = expand_cash_rollforward_specs(periods, start_order=1000)
    assert len(specs) == 5
    assert {s.family_id for s in specs} == {"cash_movement_from_flows"}
    assert all(s.semantic_key.startswith("cash_rollforward.") for s in specs)
    full = expand_cash_rollforward_specs(
        periods,
        start_order=1000,
        include_movement_difference=True,
        include_ending_from_flows=True,
        include_ending_difference=True,
    )
    assert len(full) == 20
    assert {s.family_id for s in full} == {
        "cash_movement_from_flows",
        "cash_movement_difference",
        "cash_ending_from_flows",
        "cash_ending_difference",
    }
    with pytest.raises(ValueError, match="ending difference requires"):
        expand_cash_rollforward_specs(
            periods,
            start_order=1,
            include_ending_difference=True,
        )
    with pytest.raises(ValueError, match="duplicate"):
        expand_cash_rollforward_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_cash_rollforward_specs(list(reversed(periods)), start_order=1)


def test_unique_concept_resolution_and_renamed_label():
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_cash(fin)
    sources = resolve_cash_rollforward_sources(fin)
    assert sources.operating is stored["operating"]
    assert sources.fx is stored["fx"]
    assert cash_rollforward_applicable(fin)
    assert cash_movement_difference_applicable(fin)
    assert cash_ending_from_flows_applicable(fin)
    assert cash_ending_difference_applicable(fin)
    avail = cash_rollforward_availability(fin)
    assert avail.operating is True
    assert avail.fx_ambiguous is False

    renamed = _tiny(with_cfo=False, with_payments=False)
    _add_cash(renamed, operating_concept="operating_cash_flow")
    assert (
        resolve_cash_rollforward_sources(renamed).operating.concept
        == "operating_cash_flow"
    )


def test_absent_label_only_duplicate_wrong_statement():
    absent = _tiny(with_cfo=False, with_payments=False)
    assert not cash_rollforward_applicable(absent)
    assert resolve_cash_rollforward_sources(absent).operating is None
    with pytest.raises(MissingLineError):
        compute_cash_rollforward_series(
            absent, list(canonical_fiscal_periods(absent))
        )
    assert ReferenceModelBuilder(absent).cash_rollforward_specs == ()

    label_only = _tiny(with_cfo=False, with_payments=False)
    _add_cash(label_only, operating_concept="")
    assert not cash_rollforward_applicable(label_only)
    assert resolve_cash_rollforward_sources(label_only).operating is None
    assert (
        resolve_line(
            label_only.cash_flow, OPERATING_CONCEPT, required=False
        ).item
        is None
    )
    assert ReferenceModelBuilder(label_only).cash_rollforward_specs == ()

    dup = _tiny(with_cfo=False, with_payments=False)
    _add_cash(dup, duplicate="fx")
    avail = cash_rollforward_availability(dup)
    assert avail.fx_ambiguous is True
    assert avail.fx is False
    assert not cash_rollforward_applicable(dup)
    with pytest.raises(AmbiguousLineError):
        resolve_line(dup.cash_flow, FX_CONCEPT, required=False)
    assert ReferenceModelBuilder(dup).cash_rollforward_specs == ()

    wrong = _tiny(with_cfo=False, with_payments=False)
    _add_cash(wrong, on_balance_sheet="beginning")
    assert cash_movement_from_flows_applicable(wrong)
    assert not cash_ending_from_flows_applicable(wrong)
    assert (
        resolve_line(
            wrong.balance_sheet, BEGINNING_CONCEPT, required=False
        ).item
        is not None
    )


def test_no_substitution_from_bs_cash_or_nearby_cf():
    fin = _tiny(with_cfo=False, with_payments=False)
    fin.balance_sheet.append(
        _li("Cash and cash equivalents", {P1: 50.0, P2: 60.0}, concept="cash")
    )
    fin.cash_flow.append(
        _li(
            "Cash generated from operations",
            {P1: 70.0, P2: 80.0},
            concept="cash_generated_from_operations",
        )
    )
    fin.cash_flow.append(
        _li(
            "Foreign exchange losses/(gains)",
            {P1: -3.0, P2: -4.0},
            concept="foreign_exchange_losses_gains",
        )
    )
    assert resolve_cash_rollforward_sources(fin).operating is None
    assert resolve_cash_rollforward_sources(fin).fx is None
    assert resolve_cash_rollforward_sources(fin).beginning is None
    assert not cash_rollforward_applicable(fin)


def test_signed_flows_zeros_and_nonzero_differences():
    fin = _tiny(with_cfo=False, with_payments=False)
    _add_cash(fin)
    series = compute_cash_rollforward_series(
        fin, list(canonical_fiscal_periods(fin))
    )
    assert series.cash_movement_from_flows == (25.0, 32.0)
    assert series.cash_movement_difference == (0.0, 0.0)
    assert series.cash_ending_from_flows == (125.0, 157.0)
    assert series.cash_ending_difference == (0.0, 0.0)

    signed = _tiny(with_cfo=False, with_payments=False)
    _add_cash(
        signed,
        operating=(-80.0, 10.0),
        investing=(20.0, -5.0),
        financing=(30.0, 0.0),
        fx=(5.0, -2.0),
        change=(-24.0, 4.0),
        beginning=(200.0, 175.0),
        ending=(175.0, 178.0),
    )
    series_s = compute_cash_rollforward_series(
        signed, list(canonical_fiscal_periods(signed))
    )
    assert series_s.cash_movement_from_flows == (-25.0, 3.0)
    assert series_s.cash_movement_difference == (-1.0, -1.0)
    assert series_s.cash_ending_from_flows == (175.0, 178.0)
    assert series_s.cash_ending_difference == (0.0, 0.0)

    zero = _tiny(with_cfo=False, with_payments=False)
    _add_cash(zero, fx=(0.0, 0.0), change=(30.0, 30.0), ending=(130.0, 155.0))
    series_z = compute_cash_rollforward_series(
        zero, list(canonical_fiscal_periods(zero))
    )
    assert series_z.fx == (0.0, 0.0)
    assert series_z.cash_movement_from_flows == (30.0, 30.0)

    mapped = cash_rollforward_expected_series(series)
    assert set(mapped) == {f.id for f in CASH_ROLLFORWARD_COMPONENT_CATALOG}


def test_missing_none_versus_reported_zero():
    missing = _tiny(with_cfo=False, with_payments=False)
    _add_cash(missing, missing_period="fx")
    periods = list(canonical_fiscal_periods(missing))
    with pytest.raises(MissingHistoricalValueError, match="effect_of_fx_on_cash"):
        compute_cash_rollforward_series(missing, periods)

    none_period = _tiny(with_cfo=False, with_payments=False)
    _add_cash(none_period, none_period="change")
    with pytest.raises(MissingHistoricalValueError, match="change_in_cash"):
        compute_cash_rollforward_series(
            none_period, list(canonical_fiscal_periods(none_period))
        )

    zero = _tiny(with_cfo=False, with_payments=False)
    stored = _add_cash(zero, fx=(0.0, 2.0), change=(30.0, 32.0), ending=(130.0, 157.0))
    series = compute_cash_rollforward_series(
        zero, list(canonical_fiscal_periods(zero))
    )
    assert series.fx[0] == 0.0
    assert stored["fx"].values[P1] == 0.0


def test_absent_ambiguous_sources_omit_only_dependent_families():
    no_fx = _tiny(with_cfo=False, with_payments=False)
    _add_cash(no_fx, include_fx=False)
    assert not cash_rollforward_applicable(no_fx)
    assert ReferenceModelBuilder(no_fx).cash_rollforward_specs == ()

    no_change = _tiny(with_cfo=False, with_payments=False)
    _add_cash(no_change, include_change=False)
    assert cash_rollforward_applicable(no_change)
    assert not cash_movement_difference_applicable(no_change)
    assert cash_ending_from_flows_applicable(no_change)
    series = compute_cash_rollforward_series(
        no_change, list(canonical_fiscal_periods(no_change))
    )
    assert series.cash_movement_from_flows == (25.0, 32.0)
    assert series.cash_movement_difference is None
    assert series.cash_ending_from_flows == (125.0, 157.0)
    mapped = cash_rollforward_expected_series(series)
    assert set(mapped) == {
        "cash_movement_from_flows",
        "cash_ending_from_flows",
        "cash_ending_difference",
    }
    builder = ReferenceModelBuilder(no_change)
    assert {s.family_id for s in builder.cash_rollforward_specs} == {
        "cash_movement_from_flows",
        "cash_ending_from_flows",
        "cash_ending_difference",
    }
    assert len(builder.cash_rollforward_specs) == 6

    no_beginning = _tiny(with_cfo=False, with_payments=False)
    _add_cash(no_beginning, include_beginning=False)
    assert cash_rollforward_applicable(no_beginning)
    assert cash_movement_difference_applicable(no_beginning)
    assert not cash_ending_from_flows_applicable(no_beginning)
    assert not cash_ending_difference_applicable(no_beginning)
    series_nb = compute_cash_rollforward_series(
        no_beginning, list(canonical_fiscal_periods(no_beginning))
    )
    assert series_nb.cash_ending_from_flows is None
    assert series_nb.cash_ending_difference is None
    assert len(ReferenceModelBuilder(no_beginning).cash_rollforward_specs) == 4

    no_ending = _tiny(with_cfo=False, with_payments=False)
    _add_cash(no_ending, include_ending=False)
    series_ne = compute_cash_rollforward_series(
        no_ending, list(canonical_fiscal_periods(no_ending))
    )
    assert series_ne.cash_ending_from_flows == (125.0, 157.0)
    assert series_ne.cash_ending_difference is None
    assert len(ReferenceModelBuilder(no_ending).cash_rollforward_specs) == 6

    amb_fx = _tiny(with_cfo=False, with_payments=False)
    _add_cash(amb_fx, duplicate="operating")
    avail = cash_rollforward_availability(amb_fx)
    assert avail.operating_ambiguous is True
    assert not cash_rollforward_applicable(amb_fx)

    competing = _tiny(with_cfo=False, with_payments=False)
    _add_cash(competing)
    competing.cash_flow.append(
        _li(
            "FX alias row",
            {P1: 1.0, P2: 1.0},
            concept="effect_of_exchange_rate_on_cash",
        )
    )
    avail_c = cash_rollforward_availability(competing)
    assert avail_c.fx_ambiguous is True
    assert not cash_rollforward_applicable(competing)


def test_period_ordering_and_unchanged_input():
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_cash(fin)
    reverse = [P2, P1]
    series = compute_cash_rollforward_series(fin, reverse)
    assert series.cash_movement_from_flows == (32.0, 25.0)
    assert series.cash_ending_from_flows == (157.0, 125.0)

    before = copy.deepcopy(stored["operating"].values)
    compute_cash_rollforward_series(fin, [P1, P2])
    assert stored["operating"].values == before
    assert stored["operating"].concept == OPERATING_CONCEPT
    assert stored["fx"].concept == FX_CONCEPT


def test_workbook_gating_structure_formulas_notes_and_check(tmp_path):
    data = _tiny(with_cfo=False, with_payments=False)
    _add_cash(data)
    trainer, answer = build_training_workbook(data, tmp_path / "CR_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.cash_rollforward_series is not None
    assert len(builder.cash_rollforward_specs) == 8
    assert {s.family_id for s in builder.cash_rollforward_specs} == {
        "cash_movement_from_flows",
        "cash_movement_difference",
        "cash_ending_from_flows",
        "cash_ending_difference",
    }

    smap = load_semantic_map(answer)
    cr_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in CASH_ROLLFORWARD_COMPONENT_CATALOG}
    ]
    assert len(cr_comps) == 8

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "CASH ROLL-FORWARD CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        operating_row = _dupont_row_by_label(
            ws, "Net cash from operating activities (reported)"
        )
        movement_row = _dupont_row_by_label(ws, "Cash movement from flows")
        diff_row = _dupont_row_by_label(ws, "Cash movement difference")
        ending_row = _dupont_row_by_label(ws, "Cash ending from flows")
        ending_diff_row = _dupont_row_by_label(ws, "Cash ending difference")
        assert isinstance(ws.cell(operating_row, 2).value, str)
        assert ws.cell(operating_row, 2).value.startswith("=")
        for col in (2, 3):
            movement_cell = ws.cell(movement_row, column=col)
            diff_cell = ws.cell(diff_row, column=col)
            ending_cell = ws.cell(ending_row, column=col)
            ending_diff_cell = ws.cell(ending_diff_row, column=col)
            if path == answer:
                assert isinstance(movement_cell.value, str) and movement_cell.value.startswith(
                    "="
                )
                compact = movement_cell.value.replace(" ", "")
                assert compact.count("+") == 3
                assert "-" in str(diff_cell.value)
                note = (movement_cell.comment.text or "") if movement_cell.comment else ""
                assert "fx" in note.lower()
                assert "signed" in note.lower() or "sign" in note.lower()
                diff_note = (diff_cell.comment.text or "") if diff_cell.comment else ""
                assert "statement" in diff_note.lower()
                assert "plug" in diff_note.lower()
                assert "debt" in diff_note.lower()
                ending_note = (
                    (ending_cell.comment.text or "") if ending_cell.comment else ""
                )
                assert "beginning" in ending_note.lower()
                ending_diff_note = (
                    (ending_diff_cell.comment.text or "")
                    if ending_diff_cell.comment
                    else ""
                )
                assert "plug" in ending_diff_note.lower()
            else:
                assert movement_cell.value is None
                assert diff_cell.value is None
                assert ending_cell.value is None
                assert ending_diff_cell.value is None
                assert movement_cell.comment is None
                assert diff_cell.comment is None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in cr_comps if c.family_id == "cash_movement_from_flows")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Cash movement from flows" not in dumped


def test_source_row_fidelity_round_trip_and_immutability(tmp_path):
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_cash(fin)
    original_index = fin.cash_flow.index(stored["operating"])
    resolved = resolve_line(fin.cash_flow, OPERATING_CONCEPT, required=True)
    assert resolved.index == original_index
    assert resolved.item is stored["operating"]

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt = resolve_cash_rollforward_sources(restored)
    assert rt.operating is not None
    assert rt.operating.concept == OPERATING_CONCEPT
    assert rt.operating.values == stored["operating"].values
    assert rt.fx.concept == FX_CONCEPT

    trainer, answer = build_training_workbook(fin, tmp_path / "CR_ROW.xlsx")
    wb = load_workbook(answer, data_only=False)
    ws = wb["ALT DuPont"]
    operating_row = _dupont_row_by_label(
        ws, "Net cash from operating activities (reported)"
    )
    level_f = str(ws.cell(operating_row, 2).value).replace(" ", "")
    expected_row = 7 + original_index
    assert f"'CashFlowStatement'!B{expected_row}" in level_f or (
        f"'Cash Flow Statement'!B{expected_row}" in level_f
    )
    wb.close()
    assert stored["operating"].concept == OPERATING_CONCEPT
    assert stored["fx"].values == {P1: -5.0, P2: 2.0}
    assert trainer.exists() and answer.exists()


def test_demo_omits_cash_rollforward_fast_retailing_activates(tmp_path):
    demo = HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )
    assert not cash_rollforward_applicable(demo)
    builder = ReferenceModelBuilder(demo)
    assert builder.cash_rollforward_series is None
    assert builder.cash_rollforward_specs == ()
    trainer, answer = build_training_workbook(demo, tmp_path / "CR_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312

    payload = json.loads(FR_JSON.read_text(encoding="utf-8"))
    fr = standardized_from_payload(payload)
    assert cash_rollforward_applicable(fr)
    fr_builder = ReferenceModelBuilder(fr)
    assert len(fr_builder.cash_rollforward_specs) == 20
    assert len(fr_builder.expected_specs) == 577
    assert "cash_movement_from_flows" in {
        s.family_id for s in fr_builder.expected_specs
    }

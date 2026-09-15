"""Step 9M.2 — Lululemon real-company benchmark baseline acceptance."""

from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.standardized_io import (
    standardized_from_payload,
    standardized_to_payload,
)
from core.ingestion.filing_json import load_extracted_filing
from core.ingestion.filing_validator import validate_extracted_filing
from core.model.classification import (
    UnclassifiedBalanceSheetLineError,
    classify_balance_sheet_line,
    is_balance_sheet_subtotal,
)
from core.engine.reference_model import (
    EARNINGS_QUALITY_SHEET,
    SOURCE_START_ROW,
    ReferenceModelBuilder,
)
from core.model.capex import (
    capex_applicable,
    capex_availability,
    cash_after_ppe_capex_applicable,
    compute_capex_series,
    resolve_capex_source,
    resolve_operating_cash_source,
)
from core.model.deferred_tax import (
    compute_deferred_tax_series,
    deferred_tax_applicable,
    deferred_tax_availability,
    resolve_deferred_tax_sources,
)
from core.model.financial_math import compute_anchor
from core.model.fixed_asset import fixed_asset_applicable, fixed_asset_availability
from core.model.lease_rou import (
    compute_lease_rou_series,
    lease_rou_applicable,
    lease_rou_availability,
    resolve_lease_rou_source,
)
from core.model.line_resolver import resolve_line, workbook_row_for
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import SOURCE_UNAVAILABLE
from core.model.source_values import required_period_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook

ROOT = Path(__file__).resolve().parents[2]
BENCH = ROOT / "benchmark" / "lululemon"
EXTRACTED = BENCH / "extracted"
SOURCE = BENCH / "source"
RECONCILED = BENCH / "reconciled"
STD_JSON = RECONCILED / "standardized.json"
PROV_JSON = RECONCILED / "provenance.json"
CONFLICTS_JSON = RECONCILED / "conflicts.json"

SOURCE_PDFS = {
    2022: (
        "LULU_FY2022_Annual_Report.pdf",
        "b344d1e7a710259fa06f88773dee0b3827334820ce2b881fe6b95ca2ae275e4e",
        4913067,
    ),
    2023: (
        "LULU_FY2023_Annual_Report.pdf",
        "cd47ea251d608d06a3e58b5d782f2d41d5a231a994d2f7993267a430cb13c0f1",
        5848446,
    ),
    2024: (
        "LULU_FY2024_Annual_Report.pdf",
        "9268fd530db162babdd1ec4363cf388ebce57125d83b7e097aba6f98ba0ca7ec",
        5953217,
    ),
    2025: (
        "LULU_FY2025_Annual_Report.pdf",
        "82e00f900cc912a7d79596409594156b7779c3a193783ea8fecf87bc013c71cc",
        6590658,
    ),
}

EXPECTED_PERIODS = [
    date(2023, 1, 29),
    date(2024, 1, 28),
    date(2025, 2, 2),
    date(2026, 2, 1),
]

REVENUE_ANCHORS = {
    date(2023, 1, 29): 8110518.0,
    date(2024, 1, 28): 9619278.0,
    date(2025, 2, 2): 10588126.0,
    date(2026, 2, 1): 11102600.0,
}

DILUTED_WAS_ANCHORS = {
    date(2023, 1, 29): 128017.0,
    date(2024, 1, 28): 127060.0,
    date(2025, 2, 2): 123935.0,
    date(2026, 2, 1): 119068.0,
}

BUILD_BLOCKER_LABELS: set[str] = set()


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_four_filings_validate_and_remain_source_bound():
    files = sorted(EXTRACTED.glob("LULU_FY*.json"))
    assert {int(p.stem.replace("LULU_FY", "")) for p in files} == {2022, 2023, 2024, 2025}

    for fiscal_year, (pdf_name, digest, nbytes) in SOURCE_PDFS.items():
        pdf = SOURCE / pdf_name
        assert pdf.is_file()
        data = pdf.read_bytes()
        assert len(data) == nbytes
        assert hashlib.sha256(data).hexdigest() == digest

        extracted = EXTRACTED / f"LULU_FY{fiscal_year}.json"
        filing = load_extracted_filing(extracted)
        assert filing.schema_version == "1.0"
        assert filing.filing.fiscal_year == fiscal_year
        assert filing.filing.source_file == pdf_name
        assert filing.income_statement
        assert filing.balance_sheet
        assert filing.cash_flow

        report = validate_extracted_filing(filing, source_root=SOURCE)
        assert report.ok
        assert report.computed_source_sha256 == digest
        assert report.warnings == ()


def _reconcile_cmd(out: Path) -> list[str]:
    return [
        sys.executable,
        "-m",
        "core",
        "reconcile",
        str(EXTRACTED),
        "--source-root",
        str(SOURCE),
        "-o",
        str(out),
    ]


def _subprocess_env() -> dict[str, str]:
    env = dict(os.environ)
    env["PYTHONPATH"] = str(ROOT)
    return env


ARTIFACT_NAMES = ("standardized.json", "provenance.json", "conflicts.json")


def _read_committed_artifacts() -> dict[str, bytes]:
    return {name: (RECONCILED / name).read_bytes() for name in ARTIFACT_NAMES}


def _assert_artifact_sets_match(
    generated: Path,
    expected_bytes: dict[str, bytes],
) -> None:
    """Compare generated artifacts to expected bytes; never write expected paths."""
    mismatches: list[str] = []
    for name in ARTIFACT_NAMES:
        actual = (generated / name).read_bytes()
        if actual != expected_bytes[name]:
            mismatches.append(name)
    if mismatches:
        raise AssertionError(
            "reconciliation artifact mismatch: " + ", ".join(mismatches)
        )


def _run_reconcile_pass(out: Path) -> None:
    completed = subprocess.run(
        _reconcile_cmd(out),
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
        env=_subprocess_env(),
    )
    assert "overlap_conflicts=" in completed.stdout


def _assert_inter_run_artifacts_match(out1: Path, out2: Path) -> None:
    for name in ARTIFACT_NAMES:
        assert (out1 / name).read_bytes() == (out2 / name).read_bytes()


def _assert_committed_artifacts_unchanged(committed_before: dict[str, bytes]) -> None:
    committed_after = _read_committed_artifacts()
    mutated = [
        name
        for name in ARTIFACT_NAMES
        if committed_after[name] != committed_before[name]
    ]
    if mutated:
        raise AssertionError(
            "committed reconciliation artifacts mutated: " + ", ".join(mutated)
        )


def _guarded_deterministic_reconcile(
    tmp_path: Path,
    *,
    committed_before: dict[str, bytes] | None = None,
    run_pass=_run_reconcile_pass,
    assert_inter_run=_assert_inter_run_artifacts_match,
    assert_baseline=_assert_artifact_sets_match,
    read_committed=_read_committed_artifacts,
) -> None:
    """Run both reconcile passes under try; always re-verify committed bytes in finally."""
    if committed_before is None:
        committed_before = read_committed()
    out1 = tmp_path / "r1"
    out2 = tmp_path / "r2"
    error: BaseException | None = None
    try:
        run_pass(out1)
        run_pass(out2)
        assert_inter_run(out1, out2)
        assert_baseline(out1, committed_before)
        assert_baseline(out2, committed_before)
    except BaseException as exc:
        error = exc
    finally:
        try:
            after = read_committed()
            mutated = [
                name
                for name in ARTIFACT_NAMES
                if after[name] != committed_before[name]
            ]
            if mutated:
                raise AssertionError(
                    "committed reconciliation artifacts mutated: "
                    + ", ".join(mutated)
                )
        except AssertionError:
            raise
        else:
            if error is not None:
                raise error


def test_generic_reconcile_is_deterministic(tmp_path: Path):
    committed_before = _read_committed_artifacts()
    _guarded_deterministic_reconcile(tmp_path, committed_before=committed_before)


@pytest.mark.parametrize("fail_pass", [1, 2])
def test_reconcile_immutability_verified_after_subprocess_failure(
    tmp_path: Path, fail_pass: int
):
    """Subprocess failure still triggers final committed-byte verification."""
    committed_before = _read_committed_artifacts()
    calls = {"n": 0}
    verified = {"ok": False}

    def boom_pass(out: Path) -> None:
        calls["n"] += 1
        if calls["n"] == fail_pass:
            raise subprocess.CalledProcessError(1, _reconcile_cmd(out))
        _run_reconcile_pass(out)

    def tracking_read() -> dict[str, bytes]:
        verified["ok"] = True
        return _read_committed_artifacts()

    with pytest.raises(subprocess.CalledProcessError):
        _guarded_deterministic_reconcile(
            tmp_path,
            committed_before=committed_before,
            run_pass=boom_pass,
            read_committed=tracking_read,
        )
    assert verified["ok"]
    _assert_committed_artifacts_unchanged(committed_before)


@pytest.mark.parametrize("fail_at", ["inter_run", "baseline"])
def test_reconcile_immutability_verified_after_comparison_failure(
    tmp_path: Path, fail_at: str
):
    """Comparison failure still triggers final committed-byte verification."""
    committed_before = _read_committed_artifacts()
    verified = {"ok": False}

    def boom_inter(out1: Path, out2: Path) -> None:
        raise AssertionError("injected inter-run mismatch")

    def boom_baseline(generated: Path, expected_bytes: dict[str, bytes]) -> None:
        raise AssertionError("injected baseline mismatch")

    def tracking_read() -> dict[str, bytes]:
        verified["ok"] = True
        return _read_committed_artifacts()

    kwargs: dict = {
        "committed_before": committed_before,
        "read_committed": tracking_read,
    }
    if fail_at == "inter_run":
        kwargs["assert_inter_run"] = boom_inter
        match = "injected inter-run mismatch"
    else:
        kwargs["assert_baseline"] = boom_baseline
        match = "injected baseline mismatch"

    with pytest.raises(AssertionError, match=match):
        _guarded_deterministic_reconcile(tmp_path, **kwargs)
    assert verified["ok"]
    _assert_committed_artifacts_unchanged(committed_before)


@pytest.mark.parametrize("mutated_name", ARTIFACT_NAMES)
@pytest.mark.parametrize("fail_at", ["pass1", "inter_run", "baseline"])
def test_reconcile_immutability_guard_detects_temp_mutation(
    tmp_path: Path, mutated_name: str, fail_at: str
):
    """Temp artifact mutation before injected failure is detected by the finally guard."""
    expected_dir = tmp_path / "expected"
    expected_dir.mkdir()
    for name in ARTIFACT_NAMES:
        (expected_dir / name).write_bytes((RECONCILED / name).read_bytes())

    def read_temp() -> dict[str, bytes]:
        return {name: (expected_dir / name).read_bytes() for name in ARTIFACT_NAMES}

    committed_before = read_temp()
    mutated_path = expected_dir / mutated_name
    altered = mutated_path.read_bytes() + b"\n#mutated\n"
    finally_saw_mutation = {"ok": False}

    def mutate_then_fail_pass(out: Path) -> None:
        mutated_path.write_bytes(altered)
        raise subprocess.CalledProcessError(1, ["fake"])

    def mutate_then_fail_inter(out1: Path, out2: Path) -> None:
        mutated_path.write_bytes(altered)
        raise AssertionError("injected inter-run mismatch")

    def mutate_then_fail_baseline(
        generated: Path, expected_bytes: dict[str, bytes]
    ) -> None:
        mutated_path.write_bytes(altered)
        raise AssertionError("injected baseline mismatch")

    def read_temp_tracking() -> dict[str, bytes]:
        current = read_temp()
        if current[mutated_name] != committed_before[mutated_name]:
            finally_saw_mutation["ok"] = True
        return current

    kwargs: dict = {
        "committed_before": committed_before,
        "read_committed": read_temp_tracking,
    }
    if fail_at == "pass1":
        kwargs["run_pass"] = mutate_then_fail_pass
    elif fail_at == "inter_run":
        kwargs["assert_inter_run"] = mutate_then_fail_inter
    else:
        kwargs["assert_baseline"] = mutate_then_fail_baseline

    with pytest.raises(AssertionError, match=mutated_name):
        _guarded_deterministic_reconcile(tmp_path, **kwargs)

    assert finally_saw_mutation["ok"]
    assert mutated_path.read_bytes() == altered
    for name in ARTIFACT_NAMES:
        if name == mutated_name:
            continue
        assert (expected_dir / name).read_bytes() == (RECONCILED / name).read_bytes()


@pytest.mark.parametrize("drifted_name", ARTIFACT_NAMES)
def test_reconcile_drift_is_detected_without_rewriting_expected(
    tmp_path: Path, drifted_name: str
):
    """Altering a temp expected copy must fail comparison and leave that copy intact."""
    expected_dir = tmp_path / "expected"
    expected_dir.mkdir()
    for name in ARTIFACT_NAMES:
        (expected_dir / name).write_bytes((RECONCILED / name).read_bytes())

    drifted_path = expected_dir / drifted_name
    original_expected = drifted_path.read_bytes()
    drifted_path.write_bytes(original_expected + b"\n#drift\n")
    altered_expected = drifted_path.read_bytes()
    assert altered_expected != original_expected

    expected_bytes = {name: (expected_dir / name).read_bytes() for name in ARTIFACT_NAMES}

    out = tmp_path / "generated"
    completed = subprocess.run(
        _reconcile_cmd(out),
        cwd=ROOT,
        check=True,
        capture_output=True,
        text=True,
        env=_subprocess_env(),
    )
    assert "overlap_conflicts=" in completed.stdout

    with pytest.raises(AssertionError, match=drifted_name):
        _assert_artifact_sets_match(out, expected_bytes)

    assert drifted_path.read_bytes() == altered_expected
    for name in ARTIFACT_NAMES:
        if name == drifted_name:
            continue
        assert (expected_dir / name).read_bytes() == (RECONCILED / name).read_bytes()


def test_reconciled_axis_anchors_and_conflicts():
    payload = _load_json(STD_JSON)
    provenance = _load_json(PROV_JSON)
    conflicts = _load_json(CONFLICTS_JSON)
    fin = standardized_from_payload(payload)

    assert fin.ticker == "LULU"
    assert fin.company_name == "lululemon athletica inc."
    assert fin.currency == "USD"
    assert fin.units == "USD in Thousands"
    assert fin.jurisdiction == "US"
    assert [p.end_date for p in fin.periods] == EXPECTED_PERIODS

    # Four FY2022–FY2025 filing year-ends are on the axis. Earlier comparative
    # IS dates present in extracted filings are not promoted onto the canonical axis.
    assert len(fin.periods) == 4

    revenue = next(item for item in fin.income_statement if item.concept == "revenue")
    assert {p: revenue.values[p] for p in EXPECTED_PERIODS} == REVENUE_ANCHORS

    assert fin.historical_shares is not None
    assert fin.historical_shares.scale_basis == "financial_statement_units"
    assert fin.historical_shares.basis == "reported"
    assert fin.historical_shares.diluted_weighted_average == DILUTED_WAS_ANCHORS

    assert conflicts["overlap_conflict_count"] == 3
    assert conflicts["supplemental_conflict_count"] == 0
    assert len(conflicts["conflicts"]) == 3

    source_files = provenance["source_files"]
    assert len(source_files) == 4
    by_year = {entry["filing_year"]: entry for entry in source_files}
    assert set(by_year) == {2022, 2023, 2024, 2025}
    for year, (pdf_name, digest, _) in SOURCE_PDFS.items():
        assert by_year[year]["source_file"] == pdf_name
        assert by_year[year]["source_sha256"] == digest


def test_gift_card_liability_classifies_across_all_periods():
    """G1: unchanged Lululemon gift-card row is operating WC liability in every period."""
    fin = standardized_from_payload(_load_json(STD_JSON))
    item = next(
        row
        for row in fin.balance_sheet
        if row.concept == "unredeemed_gift_card_liability"
        and row.label == "Unredeemed gift card liability"
    )
    assert set(item.values) >= set(EXPECTED_PERIODS)
    decision = classify_balance_sheet_line(item)
    assert decision.category == "Operating Working Capital Liability"
    assert decision.ambiguous is False
    assert decision.judgment_code is None
    assert decision.overridden is False
    for period in EXPECTED_PERIODS:
        assert item.values[period] is not None


def test_ppe_classifies_resolves_and_enables_fixed_asset():
    """G2: PPE row classifies/resolves with all four periods; fixed-asset unlocks."""
    fin = standardized_from_payload(_load_json(STD_JSON))
    original_index, item = next(
        (idx, row)
        for idx, row in enumerate(fin.balance_sheet)
        if row.concept == "property_plant_and_equipment"
        and row.label == "Property and equipment, net"
    )
    assert set(item.values) >= set(EXPECTED_PERIODS)
    for period in EXPECTED_PERIODS:
        assert item.values[period] is not None

    decision = classify_balance_sheet_line(item)
    assert decision.category == "Operating Long-Term Asset"
    assert decision.ambiguous is False
    assert decision.judgment_code is None
    assert decision.overridden is False

    resolved = resolve_line(
        fin.balance_sheet, "property_plant_equipment", required=True
    )
    assert resolved.index == original_index
    assert resolved.item is item
    assert resolved.item.concept == "property_plant_and_equipment"

    avail = fixed_asset_availability(fin)
    assert avail.ppe is True
    assert avail.depreciation_amortization is True
    assert fixed_asset_applicable(fin) is True


COMMON_STOCK_VALUES = {
    date(2023, 1, 29): 611.0,
    date(2024, 1, 28): 606.0,
    date(2025, 2, 2): 581.0,
    date(2026, 2, 1): 557.0,
}


def test_common_stock_classifies_across_all_periods():
    """G3: unchanged Lululemon Common stock row is Equity in every period."""
    fin = standardized_from_payload(_load_json(STD_JSON))
    item = next(
        row
        for row in fin.balance_sheet
        if row.concept == "common_stock" and row.label == "Common stock"
    )
    assert item.concept == "common_stock"
    assert item.label == "Common stock"
    assert set(item.values) >= set(EXPECTED_PERIODS)
    for period, expected in COMMON_STOCK_VALUES.items():
        assert item.values[period] == expected

    decision = classify_balance_sheet_line(item)
    assert decision.category == "Equity"
    assert decision.ambiguous is False
    assert decision.judgment_code is None
    assert decision.overridden is False


def test_non_current_income_taxes_payable_restored_sparse_axis():
    """Sparse NCIT liability is retained with explicit None for FY2026 absence."""
    fin = standardized_from_payload(_load_json(STD_JSON))
    item = next(
        row
        for row in fin.balance_sheet
        if row.concept == "non_current_income_taxes_payable"
        and row.label == "Non-current income taxes payable"
    )
    assert item.values == {
        date(2023, 1, 29): 28555.0,
        date(2024, 1, 28): 15864.0,
        date(2025, 2, 2): 0.0,
        date(2026, 2, 1): None,
    }
    provenance = _load_json(PROV_JSON)
    retained = [
        entry
        for entry in provenance["retained_sparse_axis"]
        if entry["suggested_concept"] == "non_current_income_taxes_payable"
    ]
    assert len(retained) == 1
    assert retained[0]["missing_periods"] == ["2026-02-01"]
    assert retained[0]["available_periods"] == [
        "2023-01-29",
        "2024-01-28",
        "2025-02-02",
    ]
    assert not any(
        entry["suggested_concept"] == "non_current_income_taxes_payable"
        for entry in provenance["omitted_incomplete_axis"]
    )


def test_four_period_reformulation_integrity(tmp_path: Path):
    """Sparse NCIT None is evidence-gated; all four periods reconcile."""
    from core.model.classification import (
        check_reformulation_integrity,
        reformulate_balance_sheet,
    )

    fin = standardized_from_payload(_load_json(STD_JSON))

    unclassified = set()
    for item in fin.balance_sheet:
        if is_balance_sheet_subtotal(item):
            continue
        try:
            classify_balance_sheet_line(item)
        except UnclassifiedBalanceSheetLineError:
            unclassified.add(item.label)
    assert unclassified == BUILD_BLOCKER_LABELS

    gift = next(
        row
        for row in fin.balance_sheet
        if row.label == "Unredeemed gift card liability"
        and row.concept == "unredeemed_gift_card_liability"
    )
    assert (
        classify_balance_sheet_line(gift).category
        == "Operating Working Capital Liability"
    )

    ppe = next(
        row
        for row in fin.balance_sheet
        if row.label == "Property and equipment, net"
        and row.concept == "property_plant_and_equipment"
    )
    assert (
        classify_balance_sheet_line(ppe).category == "Operating Long-Term Asset"
    )

    common = next(
        row
        for row in fin.balance_sheet
        if row.label == "Common stock" and row.concept == "common_stock"
    )
    assert classify_balance_sheet_line(common).category == "Equity"

    ncit = next(
        row
        for row in fin.balance_sheet
        if row.concept == "non_current_income_taxes_payable"
    )
    assert ncit.values == {
        date(2023, 1, 29): 28555.0,
        date(2024, 1, 28): 15864.0,
        date(2025, 2, 2): 0.0,
        date(2026, 2, 1): None,
    }

    reform = reformulate_balance_sheet(fin, EXPECTED_PERIODS)
    assert reform.asset_detail_gap == (0.0, 0.0, 0.0, 0.0)
    assert reform.liability_detail_gap == (0.0, 0.0, 0.0, 0.0)
    assert reform.equity_detail_gap == (0.0, 0.0, 0.0, 0.0)
    assert reform.equity_gap == (0.0, 0.0, 0.0, 0.0)
    check_reformulation_integrity(reform, EXPECTED_PERIODS)
    # Evidence gate must not invent a reported zero into source facts.
    assert ncit.values[date(2026, 2, 1)] is None
    assert ncit.values[date(2025, 2, 2)] == 0.0

    pretax = resolve_line(fin.income_statement, "pretax_income", required=True)
    assert pretax.item is not None
    assert pretax.item.concept == "income_before_tax"
    assert pretax.item.label == "Income before income tax expense"
    assert pretax.item.values == {
        date(2023, 1, 29): 1332571.0,
        date(2024, 1, 28): 2175735.0,
        date(2025, 2, 2): 2576077.0,
        date(2026, 2, 1): 2238967.0,
    }
    pretax_row = workbook_row_for(pretax, start_row=SOURCE_START_ROW)
    assert pretax_row == SOURCE_START_ROW + pretax.index
    assert fin.income_statement[pretax.index] is pretax.item

    tax = resolve_line(fin.income_statement, "tax_expense", required=True)
    assert tax.item is not None
    assert tax.item.label == "Income tax expense"
    assert tax.index != pretax.index

    restored = standardized_from_payload(standardized_to_payload(fin))
    pretax_rt = resolve_line(restored.income_statement, "pretax_income", required=True)
    assert pretax_rt.item is not None
    assert pretax_rt.item.concept == "income_before_tax"
    assert pretax_rt.item.label == "Income before income tax expense"
    assert pretax_rt.item.values == pretax.item.values
    assert pretax_rt.index == pretax.index

    out = tmp_path / "Lululemon"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    builder = ReferenceModelBuilder(fin)
    families = {s.family_id for s in builder.expected_specs}
    assert "nopat_fy" not in families
    assert "net_interest_fy" not in families
    assert "revenue_link" in families
    wb = load_workbook(answer, data_only=False)
    labels = {
        wb["Condensed Financials"].cell(row=r, column=1).value: r
        for r in range(1, wb["Condensed Financials"].max_row + 1)
    }
    assert "Interest Expense" not in labels
    assert "Interest Income" not in labels
    assert (
        wb["Condensed Financials"].cell(row=labels["NOPAT"], column=2).value
        == SOURCE_UNAVAILABLE
    )
    wb.close()
    assert ncit.values[date(2026, 2, 1)] is None
    assert pretax.item.values[date(2026, 2, 1)] == 2238967.0


CAPEX_REPORTED_PAYMENTS = {
    date(2023, 1, 29): -638657.0,
    date(2024, 1, 28): -651865.0,
    date(2025, 2, 2): -689232.0,
    date(2026, 2, 1): -680802.0,
}


def test_source_supported_capex_alias_four_period_diagnostics(tmp_path: Path):
    """G4: CF capital_expenditures resolves through payments_for_ppe without mutation."""
    fin = standardized_from_payload(_load_json(STD_JSON))
    original_index, stored = next(
        (idx, row)
        for idx, row in enumerate(fin.cash_flow)
        if row.concept == "capital_expenditures"
        and row.label == "Purchase of property and equipment"
    )
    assert stored.values == CAPEX_REPORTED_PAYMENTS

    assert capex_applicable(fin) is True
    avail = capex_availability(fin)
    assert avail.payments_for_ppe is True
    assert avail.ambiguous is False

    item = resolve_capex_source(fin)
    assert item is stored
    assert item.concept == "capital_expenditures"
    assert item.label == "Purchase of property and equipment"
    assert item.values == CAPEX_REPORTED_PAYMENTS

    resolved = resolve_line(fin.cash_flow, "payments_for_ppe", required=True)
    assert resolved.index == original_index
    assert resolved.item is stored
    capex_row = workbook_row_for(resolved, start_row=SOURCE_START_ROW)
    assert capex_row == SOURCE_START_ROW + original_index

    revenue = resolve_line(fin.income_statement, "revenue", required=True).item
    assert revenue is not None
    payments = []
    ppe_capex = []
    ratios = []
    for period in EXPECTED_PERIODS:
        pay = required_period_value(stored, period, field="payments_for_ppe")
        rev = required_period_value(revenue, period, field="revenue")
        capex = -pay
        payments.append(pay)
        ppe_capex.append(capex)
        ratios.append(capex / rev)
        assert pay == CAPEX_REPORTED_PAYMENTS[period]
        assert rev == REVENUE_ANCHORS[period]
    assert payments == [
        -638657.0,
        -651865.0,
        -689232.0,
        -680802.0,
    ]
    assert ppe_capex == [638657.0, 651865.0, 689232.0, 680802.0]
    assert ratios[0] == pytest.approx(638657.0 / 8110518.0)
    assert ratios[1] == pytest.approx(651865.0 / 9619278.0)
    assert ratios[2] == pytest.approx(689232.0 / 10588126.0)
    assert ratios[3] == pytest.approx(680802.0 / 11102600.0)

    cfo_item = resolve_operating_cash_source(fin)
    assert cfo_item is not None
    assert cfo_item.label == "Net cash provided by operating activities"
    assert cash_after_ppe_capex_applicable(fin) is True
    cash_after = []
    cash_margins = []
    for period in EXPECTED_PERIODS:
        cfo = required_period_value(cfo_item, period, field="operating_cash_flow")
        pay = required_period_value(stored, period, field="payments_for_ppe")
        value = cfo - (-pay)
        cash_after.append(value)
        cash_margins.append(value / REVENUE_ANCHORS[period])
    assert cash_after == [327806.0, 1644299.0, 1583481.0, 921675.0]
    series = compute_capex_series(fin, list(EXPECTED_PERIODS), compute_anchor(fin, list(EXPECTED_PERIODS)))
    assert series.cash_after_ppe_capex == tuple(cash_after)
    for j, margin in enumerate(cash_margins):
        assert series.cash_after_ppe_capex_to_revenue[j] == pytest.approx(margin)
    builder = ReferenceModelBuilder(fin)
    cash_ids = {
        s.semantic_key
        for s in builder.expected_specs
        if s.family_id
        in {"cash_after_ppe_capex", "cash_after_ppe_capex_to_revenue"}
    }
    assert len(cash_ids) == 8

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt_item = resolve_capex_source(restored)
    assert rt_item is not None
    assert rt_item.concept == "capital_expenditures"
    assert rt_item.label == "Purchase of property and equipment"
    assert rt_item.values == CAPEX_REPORTED_PAYMENTS
    assert capex_applicable(restored) is True
    rt_resolved = resolve_line(restored.cash_flow, "payments_for_ppe", required=True)
    assert rt_resolved.index == original_index

    out = tmp_path / "LululemonCapex"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    assert stored.concept == "capital_expenditures"
    assert stored.values == CAPEX_REPORTED_PAYMENTS


SBC_REPORTED = {
    date(2023, 1, 29): 78075.0,
    date(2024, 1, 28): 93560.0,
    date(2025, 2, 2): 90011.0,
    date(2026, 2, 1): 62203.0,
}
CFO_LESS_SBC = {
    date(2023, 1, 29): 888388.0,
    date(2024, 1, 28): 2202604.0,
    date(2025, 2, 2): 2182702.0,
    date(2026, 2, 1): 1540274.0,
}


def test_source_supported_sbc_four_period_diagnostics(tmp_path: Path):
    from core.model.earnings_quality import (
        compute_earnings_quality_series,
        resolve_sbc_source,
        sbc_diagnostics_applicable,
    )

    fin = standardized_from_payload(_load_json(STD_JSON))
    original_index, stored = next(
        (idx, row)
        for idx, row in enumerate(fin.cash_flow)
        if row.concept == "stock_based_compensation"
        and row.label == "Stock-based compensation expense"
    )
    assert stored.values == SBC_REPORTED
    assert sbc_diagnostics_applicable(fin) is True
    item = resolve_sbc_source(fin)
    assert item is stored
    resolved = resolve_line(fin.cash_flow, "stock_based_compensation", required=True)
    assert resolved.index == original_index
    assert resolved.item is stored
    assert workbook_row_for(resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + original_index
    )

    cfo_item = resolve_operating_cash_source(fin)
    assert cfo_item is not None
    revenue = resolve_line(fin.income_statement, "revenue", required=True).item
    assert revenue is not None
    independent_less = []
    for period in EXPECTED_PERIODS:
        sbc = required_period_value(stored, period, field="stock_based_compensation")
        cfo = required_period_value(cfo_item, period, field="operating_cash_flow")
        rev = required_period_value(revenue, period, field="revenue")
        independent_less.append(cfo - sbc)
        assert sbc == SBC_REPORTED[period]
        assert rev == REVENUE_ANCHORS[period]
        assert cfo - sbc == CFO_LESS_SBC[period]
    assert independent_less == [888388.0, 2202604.0, 2182702.0, 1540274.0]

    series = compute_earnings_quality_series(
        fin, list(EXPECTED_PERIODS), compute_anchor(fin, list(EXPECTED_PERIODS))
    )
    assert series.stock_based_compensation == (
        78075.0,
        93560.0,
        90011.0,
        62203.0,
    )
    assert series.operating_cash_flow_less_sbc == tuple(independent_less)
    for j, period in enumerate(EXPECTED_PERIODS):
        sbc = SBC_REPORTED[period]
        rev = REVENUE_ANCHORS[period]
        cfo = required_period_value(cfo_item, period, field="operating_cash_flow")
        assert series.sbc_to_revenue[j] == pytest.approx(sbc / rev)
        assert series.sbc_to_operating_cash_flow[j] == pytest.approx(sbc / cfo)

    builder = ReferenceModelBuilder(fin)
    sbc_ids = {
        s.semantic_key
        for s in builder.expected_specs
        if s.family_id
        in {
            "sbc_to_revenue",
            "sbc_to_operating_cash_flow",
            "operating_cash_flow_less_sbc",
        }
    }
    assert len(sbc_ids) == 12
    assert len(builder.expected_specs) == LEASE_DT_LULULEMON_SPECS

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt_item = resolve_sbc_source(restored)
    assert rt_item is not None
    assert rt_item.concept == "stock_based_compensation"
    assert rt_item.label == "Stock-based compensation expense"
    assert rt_item.values == SBC_REPORTED

    out = tmp_path / "LululemonSBC"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == LEASE_DT_LULULEMON_SPECS
    awb = load_workbook(answer, data_only=False)
    ws = awb[EARNINGS_QUALITY_SHEET]
    reported_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Stock-based compensation (reported)"
    )
    less_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Reported CFO less SBC add-back"
    )
    src_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
    assert src_f.startswith("='CashFlowStatement'!")
    less_f = str(ws.cell(less_row, 2).value).replace(" ", "")
    assert less_f.startswith("=") and "-" in less_f
    note_cell = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "operating_cash_flow_less_sbc"
    )
    nrow, ncol = parse_cell_ref(note_cell.cell)
    note = (ws.cell(nrow, ncol).comment.text or "") if ws.cell(nrow, ncol).comment else ""
    assert "does not restate reported CFO" in note
    assert "dilution" in note.lower()
    assert "free cash flow" in note.lower()
    assert "tax" in note.lower()
    awb.close()
    assert stored.concept == "stock_based_compensation"
    assert stored.values == SBC_REPORTED


ACQUISITION_REPORTED = {
    date(2023, 1, 29): 0.0,
    date(2024, 1, 28): 0.0,
    date(2025, 2, 2): -154146.0,
    date(2026, 2, 1): 0.0,
}
ACQUISITION_OUTFLOW = {
    date(2023, 1, 29): 0.0,
    date(2024, 1, 28): 0.0,
    date(2025, 2, 2): 154146.0,
    date(2026, 2, 1): 0.0,
}
CASH_AFTER_PPE_CAPEX_AND_ACQUISITIONS = {
    date(2023, 1, 29): 327806.0,
    date(2024, 1, 28): 1644299.0,
    date(2025, 2, 2): 1429335.0,
    date(2026, 2, 1): 921675.0,
}


def test_source_supported_acquisition_cash_four_period_diagnostics(tmp_path: Path):
    from core.model.acquisition_cash import (
        acquisition_cash_applicable,
        cash_after_ppe_capex_and_acquisitions_applicable,
        compute_acquisition_cash_series,
        resolve_acquisition_cash_source,
    )

    fin = standardized_from_payload(_load_json(STD_JSON))
    original_index, stored = next(
        (idx, row)
        for idx, row in enumerate(fin.cash_flow)
        if row.concept == "acquisition_net_of_cash_acquired"
        and row.label == "Acquisition, net of cash acquired"
    )
    assert stored.values == ACQUISITION_REPORTED
    assert acquisition_cash_applicable(fin) is True
    assert cash_after_ppe_capex_and_acquisitions_applicable(fin) is True
    item = resolve_acquisition_cash_source(fin)
    assert item is stored
    resolved = resolve_line(
        fin.cash_flow, "acquisition_net_of_cash_acquired", required=True
    )
    assert resolved.index == original_index
    assert resolved.item is stored
    assert workbook_row_for(resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + original_index
    )

    revenue = resolve_line(fin.income_statement, "revenue", required=True).item
    assert revenue is not None
    cfo_item = resolve_operating_cash_source(fin)
    assert cfo_item is not None
    capex_item = resolve_capex_source(fin)
    assert capex_item is not None
    independent_outflow = []
    independent_residual = []
    independent_ratios = []
    for period in EXPECTED_PERIODS:
        reported = required_period_value(
            stored, period, field="acquisition_net_of_cash_acquired"
        )
        rev = required_period_value(revenue, period, field="revenue")
        cfo = required_period_value(cfo_item, period, field="operating_cash_flow")
        pay = required_period_value(capex_item, period, field="payments_for_ppe")
        outflow = -reported
        residual = cfo - (-pay) - outflow
        independent_outflow.append(outflow)
        independent_residual.append(residual)
        independent_ratios.append(outflow / rev if rev else None)
        assert reported == ACQUISITION_REPORTED[period]
        assert outflow == ACQUISITION_OUTFLOW[period]
        assert residual == CASH_AFTER_PPE_CAPEX_AND_ACQUISITIONS[period]
        assert rev == REVENUE_ANCHORS[period]
    assert independent_outflow == [0.0, 0.0, 154146.0, 0.0]
    assert independent_residual == [327806.0, 1644299.0, 1429335.0, 921675.0]

    series = compute_acquisition_cash_series(
        fin, list(EXPECTED_PERIODS), compute_anchor(fin, list(EXPECTED_PERIODS))
    )
    assert series.payments_reported == (0.0, 0.0, -154146.0, 0.0)
    assert series.acquisition_cash_outflow == tuple(independent_outflow)
    assert series.cash_after_ppe_capex_and_acquisitions == tuple(independent_residual)
    for j, ratio in enumerate(independent_ratios):
        assert series.acquisition_cash_to_revenue[j] == pytest.approx(ratio)

    builder = ReferenceModelBuilder(fin)
    acq_ids = {
        s.semantic_key
        for s in builder.expected_specs
        if s.family_id
        in {
            "acquisition_cash_outflow",
            "acquisition_cash_to_revenue",
            "cash_after_ppe_capex_and_acquisitions",
        }
    }
    assert len(acq_ids) == 12
    assert len(builder.expected_specs) == LEASE_DT_LULULEMON_SPECS

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt_item = resolve_acquisition_cash_source(restored)
    assert rt_item is not None
    assert rt_item.concept == "acquisition_net_of_cash_acquired"
    assert rt_item.label == "Acquisition, net of cash acquired"
    assert rt_item.values == ACQUISITION_REPORTED

    out = tmp_path / "LululemonAcq"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == LEASE_DT_LULULEMON_SPECS
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    reported_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Acquisition, net of cash acquired (reported)"
    )
    residual_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Operating cash after PP&E capex and acquisitions"
    )
    src_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
    assert src_f.startswith("='CashFlowStatement'!") or src_f.startswith(
        "='Cash Flow Statement'!"
    )
    residual_f = str(ws.cell(residual_row, 2).value).replace(" ", "")
    assert residual_f.startswith("=") and residual_f.count("-") >= 2
    note_cell = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "cash_after_ppe_capex_and_acquisitions"
    )
    nrow, ncol = parse_cell_ref(note_cell.cell)
    note = (ws.cell(nrow, ncol).comment.text or "") if ws.cell(nrow, ncol).comment else ""
    assert "net of cash acquired" in note.lower()
    assert "free cash flow" in note.lower()
    assert "purchase-price" in note.lower() or "purchase price" in note.lower()
    assert "goodwill" in note.lower()
    awb.close()
    assert stored.concept == "acquisition_net_of_cash_acquired"
    assert stored.values == ACQUISITION_REPORTED


REPURCHASE_REPORTED = {
    date(2023, 1, 29): -444001.0,
    date(2024, 1, 28): -558652.0,
    date(2025, 2, 2): -1636879.0,
    date(2026, 2, 1): -1178349.0,
}
REPURCHASE_OUTFLOW = {
    date(2023, 1, 29): 444001.0,
    date(2024, 1, 28): 558652.0,
    date(2025, 2, 2): 1636879.0,
    date(2026, 2, 1): 1178349.0,
}
CASH_AFTER_PPE_CAPEX_ACQUISITIONS_AND_REPURCHASES = {
    date(2023, 1, 29): -116195.0,
    date(2024, 1, 28): 1085647.0,
    date(2025, 2, 2): -207544.0,
    date(2026, 2, 1): -256674.0,
}


def test_source_supported_share_repurchase_four_period_diagnostics(tmp_path: Path):
    from core.model.share_repurchase import (
        cash_after_ppe_capex_acquisitions_and_repurchases_applicable,
        compute_share_repurchase_series,
        resolve_share_repurchase_source,
        share_repurchase_applicable,
    )

    fin = standardized_from_payload(_load_json(STD_JSON))
    original_index, stored = next(
        (idx, row)
        for idx, row in enumerate(fin.cash_flow)
        if row.concept == "repurchase_of_common_stock"
        and row.label == "Repurchase of common stock"
    )
    assert stored.values == REPURCHASE_REPORTED
    assert share_repurchase_applicable(fin) is True
    assert cash_after_ppe_capex_acquisitions_and_repurchases_applicable(fin) is True
    item = resolve_share_repurchase_source(fin)
    assert item is stored
    resolved = resolve_line(
        fin.cash_flow, "repurchase_of_common_stock", required=True
    )
    assert resolved.index == original_index
    assert resolved.item is stored
    assert workbook_row_for(resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + original_index
    )

    revenue = resolve_line(fin.income_statement, "revenue", required=True).item
    assert revenue is not None
    cfo_item = resolve_operating_cash_source(fin)
    assert cfo_item is not None
    capex_item = resolve_capex_source(fin)
    assert capex_item is not None
    acq_item = next(
        row
        for row in fin.cash_flow
        if row.concept == "acquisition_net_of_cash_acquired"
    )
    independent_outflow = []
    independent_residual = []
    independent_ratios = []
    for period in EXPECTED_PERIODS:
        reported = required_period_value(
            stored, period, field="repurchase_of_common_stock"
        )
        rev = required_period_value(revenue, period, field="revenue")
        cfo = required_period_value(cfo_item, period, field="operating_cash_flow")
        pay = required_period_value(capex_item, period, field="payments_for_ppe")
        acq = required_period_value(
            acq_item, period, field="acquisition_net_of_cash_acquired"
        )
        outflow = -reported
        residual = cfo - (-pay) - (-acq) - outflow
        independent_outflow.append(outflow)
        independent_residual.append(residual)
        independent_ratios.append(outflow / rev if rev else None)
        assert reported == REPURCHASE_REPORTED[period]
        assert outflow == REPURCHASE_OUTFLOW[period]
        assert residual == CASH_AFTER_PPE_CAPEX_ACQUISITIONS_AND_REPURCHASES[period]
        assert rev == REVENUE_ANCHORS[period]
    assert independent_outflow == [444001.0, 558652.0, 1636879.0, 1178349.0]
    assert independent_residual == [-116195.0, 1085647.0, -207544.0, -256674.0]

    series = compute_share_repurchase_series(
        fin, list(EXPECTED_PERIODS), compute_anchor(fin, list(EXPECTED_PERIODS))
    )
    assert series.payments_reported == (-444001.0, -558652.0, -1636879.0, -1178349.0)
    assert series.share_repurchase_outflow == tuple(independent_outflow)
    assert series.cash_after_ppe_capex_acquisitions_and_repurchases == tuple(
        independent_residual
    )
    for j, ratio in enumerate(independent_ratios):
        assert series.share_repurchase_to_revenue[j] == pytest.approx(ratio)

    builder = ReferenceModelBuilder(fin)
    rp_ids = {
        s.semantic_key
        for s in builder.expected_specs
        if s.family_id
        in {
            "share_repurchase_outflow",
            "share_repurchase_to_revenue",
            "cash_after_ppe_capex_acquisitions_and_repurchases",
        }
    }
    assert len(rp_ids) == 12
    assert len(builder.expected_specs) == LEASE_DT_LULULEMON_SPECS

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt_item = resolve_share_repurchase_source(restored)
    assert rt_item is not None
    assert rt_item.concept == "repurchase_of_common_stock"
    assert rt_item.label == "Repurchase of common stock"
    assert rt_item.values == REPURCHASE_REPORTED

    out = tmp_path / "LululemonRP"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == LEASE_DT_LULULEMON_SPECS
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    reported_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Repurchase of common stock (reported)"
    )
    residual_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value
        == "Operating cash after PP&E capex, acquisitions and repurchases"
    )
    src_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
    assert src_f.startswith("='CashFlowStatement'!") or src_f.startswith(
        "='Cash Flow Statement'!"
    )
    residual_f = str(ws.cell(residual_row, 2).value).replace(" ", "")
    assert residual_f.startswith("=") and residual_f.count("-") >= 3
    note_cell = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "cash_after_ppe_capex_acquisitions_and_repurchases"
    )
    nrow, ncol = parse_cell_ref(note_cell.cell)
    note = (ws.cell(nrow, ncol).comment.text or "") if ws.cell(nrow, ncol).comment else ""
    assert "exceed" in note.lower()
    assert "debt" in note.lower()
    assert "free cash flow" in note.lower()
    outflow_note_cell = next(
        c for c in smap.all_ordered() if c.family_id == "share_repurchase_outflow"
    )
    orow, ocol = parse_cell_ref(outflow_note_cell.cell)
    outflow_note = (
        (ws.cell(orow, ocol).comment.text or "")
        if ws.cell(orow, ocol).comment
        else ""
    )
    assert "distribution" in outflow_note.lower()
    assert "dilution" in outflow_note.lower()
    awb.close()
    assert stored.concept == "repurchase_of_common_stock"
    assert stored.values == REPURCHASE_REPORTED


CASH_MOVEMENT_FROM_FLOWS = {
    date(2023, 1, 29): -105004.0,
    date(2024, 1, 28): 1089104.0,
    date(2025, 2, 2): -259635.0,
    date(2026, 2, 1): -177134.0,
}


def test_source_supported_cash_rollforward_four_period_diagnostics(tmp_path: Path):
    from core.model.cash_rollforward import (
        cash_rollforward_applicable,
        compute_cash_rollforward_series,
        resolve_cash_rollforward_sources,
    )

    fin = standardized_from_payload(_load_json(STD_JSON))
    original_index, stored = next(
        (idx, row)
        for idx, row in enumerate(fin.cash_flow)
        if row.concept == "net_cash_from_operating_activities"
        and row.label == "Net cash provided by operating activities"
    )
    fx_item = next(
        row for row in fin.cash_flow if row.concept == "effect_of_fx_on_cash"
    )
    change_item = next(row for row in fin.cash_flow if row.concept == "change_in_cash")
    beginning_item = next(
        row for row in fin.cash_flow if row.concept == "cash_beginning"
    )
    ending_item = next(row for row in fin.cash_flow if row.concept == "cash_ending")
    investing_item = next(
        row
        for row in fin.cash_flow
        if row.concept == "net_cash_from_investing_activities"
    )
    financing_item = next(
        row
        for row in fin.cash_flow
        if row.concept == "net_cash_from_financing_activities"
    )
    assert cash_rollforward_applicable(fin) is True
    sources = resolve_cash_rollforward_sources(fin)
    assert sources.operating is stored
    assert sources.fx is fx_item
    resolved = resolve_line(
        fin.cash_flow, "net_cash_from_operating_activities", required=True
    )
    assert resolved.index == original_index
    assert resolved.item is stored
    assert workbook_row_for(resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + original_index
    )

    independent_movement = []
    independent_move_diff = []
    independent_ending = []
    independent_end_diff = []
    for period in EXPECTED_PERIODS:
        cfo = required_period_value(
            stored, period, field="net_cash_from_operating_activities"
        )
        cfi = required_period_value(
            investing_item, period, field="net_cash_from_investing_activities"
        )
        cff = required_period_value(
            financing_item, period, field="net_cash_from_financing_activities"
        )
        fx = required_period_value(fx_item, period, field="effect_of_fx_on_cash")
        reported_change = required_period_value(
            change_item, period, field="change_in_cash"
        )
        beginning = required_period_value(
            beginning_item, period, field="cash_beginning"
        )
        ending = required_period_value(ending_item, period, field="cash_ending")
        movement = cfo + cfi + cff + fx
        independent_movement.append(movement)
        independent_move_diff.append(movement - reported_change)
        ending_from_flows = beginning + movement
        independent_ending.append(ending_from_flows)
        independent_end_diff.append(ending_from_flows - ending)
        assert movement == CASH_MOVEMENT_FROM_FLOWS[period]
    assert independent_movement == [-105004.0, 1089104.0, -259635.0, -177134.0]
    assert independent_move_diff == [0.0, 0.0, 0.0, 0.0]
    assert independent_end_diff == [0.0, 0.0, 0.0, 0.0]

    series = compute_cash_rollforward_series(fin, list(EXPECTED_PERIODS))
    assert series.cash_movement_from_flows == tuple(independent_movement)
    assert series.cash_movement_difference == tuple(independent_move_diff)
    assert series.cash_ending_from_flows == tuple(independent_ending)
    assert series.cash_ending_difference == tuple(independent_end_diff)

    builder = ReferenceModelBuilder(fin)
    cr_ids = {
        s.semantic_key
        for s in builder.expected_specs
        if s.family_id
        in {
            "cash_movement_from_flows",
            "cash_movement_difference",
            "cash_ending_from_flows",
            "cash_ending_difference",
        }
    }
    assert len(cr_ids) == 16
    assert len(builder.expected_specs) == LEASE_DT_LULULEMON_SPECS

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt = resolve_cash_rollforward_sources(restored)
    assert rt.operating is not None
    assert rt.operating.concept == "net_cash_from_operating_activities"
    assert rt.fx.concept == "effect_of_fx_on_cash"
    assert rt.change.concept == "change_in_cash"

    out = tmp_path / "LululemonCR"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == LEASE_DT_LULULEMON_SPECS
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    reported_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Net cash from operating activities (reported)"
    )
    src_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
    assert src_f.startswith("='CashFlowStatement'!") or src_f.startswith(
        "='Cash Flow Statement'!"
    )
    note_cell = next(
        c for c in smap.all_ordered() if c.family_id == "cash_movement_difference"
    )
    nrow, ncol = parse_cell_ref(note_cell.cell)
    note = (ws.cell(nrow, ncol).comment.text or "") if ws.cell(nrow, ncol).comment else ""
    assert "plug" in note.lower()
    assert "statement" in note.lower()
    awb.close()
    assert stored.concept == "net_cash_from_operating_activities"
    assert fx_item.values == {
        date(2023, 1, 29): -34043.0,
        date(2024, 1, 28): -4100.0,
        date(2025, 2, 2): -81666.0,
        date(2026, 2, 1): 91163.0,
    }


def test_source_supported_reported_margin_four_period_diagnostics(tmp_path: Path):
    from core.model.reported_margin import (
        compute_reported_margin_series,
        reported_margin_applicable,
        resolve_reported_margin_sources,
    )
    from core.model.ratio_values import UNDEFINED_RATIO

    fin = standardized_from_payload(_load_json(STD_JSON))
    original_index, stored = next(
        (idx, row)
        for idx, row in enumerate(fin.income_statement)
        if row.concept == "gross_profit" and row.label == "Gross profit"
    )
    operating_item = next(
        row for row in fin.income_statement if row.concept == "operating_income"
    )
    revenue_item = next(
        row for row in fin.income_statement if row.concept == "revenue"
    )
    assert reported_margin_applicable(fin) is True
    sources = resolve_reported_margin_sources(fin)
    assert sources.gross_profit is stored
    assert sources.operating_profit is operating_item
    resolved = resolve_line(fin.income_statement, "gross_profit", required=True)
    assert resolved.index == original_index
    assert resolved.item is stored
    assert workbook_row_for(resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + original_index
    )

    independent_gm = []
    independent_om = []
    independent_burden = []
    independent_gm_change = [None]
    independent_burden_change = [None]
    independent_recon = [None]
    for j, period in enumerate(EXPECTED_PERIODS):
        gp = required_period_value(stored, period, field="gross_profit")
        op = required_period_value(operating_item, period, field="operating_profit")
        rev = required_period_value(revenue_item, period, field="revenue")
        gm = gp / rev
        om = op / rev
        burden = (gp - op) / rev
        independent_gm.append(gm)
        independent_om.append(om)
        independent_burden.append(burden)
        if j > 0:
            independent_gm_change.append(gm - independent_gm[j - 1])
            independent_burden_change.append(burden - independent_burden[j - 1])
            independent_recon.append(
                independent_gm_change[j] - independent_burden_change[j]
            )
            assert independent_recon[j] == pytest.approx(om - independent_om[j - 1])
    assert independent_gm[-1] == pytest.approx(6284132.0 / 11102600.0)
    assert independent_om[-1] == pytest.approx(2210615.0 / 11102600.0)
    assert independent_gm[-1] == pytest.approx(0.566005, abs=5e-7)
    assert independent_om[-1] == pytest.approx(0.199108, abs=5e-7)
    assert UNDEFINED_RATIO not in independent_gm

    series = compute_reported_margin_series(fin, list(EXPECTED_PERIODS))
    assert series.gross_margin == tuple(independent_gm)
    assert series.reported_operating_margin == tuple(independent_om)
    assert series.net_operating_expense_burden == tuple(independent_burden)
    assert series.reconstructed_operating_margin_change[1:] == tuple(
        independent_recon[1:]
    )

    builder = ReferenceModelBuilder(fin)
    rm_ids = {
        s.semantic_key
        for s in builder.expected_specs
        if s.family_id
        in {
            "gross_margin",
            "reported_operating_margin",
            "net_operating_expense_burden",
            "gross_margin_change",
            "net_operating_expense_burden_change",
            "reconstructed_operating_margin_change",
        }
    }
    assert len(rm_ids) == 21
    assert len(builder.expected_specs) == LEASE_DT_LULULEMON_SPECS

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt = resolve_reported_margin_sources(restored)
    assert rt.gross_profit is not None
    assert rt.gross_profit.concept == "gross_profit"
    assert rt.operating_profit.concept == "operating_income"

    out = tmp_path / "LululemonRM"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == LEASE_DT_LULULEMON_SPECS
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    reported_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Gross profit (reported)"
    )
    src_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
    assert src_f.startswith("='IncomeStatement'!") or src_f.startswith(
        "='Income Statement'!"
    )
    note_cell = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "net_operating_expense_burden"
    )
    nrow, ncol = parse_cell_ref(note_cell.cell)
    note = (ws.cell(nrow, ncol).comment.text or "") if ws.cell(nrow, ncol).comment else ""
    assert "sg&a" in note.lower() or "intervening" in note.lower()
    awb.close()
    assert stored.concept == "gross_profit"
    assert stored.values == {
        date(2023, 1, 29): 4492340.0,
        date(2024, 1, 28): 5609405.0,
        date(2025, 2, 2): 6270811.0,
        date(2026, 2, 1): 6284132.0,
    }


def test_source_supported_inventory_analysis_four_period_diagnostics(tmp_path: Path):
    from core.model.inventory_analysis import (
        compute_inventory_analysis_series,
        inventory_analysis_applicable,
        resolve_inventory_analysis_sources,
    )

    fin = standardized_from_payload(_load_json(STD_JSON))
    original_index, stored = next(
        (idx, row)
        for idx, row in enumerate(fin.balance_sheet)
        if row.concept == "inventories" and row.label == "Inventories"
    )
    revenue_item = next(
        row for row in fin.income_statement if row.concept == "revenue"
    )
    assert inventory_analysis_applicable(fin) is True
    sources = resolve_inventory_analysis_sources(fin)
    assert sources.inventories is stored
    resolved = resolve_line(fin.balance_sheet, "inventories", required=True)
    assert resolved.index == original_index
    assert resolved.item is stored
    assert workbook_row_for(resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + original_index
    )

    independent_intensity = []
    independent_change = [None]
    independent_scale = [None]
    independent_intensity_effect = [None]
    independent_recon = [None]
    independent_implied = [None]
    independent_cf_diff = [None]
    cf_item = next(
        row for row in fin.cash_flow if row.concept == "change_in_inventories"
    )
    for j, period in enumerate(EXPECTED_PERIODS):
        inv = required_period_value(stored, period, field="inventories")
        rev = required_period_value(revenue_item, period, field="revenue")
        intensity = inv / rev
        independent_intensity.append(intensity)
        if j > 0:
            prior_inv = required_period_value(
                stored, EXPECTED_PERIODS[j - 1], field="inventories"
            )
            prior_rev = required_period_value(
                revenue_item, EXPECTED_PERIODS[j - 1], field="revenue"
            )
            prior_intensity = prior_inv / prior_rev
            change = inv - prior_inv
            implied = -change
            reported_cf = required_period_value(
                cf_item, period, field="change_in_inventories"
            )
            cf_diff = reported_cf - implied
            scale = prior_intensity * (rev - prior_rev)
            intensity_effect = rev * (intensity - prior_intensity)
            independent_change.append(change)
            independent_implied.append(implied)
            independent_cf_diff.append(cf_diff)
            independent_scale.append(scale)
            independent_intensity_effect.append(intensity_effect)
            independent_recon.append(scale + intensity_effect)
            assert independent_recon[j] == pytest.approx(change, abs=1e-8)
            assert implied + cf_diff == pytest.approx(reported_cf, abs=1e-8)
    assert independent_intensity[-1] == pytest.approx(1700753.0 / 11102600.0)
    assert independent_intensity[-1] == pytest.approx(0.15318510979410227)
    assert independent_change[-1] == pytest.approx(258672.0)
    assert independent_implied[-1] == pytest.approx(-258672.0)
    assert independent_cf_diff == [None, -57181.0, -37606.0, 69962.0]
    assert independent_implied[-1] + independent_cf_diff[-1] == pytest.approx(
        -188710.0, abs=1e-8
    )
    assert independent_scale[-1] == pytest.approx(70070.30142954475)
    assert independent_intensity_effect[-1] == pytest.approx(188601.69857045508)

    series = compute_inventory_analysis_series(fin, list(EXPECTED_PERIODS))
    assert series.inventory_intensity == tuple(independent_intensity)
    assert series.inventory_change[1:] == tuple(independent_change[1:])
    assert series.inventory_balance_implied_cf_adjustment[1:] == tuple(
        independent_implied[1:]
    )
    assert series.inventory_cf_adjustment_difference[1:] == tuple(
        independent_cf_diff[1:]
    )
    assert series.reconstructed_inventory_change[1:] == tuple(independent_recon[1:])

    builder = ReferenceModelBuilder(fin)
    inv_ids = {
        s.semantic_key
        for s in builder.expected_specs
        if s.family_id
        in {
            "inventory_intensity",
            "inventory_change",
            "inventory_revenue_scale_effect",
            "inventory_intensity_effect",
            "reconstructed_inventory_change",
            "inventory_balance_implied_cf_adjustment",
            "inventory_cf_adjustment_difference",
        }
    }
    assert len(inv_ids) == 22
    assert len(builder.expected_specs) == LEASE_DT_LULULEMON_SPECS

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt = resolve_inventory_analysis_sources(restored)
    assert rt.inventories is not None
    assert rt.inventories.concept == "inventories"

    out = tmp_path / "LululemonINV"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == LEASE_DT_LULULEMON_SPECS
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    reported_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Inventory (reported)"
    )
    src_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
    assert src_f.startswith("='BalanceSheet'!") or src_f.startswith(
        "='Balance Sheet'!"
    )
    note_cell = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "inventory_cf_adjustment_difference"
    )
    nrow, ncol = parse_cell_ref(note_cell.cell)
    note = (ws.cell(nrow, ncol).comment.text or "") if ws.cell(nrow, ncol).comment else ""
    assert "further evidence" in note.lower()
    assert "balancing plug" in note.lower()
    awb.close()
    assert stored.concept == "inventories"
    assert stored.values == {
        date(2023, 1, 29): 1447367.0,
        date(2024, 1, 28): 1323602.0,
        date(2025, 2, 2): 1442081.0,
        date(2026, 2, 1): 1700753.0,
    }


ROU_BALANCES = {
    date(2023, 1, 29): 969419.0,
    date(2024, 1, 28): 1265610.0,
    date(2025, 2, 2): 1416256.0,
    date(2026, 2, 1): 1630181.0,
}
DTA_BALANCES = {
    date(2023, 1, 29): 6402.0,
    date(2024, 1, 28): 9176.0,
    date(2025, 2, 2): 17085.0,
    date(2026, 2, 1): 24037.0,
}
DTL_BALANCES = {
    date(2023, 1, 29): 55084.0,
    date(2024, 1, 28): 29522.0,
    date(2025, 2, 2): 98188.0,
    date(2026, 2, 1): 52278.0,
}
NET_DT_POSITIONS = {
    date(2023, 1, 29): -48682.0,
    date(2024, 1, 28): -20346.0,
    date(2025, 2, 2): -81103.0,
    date(2026, 2, 1): -28241.0,
}
PRIOR_LULULEMON_SPECS = 351
LEASE_DT_LULULEMON_SPECS = 376


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _count_source_unavailable(path: Path) -> int:
    wb = load_workbook(path, data_only=False)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == SOURCE_UNAVAILABLE:
                    n += 1
    wb.close()
    return n


def test_source_supported_lease_rou_and_deferred_tax_aliases(
    tmp_path: Path, monkeypatch
):
    """G5: supplied ROU/DTA/DTL aliases activate existing diagnostics without mutation."""
    from core.model import line_resolver as lr
    from core.tests.test_normalization import _inject_formula_and_cached_value

    fin = standardized_from_payload(_load_json(STD_JSON))
    rou_index, rou_stored = next(
        (idx, row)
        for idx, row in enumerate(fin.balance_sheet)
        if row.concept == "right_of_use_lease_asset"
        and row.label == "Right-of-use lease assets"
    )
    dta_index, dta_stored = next(
        (idx, row)
        for idx, row in enumerate(fin.balance_sheet)
        if row.concept == "deferred_tax_asset"
        and row.label == "Deferred income tax assets"
    )
    dtl_index, dtl_stored = next(
        (idx, row)
        for idx, row in enumerate(fin.balance_sheet)
        if row.concept == "deferred_tax_liability"
        and row.label == "Deferred income tax liabilities"
    )
    assert rou_stored.values == ROU_BALANCES
    assert dta_stored.values == DTA_BALANCES
    assert dtl_stored.values == DTL_BALANCES

    assert lease_rou_applicable(fin) is True
    rou_avail = lease_rou_availability(fin)
    assert rou_avail.right_of_use_assets is True
    assert rou_avail.ambiguous is False
    assert resolve_lease_rou_source(fin) is rou_stored

    assert deferred_tax_applicable(fin) is True
    dt_avail = deferred_tax_availability(fin)
    assert dt_avail.deferred_tax_assets is True
    assert dt_avail.deferred_tax_liabilities is True
    assert dt_avail.ambiguous is False
    dt_sources = resolve_deferred_tax_sources(fin)
    assert dt_sources is not None
    assert dt_sources.deferred_tax_assets is dta_stored
    assert dt_sources.deferred_tax_liabilities is dtl_stored

    rou_resolved = resolve_line(fin.balance_sheet, "right_of_use_assets", required=True)
    dta_resolved = resolve_line(fin.balance_sheet, "deferred_tax_assets", required=True)
    dtl_resolved = resolve_line(
        fin.balance_sheet, "deferred_tax_liabilities", required=True
    )
    assert rou_resolved.index == rou_index and rou_resolved.item is rou_stored
    assert dta_resolved.index == dta_index and dta_resolved.item is dta_stored
    assert dtl_resolved.index == dtl_index and dtl_resolved.item is dtl_stored
    assert workbook_row_for(rou_resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + rou_index
    )
    assert workbook_row_for(dta_resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + dta_index
    )
    assert workbook_row_for(dtl_resolved, start_row=SOURCE_START_ROW) == (
        SOURCE_START_ROW + dtl_index
    )

    periods = list(canonical_fiscal_periods(fin))
    assert periods == EXPECTED_PERIODS
    rou_series = compute_lease_rou_series(fin, periods, compute_anchor(fin, periods))
    dt_series = compute_deferred_tax_series(fin, periods)
    assert rou_series.rou_assets == (
        969419.0,
        1265610.0,
        1416256.0,
        1630181.0,
    )
    assert rou_series.rou_assets_change == (None, 296191.0, 150646.0, 213925.0)
    assert rou_series.average_rou_assets[1] == pytest.approx(1117514.5)
    assert rou_series.average_rou_assets[2] == pytest.approx(1340933.0)
    assert rou_series.average_rou_assets[3] == pytest.approx(1523218.5)
    assert rou_series.rou_assets_growth[1] == pytest.approx(296191.0 / 969419.0)
    assert rou_series.rou_assets_to_revenue[1] == pytest.approx(1117514.5 / 9619278.0)
    assert rou_series.rou_assets_to_revenue[2] == pytest.approx(1340933.0 / 10588126.0)
    assert rou_series.rou_assets_to_revenue[3] == pytest.approx(1523218.5 / 11102600.0)
    assert dt_series.deferred_tax_assets == (6402.0, 9176.0, 17085.0, 24037.0)
    assert dt_series.deferred_tax_liabilities == (55084.0, 29522.0, 98188.0, 52278.0)
    assert dt_series.net_deferred_tax_position == (-48682.0, -20346.0, -81103.0, -28241.0)
    assert dt_series.deferred_tax_assets_change == (None, 2774.0, 7909.0, 6952.0)
    assert dt_series.deferred_tax_liabilities_change == (
        None,
        -25562.0,
        68666.0,
        -45910.0,
    )
    assert dt_series.net_deferred_tax_position_change == (
        None,
        28336.0,
        -60757.0,
        52862.0,
    )
    for period in EXPECTED_PERIODS:
        assert required_period_value(
            rou_stored, period, field="right_of_use_assets"
        ) == ROU_BALANCES[period]
        assert required_period_value(
            dta_stored, period, field="deferred_tax_assets"
        ) == DTA_BALANCES[period]
        assert required_period_value(
            dtl_stored, period, field="deferred_tax_liabilities"
        ) == DTL_BALANCES[period]
        assert (
            DTA_BALANCES[period] - DTL_BALANCES[period] == NET_DT_POSITIONS[period]
        )

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt_rou = resolve_lease_rou_source(restored)
    rt_dt = resolve_deferred_tax_sources(restored)
    assert rt_rou is not None and rt_rou.concept == "right_of_use_lease_asset"
    assert rt_rou.values == ROU_BALANCES
    assert rt_dt is not None
    assert rt_dt.deferred_tax_assets.concept == "deferred_tax_asset"
    assert rt_dt.deferred_tax_liabilities.concept == "deferred_tax_liability"
    assert rt_dt.deferred_tax_assets.values == DTA_BALANCES
    assert rt_dt.deferred_tax_liabilities.values == DTL_BALANCES

    with_builder = ReferenceModelBuilder(fin)
    with_keys = {s.semantic_key for s in with_builder.expected_specs}
    monkeypatch.setattr(
        lr,
        "_EXPLICIT_CONCEPT_ALIASES",
        {
            k: v
            for k, v in lr._EXPLICIT_CONCEPT_ALIASES.items()
            if k
            not in (
                "right_of_use_assets",
                "deferred_tax_assets",
                "deferred_tax_liabilities",
            )
        },
    )
    without_keys = {
        s.semantic_key for s in ReferenceModelBuilder(fin).expected_specs
    }
    monkeypatch.undo()
    assert len(without_keys) == PRIOR_LULULEMON_SPECS
    assert without_keys <= with_keys
    added = with_keys - without_keys
    assert len(with_keys) == LEASE_DT_LULULEMON_SPECS
    assert len(added) == 25
    assert {k.split(".")[0] for k in added} == {"lease_rou", "deferred_tax"}
    assert sum(1 for k in added if k.startswith("lease_rou.")) == 12
    assert sum(1 for k in added if k.startswith("deferred_tax.")) == 13

    out = tmp_path / "LululemonLeaseDT"
    trainer, answer = build_training_workbook(fin, out)
    assert trainer.exists() and answer.exists()
    smap = load_semantic_map(answer)
    module_comps = [
        c
        for c in smap.all_ordered()
        if c.semantic_key.startswith("lease_rou.")
        or c.semantic_key.startswith("deferred_tax.")
    ]
    assert len(smap.all_ordered()) == LEASE_DT_LULULEMON_SPECS
    assert len(module_comps) == 25
    assert {c.semantic_key for c in smap.all_ordered()} == with_keys

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        rou_src_f = str(
            next(
                ws.cell(r, 2).value
                for r in range(1, (ws.max_row or 1) + 1)
                if ws.cell(r, 1).value == "Right-of-use Assets"
            )
        ).replace(" ", "")
        dta_src_f = str(
            next(
                ws.cell(r, 2).value
                for r in range(1, (ws.max_row or 1) + 1)
                if ws.cell(r, 1).value == "Deferred Tax Assets"
            )
        ).replace(" ", "")
        dtl_src_f = str(
            next(
                ws.cell(r, 2).value
                for r in range(1, (ws.max_row or 1) + 1)
                if ws.cell(r, 1).value == "Deferred Tax Liabilities"
            )
        ).replace(" ", "")
        rou_row = SOURCE_START_ROW + rou_index
        dta_row = SOURCE_START_ROW + dta_index
        dtl_row = SOURCE_START_ROW + dtl_index
        assert f"'BalanceSheet'!B{rou_row}" in rou_src_f or (
            f"'Balance Sheet'!B{rou_row}" in rou_src_f
        )
        assert f"'BalanceSheet'!B{dta_row}" in dta_src_f or (
            f"'Balance Sheet'!B{dta_row}" in dta_src_f
        )
        assert f"'BalanceSheet'!B{dtl_row}" in dtl_src_f or (
            f"'Balance Sheet'!B{dtl_row}" in dtl_src_f
        )
        wb.close()

    twb = load_workbook(trainer, data_only=False)
    awb = load_workbook(answer, data_only=False)
    for comp in module_comps:
        row, col = parse_cell_ref(comp.cell)
        tcell = twb[comp.tab].cell(row=row, column=col)
        acell = awb[comp.tab].cell(row=row, column=col)
        assert tcell.value is None
        assert tcell.comment is None
        assert _fill_rgb(tcell) == "FFFF00"
        assert isinstance(acell.value, str) and acell.value.startswith("=")
        assert acell.comment is not None
        assert (acell.comment.text or "").strip()
        assert _fill_rgb(acell) == "FFFF00"
    twb.close()
    awb.close()

    assert _count_source_unavailable(trainer) == 74
    assert _count_source_unavailable(answer) == 74

    blank = check_workbook(trainer)
    assert (blank.correct, blank.incorrect, blank.blank, blank.total) == (
        0,
        0,
        LEASE_DT_LULULEMON_SPECS,
        LEASE_DT_LULULEMON_SPECS,
    )

    filled_dir = tmp_path / "filled_check"
    filled_dir.mkdir()
    filled_trainer = filled_dir / trainer.name
    shutil.copy2(trainer, filled_trainer)
    shutil.copy2(answer, filled_dir / answer.name)
    for sidecar in (
        answer.with_suffix(".component_map.json"),
        answer.with_suffix(".assumptions.json"),
        answer.with_suffix(".trainer.json"),
    ):
        if sidecar.is_file():
            shutil.copy2(sidecar, filled_dir / sidecar.name)
    wb = load_workbook(filled_trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(filled_trainer)
    wb.close()
    filled = check_workbook(filled_trainer)
    assert (filled.correct, filled.incorrect, filled.blank, filled.total) == (
        LEASE_DT_LULULEMON_SPECS,
        0,
        0,
        LEASE_DT_LULULEMON_SPECS,
    )

    bad = next(c for c in module_comps if c.family_id == "rou_assets_change")
    _inject_formula_and_cached_value(
        filled_trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(filled_trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Change in Right-of-use Assets" not in dumped

    blank_again = check_workbook(trainer)
    assert blank_again.blank == LEASE_DT_LULULEMON_SPECS
    assert rou_stored.concept == "right_of_use_lease_asset"
    assert dta_stored.concept == "deferred_tax_asset"
    assert dtl_stored.concept == "deferred_tax_liability"


def test_no_lulu_specific_production_branch():
    """Production engine must stay generic — no ticker/issuer hard-codes."""
    core_root = ROOT / "core"
    offenders: list[str] = []
    for path in core_root.rglob("*.py"):
        if "tests" in path.parts:
            continue
        text = path.read_text(encoding="utf-8")
        for token in ("LULU", "lululemon", "Lululemon"):
            if token in text:
                offenders.append(f"{path.relative_to(ROOT)}:{token}")
    assert offenders == []


def test_committed_reconciled_hashes_are_stable():
    """Lock measured baseline artifact digests for Step 9M.2.4.1."""
    assert _sha256(STD_JSON) == (
        "29852347d78387be6fd9224246ab337b15a5176218cd01b2c20a0c8c3c00b361"
    )
    assert _sha256(CONFLICTS_JSON) == (
        "d8a33012f6ea73126ac4e2ece3613e7011c11cb2b581745d8c3563e3c2e978e0"
    )
    # provenance is large; lock size + digest together
    assert PROV_JSON.stat().st_size == 699401
    assert _sha256(PROV_JSON) == (
        "a31f7b05cddc16a61df91cdc8578bb69713069651af21160ff662ea562433075"
    )

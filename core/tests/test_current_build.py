"""Current company workflow, publication safety, and public CLI contracts."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]


def test_public_namespace_and_compatibility():
    for namespace in ('bav', 'core'):
        result = subprocess.run([sys.executable, '-m', namespace, '--help'], cwd=ROOT, capture_output=True, text=True)
        assert result.returncode == 0, result.stderr
        for command in ('build', 'check', 'list', 'reconcile', 'validate-source', 'publish'):
            assert command in result.stdout


def test_public_help_is_bav_first_and_check_is_diagnostic():
    def _help(*args):
        result = subprocess.run(
            [sys.executable, '-m', 'bav', *args],
            cwd=ROOT,
            capture_output=True,
            text=True,
        )
        assert result.returncode == 0, result.stderr
        return ' '.join(result.stdout.split())

    top = _help('--help')
    assert 'BAV Excel Trainer' not in top
    assert 'BAV' in top
    assert 'build' in top
    assert 'check' in top
    assert 'publish' in top

    build = _help('build', '--help')
    assert 'professional BAV' in build
    assert 'Trainer generation' in build

    check = _help('check', '--help')
    assert 'Diagnose' in check
    assert 'BAV' in check
    assert 'without disclosing answers' in check
    assert 'Trainer' in check
    assert 'Validate every practice cell in a Trainer workbook' not in check

    listing = _help('list', '--help')
    assert 'analytical families' in listing
    assert 'BAV' in listing
    assert 'optional' in listing.lower()
    assert 'Trainer' in listing


def test_output_directory_uses_lowercase_slug(tmp_path, monkeypatch):
    from core import current_build
    from core.__main__ import main
    monkeypatch.setattr(current_build, 'OUTPUT_ROOT', tmp_path)
    legacy = tmp_path / 'Lululemon'
    legacy.mkdir()
    (legacy / 'stale.txt').write_text('stale', encoding='utf-8')
    assert main(['build', 'Lululemon']) == 0
    names = [p.name for p in tmp_path.iterdir()]
    assert names == ['lululemon']
    assert (tmp_path / 'lululemon' / 'Lululemon_BAV.xlsx').is_file()
    assert not (tmp_path / 'lululemon' / 'stale.txt').exists()


def test_aliases_and_missing_company(tmp_path):
    from core.current_build import resolve_company
    identities = [resolve_company(name) for name in ('Lululemon', 'lululemon', 'LULU')]
    assert identities[0] == identities[1] == identities[2]
    company = identities[0]
    assert company.output == ROOT / 'build/output/lululemon'
    assert company.input == ROOT / 'build/input/lululemon'
    assert company.bav.name == 'Lululemon_BAV.xlsx'
    assert company.trainer.name == 'Lululemon_BAV_Trainer.xlsx'
    assert company.answer.name == 'Lululemon_BAV.xlsx'
    with pytest.raises(ValueError, match='Unknown company.*Lululemon'):
        resolve_company('missing-company')


def snapshot(directory):
    return {str(p.relative_to(directory)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in directory.rglob('*') if p.is_file()}


def test_company_build_rebuild_failure_and_workbook_contract(tmp_path, monkeypatch, capsys):
    from core import current_build
    from core.__main__ import main
    from core.trainer.semantic_io import load_semantic_map
    from core.trainer.workbook import derive_trainer_workbook
    from core.engine.component_catalog import is_operating_kpi_source_identity
    monkeypatch.setattr(current_build, 'OUTPUT_ROOT', tmp_path)
    assert main(['build', 'Lululemon']) == 0
    company = current_build.resolve_company('LULU')
    assert company.output == tmp_path / 'lululemon'
    smap = load_semantic_map(company.bav)
    families = {c.family_id for c in smap.all_ordered()}
    assert 'geographic_operating_profit_growth_contribution' in families
    assert 'store_count_source' in families
    assert 'store_count_growth' in families
    assert 'revenue_per_store_period_end' in families
    assert any('comparable_sales' in f for f in families)
    assert any('square_foot' in f for f in families)
    assert any(f.startswith('revenue_driver_') for f in families)
    supporting = company.output / 'supporting'
    assert (supporting / 'build_status.json').is_file()
    assert (supporting / 'assumptions.json').is_file()
    assert (supporting / 'component_map.json').is_file()
    assert (supporting / 'rowmap.json').is_file()
    assert not (supporting / 'standardized.json').is_file()
    assert company.bav.is_file()
    assert not company.trainer.is_file()
    research = company.output / 'research'
    assert (research / 'Lululemon_Drivers.md').is_file()
    for name in ('Lululemon_Forecast.md', 'Lululemon_Valuation.md', 'Lululemon_Overview.md'):
        path = research / name
        assert path.is_file() and path.stat().st_size == 0
    for name in ('growth.png', 'geography.png', 'margin.png'):
        assert (company.output / 'figures' / 'drivers' / name).is_file()
    wb = load_workbook(company.bav)
    assert 'Build Status' in wb.sheetnames
    assert 'Overview' in wb.sheetnames
    assert 'Revenue Driver Analysis' in wb.sheetnames
    assert 'Trainer' not in wb.sheetnames
    rows = list(wb['Build Status'].values)
    assert any('Comparable Sales' in str(row) and 'Active / available' in str(row) for row in rows)
    rps_rows = [row for row in rows if 'Revenue per Store Analysis' in str(row)]
    assert len(rps_rows) == 1
    assert 'Active / available' in str(rps_rows[0]) and rps_rows[0][3] == 17
    assert any('Sales per Square Foot' in str(row) and 'Active / available' in str(row) for row in rows)
    driver_rows = [row for row in rows if 'Revenue Driver Analysis' in str(row)]
    assert len(driver_rows) == 1
    assert 'Active / available' in str(driver_rows[0]) and driver_rows[0][3] > 0
    opening = ' '.join(
        str(cell.value or '')
        for row in wb['Overview'].iter_rows()
        for cell in row
    )
    for term in ('Trainer', 'Answer Key', 'exercise', 'practice', 'Check'):
        assert term not in opening
    assert wb['Overview']['A1'].value == 'BAV'
    assert 'lululemon athletica inc.' in opening
    assert 'Historical coverage:' in opening
    assert 'Units:' in opening
    assert 'Historical reading' in opening
    assert 'Evidence limits' in opening
    assert 'untested' in opening.lower()
    assert 'do not establish causal drivers' in opening
    assert 'audit-only' in opening.lower()
    assert 'HISTORICAL REVENUE AND DISCLOSED STRATEGY' not in opening
    assert 'Management statement' not in opening
    assert 'Historical finding' not in opening
    assert 'Analyst inference' not in opening
    assert 'Schedules: ' not in opening
    links = {
        str(getattr(cell.hyperlink, 'target', '') or cell.hyperlink)
        for row in wb['Overview'].iter_rows()
        for cell in row
        if cell.hyperlink
    }
    assert any('Revenue Driver Analysis' in target for target in links)
    assert any('Build Status' in target for target in links)
    driver = ' '.join(
        str(cell.value or '')
        for row in wb['Revenue Driver Analysis'].iter_rows()
        for cell in row
    )
    assert 'Management statement' in driver
    assert 'exceeded revenue growth' in driver.lower()
    assert 'contributed negatively' in driver.lower()
    for comp in smap.all_ordered():
        cell = wb[comp.tab][comp.cell]
        if is_operating_kpi_source_identity(comp):
            assert cell.value is not None
        else:
            assert cell.value == comp.formula and cell.comment is not None
    assert not any(c.fill.fgColor.type == 'rgb' and c.fill.fgColor.rgb in ('FFFFFF00','00FFFF00') for ws in wb for row in ws for c in row)
    wb.close()
    assert main(['list', 'lululemon']) == 0
    before_bav = company.bav.read_bytes()
    derive_trainer_workbook(company.bav, company.trainer)
    assert company.bav.read_bytes() == before_bav
    wb = load_workbook(company.trainer)
    for comp in smap.all_ordered():
        cell = wb[comp.tab][comp.cell]
        if is_operating_kpi_source_identity(comp):
            assert cell.value is not None
        else:
            assert cell.value is None and cell.comment is None
            assert cell.fill.fgColor.rgb in ('00FFFF00', 'FFFFFF00')
    wb.close()
    assert main(['check', 'lulu']) == 0
    assert main(['build', 'LULU']) == 0
    assert sorted(p.name for p in tmp_path.iterdir()) == ['lululemon']
    assert sorted(p.name for p in company.output.glob('*.xlsx')) == ['Lululemon_BAV.xlsx']
    before = snapshot(company.output)
    def fail(*args, **kwargs):
        raise ValueError('deliberate staged validation failure')
    monkeypatch.setattr(current_build, 'verify_staged', fail)
    assert main(['build', 'Lululemon']) == 1
    assert snapshot(company.output) == before
    assert 'deliberate staged validation failure' in capsys.readouterr().err
    assert sorted(p.name for p in tmp_path.iterdir()) == ['lululemon']


def test_missing_build_errors(tmp_path, monkeypatch, capsys):
    from core import current_build
    from core.__main__ import main
    monkeypatch.setattr(current_build, 'OUTPUT_ROOT', tmp_path)
    for command in ('check','list'):
        assert main([command, 'LULU']) == 1
        assert 'python -m bav build Lululemon' in capsys.readouterr().err


def test_progressive_incomplete_parent(monkeypatch, tmp_path):
    from dataclasses import replace
    from core.engine import build_contract
    from core.current_build import resolve_company, prepare_company_input
    from core.engine.reference_model import ReferenceModelBuilder
    company = resolve_company('LULU')
    fin = prepare_company_input(company, tmp_path)
    monkeypatch.setattr(build_contract, 'BUILD_MODULES', tuple(
        replace(m, status='incomplete', complete_analysis=False) if m.id in ('geographic','operating_kpi') else m
        for m in build_contract.BUILD_MODULES))
    builder = ReferenceModelBuilder(fin, current_snapshot=True)
    families = {s.family_id for s in builder.expected_specs}
    assert 'geographic_revenue_share' in families
    assert 'store_count_growth' in families
    assert any('comparable_sales' in f for f in families)


def test_atomic_exchange_failure_preserves_current(tmp_path, monkeypatch):
    from core import current_build
    from core.__main__ import main
    monkeypatch.setattr(current_build, 'OUTPUT_ROOT', tmp_path)
    assert main(['build', 'LULU']) == 0
    current = current_build.resolve_company('LULU').output
    before = snapshot(current)
    def fail(*args):
        raise OSError('forced exchange failure')
    monkeypatch.setattr(current_build, '_exchange_directories', fail)
    assert main(['build', 'LULU']) == 1
    assert snapshot(current) == before


def test_source_failure_does_not_fallback_or_replace(tmp_path, monkeypatch):
    from core import current_build
    from core.__main__ import main
    monkeypatch.setattr(current_build, 'OUTPUT_ROOT', tmp_path)
    assert main(['build', 'LULU']) == 0
    current = current_build.resolve_company('LULU').output
    before = snapshot(current)
    def fail(*args, **kwargs):
        raise ValueError('source hash mismatch')
    monkeypatch.setattr(current_build, 'standardized_from_payload', fail)
    assert main(['build', 'LULU']) == 1
    assert snapshot(current) == before


def test_ambiguous_aliases_fail_closed(monkeypatch):
    from core import current_build
    monkeypatch.setattr(current_build, 'PROJECTS', current_build.PROJECTS + (
        ('Other', 'other', ('LULU',), (), None),))
    with pytest.raises(ValueError, match='Ambiguous.*Lululemon, Other'):
        current_build.resolve_company('lulu')


def test_current_status_uses_emitted_families():
    from core.build_status import status_rows, ACTIVE, UNAVAILABLE, INACTIVE
    from types import SimpleNamespace
    class Map:
        def all_ordered(self):
            return [SimpleNamespace(family_id='store_count_source', tab='Store Count Analysis')]
    rows = {r['family']: r for r in status_rows(Map())}
    assert rows['Store-count history']['status'] == ACTIVE
    assert rows['Store-count growth']['status'] == UNAVAILABLE
    assert rows['Comparable Sales Analysis']['status'] == UNAVAILABLE
    assert rows['Revenue per Store Analysis']['status'] == UNAVAILABLE


def test_superseded_legacy_names_removed_only_after_success(tmp_path, monkeypatch):
    from core import current_build
    from core.__main__ import main
    monkeypatch.setattr(current_build, 'OUTPUT_ROOT', tmp_path)
    old = tmp_path / 'Lululemon_Live_Trainer.xlsx'
    old.write_bytes(b'old workbook')
    other = tmp_path / 'Other_Trainer.xlsx'
    other.write_bytes(b'other workbook')
    assert main(['build', 'LULU']) == 0
    assert not old.exists()
    assert other.read_bytes() == b'other workbook'


def test_dotted_ticker_uses_company_resolver(monkeypatch, capsys):
    from core.__main__ import main
    from core import current_build
    def stop(company, assumptions):
        assert company.name == 'FastRetailing'
        raise ValueError('resolved dotted ticker')
    monkeypatch.setattr(current_build, 'build_company', stop)
    assert main(['build', '9983.T']) == 1
    assert 'resolved dotted ticker' in capsys.readouterr().err


def test_build_preserves_protected_evidence(tmp_path, monkeypatch):
    from core.__main__ import main
    from core import current_build
    protected = ['TARGET.md', 'IMPLEMENTATION.md', 'RESULT.md',
                 '.git/autocycle/resume-state', '.git/autocycle/work-state.json',
                 '.git/autocycle/implementation-baseline.json', '.git/autocycle/latest-implementation']
    files = [ROOT / p for p in protected if (ROOT / p).is_file()]
    files += [p for root in ('benchmark', 'release') for p in (ROOT / root).rglob('*') if p.is_file()]
    before = {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    monkeypatch.setattr(current_build, 'OUTPUT_ROOT', tmp_path)
    assert main(['build', 'LULU']) == 0
    assert {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in files} == before


def test_disabled_module_is_not_reported_as_missing_source(monkeypatch):
    from dataclasses import replace
    from core.engine import build_contract
    from core.build_status import status_rows, INACTIVE
    from core.engine.semantic_map import SemanticMap
    monkeypatch.setattr(build_contract, 'BUILD_MODULES', tuple(
        replace(m, status='deferred', current_ready=False) if m.id == 'operating_kpi' else m
        for m in build_contract.BUILD_MODULES))
    rows = {r['family']: r for r in status_rows(SemanticMap())}
    assert rows['Store-count history']['status'] == INACTIVE

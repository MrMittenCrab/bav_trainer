"""Step 9M.2.4.1.1.1.28 — inventory balance vs reported operating CF adjustment."""

from __future__ import annotations

import copy
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    INVENTORY_ANALYSIS_COMPONENT_CATALOG,
    expand_inventory_analysis_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.historical_expected import inventory_analysis_expected_series
from core.model.inventory_analysis import (
    CHANGE_IN_INVENTORIES_CONCEPT,
    INVENTORIES_CONCEPT,
    REVENUE_CONCEPT,
    compute_inventory_analysis_series,
    inventory_analysis_applicable,
    inventory_analysis_availability,
    inventory_balance_implied_cf_adjustment_applicable,
    inventory_cf_adjustment_difference_applicable,
    inventory_change_applicable,
    inventory_intensity_applicable,
    resolve_inventory_analysis_sources,
)
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_capex import P1, P2, _dupont_row_by_label, _tiny
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "benchmark" / "fast_retailing" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _vals(a, b, *, missing_period=False, none_period=False):
    if missing_period:
        return {P1: a}
    if none_period:
        return {P1: a, P2: None}
    return {P1: a, P2: b}


def _add_inventory(
    fin: StandardizedFinancials,
    *,
    inventories=(100.0, 140.0),
    revenue=(1000.0, 1100.0),
    cf_adjustment=(-40.0, -35.0),
    inventory_concept=INVENTORIES_CONCEPT,
    revenue_concept=REVENUE_CONCEPT,
    cf_concept=CHANGE_IN_INVENTORIES_CONCEPT,
    include_inventory=True,
    include_revenue=True,
    include_cf=False,
    duplicate=None,
    missing_period=None,
    none_period=None,
    on_income_statement=False,
    on_balance_sheet_cf=False,
):
    stored = {}
    if include_revenue:
        existing = next(
            (
                row
                for row in fin.income_statement
                if row.concept == "revenue" or row.label.lower() == "revenue"
            ),
            None,
        )
        if existing is not None:
            fin.income_statement.remove(existing)
        rev_item = _li(
            "Revenue",
            _vals(
                revenue[0],
                revenue[1],
                missing_period=missing_period == "revenue",
                none_period=none_period == "revenue",
            ),
            concept=revenue_concept,
        )
        fin.income_statement.insert(0, rev_item)
        stored["revenue"] = rev_item
        if duplicate == "revenue":
            fin.income_statement.append(
                _li(
                    "Revenue duplicate",
                    _vals(revenue[0], revenue[1]),
                    concept=revenue_concept,
                )
            )
    if include_inventory:
        inv_item = _li(
            "Inventories",
            _vals(
                inventories[0],
                inventories[1],
                missing_period=missing_period == "inventories",
                none_period=none_period == "inventories",
            ),
            concept=inventory_concept,
        )
        target = (
            fin.income_statement if on_income_statement else fin.balance_sheet
        )
        target.append(inv_item)
        stored["inventories"] = inv_item
        if not on_income_statement:
            equity = next(
                row
                for row in fin.balance_sheet
                if "equity" in row.label.lower()
            )
            for period, amount in inv_item.values.items():
                if amount is None or period not in equity.values:
                    continue
                equity.values[period] = float(equity.values[period]) + float(amount)
        if duplicate == "inventories" and not on_income_statement:
            dup_item = _li(
                "Inventories duplicate",
                _vals(inventories[0], inventories[1]),
                concept=inventory_concept,
            )
            target.append(dup_item)
            equity = next(
                row
                for row in fin.balance_sheet
                if "equity" in row.label.lower()
            )
            for period, amount in dup_item.values.items():
                equity.values[period] = float(equity.values[period]) + float(amount)
    if include_cf:
        cf_item = _li(
            "(Increase)/decrease in inventories",
            _vals(
                cf_adjustment[0],
                cf_adjustment[1],
                missing_period=missing_period == "change_in_inventories",
                none_period=none_period == "change_in_inventories",
            ),
            concept=cf_concept,
        )
        cf_target = fin.balance_sheet if on_balance_sheet_cf else fin.cash_flow
        cf_target.append(cf_item)
        stored["change_in_inventories"] = cf_item
        if duplicate == "change_in_inventories" and not on_balance_sheet_cf:
            cf_target.append(
                _li(
                    "Inventories CF duplicate",
                    _vals(cf_adjustment[0], cf_adjustment[1]),
                    concept=cf_concept,
                )
            )
    return stored


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(INVENTORY_ANALYSIS_COMPONENT_CATALOG) == 7
    assert [f.order for f in INVENTORY_ANALYSIS_COMPONENT_CATALOG] == list(
        range(145, 152)
    )
    intensity = expand_inventory_analysis_specs(
        periods,
        start_order=1000,
        include_intensity=True,
    )
    assert len(intensity) == 5
    assert {s.family_id for s in intensity} == {"inventory_intensity"}
    assert all(s.semantic_key.startswith("inventory_analysis.") for s in intensity)
    full = expand_inventory_analysis_specs(
        periods,
        start_order=1000,
        include_intensity=True,
        include_change=True,
        include_revenue_scale=True,
        include_intensity_effect=True,
        include_reconstructed=True,
        include_balance_implied=True,
        include_cf_difference=True,
    )
    assert len(full) == 29
    assert {s.family_id for s in full} == {
        "inventory_intensity",
        "inventory_change",
        "inventory_revenue_scale_effect",
        "inventory_intensity_effect",
        "reconstructed_inventory_change",
        "inventory_balance_implied_cf_adjustment",
        "inventory_cf_adjustment_difference",
    }
    assert not any(
        s.family_id != "inventory_intensity" and s.period_index == 0 for s in full
    )
    with pytest.raises(ValueError, match="revenue-scale requires"):
        expand_inventory_analysis_specs(
            periods, start_order=1, include_revenue_scale=True
        )
    with pytest.raises(ValueError, match="reconstructed change requires"):
        expand_inventory_analysis_specs(
            periods,
            start_order=1,
            include_intensity=True,
            include_revenue_scale=True,
            include_reconstructed=True,
        )
    with pytest.raises(ValueError, match="balance-implied requires"):
        expand_inventory_analysis_specs(
            periods, start_order=1, include_balance_implied=True
        )
    with pytest.raises(ValueError, match="CF difference requires"):
        expand_inventory_analysis_specs(
            periods,
            start_order=1,
            include_change=True,
            include_cf_difference=True,
        )
    with pytest.raises(ValueError, match="duplicate"):
        expand_inventory_analysis_specs(
            [periods[0], periods[0]], start_order=1, include_change=True
        )
    with pytest.raises(ValueError, match="chronological"):
        expand_inventory_analysis_specs(
            list(reversed(periods)), start_order=1, include_change=True
        )


def test_unique_concept_resolution_and_renamed_label():
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_inventory(fin)
    sources = resolve_inventory_analysis_sources(fin)
    assert sources.inventories is stored["inventories"]
    assert sources.revenue is stored["revenue"]
    assert inventory_analysis_applicable(fin)
    assert inventory_change_applicable(fin)
    assert inventory_intensity_applicable(fin)
    assert inventory_balance_implied_cf_adjustment_applicable(fin)
    assert not inventory_cf_adjustment_difference_applicable(fin)
    avail = inventory_analysis_availability(fin)
    assert avail.inventories is True
    assert avail.inventories_ambiguous is False
    assert avail.change_in_inventories is False
    assert avail.change_in_inventories_ambiguous is False

    with_cf = _tiny(with_cfo=False, with_payments=False)
    stored_cf = _add_inventory(with_cf, include_cf=True)
    sources_cf = resolve_inventory_analysis_sources(with_cf)
    assert sources_cf.change_in_inventories is stored_cf["change_in_inventories"]
    assert inventory_cf_adjustment_difference_applicable(with_cf)
    renamed_cf = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(renamed_cf, include_cf=True)
    renamed_cf.cash_flow[-1].label = "Inventories"
    assert (
        resolve_inventory_analysis_sources(renamed_cf).change_in_inventories.concept
        == CHANGE_IN_INVENTORIES_CONCEPT
    )

    renamed = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(renamed)
    renamed.balance_sheet[-1].label = "Merchandise inventories"
    assert (
        resolve_inventory_analysis_sources(renamed).inventories.concept
        == INVENTORIES_CONCEPT
    )


def test_absent_label_only_duplicate_wrong_statement():
    absent = _tiny(with_cfo=False, with_payments=False)
    assert not inventory_analysis_applicable(absent)
    assert resolve_inventory_analysis_sources(absent).inventories is None
    with pytest.raises(MissingLineError):
        compute_inventory_analysis_series(
            absent, list(canonical_fiscal_periods(absent))
        )
    assert ReferenceModelBuilder(absent).inventory_analysis_specs == ()

    label_only = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(label_only, inventory_concept="")
    assert not inventory_analysis_applicable(label_only)
    assert resolve_inventory_analysis_sources(label_only).inventories is None
    assert (
        resolve_line(
            label_only.balance_sheet, INVENTORIES_CONCEPT, required=False
        ).item
        is None
    )
    assert ReferenceModelBuilder(label_only).inventory_analysis_specs == ()

    dup = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(dup, duplicate="inventories")
    avail = inventory_analysis_availability(dup)
    assert avail.inventories_ambiguous is True
    assert avail.inventories is False
    assert not inventory_change_applicable(dup)
    with pytest.raises(AmbiguousLineError):
        resolve_line(dup.balance_sheet, INVENTORIES_CONCEPT, required=False)
    assert ReferenceModelBuilder(dup).inventory_analysis_specs == ()

    wrong = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(wrong, on_income_statement=True)
    assert not inventory_analysis_applicable(wrong)
    assert (
        resolve_line(
            wrong.income_statement, INVENTORIES_CONCEPT, required=False
        ).item
        is not None
    )
    assert ReferenceModelBuilder(wrong).inventory_analysis_specs == ()

    label_only_cf = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(label_only_cf, include_cf=True, cf_concept="")
    assert inventory_balance_implied_cf_adjustment_applicable(label_only_cf)
    assert not inventory_cf_adjustment_difference_applicable(label_only_cf)
    assert resolve_inventory_analysis_sources(label_only_cf).change_in_inventories is None
    assert (
        resolve_line(
            label_only_cf.cash_flow, CHANGE_IN_INVENTORIES_CONCEPT, required=False
        ).item
        is None
    )

    dup_cf = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(dup_cf, include_cf=True, duplicate="change_in_inventories")
    avail_cf = inventory_analysis_availability(dup_cf)
    assert avail_cf.change_in_inventories_ambiguous is True
    assert avail_cf.change_in_inventories is False
    assert inventory_balance_implied_cf_adjustment_applicable(dup_cf)
    assert not inventory_cf_adjustment_difference_applicable(dup_cf)
    with pytest.raises(AmbiguousLineError):
        resolve_line(dup_cf.cash_flow, CHANGE_IN_INVENTORIES_CONCEPT, required=False)

    wrong_cf = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(wrong_cf, include_cf=True, on_balance_sheet_cf=True)
    assert inventory_balance_implied_cf_adjustment_applicable(wrong_cf)
    assert not inventory_cf_adjustment_difference_applicable(wrong_cf)
    assert (
        resolve_line(
            wrong_cf.balance_sheet, CHANGE_IN_INVENTORIES_CONCEPT, required=False
        ).item
        is not None
    )
    assert resolve_inventory_analysis_sources(wrong_cf).change_in_inventories is None


def test_independent_family_omissions():
    no_rev = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(no_rev, include_revenue=False)
    no_rev.income_statement = [
        row
        for row in no_rev.income_statement
        if row.concept != "revenue" and row.label.lower() != "revenue"
    ]
    assert inventory_change_applicable(no_rev)
    assert not inventory_intensity_applicable(no_rev)
    assert inventory_balance_implied_cf_adjustment_applicable(no_rev)
    assert not inventory_cf_adjustment_difference_applicable(no_rev)
    series = compute_inventory_analysis_series(
        no_rev, list(canonical_fiscal_periods(no_rev))
    )
    assert series.inventory_change[1] == pytest.approx(40.0)
    assert series.inventory_balance_implied_cf_adjustment[1] == pytest.approx(-40.0)
    assert series.inventory_intensity is None
    assert series.inventory_revenue_scale_effect is None
    assert series.inventory_cf_adjustment_difference is None
    # Core historical builder requires unique revenue, so change-only
    # gating is asserted on the series/applicability surface rather than
    # a completed workbook.

    no_cf = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(no_cf)
    assert inventory_balance_implied_cf_adjustment_applicable(no_cf)
    assert not inventory_cf_adjustment_difference_applicable(no_cf)
    series_no_cf = compute_inventory_analysis_series(
        no_cf, list(canonical_fiscal_periods(no_cf))
    )
    assert series_no_cf.inventory_balance_implied_cf_adjustment[1] == pytest.approx(-40.0)
    assert series_no_cf.inventory_cf_adjustment_difference is None
    no_cf_ids = {s.family_id for s in ReferenceModelBuilder(no_cf).inventory_analysis_specs}
    assert "inventory_balance_implied_cf_adjustment" in no_cf_ids
    assert "inventory_cf_adjustment_difference" not in no_cf_ids

    no_cfo = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(no_cfo, include_cf=True)
    assert inventory_cf_adjustment_difference_applicable(no_cfo)
    series_no_cfo = compute_inventory_analysis_series(
        no_cfo, list(canonical_fiscal_periods(no_cfo))
    )
    assert series_no_cfo.inventory_cf_adjustment_difference[1] == pytest.approx(5.0)

    no_inv = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(no_inv, include_inventory=False, include_cf=True)
    assert not inventory_analysis_applicable(no_inv)
    assert not inventory_cf_adjustment_difference_applicable(no_inv)
    with pytest.raises(MissingLineError):
        compute_inventory_analysis_series(
            no_inv, list(canonical_fiscal_periods(no_inv))
        )

    amb_rev = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(amb_rev, duplicate="revenue", include_cf=True)
    avail = inventory_analysis_availability(amb_rev)
    assert avail.revenue_ambiguous is True
    assert inventory_change_applicable(amb_rev)
    assert not inventory_intensity_applicable(amb_rev)
    assert inventory_cf_adjustment_difference_applicable(amb_rev)
    series_amb = compute_inventory_analysis_series(
        amb_rev, list(canonical_fiscal_periods(amb_rev))
    )
    assert series_amb.inventory_change[1] == pytest.approx(40.0)
    assert series_amb.inventory_intensity is None
    assert series_amb.inventory_balance_implied_cf_adjustment[1] == pytest.approx(-40.0)
    assert series_amb.inventory_cf_adjustment_difference[1] == pytest.approx(5.0)
    mapped = inventory_analysis_expected_series(series_amb)
    assert set(mapped) == {
        "inventory_change",
        "inventory_balance_implied_cf_adjustment",
        "inventory_cf_adjustment_difference",
    }


def test_zeros_undefined_negative_flat_and_missing_values():
    fin = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(
        fin,
        inventories=(0.0, 50.0),
        revenue=(1000.0, 2000.0),
        include_cf=True,
        cf_adjustment=(0.0, -50.0),
    )
    series = compute_inventory_analysis_series(fin, [P1, P2])
    assert series.inventory_intensity == (0.0, pytest.approx(50.0 / 2000.0))
    assert series.inventory_change[1] == pytest.approx(50.0)
    assert series.inventory_balance_implied_cf_adjustment[1] == pytest.approx(-50.0)
    assert series.change_in_inventories[0] == pytest.approx(0.0)
    assert series.inventory_cf_adjustment_difference[1] == pytest.approx(0.0)

    decreasing = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(
        decreasing,
        inventories=(200.0, 150.0),
        revenue=(1000.0, 800.0),
        include_cf=True,
        cf_adjustment=(10.0, 40.0),
    )
    series_d = compute_inventory_analysis_series(decreasing, [P1, P2])
    assert series_d.inventory_change[1] == pytest.approx(-50.0)
    assert series_d.inventory_balance_implied_cf_adjustment[1] == pytest.approx(50.0)
    assert series_d.inventory_cf_adjustment_difference[1] == pytest.approx(-10.0)
    i0 = 200.0 / 1000.0
    i1 = 150.0 / 800.0
    scale = i0 * (800.0 - 1000.0)
    intensity_effect = 800.0 * (i1 - i0)
    assert series_d.inventory_revenue_scale_effect[1] == pytest.approx(scale)
    assert series_d.inventory_intensity_effect[1] == pytest.approx(intensity_effect)
    assert series_d.reconstructed_inventory_change[1] == pytest.approx(-50.0)

    flat = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(
        flat,
        inventories=(100.0, 100.0),
        include_cf=True,
        cf_adjustment=(-5.0, 12.0),
    )
    series_f = compute_inventory_analysis_series(flat, [P1, P2])
    assert series_f.inventory_change[1] == pytest.approx(0.0)
    assert series_f.inventory_balance_implied_cf_adjustment[1] == pytest.approx(0.0)
    assert series_f.inventory_cf_adjustment_difference[1] == pytest.approx(12.0)

    zero_rev = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(zero_rev, revenue=(0.0, 2000.0), include_cf=True)
    series_z = compute_inventory_analysis_series(zero_rev, [P1, P2])
    assert series_z.inventory_intensity[0] == UNDEFINED_RATIO
    assert series_z.inventory_change[1] == pytest.approx(40.0)
    assert series_z.inventory_balance_implied_cf_adjustment[1] == pytest.approx(-40.0)
    assert series_z.inventory_revenue_scale_effect[1] == UNDEFINED_RATIO
    assert series_z.inventory_intensity_effect[1] == UNDEFINED_RATIO
    assert series_z.reconstructed_inventory_change[1] == UNDEFINED_RATIO
    assert series_z.inventory_intensity[1] == pytest.approx(140.0 / 2000.0)
    assert series_z.inventory_cf_adjustment_difference[1] == pytest.approx(5.0)

    zero_rev_current = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(zero_rev_current, revenue=(1000.0, 0.0))
    series_zc = compute_inventory_analysis_series(zero_rev_current, [P1, P2])
    assert series_zc.inventory_intensity[1] == UNDEFINED_RATIO
    assert series_zc.inventory_change[1] == pytest.approx(40.0)
    assert series_zc.inventory_balance_implied_cf_adjustment[1] == pytest.approx(-40.0)
    assert series_zc.inventory_revenue_scale_effect[1] == pytest.approx(
        (100.0 / 1000.0) * (0.0 - 1000.0)
    )
    assert series_zc.inventory_intensity_effect[1] == UNDEFINED_RATIO
    assert series_zc.reconstructed_inventory_change[1] == UNDEFINED_RATIO
    assert series_zc.inventory_cf_adjustment_difference is None

    mapped = inventory_analysis_expected_series(series)
    assert set(mapped) == {f.id for f in INVENTORY_ANALYSIS_COMPONENT_CATALOG}

    missing = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(missing, missing_period="inventories")
    with pytest.raises(MissingHistoricalValueError):
        compute_inventory_analysis_series(missing, [P1, P2])
    none_period = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(none_period, none_period="revenue")
    with pytest.raises(MissingHistoricalValueError):
        compute_inventory_analysis_series(none_period, [P1, P2])
    missing_current_cf = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(
        missing_current_cf, include_cf=True, missing_period="change_in_inventories"
    )
    with pytest.raises(MissingHistoricalValueError):
        compute_inventory_analysis_series(missing_current_cf, [P1, P2])
    none_cf = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(none_cf, include_cf=True, none_period="change_in_inventories")
    with pytest.raises(MissingHistoricalValueError):
        compute_inventory_analysis_series(none_cf, [P1, P2])
    opening_missing_cf = _tiny(with_cfo=False, with_payments=False)
    stored_open = _add_inventory(opening_missing_cf, include_cf=True)
    stored_open["change_in_inventories"].values.pop(P1)
    series_open = compute_inventory_analysis_series(opening_missing_cf, [P1, P2])
    assert series_open.change_in_inventories[0] is None
    assert series_open.inventory_cf_adjustment_difference[1] == pytest.approx(5.0)


def test_reconstructed_change_reconciles_and_immutability():
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_inventory(fin)
    series = compute_inventory_analysis_series(fin, [P1, P2])
    i0 = 100.0 / 1000.0
    i1 = 140.0 / 1100.0
    scale = i0 * (1100.0 - 1000.0)
    intensity_effect = 1100.0 * (i1 - i0)
    assert series.inventory_intensity[0] == pytest.approx(i0)
    assert series.inventory_intensity[1] == pytest.approx(i1)
    assert series.inventory_change[1] == pytest.approx(40.0)
    assert series.inventory_balance_implied_cf_adjustment[1] == pytest.approx(-40.0)
    assert series.inventory_revenue_scale_effect[1] == pytest.approx(scale)
    assert series.inventory_intensity_effect[1] == pytest.approx(intensity_effect)
    reconstructed = series.reconstructed_inventory_change[1]
    assert reconstructed == pytest.approx(scale + intensity_effect)
    assert reconstructed == pytest.approx(40.0)
    assert series.inventory_change[0] is None
    assert series.inventory_balance_implied_cf_adjustment[0] is None
    assert series.reconstructed_inventory_change[0] is None
    assert series.inventory_cf_adjustment_difference is None

    with_cf = _tiny(with_cfo=False, with_payments=False)
    stored_cf = _add_inventory(with_cf, include_cf=True, cf_adjustment=(-12.0, -30.0))
    series_cf = compute_inventory_analysis_series(with_cf, [P1, P2])
    implied = series_cf.inventory_balance_implied_cf_adjustment[1]
    difference = series_cf.inventory_cf_adjustment_difference[1]
    reported = series_cf.change_in_inventories[1]
    assert implied == pytest.approx(-40.0)
    assert reported == pytest.approx(-30.0)
    assert difference == pytest.approx(10.0)
    assert implied + difference == pytest.approx(reported, abs=1e-8)
    assert series_cf.inventory_cf_adjustment_difference[0] is None
    assert series_cf.change_in_inventories[0] == pytest.approx(-12.0)

    reverse = [P2, P1]
    series_rev = compute_inventory_analysis_series(fin, reverse)
    assert series_rev.inventory_intensity[0] == pytest.approx(i1)
    assert series_rev.inventory_change[1] == pytest.approx(-40.0)
    assert series_rev.inventory_balance_implied_cf_adjustment[1] == pytest.approx(40.0)

    before = copy.deepcopy(stored["inventories"].values)
    compute_inventory_analysis_series(fin, [P1, P2])
    assert stored["inventories"].values == before
    assert stored["inventories"].concept == INVENTORIES_CONCEPT
    assert stored["revenue"].concept == REVENUE_CONCEPT
    before_cf = copy.deepcopy(stored_cf["change_in_inventories"].values)
    compute_inventory_analysis_series(with_cf, [P1, P2])
    assert stored_cf["change_in_inventories"].values == before_cf
    assert stored_cf["change_in_inventories"].concept == CHANGE_IN_INVENTORIES_CONCEPT


def test_workbook_gating_structure_formulas_notes_and_check(tmp_path):
    data = _tiny(with_cfo=False, with_payments=False)
    _add_inventory(data, include_cf=True)
    trainer, answer = build_training_workbook(data, tmp_path / "INV_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.inventory_analysis_series is not None
    assert len(builder.inventory_analysis_specs) == 8
    assert {s.family_id for s in builder.inventory_analysis_specs} == {
        "inventory_intensity",
        "inventory_change",
        "inventory_revenue_scale_effect",
        "inventory_intensity_effect",
        "reconstructed_inventory_change",
        "inventory_balance_implied_cf_adjustment",
        "inventory_cf_adjustment_difference",
    }

    smap = load_semantic_map(answer)
    inv_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in INVENTORY_ANALYSIS_COMPONENT_CATALOG}
    ]
    assert len(inv_comps) == 8
    assert not any(
        c.period_index == 0 and c.family_id != "inventory_intensity" for c in inv_comps
    )

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "INVENTORY GROWTH AND INTENSITY CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        inv_row = _dupont_row_by_label(ws, "Inventory (reported)")
        intensity_row = _dupont_row_by_label(ws, "Inventory intensity")
        change_row = _dupont_row_by_label(ws, "Change in inventory")
        scale_row = _dupont_row_by_label(ws, "Revenue-scale effect on inventory")
        recon_row = _dupont_row_by_label(ws, "Reconstructed change in inventory")
        cf_row = _dupont_row_by_label(ws, "Inventory CF adjustment (reported)")
        implied_row = _dupont_row_by_label(
            ws, "Balance-implied inventory CF adjustment"
        )
        diff_row = _dupont_row_by_label(ws, "Unexplained inventory CF difference")
        assert isinstance(ws.cell(inv_row, 2).value, str)
        assert ws.cell(inv_row, 2).value.startswith("=")
        assert "Balance Sheet" in str(ws.cell(inv_row, 2).value)
        assert isinstance(ws.cell(cf_row, 2).value, str)
        assert ws.cell(cf_row, 2).value.startswith("=")
        assert "Cash Flow" in str(ws.cell(cf_row, 2).value)
        assert ws.cell(change_row, 2).value in (None, "")
        assert ws.cell(implied_row, 2).value in (None, "")
        assert ws.cell(diff_row, 2).value in (None, "")
        for col in (2, 3):
            intensity_cell = ws.cell(intensity_row, column=col)
            scale_cell = ws.cell(scale_row, column=col)
            recon_cell = ws.cell(recon_row, column=col)
            implied_cell = ws.cell(implied_row, column=col)
            diff_cell = ws.cell(diff_row, column=col)
            if path == answer:
                assert isinstance(intensity_cell.value, str) and intensity_cell.value.startswith("=")
                assert "NA()" in str(intensity_cell.value)
                note = (intensity_cell.comment.text or "") if intensity_cell.comment else ""
                assert "annual revenue" in note.lower()
                assert "days" in note.lower() or "turnover" in note.lower()
                if col == 3:
                    scale_note = (
                        (scale_cell.comment.text or "") if scale_cell.comment else ""
                    )
                    assert "prior" in scale_note.lower()
                    recon_note = (
                        (recon_cell.comment.text or "") if recon_cell.comment else ""
                    )
                    assert "arithmetic" in recon_note.lower() or "decomposition" in recon_note.lower()
                    assert "cash" in recon_note.lower()
                    implied_note = (
                        (implied_cell.comment.text or "") if implied_cell.comment else ""
                    )
                    assert "negative" in implied_note.lower() or "−" in implied_note or "-" in implied_note
                    assert "reported operating cash-flow" in implied_note.lower() or "operating cash-flow reconciliation" in implied_note.lower()
                    diff_note = (
                        (diff_cell.comment.text or "") if diff_cell.comment else ""
                    )
                    assert "further evidence" in diff_note.lower()
                    assert "balancing plug" in diff_note.lower()
                    assert "write-down" in diff_note.lower() or "write-downs" in diff_note.lower()
                    assert implied_cell.value == f"=-{chr(64+col)}{change_row}" or str(
                        implied_cell.value
                    ).replace(" ", "") == f"=-{chr(64+col)}{change_row}"
                    assert str(diff_cell.value).replace(" ", "") == (
                        f"={chr(64+col)}{cf_row}-{chr(64+col)}{implied_row}"
                    )
            else:
                assert intensity_cell.value is None
                assert intensity_cell.comment is None
                assert scale_cell.comment is None
                assert implied_cell.comment is None
                assert diff_cell.comment is None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in inv_comps if c.family_id == "inventory_intensity")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Inventory intensity" not in dumped


def test_source_row_fidelity_round_trip_and_immutability(tmp_path):
    fin = _tiny(with_cfo=False, with_payments=False)
    stored = _add_inventory(fin, include_cf=True)
    original_index = fin.balance_sheet.index(stored["inventories"])
    original_cf_index = fin.cash_flow.index(stored["change_in_inventories"])
    trainer, answer = build_training_workbook(fin, tmp_path / "INV_FID.xlsx")
    smap = load_semantic_map(answer)
    awb = load_workbook(answer, data_only=False)
    ws = awb["ALT DuPont"]
    inv_row = _dupont_row_by_label(ws, "Inventory (reported)")
    src_f = str(ws.cell(inv_row, 2).value).replace(" ", "")
    assert src_f.startswith("='BalanceSheet'!") or src_f.startswith(
        "='Balance Sheet'!"
    )
    assert str(original_index + 7) in src_f
    cf_row = _dupont_row_by_label(ws, "Inventory CF adjustment (reported)")
    cf_f = str(ws.cell(cf_row, 2).value).replace(" ", "")
    assert "CashFlow" in cf_f or "Cash Flow" in str(ws.cell(cf_row, 2).value)
    assert str(original_cf_index + 7) in cf_f
    awb.close()
    restored = standardized_from_payload(standardized_to_payload(fin))
    rt = resolve_inventory_analysis_sources(restored)
    assert rt.inventories is not None
    assert rt.inventories.concept == INVENTORIES_CONCEPT
    assert rt.change_in_inventories is not None
    assert rt.change_in_inventories.concept == CHANGE_IN_INVENTORIES_CONCEPT
    assert stored["inventories"].values == {P1: 100.0, P2: 140.0}
    assert stored["change_in_inventories"].values == {P1: -40.0, P2: -35.0}
    assert len(smap.all_ordered()) == len(ReferenceModelBuilder(fin).expected_specs)


def test_demo_omits_inventory_analysis_fast_retailing_activates(tmp_path):
    demo = HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )
    assert not inventory_analysis_applicable(demo)
    builder = ReferenceModelBuilder(demo)
    assert builder.inventory_analysis_series is None
    assert builder.inventory_analysis_specs == ()
    trainer, answer = build_training_workbook(demo, tmp_path / "INV_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312

    payload = json.loads(FR_JSON.read_text(encoding="utf-8"))
    fr = standardized_from_payload(payload)
    assert inventory_analysis_applicable(fr)
    fr_builder = ReferenceModelBuilder(fr)
    assert len(fr_builder.inventory_analysis_specs) == 29
    assert len(fr_builder.expected_specs) == 577
    assert "inventory_cf_adjustment_difference" in {
        s.family_id for s in fr_builder.expected_specs
    }

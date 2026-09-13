"""Step 9G.1 — cross-company historical robustness matrix."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import DocumentManifest, DocumentType
from core.data.line_identity import line_identity, validate_financials_identities
from core.data.standardized_io import standardized_to_payload
from core.engine.reference_model import (
    DEFERRED_PLACEHOLDER,
    DEFERRED_TAB_NAMES,
    EARNINGS_QUALITY_SHEET,
    PER_SHARE_SHEET,
    WORKING_CAPITAL_SHEET,
    ReferenceModelBuilder,
)
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.ingestion.reconciler import reconcile_financials
from core.model.financial_math import compute_anchor
from core.model.historical_expected import expected_value_for_component
from core.model.normalization import compute_normalization_series, normalization_cases
from core.model.period_axis import canonical_fiscal_periods
from core.model.per_share import compute_per_share_series
from core.tests.cross_company_fixtures import (
    PERIODS,
    ALL_ROBUST_CASES,
    RobustCompanyCase,
    asset_light_services_case,
    capital_intensive_manufacturer_case,
    inventory_retail_case,
)
from core.tests.test_normalization import (
    _inject_formula_and_cached_value,
    _set_normalization_treatment,
)
from core.tests.test_per_share import DEMO_ASSUMPTIONS, DEMO_JSON, _share_enabled_demo
from core.trainer.check_context import (
    classification_overrides_for_check,
    load_check_context,
    normalization_treatments_for_check,
)
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
EARNINGS_NORMALIZATION_SHEET = "Earnings Normalization"
JUDGMENT_SHEET = "Accounting Judgment"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _cases() -> list[RobustCompanyCase]:
    return [factory() for factory in ALL_ROBUST_CASES]


def _fill_all_practice_formulas(trainer: Path, answer: Path) -> None:
    smap = load_semantic_map(answer)
    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()


def _round_trip(case: RobustCompanyCase, tmp_path: Path):
    payload = standardized_to_payload(case.financials)
    json_path = tmp_path / f"{case.key}.json"
    json_path.write_text(json.dumps(payload), encoding="utf-8")
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(json_path), doc_type=DocumentType.OTHER)]
    )


def _set_judgment_treatment_by_label(trainer_path: Path, label_substr: str, treatment: str):
    wb = load_workbook(trainer_path, data_only=False)
    ws = wb[JUDGMENT_SHEET]
    for row in range(5, (ws.max_row or 5) + 1):
        label = str(ws.cell(row, 2).value or "")
        if label_substr.lower() in label.lower():
            ws.cell(row, 6).value = treatment
            wb.save(trainer_path)
            wb.close()
            return row
    wb.close()
    raise AssertionError(f"No Accounting Judgment row matching {label_substr!r}")


@pytest.mark.parametrize(
    "factory",
    [
        asset_light_services_case,
        inventory_retail_case,
        capital_intensive_manufacturer_case,
    ],
    ids=["services", "retail", "manufacturer"],
)
def test_fixture_construction_and_integrity(factory):
    case = factory()
    fin = case.financials
    assert len(fin.periods) == 5
    assert [p.end_date for p in fin.periods] == list(PERIODS)
    identities = [line_identity(item).key() for item in fin.income_statement]
    identities += [line_identity(item).key() for item in fin.balance_sheet]
    identities += [line_identity(item).key() for item in fin.cash_flow]
    assert len(identities) == len(set(identities))

    if case.expect_per_share:
        assert fin.historical_shares is not None
        assert fin.historical_shares.scale_basis == "financial_statement_units"
        assert set(fin.historical_shares.diluted_weighted_average) == set(PERIODS)
    else:
        assert fin.historical_shares is None

    if case.expect_normalization:
        assert case.assumptions.get("normalizationCandidates")
    else:
        assert not (case.assumptions.get("normalizationCandidates") or [])

    validate_financials_identities(fin)
    report = reconcile_financials(fin)
    assert all(report.checksums.values())
    ReferenceModelBuilder(fin, case.assumptions)


@pytest.mark.parametrize("case", _cases(), ids=lambda c: c.key)
def test_round_trip_ingest_and_surface_matrix(case: RobustCompanyCase, tmp_path):
    ingested = _round_trip(case, tmp_path)
    assert ingested.ticker == case.financials.ticker
    assert list(canonical_fiscal_periods(ingested)) == list(PERIODS)
    if case.expect_per_share:
        assert ingested.historical_shares is not None
        assert ingested.historical_shares.diluted_weighted_average == (
            case.financials.historical_shares.diluted_weighted_average
        )
    else:
        assert ingested.historical_shares is None

    trainer, answer = build_training_workbook(
        ingested,
        tmp_path / f"{case.key}_Trainer.xlsx",
        case.assumptions,
    )
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == case.expected_families
    assert len(smap.all_ordered()) == case.expected_cells

    wb = load_workbook(answer)
    assert EARNINGS_QUALITY_SHEET in wb.sheetnames
    assert WORKING_CAPITAL_SHEET in wb.sheetnames
    if case.expect_normalization:
        assert EARNINGS_NORMALIZATION_SHEET in wb.sheetnames
    else:
        assert EARNINGS_NORMALIZATION_SHEET not in wb.sheetnames
    if case.expect_per_share:
        assert PER_SHARE_SHEET in wb.sheetnames
    else:
        assert PER_SHARE_SHEET not in wb.sheetnames

    for name in DEFERRED_TAB_NAMES:
        assert name in wb.sheetnames
        assert wb[name].sheet_state == "hidden"
        assert wb[name]["A1"].value == DEFERRED_PLACEHOLDER
    wb.close()

    summary = check_workbook(trainer)
    assert summary.total == case.expected_cells
    assert summary.blank == case.expected_cells
    assert summary.correct == 0
    assert summary.incorrect == 0

    _fill_all_practice_formulas(trainer, answer)
    filled = check_workbook(trainer)
    assert filled.correct == case.expected_cells
    assert filled.blank == 0
    assert filled.incorrect == 0


def test_services_omits_asset_scaled_quality(tmp_path):
    case = asset_light_services_case()
    trainer, answer = build_training_workbook(
        case.financials, tmp_path / "SVC.xlsx", case.assumptions
    )
    smap = load_semantic_map(answer)
    families = {c.family_id for c in smap.all_ordered()}
    assert "operating_cash_flow_link" in families
    assert "cash_conversion_ratio" in families
    assert "total_accruals" in families
    assert "average_total_assets" not in families
    assert "accrual_ratio" not in families
    assert "accrual_ratio_change" not in families
    assert WORKING_CAPITAL_SHEET in load_workbook(answer).sheetnames
    assert check_workbook(trainer).blank == 248


def test_retail_declining_revenue_working_capital(tmp_path):
    case = inventory_retail_case()
    _, answer = build_training_workbook(
        case.financials, tmp_path / "RET_WC.xlsx", case.assumptions
    )
    smap = load_semantic_map(answer)
    rev = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "revenue_change" and c.period_end == "2024-12-31"
    )
    assert float(rev.expected_value) < 0
    for family_id in (
        "incremental_nowc_to_revenue_change",
        "incremental_owca_to_revenue_change",
        "incremental_owcl_to_revenue_change",
    ):
        comp = next(
            c
            for c in smap.all_ordered()
            if c.family_id == family_id and c.period_end == "2024-12-31"
        )
        upper = comp.formula.upper()
        assert "ABS(" not in upper
        assert "IFERROR" not in upper


def test_retail_sti_judgment_live_check(tmp_path):
    case = inventory_retail_case()
    trainer, answer = build_training_workbook(
        case.financials, tmp_path / "RET_STI.xlsx", case.assumptions
    )
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "nowc_to_revenue" and c.period_index == 4
    )
    ref = float(comp.expected_value)
    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula=f"={ref}", cached_value=ref
    )
    assert check_workbook(trainer).correct == 1

    _set_judgment_treatment_by_label(
        trainer, "Short-term investment", "Operating Working Capital Asset"
    )

    ctx = load_check_context(answer)
    wb = load_workbook(trainer, data_only=False)
    overrides = classification_overrides_for_check(wb, ctx)
    wb.close()
    periods = list(canonical_fiscal_periods(case.financials))
    alt_anchor = compute_anchor(
        case.financials, periods, classification_overrides=overrides
    )
    alt = float(expected_value_for_component(alt_anchor, comp))
    assert alt != pytest.approx(ref)

    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula="=999", cached_value=ref
    )
    stale = check_workbook(trainer)
    assert stale.incorrect == 1

    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula=f"={alt}", cached_value=alt
    )
    assert check_workbook(trainer).correct == 1


def test_retail_normalization_changes_normalized_eps_not_reported(tmp_path):
    case = inventory_retail_case()
    trainer, answer = build_training_workbook(
        case.financials, tmp_path / "RET_NORM.xlsx", case.assumptions
    )
    smap = load_semantic_map(answer)
    # FY2023 = period_index 2 (restructuring year)
    reported = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "reported_diluted_eps" and c.period_index == 2
    )
    normalized = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "normalized_diluted_eps" and c.period_index == 2
    )
    ref_reported = float(reported.expected_value)
    ref_norm = float(normalized.expected_value)
    assert ref_norm != pytest.approx(ref_reported)

    _inject_formula_and_cached_value(
        trainer,
        normalized.tab,
        normalized.cell,
        formula=f"={ref_norm}",
        cached_value=ref_norm,
    )
    assert check_workbook(trainer).correct == 1

    _set_normalization_treatment(trainer, "Recurring")
    ctx = load_check_context(answer)
    wb = load_workbook(trainer, data_only=False)
    treatments = normalization_treatments_for_check(wb, ctx)
    wb.close()
    periods = list(canonical_fiscal_periods(case.financials))
    anchor = compute_anchor(case.financials, periods)
    cases = normalization_cases(case.financials, periods, case.assumptions)
    alt_norm_series = compute_normalization_series(
        case.financials, periods, anchor, cases, treatments
    )
    alt_ps = compute_per_share_series(case.financials, periods, anchor)
    alt_reported = float(
        expected_value_for_component(
            anchor, reported, normalization=alt_norm_series, per_share=alt_ps
        )
    )
    alt_normalized = float(
        expected_value_for_component(
            anchor, normalized, normalization=alt_norm_series, per_share=alt_ps
        )
    )
    assert alt_reported == pytest.approx(ref_reported)
    assert alt_normalized == pytest.approx(ref_reported)
    assert alt_normalized != pytest.approx(ref_norm)

    _inject_formula_and_cached_value(
        trainer,
        normalized.tab,
        normalized.cell,
        formula="=999",
        cached_value=ref_norm,
    )
    assert check_workbook(trainer).incorrect == 1

    _inject_formula_and_cached_value(
        trainer,
        normalized.tab,
        normalized.cell,
        formula=f"={alt_normalized}",
        cached_value=alt_normalized,
    )
    assert check_workbook(trainer).correct == 1


def test_manufacturer_lease_and_falling_share_effects(tmp_path):
    case = capital_intensive_manufacturer_case()
    trainer, answer = build_training_workbook(
        case.financials, tmp_path / "MFR.xlsx", case.assumptions
    )
    smap = load_semantic_map(answer)

    # Falling shares + positive NI -> positive share-count effect on EPS change
    for period_end in ("2023-12-31", "2024-12-31", "2025-12-31"):
        effect = next(
            c
            for c in smap.all_ordered()
            if c.family_id == "share_count_effect_on_diluted_eps_change"
            and c.period_end == period_end
        )
        assert float(effect.expected_value) > 0

    # No normalized-per-share families without normalization
    assert not any(c.category == "normalized_per_share" for c in smap.all_ordered())

    net_debt = next(
        c for c in smap.all_ordered() if c.family_id == "net_debt" and c.period_index == 4
    )
    reported = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "reported_diluted_eps" and c.period_index == 4
    )
    ref_nd = float(net_debt.expected_value)
    ref_eps = float(reported.expected_value)

    _set_judgment_treatment_by_label(trainer, "lease", "Financial Liability")
    ctx = load_check_context(answer)
    wb = load_workbook(trainer, data_only=False)
    overrides = classification_overrides_for_check(wb, ctx)
    wb.close()
    periods = list(canonical_fiscal_periods(case.financials))
    alt_anchor = compute_anchor(
        case.financials, periods, classification_overrides=overrides
    )
    alt_ps = compute_per_share_series(case.financials, periods, alt_anchor)
    alt_nd = float(expected_value_for_component(alt_anchor, net_debt))
    alt_eps = float(
        expected_value_for_component(alt_anchor, reported, per_share=alt_ps)
    )
    assert alt_nd != pytest.approx(ref_nd)
    assert alt_eps == pytest.approx(ref_eps)


@pytest.mark.parametrize("case", _cases(), ids=lambda c: c.key)
def test_trusted_source_tamper_before_recolor(case: RobustCompanyCase, tmp_path):
    trainer, answer = build_training_workbook(
        case.financials, tmp_path / f"{case.key}_TAMPER.xlsx", case.assumptions
    )
    smap = load_semantic_map(answer)
    comp = smap.all_ordered()[0]
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    # Tamper first-period Revenue on Income Statement
    ws = wb["Income Statement"]
    # Find revenue row
    rev_row = None
    for r in range(1, (ws.max_row or 1) + 1):
        if str(ws.cell(r, 1).value or "").lower() == "revenue":
            rev_row = r
            break
    assert rev_row is not None
    ws.cell(rev_row, 2).value = 1
    wb.save(trainer)
    wb.close()
    with pytest.raises(ValueError, match="Trusted workbook cell was modified"):
        check_workbook(trainer)
    wb = load_workbook(trainer, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


def test_ordinary_demo_and_share_enabled_surfaces_preserved(tmp_path):
    adapter = HKManualDocumentAdapter()
    data = adapter.ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )
    trainer, answer = build_training_workbook(data, tmp_path / "DEMO_BASE.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 70
    assert len(smap.all_ordered()) == 294
    assert check_workbook(trainer).blank == 294

    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        data, tmp_path / "DEMO_NORM.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(group_components_by_family(smap_n)) == 74
    assert len(smap_n.all_ordered()) == 314
    assert check_workbook(trainer_n).blank == 314

    shares = _share_enabled_demo()
    trainer_s, answer_s = build_training_workbook(shares, tmp_path / "SHARE.xlsx")
    smap_s = load_semantic_map(answer_s)
    assert len(group_components_by_family(smap_s)) == 78
    assert len(smap_s.all_ordered()) == 328

    trainer_sn, answer_sn = build_training_workbook(
        shares, tmp_path / "SHARE_NORM.xlsx", assumptions
    )
    smap_sn = load_semantic_map(answer_sn)
    assert len(group_components_by_family(smap_sn)) == 86
    assert len(smap_sn.all_ordered()) == 366

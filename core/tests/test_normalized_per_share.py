"""Step 9F.3 — historical normalized diluted EPS bridge."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import FinancialPeriod, LineItem, StandardizedFinancials
from core.engine.component_catalog import (
    NORMALIZED_PER_SHARE_COMPONENT_CATALOG,
    PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG,
    PER_SHARE_COMPONENT_CATALOG,
    expand_normalized_per_share_specs,
)
from core.engine.reference_model import PER_SHARE_SHEET, ReferenceModelBuilder
from core.model.historical_expected import (
    expected_value_for_component,
    normalized_per_share_expected_series,
)
from core.model.normalization import NormalizationSeries
from core.model.normalized_per_share import compute_normalized_per_share_series
from core.model.per_share import PerShareSeries
from core.model.ratio_values import UNDEFINED_RATIO
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family
from core.tests.test_normalization import (
    _inject_formula_and_cached_value,
    _set_normalization_treatment,
)
from core.tests.test_per_share import (
    DEMO_ASSUMPTIONS,
    _fill_rgb,
    _ingest_demo,
    _share_enabled_demo,
)

ROOT = Path(__file__).resolve().parents[2]


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _norm_series(*, after_tax, normalized_ni, pretax=None, nopat=None):
    n = len(after_tax)
    return NormalizationSeries(
        pretax_adjustment=tuple(pretax if pretax is not None else (0.0,) * n),
        after_tax_adjustment=tuple(after_tax),
        normalized_nopat=tuple(nopat if nopat is not None else normalized_ni),
        normalized_net_income=tuple(normalized_ni),
    )


def _ps(*, shares, reported_eps, eps_change=None):
    if eps_change is None:
        change: list[float | None] = [None]
        for i in range(1, len(reported_eps)):
            change.append(reported_eps[i] - reported_eps[i - 1])
        eps_change = tuple(change)
    return PerShareSeries(
        diluted_weighted_average_shares=shares,
        reported_diluted_eps=reported_eps,
        nopat_per_diluted_share=reported_eps,
        diluted_eps_change=eps_change,
        diluted_share_count_change=tuple(
            None if i == 0 else shares[i] - shares[i - 1]
            for i in range(len(shares))
        ),
    )


def test_ordinary_normalized_per_share_bridge():
    shares = (100.0, 100.0, 110.0)
    reported_eps = (1.0, 1.2, 1.2)
    after_tax = (0.0, 10.0, 11.0)
    normalized_ni = (100.0, 130.0, 143.0)
    series = compute_normalized_per_share_series(
        _norm_series(after_tax=after_tax, normalized_ni=normalized_ni),
        _ps(shares=shares, reported_eps=reported_eps),
    )
    assert series.normalization_adjustment_per_diluted_share == pytest.approx(
        (0.0, 0.1, 0.1)
    )
    assert series.normalized_diluted_eps == pytest.approx((1.0, 1.3, 1.3))
    assert series.normalized_diluted_eps_change[0] is None
    assert series.normalization_effect_on_diluted_eps_change[0] is None
    assert series.normalized_diluted_eps_change[1] == pytest.approx(0.3)
    assert series.normalization_effect_on_diluted_eps_change[1] == pytest.approx(0.1)
    assert series.normalized_diluted_eps_change[2] == pytest.approx(0.0)
    assert series.normalization_effect_on_diluted_eps_change[2] == pytest.approx(0.0)

    for i in range(3):
        assert reported_eps[i] + float(
            series.normalization_adjustment_per_diluted_share[i]
        ) == pytest.approx(float(series.normalized_diluted_eps[i]))
    for i in (1, 2):
        reported_chg = reported_eps[i] - reported_eps[i - 1]
        assert reported_chg + float(
            series.normalization_effect_on_diluted_eps_change[i]
        ) == pytest.approx(float(series.normalized_diluted_eps_change[i]))


def test_normalized_per_share_edge_cases():
    shares = (100.0, 100.0)
    # Zero adjustment -> normalized EPS = reported EPS
    s0 = compute_normalized_per_share_series(
        _norm_series(after_tax=(0.0, 0.0), normalized_ni=(100.0, 120.0)),
        _ps(shares=shares, reported_eps=(1.0, 1.2)),
    )
    assert s0.normalization_adjustment_per_diluted_share == pytest.approx((0.0, 0.0))
    assert s0.normalized_diluted_eps == pytest.approx((1.0, 1.2))
    assert s0.normalization_effect_on_diluted_eps_change[1] == pytest.approx(0.0)

    # Negative adjustment retains sign
    sn = compute_normalized_per_share_series(
        _norm_series(after_tax=(0.0, -10.0), normalized_ni=(100.0, 110.0)),
        _ps(shares=shares, reported_eps=(1.0, 1.2)),
    )
    assert sn.normalization_adjustment_per_diluted_share[1] == pytest.approx(-0.1)

    # Nonzero #N/A after-tax adjustment
    sna = compute_normalized_per_share_series(
        _norm_series(
            after_tax=(0.0, UNDEFINED_RATIO),
            normalized_ni=(100.0, UNDEFINED_RATIO),
        ),
        _ps(shares=shares, reported_eps=(1.0, 1.2)),
    )
    assert sna.normalization_adjustment_per_diluted_share[1] == UNDEFINED_RATIO
    assert sna.normalized_diluted_eps[1] == UNDEFINED_RATIO
    assert sna.normalized_diluted_eps_change[1] == UNDEFINED_RATIO
    assert sna.normalization_effect_on_diluted_eps_change[1] == UNDEFINED_RATIO

    # Zero after-tax with undefined tax semantics still numeric 0
    sz = compute_normalized_per_share_series(
        _norm_series(after_tax=(0.0, 0.0), normalized_ni=(100.0, 120.0)),
        _ps(shares=shares, reported_eps=(1.0, 1.2)),
    )
    assert sz.normalization_adjustment_per_diluted_share[1] == pytest.approx(0.0)
    assert sz.normalized_diluted_eps[1] == pytest.approx(1.2)

    # Unchanged adjustment/share -> effect 0
    su = compute_normalized_per_share_series(
        _norm_series(after_tax=(10.0, 10.0), normalized_ni=(110.0, 130.0)),
        _ps(shares=shares, reported_eps=(1.0, 1.2)),
    )
    assert su.normalization_effect_on_diluted_eps_change[1] == pytest.approx(0.0)

    # Length mismatch
    with pytest.raises(ValueError, match="length mismatch"):
        compute_normalized_per_share_series(
            _norm_series(after_tax=(0.0,), normalized_ni=(100.0,)),
            _ps(shares=(100.0, 110.0), reported_eps=(1.0, 1.1)),
        )

    # Non-positive shares
    with pytest.raises(ValueError, match="positive diluted"):
        compute_normalized_per_share_series(
            _norm_series(after_tax=(0.0, 0.0), normalized_ni=(100.0, 120.0)),
            _ps(shares=(100.0, 0.0), reported_eps=(1.0, 0.0)),
        )

    # Inconsistent level bridge
    bad_level = PerShareSeries(
        diluted_weighted_average_shares=(100.0, 100.0),
        reported_diluted_eps=(1.0, 1.2),
        nopat_per_diluted_share=(1.0, 1.2),
        diluted_eps_change=(None, 0.2),
        diluted_share_count_change=(None, 0.0),
    )
    with pytest.raises(ValueError, match="level bridge"):
        compute_normalized_per_share_series(
            _norm_series(after_tax=(0.0, 10.0), normalized_ni=(100.0, 999.0)),
            bad_level,
        )

    # Inconsistent change bridge
    bad_chg = _ps(shares=shares, reported_eps=(1.0, 1.2), eps_change=(None, 9.0))
    with pytest.raises(ValueError, match="change bridge"):
        compute_normalized_per_share_series(
            _norm_series(after_tax=(0.0, 10.0), normalized_ni=(100.0, 130.0)),
            bad_chg,
        )


def test_catalog_expand_and_gating(tmp_path):
    assert [f.order for f in NORMALIZED_PER_SHARE_COMPONENT_CATALOG] == [75, 76, 77, 78]
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    specs = expand_normalized_per_share_specs(periods, start_order=294)
    assert len(specs) == 18

    # no shares + no normalization
    base = _ingest_demo()
    b0 = ReferenceModelBuilder(base)
    assert b0.normalized_per_share_specs == ()
    assert b0.normalized_per_share_series is None

    # shares + no normalization
    shares_only = _share_enabled_demo()
    b1 = ReferenceModelBuilder(shares_only)
    assert b1.per_share_series is not None
    assert b1.normalized_per_share_specs == ()
    assert b1.normalized_per_share_series is None
    trainer1, answer1 = build_training_workbook(shares_only, tmp_path / "SO.xlsx")
    smap1 = load_semantic_map(answer1)
    assert len(smap1.all_ordered()) == 328
    assert len(group_components_by_family(smap1)) == 78
    assert not any(c.category == "normalized_per_share" for c in smap1.all_ordered())

    # normalization + no shares
    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        base, tmp_path / "NORM_ONLY.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(smap_n.all_ordered()) == 314
    assert PER_SHARE_SHEET not in load_workbook(answer_n).sheetnames
    assert not any(c.category == "normalized_per_share" for c in smap_n.all_ordered())

    # shares + normalization
    both = _share_enabled_demo()
    trainer_b, answer_b = build_training_workbook(
        both, tmp_path / "BOTH.xlsx", assumptions
    )
    smap_b = load_semantic_map(answer_b)
    assert len(smap_b.all_ordered()) == 366
    assert len(group_components_by_family(smap_b)) == 86
    nps = [c for c in smap_b.all_ordered() if c.category == "normalized_per_share"]
    assert len(nps) == 18
    assert {c.family_id for c in nps} == {
        f.id for f in NORMALIZED_PER_SHARE_COMPONENT_CATALOG
    }
    builder = ReferenceModelBuilder(both, assumptions)
    assert builder.normalized_per_share_series is not None
    series = normalized_per_share_expected_series(
        builder.normalization_series, builder.per_share_series
    )
    assert set(series) == {f.id for f in NORMALIZED_PER_SHARE_COMPONENT_CATALOG}
    assert check_workbook(trainer_b).blank == 366


def test_normalized_bridge_sheet_layout_and_trust(tmp_path):
    data = _share_enabled_demo()
    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer, answer = build_training_workbook(data, tmp_path / "LAYOUT.xlsx", assumptions)
    smap = load_semantic_map(answer)
    nps = [c for c in smap.all_ordered() if c.category == "normalized_per_share"]
    assert len(nps) == 18

    wb = load_workbook(answer)
    ws = wb[PER_SHARE_SHEET]
    assert ws.cell(15, 1).value == "DILUTED EPS CHANGE ATTRIBUTION"
    assert ws.cell(20, 1).value == "DILUTED EPS CHANGE ATTRIBUTION CHECK"
    assert ws.cell(22, 1).value == "NORMALIZED DILUTED EPS BRIDGE"
    assert ws.cell(23, 1).value == "After-Tax Normalization Adjustment"
    assert ws.cell(24, 1).value == "Normalized Net Income"
    assert ws.cell(25, 1).value == "Normalization Adjustment per Diluted Share"
    assert ws.cell(26, 1).value == "Normalized Diluted EPS"
    assert ws.cell(27, 1).value == "NORMALIZED EPS LEVEL CHECK"
    assert ws.cell(29, 1).value == "Change in Normalized Diluted EPS"
    assert ws.cell(30, 1).value == "Normalization Effect on Change in Diluted EPS"
    assert ws.cell(31, 1).value == "NORMALIZED EPS CHANGE CHECK"
    assert ws.cell(29, 2).value == "N/A"
    assert ws.cell(30, 2).value == "N/A"
    assert ws.cell(31, 2).value == "N/A"
    practice = {(c.tab, c.cell) for c in smap.all_ordered()}
    assert (PER_SHARE_SHEET, "B23") not in practice
    assert (PER_SHARE_SHEET, "B24") not in practice
    assert (PER_SHARE_SHEET, "B27") not in practice
    assert (PER_SHARE_SHEET, "C31") not in practice
    for comp in nps:
        assert "ABS" not in comp.formula.upper()
        assert "IFERROR" not in comp.formula.upper()

    wb_t = load_workbook(trainer)
    for comp in nps:
        row, col = parse_cell_ref(comp.cell)
        cell = wb_t[comp.tab].cell(row=row, column=col)
        assert cell.value is None
        assert _fill_rgb(cell) == "FFFF00"
    wb.close()
    wb_t.close()

    # Exact formula green
    comp = next(c for c in nps if c.family_id == "normalized_diluted_eps")
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    # Trusted link tamper
    trainer2, answer2 = build_training_workbook(
        data, tmp_path / "TAMPER_LINK.xlsx", assumptions
    )
    smap2 = load_semantic_map(answer2)
    comp2 = next(
        c for c in smap2.all_ordered() if c.family_id == "normalized_diluted_eps"
    )
    wb = load_workbook(trainer2, data_only=False)
    row, col = parse_cell_ref(comp2.cell)
    wb[comp2.tab].cell(row=row, column=col).value = comp2.formula
    wb[PER_SHARE_SHEET].cell(23, 2).value = 1
    wb.save(trainer2)
    wb.close()
    with pytest.raises(
        ValueError, match="Trusted workbook cell was modified: Per Share Analysis"
    ):
        check_workbook(trainer2)
    wb = load_workbook(trainer2, data_only=False)
    assert _fill_rgb(wb[comp2.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

    # Level check tamper
    trainer3, answer3 = build_training_workbook(
        data, tmp_path / "TAMPER_CHECK.xlsx", assumptions
    )
    smap3 = load_semantic_map(answer3)
    comp3 = next(
        c for c in smap3.all_ordered() if c.family_id == "normalized_diluted_eps"
    )
    wb = load_workbook(trainer3, data_only=False)
    row, col = parse_cell_ref(comp3.cell)
    wb[comp3.tab].cell(row=row, column=col).value = comp3.formula
    wb[PER_SHARE_SHEET].cell(27, 2).value = '="TAMPERED"'
    wb.save(trainer3)
    wb.close()
    with pytest.raises(
        ValueError, match="Trusted workbook cell was modified: Per Share Analysis"
    ):
        check_workbook(trainer3)


def test_live_normalization_judgment_changes_normalized_eps(tmp_path):
    data = _share_enabled_demo()
    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer, answer = build_training_workbook(data, tmp_path / "LIVE.xlsx", assumptions)
    smap = load_semantic_map(answer)
    n_eps = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "normalized_diluted_eps" and c.period_index == 2
    )
    reported = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "reported_diluted_eps" and c.period_index == 2
    )
    ref_norm = float(n_eps.expected_value)
    ref_reported = float(reported.expected_value)
    assert ref_norm != pytest.approx(ref_reported)

    _inject_formula_and_cached_value(
        trainer,
        n_eps.tab,
        n_eps.cell,
        formula=f"={ref_norm}",
        cached_value=ref_norm,
    )
    assert check_workbook(trainer).correct == 1

    _set_normalization_treatment(trainer, "Recurring")

    from core.data.standardized_io import standardized_from_payload
    from core.model.financial_math import compute_anchor
    from core.model.normalization import compute_normalization_series, normalization_cases
    from core.model.period_axis import canonical_fiscal_periods
    from core.model.per_share import compute_per_share_series
    from core.trainer.check_context import (
        load_check_context,
        normalization_treatments_for_check,
    )

    ctx = load_check_context(answer)
    wb = load_workbook(trainer, data_only=False)
    treatments = normalization_treatments_for_check(wb, ctx)
    wb.close()
    fin = standardized_from_payload(ctx.source_payload)
    periods = list(canonical_fiscal_periods(fin))
    anchor = compute_anchor(fin, periods)
    cases = normalization_cases(fin, periods, assumptions)
    alt_norm = compute_normalization_series(fin, periods, anchor, cases, treatments)
    alt_ps = compute_per_share_series(fin, periods, anchor)
    alt_expected = float(
        expected_value_for_component(
            anchor, n_eps, normalization=alt_norm, per_share=alt_ps
        )
    )
    assert alt_expected == pytest.approx(ref_reported)
    assert alt_expected != pytest.approx(ref_norm)

    # Stale Non-recurring cache under Recurring -> red
    _inject_formula_and_cached_value(
        trainer,
        n_eps.tab,
        n_eps.cell,
        formula="=999",
        cached_value=ref_norm,
    )
    summary_stale = check_workbook(trainer)
    assert summary_stale.incorrect == 1
    assert summary_stale.correct == 0

    _inject_formula_and_cached_value(
        trainer,
        n_eps.tab,
        n_eps.cell,
        formula=f"={alt_expected}",
        cached_value=alt_expected,
    )
    assert check_workbook(trainer).correct == 1
    # Reported EPS expected unchanged by normalization treatment
    assert float(reported.expected_value) == pytest.approx(ref_reported)


def test_normalized_eps_na_check(tmp_path):
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)

    def vals(a, b):
        return {d1: a, d2: b}

    fin = StandardizedFinancials(
        ticker="NA",
        company_name="NA Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=[
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ],
        income_statement=[
            _li("Revenue", vals(1000, 1100)),
            _li("Finance costs", vals(-10, -11)),
            _li("Finance income", vals(0, 0)),
            _li("Restructuring expense", vals(0, -20), concept="restructuring_expense"),
            # Pretax zero in period 2 -> undefined ETR / after-tax adj #N/A
            _li("Profit before tax", vals(200, 0)),
            _li("Income tax expense", vals(-30, 0)),
            _li("Profit for the year", vals(170, 0)),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", vals(50, 55)),
            _li("Trade receivables", vals(40, 45)),
            _li("Property, plant and equipment", vals(400, 420)),
            _li("Trade payables", vals(30, 35)),
            _li("Bank borrowings", vals(200, 210)),
            _li("Total equity", vals(260, 275)),
        ],
        cash_flow=[_li("Net cash from operating activities", vals(50, 60))],
    )
    from core.data.interface import HistoricalShareData
    from core.model.per_share import SUPPORTED_SHARE_SCALE_BASIS

    fin.historical_shares = HistoricalShareData(
        scale_basis=SUPPORTED_SHARE_SCALE_BASIS,
        diluted_weighted_average={d1: 100.0, d2: 105.0},
    )
    assumptions = {
        "normalizationCandidates": [
            {
                "selector": "concept:restructuring_expense",
                "referenceTreatment": "Non-recurring",
                "scope": "operating_pretax_effective_tax",
                "topic": "Restructuring",
                "referenceRationale": "One-off.",
                "consequenceNote": "Affects normalized earnings.",
            }
        ]
    }
    trainer, answer = build_training_workbook(fin, tmp_path / "NA.xlsx", assumptions)
    smap = load_semantic_map(answer)
    n_eps = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "normalized_diluted_eps" and c.period_index == 1
    )
    assert n_eps.expected_value == UNDEFINED_RATIO
    adj = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "normalization_adjustment_per_diluted_share"
        and c.period_index == 1
    )
    assert adj.expected_value == UNDEFINED_RATIO

    _inject_formula_and_cached_value(
        trainer,
        n_eps.tab,
        n_eps.cell,
        formula="=NA()",
        cached_value=UNDEFINED_RATIO,
    )
    assert check_workbook(trainer).correct == 1

    trainer2, _ = build_training_workbook(fin, tmp_path / "NA0.xlsx", assumptions)
    smap2 = load_semantic_map(answer)
    n_eps2 = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "normalized_diluted_eps" and c.period_index == 1
    )
    _inject_formula_and_cached_value(
        trainer2,
        n_eps2.tab,
        n_eps2.cell,
        formula="=0",
        cached_value=0.0,
    )
    assert check_workbook(trainer2).incorrect == 1

    # Zero adjustment with zero pretax remains numeric (Recurring => adj 0)
    fin0 = StandardizedFinancials(
        ticker="Z",
        company_name="Z Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=fin.periods,
        income_statement=[
            _li("Revenue", vals(1000, 1100)),
            _li("Finance costs", vals(-10, -11)),
            _li("Finance income", vals(0, 0)),
            _li("Restructuring expense", vals(0, -20), concept="restructuring_expense"),
            _li("Profit before tax", vals(200, 0)),
            _li("Income tax expense", vals(-30, 0)),
            _li("Profit for the year", vals(170, 0)),
        ],
        balance_sheet=fin.balance_sheet,
        cash_flow=fin.cash_flow,
        historical_shares=fin.historical_shares,
    )
    assumptions0 = {
        "normalizationCandidates": [
            {
                "selector": "concept:restructuring_expense",
                "referenceTreatment": "Recurring",
                "scope": "operating_pretax_effective_tax",
                "topic": "Restructuring",
                "referenceRationale": "Keep in earnings.",
                "consequenceNote": "No normalization adjustment.",
            }
        ]
    }
    trainer_z, answer_z = build_training_workbook(
        fin0, tmp_path / "Z.xlsx", assumptions0
    )
    smap_z = load_semantic_map(answer_z)
    n_eps_z = next(
        c
        for c in smap_z.all_ordered()
        if c.family_id == "normalized_diluted_eps" and c.period_index == 1
    )
    reported_z = next(
        c
        for c in smap_z.all_ordered()
        if c.family_id == "reported_diluted_eps" and c.period_index == 1
    )
    assert n_eps_z.expected_value != UNDEFINED_RATIO
    assert float(n_eps_z.expected_value) == pytest.approx(float(reported_z.expected_value))
    adj_z = next(
        c
        for c in smap_z.all_ordered()
        if c.family_id == "normalization_adjustment_per_diluted_share"
        and c.period_index == 1
    )
    assert float(adj_z.expected_value) == pytest.approx(0.0)


def test_expected_requires_both_series(tmp_path):
    data = _share_enabled_demo()
    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    _, answer = build_training_workbook(data, tmp_path / "REQ.xlsx", assumptions)
    smap = load_semantic_map(answer)
    comp = next(c for c in smap.all_ordered() if c.family_id == "normalized_diluted_eps")
    builder = ReferenceModelBuilder(data, assumptions)
    with pytest.raises(ValueError, match="requires a NormalizationSeries"):
        expected_value_for_component(
            builder.anchor, comp, normalization=None, per_share=builder.per_share_series
        )
    with pytest.raises(ValueError, match="requires a PerShareSeries"):
        expected_value_for_component(
            builder.anchor,
            comp,
            normalization=builder.normalization_series,
            per_share=None,
        )

"""Step 9F.2 — diluted EPS earnings / share-count attribution."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from core.engine.component_catalog import (
    PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG,
    PER_SHARE_COMPONENT_CATALOG,
    expand_per_share_attribution_specs,
)
from core.engine.reference_model import PER_SHARE_SHEET, ReferenceModelBuilder
from core.model.historical_expected import (
    expected_value_for_component,
    per_share_attribution_expected_series,
)
from core.model.per_share import PerShareSeries
from core.model.per_share_attribution import compute_per_share_attribution_series
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.tests.test_per_share import (
    DEMO_ASSUMPTIONS,
    _fill_rgb,
    _ingest_demo,
    _share_enabled_demo,
)

ROOT = Path(__file__).resolve().parents[2]


def _anchor(net_income: tuple[float, ...]):
    return SimpleNamespace(historical=SimpleNamespace(net_income=list(net_income)))


def _per_share(
    *,
    shares: tuple[float, ...],
    net_income: tuple[float, ...] | None = None,
    eps: tuple[float, ...] | None = None,
    eps_change: tuple[float | None, ...] | None = None,
) -> PerShareSeries:
    if net_income is None:
        # Default NI so EPS = NI/shares when eps not overridden
        net_income = tuple(1.0 * s for s in shares)
    if eps is None:
        eps = tuple(net_income[i] / shares[i] for i in range(len(shares)))
    if eps_change is None:
        change: list[float | None] = [None]
        for i in range(1, len(eps)):
            change.append(eps[i] - eps[i - 1])
        eps_change = tuple(change)
    return PerShareSeries(
        diluted_weighted_average_shares=shares,
        reported_diluted_eps=eps,
        nopat_per_diluted_share=tuple(eps),
        diluted_eps_change=eps_change,
        diluted_share_count_change=tuple(
            None if i == 0 else shares[i] - shares[i - 1] for i in range(len(shares))
        ),
        earnings_numerator=net_income,
    )


def test_ordinary_midpoint_attribution():
    ni = (100.0, 120.0)
    shares = (100.0, 110.0)
    eps = (ni[0] / shares[0], ni[1] / shares[1])
    ps = _per_share(shares=shares, net_income=ni, eps=eps)
    series = compute_per_share_attribution_series(_anchor(ni), ps)

    assert series.reported_net_income_change[0] is None
    assert series.earnings_effect_on_diluted_eps_change[0] is None
    assert series.share_count_effect_on_diluted_eps_change[0] is None
    assert series.diluted_eps_change_from_drivers[0] is None

    assert series.reported_net_income_change[1] == pytest.approx(20.0)
    q_t, q_p = 1.0 / 110.0, 1.0 / 100.0
    earnings = 20.0 * (q_t + q_p) / 2.0
    share_fx = (q_t - q_p) * (120.0 + 100.0) / 2.0
    assert series.earnings_effect_on_diluted_eps_change[1] == pytest.approx(earnings)
    assert series.share_count_effect_on_diluted_eps_change[1] == pytest.approx(share_fx)
    assert series.share_count_effect_on_diluted_eps_change[1] == pytest.approx(-0.1)
    assert series.diluted_eps_change_from_drivers[1] == pytest.approx(eps[1] - eps[0])
    assert series.diluted_eps_change_from_drivers[1] == pytest.approx(
        earnings + share_fx
    )


def test_attribution_edge_cases():
    # Unchanged NI -> earnings effect 0; unchanged shares -> share effect 0
    ni = (100.0, 100.0, 120.0)
    shares = (100.0, 100.0, 100.0)
    ps = _per_share(shares=shares, net_income=ni)
    s0 = compute_per_share_attribution_series(_anchor(ni), ps)
    assert s0.earnings_effect_on_diluted_eps_change[1] == pytest.approx(0.0)
    assert s0.share_count_effect_on_diluted_eps_change[1] == pytest.approx(0.0)
    assert s0.earnings_effect_on_diluted_eps_change[2] == pytest.approx(0.2)
    assert s0.share_count_effect_on_diluted_eps_change[2] == pytest.approx(0.0)

    # Positive earnings + rising shares -> negative share-count effect
    ni_p = (100.0, 100.0)
    shares_p = (100.0, 110.0)
    sp = compute_per_share_attribution_series(
        _anchor(ni_p), _per_share(shares=shares_p, net_income=ni_p)
    )
    assert sp.share_count_effect_on_diluted_eps_change[1] < 0

    # Negative earnings + rising shares -> share-count effect may be positive
    ni_n = (-100.0, -100.0)
    shares_n = (100.0, 110.0)
    sn = compute_per_share_attribution_series(
        _anchor(ni_n), _per_share(shares=shares_n, net_income=ni_n)
    )
    assert sn.share_count_effect_on_diluted_eps_change[1] > 0

    # Falling shares with positive earnings -> positive share-count effect
    ni_f = (100.0, 100.0)
    shares_f = (110.0, 100.0)
    sf = compute_per_share_attribution_series(
        _anchor(ni_f), _per_share(shares=shares_f, net_income=ni_f)
    )
    assert sf.share_count_effect_on_diluted_eps_change[1] > 0

    # Negative NI change retains sign on earnings effect
    ni_d = (120.0, 100.0)
    shares_d = (100.0, 100.0)
    sd = compute_per_share_attribution_series(
        _anchor(ni_d), _per_share(shares=shares_d, net_income=ni_d)
    )
    assert sd.reported_net_income_change[1] == pytest.approx(-20.0)
    assert sd.earnings_effect_on_diluted_eps_change[1] < 0

    # Zero NI still numeric (not #N/A)
    ni_z = (0.0, 0.0)
    shares_z = (100.0, 110.0)
    sz = compute_per_share_attribution_series(
        _anchor(ni_z), _per_share(shares=shares_z, net_income=ni_z)
    )
    assert sz.earnings_effect_on_diluted_eps_change[1] == pytest.approx(0.0)
    assert sz.share_count_effect_on_diluted_eps_change[1] == pytest.approx(0.0)
    assert sz.diluted_eps_change_from_drivers[1] == pytest.approx(0.0)

    # Non-positive shares
    with pytest.raises(ValueError, match="positive diluted"):
        compute_per_share_attribution_series(
            _anchor((100.0, 120.0)),
            _per_share(shares=(100.0, 0.0), net_income=(100.0, 120.0), eps=(1.0, 0.0)),
        )

    # Length mismatch
    bad = PerShareSeries(
        diluted_weighted_average_shares=(100.0,),
        reported_diluted_eps=(1.0, 1.1),
        nopat_per_diluted_share=(1.0, 1.1),
        diluted_eps_change=(None, 0.1),
        diluted_share_count_change=(None, 10.0),
        earnings_numerator=(100.0,),
    )
    with pytest.raises(ValueError, match="length mismatch"):
        compute_per_share_attribution_series(_anchor((100.0, 120.0)), bad)

    # Inconsistent EPS level
    bad_eps = _per_share(shares=(100.0, 110.0), net_income=(100.0, 120.0))
    bad_eps = PerShareSeries(
        diluted_weighted_average_shares=bad_eps.diluted_weighted_average_shares,
        reported_diluted_eps=(1.0, 9.0),
        nopat_per_diluted_share=bad_eps.nopat_per_diluted_share,
        diluted_eps_change=bad_eps.diluted_eps_change,
        diluted_share_count_change=bad_eps.diluted_share_count_change,
        earnings_numerator=bad_eps.earnings_numerator,
    )
    with pytest.raises(ValueError, match="does not reconcile before attribution"):
        compute_per_share_attribution_series(_anchor((100.0, 120.0)), bad_eps)

    # Inconsistent direct EPS change
    good = _per_share(shares=(100.0, 110.0), net_income=(100.0, 120.0))
    bad_chg = PerShareSeries(
        diluted_weighted_average_shares=good.diluted_weighted_average_shares,
        reported_diluted_eps=good.reported_diluted_eps,
        nopat_per_diluted_share=good.nopat_per_diluted_share,
        diluted_eps_change=(None, 99.0),
        diluted_share_count_change=good.diluted_share_count_change,
        earnings_numerator=good.earnings_numerator,
    )
    with pytest.raises(ValueError, match="does not reconcile"):
        compute_per_share_attribution_series(_anchor((100.0, 120.0)), bad_chg)


def test_catalog_expand_and_expected_keys():
    assert len(PER_SHARE_COMPONENT_CATALOG) == 4
    assert [f.id for f in PER_SHARE_COMPONENT_CATALOG] == [
        "reported_diluted_eps",
        "nopat_per_diluted_share",
        "diluted_eps_change",
        "diluted_share_count_change",
    ]
    assert [f.order for f in PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG] == [71, 72, 73, 74]
    assert all(
        f.period_scope == "comparable" for f in PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG
    )
    assert len({f.id for f in PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG}) == 4

    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    specs = expand_per_share_attribution_specs(periods, start_order=278)
    assert len(specs) == 16

    data = _share_enabled_demo()
    builder = ReferenceModelBuilder(data)
    assert builder.per_share_attribution_series is not None
    assert len(builder.per_share_attribution_specs) == 16
    series = per_share_attribution_expected_series(
        builder.anchor, builder.per_share_series
    )
    assert set(series) == {f.id for f in PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG}
    for i in range(1, 5):
        assert series["diluted_eps_change_from_drivers"][i] == pytest.approx(
            builder.per_share_series.diluted_eps_change[i]
        )


def test_no_share_omits_attribution(tmp_path):
    data = _ingest_demo()
    builder = ReferenceModelBuilder(data)
    assert builder.per_share_attribution_series is None
    assert builder.per_share_attribution_specs == ()
    trainer, answer = build_training_workbook(data, tmp_path / "NO_ATTR.xlsx")
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == 312
    assert len(group_components_by_family(smap)) == 74
    assert not any(c.category == "per_share_attribution" for c in smap.all_ordered())
    wb = load_workbook(answer)
    assert PER_SHARE_SHEET not in wb.sheetnames
    wb.close()
    assert check_workbook(trainer).blank == 312


def test_share_enabled_attribution_surface(tmp_path):
    data = _share_enabled_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "ATTR_BASE.xlsx")
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == 346
    assert len(group_components_by_family(smap)) == 82
    attr = [c for c in smap.all_ordered() if c.category == "per_share_attribution"]
    ps = [c for c in smap.all_ordered() if c.category in ("per_share", "per_share_attribution")]
    assert len(attr) == 16
    assert len(ps) == 34
    assert {c.family_id for c in attr} == {
        f.id for f in PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG
    }

    wb = load_workbook(answer)
    ws = wb[PER_SHARE_SHEET]
    assert ws.cell(5, 1).value == "Reported Net Income"
    assert ws.cell(7, 1).value == "Diluted Weighted-Average Shares"
    assert ws.cell(12, 1).value == "Change in Diluted EPS"
    assert ws.cell(15, 1).value == "DILUTED EPS CHANGE ATTRIBUTION"
    assert ws.cell(16, 1).value == "Change in Reported Net Income"
    assert ws.cell(17, 1).value == "Earnings Effect on Change in Diluted EPS"
    assert ws.cell(18, 1).value == "Share-Count Effect on Change in Diluted EPS"
    assert ws.cell(19, 1).value == "Diluted EPS Change from Drivers"
    assert ws.cell(20, 1).value == "DILUTED EPS CHANGE ATTRIBUTION CHECK"
    for row in range(16, 21):
        assert ws.cell(row, 2).value == "N/A"
    practice = {(c.tab, c.cell) for c in smap.all_ordered()}
    assert (PER_SHARE_SHEET, "C20") not in practice
    assert (PER_SHARE_SHEET, "B7") not in practice

    for comp in attr:
        assert "ABS" not in comp.formula.upper()
        if comp.family_id == "share_count_effect_on_diluted_eps_change":
            assert "1/" in comp.formula.replace(" ", "")
            assert "-(" in comp.formula.replace(" ", "") or ")-(1/" in comp.formula.replace(
                " ", ""
            )

    # Driver reconciles to direct diluted_eps_change expected
    for j in range(1, 5):
        driver = next(
            c
            for c in attr
            if c.family_id == "diluted_eps_change_from_drivers" and c.period_index == j
        )
        direct = next(
            c
            for c in smap.all_ordered()
            if c.family_id == "diluted_eps_change" and c.period_index == j
        )
        assert float(driver.expected_value) == pytest.approx(float(direct.expected_value))

    wb_t = load_workbook(trainer)
    for j in range(5):
        assert wb_t[PER_SHARE_SHEET].cell(7, 2 + j).value == pytest.approx(
            1000.0 + 25.0 * j
        )
    for comp in attr:
        row, col = parse_cell_ref(comp.cell)
        cell = wb_t[comp.tab].cell(row=row, column=col)
        assert cell.value is None
        assert _fill_rgb(cell) == "FFFF00"
    wb.close()
    wb_t.close()

    assert check_workbook(trainer).blank == 346

    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        data, tmp_path / "ATTR_NORM.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(smap_n.all_ordered()) == 384
    assert len(group_components_by_family(smap_n)) == 90
    assert check_workbook(trainer_n).blank == 384


def test_attribution_check_trust_and_colors(tmp_path):
    data = _share_enabled_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "ATTR_Check.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "earnings_effect_on_diluted_eps_change"
        and c.period_index == 1
    )
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    trainer2, answer2 = build_training_workbook(data, tmp_path / "ATTR_Eq.xlsx")
    smap2 = load_semantic_map(answer2)
    comp2 = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "earnings_effect_on_diluted_eps_change"
        and c.period_index == 1
    )
    expected = float(comp2.expected_value)
    _inject_formula_and_cached_value(
        trainer2,
        comp2.tab,
        comp2.cell,
        formula=f"={expected}",
        cached_value=expected,
    )
    assert check_workbook(trainer2).correct == 1

    trainer3, answer3 = build_training_workbook(data, tmp_path / "ATTR_Bad.xlsx")
    smap3 = load_semantic_map(answer3)
    comp3 = next(
        c
        for c in smap3.all_ordered()
        if c.family_id == "earnings_effect_on_diluted_eps_change"
        and c.period_index == 1
    )
    _inject_formula_and_cached_value(
        trainer3,
        comp3.tab,
        comp3.cell,
        formula="=1+1",
        cached_value=999.0,
    )
    assert check_workbook(trainer3).incorrect == 1

    # Attribution check-row tamper
    trainer4, answer4 = build_training_workbook(data, tmp_path / "ATTR_Tamper.xlsx")
    smap4 = load_semantic_map(answer4)
    comp4 = next(
        c
        for c in smap4.all_ordered()
        if c.family_id == "earnings_effect_on_diluted_eps_change"
        and c.period_index == 1
    )
    wb = load_workbook(trainer4, data_only=False)
    row, col = parse_cell_ref(comp4.cell)
    wb[comp4.tab].cell(row=row, column=col).value = comp4.formula
    wb[PER_SHARE_SHEET].cell(20, 3).value = '="TAMPERED"'
    wb.save(trainer4)
    wb.close()
    with pytest.raises(
        ValueError, match="Trusted workbook cell was modified: Per Share Analysis"
    ):
        check_workbook(trainer4)
    wb = load_workbook(trainer4, data_only=False)
    assert _fill_rgb(wb[comp4.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

    # Share-row tamper still fails closed
    trainer5, answer5 = build_training_workbook(data, tmp_path / "ATTR_ShareTamper.xlsx")
    smap5 = load_semantic_map(answer5)
    comp5 = next(
        c
        for c in smap5.all_ordered()
        if c.family_id == "reported_net_income_change" and c.period_index == 1
    )
    wb = load_workbook(trainer5, data_only=False)
    row, col = parse_cell_ref(comp5.cell)
    wb[comp5.tab].cell(row=row, column=col).value = comp5.formula
    wb[PER_SHARE_SHEET].cell(7, 2).value = 1
    wb.save(trainer5)
    wb.close()
    with pytest.raises(
        ValueError, match="Trusted workbook cell was modified: Per Share Analysis"
    ):
        check_workbook(trainer5)

    # Live judgment coexistence
    trainer6, answer6 = build_training_workbook(data, tmp_path / "ATTR_Live.xlsx")
    smap6 = load_semantic_map(answer6)
    attr_comp = next(
        c
        for c in smap6.all_ordered()
        if c.family_id == "diluted_eps_change_from_drivers" and c.period_index == 4
    )
    wb = load_workbook(trainer6, data_only=False)
    wb["Accounting Judgment"].cell(5, 6).value = "Financial Liability"
    wb.save(trainer6)
    wb.close()
    _inject_formula_and_cached_value(
        trainer6,
        attr_comp.tab,
        attr_comp.cell,
        formula=attr_comp.formula,
        cached_value=float(attr_comp.expected_value),
    )
    summary = check_workbook(trainer6)
    assert summary.correct == 1
    assert summary.incorrect == 0


def test_expected_requires_per_share_for_attribution(tmp_path):
    data = _share_enabled_demo()
    _, answer = build_training_workbook(data, tmp_path / "ATTR_Req.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c for c in smap.all_ordered() if c.family_id == "reported_net_income_change"
    )
    builder = ReferenceModelBuilder(data)
    with pytest.raises(ValueError, match="requires a PerShareSeries"):
        expected_value_for_component(builder.anchor, comp, per_share=None)

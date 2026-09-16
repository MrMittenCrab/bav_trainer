"""Manual builds must preserve the canonical model and write only local outputs."""

import copy
import json
import os
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from core.__main__ import main
from core.data.interface import (
    HistoricalLeaseData, HistoricalOperatingKpiData,
    HistoricalOperatingKpiObservation, HistoricalShareData,
)
from core.data.standardized_io import standardized_to_payload
from core.tests.test_capex import P1, P2
from core.tests.test_geographic_segment_workbook import _geo_tiny
from core.tests.test_historical_segment import _snapshot
from core.trainer.check_context import load_check_context
from core.trainer.semantic_io import load_semantic_map


@pytest.fixture
def payload():
    fin = _geo_tiny(_snapshot(P1), _snapshot(P2))
    fin.ticker, fin.company_name, fin.jurisdiction = "ACME", "Acme", "US"
    fin.historical_shares = HistoricalShareData(
        scale_basis="financial_statement_units",
        diluted_weighted_average={P1: 10, P2: 11},
    )
    fin.historical_lease = HistoricalLeaseData(lease_interest_expense={P1: 2, P2: 3})
    fin.historical_operating_kpis = HistoricalOperatingKpiData(observations=[
        HistoricalOperatingKpiObservation("store_count", "company_operated", P2, 12, "stores"),
        HistoricalOperatingKpiObservation("store_count", "company_operated", P1, 10, "stores"),
    ])
    return standardized_to_payload(fin)


def _write_input(tmp_path, payload):
    path = tmp_path / "build/input/ACME.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload), encoding="utf-8")
    return path


def test_module_command_preserves_all_modules_and_only_writes_outputs(tmp_path, payload):
    source = _write_input(tmp_path, payload)
    for name in ("TARGET.md", "IMPLEMENTATION.md", "RESULT.md", "benchmark/source.pdf",
                 "benchmark/extracted/source.json", "benchmark/reconciled/standardized.json",
                 "release/answer.xlsx"):
        p = tmp_path / name
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(b"protected")
    before = {p: p.read_bytes() for p in tmp_path.rglob("*") if p.is_file()}
    env = dict(os.environ, PYTHONPATH=str(Path(__file__).resolve().parents[2]))
    result = subprocess.run(
        [sys.executable, "-m", "core", "build", str(source), "-o", "build/output/Acme"],
        cwd=tmp_path, env=env, capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stderr + result.stdout
    for p, data in before.items():
        assert p.read_bytes() == data
    output = tmp_path / "build/output"
    assert {p.name for p in output.iterdir()} == {
        "Acme_Trainer.xlsx", "Acme_Answer_Key.xlsx", "rowmap.json",
        "Acme_Answer_Key.component_map.json", "Acme_Answer_Key.assumptions.json",
    }
    assert all(p in before or p.is_relative_to(output)
               for p in tmp_path.rglob("*") if p.is_file())
    answer = output / "Acme_Answer_Key.xlsx"
    assert load_check_context(answer).source_payload == payload
    wb = load_workbook(answer)
    assert "Geographic Segment Analysis" in wb.sheetnames
    assert "Per Share Analysis" in wb.sheetnames
    lease_row = next(row for row in wb["Income Statement"].iter_rows()
                     if row[0].value == "Lease interest expense (reported note)")
    assert [cell.value for cell in lease_row[1:3]] == [2, 3]
    wb.close()
    assert any(c.family_id == "reported_diluted_eps"
               for c in load_semantic_map(answer).all_ordered())


@pytest.mark.parametrize("keys,value", [
    (("future_module",), {}),
    (("historical_shares", "unknown"), 1),
    (("historical_lease", "unknown"), 1),
    (("historical_segment", "periods", 0, "unknown"), 1),
    (("historical_operating_kpis", "observations", 0, "unknown"), 1),
    (("income_statement", 0, "unknown"), 1),
    (("periods", 0, "end_date"), "2024-12-31junk"),
    (("periods", 0, "is_interim"), "false"),
    (("historical_lease", "lease_interest_expense", "2024-12-31"), True),
    (("income_statement", 0, "values", "2024-12-31"), "120"),
    (("income_statement", 0, "values", "2024-12-31"), 10 ** 400),
    (("balance_sheet",), {}),
    (("historical_operating_kpis", "observations", 0, "metric"), "unsupported"),
])
def test_invalid_model_fails_without_outputs(tmp_path, payload, capsys, keys, value):
    bad = copy.deepcopy(payload)
    target = bad
    for key in keys[:-1]:
        target = target[key]
    target[keys[-1]] = value
    source = _write_input(tmp_path, bad)
    output = tmp_path / "build/output/Acme"
    assert main(["build", str(source), "-o", str(output)]) != 0
    assert "error:" in capsys.readouterr().err
    assert not output.parent.exists()


@pytest.mark.parametrize("raw", ["{", "[]", "{}", '{"ticker":"A","ticker":"B"}',
                                   '{"ticker": NaN}', '{"ticker": Infinity}'])
def test_invalid_json_fails_clearly(tmp_path, capsys, raw):
    source = _write_input(tmp_path, {})
    source.write_text(raw)
    assert main(["build", str(source), "-o", str(tmp_path / "build/output/Acme")]) != 0
    assert "error:" in capsys.readouterr().err
    assert not (tmp_path / "build/output").exists()


@pytest.mark.parametrize("destination", ["benchmark/test", "release/test"])
def test_protected_output_rejected(tmp_path, payload, capsys, monkeypatch, destination):
    monkeypatch.chdir(tmp_path)
    source = _write_input(tmp_path, payload)
    assert main(["build", str(source), "-o", destination]) != 0
    assert "error:" in capsys.readouterr().err
    assert not (tmp_path / destination).parent.exists()


def test_optional_modules_may_be_absent_and_assumptions_still_work(tmp_path, payload):
    for key in ("historical_shares", "historical_lease", "historical_segment",
                "historical_operating_kpis"):
        payload.pop(key)
    source = _write_input(tmp_path, payload)
    assumptions = source.with_name("assumptions.json")
    assumptions.write_text('{"includeDeferredForecast": false}')
    output = tmp_path / "build/output/Acme_Trainer.xlsx"
    assert main(["build", str(source), "-a", str(assumptions), "-o", str(output)]) == 0
    assert output.exists()
    answer = output.with_name("Acme_Answer_Key.xlsx")
    context = load_check_context(answer)
    assert context.source_payload["historical_shares"] is None
    assert "historical_operating_kpis" not in context.source_payload
    wb = load_workbook(answer)
    assert "Per Share Analysis" not in wb.sheetnames
    assert "Geographic Segment Analysis" not in wb.sheetnames
    wb.close()


def test_output_symlink_cannot_overwrite_input(tmp_path, payload, capsys):
    source = _write_input(tmp_path, payload)
    before = source.read_bytes()
    output = tmp_path / "build/output"
    output.mkdir()
    (output / "rowmap.json").symlink_to(source)
    assert main(["build", str(source), "-o", str(output / "Acme")]) != 0
    assert "error:" in capsys.readouterr().err
    assert source.read_bytes() == before
    assert not (output / "Acme_Answer_Key.xlsx").exists()


def test_excel_input_still_builds(tmp_path, payload):
    from openpyxl import Workbook

    source = tmp_path / "Acme.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    for name, key in (("Income Statement", "income_statement"),
                      ("Balance Sheet", "balance_sheet"),
                      ("Cash Flow", "cash_flow")):
        ws = wb.create_sheet(name)
        ws.append(["Concept", "Label", "2024-12-31", "2025-12-31"])
        for item in payload[key]:
            ws.append([item["concept"], item["label"],
                       item["values"].get("2024-12-31"), item["values"].get("2025-12-31")])
    wb.save(source)
    assert main(["build", str(source), "-o", str(tmp_path / "build/output/Acme")]) == 0
    assert (tmp_path / "build/output/Acme_Trainer.xlsx").exists()


def test_missing_input_fails_clearly(tmp_path, capsys):
    assert main(["build", str(tmp_path / "missing.json"),
                 "-o", str(tmp_path / "build/output/Acme")]) != 0
    assert "error:" in capsys.readouterr().err
    assert not (tmp_path / "build/output").exists()

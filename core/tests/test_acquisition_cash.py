"""Step 9M.2.4.1.1.1.23 — source-supported acquisition cash and cash-use diagnostics."""

from __future__ import annotations

import copy
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    ACQUISITION_CASH_COMPONENT_CATALOG,
    expand_acquisition_cash_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.acquisition_cash import (
    acquisition_cash_applicable,
    acquisition_cash_availability,
    cash_after_ppe_capex_and_acquisitions_applicable,
    compute_acquisition_cash_series,
    resolve_acquisition_cash_source,
)
from core.model.financial_math import compute_anchor
from core.model.historical_expected import acquisition_cash_expected_series
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_capex import P1, P2, _anchor, _dupont_row_by_label, _tiny
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "benchmark" / "fast_retailing" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _add_acq(
    fin: StandardizedFinancials,
    *,
    values=(-10.0, -20.0),
    label="Acquisition, net of cash acquired",
    concept="acquisition_net_of_cash_acquired",
    duplicate: bool = False,
    missing_period: bool = False,
    none_period: bool = False,
    on_balance_sheet: bool = False,
) -> LineItem:
    if missing_period:
        pay_values = {P1: values[0]}
    elif none_period:
        pay_values = {P1: values[0], P2: None}
    else:
        pay_values = {P1: values[0], P2: values[1]}
    item = _li(label, pay_values, concept=concept)
    target = fin.balance_sheet if on_balance_sheet else fin.cash_flow
    target.append(item)
    if duplicate and not on_balance_sheet:
        fin.cash_flow.append(
            _li("Acquisition duplicate", pay_values, concept=concept)
        )
    return item


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(ACQUISITION_CASH_COMPONENT_CATALOG) == 3
    assert [f.order for f in ACQUISITION_CASH_COMPONENT_CATALOG] == [129, 130, 131]
    specs = expand_acquisition_cash_specs(periods, start_order=1000)
    assert len(specs) == 10
    assert {s.family_id for s in specs} == {
        "acquisition_cash_outflow",
        "acquisition_cash_to_revenue",
    }
    assert all(s.semantic_key.startswith("acquisition_cash.") for s in specs)
    residual = expand_acquisition_cash_specs(
        periods, start_order=1000, include_cash_after_capex=True
    )
    assert len(residual) == 15
    assert {s.family_id for s in residual} == {
        "acquisition_cash_outflow",
        "acquisition_cash_to_revenue",
        "cash_after_ppe_capex_and_acquisitions",
    }
    with pytest.raises(ValueError, match="duplicate"):
        expand_acquisition_cash_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_acquisition_cash_specs(list(reversed(periods)), start_order=1)


def test_unique_concept_resolution_and_renamed_label():
    fin = _tiny()
    stored = _add_acq(fin)
    item = resolve_acquisition_cash_source(fin)
    assert item is stored
    assert item.concept == "acquisition_net_of_cash_acquired"
    assert item.label == "Acquisition, net of cash acquired"
    assert acquisition_cash_applicable(fin)
    avail = acquisition_cash_availability(fin)
    assert avail.acquisition_net_of_cash_acquired is True
    assert avail.ambiguous is False
    assert cash_after_ppe_capex_and_acquisitions_applicable(fin)

    renamed = _tiny()
    _add_acq(renamed, label="Business combination, net of cash")
    assert resolve_acquisition_cash_source(renamed).label == (
        "Business combination, net of cash"
    )


def test_absent_label_only_duplicate_wrong_statement():
    absent = _tiny()
    assert not acquisition_cash_applicable(absent)
    assert resolve_acquisition_cash_source(absent) is None
    assert acquisition_cash_availability(absent).acquisition_net_of_cash_acquired is False
    with pytest.raises(MissingLineError):
        compute_acquisition_cash_series(
            absent, list(canonical_fiscal_periods(absent)), _anchor(absent)
        )
    assert ReferenceModelBuilder(absent).acquisition_cash_specs == ()

    label_only = _tiny()
    _add_acq(label_only, concept="")
    assert not acquisition_cash_applicable(label_only)
    assert resolve_acquisition_cash_source(label_only) is None
    assert (
        resolve_line(
            label_only.cash_flow, "acquisition_net_of_cash_acquired", required=False
        ).item
        is None
    )
    assert ReferenceModelBuilder(label_only).acquisition_cash_specs == ()

    dup = _tiny()
    _add_acq(dup, duplicate=True)
    avail = acquisition_cash_availability(dup)
    assert avail.ambiguous is True
    assert avail.acquisition_net_of_cash_acquired is False
    assert not acquisition_cash_applicable(dup)
    assert resolve_acquisition_cash_source(dup) is None
    with pytest.raises(AmbiguousLineError):
        resolve_line(
            dup.cash_flow, "acquisition_net_of_cash_acquired", required=False
        )
    assert ReferenceModelBuilder(dup).acquisition_cash_specs == ()

    wrong = _tiny()
    _add_acq(wrong, on_balance_sheet=True)
    assert not acquisition_cash_applicable(wrong)
    assert resolve_acquisition_cash_source(wrong) is None
    assert (
        resolve_line(
            wrong.balance_sheet, "acquisition_net_of_cash_acquired", required=False
        ).item
        is not None
    )


def test_no_substitution_from_goodwill_intangibles_securities_or_rou():
    fin = _tiny()
    fin.balance_sheet.append(
        _li("Goodwill", {P1: 80.0, P2: 90.0}, concept="goodwill")
    )
    fin.cash_flow.append(
        _li(
            "Payments for intangible assets",
            {P1: -5.0, P2: -6.0},
            concept="payments_for_intangible_assets",
        )
    )
    fin.cash_flow.append(
        _li(
            "Payments for acquisition of right-of-use assets",
            {P1: -7.0, P2: -8.0},
            concept="payments_for_rou_assets",
        )
    )
    fin.cash_flow.append(
        _li(
            "Purchases of marketable securities",
            {P1: -15.0, P2: -16.0},
            concept="purchases_of_marketable_securities",
        )
    )
    assert resolve_acquisition_cash_source(fin) is None
    assert not acquisition_cash_applicable(fin)
    assert acquisition_cash_availability(fin).acquisition_net_of_cash_acquired is False


def test_sign_conversion_ratios_positive_negative_mixed_zero():
    fin = _tiny()
    _add_acq(fin, values=(-50.0, -80.0))
    series = compute_acquisition_cash_series(
        fin, list(canonical_fiscal_periods(fin)), _anchor(fin)
    )
    assert series.payments_reported == (-50.0, -80.0)
    assert series.acquisition_cash_outflow == (50.0, 80.0)
    assert series.acquisition_cash_to_revenue[0] == pytest.approx(50.0 / 1000.0)
    assert series.cash_after_ppe_capex_and_acquisitions == (
        80.0 - 100.0 - 50.0,
        90.0 - 120.0 - 80.0,
    )

    inflow = _tiny()
    _add_acq(inflow, values=(50.0, 80.0))
    series_in = compute_acquisition_cash_series(
        inflow, list(canonical_fiscal_periods(inflow)), _anchor(inflow)
    )
    assert series_in.payments_reported == (50.0, 80.0)
    assert series_in.acquisition_cash_outflow == (-50.0, -80.0)

    zero = _tiny()
    _add_acq(zero, values=(0.0, -10.0))
    series_z = compute_acquisition_cash_series(
        zero, list(canonical_fiscal_periods(zero)), _anchor(zero)
    )
    assert series_z.payments_reported == (0.0, -10.0)
    assert series_z.acquisition_cash_outflow == (0.0, 10.0)
    assert series_z.acquisition_cash_to_revenue[0] == pytest.approx(0.0)

    mapped = acquisition_cash_expected_series(series)
    assert set(mapped) == {f.id for f in ACQUISITION_CASH_COMPONENT_CATALOG}


def test_zero_revenue_undefined_ratio():
    fin = _tiny(revenue=(0.0, 2000.0))
    _add_acq(fin, values=(-50.0, -80.0))
    series = compute_acquisition_cash_series(
        fin, list(canonical_fiscal_periods(fin)), _anchor(fin)
    )
    assert series.acquisition_cash_outflow == (50.0, 80.0)
    assert series.acquisition_cash_to_revenue[0] == UNDEFINED_RATIO
    assert series.acquisition_cash_to_revenue[1] == pytest.approx(80.0 / 2000.0)


def test_missing_none_versus_reported_zero():
    missing = _tiny()
    _add_acq(missing, missing_period=True)
    periods = list(canonical_fiscal_periods(missing))
    with pytest.raises(MissingHistoricalValueError, match="acquisition_net_of_cash_acquired"):
        compute_acquisition_cash_series(missing, periods, _anchor(missing))

    none_period = _tiny()
    _add_acq(none_period, none_period=True)
    periods_n = list(canonical_fiscal_periods(none_period))
    with pytest.raises(MissingHistoricalValueError, match="acquisition_net_of_cash_acquired"):
        compute_acquisition_cash_series(none_period, periods_n, _anchor(none_period))

    zero = _tiny()
    stored = _add_acq(zero, values=(0.0, -10.0))
    series = compute_acquisition_cash_series(
        zero, list(canonical_fiscal_periods(zero)), _anchor(zero)
    )
    assert series.payments_reported[0] == 0.0
    assert stored.values[P1] == 0.0


def test_absent_ambiguous_cfo_and_capex_omit_only_residual():
    no_cfo = _tiny(with_cfo=False)
    _add_acq(no_cfo)
    assert acquisition_cash_applicable(no_cfo)
    assert not cash_after_ppe_capex_and_acquisitions_applicable(no_cfo)
    series = compute_acquisition_cash_series(
        no_cfo, list(canonical_fiscal_periods(no_cfo)), _anchor(no_cfo)
    )
    assert series.acquisition_cash_outflow == (10.0, 20.0)
    assert series.cash_after_ppe_capex_and_acquisitions is None
    mapped = acquisition_cash_expected_series(series)
    assert set(mapped) == {"acquisition_cash_outflow", "acquisition_cash_to_revenue"}
    builder = ReferenceModelBuilder(no_cfo)
    assert {s.family_id for s in builder.acquisition_cash_specs} == {
        "acquisition_cash_outflow",
        "acquisition_cash_to_revenue",
    }
    assert len(builder.acquisition_cash_specs) == 4

    no_capex = _tiny(with_payments=False)
    _add_acq(no_capex)
    assert acquisition_cash_applicable(no_capex)
    assert not cash_after_ppe_capex_and_acquisitions_applicable(no_capex)
    series_nc = compute_acquisition_cash_series(
        no_capex, list(canonical_fiscal_periods(no_capex)), _anchor(no_capex)
    )
    assert series_nc.cash_after_ppe_capex_and_acquisitions is None
    assert len(ReferenceModelBuilder(no_capex).acquisition_cash_specs) == 4

    amb_cfo = _tiny(duplicate_cfo=True, cfo_concept="operating_cash_flow")
    _add_acq(amb_cfo)
    avail = acquisition_cash_availability(amb_cfo)
    assert avail.operating_cash_flow_ambiguous is True
    assert acquisition_cash_applicable(amb_cfo)
    assert not cash_after_ppe_capex_and_acquisitions_applicable(amb_cfo)
    series_a = compute_acquisition_cash_series(
        amb_cfo, list(canonical_fiscal_periods(amb_cfo)), _anchor(amb_cfo)
    )
    assert series_a.cash_after_ppe_capex_and_acquisitions is None

    amb_capex = _tiny(duplicate=True)
    _add_acq(amb_capex)
    avail_c = acquisition_cash_availability(amb_capex)
    assert avail_c.ppe_capex_ambiguous is True
    assert acquisition_cash_applicable(amb_capex)
    assert not cash_after_ppe_capex_and_acquisitions_applicable(amb_capex)


def test_independent_gating_without_goodwill():
    fin = _tiny()
    _add_acq(fin)
    assert "goodwill" not in {item.concept for item in fin.balance_sheet}
    assert acquisition_cash_applicable(fin)
    builder = ReferenceModelBuilder(fin)
    assert len(builder.acquisition_cash_specs) == 6
    assert builder.goodwill_intangibles_specs == ()


def test_period_ordering_and_unchanged_input():
    fin = _tiny()
    stored = _add_acq(fin, values=(-10.0, -20.0))
    reverse = [P2, P1]
    series = compute_acquisition_cash_series(
        fin, reverse, compute_anchor(fin, reverse)
    )
    assert series.payments_reported == (-20.0, -10.0)
    assert series.acquisition_cash_outflow == (20.0, 10.0)

    before = copy.deepcopy(stored.values)
    compute_acquisition_cash_series(fin, [P1, P2], _anchor(fin))
    assert stored.values == before
    assert stored.concept == "acquisition_net_of_cash_acquired"


def test_workbook_gating_structure_formulas_notes_and_check(tmp_path):
    data = _tiny()
    _add_acq(data)
    trainer, answer = build_training_workbook(data, tmp_path / "ACQ_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.acquisition_cash_series is not None
    assert len(builder.acquisition_cash_specs) == 6
    assert {s.family_id for s in builder.acquisition_cash_specs} == {
        "acquisition_cash_outflow",
        "acquisition_cash_to_revenue",
        "cash_after_ppe_capex_and_acquisitions",
    }

    smap = load_semantic_map(answer)
    acq_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in ACQUISITION_CASH_COMPONENT_CATALOG}
    ]
    assert len(acq_comps) == 6

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "ACQUISITION CASH CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        reported_row = _dupont_row_by_label(
            ws, "Acquisition, net of cash acquired (reported)"
        )
        outflow_row = _dupont_row_by_label(ws, "Acquisition cash outflow (−reported)")
        ratio_row = _dupont_row_by_label(ws, "Acquisition cash / Revenue")
        residual_row = _dupont_row_by_label(
            ws, "Operating cash after PP&E capex and acquisitions"
        )
        assert isinstance(ws.cell(reported_row, 2).value, str)
        assert ws.cell(reported_row, 2).value.startswith("=")
        for col in (2, 3):
            outflow_cell = ws.cell(outflow_row, column=col)
            ratio_cell = ws.cell(ratio_row, column=col)
            residual_cell = ws.cell(residual_row, column=col)
            if path == answer:
                assert isinstance(outflow_cell.value, str) and outflow_cell.value.startswith(
                    "="
                )
                assert outflow_cell.value.replace(" ", "").startswith("=-")
                assert "IF(" in str(ratio_cell.value) and "NA()" in str(ratio_cell.value)
                assert residual_cell.value.replace(" ", "").count("-") >= 2
                note = (outflow_cell.comment.text or "") if outflow_cell.comment else ""
                assert "net of cash acquired" in note.lower()
                assert "free cash flow" in note.lower()
                assert "purchase-price" in note.lower() or "purchase price" in note.lower()
                assert "goodwill" in note.lower()
                assert residual_cell.comment is not None
                residual_note = residual_cell.comment.text or ""
                assert "mechanical" in residual_note.lower()
            else:
                assert outflow_cell.value is None
                assert ratio_cell.value is None
                assert residual_cell.value is None
                assert outflow_cell.comment is None
                assert ratio_cell.comment is None
                assert residual_cell.comment is None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in acq_comps if c.family_id == "acquisition_cash_outflow")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Acquisition cash outflow" not in dumped


def test_source_row_fidelity_round_trip_and_immutability(tmp_path):
    fin = _tiny()
    stored = _add_acq(fin, values=(-10.0, -20.0))
    original_index = fin.cash_flow.index(stored)
    resolved = resolve_line(
        fin.cash_flow, "acquisition_net_of_cash_acquired", required=True
    )
    assert resolved.index == original_index
    assert resolved.item is stored

    restored = standardized_from_payload(standardized_to_payload(fin))
    rt = resolve_acquisition_cash_source(restored)
    assert rt is not None
    assert rt.concept == "acquisition_net_of_cash_acquired"
    assert rt.values == stored.values

    trainer, answer = build_training_workbook(fin, tmp_path / "ACQ_ROW.xlsx")
    wb = load_workbook(answer, data_only=False)
    ws = wb["ALT DuPont"]
    reported_row = _dupont_row_by_label(
        ws, "Acquisition, net of cash acquired (reported)"
    )
    level_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
    expected_row = 7 + original_index
    assert f"'CashFlowStatement'!B{expected_row}" in level_f or (
        f"'Cash Flow Statement'!B{expected_row}" in level_f
    )
    wb.close()
    assert stored.concept == "acquisition_net_of_cash_acquired"
    assert stored.values == {P1: -10.0, P2: -20.0}
    assert trainer.exists() and answer.exists()


def test_demo_and_fast_retailing_omit_acquisition_cash(tmp_path):
    demo = HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )
    assert not acquisition_cash_applicable(demo)
    builder = ReferenceModelBuilder(demo)
    assert builder.acquisition_cash_series is None
    assert builder.acquisition_cash_specs == ()
    trainer, answer = build_training_workbook(demo, tmp_path / "ACQ_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312

    payload = json.loads(FR_JSON.read_text(encoding="utf-8"))
    fr = standardized_from_payload(payload)
    assert not acquisition_cash_applicable(fr)
    fr_builder = ReferenceModelBuilder(fr)
    assert fr_builder.acquisition_cash_specs == ()
    assert len(fr_builder.expected_specs) == 501
    assert "acquisition_cash_outflow" not in {
        s.family_id for s in fr_builder.expected_specs
    }

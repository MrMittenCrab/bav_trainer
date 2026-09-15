"""Step 9M.6 — lease ROU-asset intensity / trend diagnostics."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    LEASE_ROU_COMPONENT_CATALOG,
    expand_lease_rou_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.historical_expected import lease_rou_expected_series
from core.model.lease_liability import lease_liability_applicable
from core.model.lease_rou import (
    compute_lease_rou_series,
    lease_rou_applicable,
    lease_rou_availability,
    resolve_lease_rou_source,
)
from core.model.line_resolver import AmbiguousLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"


def _ingest_demo():
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _tiny(
    *,
    rou=(100.0, 120.0),
    revenue=(1000.0, 1100.0),
    with_rou: bool = True,
    label_only: bool = False,
    duplicate: bool = False,
    missing_period: bool = False,
    none_period: bool = False,
    with_lease_liability: bool = False,
    rou_first: bool = False,
    single_period: bool = False,
    rou_concept: str = "right_of_use_assets",
    rou_label: str = "Right-of-use assets",
):
    if single_period:
        d1 = date(2025, 12, 31)
        periods = [FinancialPeriod(end_date=d1, label="FY2025")]

        def vals(a, b=None):
            return {d1: a}

        rev = (revenue[0],)
        rou_vals = (rou[0],)
    else:
        d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
        periods = [
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ]

        def vals(a, b=None):
            return {d1: a, d2: b if b is not None else a}

        rev = revenue
        rou_vals = rou

    cash, ar, ap, bank = 50.0, 40.0, 30.0, 20.0
    rou0 = rou_vals[0] if with_rou else 0.0
    rou1 = rou_vals[-1] if with_rou else 0.0
    lease0 = lease1 = 0.0
    if with_lease_liability:
        lease0, lease1 = 80.0, 90.0
    eq0 = cash + ar + rou0 - ap - bank - lease0
    eq1 = cash + ar + rou1 - ap - bank - lease1

    rou_item = None
    if with_rou:
        concept = "" if label_only else rou_concept
        if missing_period and not single_period:
            rou_values = {date(2024, 12, 31): rou_vals[0]}
        elif none_period and not single_period:
            rou_values = {date(2024, 12, 31): rou_vals[0], date(2025, 12, 31): None}
        elif single_period:
            rou_values = vals(rou_vals[0])
        else:
            rou_values = vals(*rou_vals)
        rou_item = _li(rou_label, rou_values, concept=concept)

    bs = [
        _li("Cash and cash equivalents", vals(cash, cash)),
        _li("Trade receivables", vals(ar, ar)),
        _li("Trade payables", vals(ap, ap)),
        _li("Bank borrowings", vals(bank, bank)),
    ]
    if rou_item is not None:
        if rou_first:
            bs.insert(0, rou_item)
        else:
            bs.insert(2, rou_item)
        if duplicate:
            bs.insert(
                3,
                _li(
                    "Right-of-use assets duplicate",
                    vals(*rou_vals) if not single_period else vals(rou_vals[0]),
                    concept=rou_concept,
                ),
            )
    if with_lease_liability:
        bs.append(
            _li(
                "Operating lease liabilities",
                vals(lease0, lease1),
                concept="lease_liability",
            )
        )
    bs.append(_li("Total equity", vals(eq0, eq1)))

    is_rows = [
        _li("Revenue", vals(*rev) if not single_period else vals(rev[0])),
        _li("Finance costs", vals(-10, -11)),
        _li("Finance income", vals(0, 0)),
        _li("Profit before tax", vals(200, 220)),
        _li("Income tax expense", vals(-30, -33)),
        _li("Profit for the year", vals(170, 187)),
    ]
    if single_period:
        is_rows = [
            _li("Revenue", vals(rev[0])),
            _li("Finance costs", vals(-10)),
            _li("Finance income", vals(0)),
            _li("Profit before tax", vals(200)),
            _li("Income tax expense", vals(-30)),
            _li("Profit for the year", vals(170)),
        ]

    return StandardizedFinancials(
        ticker="ROU",
        company_name="ROU Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=is_rows,
        balance_sheet=bs,
        cash_flow=[
            _li(
                "Net cash from operating activities",
                vals(80, 90) if not single_period else vals(80),
            )
        ],
    )


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(LEASE_ROU_COMPONENT_CATALOG) == 4
    assert [f.order for f in LEASE_ROU_COMPONENT_CATALOG] == list(range(112, 116))
    specs = expand_lease_rou_specs(periods, start_order=1000)
    assert len(specs) == 16
    assert {s.family_order for s in specs} == set(range(112, 116))
    with pytest.raises(ValueError, match="duplicate"):
        expand_lease_rou_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_lease_rou_specs(list(reversed(periods)), start_order=1)


def test_ordinary_two_period_math():
    fin = _tiny()
    periods = list(canonical_fiscal_periods(fin))
    anchor = compute_anchor(fin, periods)
    assert lease_rou_applicable(fin)
    series = compute_lease_rou_series(fin, periods, anchor)
    assert series.rou_assets == (100.0, 120.0)
    assert series.rou_assets_change == (None, 20.0)
    assert series.rou_assets_growth[0] is None
    assert series.rou_assets_growth[1] == pytest.approx(0.20)
    assert series.average_rou_assets[0] is None
    assert series.average_rou_assets[1] == pytest.approx(110.0)
    assert series.rou_assets_to_revenue[0] is None
    assert series.rou_assets_to_revenue[1] == pytest.approx(110.0 / 1100.0)
    mapped = lease_rou_expected_series(series)
    assert set(mapped) == {f.id for f in LEASE_ROU_COMPONENT_CATALOG}


def test_zero_prior_zero_revenue_and_one_period():
    fin = _tiny(rou=(0.0, 10.0), revenue=(1000.0, 0.0))
    periods = list(canonical_fiscal_periods(fin))
    series = compute_lease_rou_series(fin, periods, compute_anchor(fin, periods))
    assert series.rou_assets_growth[1] == UNDEFINED_RATIO
    assert series.rou_assets_to_revenue[1] == UNDEFINED_RATIO

    one = _tiny(single_period=True)
    periods = list(canonical_fiscal_periods(one))
    series_one = compute_lease_rou_series(one, periods, compute_anchor(one, periods))
    assert series_one.rou_assets == (100.0,)
    assert series_one.rou_assets_change == (None,)
    assert series_one.average_rou_assets == (None,)
    builder = ReferenceModelBuilder(one)
    assert builder.lease_rou_specs == ()


def test_absent_label_only_duplicate_incomplete():
    absent = _tiny(with_rou=False)
    assert not lease_rou_applicable(absent)
    assert lease_rou_availability(absent).right_of_use_assets is False
    assert resolve_lease_rou_source(absent) is None
    assert ReferenceModelBuilder(absent).lease_rou_specs == ()

    label_only = _tiny(label_only=True)
    assert not lease_rou_applicable(label_only)
    assert resolve_lease_rou_source(label_only) is None
    assert ReferenceModelBuilder(label_only).lease_rou_specs == ()

    dup = _tiny(duplicate=True)
    avail = lease_rou_availability(dup)
    assert avail.ambiguous is True
    assert avail.right_of_use_assets is False
    assert not lease_rou_applicable(dup)
    assert resolve_lease_rou_source(dup) is None

    missing = _tiny(missing_period=True)
    periods = list(canonical_fiscal_periods(missing))
    with pytest.raises(MissingHistoricalValueError):
        compute_lease_rou_series(missing, periods, compute_anchor(missing, periods))


def test_independent_of_lease_liability():
    without_liab = _tiny(with_lease_liability=False)
    assert lease_rou_applicable(without_liab)
    assert not lease_liability_applicable(without_liab)
    builder = ReferenceModelBuilder(without_liab)
    assert builder.lease_rou_series is not None
    assert len(builder.lease_rou_specs) == 4
    assert builder.lease_liability_specs == ()

    with_both = _tiny(with_lease_liability=True)
    assert lease_rou_applicable(with_both)
    assert lease_liability_applicable(with_both)
    builder_both = ReferenceModelBuilder(with_both)
    assert builder_both.lease_rou_specs
    assert builder_both.lease_liability_specs


def test_semantic_mapping_survives_source_row_reordering(tmp_path):
    for rou_first in (False, True):
        fin = _tiny(rou_first=rou_first)
        trainer, answer = build_training_workbook(
            fin, tmp_path / f"ROU_REORDER_{rou_first}.xlsx"
        )
        idx = next(
            i
            for i, item in enumerate(fin.balance_sheet)
            if (item.concept or "") == "right_of_use_assets"
        )
        expected_row = 7 + idx
        smap = load_semantic_map(answer)
        change = next(c for c in smap.all_ordered() if c.family_id == "rou_assets_change")
        # Balance is populated; change formula references the level row which links
        # to the Balance Sheet source row via the level formula on ALT DuPont.
        wb = load_workbook(answer, data_only=False)
        ws = wb["ALT DuPont"]
        level_row = next(
            r
            for r in range(1, (ws.max_row or 1) + 1)
            if ws.cell(r, 1).value == "Right-of-use Assets"
        )
        level_f = str(ws.cell(level_row, 2).value).replace(" ", "")
        assert f"'BalanceSheet'!B{expected_row}" in level_f or (
            f"'Balance Sheet'!B{expected_row}" in level_f
        )
        assert change.period_index == 1
        wb.close()


def _dupont_row_by_label(ws, label: str) -> int:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"ALT DuPont label not found: {label!r}")


def test_workbook_first_period_blank_and_check(tmp_path):
    data = _tiny()
    trainer, answer = build_training_workbook(data, tmp_path / "ROU_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.lease_rou_series is not None
    assert len(builder.lease_rou_specs) == 4
    assert {s.family_id for s in builder.lease_rou_specs} == {
        f.id for f in LEASE_ROU_COMPONENT_CATALOG
    }

    smap = load_semantic_map(answer)
    practice_cells = {(c.tab, c.cell) for c in smap.all_ordered()}
    rou_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in LEASE_ROU_COMPONENT_CATALOG}
    ]
    assert len(rou_comps) == 4

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "LEASE ROU-ASSET CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        level_row = _dupont_row_by_label(ws, "Right-of-use Assets")
        diag_rows = (
            _dupont_row_by_label(ws, "Change in Right-of-use Assets"),
            _dupont_row_by_label(ws, "Right-of-use Assets Growth"),
            _dupont_row_by_label(ws, "Average Right-of-use Assets"),
            _dupont_row_by_label(ws, "Average Right-of-use Assets / Revenue"),
        )
        assert ws.cell(row=level_row, column=2).value is not None
        for row in diag_rows:
            cell = ws.cell(row=row, column=2)
            assert cell.value is None
            assert ("ALT DuPont", cell.coordinate) not in practice_cells
            later = ws.cell(row=row, column=3)
            if path == answer:
                assert isinstance(later.value, str) and later.value.startswith("=")
                assert later.comment is not None
                assert (later.comment.text or "").strip()
            else:
                assert later.value is None
                assert later.comment is None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in rou_comps if c.family_id == "rou_assets_change")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Change in Right-of-use Assets" not in dumped


def test_demo_omits_and_practice_counts_unchanged(tmp_path):
    data = _ingest_demo()
    assert not lease_rou_applicable(data)
    builder = ReferenceModelBuilder(data)
    assert builder.lease_rou_series is None
    assert builder.lease_rou_specs == ()

    trainer, answer = build_training_workbook(data, tmp_path / "ROU_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312
    assert check_workbook(trainer).blank == 312


def _alias_kwargs():
    return {
        "rou_concept": "right_of_use_lease_asset",
        "rou_label": "Right-of-use lease assets",
    }


def test_right_of_use_lease_asset_alias_resolution_zero_missing_and_none():
    alias = _tiny(**_alias_kwargs())
    item = resolve_lease_rou_source(alias)
    assert item is not None
    assert item.concept == "right_of_use_lease_asset"
    assert item.label == "Right-of-use lease assets"
    assert lease_rou_applicable(alias)
    avail = lease_rou_availability(alias)
    assert avail.right_of_use_assets is True and avail.ambiguous is False

    series = compute_lease_rou_series(
        alias, list(canonical_fiscal_periods(alias)), compute_anchor(alias, list(canonical_fiscal_periods(alias)))
    )
    assert series.rou_assets == (100.0, 120.0)
    assert series.rou_assets_change == (None, 20.0)
    assert series.average_rou_assets[1] == pytest.approx(110.0)

    canonical = _tiny()
    canon_series = compute_lease_rou_series(
        canonical,
        list(canonical_fiscal_periods(canonical)),
        compute_anchor(canonical, list(canonical_fiscal_periods(canonical))),
    )
    assert series.rou_assets == canon_series.rou_assets
    assert series.rou_assets_change == canon_series.rou_assets_change

    zero = _tiny(rou=(0.0, 10.0), **_alias_kwargs())
    series_z = compute_lease_rou_series(
        zero, list(canonical_fiscal_periods(zero)), compute_anchor(zero, list(canonical_fiscal_periods(zero)))
    )
    assert series_z.rou_assets == (0.0, 10.0)
    assert series_z.rou_assets_change == (None, 10.0)

    missing = _tiny(missing_period=True, **_alias_kwargs())
    with pytest.raises(MissingHistoricalValueError):
        compute_lease_rou_series(
            missing,
            list(canonical_fiscal_periods(missing)),
            compute_anchor(missing, list(canonical_fiscal_periods(missing))),
        )

    none_period = _tiny(none_period=True, **_alias_kwargs())
    with pytest.raises(MissingHistoricalValueError):
        compute_lease_rou_series(
            none_period,
            list(canonical_fiscal_periods(none_period)),
            compute_anchor(none_period, list(canonical_fiscal_periods(none_period))),
        )


def test_right_of_use_lease_asset_alias_ambiguity_and_label_only():
    both = _tiny()
    extra_vals = {date(2024, 12, 31): 100.0, date(2025, 12, 31): 120.0}
    both.balance_sheet.insert(
        3,
        _li(
            "Right-of-use lease assets",
            extra_vals,
            concept="right_of_use_lease_asset",
        ),
    )
    equity = next(item for item in both.balance_sheet if item.label == "Total equity")
    for period, amount in extra_vals.items():
        equity.values[period] = float(equity.values[period]) + amount
    avail = lease_rou_availability(both)
    assert avail.ambiguous is True
    assert avail.right_of_use_assets is False
    assert not lease_rou_applicable(both)
    assert resolve_lease_rou_source(both) is None
    with pytest.raises(AmbiguousLineError):
        resolve_line(both.balance_sheet, "right_of_use_assets", required=False)
    assert ReferenceModelBuilder(both).lease_rou_specs == ()

    dup_alias = _tiny(duplicate=True, **_alias_kwargs())
    assert lease_rou_availability(dup_alias).ambiguous is True
    assert not lease_rou_applicable(dup_alias)

    label_only = _tiny(label_only=True, rou_label="Right-of-use lease assets")
    assert not lease_rou_applicable(label_only)
    assert resolve_lease_rou_source(label_only) is None

    absent = _tiny(with_rou=False)
    assert not lease_rou_applicable(absent)
    assert resolve_lease_rou_source(absent) is None


def test_right_of_use_lease_asset_alias_round_trip_preserves_stored_identity():
    fin = _tiny(**_alias_kwargs())
    original = resolve_lease_rou_source(fin)
    assert original is not None
    restored = standardized_from_payload(standardized_to_payload(fin))
    item = resolve_lease_rou_source(restored)
    assert item is not None
    assert item.concept == "right_of_use_lease_asset"
    assert item.label == original.label
    assert item.values == original.values
    assert lease_rou_applicable(restored)
    series = compute_lease_rou_series(
        restored,
        list(canonical_fiscal_periods(restored)),
        compute_anchor(restored, list(canonical_fiscal_periods(restored))),
    )
    assert series.rou_assets == (100.0, 120.0)


def test_right_of_use_lease_asset_alias_python_excel_source_identity(tmp_path):
    for rou_first in (False, True):
        fin = _tiny(rou_first=rou_first, **_alias_kwargs())
        resolved = resolve_line(fin.balance_sheet, "right_of_use_assets", required=True)
        assert resolved.item is not None
        assert resolved.item.concept == "right_of_use_lease_asset"
        expected_row = 7 + resolved.index
        python_levels = compute_lease_rou_series(
            fin,
            list(canonical_fiscal_periods(fin)),
            compute_anchor(fin, list(canonical_fiscal_periods(fin))),
        ).rou_assets
        assert python_levels == (100.0, 120.0)

        trainer, answer = build_training_workbook(
            fin, tmp_path / f"ROU_ALIAS_{rou_first}.xlsx"
        )
        builder = ReferenceModelBuilder(fin)
        assert builder.lease_rou_series is not None
        assert builder.lease_rou_series.rou_assets == python_levels
        assert len(builder.lease_rou_specs) == 4

        wb = load_workbook(answer, data_only=False)
        ws = wb["ALT DuPont"]
        level_row = _dupont_row_by_label(ws, "Right-of-use Assets")
        level_f = str(ws.cell(level_row, 2).value).replace(" ", "")
        assert f"'BalanceSheet'!B{expected_row}" in level_f or (
            f"'Balance Sheet'!B{expected_row}" in level_f
        )
        wb.close()

        blank = check_workbook(trainer)
        assert blank.incorrect == 0
        assert blank.blank == blank.total
        assert blank.correct == 0

        smap = load_semantic_map(answer)
        wb = load_workbook(trainer, data_only=False)
        for comp in smap.all_ordered():
            row, col = parse_cell_ref(comp.cell)
            wb[comp.tab].cell(row=row, column=col).value = comp.formula
        wb.save(trainer)
        wb.close()
        filled = check_workbook(trainer)
        assert filled.correct == filled.total
        assert filled.incorrect == 0
        assert filled.blank == 0

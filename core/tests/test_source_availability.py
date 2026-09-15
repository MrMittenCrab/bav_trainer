"""Generic source-availability gating for interest-dependent historical analysis."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import FinancialPeriod, LineItem, StandardizedFinancials
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.reference_model import ReferenceModelBuilder
from core.model.financial_math import compute_anchor
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import SOURCE_UNAVAILABLE, UNDEFINED_RATIO, is_source_unavailable
from core.model.source_availability import (
    REASON_ABSENT_LINE,
    REASON_MISSING_PERIOD_VALUE,
    assess_concept_availability,
    assess_interest_availability,
)
from core.model.source_values import MissingHistoricalValueError
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook
from core.tests.test_normalization import _inject_formula_and_cached_value

ROOT = Path(__file__).resolve().parents[2]
LULU_STD = ROOT / "benchmark" / "lululemon" / "reconciled" / "standardized.json"

P1 = date(2024, 12, 31)
P2 = date(2025, 12, 31)
P3 = date(2026, 12, 31)


def _li(label, v1, v2, concept="", *, v3=None):
    values = {P1: v1, P2: v2}
    if v3 is not None:
        values[P3] = v3
    return LineItem(label=label, values=values, concept=concept)


def _periods(*dates: date) -> list[FinancialPeriod]:
    return [
        FinancialPeriod(end_date=d, label=f"FY{d.year}")
        for d in dates
    ]


def _base_fin(**overrides) -> StandardizedFinancials:
    periods = _periods(P1, P2)
    kwargs = dict(
        ticker="SYN",
        company_name="Synthetic Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -40, -50),
            _li("Finance income", 5, 6),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 80, 90),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 50, 55),
            _li("Bank borrowings", 200, 210),
            _li("Total equity", 330, 355),
        ],
        cash_flow=[_li("Net cash from operating activities", 50, 60)],
    )
    kwargs.update(overrides)
    return StandardizedFinancials(**kwargs)


def _labels(ws) -> dict[str, int]:
    return {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}


def test_each_interest_concept_missing_independently_gates(tmp_path: Path):
    periods = [P1, P2]
    missing_expense = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance income", 5, 6),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    missing_income = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -40, -50),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    for fin, absent in (
        (missing_expense, "interest_expense"),
        (missing_income, "interest_income"),
    ):
        avail = assess_interest_availability(fin, periods)
        assert avail.interest_expense.line_present is (absent != "interest_expense")
        assert avail.interest_income.line_present is (absent != "interest_income")
        for period in periods:
            assert avail.output("net_interest", period).available is False
            assert absent in avail.output("net_interest", period).missing_concepts
            assert avail.output("nopat", period).available is False
        with pytest.raises(MissingLineError, match=absent):
            resolve_line(fin.income_statement, absent, required=True)
        anchor = compute_anchor(fin, periods)
        assert all(is_source_unavailable(v) for v in anchor.historical.net_interest)
        assert all(is_source_unavailable(v) for v in anchor.historical.nopat)
        assert is_source_unavailable(anchor.hist_avg_after_tax_cod)
        assert anchor.historical.revenue == [1000.0, 1100.0]
        assert anchor.historical.net_income == [170.0, 187.0]
        builder = ReferenceModelBuilder(fin)
        _, answer = build_training_workbook(fin, tmp_path / f"Miss_{absent}_Trainer.xlsx")
        wb = load_workbook(answer, data_only=False)
        labels = _labels(wb["Condensed Financials"])
        present = "Interest Income" if absent == "interest_expense" else "Interest Expense"
        missing = "Interest Expense" if absent == "interest_expense" else "Interest Income"
        assert present in labels
        assert missing not in labels
        assert wb["Condensed Financials"].cell(
            row=labels["Net Interest"], column=2
        ).value == SOURCE_UNAVAILABLE
        assert wb["Condensed Financials"].cell(row=labels["NOPAT"], column=2).value == SOURCE_UNAVAILABLE
        families = {s.family_id for s in builder.expected_specs}
        assert "net_interest_fy" not in families
        assert "nopat_fy" not in families
        assert "revenue_link" in families
        assert "flev" in families
        smap = load_semantic_map(answer)
        assert {c.family_id for c in smap.all_ordered()} == families
        wb.close()


def test_both_interest_concepts_absent_gates_dependents(tmp_path: Path):
    fin = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    periods = [P1, P2]
    avail = assess_interest_availability(fin, periods)
    assert avail.interest_expense.line_present is False
    assert avail.interest_income.line_present is False
    assert avail.interest_expense.for_period(P1).reason == REASON_ABSENT_LINE
    anchor = compute_anchor(fin, periods)
    assert anchor.dupont["Sales Growth"][1] == pytest.approx(0.1)
    assert not is_source_unavailable(anchor.dupont["FLEV"][1])
    assert not is_source_unavailable(anchor.dupont["Actual ROE"][1])
    assert is_source_unavailable(anchor.dupont["NOPAT Margin"][0])
    assert is_source_unavailable(anchor.dupont["RNOA"][1])
    assert is_source_unavailable(anchor.dupont["After-tax CoD"][1])
    assert is_source_unavailable(anchor.dupont["Spread"][1])
    assert is_source_unavailable(anchor.dupont["ROE (decomposed)"][1])
    builder = ReferenceModelBuilder(fin)
    trainer, answer = build_training_workbook(fin, tmp_path / "BothAbsent_Trainer.xlsx")
    families = {s.family_id for s in builder.expected_specs}
    for family in (
        "net_interest_fy",
        "net_interest_after_tax_fy",
        "nopat_fy",
        "nopat_margin",
        "rnoa",
        "after_tax_cod",
        "spread",
        "roe_decomp",
        "rnoa_margin_turnover",
    ):
        assert family not in families
    for family in ("revenue_link", "sales_growth", "flev", "actual_roe", "noa_turnover"):
        assert family in families
    blank = check_workbook(trainer)
    assert blank.blank == blank.total
    assert blank.correct == 0
    assert blank.incorrect == 0
    smap = load_semantic_map(answer)
    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        from core.trainer.semantic_io import parse_cell_ref

        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0


def test_partial_period_absence_gates_only_dependent_periods(tmp_path: Path):
    fin = _base_fin()
    item = next(i for i in fin.income_statement if i.label == "Finance costs")
    del item.values[P2]
    periods = [P1, P2]
    avail = assess_concept_availability(fin.income_statement, "interest_expense", periods)
    assert avail.available_on(P1) is True
    assert avail.available_on(P2) is False
    assert avail.for_period(P2).reason == REASON_MISSING_PERIOD_VALUE
    anchor = compute_anchor(fin, periods)
    assert anchor.historical.net_interest[0] == pytest.approx(35.0)
    assert is_source_unavailable(anchor.historical.net_interest[1])
    assert not is_source_unavailable(anchor.historical.nopat[0])
    assert is_source_unavailable(anchor.historical.nopat[1])
    assert is_source_unavailable(anchor.hist_avg_after_tax_cod)
    with pytest.raises(MissingHistoricalValueError, match="interest_expense"):
        from core.model.source_values import required_period_value

        required_period_value(item, P2, field="interest_expense")
    builder = ReferenceModelBuilder(fin)
    _, answer = build_training_workbook(fin, tmp_path / "Partial_Trainer.xlsx")
    wb = load_workbook(answer, data_only=False)
    labels = _labels(wb["Condensed Financials"])
    ws = wb["Condensed Financials"]
    assert ws.cell(row=labels["Net Interest"], column=2).value not in (None, SOURCE_UNAVAILABLE)
    assert str(ws.cell(row=labels["Net Interest"], column=2).value).startswith("=")
    assert ws.cell(row=labels["Net Interest"], column=3).value == SOURCE_UNAVAILABLE
    families = {(s.family_id, s.period_index) for s in builder.expected_specs}
    assert ("net_interest_fy", 0) in families
    assert ("net_interest_fy", 1) not in families
    wb.close()


def test_reported_zero_is_available_not_gated(tmp_path: Path):
    fin = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", 0, 0),
            _li("Finance income", 0, 0),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    periods = [P1, P2]
    avail = assess_interest_availability(fin, periods)
    assert avail.output("net_interest", P1).available is True
    anchor = compute_anchor(fin, periods)
    assert anchor.historical.net_interest == [0.0, 0.0]
    assert anchor.historical.net_interest_after_tax == [0.0, 0.0]
    assert anchor.historical.nopat == [170.0, 187.0]
    _, answer = build_training_workbook(fin, tmp_path / "Zero_Trainer.xlsx")
    wb = load_workbook(answer, data_only=False)
    labels = _labels(wb["Condensed Financials"])
    assert "Interest Expense" in labels
    assert "Interest Income" in labels
    assert wb["Condensed Financials"].cell(row=labels["NOPAT"], column=2).value != SOURCE_UNAVAILABLE
    wb.close()


def test_alias_and_explicit_concept_precedence():
    alias = _base_fin()
    periods = [P1, P2]
    exp = resolve_line(alias.income_statement, "interest_expense", required=True)
    inc = resolve_line(alias.income_statement, "interest_income", required=True)
    assert exp.item is not None and exp.item.label == "Finance costs"
    assert inc.item is not None and inc.item.label == "Finance income"
    explicit = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Non-operating finance line", -40, -50, concept="interest_expense"),
            _li("Non-operating income line", 5, 6, concept="interest_income"),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    exp_x = resolve_line(explicit.income_statement, "interest_expense", required=True)
    assert exp_x.item is not None
    assert exp_x.item.concept == "interest_expense"
    assert exp_x.item.label == "Non-operating finance line"
    mixed = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -1, -1),
            _li("Bond coupon", -40, -50, concept="interest_expense"),
            _li("Finance income", 5, 6),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    chosen = resolve_line(mixed.income_statement, "interest_expense", required=True)
    assert chosen.item is not None
    assert chosen.item.concept == "interest_expense"
    assert chosen.item.label == "Bond coupon"
    anchor = compute_anchor(mixed, periods)
    assert anchor.historical.net_interest == [35.0, 44.0]


def test_ambiguity_and_malformed_inputs_still_raise():
    dup = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Finance costs", -40, -50),
            _li("Interest expense", -10, -12),
            _li("Finance income", 5, 6),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    with pytest.raises(AmbiguousLineError):
        assess_concept_availability(dup.income_statement, "interest_expense", [P1, P2])
    with pytest.raises(AmbiguousLineError):
        compute_anchor(dup, [P1, P2])
    malformed = _base_fin()
    item = next(i for i in malformed.income_statement if i.label == "Finance costs")
    item.values[P1] = "not-a-number"  # type: ignore[assignment]
    with pytest.raises(ValueError, match="invalid source value"):
        assess_concept_availability(malformed.income_statement, "interest_expense", [P1, P2])
    with pytest.raises(ValueError, match="invalid source value"):
        compute_anchor(malformed, [P1, P2])


def test_unrelated_required_facts_still_fail_closed():
    fin = _base_fin()
    ni = next(i for i in fin.income_statement if i.label == "Profit for the year")
    del ni.values[P2]
    with pytest.raises(MissingHistoricalValueError, match="net_income"):
        compute_anchor(fin, [P1, P2])
    missing_rev = _base_fin(
        income_statement=[
            _li("Finance costs", -40, -50),
            _li("Finance income", 5, 6),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    with pytest.raises(MissingLineError, match="revenue"):
        compute_anchor(missing_rev, [P1, P2])


def test_adjacent_period_dependencies_and_undefined_ratio_distinct():
    three = StandardizedFinancials(
        ticker="SYN3",
        company_name="Three Period Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=_periods(P1, P2, P3),
        income_statement=[
            _li("Revenue", 1000, 1100, v3=1210),
            _li("Finance costs", -40, -50, v3=-55),
            _li("Finance income", 5, 6, v3=7),
            _li("Profit before tax", 200, 220, v3=240),
            _li("Income tax expense", -30, -33, v3=-36),
            _li("Profit for the year", 170, 187, v3=204),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110, v3=120),
            _li("Trade receivables", 80, 90, v3=100),
            _li("Property, plant and equipment", 400, 420, v3=440),
            _li("Trade payables", 50, 55, v3=60),
            _li("Bank borrowings", 200, 210, v3=220),
            _li("Total equity", 330, 355, v3=380),
        ],
        cash_flow=[_li("Net cash from operating activities", 50, 60, v3=70)],
    )
    expense = next(i for i in three.income_statement if i.label == "Finance costs")
    del expense.values[P1]
    periods = [P1, P2, P3]
    anchor = compute_anchor(three, periods)
    assert is_source_unavailable(anchor.historical.nopat[0])
    assert not is_source_unavailable(anchor.historical.nopat[1])
    assert not is_source_unavailable(anchor.historical.nopat[2])
    assert is_source_unavailable(anchor.dupont["NOPAT Margin"][0])
    assert not is_source_unavailable(anchor.dupont["RNOA"][1])
    assert not is_source_unavailable(anchor.hist_avg_after_tax_cod)

    zero_debt = _base_fin(
        balance_sheet=[
            _li("Cash and cash equivalents", 100, 110),
            _li("Trade receivables", 80, 90),
            _li("Property, plant and equipment", 400, 420),
            _li("Trade payables", 50, 55),
            _li("Bank borrowings", 100, 110),
            _li("Total equity", 430, 455),
        ]
    )
    z_anchor = compute_anchor(zero_debt, [P1, P2])
    assert z_anchor.dupont["After-tax CoD"][1] == UNDEFINED_RATIO
    assert not is_source_unavailable(z_anchor.hist_avg_after_tax_cod)
    assert z_anchor.hist_avg_after_tax_cod == pytest.approx(0.04)


def test_python_excel_export_reload_and_check_agreement(tmp_path: Path):
    fin = _base_fin(
        income_statement=[
            _li("Revenue", 1000, 1100),
            _li("Profit before tax", 200, 220),
            _li("Income tax expense", -30, -33),
            _li("Profit for the year", 170, 187),
        ]
    )
    restored = standardized_from_payload(standardized_to_payload(fin))
    periods = canonical_fiscal_periods(restored)
    original = compute_anchor(fin, periods)
    round_trip = compute_anchor(restored, periods)
    assert original.historical.revenue == round_trip.historical.revenue
    assert original.historical.net_interest == round_trip.historical.net_interest
    trainer, answer = build_training_workbook(restored, tmp_path / "Reload_Trainer.xlsx")
    smap = load_semantic_map(answer)
    families = {c.family_id for c in smap.all_ordered()}
    assert "nopat_fy" not in families
    answer_wb = load_workbook(answer, data_only=False)
    trainer_wb = load_workbook(trainer, data_only=False)
    a_labels = _labels(answer_wb["Condensed Financials"])
    t_labels = _labels(trainer_wb["Condensed Financials"])
    assert a_labels["NOPAT"] == t_labels["NOPAT"]
    assert answer_wb["Condensed Financials"].cell(row=a_labels["NOPAT"], column=2).value == SOURCE_UNAVAILABLE
    assert trainer_wb["Condensed Financials"].cell(row=t_labels["NOPAT"], column=2).value == SOURCE_UNAVAILABLE
    for comp in smap.all_ordered():
        assert not is_source_unavailable(comp.expected_value)
        assert str(comp.formula).startswith("=")
    answer_wb.close()
    trainer_wb.close()
    blank = check_workbook(trainer)
    assert blank.blank == blank.total
    filled_wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        filled_wb[comp.tab].cell(row=row, column=col).value = comp.formula
    filled_wb.save(trainer)
    filled_wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    target = next(c for c in smap.all_ordered() if c.family_id == "revenue_link")
    _inject_formula_and_cached_value(
        trainer, target.tab, target.cell, formula="=0", cached_value=0.0
    )
    incorrect = check_workbook(trainer)
    assert incorrect.incorrect == 1


def test_lululemon_unchanged_facts_gate_interest_and_keep_supported_outputs(tmp_path: Path):
    payload = json.loads(LULU_STD.read_text(encoding="utf-8"))
    fin = standardized_from_payload(payload)
    periods = canonical_fiscal_periods(fin)
    assert periods == [
        date(2023, 1, 29),
        date(2024, 1, 28),
        date(2025, 2, 2),
        date(2026, 2, 1),
    ]
    avail = assess_interest_availability(fin, periods)
    assert avail.interest_expense.line_present is False
    assert avail.interest_income.line_present is False
    with pytest.raises(MissingLineError, match="interest_expense"):
        resolve_line(fin.income_statement, "interest_expense", required=True)
    anchor = compute_anchor(fin, periods)
    assert anchor.historical.revenue == [8110518.0, 9619278.0, 10588126.0, 11102600.0]
    assert all(is_source_unavailable(v) for v in anchor.historical.net_interest)
    assert all(is_source_unavailable(v) for v in anchor.historical.nopat)
    assert is_source_unavailable(anchor.hist_avg_after_tax_cod)
    ncit = next(
        i for i in fin.balance_sheet if i.concept == "non_current_income_taxes_payable"
    )
    assert ncit.values[date(2023, 1, 29)] == 28555.0
    assert ncit.values[date(2024, 1, 28)] == 15864.0
    assert ncit.values[date(2025, 2, 2)] == 0.0
    assert ncit.values[date(2026, 2, 1)] is None
    common = next(i for i in fin.balance_sheet if i.label == "Common stock")
    assert common.values == {
        date(2023, 1, 29): 611.0,
        date(2024, 1, 28): 606.0,
        date(2025, 2, 2): 581.0,
        date(2026, 2, 1): 557.0,
    }
    builder = ReferenceModelBuilder(fin)
    trainer, answer = build_training_workbook(fin, tmp_path / "LululemonAvail")
    assert trainer.exists() and answer.exists()
    families = {s.family_id for s in builder.expected_specs}
    assert "nopat_fy" not in families
    assert "net_interest_fy" not in families
    assert "after_tax_cod" not in families
    assert "revenue_link" in families
    assert "ppe_capex" in families
    wb = load_workbook(answer, data_only=False)
    labels = _labels(wb["Condensed Financials"])
    assert "Interest Expense" not in labels
    assert "Interest Income" not in labels
    assert wb["Condensed Financials"].cell(row=labels["NOPAT"], column=2).value == SOURCE_UNAVAILABLE
    assert ncit.values[date(2026, 2, 1)] is None
    blank = check_workbook(trainer)
    assert blank.blank == blank.total == len(builder.expected_specs)
    assert blank.incorrect == 0
    smap = load_semantic_map(answer)
    twb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        twb[comp.tab].cell(row=row, column=col).value = comp.formula
    twb.save(trainer)
    twb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    wb.close()

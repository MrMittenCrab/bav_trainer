"""Step 9L.1 — lease-liability intensity / trend diagnostics."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.data.line_identity import line_identity, validate_statement_identities
from core.engine.component_catalog import (
    LEASE_LIABILITY_COMPONENT_CATALOG,
    expand_lease_liability_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.classification import (
    AmbiguousClassificationOverrideError,
    reformulate_balance_sheet,
    resolve_classification_overrides,
)
from core.model.financial_math import compute_anchor
from core.model.historical_expected import (
    expected_value_for_component,
    lease_liability_expected_series,
)
from core.model.judgment import classification_judgment_cases
from core.model.lease_liability import (
    compute_lease_liability_series,
    lease_liability_applicable,
    lease_liability_availability,
)
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.check_context import (
    classification_overrides_for_check,
    load_check_context,
)
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _tiny(
    *,
    lease=(100.0, 120.0),
    revenue=(1000.0, 1100.0),
    split: bool = False,
    missing_period: bool = False,
):
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)

    def vals(a, b):
        return {d1: a, d2: b}

    cash, ar, ap, bank = 50.0, 40.0, 30.0, 20.0
    if split:
        cur, noncur = (40.0, 50.0), (60.0, 70.0)
        lease_total0 = cur[0] + noncur[0]
        lease_total1 = cur[1] + noncur[1]
        lease_items = [
            _li("Current lease liabilities", vals(*cur), concept="lease_liability"),
            _li(
                "Non-current lease liabilities",
                vals(*noncur),
                concept="lease_liability",
            ),
        ]
        lease0, lease1 = lease_total0, lease_total1
    else:
        lease0, lease1 = lease
        lease_items = [
            _li("Operating lease liabilities", vals(lease0, lease1), concept="lease_liability")
        ]
    eq0 = cash + ar - ap - bank - lease0
    eq1 = cash + ar - ap - bank - lease1
    bs = [
        _li("Cash and cash equivalents", vals(cash, cash)),
        _li("Trade receivables", vals(ar, ar)),
        _li("Trade payables", vals(ap, ap)),
        _li("Bank borrowings", vals(bank, bank)),
        *lease_items,
        _li("Total equity", vals(eq0, eq1)),
    ]
    if missing_period:
        lease_items[0].values = {d1: lease0}
    return StandardizedFinancials(
        ticker="LL",
        company_name="Lease Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=[
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ],
        income_statement=[
            _li("Revenue", vals(*revenue)),
            _li("Finance costs", vals(-10, -11)),
            _li("Finance income", vals(0, 0)),
            _li("Profit before tax", vals(200, 220)),
            _li("Income tax expense", vals(-30, -33)),
            _li("Profit for the year", vals(170, 187)),
        ],
        balance_sheet=bs,
        cash_flow=[_li("Net cash from operating activities", vals(80, 90))],
    )


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(LEASE_LIABILITY_COMPONENT_CATALOG) == 4
    assert [f.order for f in LEASE_LIABILITY_COMPONENT_CATALOG] == list(range(87, 91))
    specs = expand_lease_liability_specs(periods, start_order=1000)
    assert len(specs) == 18
    assert {s.family_order for s in specs} == set(range(87, 91))
    with pytest.raises(ValueError, match="duplicate"):
        expand_lease_liability_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_lease_liability_specs(list(reversed(periods)), start_order=1)


def test_ordinary_two_period_math():
    fin = _tiny()
    periods = list(canonical_fiscal_periods(fin))
    anchor = compute_anchor(fin, periods)
    assert lease_liability_applicable(fin)
    series = compute_lease_liability_series(fin, periods, anchor)
    assert series.lease_liability == (100.0, 120.0)
    assert series.lease_liability_to_revenue[0] == pytest.approx(0.10)
    assert series.lease_liability_to_revenue[1] == pytest.approx(120.0 / 1100.0)
    assert series.lease_liability_change == (None, 20.0)
    assert series.lease_liability_growth[0] is None
    assert series.lease_liability_growth[1] == pytest.approx(0.20)
    mapped = lease_liability_expected_series(series)
    assert set(mapped) == {f.id for f in LEASE_LIABILITY_COMPONENT_CATALOG}


def test_zero_denominator_and_sign_preservation():
    fin = _tiny(lease=(0.0, -10.0), revenue=(0.0, 100.0))
    periods = list(canonical_fiscal_periods(fin))
    series = compute_lease_liability_series(fin, periods, compute_anchor(fin, periods))
    assert series.lease_liability_to_revenue[0] == UNDEFINED_RATIO
    assert series.lease_liability[1] == -10.0
    assert series.lease_liability_growth[1] == UNDEFINED_RATIO

    fin2 = _tiny(lease=(50.0, 0.0), revenue=(1000.0, 1100.0))
    series2 = compute_lease_liability_series(
        fin2, list(canonical_fiscal_periods(fin2)), compute_anchor(fin2, list(canonical_fiscal_periods(fin2)))
    )
    assert series2.lease_liability_growth[1] == pytest.approx(-1.0)


def test_missing_period_fails_closed():
    fin = _tiny(missing_period=True)
    periods = list(canonical_fiscal_periods(fin))
    with pytest.raises(MissingHistoricalValueError):
        compute_lease_liability_series(fin, periods, compute_anchor(fin, periods))


def test_split_rows_are_ambiguous_and_omit_module():
    fin = _tiny(split=True)
    avail = lease_liability_availability(fin)
    assert avail.ambiguous is True
    assert avail.lease_liability is False
    assert not lease_liability_applicable(fin)
    builder = ReferenceModelBuilder(fin)
    assert builder.lease_liability_specs == ()


def test_canonical_demo_includes_lease_section(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "LL_BASE.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312
    lease = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in LEASE_LIABILITY_COMPONENT_CATALOG}
    ]
    assert len({c.family_id for c in lease}) == 4
    assert len(lease) == 18
    wb = load_workbook(answer)
    labels = [
        wb["ALT DuPont"].cell(row=r, column=1).value
        for r in range(1, (wb["ALT DuPont"].max_row or 1) + 1)
    ]
    assert "LEASE LIABILITY CONTEXT" in labels
    assert "Lease Analysis" not in wb.sheetnames
    wb.close()
    assert check_workbook(trainer).blank == 312


def test_check_accepts_formula_and_rejects_wrong(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "LL_CHK.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "lease_liability_to_revenue"
        and isinstance(c.expected_value, (int, float))
    )
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct >= 1

    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad = check_workbook(trainer)
    assert bad.incorrect >= 1


def test_trusted_lease_source_tamper_fails_before_recolor(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "LL_TAMPER.xlsx")
    smap = load_semantic_map(answer)
    comp = next(c for c in smap.all_ordered() if c.family_id == "lease_liability_source_link")
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    ws = wb["Balance Sheet"]
    lease_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "Operating lease liabilities"
    )
    ws.cell(lease_row, 2).value = 1
    wb.save(trainer)
    wb.close()
    with pytest.raises(ValueError, match="Trusted workbook cell was modified"):
        check_workbook(trainer)
    wb = load_workbook(trainer, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


def test_classification_switch_changes_reformulation_not_raw_lease_ratio(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "LL_JUDG.xlsx")
    smap = load_semantic_map(answer)
    ratio = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "lease_liability_to_revenue" and c.period_index == 4
    )
    noa = next(
        c for c in smap.all_ordered() if c.family_id == "noa_agg" and c.period_index == 4
    )
    net_debt = next(
        c for c in smap.all_ordered() if c.family_id == "net_debt" and c.period_index == 4
    )
    ref_ratio = float(ratio.expected_value)
    ref_noa = float(noa.expected_value)
    ref_nd = float(net_debt.expected_value)

    wb = load_workbook(trainer, data_only=False)
    ws = wb["Accounting Judgment"]
    for r in range(5, (ws.max_row or 5) + 1):
        if "lease" in str(ws.cell(r, 2).value or "").lower():
            ws.cell(r, 6).value = "Financial Liability"
            break
    else:
        raise AssertionError("lease judgment row not found")
    wb.save(trainer)
    wb.close()

    ctx = load_check_context(answer)
    wb = load_workbook(trainer, data_only=False)
    overrides = classification_overrides_for_check(wb, ctx)
    wb.close()
    periods = list(canonical_fiscal_periods(data))
    alt_anchor = compute_anchor(data, periods, classification_overrides=overrides)
    builder = ReferenceModelBuilder(data)
    alt_ratio = float(
        expected_value_for_component(
            alt_anchor,
            ratio,
            lease_liability=builder.lease_liability_series,
        )
    )
    alt_noa = float(expected_value_for_component(alt_anchor, noa))
    alt_nd = float(expected_value_for_component(alt_anchor, net_debt))
    assert alt_ratio == pytest.approx(ref_ratio)
    assert alt_noa != pytest.approx(ref_noa)
    assert alt_nd != pytest.approx(ref_nd)


def test_split_lease_judgment_without_diagnostics(tmp_path):
    fin = _tiny(split=True)
    validate_statement_identities(fin.balance_sheet, "balance_sheet")
    periods = list(canonical_fiscal_periods(fin))
    reform = reformulate_balance_sheet(fin, periods)
    cases = classification_judgment_cases(fin, periods, reform)
    lease_cases = [c for c in cases if "lease" in c.label.lower()]
    assert len(lease_cases) == 2
    assert all(c.override_selector.startswith("identity:") for c in lease_cases)
    assert lease_cases[0].override_selector != lease_cases[1].override_selector

    with pytest.raises(AmbiguousClassificationOverrideError):
        resolve_classification_overrides(
            [fin.balance_sheet[i] for i in reform.detail_indices],
            {"concept:lease_liability": "Financial Liability"},
        )

    trainer, answer = build_training_workbook(fin, tmp_path / "SPLIT.xlsx")
    smap = load_semantic_map(answer)
    assert not any(
        c.family_id in {f.id for f in LEASE_LIABILITY_COMPONENT_CATALOG}
        for c in smap.all_ordered()
    )

    wb = load_workbook(trainer, data_only=False)
    ws = wb["Accounting Judgment"]
    for r in range(5, (ws.max_row or 5) + 1):
        label = str(ws.cell(r, 2).value or "")
        if "Current lease" in label:
            ws.cell(r, 6).value = "Financial Liability"
        elif "Non-current lease" in label:
            ws.cell(r, 6).value = "Operating Long-Term Liability"
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    summary = check_workbook(trainer)
    assert summary.incorrect == 0
    assert summary.blank == 0

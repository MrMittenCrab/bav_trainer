"""Disclosure-led historical revenue-driver tests and workbook presentation."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
import shutil

import pytest
from openpyxl import load_workbook

from core.data.historical_strategy import (
    ROLE_OBJECTIVE,
    ROLE_OPERATING_USE,
    ROLE_STRATEGY,
    THEME_COMPARABLE_SALES,
    THEME_GEOGRAPHIC_GROWTH,
    THEME_PRODUCTIVITY,
    THEME_STORE_EXPANSION,
    HistoricalStrategyData,
    HistoricalStrategyDisclosure,
    load_strategy_disclosures,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.data.interface import (
    HistoricalManagementKpiDeferredDisagreement,
    HistoricalManagementKpiDeferredMember,
)
from core.engine.component_catalog import (
    REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_FAMILY_ID,
    REVENUE_DRIVER_COMPONENT_CATALOG,
    REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
    REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_RPS_FAMILY_ID,
    REVENUE_DRIVER_SHEET_NAME,
    REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
    SemanticCellRef,
    expand_revenue_driver_specs,
    is_operating_kpi_source_identity,
    resolve_revenue_driver_link_formula,
    revenue_driver_component_id,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.management_kpi_identity import (
    FAMILY_COMPARABLE_SALES_GROWTH,
    FAMILY_SALES_PER_SQUARE_FOOT,
    POP_COMPANY_OPERATED_STORES,
    POP_STORES_AND_DTC,
    POP_STORES_AND_ECOMMERCE,
)
from core.model.line_resolver import MissingLineError
from core.model.period_axis import PeriodAxisError
from core.model.revenue_driver import (
    ADDITIONAL_SPSF,
    FAILED_COMPSALES,
    FAILED_SPSF_GROWTH,
    FAILED_STORE_GROWTH,
    SCOPE_NOTE,
    VERDICT_CONTRADICTED,
    VERDICT_INSUFFICIENT,
    VERDICT_MIXED,
    VERDICT_SUPPORTED,
    compute_revenue_driver_analysis,
    revenue_driver_applicable,
)
from core.model.revenue_strategy_synthesis import (
    DEFERRED_SPSF_LINK,
    PROFESSIONAL_FALLBACK,
    UNTESTED_INITIATIVES,
    WHAT_HISTORY_ESTABLISHES,
    compute_historical_strategy_synthesis,
    disclosure_locator,
    strategy_synthesis_applicable,
)
from core.tests.test_capex import P1, P2, _tiny
from core.tests.test_historical_segment import _corp_values, _snapshot
from core.tests.test_learner_ready_presentation import (
    _assert_answer_key_no_yellow,
    assert_bav_has_no_exercise_framing,
)
from core.tests.test_geographic_segment_workbook import _geo_tiny
from core.tests.test_operating_kpi_facts import _kpi_model_observation
from core.tests.test_operating_kpi_management_history import _compsales, _spsf
from core.tests.test_operating_kpi_relationships import _fin_with_relationship
from core.trainer.checker import check_workbook
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, derive_trainer_workbook

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "build" / "input" / "fast_retailing" / "reconciled" / "standardized.json"
LULU_JSON = ROOT / "core" / "tests" / "fixtures" / "ordinary_reconcile" / "lululemon" / "standardized.json"
LULU_DISCLOSURES = (
    ROOT / "core" / "tests" / "fixtures" / "strategy" / "lululemon_management_disclosures.json"
)
P0 = date(2023, 12, 31)


def _disclosure(
    theme: str,
    *,
    role: str = ROLE_STRATEGY,
    text: str = "We open stores.",
    period: date = P2,
) -> HistoricalStrategyDisclosure:
    return HistoricalStrategyDisclosure(
        theme=theme,
        role=role,
        text=text,
        period=period,
        source_file="example.pdf",
        page_reference="Form 10-K p. 1",
        section="Item 1",
    )


def _with_strategy(fin, *disclosures: HistoricalStrategyDisclosure):
    fin.historical_strategy = HistoricalStrategyData(disclosures=disclosures)
    return fin


def _store_fin(
    counts: dict[date, float],
    revenue: dict[date, float | None],
    *disclosures: HistoricalStrategyDisclosure,
    management=None,
):
    observations = tuple(
        _kpi_model_observation(period, value) for period, value in counts.items()
    )
    fin = _fin_with_relationship(
        *observations,
        revenue=revenue,
        extra_periods=list(counts),
        management=management,
    )
    return _with_strategy(fin, *disclosures)


def _deferred_member(
    *,
    locator: str,
    definition_text: str,
    presentation_role: str = "current",
    extraction_document: str = "example.json",
    page_reference: str = "Form 10-K p. 10",
    physical_page_mapping: str = "10",
    reported_value: float | None = 1580.0,
) -> HistoricalManagementKpiDeferredMember:
    return HistoricalManagementKpiDeferredMember(
        locator=locator,
        extraction_document=extraction_document,
        page_reference=page_reference,
        physical_page_mapping=physical_page_mapping,
        presentation_role=presentation_role,
        definition_text=definition_text,
        population=POP_COMPANY_OPERATED_STORES,
        unit="USD_per_square_foot",
        basis="reported",
        calendar_week_adjustment="included",
        calendar_reporting_basis="52_week",
        reported_value=reported_value,
    )


def _deferred_disagreement(
    *,
    family: str = FAMILY_SALES_PER_SQUARE_FOOT,
    period: date = P2,
    reasons: tuple[str, ...] = ("definition_mismatch", "ordinary_disagreement"),
    members: tuple[HistoricalManagementKpiDeferredMember, ...] | None = None,
) -> HistoricalManagementKpiDeferredDisagreement:
    return HistoricalManagementKpiDeferredDisagreement(
        family=family,
        period=period,
        reasons=list(reasons),
        members=list(
            members
            or (
                _deferred_member(
                    locator=f"{family}:current:{period.isoformat()}",
                    definition_text="average during the year",
                    presentation_role="current",
                ),
                _deferred_member(
                    locator=f"{family}:prior:{period.isoformat()}",
                    definition_text="average ending square footage",
                    presentation_role="prior",
                    page_reference="Form 10-K p. 11",
                    physical_page_mapping="11",
                ),
            )
        ),
    )


def _attach_deferred(fin, *items: HistoricalManagementKpiDeferredDisagreement):
    data = fin.historical_operating_kpis
    assert data is not None
    data.deferred_disagreements = list(items)
    return fin


def test_catalog_orders_and_link_formulas():
    assert REVENUE_DRIVER_SHEET_NAME == "Revenue Driver Analysis"
    assert [family.order for family in REVENUE_DRIVER_COMPONENT_CATALOG] == list(
        range(177, 184)
    )
    assert [family.id for family in REVENUE_DRIVER_COMPONENT_CATALOG] == [
        REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
        REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
        REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
        REVENUE_DRIVER_COMPSALES_FAMILY_ID,
        REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
        REVENUE_DRIVER_RPS_FAMILY_ID,
        REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
    ]
    source = SemanticCellRef(
        id="src",
        semantic_key="operating_kpi.store_count.growth.2025-12-31",
        period_end=P2.isoformat(),
        cell="C12",
        tab="Store Count Analysis",
    )
    assert (
        resolve_revenue_driver_link_formula(source, from_tab=REVENUE_DRIVER_SHEET_NAME)
        == "='Store Count Analysis'!C12"
    )
    specs = expand_revenue_driver_specs(
        [P1, P2],
        start_order=1,
        store_growth_periods=(P2,),
    )
    assert [spec.id for spec in specs] == [
        revenue_driver_component_id(REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID, P2)
    ]


def test_opening_without_strategy_stays_professional_fallback(tmp_path):
    fin = _tiny()
    assert not strategy_synthesis_applicable(fin)
    trainer, answer = build_training_workbook(fin, tmp_path / "NO_STRATEGY.xlsx")
    awb = load_workbook(answer, data_only=False)
    opening = " ".join(
        str(cell.value or "") for row in awb["Overview"].iter_rows() for cell in row
    )
    assert awb["Overview"]["A1"].value == "BAV"
    assert PROFESSIONAL_FALLBACK in opening
    assert "Historical reading" in opening
    assert "HISTORICAL REVENUE AND DISCLOSED STRATEGY" not in opening
    assert "Schedules: " not in opening
    assert "Management statement" not in opening
    assert "Historical finding" not in opening
    assert any(
        cell.hyperlink
        and "Build Status" in str(getattr(cell.hyperlink, "target", "") or cell.hyperlink)
        for row in awb["Overview"].iter_rows()
        for cell in row
        if cell.hyperlink
    )
    awb.close()
    derive_trainer_workbook(answer, trainer)


def test_missing_disclosures_skip_analysis():
    fin = _fin_with_relationship(
        _kpi_model_observation(P1, 10),
        _kpi_model_observation(P2, 12),
        revenue={P1: 100.0, P2: 130.0},
    )
    assert fin.historical_strategy is None
    assert not revenue_driver_applicable(fin)
    with pytest.raises(MissingLineError, match="revenue-driver disclosures"):
        compute_revenue_driver_analysis(fin)
    payload = standardized_to_payload(fin)
    assert "historical_strategy" not in payload
    restored = standardized_from_payload(payload)
    assert restored.historical_strategy is None


def test_store_expansion_supported_mixed_contradicted_and_insufficient():
    supported = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_STORE_EXPANSION, role=ROLE_OBJECTIVE, text="We plan to open stores."),
    )
    result = compute_revenue_driver_analysis(supported)
    assert result.scope_note == SCOPE_NOTE
    assert result.calculation_kind == "analyst-derived"
    test = result.tests[0]
    assert test.theme == THEME_STORE_EXPANSION
    assert test.verdict == VERDICT_SUPPORTED
    assert test.sample_size == 2
    assert test.failed_requirement == ""
    assert any(item.role == ROLE_OBJECTIVE for item in test.disclosures)
    assert "not treated as achieved historical outcomes" in " ".join(test.limitations)
    assert "descriptive" in " ".join(test.limitations).lower()
    assert "divided by company-operated stores" in " ".join(test.identity_notes).lower()
    notes = " ".join(item.note for item in test.observations)
    assert "Both consolidated revenue and company-operated store counts grew." in notes

    mixed = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 120.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    mixed_test = compute_revenue_driver_analysis(mixed).tests[0]
    assert mixed_test.verdict == VERDICT_MIXED
    assert mixed_test.sample_size == 2
    assert any(
        "did not" in item.note for item in mixed_test.observations if item.consistent is False
    )

    contradicted = _store_fin(
        {P1: 10, P2: 12},
        {P1: 130.0, P2: 100.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    contradicted_test = compute_revenue_driver_analysis(contradicted).tests[0]
    assert contradicted_test.verdict == VERDICT_CONTRADICTED
    assert contradicted_test.sample_size == 1

    missing_history = _with_strategy(
        _tiny(with_payments=False),
        _disclosure(THEME_STORE_EXPANSION),
    )
    missing_test = compute_revenue_driver_analysis(missing_history).tests[0]
    assert missing_test.verdict == VERDICT_INSUFFICIENT
    assert missing_test.sample_size == 0
    assert missing_test.failed_requirement == FAILED_STORE_GROWTH


def test_missing_adjacent_inputs_do_not_force_a_verdict():
    fin = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: None},
        _disclosure(THEME_STORE_EXPANSION),
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    assert test.verdict == VERDICT_INSUFFICIENT
    assert test.sample_size == 0
    assert test.observations[0].consistent is None
    assert test.failed_requirement == FAILED_STORE_GROWTH


def test_comparable_sales_keeps_identities_and_descriptive_differences():
    fin = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: 130.0},
        _disclosure(THEME_COMPARABLE_SALES, role=ROLE_OPERATING_USE),
        management=[
            _compsales(period=P1, value=4.0, geography="global", basis="reported"),
            _compsales(period=P2, value=5.0, geography="global", basis="reported"),
            _compsales(period=P2, value=1.5, geography="global", basis="constant_dollar"),
        ],
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    assert test.theme == THEME_COMPARABLE_SALES
    assert test.verdict == VERDICT_SUPPORTED
    identities = {item.inputs["identity"] for item in test.observations if item.inputs.get("identity")}
    assert identities
    bases = {item.inputs.get("basis") for item in test.observations if item.consistent is not None}
    assert "reported" in bases
    assert "constant_dollar" not in bases
    joined_limits = " ".join(test.limitations).lower()
    assert "not new-store contribution" in joined_limits
    assert "fy2022" not in joined_limits
    assert "2023-01-29" not in joined_limits
    assert "store-only" not in joined_limits
    round_trip = standardized_from_payload(standardized_to_payload(fin))
    assert round_trip.historical_strategy == fin.historical_strategy

    absent = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: 130.0},
        _disclosure(THEME_COMPARABLE_SALES),
    )
    absent_test = compute_revenue_driver_analysis(absent).tests[0]
    assert absent_test.verdict == VERDICT_INSUFFICIENT
    assert absent_test.failed_requirement == FAILED_COMPSALES


def test_productivity_insufficient_without_adjacent_spsf_growth():
    fin = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: 130.0},
        _disclosure(THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE),
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    assert test.verdict == VERDICT_INSUFFICIENT
    assert test.failed_requirement == FAILED_SPSF_GROWTH
    assert ADDITIONAL_SPSF in test.additional_evidence
    assert "divided by company-operated stores" in " ".join(test.identity_notes).lower()
    joined_limits = " ".join(test.limitations)
    assert "2023-01-29" not in joined_limits
    assert "definition disagreement" not in joined_limits.lower()


def test_comparable_sales_limitations_follow_admitted_populations():
    fin = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: 130.0},
        _disclosure(THEME_COMPARABLE_SALES),
        management=[
            _compsales(
                period=P1,
                value=4.0,
                geography="",
                basis="reported",
                population=POP_COMPANY_OPERATED_STORES,
            ),
            _compsales(
                period=P2,
                value=5.0,
                geography="global",
                basis="reported",
                population=POP_STORES_AND_ECOMMERCE,
            ),
            _compsales(
                period=P1,
                value=6.0,
                geography="global",
                basis="reported",
                population=POP_STORES_AND_DTC,
            ),
        ],
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    joined = " ".join(test.limitations)
    assert POP_COMPANY_OPERATED_STORES in joined
    assert POP_STORES_AND_ECOMMERCE in joined
    assert POP_STORES_AND_DTC in joined
    assert P1.isoformat() in joined
    assert P2.isoformat() in joined
    assert "FY2022" not in joined
    assert "store-only and later stores-plus-DTC" not in joined


def test_productivity_limitations_follow_missing_spsf_periods():
    fin = _store_fin(
        {P0: 10, P1: 12, P2: 15},
        {P0: 100.0, P1: 130.0, P2: 160.0},
        _disclosure(THEME_PRODUCTIVITY),
        management=[
            _spsf(period=P1, value=1600),
            _spsf(period=P2, value=1500),
        ],
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    joined = " ".join(test.limitations)
    assert P0.isoformat() in joined
    assert "not bridged" in joined.lower()
    assert "2023-01-29" not in joined
    assert "definition disagreement" not in joined.lower()
    assert "ordinary_disagreement" not in joined


def test_productivity_discloses_evidence_derived_deferred_disagreement():
    period = P1
    disagreement = _deferred_disagreement(period=period)
    fin = _attach_deferred(
        _store_fin(
            {P0: 10, P1: 12, P2: 15},
            {P0: 100.0, P1: 130.0, P2: 160.0},
            _disclosure(THEME_PRODUCTIVITY),
            management=[_spsf(period=P2, value=1500)],
        ),
        disagreement,
    )
    test = compute_revenue_driver_analysis(fin).tests[0]
    joined = " ".join(test.limitations)
    assert period.isoformat() in joined
    assert "definition_mismatch" in joined
    assert "ordinary_disagreement" in joined
    assert "average during the year" in joined
    assert "average ending square footage" in joined
    assert disagreement.members[0].locator in joined
    assert disagreement.members[1].locator in joined
    assert "example.json" in joined
    assert "Form 10-K p. 10" in joined
    assert "audit-only" in joined.lower()
    assert "not admitted" in joined.lower()
    assert "does not create an admitted observation" in joined
    assert "2023-01-29" not in joined
    round_trip = standardized_from_payload(standardized_to_payload(fin))
    restored = round_trip.historical_operating_kpis.deferred_disagreements
    assert len(restored) == 1
    assert restored[0].period == period
    assert restored[0].reasons == list(disagreement.reasons)
    assert [item.locator for item in restored[0].members] == [
        item.locator for item in disagreement.members
    ]
    assert all(
        item.period != period
        for item in round_trip.historical_operating_kpis.management_observations
        if item.family == FAMILY_SALES_PER_SQUARE_FOOT
    )


def test_unrelated_absent_and_incompatible_inputs_do_not_assert_disagreement():
    absent = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: 130.0},
        _disclosure(THEME_PRODUCTIVITY),
    )
    absent_joined = " ".join(compute_revenue_driver_analysis(absent).tests[0].limitations)
    assert "ordinary_disagreement" not in absent_joined
    assert "definition_mismatch" not in absent_joined
    assert "audit-only" not in absent_joined.lower()

    unrelated = _attach_deferred(
        _store_fin(
            {P1: 10, P2: 12},
            {P1: 100.0, P2: 130.0},
            _disclosure(THEME_PRODUCTIVITY),
            _disclosure(THEME_COMPARABLE_SALES),
            management=[
                _compsales(period=P1, value=4.0, geography="global", basis="reported"),
                _compsales(period=P2, value=5.0, geography="global", basis="reported"),
            ],
        ),
        _deferred_disagreement(
            family=FAMILY_COMPARABLE_SALES_GROWTH,
            period=P2,
            members=(
                _deferred_member(
                    locator="compsales:current",
                    definition_text="stores only",
                    presentation_role="current",
                ),
                _deferred_member(
                    locator="compsales:prior",
                    definition_text="stores plus digital",
                    presentation_role="prior",
                ),
            ),
        ),
    )
    by_theme = {
        test.theme: test for test in compute_revenue_driver_analysis(unrelated).tests
    }
    productivity_joined = " ".join(by_theme[THEME_PRODUCTIVITY].limitations)
    compsales_joined = " ".join(by_theme[THEME_COMPARABLE_SALES].limitations)
    assert "ordinary_disagreement" not in productivity_joined
    assert "stores plus digital" not in productivity_joined
    assert "ordinary_disagreement" in compsales_joined
    assert "stores plus digital" in compsales_joined

    incompatible = _attach_deferred(
        _store_fin(
            {P1: 10, P2: 12},
            {P1: 100.0, P2: 130.0},
            _disclosure(THEME_PRODUCTIVITY),
        ),
        _deferred_disagreement(
            family=FAMILY_COMPARABLE_SALES_GROWTH,
            period=P2,
        ),
    )
    incompatible_joined = " ".join(
        compute_revenue_driver_analysis(incompatible).tests[0].limitations
    )
    assert "ordinary_disagreement" not in incompatible_joined
    assert "average during the year" not in incompatible_joined


def test_geographic_mixed_when_a_segment_subtracts():
    from core.tests.test_geographic_segment_workbook import _geo_tiny

    fin = _geo_tiny(
        _snapshot(P1, values=_corp_values(rev=(80.0, 25.0, 15.0))),
        _snapshot(P2, values=_corp_values(rev=(110.0, 15.0, 15.0))),
    )
    _with_strategy(fin, _disclosure(THEME_GEOGRAPHIC_GROWTH))
    test = compute_revenue_driver_analysis(fin).tests[0]
    assert test.theme == THEME_GEOGRAPHIC_GROWTH
    assert test.verdict == VERDICT_MIXED
    assert test.sample_size == 1
    assert "not organic, constant-currency, or causal" in " ".join(test.limitations)
    assert any("negatively" in item.note for item in test.observations)


def test_strategy_synthesis_connects_findings_without_claiming_outcomes():
    from core.tests.test_geographic_segment_workbook import _geo_tiny

    fin = _geo_tiny(
        _snapshot(P1, values=_corp_values(rev=(80.0, 25.0, 15.0))),
        _snapshot(P2, values=_corp_values(rev=(110.0, 15.0, 15.0))),
    )
    store = _store_fin(
        {P1: 10, P2: 12},
        {P1: 120.0, P2: 140.0},
        _disclosure(
            THEME_STORE_EXPANSION,
            role=ROLE_OBJECTIVE,
            text="We plan to open stores.",
        ),
        management=[
            _compsales(period=P1, value=4.0, geography="global", basis="reported"),
            _compsales(period=P2, value=5.0, geography="global", basis="reported"),
        ],
    )
    fin.historical_operating_kpis = store.historical_operating_kpis
    store_statement = _disclosure(
        THEME_STORE_EXPANSION,
        role=ROLE_OBJECTIVE,
        text="We plan to open stores.",
    )
    compsales_statement = _disclosure(
        THEME_COMPARABLE_SALES, role=ROLE_OPERATING_USE
    )
    productivity_statement = _disclosure(
        THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE
    )
    geo_statement = _disclosure(THEME_GEOGRAPHIC_GROWTH)
    fin.historical_strategy = HistoricalStrategyData(
        disclosures=(
            store_statement,
            compsales_statement,
            productivity_statement,
            geo_statement,
        )
    )
    analysis = compute_revenue_driver_analysis(fin)
    synthesis = compute_historical_strategy_synthesis(fin, analysis)
    assert strategy_synthesis_applicable(fin)
    assert "historical growth pattern" in synthesis.lead
    assert "do not establish causal drivers" in synthesis.lead
    assert "comprehensive strategy execution" in synthesis.lead
    assert synthesis.untested == UNTESTED_INITIATIVES
    assert synthesis.limits == WHAT_HISTORY_ESTABLISHES
    by_theme = {item.theme: item for item in synthesis.interpretations}
    assert set(by_theme) == {
        THEME_STORE_EXPANSION,
        THEME_COMPARABLE_SALES,
        THEME_PRODUCTIVITY,
        THEME_GEOGRAPHIC_GROWTH,
    }
    store_row = by_theme[THEME_STORE_EXPANSION]
    assert store_statement.text in store_row.management_statement
    assert disclosure_locator(store_statement) in store_row.management_statement
    assert analysis.tests[0].finding in store_row.finding
    assert REVENUE_DRIVER_SHEET_NAME in store_row.finding
    assert "not achieved historical outcomes" in store_row.inference
    assert "not new-store contribution" in store_row.inference
    geo_row = by_theme[THEME_GEOGRAPHIC_GROWTH]
    assert "contributed negatively" in geo_row.inference
    assert "arithmetic decomposition" in geo_row.inference
    productivity = by_theme[THEME_PRODUCTIVITY]
    assert "cannot test" in productivity.inference
    assert "not independent productivity evidence" in synthesis.productivity_gap
    assert DEFERRED_SPSF_LINK not in synthesis.productivity_gap
    assert "average during the year" not in synthesis.productivity_gap
    assert "ordinary_disagreement" not in synthesis.productivity_gap


def test_strategy_synthesis_links_deferred_spsf_without_promoting_it():
    period = P1
    disagreement = _deferred_disagreement(period=period)
    fin = _attach_deferred(
        _store_fin(
            {P0: 10, P1: 12, P2: 15},
            {P0: 100.0, P1: 130.0, P2: 160.0},
            _disclosure(THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE),
            management=[_spsf(period=P2, value=1500)],
        ),
        disagreement,
    )
    synthesis = compute_historical_strategy_synthesis(fin)
    assert DEFERRED_SPSF_LINK in synthesis.productivity_gap
    assert "sample size 0" in synthesis.productivity_gap
    assert disagreement.members[0].locator not in synthesis.productivity_gap
    assert disagreement.members[0].definition_text not in synthesis.productivity_gap
    assert "average during the year" not in synthesis.lead
    productivity = next(
        item for item in synthesis.interpretations if item.theme == THEME_PRODUCTIVITY
    )
    assert disagreement.members[0].locator not in productivity.finding
    assert disagreement.members[0].locator not in productivity.inference


def test_interim_axis_is_rejected_when_operating_history_exists():
    fin = _store_fin(
        {P1: 10, P2: 12},
        {P1: 100.0, P2: 130.0},
        _disclosure(THEME_STORE_EXPANSION),
    )
    for period in fin.periods:
        period.is_interim = True
    with pytest.raises(PeriodAxisError, match="annual fiscal periods"):
        compute_revenue_driver_analysis(fin)


DRIVER_FAMILY_IDS = (
    REVENUE_DRIVER_STORE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_REVENUE_GROWTH_FAMILY_ID,
    REVENUE_DRIVER_STORE_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_FAMILY_ID,
    REVENUE_DRIVER_COMPSALES_DIFFERENCE_FAMILY_ID,
    REVENUE_DRIVER_RPS_FAMILY_ID,
    REVENUE_DRIVER_GEO_CONTRIBUTION_FAMILY_ID,
)


def _copy_workbooks(trainer: Path, answer: Path, dest: Path) -> tuple[Path, Path]:
    dest.mkdir()
    copied_trainer = dest / trainer.name
    copied_answer = dest / answer.name
    shutil.copy2(trainer, copied_trainer)
    shutil.copy2(answer, copied_answer)
    for sidecar in (
        answer.with_suffix(".component_map.json"),
        answer.with_suffix(".assumptions.json"),
        answer.with_suffix(".trainer.json"),
    ):
        if sidecar.is_file():
            shutil.copy2(sidecar, dest / sidecar.name)
    return copied_trainer, copied_answer


def _assert_readable_driver_layout(sheet) -> None:
    long_cells = 0
    for row in sheet.iter_rows(min_col=1, max_col=2, max_row=sheet.max_row or 1):
        for cell in row:
            value = cell.value
            if not isinstance(value, str) or value.startswith("="):
                continue
            if len(value) < 80:
                continue
            long_cells += 1
            assert cell.alignment.wrap_text is True, cell.coordinate
            height = sheet.row_dimensions[cell.row].height or 15
            assert height > 15, (cell.coordinate, len(value), height)
    assert long_cells >= 1


def test_workbook_links_notes_trainer_check_and_skips_without_disclosures(tmp_path):
    fin = _geo_tiny(
        _snapshot(P1, values=_corp_values(rev=(80.0, 25.0, 15.0))),
        _snapshot(P2, values=_corp_values(rev=(110.0, 15.0, 15.0))),
    )
    store = _store_fin(
        {P1: 10, P2: 12},
        {P1: 120.0, P2: 140.0},
        _disclosure(THEME_STORE_EXPANSION, role=ROLE_OBJECTIVE),
        management=[
            _compsales(period=P1, value=4.0, geography="global", basis="reported"),
            _compsales(period=P2, value=5.0, geography="global", basis="reported"),
        ],
    )
    fin.historical_operating_kpis = store.historical_operating_kpis
    fin.historical_strategy = HistoricalStrategyData(
        disclosures=(
            _disclosure(THEME_STORE_EXPANSION, role=ROLE_OBJECTIVE),
            _disclosure(THEME_COMPARABLE_SALES, role=ROLE_OPERATING_USE),
            _disclosure(THEME_PRODUCTIVITY, role=ROLE_OPERATING_USE),
            _disclosure(THEME_GEOGRAPHIC_GROWTH),
        )
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "DRIVER.xlsx")
    smap = load_semantic_map(answer)
    families = {comp.family_id for comp in smap.all_ordered()}
    for family_id in DRIVER_FAMILY_IDS:
        assert family_id in families
    awb = load_workbook(answer, data_only=False)
    assert REVENUE_DRIVER_SHEET_NAME in awb.sheetnames
    sheet = awb[REVENUE_DRIVER_SHEET_NAME]
    values = " ".join(
        str(cell.value or "")
        for row in sheet.iter_rows()
        for cell in row
    )
    assert "Analyst hypothesis" in values
    assert "Management statement" in values
    assert "Period-specific interpretation" in values
    assert "We plan to open stores." not in values or "objective" in values.lower()
    assert "descriptive" in values.lower()
    assert "NOTES" in values
    assert "not new-store contribution" in values.lower() or "causal attribution" in values.lower()
    assert "contributed negatively" in values.lower()
    assert "FY2022 store-only" not in values
    assert "2023-01-29 SPSF" not in values
    assert "DEFERRED DISAGREEMENT" not in values
    assert any(
        isinstance(cell.value, str) and cell.value.startswith("=")
        for row in sheet.iter_rows()
        for cell in row
    )
    _assert_readable_driver_layout(sheet)
    overview = awb["Overview"]
    opening = " ".join(
        str(cell.value or "") for row in overview.iter_rows() for cell in row
    )
    assert overview["A1"].value == "BAV"
    assert "Historical reading" in opening
    assert "Evidence limits" in opening
    assert "untested" in opening.lower()
    assert "do not establish causal drivers" in opening
    assert PROFESSIONAL_FALLBACK not in opening
    assert "HISTORICAL REVENUE AND DISCLOSED STRATEGY" not in opening
    assert "Management statement" not in opening
    assert "Historical finding" not in opening
    assert "Analyst inference" not in opening
    assert "We open stores." not in opening
    assert "Schedules: " not in opening
    assert any(
        cell.hyperlink
        and "Revenue Driver Analysis" in str(getattr(cell.hyperlink, "target", "") or cell.hyperlink)
        for row in overview.iter_rows()
        for cell in row
        if cell.hyperlink
    )
    assert any(
        cell.hyperlink
        and "Build Status" in str(getattr(cell.hyperlink, "target", "") or cell.hyperlink)
        for row in overview.iter_rows()
        for cell in row
        if cell.hyperlink
    )
    _assert_readable_driver_layout(overview)
    awb.close()
    _assert_answer_key_no_yellow(answer)
    assert_bav_has_no_exercise_framing(answer)
    before = answer.read_bytes()
    derive_trainer_workbook(answer, trainer)
    assert answer.read_bytes() == before
    twb = load_workbook(trainer, data_only=False)
    tsheet = twb[REVENUE_DRIVER_SHEET_NAME]
    driver_comps = [c for c in smap.all_ordered() if c.family_id.startswith("revenue_driver_")]
    for comp in driver_comps:
        cell = twb[comp.tab][comp.cell]
        assert cell.value is None
        assert cell.comment is None
        assert cell.fill.fgColor.rgb in ("00FFFF00", "FFFFFF00")
    assert tsheet["A1"].value is not None
    twb.close()
    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.correct == 0
    assert blank.blank == blank.total
    dumped_blank = repr(blank)
    assert "999" not in dumped_blank

    practice = [
        comp for comp in smap.all_ordered() if not is_operating_kpi_source_identity(comp)
    ]
    filled_trainer, filled_answer = _copy_workbooks(
        trainer, answer, tmp_path / "filled_check"
    )
    filled_wb = load_workbook(filled_trainer, data_only=False)
    for comp in practice:
        row, col = parse_cell_ref(comp.cell)
        filled_wb[comp.tab].cell(row=row, column=col).value = comp.formula
    filled_wb.save(filled_trainer)
    filled_wb.close()
    filled = check_workbook(filled_trainer)
    assert filled.incorrect == 0
    assert filled.blank == 0
    assert filled.correct == filled.total
    assert answer.read_bytes() == before
    assert filled_answer.read_bytes() == before

    for family_id in DRIVER_FAMILY_IDS:
        case_trainer, case_answer = _copy_workbooks(
            filled_trainer, filled_answer, tmp_path / f"incorrect_{family_id}"
        )
        bad = next(comp for comp in driver_comps if comp.family_id == family_id)
        _inject_formula_and_cached_value(
            case_trainer,
            bad.tab,
            bad.cell,
            formula="=999",
            cached_value=999.0,
        )
        summary = check_workbook(case_trainer)
        assert summary.incorrect == 1
        assert summary.blank == 0
        assert summary.correct == summary.total - 1
        dumped = repr(summary)
        assert "=999" not in dumped
        assert bad.formula not in dumped
        assert case_answer.read_bytes() == before
        assert answer.read_bytes() == before

    skipped = ReferenceModelBuilder(_tiny(with_payments=False))
    assert skipped.revenue_driver_schedule is False
    assert skipped.revenue_driver_specs == ()


def test_protected_standardized_payloads_do_not_carry_strategy():
    lulu = standardized_from_payload(json.loads(LULU_JSON.read_text(encoding="utf-8")))
    fr = standardized_from_payload(json.loads(FR_JSON.read_text(encoding="utf-8")))
    assert lulu.historical_strategy is None
    assert fr.historical_strategy is None
    assert not revenue_driver_applicable(lulu)
    assert not revenue_driver_applicable(fr)
    assert ReferenceModelBuilder(fr).revenue_driver_schedule is False
    assert ReferenceModelBuilder(fr).revenue_driver_specs == ()


def test_lululemon_ordinary_disclosures_test_admitted_history(tmp_path):
    from core.current_build import prepare_company_input, resolve_company

    company = resolve_company("Lululemon")
    fin = prepare_company_input(company, tmp_path)
    fixture = load_strategy_disclosures(LULU_DISCLOSURES)
    assert fin.historical_strategy == fixture
    analysis = compute_revenue_driver_analysis(fin)
    by_theme = {test.theme: test for test in analysis.tests}
    assert set(by_theme) == {
        THEME_STORE_EXPANSION,
        THEME_COMPARABLE_SALES,
        THEME_PRODUCTIVITY,
        THEME_GEOGRAPHIC_GROWTH,
    }
    store = by_theme[THEME_STORE_EXPANSION]
    assert store.verdict == VERDICT_SUPPORTED
    assert store.sample_size == 4
    assert [period.isoformat() for period in store.periods_tested] == [
        "2023-01-29",
        "2024-01-28",
        "2025-02-02",
        "2026-02-01",
    ]
    latest = next(item for item in store.observations if item.period.isoformat() == "2026-02-01")
    assert latest.consistent is True
    assert "Store-count growth exceeded revenue growth" in latest.note or "exceeded revenue growth" in latest.note
    assert "2026-02-01" in latest.note
    assert "Period-end Revenue per Store declined" in latest.note or "period-end Revenue per Store declined" in latest.note
    assert "exceeded revenue growth" in store.finding
    assert any(item.role == ROLE_OBJECTIVE for item in store.disclosures)
    compsales = by_theme[THEME_COMPARABLE_SALES]
    assert compsales.verdict in {VERDICT_SUPPORTED, VERDICT_MIXED}
    compsales_limits = " ".join(compsales.limitations)
    assert "ineligible" in compsales_limits
    assert "company_operated_stores" in compsales_limits
    assert "company_operated_stores_and_direct_to_consumer" in compsales_limits or "direct_to_consumer" in compsales_limits
    assert "2023-01-29" in compsales_limits
    assert "FY2022 store-only" not in compsales_limits
    identities = {
        str(item.inputs.get("identity"))
        for item in compsales.observations
        if item.inputs.get("identity")
    }
    assert any("store" in identity.lower() or "dtc" in identity.lower() or identity for identity in identities)
    productivity = by_theme[THEME_PRODUCTIVITY]
    assert productivity.verdict == VERDICT_INSUFFICIENT
    assert productivity.failed_requirement == FAILED_SPSF_GROWTH
    productivity_limits = " ".join(productivity.limitations)
    assert "2023-01-29" in productivity_limits
    assert "not bridged" in productivity_limits.lower()
    assert "definition_mismatch" in productivity_limits
    assert "ordinary_disagreement" in productivity_limits
    assert "audit-only" in productivity_limits.lower()
    assert "not admitted" in productivity_limits.lower()
    deferred = fin.historical_operating_kpis.deferred_disagreements
    spsf_deferred = [
        item for item in deferred if item.family == FAMILY_SALES_PER_SQUARE_FOOT
    ]
    assert len(spsf_deferred) == 1
    assert spsf_deferred[0].period.isoformat() == "2023-01-29"
    assert "definition_mismatch" in spsf_deferred[0].reasons
    assert "ordinary_disagreement" in spsf_deferred[0].reasons
    locators = {member.locator for member in spsf_deferred[0].members}
    definitions = {member.definition_text for member in spsf_deferred[0].members}
    assert len(spsf_deferred[0].members) == 2
    assert locators
    assert any("during" in text.lower() or "ending" in text.lower() for text in definitions)
    assert all(locator in productivity_limits for locator in locators)
    for member in spsf_deferred[0].members:
        assert member.locator in productivity_limits
        assert member.extraction_document in productivity_limits
        if member.page_reference:
            assert member.page_reference in productivity_limits
        assert member.definition_text in productivity_limits
    spsf_admitted = {
        item.period.isoformat()
        for item in fin.historical_operating_kpis.management_observations
        if item.family == FAMILY_SALES_PER_SQUARE_FOOT
    }
    assert "2023-01-29" not in spsf_admitted
    assert len(spsf_admitted) == 3
    geographic = by_theme[THEME_GEOGRAPHIC_GROWTH]
    assert geographic.verdict == VERDICT_MIXED
    assert geographic.sample_size >= 1
    assert any("contributed negatively" in item.note for item in geographic.observations)
    assert "contributed negatively" in geographic.finding
    assert "2026-02-01" in geographic.finding or any(
        "2026-02-01" in item.note for item in geographic.observations
    )
    builder = ReferenceModelBuilder(fin)
    assert builder.revenue_driver_schedule is True
    assert builder.revenue_driver_specs
    exported = standardized_from_payload(standardized_to_payload(fin))
    assert exported.historical_strategy == fin.historical_strategy
    trainer, answer = build_training_workbook(fin, tmp_path / "LULU_DRIVER.xlsx")
    awb = load_workbook(answer, data_only=False)
    sheet = awb[REVENUE_DRIVER_SHEET_NAME]
    values = " ".join(str(cell.value or "") for row in sheet.iter_rows() for cell in row)
    assert "Period-specific interpretation" in values
    assert "exceeded revenue growth" in values.lower()
    assert "contributed negatively" in values.lower()
    assert "2026-02-01" in values
    assert "FY2022 store-only" not in values
    assert "definition_mismatch" in values
    assert "ordinary_disagreement" in values
    assert "2023-01-29" in values
    assert "DEFERRED DISAGREEMENT" in values
    for member in spsf_deferred[0].members:
        assert member.locator in values
        assert member.definition_text in values
    _assert_readable_driver_layout(sheet)
    overview = awb["Overview"]
    opening = " ".join(
        str(cell.value or "") for row in overview.iter_rows() for cell in row
    )
    assert overview["A1"].value == "BAV"
    assert "Historical reading" in opening
    assert "Evidence limits" in opening
    assert "cannot test" in opening.lower() or "cannot be treated" in opening.lower()
    assert "audit-only" in opening.lower()
    assert DEFERRED_SPSF_LINK in opening
    assert "untested" in opening.lower()
    assert "do not establish causal drivers" in opening
    assert "HISTORICAL REVENUE AND DISCLOSED STRATEGY" not in opening
    assert "Management statement" not in opening
    assert store.finding not in opening
    assert "Schedules: " not in opening
    for disclosure in fixture.disclosures:
        assert disclosure_locator(disclosure) in values
        if disclosure.role != ROLE_OBJECTIVE:
            assert disclosure.text in values
        assert disclosure_locator(disclosure) not in opening
    assert store.finding in values
    assert "exceeded revenue growth" in values.lower()
    assert "2026-02-01" in values
    assert "contributed negatively" in values.lower()
    assert "arithmetic decomposition" in values.lower()
    assert "not achieved outcomes" in values
    for member in spsf_deferred[0].members:
        assert member.locator not in opening
        assert member.definition_text not in opening
    assert any(
        cell.hyperlink
        and "Revenue Driver Analysis" in str(getattr(cell.hyperlink, "target", "") or cell.hyperlink)
        for row in overview.iter_rows()
        for cell in row
        if cell.hyperlink
    )
    _assert_readable_driver_layout(overview)
    awb.close()
    _assert_answer_key_no_yellow(answer)
    assert_bav_has_no_exercise_framing(answer)
    before = answer.read_bytes()
    derive_trainer_workbook(answer, trainer)
    assert answer.read_bytes() == before


def test_lululemon_revenue_reconstruction_and_seven_part_validation():
    from core.current_build import prepare_company_input, resolve_company

    fin = prepare_company_input(resolve_company("Lululemon"), Path("/tmp/unused"))
    analysis = compute_revenue_driver_analysis(fin)
    geo = analysis.geographic_reconstruction
    footprint = analysis.footprint_identity
    assert geo is not None
    assert footprint is not None
    assert all(value == 0.0 for value in geo.residual)
    assert all(
        value == 0.0 or value is None for value in footprint.reconstruction_residual
    )
    for value in footprint.change_residual:
        if value is None:
            continue
        assert abs(value) < 1e-6
    assert "non-store" in footprint.scope_note.lower() or "outside those stores" in footprint.scope_note.lower()
    names = {item.name for item in analysis.assessments}
    assert "geographic revenue reconstruction" in names
    assert "footprint and intensity identity" in names
    assert "comparable-sales coincidence" in names
    assert "sales-per-square-foot productivity" in names
    spsf = next(
        item
        for item in analysis.assessments
        if item.name.startswith("sales-per-square-foot")
    )
    assert spsf.established is False
    assert spsf.kind == "unestablished_inference"
    for test in analysis.tests:
        assert test.assessment is not None

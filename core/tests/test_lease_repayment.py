"""Step 9N.3 — Lease-repayment practice, ratios, and workbook Check."""

from __future__ import annotations

import copy
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload
from core.engine.component_catalog import (
    LEASE_REPAYMENT_COMPONENT_CATALOG,
    expand_lease_repayment_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.historical_expected import lease_repayment_expected_series
from core.model.lease_repayment import (
    compute_lease_repayment_series,
    lease_repayment_applicable,
    lease_repayment_availability,
    resolve_lease_repayment_source,
)
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
STD_JSON = ROOT / "benchmark" / "fast_retailing" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"

P1 = date(2024, 12, 31)
P2 = date(2025, 12, 31)


def _ingest_demo():
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _tiny(
    *,
    repayments=(-100.0, -120.0),
    revenue=(1000.0, 1100.0),
    with_repayments: bool = True,
    label_only: bool = False,
    duplicate: bool = False,
    missing_period: bool = False,
    none_period: bool = False,
    missing_revenue: bool = False,
    none_revenue: bool = False,
    repayments_first: bool = False,
    on_balance_sheet: bool = False,
    single_period: bool = False,
):
    if single_period:
        d1 = date(2025, 12, 31)
        periods = [FinancialPeriod(end_date=d1, label="FY2025")]

        def vals(a, b=None):
            return {d1: a}

        repay_vals = (repayments[0],)
        rev_vals = (revenue[0],)
    else:
        d1, d2 = P1, P2
        periods = [
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ]

        def vals(a, b=None):
            return {d1: a, d2: b if b is not None else a}

        repay_vals = repayments
        rev_vals = revenue

    concept = "" if label_only else "repayments_of_lease_liabilities"
    if missing_period and not single_period:
        repay_values = {P1: repay_vals[0]}
    elif none_period and not single_period:
        repay_values = {P1: repay_vals[0], P2: None}
    elif single_period:
        repay_values = vals(repay_vals[0])
    else:
        repay_values = vals(*repay_vals)

    repay_item = _li(
        "Repayments of lease liabilities",
        repay_values,
        concept=concept,
    )

    if missing_revenue and not single_period:
        rev_values = {P1: rev_vals[0]}
    elif none_revenue and not single_period:
        rev_values = {P1: rev_vals[0], P2: None}
    elif single_period:
        rev_values = vals(rev_vals[0])
    else:
        rev_values = vals(*rev_vals)

    cash, ar, ap, bank = 50.0, 40.0, 30.0, 20.0
    eq0 = cash + ar - ap - bank
    eq1 = eq0
    bs = [
        _li("Cash and cash equivalents", vals(cash, cash)),
        _li("Trade receivables", vals(ar, ar)),
        _li("Trade payables", vals(ap, ap)),
        _li("Bank borrowings", vals(bank, bank)),
        _li("Total equity", vals(eq0, eq1)),
    ]
    if single_period:
        is_rows = [
            _li("Revenue", rev_values, concept="revenue"),
            _li("Finance costs", vals(-10)),
            _li("Finance income", vals(0)),
            _li("Profit before tax", vals(200)),
            _li("Income tax expense", vals(-30)),
            _li("Profit for the year", vals(170)),
        ]
        cf = [
            _li("Net cash from operating activities", vals(80)),
        ]
    else:
        is_rows = [
            _li("Revenue", rev_values, concept="revenue"),
            _li("Finance costs", vals(-10, -11)),
            _li("Finance income", vals(0, 0)),
            _li("Profit before tax", vals(200, 220)),
            _li("Income tax expense", vals(-30, -33)),
            _li("Profit for the year", vals(170, 187)),
        ]
        cf = [
            _li("Net cash from operating activities", vals(80, 90)),
        ]

    if with_repayments:
        target = bs if on_balance_sheet else cf
        if repayments_first:
            target.insert(0, repay_item)
        else:
            target.append(repay_item)
        if duplicate and not on_balance_sheet:
            cf.append(
                _li(
                    "Repayments of lease liabilities duplicate",
                    vals(*repay_vals) if not single_period else vals(repay_vals[0]),
                    concept="repayments_of_lease_liabilities",
                )
            )

    return StandardizedFinancials(
        ticker="LEASE_REPAY",
        company_name="Lease Repay Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=is_rows,
        balance_sheet=bs,
        cash_flow=cf,
    )


def _anchor(fin: StandardizedFinancials):
    return compute_anchor(fin, list(canonical_fiscal_periods(fin)))


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(LEASE_REPAYMENT_COMPONENT_CATALOG) == 2
    assert [f.order for f in LEASE_REPAYMENT_COMPONENT_CATALOG] == [122, 123]
    specs = expand_lease_repayment_specs(periods, start_order=1000)
    assert len(specs) == 10
    assert {s.family_id for s in specs} == {
        "lease_repayments",
        "lease_repayments_to_revenue",
    }
    assert all(s.semantic_key.startswith("lease_repayment.") for s in specs)
    assert all(s.category == "lease_repayment" for s in specs)
    with pytest.raises(ValueError, match="duplicate"):
        expand_lease_repayment_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_lease_repayment_specs(list(reversed(periods)), start_order=1)


def test_unique_concept_resolution_and_renamed_label():
    fin = _tiny()
    item = resolve_lease_repayment_source(fin)
    assert item is not None
    assert item.concept == "repayments_of_lease_liabilities"
    assert item.label == "Repayments of lease liabilities"
    assert lease_repayment_applicable(fin)
    avail = lease_repayment_availability(fin)
    assert (
        avail.repayments_of_lease_liabilities is True and avail.ambiguous is False
    )

    renamed = _tiny()
    renamed.cash_flow[-1] = _li(
        "Lease liability principal repaid (renamed)",
        {P1: -100.0, P2: -120.0},
        concept="repayments_of_lease_liabilities",
    )
    assert resolve_lease_repayment_source(renamed) is not None
    assert (
        resolve_lease_repayment_source(renamed).label
        == "Lease liability principal repaid (renamed)"
    )


def test_reordered_rows_still_resolve():
    first = _tiny(repayments_first=True)
    last = _tiny(repayments_first=False)
    assert resolve_lease_repayment_source(first) is not None
    assert resolve_lease_repayment_source(last) is not None
    assert (
        resolve_lease_repayment_source(first).values
        == resolve_lease_repayment_source(last).values
    )


def test_absent_label_only_duplicate_wrong_statement():
    absent = _tiny(with_repayments=False)
    assert not lease_repayment_applicable(absent)
    assert resolve_lease_repayment_source(absent) is None
    assert lease_repayment_availability(absent).repayments_of_lease_liabilities is False
    with pytest.raises(MissingLineError):
        compute_lease_repayment_series(
            absent, list(canonical_fiscal_periods(absent)), _anchor(absent)
        )
    assert ReferenceModelBuilder(absent).lease_repayment_specs == ()

    label_only = _tiny(label_only=True)
    assert not lease_repayment_applicable(label_only)
    assert resolve_lease_repayment_source(label_only) is None
    assert (
        resolve_line(
            label_only.cash_flow, "repayments_of_lease_liabilities", required=False
        ).item
        is None
    )
    assert ReferenceModelBuilder(label_only).lease_repayment_specs == ()

    dup = _tiny(duplicate=True)
    avail = lease_repayment_availability(dup)
    assert avail.ambiguous is True
    assert avail.repayments_of_lease_liabilities is False
    assert not lease_repayment_applicable(dup)
    assert resolve_lease_repayment_source(dup) is None
    with pytest.raises(AmbiguousLineError):
        resolve_line(dup.cash_flow, "repayments_of_lease_liabilities", required=False)
    assert ReferenceModelBuilder(dup).lease_repayment_specs == ()

    wrong = _tiny(on_balance_sheet=True)
    assert not lease_repayment_applicable(wrong)
    assert resolve_lease_repayment_source(wrong) is None
    assert (
        resolve_line(
            wrong.balance_sheet, "repayments_of_lease_liabilities", required=False
        ).item
        is not None
    )


def test_sign_conversion_ratios_positive_negative_mixed_zero():
    negative = _tiny(repayments=(-50.0, -80.0), revenue=(1000.0, 2000.0))
    series_neg = compute_lease_repayment_series(
        negative, list(canonical_fiscal_periods(negative)), _anchor(negative)
    )
    assert series_neg.repayments_reported == (-50.0, -80.0)
    assert series_neg.lease_repayments == (50.0, 80.0)
    assert series_neg.lease_repayments_to_revenue[0] == pytest.approx(50.0 / 1000.0)
    assert series_neg.lease_repayments_to_revenue[1] == pytest.approx(80.0 / 2000.0)

    positive = _tiny(repayments=(50.0, 80.0))
    series_pos = compute_lease_repayment_series(
        positive, list(canonical_fiscal_periods(positive)), _anchor(positive)
    )
    assert series_pos.repayments_reported == (50.0, 80.0)
    assert series_pos.lease_repayments == (-50.0, -80.0)
    assert series_pos.lease_repayments_to_revenue[0] == pytest.approx(-50.0 / 1000.0)

    mixed = _tiny(repayments=(-50.0, 80.0))
    series_mix = compute_lease_repayment_series(
        mixed, list(canonical_fiscal_periods(mixed)), _anchor(mixed)
    )
    assert series_mix.repayments_reported == (-50.0, 80.0)
    assert series_mix.lease_repayments == (50.0, -80.0)

    zero = _tiny(repayments=(0.0, -10.0))
    series_z = compute_lease_repayment_series(
        zero, list(canonical_fiscal_periods(zero)), _anchor(zero)
    )
    assert series_z.repayments_reported == (0.0, -10.0)
    assert series_z.lease_repayments == (0.0, 10.0)
    assert series_z.lease_repayments_to_revenue[0] == pytest.approx(0.0)

    mapped = lease_repayment_expected_series(series_neg)
    assert set(mapped) == {f.id for f in LEASE_REPAYMENT_COMPONENT_CATALOG}


def test_zero_revenue_undefined_ratio():
    fin = _tiny(repayments=(-50.0, -80.0), revenue=(0.0, 2000.0))
    series = compute_lease_repayment_series(
        fin, list(canonical_fiscal_periods(fin)), _anchor(fin)
    )
    assert series.lease_repayments == (50.0, 80.0)
    assert series.lease_repayments_to_revenue[0] == UNDEFINED_RATIO
    assert series.lease_repayments_to_revenue[1] == pytest.approx(80.0 / 2000.0)


def test_missing_none_inputs_raise():
    missing = _tiny(missing_period=True)
    periods = list(canonical_fiscal_periods(missing))
    anchor = compute_anchor(missing, periods)
    with pytest.raises(MissingHistoricalValueError):
        compute_lease_repayment_series(missing, periods, anchor)

    none_period = _tiny(none_period=True)
    periods_n = list(canonical_fiscal_periods(none_period))
    anchor_n = compute_anchor(none_period, periods_n)
    with pytest.raises(MissingHistoricalValueError):
        compute_lease_repayment_series(none_period, periods_n, anchor_n)

    missing_rev = _tiny(missing_revenue=True)
    periods_r = list(canonical_fiscal_periods(missing_rev))
    with pytest.raises(MissingHistoricalValueError):
        compute_anchor(missing_rev, periods_r)

    none_rev = _tiny(none_revenue=True)
    periods_nr = list(canonical_fiscal_periods(none_rev))
    with pytest.raises(MissingHistoricalValueError):
        compute_anchor(none_rev, periods_nr)


def test_period_ordering_unchanged_input_and_single_period():
    fin = _tiny(repayments=(-10.0, -20.0))
    reverse = [P2, P1]
    anchor = compute_anchor(fin, reverse)
    series = compute_lease_repayment_series(fin, reverse, anchor)
    assert series.repayments_reported == (-20.0, -10.0)
    assert series.lease_repayments == (20.0, 10.0)

    before = copy.deepcopy(resolve_lease_repayment_source(fin).values)
    compute_lease_repayment_series(fin, [P1, P2], compute_anchor(fin, [P1, P2]))
    after = resolve_lease_repayment_source(fin).values
    assert after == before

    single = _tiny(single_period=True, repayments=(-40.0, -50.0), revenue=(800.0, 900.0))
    assert lease_repayment_applicable(single)
    builder = ReferenceModelBuilder(single)
    assert builder.lease_repayment_series is not None
    assert builder.lease_repayment_series.lease_repayments == (40.0,)
    assert len(builder.lease_repayment_specs) == 2


def test_independent_of_liability_rou_interest():
    """Repayment activation does not require liability, ROU, or lease interest."""
    fin = _tiny()
    assert lease_repayment_applicable(fin)
    assert fin.historical_lease is None
    builder = ReferenceModelBuilder(fin)
    assert builder.lease_repayment_specs
    assert builder.lease_liability_specs == ()
    assert builder.lease_rou_specs == ()


def _dupont_row_by_label(ws, label: str) -> int:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"ALT DuPont label not found: {label!r}")


def test_workbook_gating_structure_formulas_notes_and_check(tmp_path):
    data = _tiny()
    trainer, answer = build_training_workbook(data, tmp_path / "LEASE_REPAY_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.lease_repayment_series is not None
    assert len(builder.lease_repayment_specs) == 4
    assert {s.family_id for s in builder.lease_repayment_specs} == {
        "lease_repayments",
        "lease_repayments_to_revenue",
    }

    smap = load_semantic_map(answer)
    repay_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in LEASE_REPAYMENT_COMPONENT_CATALOG}
    ]
    assert len(repay_comps) == 4

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "LEASE REPAYMENT CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        reported_row = _dupont_row_by_label(
            ws, "Repayments of Lease Liabilities (reported)"
        )
        repay_row = _dupont_row_by_label(ws, "Lease Repayments (−reported)")
        ratio_row = _dupont_row_by_label(ws, "Lease Repayments / Revenue")
        assert isinstance(ws.cell(reported_row, 2).value, str)
        assert ws.cell(reported_row, 2).value.startswith("=")
        for col in (2, 3):
            reported = ws.cell(reported_row, column=col)
            repay_cell = ws.cell(repay_row, column=col)
            ratio_cell = ws.cell(ratio_row, column=col)
            if path == answer:
                assert isinstance(repay_cell.value, str) and repay_cell.value.startswith(
                    "="
                )
                assert repay_cell.value.replace(" ", "").startswith("=-")
                assert "IF(" in str(ratio_cell.value) and "NA()" in str(ratio_cell.value)
                assert repay_cell.comment is not None
                note = (repay_cell.comment.text or "").strip().lower()
                assert note
                assert "rou" in note or "liability" in note or "total lease" in note
                assert ratio_cell.comment is not None
                assert (ratio_cell.comment.text or "").strip()
            else:
                assert repay_cell.value is None
                assert ratio_cell.value is None
                assert repay_cell.comment is None
                assert ratio_cell.comment is None
                assert reported.value is not None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in repay_comps if c.family_id == "lease_repayments")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Lease Repayments" not in dumped


def test_workbook_renamed_reordered_source_links(tmp_path):
    for repayments_first in (False, True):
        fin = _tiny(repayments_first=repayments_first)
        fin.cash_flow[
            next(
                i
                for i, item in enumerate(fin.cash_flow)
                if (item.concept or "") == "repayments_of_lease_liabilities"
            )
        ] = _li(
            "Lease liability principal repaid (renamed)",
            {P1: -100.0, P2: -120.0},
            concept="repayments_of_lease_liabilities",
        )
        trainer, answer = build_training_workbook(
            fin, tmp_path / f"LEASE_REPAY_REORDER_{repayments_first}.xlsx"
        )
        idx = next(
            i
            for i, item in enumerate(fin.cash_flow)
            if (item.concept or "") == "repayments_of_lease_liabilities"
        )
        expected_row = 7 + idx
        wb = load_workbook(answer, data_only=False)
        ws = wb["ALT DuPont"]
        reported_row = _dupont_row_by_label(
            ws, "Repayments of Lease Liabilities (reported)"
        )
        level_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
        assert f"'CashFlowStatement'!B{expected_row}" in level_f or (
            f"'Cash Flow Statement'!B{expected_row}" in level_f
        )
        wb.close()


def test_undefined_ratio_check_accepts_na(tmp_path):
    data = _tiny(repayments=(-50.0, -80.0), revenue=(0.0, 2000.0))
    trainer, answer = build_training_workbook(data, tmp_path / "LEASE_REPAY_NA.xlsx")
    smap = load_semantic_map(answer)
    undef = [
        c
        for c in smap.all_ordered()
        if c.family_id == "lease_repayments_to_revenue"
        and c.expected_value == UNDEFINED_RATIO
    ]
    assert undef
    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    for comp in undef:
        _inject_formula_and_cached_value(
            trainer,
            comp.tab,
            comp.cell,
            formula=comp.formula,
            cached_value="#N/A",
        )
    filled = check_workbook(trainer)
    assert filled.incorrect == 0
    assert filled.blank == 0
    assert filled.correct == filled.total


def test_demo_omits_and_practice_counts_unchanged(tmp_path):
    data = _ingest_demo()
    assert not lease_repayment_applicable(data)
    builder = ReferenceModelBuilder(data)
    assert builder.lease_repayment_series is None
    assert builder.lease_repayment_specs == ()
    trainer, answer = build_training_workbook(data, tmp_path / "LEASE_REPAY_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312
    assert check_workbook(trainer).blank == 312


def test_fast_retailing_fy2021_fy2025_repayment_anchors_and_ratios():
    payload = json.loads(STD_JSON.read_text(encoding="utf-8"))
    fin = standardized_from_payload(payload)
    assert lease_repayment_applicable(fin)
    periods = list(canonical_fiscal_periods(fin))
    assert [p.isoformat() for p in periods] == [
        "2021-08-31",
        "2022-08-31",
        "2023-08-31",
        "2024-08-31",
        "2025-08-31",
    ]
    anchor = compute_anchor(fin, periods)
    series = compute_lease_repayment_series(fin, periods, anchor)
    assert series.repayments_reported == (
        -148248.0,
        -136889.0,
        -140646.0,
        -146403.0,
        -140483.0,
    )
    assert series.lease_repayments == (
        148248.0,
        136889.0,
        140646.0,
        146403.0,
        140483.0,
    )
    revenues = (2132992.0, 2301122.0, 2766557.0, 3103836.0, 3400539.0)
    assert list(anchor.historical.revenue) == list(revenues)
    for j, rev in enumerate(revenues):
        assert series.lease_repayments_to_revenue[j] == pytest.approx(
            series.lease_repayments[j] / rev
        )

    builder = ReferenceModelBuilder(fin)
    assert len(builder.lease_repayment_specs) == 10
    assert len(builder.capex_specs) == 20
    assert len(builder.lease_liability_specs) == 18
    assert len(builder.lease_rou_specs) == 16
    assert len(builder.expected_specs) == 577

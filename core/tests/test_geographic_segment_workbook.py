"""Step 9M.2.4.1.1.1.35 — geographic workbook schedules and learner/Check."""

from __future__ import annotations

import copy
import hashlib
import json
import shutil
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from core.data.historical_segments import (
    FAMILY_CORPORATE,
    FAMILY_ITEMIZED,
    GEOGRAPHIC_SEGMENT_NAMESPACE,
    IFOP_CORPORATE,
    OP_ADD,
    OP_SUBTRACT,
    SEGMENTS,
)
from core.data.interface import DocumentManifest, DocumentType, FinancialPeriod, HistoricalSegmentData, LineItem
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    GEOGRAPHIC_SEGMENT_COMPONENT_CATALOG,
    GEOGRAPHIC_SEGMENT_IDENTITIES,
    GEOGRAPHIC_SHEET_NAME,
    SemanticCellRef,
    expand_geographic_segment_specs,
    geographic_component_id,
    geographic_identity_label,
    geographic_spec_identity,
    resolve_geographic_adjacent_change_formula,
    resolve_geographic_consolidated_operating_margin_formula,
    resolve_geographic_consolidated_revenue_growth_formula,
    resolve_geographic_operating_margin_contribution_change_residual_formula,
    resolve_geographic_operating_margin_contribution_formula,
    resolve_geographic_operating_margin_contribution_residual_formula,
    resolve_geographic_operating_margin_mix_effect_formula,
    resolve_geographic_operating_margin_mix_within_residual_formula,
    resolve_geographic_operating_margin_within_segment_effect_formula,
    resolve_geographic_operating_profit_amount_change_residual_formula,
    resolve_geographic_operating_profit_incremental_margin_formula,
    resolve_geographic_operating_profit_margin_effect_formula,
    resolve_geographic_operating_profit_revenue_effect_formula,
    resolve_geographic_reconciling_operating_margin_contribution_formula,
    resolve_geographic_reconciling_operating_profit_amount_change_formula,
    resolve_geographic_revenue_growth_contribution_formula,
    resolve_geographic_revenue_growth_contribution_residual_formula,
    semantic_formula_cell,
)
from core.engine.reference_model import (
    GEOGRAPHIC_SHEET,
    JUDGMENT_SHEET,
    NORMALIZATION_JUDGMENT_SHEET,
    ReferenceModelBuilder,
)
from core.ingestion.filing_reconciler import reconcile_filings
from core.ingestion.filing_standardizer import standardize_reconciled
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.geographic_segment import (
    GEOGRAPHIC_RATIO_TOLERANCE,
    compute_geographic_segment_series,
    geographic_segment_applicable,
)
from core.model.historical_expected import geographic_expected_value_for_component
from core.model.line_resolver import MissingLineError
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import SOURCE_UNAVAILABLE, UNDEFINED_RATIO
from core.tests.test_capex import P1, P2, _tiny
from core.tests.test_geographic_segment_analysis import (
    ADMIT_2022,
    FY2026,
    _assert_series_matches,
    _independent_from_snapshots,
    _reload_admitted_lululemon,
)
from core.tests.test_geographic_segment_facts import _mutate_fy2025_prior_2025
from core.tests.test_historical_segment import (
    _corp_values,
    _itemized_values,
    _snapshot,
)
from core.tests.test_learner_ready_presentation import (
    WHITE_RGBS,
    _assert_answer_key_no_yellow,
    _assert_fresh_visible_style,
    _judgment_response_keys,
)
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import (
    BLANK_RGB,
    CORRECT_RGB,
    INCORRECT_RGB,
    check_workbook,
)
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook

ROOT = Path(__file__).resolve().parents[2]
FR_JSON = ROOT / "benchmark" / "fast_retailing" / "reconciled" / "standardized.json"
LULU_JSON = ROOT / "benchmark" / "lululemon" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
P0 = date(2023, 12, 31)
LEASE_DT_LULULEMON_SPECS = 486
FAST_RETAILING_SPECS = 577
GEOGRAPHIC_LULULEMON_SPECS_BASELINE = 74
GEOGRAPHIC_CONTRIBUTION_SPECS = 20
GEOGRAPHIC_MARGIN_BRIDGE_SPECS = 54
GEOGRAPHIC_MIX_WITHIN_SPECS = 28
GEOGRAPHIC_AMOUNT_CHANGE_SPECS = 24
GEOGRAPHIC_PROFIT_EFFECT_SPECS = 24
GEOGRAPHIC_INCREMENTAL_MARGIN_SPECS = 16
GEOGRAPHIC_LULULEMON_SPECS = (
    GEOGRAPHIC_LULULEMON_SPECS_BASELINE
    + GEOGRAPHIC_CONTRIBUTION_SPECS
    + GEOGRAPHIC_MARGIN_BRIDGE_SPECS
    + GEOGRAPHIC_MIX_WITHIN_SPECS
    + GEOGRAPHIC_AMOUNT_CHANGE_SPECS
    + GEOGRAPHIC_PROFIT_EFFECT_SPECS
    + GEOGRAPHIC_INCREMENTAL_MARGIN_SPECS
)
LULULEMON_UNAVAILABLE_DISPLAYS = 101
GEOGRAPHIC_A2 = (
    "Source-supported geographic revenue mix, adjacent-period growth, "
    "reported operating margins, consolidated bridges, percentage-point "
    "contributions to consolidated revenue growth, an arithmetic "
    "decomposition of consolidated reported operating margin and its "
    "change, a separate mix and within-segment decomposition of "
    "that adjacent change, an adjacent operating-profit amount-"
    "change bridge, and incremental reported operating margins. "
    "Reported operating margin is distinct from BAV NOPAT margin. "
    "Revenue-growth contributions are an arithmetic decomposition of "
    "reported geographic revenue changes, not organic, "
    "constant-currency, or causal growth. Operating-margin "
    "contributions are a direct contribution bridge, distinct from the "
    "mix and within-segment decomposition, the monetary amount-change "
    "bridge, midpoint revenue and margin effects on operating-profit "
    "change, incremental reported operating margins, normalization, "
    "or a causal explanation. Mix and within-segment effects use a "
    "symmetric midpoint convention and remain arithmetic only. "
    "Operating-profit amount changes are monetary differences, "
    "distinct from the percentage-point margin bridges. Revenue and "
    "margin effects allocate each eligible segment operating-profit "
    "amount change by a monetary midpoint convention. Incremental "
    "reported operating margins divide adjacent operating-profit "
    "amount changes by adjacent revenue changes for each segment and "
    "for reported consolidated totals, including signed reconciling "
    "items in consolidated profit, and are distinct from reported "
    "operating margin."
)
GEOGRAPHIC_A4 = (
    "Calculated segment totals are distinct from any reported segment_total. "
    "Sparse unavailable amounts remain unavailable; opening growth, "
    "opening revenue-growth contributions, opening operating-margin "
    "contribution changes, opening mix and within-segment effects, "
    "opening operating-profit amount changes, opening revenue "
    "and margin effects on operating-profit change, and opening "
    "incremental reported operating margins "
    "are not practiced. Signed residuals are not forced to zero."
)


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _attach_geo(fin, *snapshots):
    fin.historical_segment = HistoricalSegmentData(
        namespace=GEOGRAPHIC_SEGMENT_NAMESPACE,
        periods=list(snapshots),
    )
    return fin


def _geo_tiny(*snapshots, extra_opening: date | None = None):
    fin = _tiny(with_payments=False)
    if extra_opening is not None:
        fin.periods.insert(
            0, FinancialPeriod(end_date=extra_opening, label="FY2023")
        )
        for item in (
            *fin.income_statement,
            *fin.balance_sheet,
            *fin.cash_flow,
        ):
            first = fin.periods[1].end_date
            if first in item.values:
                item.values[extra_opening] = item.values[first]
    if snapshots:
        _attach_geo(fin, *snapshots)
        revenue = next(item for item in fin.income_statement if item.concept == "revenue")
        ifop_values: dict[date, float] = {}
        for snap in snapshots:
            revenue.values[snap.period] = float(snap.values["net_revenue.consolidated"])
            ifop_values[snap.period] = float(
                snap.values["income_from_operations.consolidated"]
            )
        if extra_opening is not None and P1 in ifop_values:
            ifop_values[extra_opening] = ifop_values[P1]
        fin.income_statement.insert(
            1,
            LineItem(
                label="Income from operations",
                concept="operating_income",
                values=ifop_values,
            ),
        )
    return fin


def _row_by_label(ws, label: str) -> int:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row, 1).value == label:
            return row
    raise AssertionError(f"missing geographic label {label!r}")


# Visible-cell teaching text that previously disclosed active geographic practice.
_DISCLOSED_GEOGRAPHIC_ARITHMETIC = (
    "income from operations / net revenue",
    "income from operations / segment net revenue",
    "segment net revenue / reported consolidated",
    "(current − prior) / prior",
    "(current - prior) / prior",
    "corporate-column add",
    "itemized subtract",
    "100 × (current segment",
    "100*(",
    "100 × segment income",
    "segment income from operations / reported consolidated",
)


def _visible_cell_texts(wb) -> list[tuple[str, str, str]]:
    found: list[tuple[str, str, str]] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        max_row = ws.max_row or 1
        max_col = ws.max_column or 1
        for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                value = cell.value
                if isinstance(value, str) and not value.startswith("="):
                    found.append((ws.title, cell.coordinate, value))
                comment = cell.comment.text if cell.comment is not None else ""
                if comment:
                    found.append((ws.title, f"{cell.coordinate}#comment", comment))
    return found


def _assert_no_disclosed_geographic_arithmetic(wb) -> None:
    for sheet, coord, text in _visible_cell_texts(wb):
        lowered = text.lower()
        for fragment in _DISCLOSED_GEOGRAPHIC_ARITHMETIC:
            assert fragment not in lowered, (
                f"{sheet}!{coord} discloses geographic practice arithmetic: {text!r}"
            )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _count_source_unavailable(path: Path) -> int:
    wb = load_workbook(path, data_only=False)
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == SOURCE_UNAVAILABLE:
                    n += 1
    wb.close()
    return n


def _visible_inventory(wb) -> tuple[list[str], int, list[tuple[str, str, str]]]:
    sheets: list[str] = []
    cells_scanned = 0
    texts = _visible_cell_texts(wb)
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        sheets.append(ws.title)
        max_row = ws.max_row or 1
        max_col = ws.max_column or 1
        cells_scanned += max_row * max_col
    return sheets, cells_scanned, texts


def _geographic_answer_guidance() -> tuple[str, ...]:
    texts: list[str] = []
    for family in GEOGRAPHIC_SEGMENT_COMPONENT_CATALOG:
        if family.short_hint:
            texts.append(family.short_hint)
        texts.extend(family.hints)
    return tuple(texts)


def _assert_trainer_complete_text_undisclosed(wb) -> tuple[list[str], int, int]:
    sheets, cells_scanned, texts = _visible_inventory(wb)
    assert GEOGRAPHIC_SHEET in sheets
    assert "Trainer" in sheets
    assert cells_scanned > 0
    blob = "\n".join(text for _sheet, _coord, text in texts)
    lowered_blob = blob.lower()
    _assert_no_disclosed_geographic_arithmetic(wb)
    for sheet, coord, text in texts:
        if coord.endswith("#comment"):
            raise AssertionError(
                f"{sheet}!{coord} retains a visible Note on the Trainer: {text!r}"
            )
    for guidance in _geographic_answer_guidance():
        assert guidance.lower() not in lowered_blob, (
            f"Trainer visible text discloses Answer-Key guidance: {guidance!r}"
        )
    return sheets, cells_scanned, len(texts)


def _assert_margin_note_explains_reported_basis(note: str) -> None:
    text = (note or "").strip()
    assert text, "margin Note is empty"
    lowered = text.lower()
    assert lowered != "nopat"
    assert "nopat" in lowered
    assert lowered.replace("nopat", "").strip()
    assert "income from operations" in lowered
    assert "net revenue" in lowered
    assert "/" in text or "divided" in lowered
    assert "bav" in lowered
    assert "distinct" in lowered or "not" in lowered


def _copy_pair(trainer: Path, answer: Path, dest: Path) -> tuple[Path, Path]:
    dest.mkdir()
    copied_trainer = dest / trainer.name
    copied_answer = dest / answer.name
    shutil.copy2(trainer, copied_trainer)
    shutil.copy2(answer, copied_answer)
    for sidecar in (
        answer.with_suffix(".component_map.json"),
        answer.with_suffix(".assumptions.json"),
        answer.with_suffix(".trainer.json"),
    ):
        if sidecar.is_file():
            shutil.copy2(sidecar, dest / sidecar.name)
    return copied_trainer, copied_answer


def _assert_visible_parity(
    trainer: Path,
    answer: Path,
    practice_cells: set[tuple[str, str]] | None = None,
) -> None:
    twb = load_workbook(trainer, data_only=False)
    awb = load_workbook(answer, data_only=False)
    skip = set(practice_cells or ())
    skip |= _judgment_response_keys(twb)
    skip |= _judgment_response_keys(awb)
    t_visible = [ws.title for ws in twb.worksheets if ws.sheet_state == "visible"]
    a_visible = [ws.title for ws in awb.worksheets if ws.sheet_state == "visible"]
    assert t_visible == a_visible
    for name in t_visible:
        tws = twb[name]
        aws = awb[name]
        max_row = max(tws.max_row or 1, aws.max_row or 1)
        max_col = max(tws.max_column or 1, aws.max_column or 1)
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                tcell = tws.cell(row, col)
                acell = aws.cell(row, col)
                if (name, tcell.coordinate) in skip:
                    continue
                assert tcell.value == acell.value, f"{name}!{tcell.coordinate}"
                t_note = tcell.comment.text if tcell.comment is not None else None
                a_note = acell.comment.text if acell.comment is not None else None
                assert t_note == a_note, f"{name}!{tcell.coordinate} note"
    twb.close()
    awb.close()


def test_catalog_orders_and_expand_identities():
    assert GEOGRAPHIC_SHEET == GEOGRAPHIC_SHEET_NAME
    assert [family.order for family in GEOGRAPHIC_SEGMENT_COMPONENT_CATALOG] == list(
        range(152, 183)
    )
    assert [family.id for family in GEOGRAPHIC_SEGMENT_COMPONENT_CATALOG] == [
        "geographic_revenue_share",
        "geographic_revenue_growth",
        "geographic_reported_operating_margin",
        "geographic_calculated_segment_revenue_total",
        "geographic_calculated_segment_operating_profit_total",
        "geographic_signed_reconciling_contribution",
        "geographic_reconstructed_consolidated_operating_profit",
        "geographic_consolidated_revenue_difference",
        "geographic_consolidated_operating_profit_difference",
        "geographic_revenue_growth_contribution",
        "geographic_consolidated_revenue_growth",
        "geographic_revenue_growth_contribution_residual",
        "geographic_operating_margin_contribution",
        "geographic_reconciling_operating_margin_contribution",
        "geographic_consolidated_operating_margin",
        "geographic_operating_margin_contribution_residual",
        "geographic_operating_margin_contribution_change",
        "geographic_reconciling_operating_margin_contribution_change",
        "geographic_consolidated_operating_margin_change",
        "geographic_operating_margin_contribution_change_residual",
        "geographic_operating_margin_mix_effect",
        "geographic_operating_margin_within_segment_effect",
        "geographic_operating_margin_mix_within_residual",
        "geographic_operating_profit_amount_change",
        "geographic_reconciling_operating_profit_amount_change",
        "geographic_consolidated_operating_profit_amount_change",
        "geographic_operating_profit_amount_change_residual",
        "geographic_operating_profit_revenue_effect",
        "geographic_operating_profit_margin_effect",
        "geographic_operating_profit_incremental_margin",
        "geographic_consolidated_operating_profit_incremental_margin",
    ]
    specs = expand_geographic_segment_specs(
        [P1, P2],
        start_order=10,
        available_periods=(P1, P2),
        growth_identities={P2: SEGMENTS},
        bridge_identities={
            P1: (
                "ifop_reconciling.amortization_of_intangible_assets",
                "ifop_reconciling.general_corporate_expenses",
            ),
            P2: (IFOP_CORPORATE,),
        },
        contribution_identities={P2: SEGMENTS},
        consolidated_growth_periods=(P2,),
        margin_change_periods=(P2,),
    )
    assert geographic_component_id("geographic_revenue_share", P1, "americas") in {
        s.id for s in specs
    }
    assert not any(
        s.family_id == "geographic_revenue_growth" and s.period_end == P1.isoformat()
        for s in specs
    )
    assert not any(
        s.family_id == "geographic_revenue_growth_contribution"
        and s.period_end == P1.isoformat()
        for s in specs
    )
    assert geographic_component_id(
        "geographic_revenue_growth_contribution", P2, "americas"
    ) in {s.id for s in specs}
    assert geographic_component_id(
        "geographic_operating_margin_contribution", P1, "americas"
    ) in {s.id for s in specs}
    assert not any(
        s.family_id == "geographic_operating_margin_contribution_change"
        and s.period_end == P1.isoformat()
        for s in specs
    )
    assert geographic_component_id(
        "geographic_operating_margin_contribution_change", P2, "americas"
    ) in {s.id for s in specs}
    assert geographic_component_id(
        "geographic_operating_margin_mix_effect", P2, "americas"
    ) in {s.id for s in specs}
    assert geographic_component_id(
        "geographic_operating_profit_amount_change", P2, "americas"
    ) in {s.id for s in specs}
    assert geographic_component_id(
        "geographic_operating_profit_revenue_effect", P2, "americas"
    ) in {s.id for s in specs}
    assert geographic_component_id(
        "geographic_operating_profit_margin_effect", P2, "americas"
    ) in {s.id for s in specs}
    assert geographic_component_id(
        "geographic_operating_profit_incremental_margin", P2, "americas"
    ) in {s.id for s in specs}
    assert geographic_component_id(
        "geographic_consolidated_operating_profit_incremental_margin", P2
    ) in {s.id for s in specs}
    assert geographic_component_id(
        "geographic_operating_margin_within_segment_effect", P2, "americas"
    ) in {s.id for s in specs}
    assert not any(
        s.family_id == "geographic_operating_margin_mix_effect"
        and s.period_end == P1.isoformat()
        for s in specs
    )
    assert not any(
        s.family_id == "geographic_operating_profit_amount_change"
        and s.period_end == P1.isoformat()
        for s in specs
    )
    assert not any(
        s.family_id == "geographic_operating_profit_revenue_effect"
        and s.period_end == P1.isoformat()
        for s in specs
    )
    assert not any(
        s.family_id == "geographic_operating_profit_incremental_margin"
        and s.period_end == P1.isoformat()
        for s in specs
    )
    assert not any(
        s.family_id == "geographic_consolidated_operating_profit_incremental_margin"
        and s.period_end == P1.isoformat()
        for s in specs
    )
    mix_residual = next(
        s
        for s in specs
        if s.family_id == "geographic_operating_margin_mix_within_residual"
    )
    assert mix_residual.period_end == P2.isoformat()
    amount_residual = next(
        s
        for s in specs
        if s.family_id == "geographic_operating_profit_amount_change_residual"
    )
    assert amount_residual.period_end == P2.isoformat()
    assert geographic_component_id(
        "geographic_consolidated_operating_profit_amount_change", P2
    ) in amount_residual.depends_on
    assert geographic_component_id(
        "geographic_reconciling_operating_profit_amount_change", P2
    ) in amount_residual.depends_on
    assert geographic_component_id(
        "geographic_consolidated_operating_margin_change", P2
    ) in mix_residual.depends_on
    assert geographic_component_id(
        "geographic_reconciling_operating_margin_contribution_change", P2
    ) in mix_residual.depends_on
    change_residual = next(
        s
        for s in specs
        if s.family_id == "geographic_operating_margin_contribution_change_residual"
    )
    assert change_residual.period_end == P2.isoformat()
    assert geographic_component_id(
        "geographic_consolidated_operating_margin_change", P2
    ) in change_residual.depends_on
    residual = next(
        s
        for s in specs
        if s.family_id == "geographic_revenue_growth_contribution_residual"
    )
    assert residual.period_end == P2.isoformat()
    assert geographic_component_id(
        "geographic_consolidated_revenue_growth", P2
    ) in residual.depends_on
    signed = [
        s
        for s in specs
        if s.family_id == "geographic_signed_reconciling_contribution"
    ]
    assert {s.id for s in signed} == {
        geographic_component_id(
            "geographic_signed_reconciling_contribution",
            P1,
            "ifop_reconciling.amortization_of_intangible_assets",
        ),
        geographic_component_id(
            "geographic_signed_reconciling_contribution",
            P1,
            "ifop_reconciling.general_corporate_expenses",
        ),
        geographic_component_id(
            "geographic_signed_reconciling_contribution",
            P2,
            IFOP_CORPORATE,
        ),
    }
    with pytest.raises(ValueError, match="duplicate fiscal periods"):
        expand_geographic_segment_specs(
            [P1, P1],
            start_order=1,
            available_periods=(P1,),
            growth_identities={},
            bridge_identities={},
        )
    with pytest.raises(ValueError, match="strictly chronological"):
        expand_geographic_segment_specs(
            [P2, P1],
            start_order=1,
            available_periods=(P1, P2),
            growth_identities={},
            bridge_identities={},
        )


def test_absent_null_payloads_add_no_schedule_or_practice(tmp_path):
    absent = _tiny(with_payments=False)
    assert not geographic_segment_applicable(absent)
    builder = ReferenceModelBuilder(absent)
    assert builder.geographic_series is None
    assert builder.geographic_specs == ()
    trainer, answer = build_training_workbook(absent, tmp_path / "GEO_ABSENT.xlsx")
    for path in (trainer, answer):
        wb = load_workbook(path)
        assert GEOGRAPHIC_SHEET not in wb.sheetnames
        wb.close()
    smap = load_semantic_map(answer)
    assert not any(c.category == "geographic_segment" for c in smap.all_ordered())

    payload = standardized_to_payload(_tiny(with_payments=False))
    payload["historical_segment"] = None
    nulled = standardized_from_payload(payload)
    assert nulled.historical_segment is None
    assert ReferenceModelBuilder(nulled).geographic_specs == ()


def test_invalid_contract_fails_closed_without_mutation():
    fin = _geo_tiny(_snapshot(P1), _snapshot(P2))
    original = copy.deepcopy(fin.historical_segment)
    fin.historical_segment.periods[0].bridge_operations = {IFOP_CORPORATE: OP_SUBTRACT}
    mutated = copy.deepcopy(fin.historical_segment)
    with pytest.raises(ValueError):
        ReferenceModelBuilder(fin)
    assert fin.historical_segment == mutated
    fin.historical_segment = copy.deepcopy(original)
    with pytest.raises(MissingLineError):
        compute_geographic_segment_series(_tiny(with_payments=False))


def test_workbook_gating_formulas_notes_check_and_families(tmp_path):
    fin = _geo_tiny(
        _snapshot(P1, family=FAMILY_ITEMIZED, values=_itemized_values()),
        _snapshot(P2, family=FAMILY_CORPORATE, values=_corp_values()),
    )
    builder = ReferenceModelBuilder(fin)
    assert builder.geographic_series is not None
    geo_ids = {s.family_id for s in builder.geographic_specs}
    assert geo_ids == {f.id for f in GEOGRAPHIC_SEGMENT_COMPONENT_CATALOG}
    existing = [
        s.id
        for s in builder.expected_specs
        if s.category != "geographic_segment"
    ]
    assert existing
    trainer, answer = build_training_workbook(fin, tmp_path / "GEO_BASE.xlsx")
    smap = load_semantic_map(answer)
    geo_comps = [
        c for c in smap.all_ordered() if c.category == "geographic_segment"
    ]
    assert {c.id for c in geo_comps} == {s.id for s in builder.geographic_specs}
    assert not any(
        c.family_id == "geographic_revenue_growth" and c.period_index == 0
        for c in geo_comps
    )
    assert not any(
        c.family_id == "geographic_revenue_growth_contribution" and c.period_index == 0
        for c in geo_comps
    )
    assert any(
        c.family_id == "geographic_operating_margin_contribution" and c.period_index == 0
        for c in geo_comps
    )
    assert not any(
        c.family_id == "geographic_operating_margin_contribution_change"
        and c.period_index == 0
        for c in geo_comps
    )
    assert not any(
        c.family_id == "geographic_operating_margin_mix_effect"
        and c.period_index == 0
        for c in geo_comps
    )
    assert any(
        c.family_id == "geographic_operating_margin_mix_effect" and c.period_index == 1
        for c in geo_comps
    )
    assert not any(
        c.family_id == "geographic_operating_profit_amount_change"
        and c.period_index == 0
        for c in geo_comps
    )
    assert any(
        c.family_id == "geographic_operating_profit_amount_change"
        and c.period_index == 1
        for c in geo_comps
    )
    assert not any(
        c.family_id == "geographic_operating_profit_revenue_effect"
        and c.period_index == 0
        for c in geo_comps
    )
    assert any(
        c.family_id == "geographic_operating_profit_revenue_effect"
        and c.period_index == 1
        for c in geo_comps
    )
    assert any(
        c.family_id == "geographic_operating_profit_margin_effect"
        and c.period_index == 1
        for c in geo_comps
    )
    assert not any(
        c.family_id == "geographic_operating_profit_incremental_margin"
        and c.period_index == 0
        for c in geo_comps
    )
    assert any(
        c.family_id == "geographic_operating_profit_incremental_margin"
        and c.period_index == 1
        for c in geo_comps
    )
    assert any(
        c.family_id == "geographic_consolidated_operating_profit_incremental_margin"
        and c.period_index == 1
        for c in geo_comps
    )

    awb = load_workbook(answer, data_only=False)
    twb = load_workbook(trainer, data_only=False)
    assert GEOGRAPHIC_SHEET in awb.sheetnames
    assert GEOGRAPHIC_SHEET in twb.sheetnames
    aws = awb[GEOGRAPHIC_SHEET]
    tws = twb[GEOGRAPHIC_SHEET]
    assert "Americas" in str(aws.cell(_row_by_label(aws, "Americas net revenue"), 1).value)
    assert "China Mainland" in str(
        aws.cell(_row_by_label(aws, "China Mainland net revenue"), 1).value
    )
    assert "Rest of World" in str(
        aws.cell(_row_by_label(aws, "Rest of World net revenue"), 1).value
    )
    assert aws["A2"].value == tws["A2"].value
    assert aws["A4"].value == tws["A4"].value
    assert "NOPAT" in str(aws["A2"].value)
    assert "income from operations / net revenue" not in str(aws["A2"].value).lower()
    assert "add" not in str(aws["A4"].value).lower()
    assert "subtract" not in str(aws["A4"].value).lower()
    _assert_no_disclosed_geographic_arithmetic(twb)
    assert str(aws["A3"].value).startswith("Units:")
    family_row = _row_by_label(aws, "Presentation family")
    assert aws.cell(family_row, 2).value == FAMILY_ITEMIZED
    assert aws.cell(family_row, 3).value == FAMILY_CORPORATE

    americas_rev = _row_by_label(aws, "Americas net revenue")
    assert aws.cell(americas_rev, 2).value == 50.0
    assert tws.cell(americas_rev, 2).value == 50.0
    calc_rev = _row_by_label(aws, "Calculated segment revenue total")
    assert str(aws.cell(calc_rev, 2).value).startswith("=")
    assert tws.cell(calc_rev, 2).value in (None, "")
    growth_row = _row_by_label(aws, "Americas adjacent-period revenue growth")
    assert aws.cell(growth_row, 2).value == "N/A"
    assert tws.cell(growth_row, 2).value == "N/A"
    assert str(aws.cell(growth_row, 3).value).startswith("=")
    assert "NA()" in str(aws.cell(growth_row, 3).value)
    contrib_row = _row_by_label(
        aws,
        "Americas contribution to consolidated revenue growth (percentage points)",
    )
    assert aws.cell(contrib_row, 2).value == "N/A"
    assert tws.cell(contrib_row, 2).value == "N/A"
    assert str(aws.cell(contrib_row, 3).value).startswith("=")
    assert "100*" in str(aws.cell(contrib_row, 3).value).replace(" ", "")
    cons_growth_row = _row_by_label(aws, "Consolidated revenue growth")
    assert aws.cell(cons_growth_row, 2).value == "N/A"
    residual_row = _row_by_label(aws, "Contribution residual (percentage points)")
    assert aws.cell(residual_row, 2).value == "N/A"
    cs_row = _row_by_label(
        aws,
        "Americas contribution to consolidated operating margin (percentage points)",
    )
    assert str(aws.cell(cs_row, 2).value).startswith("=")
    assert "100*" in str(aws.cell(cs_row, 2).value).replace(" ", "")
    assert tws.cell(cs_row, 2).value in (None, "")
    dcs_row = _row_by_label(
        aws,
        "Americas change in contribution to consolidated operating margin "
        "(percentage points)",
    )
    assert aws.cell(dcs_row, 2).value == "N/A"
    assert tws.cell(dcs_row, 2).value == "N/A"
    assert str(aws.cell(dcs_row, 3).value).startswith("=")
    mix_row = _row_by_label(
        aws,
        "Americas mix effect on consolidated operating-margin change "
        "(percentage points)",
    )
    assert aws.cell(mix_row, 2).value == "N/A"
    assert tws.cell(mix_row, 2).value == "N/A"
    assert str(aws.cell(mix_row, 3).value).startswith("=")
    assert "100*" in str(aws.cell(mix_row, 3).value).replace(" ", "")
    within_row = _row_by_label(
        aws,
        "Americas within-segment margin effect on consolidated "
        "operating-margin change (percentage points)",
    )
    assert aws.cell(within_row, 2).value == "N/A"
    assert str(aws.cell(within_row, 3).value).startswith("=")
    mix_res_row = _row_by_label(
        aws,
        "Operating-margin mix and within-segment decomposition residual "
        "(percentage points)",
    )
    assert aws.cell(mix_res_row, 2).value == "N/A"
    dp_row = _row_by_label(aws, "Americas operating-profit amount change")
    assert aws.cell(dp_row, 2).value == "N/A"
    assert tws.cell(dp_row, 2).value == "N/A"
    assert str(aws.cell(dp_row, 3).value).startswith("=")
    assert "100*" not in str(aws.cell(dp_row, 3).value).replace(" ", "")
    dbr_row = _row_by_label(aws, "Aggregate reconciling operating-profit amount change")
    assert aws.cell(dbr_row, 2).value == "N/A"
    assert str(aws.cell(dbr_row, 3).value).startswith("=")
    dpc_row = _row_by_label(aws, "Consolidated operating-profit amount change")
    assert aws.cell(dpc_row, 2).value == "N/A"
    dpr_row = _row_by_label(aws, "Operating-profit amount-change residual")
    assert aws.cell(dpr_row, 2).value == "N/A"
    rev_eff_row = _row_by_label(aws, "Americas revenue effect on operating-profit change")
    assert aws.cell(rev_eff_row, 2).value == "N/A"
    assert tws.cell(rev_eff_row, 2).value == "N/A"
    assert str(aws.cell(rev_eff_row, 3).value).startswith("=")
    assert "100*" not in str(aws.cell(rev_eff_row, 3).value).replace(" ", "")
    mgn_eff_row = _row_by_label(aws, "Americas margin effect on operating-profit change")
    assert aws.cell(mgn_eff_row, 2).value == "N/A"
    assert str(aws.cell(mgn_eff_row, 3).value).startswith("=")
    assert "100*" not in str(aws.cell(mgn_eff_row, 3).value).replace(" ", "")
    inc_row = _row_by_label(aws, "Americas incremental reported operating margin")
    assert aws.cell(inc_row, 2).value == "N/A"
    assert tws.cell(inc_row, 2).value == "N/A"
    assert str(aws.cell(inc_row, 3).value).startswith("=")
    assert "NA()" in str(aws.cell(inc_row, 3).value)
    assert "100*" not in str(aws.cell(inc_row, 3).value).replace(" ", "")
    cons_inc_row = _row_by_label(aws, "Consolidated incremental reported operating margin")
    assert aws.cell(cons_inc_row, 2).value == "N/A"
    assert str(aws.cell(cons_inc_row, 3).value).startswith("=")
    assert "NA()" in str(aws.cell(cons_inc_row, 3).value)
    mix_note = (
        (aws.cell(mix_row, 3).comment.text or "")
        if aws.cell(mix_row, 3).comment
        else ""
    )
    lowered_mix = mix_note.lower()
    assert "midpoint" in lowered_mix
    assert "equally" in lowered_mix
    assert "contribution" in lowered_mix
    assert "nopat" in lowered_mix
    assert "price" in lowered_mix
    assert "volume" in lowered_mix
    assert "organic" in lowered_mix
    assert "causal" in lowered_mix
    assert "normalization" in lowered_mix
    within_note = (
        (aws.cell(within_row, 3).comment.text or "")
        if aws.cell(within_row, 3).comment
        else ""
    )
    assert "midpoint" in within_note.lower()
    mix_res_note = (
        (aws.cell(mix_res_row, 3).comment.text or "")
        if aws.cell(mix_res_row, 3).comment
        else ""
    )
    assert "residual" in mix_res_note.lower()
    assert "contribution" in mix_res_note.lower()
    dp_note = (
        (aws.cell(dp_row, 3).comment.text or "")
        if aws.cell(dp_row, 3).comment
        else ""
    )
    lowered_dp = dp_note.lower()
    assert "amount" in lowered_dp
    assert "monetary" in lowered_dp or "current minus" in lowered_dp
    assert "percentage-point" in lowered_dp or "margin" in lowered_dp
    assert "mix" in lowered_dp
    assert "nopat" in lowered_dp
    assert "organic" in lowered_dp
    assert "causal" in lowered_dp or "causality" in lowered_dp
    assert "normalization" in lowered_dp
    dpr_note = (
        (aws.cell(dpr_row, 3).comment.text or "")
        if aws.cell(dpr_row, 3).comment
        else ""
    )
    lowered_dpr = dpr_note.lower()
    assert "residual" in lowered_dpr
    assert "amount" in lowered_dpr
    assert "reconstructed" in lowered_dpr
    assert "reported" in lowered_dpr
    assert "−(d current − d prior)" in lowered_dpr or "- (d current" in lowered_dpr
    assert "not the positive" in lowered_dpr
    assert "nopat" in lowered_dpr
    assert "normalization" in lowered_dpr
    assert "organic" in lowered_dpr
    assert "causal" in lowered_dpr or "causality" in lowered_dpr
    rev_eff_note = (
        (aws.cell(rev_eff_row, 3).comment.text or "")
        if aws.cell(rev_eff_row, 3).comment
        else ""
    )
    lowered_rev_eff = rev_eff_note.lower()
    assert "midpoint" in lowered_rev_eff
    assert "monetary" in lowered_rev_eff
    assert "amount change" in lowered_rev_eff
    assert "nopat" in lowered_rev_eff
    assert "volume" in lowered_rev_eff
    assert "price" in lowered_rev_eff
    assert "organic" in lowered_rev_eff
    assert "causal" in lowered_rev_eff or "causality" in lowered_rev_eff
    assert "normalization" in lowered_rev_eff
    mgn_eff_note = (
        (aws.cell(mgn_eff_row, 3).comment.text or "")
        if aws.cell(mgn_eff_row, 3).comment
        else ""
    )
    lowered_mgn_eff = mgn_eff_note.lower()
    assert "midpoint" in lowered_mgn_eff
    assert "monetary" in lowered_mgn_eff
    assert "amount change" in lowered_mgn_eff
    assert "nopat" in lowered_mgn_eff
    assert "volume" in lowered_mgn_eff
    assert "organic" in lowered_mgn_eff
    inc_note = (
        (aws.cell(inc_row, 3).comment.text or "")
        if aws.cell(inc_row, 3).comment
        else ""
    )
    lowered_inc = inc_note.lower()
    assert "change in reported operating profit" in lowered_inc
    assert "revenue change" in lowered_inc
    assert "reported operating margin" in lowered_inc
    assert "small" in lowered_inc
    assert "signed" in lowered_inc or "decline" in lowered_inc
    assert "nopat" in lowered_inc
    assert "organic" in lowered_inc
    assert "marginal cost" in lowered_inc
    assert "operating leverage" in lowered_inc
    assert "normalization" in lowered_inc
    cons_inc_note = (
        (aws.cell(cons_inc_row, 3).comment.text or "")
        if aws.cell(cons_inc_row, 3).comment
        else ""
    )
    lowered_cons_inc = cons_inc_note.lower()
    assert "reconcil" in lowered_cons_inc
    assert "segment" in lowered_cons_inc
    assert "nopat" in lowered_cons_inc
    assert "reported operating margin" in lowered_cons_inc
    cs_note = (
        (aws.cell(cs_row, 2).comment.text or "")
        if aws.cell(cs_row, 2).comment
        else ""
    )
    assert "percentage points" in cs_note.lower() or "100" in cs_note
    assert "consolidated" in cs_note.lower()
    assert "nopat" in cs_note.lower()
    assert "mix" in cs_note.lower()
    assert "cause" not in cs_note.lower() or "not" in cs_note.lower()
    e_row = _row_by_label(
        aws, "Operating-margin contribution residual (percentage points)"
    )
    e_note = (
        (aws.cell(e_row, 2).comment.text or "")
        if aws.cell(e_row, 2).comment
        else ""
    )
    assert "residual" in e_note.lower()
    assert "cause" not in e_note.lower() or "not" in e_note.lower()
    residual_note = (
        (aws.cell(residual_row, 3).comment.text or "")
        if aws.cell(residual_row, 3).comment
        else ""
    )
    assert "percentage points" in residual_note.lower() or "100" in residual_note
    assert "cause" not in residual_note.lower()
    contrib_note = (
        (aws.cell(contrib_row, 3).comment.text or "")
        if aws.cell(contrib_row, 3).comment
        else ""
    )
    assert "percentage points" in contrib_note.lower() or "100" in contrib_note
    assert "organic" in contrib_note.lower()
    assert "cause" not in contrib_note.lower() or "not" in contrib_note.lower()

    itemized_signed = _row_by_label(
        aws, "Amortization Of Intangible Assets signed contribution"
    )
    assert str(aws.cell(itemized_signed, 2).value).replace(" ", "").startswith("=-")
    assert aws.cell(itemized_signed, 3).value in (None, "")
    corp_signed = _row_by_label(aws, "Corporate Unallocated signed contribution")
    assert str(aws.cell(corp_signed, 3).value).replace(" ", "") == (
        f"={chr(67)}{_row_by_label(aws, 'Corporate Unallocated (reported)')}"
    ).replace(" ", "") or str(aws.cell(corp_signed, 3).value).startswith("=")
    corp_formula = str(aws.cell(corp_signed, 3).value).replace(" ", "")
    assert corp_formula.startswith("=") and not corp_formula.startswith("=-")
    recon = _row_by_label(aws, "Reconstructed consolidated operating profit")
    assert "+" in str(aws.cell(recon, 2).value)
    margin_row = _row_by_label(aws, "Americas reported operating margin")
    note = (aws.cell(margin_row, 2).comment.text or "") if aws.cell(margin_row, 2).comment else ""
    assert "NOPAT" in note
    mix_row = _row_by_label(aws, "Americas revenue mix")
    mix_note = (aws.cell(mix_row, 2).comment.text or "") if aws.cell(mix_row, 2).comment else ""
    assert "mix" in mix_note.lower()
    assert tws.cell(mix_row, 2).comment is None
    assert tws.cell(mix_row, 2).value is None
    awb.close()
    twb.close()

    for comp in geo_comps:
        looked = geographic_expected_value_for_component(
            builder.geographic_series, comp
        )
        if isinstance(looked, float) and isinstance(comp.expected_value, float):
            assert looked == pytest.approx(comp.expected_value, abs=GEOGRAPHIC_RATIO_TOLERANCE)
        else:
            assert looked == comp.expected_value

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in geo_comps if c.family_id == "geographic_revenue_share")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Americas revenue mix" not in dumped
    assert "0.5" not in dumped


def test_sparse_undefined_negative_reorder_and_source_edit(tmp_path):
    zero_itemized = _itemized_values()
    zero_itemized["net_revenue.americas"] = 0.0
    zero_itemized["net_revenue.consolidated"] = 50.0
    neg_corp = _corp_values(ifop=(-4.0, 2.0, 1.0), corporate=-10.0)
    fin = _geo_tiny(
        _snapshot(P1, family=FAMILY_ITEMIZED, values=zero_itemized),
        _snapshot(P2, family=FAMILY_CORPORATE, values=neg_corp),
        extra_opening=P0,
    )
    builder = ReferenceModelBuilder(fin)
    series = builder.geographic_series
    assert series is not None
    axis = canonical_fiscal_periods(fin)
    assert axis[0] == P0
    assert series.presentation_family[P0] == SOURCE_UNAVAILABLE
    assert series.revenue_share[P0]["americas"] == SOURCE_UNAVAILABLE
    assert series.revenue_growth[P0]["americas"] is None
    assert series.revenue_growth_contribution[P0]["americas"] is None
    assert series.operating_margin_contribution[P0]["americas"] == SOURCE_UNAVAILABLE
    assert series.operating_margin_contribution_change[P0]["americas"] is None
    assert series.operating_margin_contribution_change[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.operating_margin_mix_effect[P0]["americas"] is None
    assert series.operating_margin_mix_effect[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.operating_margin_within_segment_effect[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.operating_margin_mix_within_residual[P1] == SOURCE_UNAVAILABLE
    assert series.operating_profit_amount_change[P0]["americas"] is None
    assert series.operating_profit_amount_change[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.operating_profit_amount_change_residual[P1] == SOURCE_UNAVAILABLE
    assert series.operating_profit_revenue_effect[P0]["americas"] is None
    assert series.operating_profit_revenue_effect[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.operating_profit_margin_effect[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.operating_profit_incremental_margin[P0]["americas"] is None
    assert series.operating_profit_incremental_margin[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.consolidated_operating_profit_incremental_margin[P1] == SOURCE_UNAVAILABLE
    assert series.revenue_growth[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.revenue_growth_contribution[P1]["americas"] == SOURCE_UNAVAILABLE
    assert series.reported_operating_margin[P1]["americas"] == UNDEFINED_RATIO
    assert series.income_from_operations[P2]["americas"] == -4.0
    assert not any(
        s.period_end == P0.isoformat() for s in builder.geographic_specs
    )
    assert not any(
        s.family_id == "geographic_revenue_growth" and s.period_end == P1.isoformat()
        for s in builder.geographic_specs
    )
    assert not any(
        s.family_id == "geographic_revenue_growth_contribution"
        and s.period_end == P1.isoformat()
        for s in builder.geographic_specs
    )
    assert not any(
        s.family_id == "geographic_operating_margin_contribution"
        and s.period_end == P0.isoformat()
        for s in builder.geographic_specs
    )
    assert not any(
        s.family_id == "geographic_operating_margin_contribution_change"
        and s.period_end == P1.isoformat()
        for s in builder.geographic_specs
    )
    assert not any(
        s.family_id == "geographic_operating_margin_mix_effect"
        and s.period_end == P1.isoformat()
        for s in builder.geographic_specs
    )
    assert not any(
        s.family_id == "geographic_operating_profit_amount_change"
        and s.period_end == P1.isoformat()
        for s in builder.geographic_specs
    )
    assert not any(
        s.family_id == "geographic_operating_profit_revenue_effect"
        and s.period_end == P1.isoformat()
        for s in builder.geographic_specs
    )
    assert not any(
        s.family_id == "geographic_operating_profit_incremental_margin"
        and s.period_end == P1.isoformat()
        for s in builder.geographic_specs
    )
    assert any(
        s.family_id == "geographic_operating_margin_contribution"
        and s.period_end == P1.isoformat()
        for s in builder.geographic_specs
    )

    reversed_snaps = (
        _snapshot(P2, family=FAMILY_CORPORATE, values=neg_corp),
        _snapshot(P1, family=FAMILY_ITEMIZED, values=zero_itemized),
    )
    reordered = _geo_tiny(*reversed_snaps, extra_opening=P0)
    reordered.historical_segment.periods = list(
        reversed(list(reordered.historical_segment.periods))
    )
    restored = standardized_from_payload(standardized_to_payload(reordered))
    rebuilt = compute_geographic_segment_series(restored)
    assert rebuilt.reported_operating_margin[P1]["americas"] == UNDEFINED_RATIO
    assert rebuilt.income_from_operations[P2]["americas"] == -4.0

    trainer, answer = build_training_workbook(fin, tmp_path / "GEO_SPARSE.xlsx")
    smap = load_semantic_map(answer)
    geo_comps = [
        c for c in smap.all_ordered() if c.category == "geographic_segment"
    ]
    awb = load_workbook(answer, data_only=False)
    aws = awb[GEOGRAPHIC_SHEET]
    share_row = _row_by_label(aws, "Americas revenue mix")
    assert SOURCE_UNAVAILABLE in {
        aws.cell(share_row, 2).value,
        aws.cell(_row_by_label(aws, "Americas net revenue"), 2).value,
    }
    growth_row = _row_by_label(aws, "Americas adjacent-period revenue growth")
    assert aws.cell(growth_row, 2).value == "N/A"
    assert aws.cell(growth_row, 3).value == SOURCE_UNAVAILABLE
    contrib_row = _row_by_label(
        aws,
        "Americas contribution to consolidated revenue growth (percentage points)",
    )
    assert aws.cell(contrib_row, 2).value == "N/A"
    assert aws.cell(contrib_row, 3).value == SOURCE_UNAVAILABLE
    cs_row = _row_by_label(
        aws,
        "Americas contribution to consolidated operating margin (percentage points)",
    )
    assert aws.cell(cs_row, 2).value == SOURCE_UNAVAILABLE
    assert "NA()" in str(aws.cell(cs_row, 3).value)
    dcs_row = _row_by_label(
        aws,
        "Americas change in contribution to consolidated operating margin "
        "(percentage points)",
    )
    assert aws.cell(dcs_row, 2).value == "N/A"
    assert aws.cell(dcs_row, 3).value == SOURCE_UNAVAILABLE
    mix_row = _row_by_label(
        aws,
        "Americas mix effect on consolidated operating-margin change "
        "(percentage points)",
    )
    assert aws.cell(mix_row, 2).value == "N/A"
    assert aws.cell(mix_row, 3).value == SOURCE_UNAVAILABLE
    inc_row = _row_by_label(aws, "Americas incremental reported operating margin")
    assert aws.cell(inc_row, 2).value == "N/A"
    assert aws.cell(inc_row, 3).value == SOURCE_UNAVAILABLE
    cons_inc_row = _row_by_label(aws, "Consolidated incremental reported operating margin")
    assert aws.cell(cons_inc_row, 2).value == "N/A"
    assert aws.cell(cons_inc_row, 3).value == SOURCE_UNAVAILABLE
    margin_row = _row_by_label(aws, "Americas reported operating margin")
    assert "NA()" in str(aws.cell(margin_row, 3).value)
    awb.close()

    source_row = _row_by_label(
        load_workbook(answer, data_only=False)[GEOGRAPHIC_SHEET],
        "Americas net revenue",
    )
    wb = load_workbook(trainer, data_only=False)
    wb[GEOGRAPHIC_SHEET].cell(source_row, 3).value = 1
    wb.save(trainer)
    wb.close()
    with pytest.raises(
        ValueError,
        match="Trusted workbook cell was modified: Geographic Segment Analysis",
    ):
        check_workbook(trainer)

    trainer2, _answer2 = build_training_workbook(fin, tmp_path / "GEO_RELOC.xlsx")
    smap2 = load_semantic_map(_answer2)
    share = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "geographic_revenue_share" and "americas" in c.id
    )
    wb = load_workbook(trainer2, data_only=False)
    row, col = parse_cell_ref(share.cell)
    wb[share.tab].cell(row=row, column=col).value = share.formula
    wb.save(trainer2)
    wb.close()
    relocated = check_workbook(trainer2)
    assert relocated.incorrect == 0
    assert geo_comps


MOVED_GEO_REVENUE_SOURCE_PLACEMENT = {
    ("americas", 0): (28, 8),
    ("china_mainland", 0): (29, 8),
    ("rest_of_world", 0): (30, 8),
    ("consolidated", 0): (31, 8),
    ("americas", 1): (28, 12, "Income Statement"),
    ("china_mainland", 1): (29, 12, "Income Statement"),
    ("rest_of_world", 1): (32, 12, "Income Statement"),
    ("consolidated", 1): (33, 12, "Income Statement"),
}


def test_contribution_formulas_follow_relocated_sources(tmp_path, monkeypatch):
    def _moved(self, identity, period_index, default_row, default_col):
        assert default_col == 2 + period_index
        return MOVED_GEO_REVENUE_SOURCE_PLACEMENT[(identity, period_index)]

    monkeypatch.setattr(
        ReferenceModelBuilder, "_geographic_revenue_source_placement", _moved
    )
    fin = _geo_tiny(
        _snapshot(P1, family=FAMILY_ITEMIZED, values=_itemized_values()),
        _snapshot(P2, family=FAMILY_CORPORATE, values=_corp_values()),
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "GEO_MOVED.xlsx")
    smap = load_semantic_map(answer)
    contrib = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_revenue_growth_contribution"
        and geographic_spec_identity(c) == "americas"
    )
    cons_growth = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_consolidated_revenue_growth"
    )
    residual = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_revenue_growth_contribution_residual"
    )
    prior = MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 0)]
    curr = MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 1)]
    prior_cons = MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("consolidated", 0)]
    curr_cons = MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("consolidated", 1)]

    def _ref(placed, stamp: str) -> SemanticCellRef:
        row, col = placed[0], placed[1]
        tab = placed[2] if len(placed) == 3 else GEOGRAPHIC_SHEET
        return SemanticCellRef(
            stamp, stamp, stamp, f"{get_column_letter(col)}{row}", tab
        )

    compact = contrib.formula.replace(" ", "")
    assert compact == resolve_geographic_revenue_growth_contribution_formula(
        _ref(curr, "curr"),
        _ref(prior, "prior"),
        _ref(prior_cons, "cons"),
        from_tab=GEOGRAPHIC_SHEET,
    ).replace(" ", "")
    aws = load_workbook(answer, data_only=False)[GEOGRAPHIC_SHEET]
    default_amer = _row_by_label(aws, "Americas net revenue")
    adjacent = (
        f"{get_column_letter(parse_cell_ref(contrib.cell)[1] - 1)}{default_amer}"
    )
    assert adjacent not in compact
    assert cons_growth.formula.replace(" ", "") == (
        resolve_geographic_consolidated_revenue_growth_formula(
            _ref(curr_cons, "curr_cons"),
            _ref(prior_cons, "prior_cons"),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    contrib_refs = []
    for identity in GEOGRAPHIC_SEGMENT_IDENTITIES:
        item = next(
            c
            for c in smap.all_ordered()
            if c.family_id == "geographic_revenue_growth_contribution"
            and geographic_spec_identity(c) == identity
        )
        contrib_refs.append(
            SemanticCellRef(
                item.id, item.semantic_key, item.period_end, item.cell, item.tab
            )
        )
    assert residual.formula.replace(" ", "") == (
        resolve_geographic_revenue_growth_contribution_residual_formula(
            SemanticCellRef(
                cons_growth.id,
                cons_growth.semantic_key,
                cons_growth.period_end,
                cons_growth.cell,
                cons_growth.tab,
            ),
            tuple(contrib_refs),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    awb = load_workbook(answer, data_only=False)
    row, col = MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 0)][:2]
    assert awb[GEOGRAPHIC_SHEET].cell(row, col).value not in (None, "")
    curr_row, curr_col, curr_tab = MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 1)]
    assert awb[curr_tab].cell(curr_row, curr_col).value not in (None, "")
    awb.close()
    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.incorrect == 0
    assert filled.correct == filled.total


MOVED_GEO_IFOP_SOURCE_PLACEMENT = {
    ("americas", 0): (40, 8),
    ("china_mainland", 0): (41, 8),
    ("rest_of_world", 0): (42, 8),
    ("consolidated", 0): (43, 8),
    ("americas", 1): (40, 12, "Income Statement"),
    ("china_mainland", 1): (41, 12, "Income Statement"),
    ("rest_of_world", 1): (42, 12, "Income Statement"),
    ("consolidated", 1): (43, 12, "Income Statement"),
}
MOVED_GEO_RECONCILING_SOURCE_PLACEMENT = {
    ("ifop_reconciling.amortization_of_intangible_assets", 0): (44, 8),
    ("ifop_reconciling.general_corporate_expenses", 0): (45, 8),
    (IFOP_CORPORATE, 1): (44, 12, "Income Statement"),
}


def test_margin_bridge_formulas_follow_relocated_sources(tmp_path, monkeypatch):
    def _moved_revenue(self, identity, period_index, default_row, default_col):
        assert default_col == 2 + period_index
        return MOVED_GEO_REVENUE_SOURCE_PLACEMENT[(identity, period_index)]

    def _moved_ifop(self, identity, period_index, default_row, default_col):
        assert default_col == 2 + period_index
        return MOVED_GEO_IFOP_SOURCE_PLACEMENT[(identity, period_index)]

    def _moved_reconciling(self, identity, period_index, default_row, default_col):
        assert default_col == 2 + period_index
        return MOVED_GEO_RECONCILING_SOURCE_PLACEMENT[(identity, period_index)]

    monkeypatch.setattr(
        ReferenceModelBuilder, "_geographic_revenue_source_placement", _moved_revenue
    )
    monkeypatch.setattr(
        ReferenceModelBuilder, "_geographic_ifop_source_placement", _moved_ifop
    )
    monkeypatch.setattr(
        ReferenceModelBuilder,
        "_geographic_reconciling_source_placement",
        _moved_reconciling,
    )
    fin = _geo_tiny(
        _snapshot(P1, family=FAMILY_ITEMIZED, values=_itemized_values()),
        _snapshot(P2, family=FAMILY_CORPORATE, values=_corp_values()),
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "GEO_MARGIN_MOVED.xlsx")
    smap = load_semantic_map(answer)

    def _ref(placed, stamp: str) -> SemanticCellRef:
        row, col = placed[0], placed[1]
        tab = placed[2] if len(placed) == 3 else GEOGRAPHIC_SHEET
        return SemanticCellRef(
            stamp, stamp, stamp, f"{get_column_letter(col)}{row}", tab
        )

    cs = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_margin_contribution"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    assert cs.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_contribution_formula(
            _ref(MOVED_GEO_IFOP_SOURCE_PLACEMENT[("americas", 1)], "ifop"),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("consolidated", 1)], "cons"),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    aws = load_workbook(answer, data_only=False)[GEOGRAPHIC_SHEET]
    default_ifop = _row_by_label(aws, "Americas income from operations")
    adjacent = (
        f"{get_column_letter(parse_cell_ref(cs.cell)[1])}{default_ifop}"
    )
    assert adjacent not in cs.formula.replace(" ", "")
    m = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_consolidated_operating_margin"
        and c.period_index == 1
    )
    assert m.formula.replace(" ", "") == (
        resolve_geographic_consolidated_operating_margin_formula(
            _ref(MOVED_GEO_IFOP_SOURCE_PLACEMENT[("consolidated", 1)], "cons_ifop"),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("consolidated", 1)], "cons"),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    signed = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_signed_reconciling_contribution"
        and c.period_index == 1
    )
    b = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_reconciling_operating_margin_contribution"
        and c.period_index == 1
    )
    assert b.formula.replace(" ", "") == (
        resolve_geographic_reconciling_operating_margin_contribution_formula(
            (
                SemanticCellRef(
                    signed.id, signed.semantic_key, signed.period_end, signed.cell, signed.tab
                ),
            ),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("consolidated", 1)], "cons"),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    compact_signed = signed.formula.replace(" ", "")
    moved_corp = MOVED_GEO_RECONCILING_SOURCE_PLACEMENT[(IFOP_CORPORATE, 1)]
    expected_signed = semantic_formula_cell(
        _ref(moved_corp, "corp"), from_tab=GEOGRAPHIC_SHEET
    )
    assert compact_signed == f"={expected_signed}".replace(" ", "")
    e = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_margin_contribution_residual"
        and c.period_index == 1
    )
    cs_refs = [
        next(
            c
            for c in smap.all_ordered()
            if c.family_id == "geographic_operating_margin_contribution"
            and geographic_spec_identity(c) == identity
            and c.period_index == 1
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    assert e.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_contribution_residual_formula(
            SemanticCellRef(m.id, m.semantic_key, m.period_end, m.cell, m.tab),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in cs_refs
            ),
            SemanticCellRef(b.id, b.semantic_key, b.period_end, b.cell, b.tab),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    dcs = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_margin_contribution_change"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    prior_cs = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_margin_contribution"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 0
    )
    assert dcs.formula.replace(" ", "") == (
        resolve_geographic_adjacent_change_formula(
            SemanticCellRef(cs.id, cs.semantic_key, cs.period_end, cs.cell, cs.tab),
            SemanticCellRef(
                prior_cs.id,
                prior_cs.semantic_key,
                prior_cs.period_end,
                prior_cs.cell,
                prior_cs.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    de = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_margin_contribution_change_residual"
        and c.period_index == 1
    )
    dm = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_consolidated_operating_margin_change"
        and c.period_index == 1
    )
    db = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_reconciling_operating_margin_contribution_change"
        and c.period_index == 1
    )
    dcs_refs = [
        next(
            c
            for c in smap.all_ordered()
            if c.family_id == "geographic_operating_margin_contribution_change"
            and geographic_spec_identity(c) == identity
            and c.period_index == 1
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    assert de.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_contribution_change_residual_formula(
            SemanticCellRef(dm.id, dm.semantic_key, dm.period_end, dm.cell, dm.tab),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in dcs_refs
            ),
            SemanticCellRef(db.id, db.semantic_key, db.period_end, db.cell, db.tab),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    mix = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_margin_mix_effect"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    prior_share = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_revenue_share"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 0
    )
    current_share = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_revenue_share"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    prior_margin = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_reported_operating_margin"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 0
    )
    current_margin = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_reported_operating_margin"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    assert mix.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_mix_effect_formula(
            SemanticCellRef(
                current_share.id,
                current_share.semantic_key,
                current_share.period_end,
                current_share.cell,
                current_share.tab,
            ),
            SemanticCellRef(
                prior_share.id,
                prior_share.semantic_key,
                prior_share.period_end,
                prior_share.cell,
                prior_share.tab,
            ),
            SemanticCellRef(
                current_margin.id,
                current_margin.semantic_key,
                current_margin.period_end,
                current_margin.cell,
                current_margin.tab,
            ),
            SemanticCellRef(
                prior_margin.id,
                prior_margin.semantic_key,
                prior_margin.period_end,
                prior_margin.cell,
                prior_margin.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    within = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_margin_within_segment_effect"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    assert within.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_within_segment_effect_formula(
            SemanticCellRef(
                current_share.id,
                current_share.semantic_key,
                current_share.period_end,
                current_share.cell,
                current_share.tab,
            ),
            SemanticCellRef(
                prior_share.id,
                prior_share.semantic_key,
                prior_share.period_end,
                prior_share.cell,
                prior_share.tab,
            ),
            SemanticCellRef(
                current_margin.id,
                current_margin.semantic_key,
                current_margin.period_end,
                current_margin.cell,
                current_margin.tab,
            ),
            SemanticCellRef(
                prior_margin.id,
                prior_margin.semantic_key,
                prior_margin.period_end,
                prior_margin.cell,
                prior_margin.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    mix_res = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_margin_mix_within_residual"
        and c.period_index == 1
    )
    mix_refs = [
        next(
            c
            for c in smap.all_ordered()
            if c.family_id == "geographic_operating_margin_mix_effect"
            and geographic_spec_identity(c) == identity
            and c.period_index == 1
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    within_refs = [
        next(
            c
            for c in smap.all_ordered()
            if c.family_id == "geographic_operating_margin_within_segment_effect"
            and geographic_spec_identity(c) == identity
            and c.period_index == 1
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    assert mix_res.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_mix_within_residual_formula(
            SemanticCellRef(dm.id, dm.semantic_key, dm.period_end, dm.cell, dm.tab),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in mix_refs
            ),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in within_refs
            ),
            SemanticCellRef(db.id, db.semantic_key, db.period_end, db.cell, db.tab),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    dp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_profit_amount_change"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    assert dp.formula.replace(" ", "") == (
        resolve_geographic_adjacent_change_formula(
            _ref(MOVED_GEO_IFOP_SOURCE_PLACEMENT[("americas", 1)], "ifop_curr"),
            _ref(MOVED_GEO_IFOP_SOURCE_PLACEMENT[("americas", 0)], "ifop_prior"),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    default_ifop_row = default_ifop
    adjacent_ifop = f"{get_column_letter(parse_cell_ref(dp.cell)[1])}{default_ifop_row}"
    assert adjacent_ifop not in dp.formula.replace(" ", "")
    dpc = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_consolidated_operating_profit_amount_change"
        and c.period_index == 1
    )
    assert dpc.formula.replace(" ", "") == (
        resolve_geographic_adjacent_change_formula(
            _ref(MOVED_GEO_IFOP_SOURCE_PLACEMENT[("consolidated", 1)], "cons_curr"),
            _ref(MOVED_GEO_IFOP_SOURCE_PLACEMENT[("consolidated", 0)], "cons_prior"),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    prior_signed = [
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_signed_reconciling_contribution"
        and c.period_index == 0
    ]
    current_signed = [
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_signed_reconciling_contribution"
        and c.period_index == 1
    ]
    dbr = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_reconciling_operating_profit_amount_change"
        and c.period_index == 1
    )
    assert dbr.formula.replace(" ", "") == (
        resolve_geographic_reconciling_operating_profit_amount_change_formula(
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in current_signed
            ),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in prior_signed
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert "100*" not in dbr.formula.replace(" ", "")
    dpr = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_profit_amount_change_residual"
        and c.period_index == 1
    )
    dp_refs = [
        next(
            c
            for c in smap.all_ordered()
            if c.family_id == "geographic_operating_profit_amount_change"
            and geographic_spec_identity(c) == identity
            and c.period_index == 1
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    assert dpr.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_amount_change_residual_formula(
            SemanticCellRef(dpc.id, dpc.semantic_key, dpc.period_end, dpc.cell, dpc.tab),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in dp_refs
            ),
            SemanticCellRef(dbr.id, dbr.semantic_key, dbr.period_end, dbr.cell, dbr.tab),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    rev_eff = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_profit_revenue_effect"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    current_margin = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_reported_operating_margin"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    prior_margin = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_reported_operating_margin"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 0
    )
    assert rev_eff.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_revenue_effect_formula(
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 1)], "rev_curr"),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 0)], "rev_prior"),
            SemanticCellRef(
                current_margin.id,
                current_margin.semantic_key,
                current_margin.period_end,
                current_margin.cell,
                current_margin.tab,
            ),
            SemanticCellRef(
                prior_margin.id,
                prior_margin.semantic_key,
                prior_margin.period_end,
                prior_margin.cell,
                prior_margin.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    default_rev_row = _row_by_label(aws, "Americas net revenue")
    adjacent_rev = (
        f"{get_column_letter(parse_cell_ref(rev_eff.cell)[1])}{default_rev_row}"
    )
    assert adjacent_rev not in rev_eff.formula.replace(" ", "")
    mgn_eff = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_profit_margin_effect"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    assert mgn_eff.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_margin_effect_formula(
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 1)], "rev_curr"),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 0)], "rev_prior"),
            SemanticCellRef(
                current_margin.id,
                current_margin.semantic_key,
                current_margin.period_end,
                current_margin.cell,
                current_margin.tab,
            ),
            SemanticCellRef(
                prior_margin.id,
                prior_margin.semantic_key,
                prior_margin.period_end,
                prior_margin.cell,
                prior_margin.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert "100*" not in rev_eff.formula.replace(" ", "")
    assert "100*" not in mgn_eff.formula.replace(" ", "")
    inc_mgn = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_operating_profit_incremental_margin"
        and geographic_spec_identity(c) == "americas"
        and c.period_index == 1
    )
    assert inc_mgn.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_incremental_margin_formula(
            SemanticCellRef(dp.id, dp.semantic_key, dp.period_end, dp.cell, dp.tab),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 1)], "rev_curr"),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("americas", 0)], "rev_prior"),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert adjacent_rev not in inc_mgn.formula.replace(" ", "")
    assert "NA()" in inc_mgn.formula.replace(" ", "")
    cons_inc_mgn = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "geographic_consolidated_operating_profit_incremental_margin"
        and c.period_index == 1
    )
    assert cons_inc_mgn.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_incremental_margin_formula(
            SemanticCellRef(dpc.id, dpc.semantic_key, dpc.period_end, dpc.cell, dpc.tab),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("consolidated", 1)], "cons_curr"),
            _ref(MOVED_GEO_REVENUE_SOURCE_PLACEMENT[("consolidated", 0)], "cons_prior"),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    awb = load_workbook(answer, data_only=False)
    row, col = MOVED_GEO_IFOP_SOURCE_PLACEMENT[("americas", 0)][:2]
    assert awb[GEOGRAPHIC_SHEET].cell(row, col).value not in (None, "")
    curr_row, curr_col, curr_tab = MOVED_GEO_IFOP_SOURCE_PLACEMENT[("americas", 1)]
    assert awb[curr_tab].cell(curr_row, curr_col).value not in (None, "")
    awb.close()
    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.incorrect == 0
    assert filled.correct == filled.total


def test_committed_lululemon_and_fast_retailing_identities_unchanged():
    lulu = standardized_from_payload(
        json.loads(LULU_JSON.read_text(encoding="utf-8"))
    )
    assert lulu.historical_segment is None
    lulu_builder = ReferenceModelBuilder(lulu)
    assert lulu_builder.geographic_specs == ()
    assert len(lulu_builder.expected_specs) == LEASE_DT_LULULEMON_SPECS

    fr = standardized_from_payload(json.loads(FR_JSON.read_text(encoding="utf-8")))
    assert fr.historical_segment is None
    fr_builder = ReferenceModelBuilder(fr)
    assert fr_builder.geographic_specs == ()
    assert len(fr_builder.expected_specs) == FAST_RETAILING_SPECS

    demo = HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )
    assert not geographic_segment_applicable(demo)
    assert ReferenceModelBuilder(demo).geographic_specs == ()


def test_lululemon_five_period_temporary_pair_matches_selected_facts(tmp_path):
    reconciled, _fin, restored = _reload_admitted_lululemon()
    prefix = "segment.geo.q4_2023."
    selected_on_axis = {
        (item.period, item.fact_type[len(prefix) :]): item.value
        for item in reconciled.selected_geographic_facts
        if item.period in set(reconciled.periods)
    }
    model_values = {
        (snap.period, identity): value
        for snap in restored.historical_segment.periods
        for identity, value in snap.values.items()
    }
    assert model_values == selected_on_axis
    assert len(model_values) == 56

    snapshots = {snap.period: snap for snap in restored.historical_segment.periods}
    series = compute_geographic_segment_series(restored)
    expected = _independent_from_snapshots(list(series.periods), snapshots)
    _assert_series_matches(series, expected)
    assert 3607682 - 1397067 == 2210615
    assert series.calculated_segment_operating_profit_total[FY2026] == 3607682
    assert series.signed_reconciling_contributions[FY2026] == (
        (IFOP_CORPORATE, -1397067.0),
    )
    assert series.reconstructed_consolidated_operating_profit[FY2026] == 2210615
    for period in series.periods:
        assert series.consolidated_revenue_difference[period] == 0.0
        assert series.consolidated_operating_profit_difference[period] == 0.0
        share = series.revenue_share[period]
        if all(isinstance(value, float) for value in share.values()):
            assert abs(sum(share.values()) - 1.0) <= GEOGRAPHIC_RATIO_TOLERANCE

    builder = ReferenceModelBuilder(restored)
    geo_n = len(builder.geographic_specs)
    assert geo_n == GEOGRAPHIC_LULULEMON_SPECS
    preserved = [
        s
        for s in builder.expected_specs
        if s.category != "geographic_segment"
    ]
    assert len(preserved) == LEASE_DT_LULULEMON_SPECS
    assert {s.id for s in preserved} == {
        s.id
        for s in ReferenceModelBuilder(
            standardized_from_payload(
                json.loads(LULU_JSON.read_text(encoding="utf-8"))
            )
        ).expected_specs
    }

    trainer, answer = build_training_workbook(
        restored, tmp_path / "LULU_GEO_TMP.xlsx"
    )
    reopened = tmp_path / "reopened"
    trainer, answer = _copy_pair(trainer, answer, reopened)
    smap = load_semantic_map(answer)
    all_comps = list(smap.all_ordered())
    geo_comps = [c for c in all_comps if c.category == "geographic_segment"]
    assert len(geo_comps) == GEOGRAPHIC_LULULEMON_SPECS
    assert len(all_comps) == LEASE_DT_LULULEMON_SPECS + GEOGRAPHIC_LULULEMON_SPECS
    practice = {(c.tab, c.cell) for c in all_comps}

    awb = load_workbook(answer, data_only=False)
    twb = load_workbook(trainer, data_only=False)
    aws = awb[GEOGRAPHIC_SHEET]
    tws = twb[GEOGRAPHIC_SHEET]
    assert aws["A2"].value == GEOGRAPHIC_A2
    assert tws["A2"].value == GEOGRAPHIC_A2
    assert aws["A4"].value == GEOGRAPHIC_A4
    assert tws["A4"].value == GEOGRAPHIC_A4
    assert "income from operations / net revenue" not in GEOGRAPHIC_A2.lower()
    assert "add" not in GEOGRAPHIC_A4.lower()
    assert "subtract" not in GEOGRAPHIC_A4.lower()
    sheets, cells_scanned, text_n = _assert_trainer_complete_text_undisclosed(twb)
    assert sheets
    assert cells_scanned > 0
    assert text_n > 0

    margin_comps = [
        c
        for c in geo_comps
        if c.family_id == "geographic_reported_operating_margin"
    ]
    expected_margin_ids = {
        geographic_component_id(
            "geographic_reported_operating_margin", period, identity
        )
        for period in series.periods
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    }
    assert len(margin_comps) == 15
    assert {c.id for c in margin_comps} == expected_margin_ids
    period_cols = {period: 2 + index for index, period in enumerate(builder.periods)}
    for identity in GEOGRAPHIC_SEGMENT_IDENTITIES:
        label = geographic_identity_label(identity)
        ifop_row = _row_by_label(aws, f"{label} income from operations")
        rev_row = _row_by_label(aws, f"{label} net revenue")
        for period in series.periods:
            col_idx = period_cols[period]
            col = get_column_letter(col_idx)
            comp = next(
                c
                for c in margin_comps
                if c.period_end == period.isoformat()
                and geographic_spec_identity(c) == identity
            )
            row, col_parsed = parse_cell_ref(comp.cell)
            assert col_parsed == col_idx
            trainer_cell = twb[comp.tab].cell(row=row, column=col_idx)
            answer_cell = awb[comp.tab].cell(row=row, column=col_idx)
            assert trainer_cell.value is None
            assert trainer_cell.comment is None
            assert _fill_rgb(trainer_cell) == "FFFF00"
            expected_formula = (
                f"=IF({col}{rev_row}=0,NA(),{col}{ifop_row}/{col}{rev_row})"
            )
            assert answer_cell.value.replace(" ", "") == expected_formula.replace(
                " ", ""
            )
            assert answer_cell.value == comp.formula
            assert _fill_rgb(answer_cell) in WHITE_RGBS
            _assert_margin_note_explains_reported_basis(
                answer_cell.comment.text if answer_cell.comment else ""
            )

    americas_rev = _row_by_label(aws, "Americas net revenue")
    china_rev = _row_by_label(aws, "China Mainland net revenue")
    row_rev = _row_by_label(aws, "Rest of World net revenue")
    cons_rev = _row_by_label(aws, "Reported consolidated revenue")
    fy2026_col = period_cols[FY2026]
    assert aws.cell(americas_rev, fy2026_col).value == snapshots[FY2026].values[
        "net_revenue.americas"
    ]
    assert tws.cell(americas_rev, fy2026_col).value == aws.cell(
        americas_rev, fy2026_col
    ).value
    assert aws.cell(china_rev, fy2026_col).value == snapshots[FY2026].values[
        "net_revenue.china_mainland"
    ]
    assert aws.cell(row_rev, fy2026_col).value == snapshots[FY2026].values[
        "net_revenue.rest_of_world"
    ]
    mix = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_revenue_share"
        and c.period_end == FY2026.isoformat()
        and "americas" in c.id
    )
    assert "NA()" in mix.formula
    assert mix.formula.replace(" ", "") == (
        f"=IF({get_column_letter(fy2026_col)}{cons_rev}=0,NA(),"
        f"{get_column_letter(fy2026_col)}{americas_rev}/"
        f"{get_column_letter(fy2026_col)}{cons_rev})"
    ).replace(" ", "")
    prior_fy2026 = builder.periods[builder.periods.index(FY2026) - 1]
    prior_col = period_cols[prior_fy2026]
    amer_contrib = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_revenue_growth_contribution"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    assert amer_contrib.formula.replace(" ", "") == (
        resolve_geographic_revenue_growth_contribution_formula(
            SemanticCellRef(
                "curr", "k", FY2026.isoformat(),
                f"{get_column_letter(fy2026_col)}{americas_rev}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "prior", "k", prior_fy2026.isoformat(),
                f"{get_column_letter(prior_col)}{americas_rev}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "cons", "k", prior_fy2026.isoformat(),
                f"{get_column_letter(prior_col)}{cons_rev}",
                GEOGRAPHIC_SHEET,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    prior_cons = float(snapshots[prior_fy2026].values["net_revenue.consolidated"])
    curr_amer = float(snapshots[FY2026].values["net_revenue.americas"])
    prior_amer = float(snapshots[prior_fy2026].values["net_revenue.americas"])
    independent_amer = 100.0 * (curr_amer - prior_amer) / prior_cons
    assert amer_contrib.expected_value == pytest.approx(
        independent_amer, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    contrib_n = [
        c
        for c in geo_comps
        if c.family_id
        in {
            "geographic_revenue_growth_contribution",
            "geographic_consolidated_revenue_growth",
            "geographic_revenue_growth_contribution_residual",
        }
    ]
    assert len(contrib_n) == GEOGRAPHIC_CONTRIBUTION_SPECS
    margin_bridge_n = [
        c
        for c in geo_comps
        if c.family_id
        in {
            "geographic_operating_margin_contribution",
            "geographic_reconciling_operating_margin_contribution",
            "geographic_consolidated_operating_margin",
            "geographic_operating_margin_contribution_residual",
            "geographic_operating_margin_contribution_change",
            "geographic_reconciling_operating_margin_contribution_change",
            "geographic_consolidated_operating_margin_change",
            "geographic_operating_margin_contribution_change_residual",
        }
    ]
    assert len(margin_bridge_n) == GEOGRAPHIC_MARGIN_BRIDGE_SPECS
    mix_within_n = [
        c
        for c in geo_comps
        if c.family_id
        in {
            "geographic_operating_margin_mix_effect",
            "geographic_operating_margin_within_segment_effect",
            "geographic_operating_margin_mix_within_residual",
        }
    ]
    assert len(mix_within_n) == GEOGRAPHIC_MIX_WITHIN_SPECS
    amount_change_n = [
        c
        for c in geo_comps
        if c.family_id
        in {
            "geographic_operating_profit_amount_change",
            "geographic_reconciling_operating_profit_amount_change",
            "geographic_consolidated_operating_profit_amount_change",
            "geographic_operating_profit_amount_change_residual",
        }
    ]
    assert len(amount_change_n) == GEOGRAPHIC_AMOUNT_CHANGE_SPECS
    profit_effect_n = [
        c
        for c in geo_comps
        if c.family_id
        in {
            "geographic_operating_profit_revenue_effect",
            "geographic_operating_profit_margin_effect",
        }
    ]
    assert len(profit_effect_n) == GEOGRAPHIC_PROFIT_EFFECT_SPECS
    incremental_n = [
        c
        for c in geo_comps
        if c.family_id
        in {
            "geographic_operating_profit_incremental_margin",
            "geographic_consolidated_operating_profit_incremental_margin",
        }
    ]
    assert len(incremental_n) == GEOGRAPHIC_INCREMENTAL_MARGIN_SPECS
    fy2026_cs = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_margin_contribution"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    fy2026_ifop_amer = float(
        snapshots[FY2026].values["income_from_operations.americas"]
    )
    fy2026_cons_rev = float(snapshots[FY2026].values["net_revenue.consolidated"])
    independent_cs = 100.0 * fy2026_ifop_amer / fy2026_cons_rev
    assert fy2026_cs.expected_value == pytest.approx(
        independent_cs, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    americas_ifop = _row_by_label(aws, "Americas income from operations")
    assert fy2026_cs.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_contribution_formula(
            SemanticCellRef(
                "ifop", "k", FY2026.isoformat(),
                f"{get_column_letter(fy2026_col)}{americas_ifop}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "cons", "k", FY2026.isoformat(),
                f"{get_column_letter(fy2026_col)}{cons_rev}",
                GEOGRAPHIC_SHEET,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    fy2026_m = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_consolidated_operating_margin"
        and c.period_end == FY2026.isoformat()
    )
    fy2026_b = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_reconciling_operating_margin_contribution"
        and c.period_end == FY2026.isoformat()
    )
    fy2026_e = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_margin_contribution_residual"
        and c.period_end == FY2026.isoformat()
    )
    cs_comps = [
        next(
            c
            for c in geo_comps
            if c.family_id == "geographic_operating_margin_contribution"
            and c.period_end == FY2026.isoformat()
            and geographic_spec_identity(c) == identity
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    independent_e = (
        float(fy2026_m.expected_value)
        - sum(float(item.expected_value) for item in cs_comps)
        - float(fy2026_b.expected_value)
    )
    assert fy2026_e.expected_value == pytest.approx(
        independent_e, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    assert fy2026_e.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_contribution_residual_formula(
            SemanticCellRef(
                fy2026_m.id,
                fy2026_m.semantic_key,
                fy2026_m.period_end,
                fy2026_m.cell,
                fy2026_m.tab,
            ),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in cs_comps
            ),
            SemanticCellRef(
                fy2026_b.id,
                fy2026_b.semantic_key,
                fy2026_b.period_end,
                fy2026_b.cell,
                fy2026_b.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    opening = series.periods[0]
    assert not any(
        c.family_id == "geographic_operating_margin_contribution_change"
        and c.period_end == opening.isoformat()
        for c in geo_comps
    )
    assert not any(
        c.family_id
        in {
            "geographic_operating_margin_mix_effect",
            "geographic_operating_margin_within_segment_effect",
            "geographic_operating_margin_mix_within_residual",
            "geographic_operating_profit_amount_change",
            "geographic_reconciling_operating_profit_amount_change",
            "geographic_consolidated_operating_profit_amount_change",
            "geographic_operating_profit_amount_change_residual",
            "geographic_operating_profit_revenue_effect",
            "geographic_operating_profit_margin_effect",
            "geographic_operating_profit_incremental_margin",
            "geographic_consolidated_operating_profit_incremental_margin",
        }
        and c.period_end == opening.isoformat()
        for c in geo_comps
    )
    fy2026_de = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_margin_contribution_change_residual"
        and c.period_end == FY2026.isoformat()
    )
    fy2026_dm = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_consolidated_operating_margin_change"
        and c.period_end == FY2026.isoformat()
    )
    fy2026_db = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_reconciling_operating_margin_contribution_change"
        and c.period_end == FY2026.isoformat()
    )
    dcs_comps = [
        next(
            c
            for c in geo_comps
            if c.family_id == "geographic_operating_margin_contribution_change"
            and c.period_end == FY2026.isoformat()
            and geographic_spec_identity(c) == identity
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    independent_de = (
        float(fy2026_dm.expected_value)
        - sum(float(item.expected_value) for item in dcs_comps)
        - float(fy2026_db.expected_value)
    )
    assert fy2026_de.expected_value == pytest.approx(
        independent_de, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    fy2026_mix = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_margin_mix_effect"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    fy2026_within = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_margin_within_segment_effect"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    fy2026_mix_res = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_margin_mix_within_residual"
        and c.period_end == FY2026.isoformat()
    )
    prior_fy2026_snap = snapshots[prior_fy2026]
    w_c = float(snapshots[FY2026].values["net_revenue.americas"]) / fy2026_cons_rev
    w_p = float(prior_fy2026_snap.values["net_revenue.americas"]) / float(
        prior_fy2026_snap.values["net_revenue.consolidated"]
    )
    m_c = float(snapshots[FY2026].values["income_from_operations.americas"]) / float(
        snapshots[FY2026].values["net_revenue.americas"]
    )
    m_p = float(prior_fy2026_snap.values["income_from_operations.americas"]) / float(
        prior_fy2026_snap.values["net_revenue.americas"]
    )
    independent_mix = 100.0 * (w_c - w_p) * (m_c + m_p) / 2.0
    independent_within = 100.0 * (m_c - m_p) * (w_c + w_p) / 2.0
    assert fy2026_mix.expected_value == pytest.approx(
        independent_mix, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    assert fy2026_within.expected_value == pytest.approx(
        independent_within, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    assert fy2026_mix.expected_value + fy2026_within.expected_value == pytest.approx(
        next(
            c
            for c in dcs_comps
            if geographic_spec_identity(c) == "americas"
        ).expected_value,
        abs=GEOGRAPHIC_RATIO_TOLERANCE,
    )
    mix_comps = [
        next(
            c
            for c in geo_comps
            if c.family_id == "geographic_operating_margin_mix_effect"
            and c.period_end == FY2026.isoformat()
            and geographic_spec_identity(c) == identity
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    within_comps = [
        next(
            c
            for c in geo_comps
            if c.family_id == "geographic_operating_margin_within_segment_effect"
            and c.period_end == FY2026.isoformat()
            and geographic_spec_identity(c) == identity
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    independent_mix_res = (
        float(fy2026_dm.expected_value)
        - sum(float(item.expected_value) for item in mix_comps)
        - sum(float(item.expected_value) for item in within_comps)
        - float(fy2026_db.expected_value)
    )
    assert fy2026_mix_res.expected_value == pytest.approx(
        independent_mix_res, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    assert fy2026_mix_res.expected_value == pytest.approx(
        fy2026_de.expected_value, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    fy2026_share = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_revenue_share"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    prior_share_comp = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_revenue_share"
        and c.period_end == prior_fy2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    fy2026_seg_m = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_reported_operating_margin"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    prior_seg_m = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_reported_operating_margin"
        and c.period_end == prior_fy2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    assert fy2026_mix.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_mix_effect_formula(
            SemanticCellRef(
                fy2026_share.id,
                fy2026_share.semantic_key,
                fy2026_share.period_end,
                fy2026_share.cell,
                fy2026_share.tab,
            ),
            SemanticCellRef(
                prior_share_comp.id,
                prior_share_comp.semantic_key,
                prior_share_comp.period_end,
                prior_share_comp.cell,
                prior_share_comp.tab,
            ),
            SemanticCellRef(
                fy2026_seg_m.id,
                fy2026_seg_m.semantic_key,
                fy2026_seg_m.period_end,
                fy2026_seg_m.cell,
                fy2026_seg_m.tab,
            ),
            SemanticCellRef(
                prior_seg_m.id,
                prior_seg_m.semantic_key,
                prior_seg_m.period_end,
                prior_seg_m.cell,
                prior_seg_m.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert fy2026_mix_res.formula.replace(" ", "") == (
        resolve_geographic_operating_margin_mix_within_residual_formula(
            SemanticCellRef(
                fy2026_dm.id,
                fy2026_dm.semantic_key,
                fy2026_dm.period_end,
                fy2026_dm.cell,
                fy2026_dm.tab,
            ),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in mix_comps
            ),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in within_comps
            ),
            SemanticCellRef(
                fy2026_db.id,
                fy2026_db.semantic_key,
                fy2026_db.period_end,
                fy2026_db.cell,
                fy2026_db.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    fy2026_dp = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_profit_amount_change"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    fy2026_dpc = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_consolidated_operating_profit_amount_change"
        and c.period_end == FY2026.isoformat()
    )
    fy2026_dbr = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_reconciling_operating_profit_amount_change"
        and c.period_end == FY2026.isoformat()
    )
    fy2026_dpr = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_profit_amount_change_residual"
        and c.period_end == FY2026.isoformat()
    )
    independent_dp = (
        float(snapshots[FY2026].values["income_from_operations.americas"])
        - float(prior_fy2026_snap.values["income_from_operations.americas"])
    )
    independent_dpc = (
        float(snapshots[FY2026].values["income_from_operations.consolidated"])
        - float(prior_fy2026_snap.values["income_from_operations.consolidated"])
    )
    curr_recon = sum(
        amount
        for _, amount in series.signed_reconciling_contributions[FY2026]
    )
    prior_recon = sum(
        amount
        for _, amount in series.signed_reconciling_contributions[prior_fy2026]
    )
    independent_dbr = curr_recon - prior_recon
    assert fy2026_dp.expected_value == pytest.approx(independent_dp)
    assert fy2026_dpc.expected_value == pytest.approx(independent_dpc)
    assert fy2026_dbr.expected_value == pytest.approx(independent_dbr)
    dp_comps = [
        next(
            c
            for c in geo_comps
            if c.family_id == "geographic_operating_profit_amount_change"
            and c.period_end == FY2026.isoformat()
            and geographic_spec_identity(c) == identity
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    independent_dpr = (
        float(fy2026_dpc.expected_value)
        - sum(float(item.expected_value) for item in dp_comps)
        - float(fy2026_dbr.expected_value)
    )
    assert fy2026_dpr.expected_value == pytest.approx(independent_dpr)
    cons_ifop_row = _row_by_label(aws, "Reported consolidated operating profit")
    fy2026_col_letter = get_column_letter(fy2026_col)
    prior_col_letter = get_column_letter(prior_col)
    assert fy2026_dp.formula.replace(" ", "") == (
        resolve_geographic_adjacent_change_formula(
            SemanticCellRef(
                "ifop", "k", FY2026.isoformat(),
                f"{fy2026_col_letter}{americas_ifop}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "ifop_p", "k", prior_fy2026.isoformat(),
                f"{prior_col_letter}{americas_ifop}",
                GEOGRAPHIC_SHEET,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert fy2026_dpc.formula.replace(" ", "") == (
        resolve_geographic_adjacent_change_formula(
            SemanticCellRef(
                "cons", "k", FY2026.isoformat(),
                f"{fy2026_col_letter}{cons_ifop_row}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "cons_p", "k", prior_fy2026.isoformat(),
                f"{prior_col_letter}{cons_ifop_row}",
                GEOGRAPHIC_SHEET,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    curr_signed_comps = [
        c
        for c in geo_comps
        if c.family_id == "geographic_signed_reconciling_contribution"
        and c.period_end == FY2026.isoformat()
    ]
    prior_signed_comps = [
        c
        for c in geo_comps
        if c.family_id == "geographic_signed_reconciling_contribution"
        and c.period_end == prior_fy2026.isoformat()
    ]
    assert fy2026_dbr.formula.replace(" ", "") == (
        resolve_geographic_reconciling_operating_profit_amount_change_formula(
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in curr_signed_comps
            ),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in prior_signed_comps
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert "100*" not in fy2026_dbr.formula.replace(" ", "")
    assert fy2026_dpr.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_amount_change_residual_formula(
            SemanticCellRef(
                fy2026_dpc.id,
                fy2026_dpc.semantic_key,
                fy2026_dpc.period_end,
                fy2026_dpc.cell,
                fy2026_dpc.tab,
            ),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in dp_comps
            ),
            SemanticCellRef(
                fy2026_dbr.id,
                fy2026_dbr.semantic_key,
                fy2026_dbr.period_end,
                fy2026_dbr.cell,
                fy2026_dbr.tab,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    fy2026_rev_eff = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_profit_revenue_effect"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    fy2026_mgn_eff = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_profit_margin_effect"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    fy2026_v = float(snapshots[FY2026].values["net_revenue.americas"])
    prior_v = float(prior_fy2026_snap.values["net_revenue.americas"])
    fy2026_p = float(snapshots[FY2026].values["income_from_operations.americas"])
    prior_p = float(prior_fy2026_snap.values["income_from_operations.americas"])
    fy2026_m_amer = fy2026_p / fy2026_v
    prior_m_amer = prior_p / prior_v
    independent_rev_eff = (fy2026_v - prior_v) * (fy2026_m_amer + prior_m_amer) / 2.0
    independent_mgn_eff = (fy2026_m_amer - prior_m_amer) * (fy2026_v + prior_v) / 2.0
    assert fy2026_rev_eff.expected_value == pytest.approx(independent_rev_eff)
    assert fy2026_mgn_eff.expected_value == pytest.approx(independent_mgn_eff)
    assert float(fy2026_rev_eff.expected_value) + float(
        fy2026_mgn_eff.expected_value
    ) == pytest.approx(float(fy2026_dp.expected_value))
    americas_rev = _row_by_label(aws, "Americas net revenue")
    americas_margin = _row_by_label(aws, "Americas reported operating margin")
    assert fy2026_rev_eff.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_revenue_effect_formula(
            SemanticCellRef(
                "rev", "k", FY2026.isoformat(),
                f"{fy2026_col_letter}{americas_rev}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "rev_p", "k", prior_fy2026.isoformat(),
                f"{prior_col_letter}{americas_rev}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "m", "k", FY2026.isoformat(),
                f"{fy2026_col_letter}{americas_margin}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "m_p", "k", prior_fy2026.isoformat(),
                f"{prior_col_letter}{americas_margin}",
                GEOGRAPHIC_SHEET,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert fy2026_mgn_eff.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_margin_effect_formula(
            SemanticCellRef(
                "rev", "k", FY2026.isoformat(),
                f"{fy2026_col_letter}{americas_rev}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "rev_p", "k", prior_fy2026.isoformat(),
                f"{prior_col_letter}{americas_rev}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "m", "k", FY2026.isoformat(),
                f"{fy2026_col_letter}{americas_margin}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "m_p", "k", prior_fy2026.isoformat(),
                f"{prior_col_letter}{americas_margin}",
                GEOGRAPHIC_SHEET,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert "100*" not in fy2026_rev_eff.formula.replace(" ", "")
    fy2026_inc = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_operating_profit_incremental_margin"
        and c.period_end == FY2026.isoformat()
        and geographic_spec_identity(c) == "americas"
    )
    fy2026_cons_inc = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_consolidated_operating_profit_incremental_margin"
        and c.period_end == FY2026.isoformat()
    )
    independent_inc = (fy2026_p - prior_p) / (fy2026_v - prior_v)
    assert fy2026_inc.expected_value == pytest.approx(independent_inc)
    assert float(fy2026_inc.expected_value) * (fy2026_v - prior_v) == pytest.approx(
        float(fy2026_dp.expected_value)
    )
    fy2026_cons_v = float(snapshots[FY2026].values["net_revenue.consolidated"])
    prior_cons_v = float(prior_fy2026_snap.values["net_revenue.consolidated"])
    fy2026_cons_p = float(snapshots[FY2026].values["income_from_operations.consolidated"])
    prior_cons_p = float(prior_fy2026_snap.values["income_from_operations.consolidated"])
    independent_cons_inc = (fy2026_cons_p - prior_cons_p) / (fy2026_cons_v - prior_cons_v)
    assert fy2026_cons_inc.expected_value == pytest.approx(independent_cons_inc)
    assert fy2026_inc.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_incremental_margin_formula(
            SemanticCellRef(
                fy2026_dp.id,
                fy2026_dp.semantic_key,
                fy2026_dp.period_end,
                fy2026_dp.cell,
                fy2026_dp.tab,
            ),
            SemanticCellRef(
                "rev", "k", FY2026.isoformat(),
                f"{fy2026_col_letter}{americas_rev}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "rev_p", "k", prior_fy2026.isoformat(),
                f"{prior_col_letter}{americas_rev}",
                GEOGRAPHIC_SHEET,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    cons_rev_row = cons_rev
    assert fy2026_cons_inc.formula.replace(" ", "") == (
        resolve_geographic_operating_profit_incremental_margin_formula(
            SemanticCellRef(
                fy2026_dpc.id,
                fy2026_dpc.semantic_key,
                fy2026_dpc.period_end,
                fy2026_dpc.cell,
                fy2026_dpc.tab,
            ),
            SemanticCellRef(
                "rev", "k", FY2026.isoformat(),
                f"{fy2026_col_letter}{cons_rev_row}",
                GEOGRAPHIC_SHEET,
            ),
            SemanticCellRef(
                "rev_p", "k", prior_fy2026.isoformat(),
                f"{prior_col_letter}{cons_rev_row}",
                GEOGRAPHIC_SHEET,
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    assert "NA()" in fy2026_inc.formula.replace(" ", "")
    max_effect_error = 0.0
    max_incremental_error = 0.0
    incremental_numeric = 0
    incremental_undefined = 0
    for identity in GEOGRAPHIC_SEGMENT_IDENTITIES:
        for current, prior in (
            (date(2023, 1, 29), date(2022, 1, 30)),
            (date(2024, 1, 28), date(2023, 1, 29)),
            (date(2025, 2, 2), date(2024, 1, 28)),
            (FY2026, prior_fy2026),
        ):
            curr_snap = snapshots[current]
            prior_snap = snapshots[prior]
            v_t = float(curr_snap.values[f"net_revenue.{identity}"])
            v_p = float(prior_snap.values[f"net_revenue.{identity}"])
            p_t = float(curr_snap.values[f"income_from_operations.{identity}"])
            p_p = float(prior_snap.values[f"income_from_operations.{identity}"])
            m_t = p_t / v_t
            m_p = p_p / v_p
            expected_rev = (v_t - v_p) * (m_t + m_p) / 2.0
            expected_mgn = (m_t - m_p) * (v_t + v_p) / 2.0
            expected_dp = p_t - p_p
            rev_comp = next(
                c
                for c in geo_comps
                if c.family_id == "geographic_operating_profit_revenue_effect"
                and c.period_end == current.isoformat()
                and geographic_spec_identity(c) == identity
            )
            mgn_comp = next(
                c
                for c in geo_comps
                if c.family_id == "geographic_operating_profit_margin_effect"
                and c.period_end == current.isoformat()
                and geographic_spec_identity(c) == identity
            )
            max_effect_error = max(
                max_effect_error,
                abs(float(rev_comp.expected_value) - expected_rev),
                abs(float(mgn_comp.expected_value) - expected_mgn),
                abs(
                    float(rev_comp.expected_value)
                    + float(mgn_comp.expected_value)
                    - expected_dp
                ),
            )
            inc_comp = next(
                c
                for c in geo_comps
                if c.family_id == "geographic_operating_profit_incremental_margin"
                and c.period_end == current.isoformat()
                and geographic_spec_identity(c) == identity
            )
            dv = v_t - v_p
            if dv == 0.0:
                assert inc_comp.expected_value == UNDEFINED_RATIO
                incremental_undefined += 1
            else:
                expected_inc = expected_dp / dv
                max_incremental_error = max(
                    max_incremental_error,
                    abs(float(inc_comp.expected_value) - expected_inc),
                    abs(float(inc_comp.expected_value) * dv - expected_dp),
                )
                incremental_numeric += 1
    assert max_effect_error <= max(GEOGRAPHIC_RATIO_TOLERANCE, 1e-6)
    for current, prior in (
        (date(2023, 1, 29), date(2022, 1, 30)),
        (date(2024, 1, 28), date(2023, 1, 29)),
        (date(2025, 2, 2), date(2024, 1, 28)),
        (FY2026, prior_fy2026),
    ):
        curr_snap = snapshots[current]
        prior_snap = snapshots[prior]
        v_t = float(curr_snap.values["net_revenue.consolidated"])
        v_p = float(prior_snap.values["net_revenue.consolidated"])
        p_t = float(curr_snap.values["income_from_operations.consolidated"])
        p_p = float(prior_snap.values["income_from_operations.consolidated"])
        cons_inc_comp = next(
            c
            for c in geo_comps
            if c.family_id
            == "geographic_consolidated_operating_profit_incremental_margin"
            and c.period_end == current.isoformat()
        )
        dv = v_t - v_p
        dp_cons = p_t - p_p
        if dv == 0.0:
            assert cons_inc_comp.expected_value == UNDEFINED_RATIO
            incremental_undefined += 1
        else:
            expected_inc = dp_cons / dv
            max_incremental_error = max(
                max_incremental_error,
                abs(float(cons_inc_comp.expected_value) - expected_inc),
                abs(float(cons_inc_comp.expected_value) * dv - dp_cons),
            )
            incremental_numeric += 1
    assert incremental_numeric + incremental_undefined == GEOGRAPHIC_INCREMENTAL_MARGIN_SPECS
    assert incremental_numeric == GEOGRAPHIC_INCREMENTAL_MARGIN_SPECS
    assert incremental_undefined == 0
    assert max_incremental_error <= max(GEOGRAPHIC_RATIO_TOLERANCE, 1e-6)
    cons_growth_comp = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_consolidated_revenue_growth"
        and c.period_end == FY2026.isoformat()
    )
    residual_comp = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_revenue_growth_contribution_residual"
        and c.period_end == FY2026.isoformat()
    )
    contrib_comps = [
        next(
            c
            for c in geo_comps
            if c.family_id == "geographic_revenue_growth_contribution"
            and c.period_end == FY2026.isoformat()
            and geographic_spec_identity(c) == identity
        )
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES
    ]
    assert residual_comp.formula.replace(" ", "") == (
        resolve_geographic_revenue_growth_contribution_residual_formula(
            SemanticCellRef(
                cons_growth_comp.id,
                cons_growth_comp.semantic_key,
                cons_growth_comp.period_end,
                cons_growth_comp.cell,
                cons_growth_comp.tab,
            ),
            tuple(
                SemanticCellRef(
                    item.id, item.semantic_key, item.period_end, item.cell, item.tab
                )
                for item in contrib_comps
            ),
            from_tab=GEOGRAPHIC_SHEET,
        ).replace(" ", "")
    )
    independent_residual = 100.0 * float(cons_growth_comp.expected_value) - sum(
        float(item.expected_value) for item in contrib_comps
    )
    assert residual_comp.expected_value == pytest.approx(
        independent_residual, abs=GEOGRAPHIC_RATIO_TOLERANCE
    )
    corp_signed = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_signed_reconciling_contribution"
        and c.period_end == FY2026.isoformat()
    )
    assert corp_signed.expected_value == -1397067.0
    assert corp_signed.formula.replace(" ", "").startswith("=")
    assert not corp_signed.formula.replace(" ", "").startswith("=-")
    recon = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_reconstructed_consolidated_operating_profit"
        and c.period_end == FY2026.isoformat()
    )
    assert recon.expected_value == 2210615.0
    assert "+" in recon.formula
    itemized = next(
        c
        for c in geo_comps
        if c.family_id == "geographic_signed_reconciling_contribution"
        and c.period_end == ADMIT_2022.isoformat()
        and "acquisition_related_expenses" in c.id
    )
    assert itemized.formula.replace(" ", "").startswith("=-")
    for period in series.periods:
        family = series.presentation_family[period]
        col_idx = period_cols[period]
        col = get_column_letter(col_idx)
        signed_comps = [
            c
            for c in geo_comps
            if c.family_id == "geographic_signed_reconciling_contribution"
            and c.period_end == period.isoformat()
        ]
        snapshot = snapshots[period]
        for comp in signed_comps:
            identity = geographic_spec_identity(comp)
            operation = snapshot.bridge_operations[identity]
            compact = comp.formula.replace(" ", "")
            if operation == OP_ADD:
                assert family == FAMILY_CORPORATE or operation == OP_ADD
                assert compact.startswith("=") and not compact.startswith("=-")
            else:
                assert operation == OP_SUBTRACT
                assert compact.startswith("=-")
            source_row = _row_by_label(
                aws, f"{geographic_identity_label(identity)} (reported)"
            )
            if operation == OP_ADD:
                assert compact == f"={col}{source_row}"
            else:
                assert compact == f"=-{col}{source_row}"
        for identity in GEOGRAPHIC_SEGMENT_IDENTITIES:
            label = geographic_identity_label(identity)
            rev_row = _row_by_label(aws, f"{label} net revenue")
            ifop_row = _row_by_label(aws, f"{label} income from operations")
            assert aws.cell(rev_row, col_idx).value == snapshot.values[
                f"net_revenue.{identity}"
            ]
            assert tws.cell(rev_row, col_idx).value == aws.cell(rev_row, col_idx).value
            assert aws.cell(ifop_row, col_idx).value == snapshot.values[
                f"income_from_operations.{identity}"
            ]
            assert tws.cell(ifop_row, col_idx).value == aws.cell(
                ifop_row, col_idx
            ).value

    for comp in all_comps:
        row, col = parse_cell_ref(comp.cell)
        trainer_cell = twb[comp.tab].cell(row=row, column=col)
        answer_cell = awb[comp.tab].cell(row=row, column=col)
        assert trainer_cell.value is None
        assert trainer_cell.comment is None
        assert _fill_rgb(trainer_cell) == "FFFF00"
        assert answer_cell.value == comp.formula
        assert _fill_rgb(answer_cell) in WHITE_RGBS
        assert answer_cell.comment is not None
        note = (answer_cell.comment.text or "").strip()
        assert note
        if comp.category == "geographic_segment":
            assert "cause" not in note.lower()
        if comp.family_id in {
            "geographic_revenue_share",
            "geographic_revenue_growth",
            "geographic_signed_reconciling_contribution",
            "geographic_reconstructed_consolidated_operating_profit",
            "geographic_revenue_growth_contribution",
            "geographic_consolidated_revenue_growth",
            "geographic_revenue_growth_contribution_residual",
            "geographic_operating_margin_contribution",
            "geographic_reconciling_operating_margin_contribution",
            "geographic_consolidated_operating_margin",
            "geographic_operating_margin_contribution_residual",
            "geographic_operating_margin_contribution_change",
            "geographic_reconciling_operating_margin_contribution_change",
            "geographic_consolidated_operating_margin_change",
            "geographic_operating_margin_contribution_change_residual",
            "geographic_operating_margin_mix_effect",
            "geographic_operating_margin_within_segment_effect",
            "geographic_operating_margin_mix_within_residual",
            "geographic_operating_profit_amount_change",
            "geographic_reconciling_operating_profit_amount_change",
            "geographic_consolidated_operating_profit_amount_change",
            "geographic_operating_profit_amount_change_residual",
            "geographic_operating_profit_revenue_effect",
            "geographic_operating_profit_margin_effect",
            "geographic_operating_profit_incremental_margin",
            "geographic_consolidated_operating_profit_incremental_margin",
        }:
            assert trainer_cell.comment is None
            assert answer_cell.comment is not None
    awb.close()
    twb.close()

    _assert_visible_parity(trainer, answer, practice)
    _assert_fresh_visible_style(trainer, practice_cells=practice, role="trainer")
    _assert_fresh_visible_style(answer, practice_cells=practice, role="answer_key")
    _assert_answer_key_no_yellow(answer)
    twb = load_workbook(trainer, data_only=False)
    awb = load_workbook(answer, data_only=False)
    assert JUDGMENT_SHEET in twb.sheetnames
    assert JUDGMENT_SHEET in awb.sheetnames
    judgment_rows = _judgment_response_keys(awb)
    assert judgment_rows
    for sheet, coord in judgment_rows:
        answer_cell = awb[sheet][coord]
        trainer_cell = twb[sheet][coord]
        assert answer_cell.value not in (None, "")
        assert _fill_rgb(answer_cell) in WHITE_RGBS
        assert trainer_cell.value is None
        assert trainer_cell.comment is None
        assert _fill_rgb(trainer_cell) == "FFFF00"
    assert NORMALIZATION_JUDGMENT_SHEET not in awb.sheetnames
    twb.close()
    awb.close()
    assert _count_source_unavailable(trainer) == LULULEMON_UNAVAILABLE_DISPLAYS
    assert _count_source_unavailable(answer) == LULULEMON_UNAVAILABLE_DISPLAYS

    answer_hash = _sha256(answer)
    blank = check_workbook(trainer)
    assert (blank.total, blank.blank, blank.correct, blank.incorrect) == (
        LEASE_DT_LULULEMON_SPECS + GEOGRAPHIC_LULULEMON_SPECS,
        LEASE_DT_LULULEMON_SPECS + GEOGRAPHIC_LULULEMON_SPECS,
        0,
        0,
    )
    twb = load_workbook(trainer, data_only=False)
    for comp in all_comps:
        row, col = parse_cell_ref(comp.cell)
        cell = twb[comp.tab].cell(row=row, column=col)
        assert cell.value is None
        assert cell.comment is None
        assert _fill_rgb(cell) == BLANK_RGB
    twb.close()
    dumped_blank = repr(blank)
    assert "income from operations / net revenue" not in dumped_blank.lower()
    assert "NOPAT" not in dumped_blank

    filled_trainer, filled_answer = _copy_pair(
        trainer, answer, tmp_path / "filled_check"
    )
    wb = load_workbook(filled_trainer, data_only=False)
    for comp in all_comps:
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(filled_trainer)
    wb.close()
    filled = check_workbook(filled_trainer)
    assert (filled.total, filled.correct, filled.incorrect, filled.blank) == (
        LEASE_DT_LULULEMON_SPECS + GEOGRAPHIC_LULULEMON_SPECS,
        LEASE_DT_LULULEMON_SPECS + GEOGRAPHIC_LULULEMON_SPECS,
        0,
        0,
    )
    wb = load_workbook(filled_trainer, data_only=False)
    sample = next(c for c in margin_comps)
    sample_row, sample_col = parse_cell_ref(sample.cell)
    assert wb[sample.tab].cell(sample_row, sample_col).value == sample.formula
    assert _fill_rgb(wb[sample.tab].cell(sample_row, sample_col)) == CORRECT_RGB
    wb.close()

    incorrect_trainer, _incorrect_answer = _copy_pair(
        filled_trainer, filled_answer, tmp_path / "incorrect_check"
    )
    bad = next(
        c for c in margin_comps if "americas" in c.id and c.period_end == FY2026.isoformat()
    )
    _inject_formula_and_cached_value(
        incorrect_trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(incorrect_trainer)
    assert bad_summary.incorrect == 1
    assert bad_summary.correct == bad_summary.total - 1
    assert bad_summary.blank == 0
    wb = load_workbook(incorrect_trainer, data_only=False)
    bad_row, bad_col = parse_cell_ref(bad.cell)
    bad_cell = wb[bad.tab].cell(bad_row, bad_col)
    assert bad_cell.value == "=999"
    assert bad_cell.comment is None
    assert _fill_rgb(bad_cell) == INCORRECT_RGB
    ok_row, ok_col = parse_cell_ref(sample.cell)
    if sample.id != bad.id:
        assert _fill_rgb(wb[sample.tab].cell(ok_row, ok_col)) == CORRECT_RGB
    wb.close()
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert bad.formula not in dumped
    assert "income from operations" not in dumped.lower()
    assert _sha256(answer) == answer_hash
    assert _sha256(filled_answer) == answer_hash
    original_blank = check_workbook(trainer)
    assert original_blank.blank == original_blank.total


def test_prior_presentation_mutation_survives_workbook_series():
    from core.ingestion.filing_json import load_extracted_filing
    from core.ingestion.filing_validator import validate_extracted_filing
    from core.tests.test_geographic_segment_analysis import EXTRACTED, SOURCE
    from core.tests.test_geographic_segment_analysis import FY2025_AMERICAS

    _baseline_reconciled, _baseline_fin, baseline = _reload_admitted_lululemon()
    fy2022 = load_extracted_filing(EXTRACTED / "LULU_FY2022.json")
    fy2023 = load_extracted_filing(EXTRACTED / "LULU_FY2023.json")
    fy2024 = load_extracted_filing(EXTRACTED / "LULU_FY2024.json")
    fy2025 = _mutate_fy2025_prior_2025(
        load_extracted_filing(EXTRACTED / "LULU_FY2025.json")
    )
    validated = []
    for filing in (fy2022, fy2023, fy2024, fy2025):
        report = validate_extracted_filing(filing, source_root=SOURCE)
        assert report.ok
        validated.append((filing, report))
    reconciled = reconcile_filings(validated, admit_periods=(ADMIT_2022,))
    restored = standardized_from_payload(
        json.loads(json.dumps(standardized_to_payload(standardize_reconciled(reconciled))))
    )
    series = compute_geographic_segment_series(restored)
    baseline_series = compute_geographic_segment_series(baseline)
    assert series.net_revenue[FY2025_AMERICAS]["americas"] == 7928156
    assert (
        series.net_revenue[FY2025_AMERICAS]["americas"]
        == baseline_series.net_revenue[FY2025_AMERICAS]["americas"]
    )
    payload = standardized_to_payload(restored)
    assert "7928256" not in json.dumps(payload.get("historical_segment"))
    builder = ReferenceModelBuilder(restored)
    assert builder.geographic_series.net_revenue[FY2025_AMERICAS]["americas"] == 7928156
    assert any(
        s.family_id == "geographic_calculated_segment_revenue_total"
        and s.period_end == FY2025_AMERICAS.isoformat()
        for s in builder.geographic_specs
    )

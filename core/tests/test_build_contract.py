"""Complete Build is selected once, and verified by semantic identity."""

from dataclasses import replace
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.__main__ import main
from core.data.standardized_io import standardized_from_payload
from core.engine.reference_model import ReferenceModelBuilder
from core.trainer.semantic_io import load_semantic_map
from core.trainer.workbook import build_training_workbook

ROOT = Path(__file__).resolve().parents[2]


def financials(company="lululemon"):
    return standardized_from_payload(json.loads(
        (ROOT / "benchmark" / company / "reconciled/standardized.json").read_text()
    ))


def extension(**changes):
    from core.engine.build_contract import BuildModule, WorkbookWriter
    from core.engine.component_catalog import ComponentSpec

    def prepare(builder, start_order):
        return (ComponentSpec(
            id="test_complete", family_id="revenue_link", order=start_order,
            family_order=1, title="Test completed module", short_hint="Link revenue",
            semantic_key="test.complete", category="historical",
            tab_template="Test Complete", period_index=0,
            period_end=builder.periods[0].isoformat(),
        ),)

    def write(builder, wb):
        ws = wb.create_sheet("Test Complete")
        # New workbook module using an already-supported analytical calculation.
        revenue = next(c for c in builder.semantic_map.all_ordered()
                       if c.family_id == "revenue_link" and c.period_index == 0)
        ws["A1"] = f"='{revenue.tab}'!{revenue.cell}"
        builder.semantic_map.register(
            builder.module_specs["test_complete"][0], ws.title, 1, 1,
            ws["A1"].value, revenue.expected_value,
        )

    return replace(BuildModule(
        id="test_complete", status="complete", workbook_capable=True,
        complete_analysis=True, prepare=prepare,
        writers=(WorkbookWriter("test_complete", 100, write),),
        depends_on=("historical",),
    ), **changes)


def register(monkeypatch, module):
    from core.engine import build_contract
    monkeypatch.setattr(build_contract, "BUILD_MODULES", build_contract.BUILD_MODULES + (module,))


def test_new_completed_module_reaches_cli_without_dispatch_changes(monkeypatch, tmp_path):
    register(monkeypatch, extension())
    source = ROOT / "benchmark/lululemon/reconciled/standardized.json"
    assert main(["build", str(source), "-o", str(tmp_path / "Example")]) == 0
    answer = tmp_path / "Example_BAV.xlsx"
    assert load_semantic_map(answer).all_ordered()[-1].semantic_key == "test.complete"
    wb = load_workbook(answer)
    assert wb["Test Complete"]["A1"].value.startswith("='Condensed Financials'!")
    wb.close()


@pytest.mark.parametrize("changes", [
    {"status": "incomplete"}, {"status": "deferred"},
    {"workbook_capable": False}, {"complete_analysis": False},
    {"applicable": lambda builder: False},
])
def test_ineligible_modules_do_not_reach_workbook(monkeypatch, tmp_path, changes):
    register(monkeypatch, extension(**changes))
    _, answer = build_training_workbook(financials(), tmp_path / "Excluded")
    assert "test.complete" not in {c.semantic_key for c in load_semantic_map(answer).all_ordered()}
    wb = load_workbook(answer)
    assert "Test Complete" not in wb.sheetnames
    wb.close()


def test_required_package_checked_before_applicability(monkeypatch, tmp_path):
    from core.engine.build_contract import RequiredInput
    register(monkeypatch, extension(
        required_inputs=(RequiredInput("management-KPI", lambda fin: fin.historical_operating_kpis is not None),),
        applicable=lambda builder: False,
    ))
    with pytest.raises(ValueError, match="test_complete.*management-KPI"):
        build_training_workbook(financials(), tmp_path / "Missing")
    assert not list(tmp_path.glob("*.xlsx"))


@pytest.mark.parametrize("company", ["lululemon", "fast_retailing"])
def test_all_registered_specs_emitted_in_repeatable_order(company, tmp_path):
    from core.engine.build_contract import verify_complete_build
    builders = [ReferenceModelBuilder(financials(company)) for _ in range(2)]
    maps = [b.build(tmp_path / str(i) / "answer.xlsx") for i, b in enumerate(builders)]
    for builder, smap in zip(builders, maps):
        verify_complete_build(builder.expected_specs, smap)
        expected = [s.semantic_key for specs in builder.module_specs.values() for s in specs]
        assert [c.semantic_key for c in smap.all_ordered()] == expected
    assert [c.to_dict() for c in maps[0].all_ordered()] == [c.to_dict() for c in maps[1].all_ordered()]


def test_verification_rejects_same_count_wrong_identity(tmp_path):
    from core.engine.build_contract import verify_complete_build
    builder = ReferenceModelBuilder(financials())
    smap = builder.build(tmp_path / "answer.xlsx")
    smap.all_ordered()[0].semantic_key = "substituted"
    with pytest.raises(ValueError, match="semantic"):
        verify_complete_build(builder.expected_specs, smap)


@pytest.mark.parametrize("company", ["lululemon", "fast_retailing"])
def test_release_build_and_verification_accept_registered_extension(monkeypatch, tmp_path, company):
    register(monkeypatch, extension())
    source = ROOT / "benchmark" / company / "reconciled/standardized.json"
    fin = financials(company)
    _, generic_answer = build_training_workbook(fin, tmp_path / "generic/Example")
    if company == "lululemon":
        from scripts import build_lululemon_release as release
        trainer, answer = release.build_workbooks(source, tmp_path / "release")
        release.verify_staged(trainer, answer, source)
    else:
        from scripts import build_fast_retailing_release as release
        monkeypatch.setattr(release, "RELEASE", tmp_path / "release")
        monkeypatch.setattr(release, "SUPPORTING", source.parent)
        trainer, answer = release.build_workbooks(source)
        release.verify_release(trainer, answer)
    generic = load_semantic_map(generic_answer).all_ordered()
    released = load_semantic_map(answer).all_ordered()
    assert [c.semantic_key for c in released] == [c.semantic_key for c in generic]
    assert released[-1].semantic_key == "test.complete"


def test_supplied_kpi_package_with_missing_period_preserves_unavailability(monkeypatch, tmp_path):
    from core.data.interface import HistoricalOperatingKpiData, HistoricalOperatingKpiObservation
    from core.engine.build_contract import RequiredInput
    from core.model.operating_kpi import compute_operating_kpi_series
    from core.model.ratio_values import SOURCE_UNAVAILABLE

    register(monkeypatch, extension(required_inputs=(
        RequiredInput("management-KPI", lambda fin: fin.historical_operating_kpis is not None),
    )))
    fin = financials()
    last = max(fin.fiscal_years())
    fin.historical_operating_kpis = HistoricalOperatingKpiData(observations=[
        HistoricalOperatingKpiObservation("store_count", "company_operated", last, 10, "stores"),
    ])
    build_training_workbook(fin, tmp_path / "Partial")
    series = compute_operating_kpi_series(fin)
    assert series.period_end_count[last] == 10
    assert series.net_count_change[last] == SOURCE_UNAVAILABLE
    assert series.period_end_count[min(fin.fiscal_years())] == SOURCE_UNAVAILABLE


def test_undisclosed_interest_retains_workbook_source_unavailable(tmp_path):
    from core.model.ratio_values import SOURCE_UNAVAILABLE
    builder = ReferenceModelBuilder(financials())
    answer = tmp_path / "answer.xlsx"
    smap = builder.build(answer)
    assert "nopat_fy" not in {c.family_id for c in smap.all_ordered()}
    wb = load_workbook(answer)
    row = builder.rowmap["condensed_nopat_row"]
    assert [wb["Condensed Financials"].cell(row, i + 2).value
            for i in range(len(builder.periods))] == [SOURCE_UNAVAILABLE] * len(builder.periods)
    wb.close()


@pytest.mark.parametrize("changes,match", [
    ({"prepare": None}, "integration"),
    ({"writers": ()}, "integration"),
    ({"depends_on": ("not_registered",)}, "dependency"),
    ({"id": "historical"}, "Duplicate Build module"),
])
def test_invalid_registration_fails_before_output(monkeypatch, tmp_path, changes, match):
    register(monkeypatch, extension(**changes))
    with pytest.raises(ValueError, match=match):
        build_training_workbook(financials(), tmp_path / "Invalid")
    assert not list(tmp_path.glob("*.xlsx"))


def test_missing_input_is_a_clear_cli_error(monkeypatch, tmp_path, capsys):
    from core.engine.build_contract import RequiredInput
    register(monkeypatch, extension(required_inputs=(
        RequiredInput("management-KPI", lambda fin: fin.historical_operating_kpis is not None),
    )))
    source = ROOT / "benchmark/lululemon/reconciled/standardized.json"
    assert main(["build", str(source), "-o", str(tmp_path / "Missing")]) == 1
    error = capsys.readouterr().err
    assert "test_complete: missing required input: management-KPI" in error
    assert not list(tmp_path.glob("*.xlsx"))


@pytest.mark.parametrize("company", ["lululemon", "fast_retailing"])
def test_release_rejects_workbook_from_older_contract(monkeypatch, tmp_path, company):
    source = ROOT / "benchmark" / company / "reconciled/standardized.json"
    trainer, answer = build_training_workbook(financials(company), tmp_path / "Old")
    register(monkeypatch, extension())
    with pytest.raises((ValueError, RuntimeError), match="semantic component mismatch"):
        if company == "lululemon":
            from scripts.build_lululemon_release import verify_staged
            verify_staged(trainer, answer, source)
        else:
            from scripts import build_fast_retailing_release as release
            monkeypatch.setattr(release, "SUPPORTING", source.parent)
            release.verify_release(trainer, answer)


def test_writer_omitting_registered_component_fails(monkeypatch, tmp_path):
    from core.engine.build_contract import WorkbookWriter
    register(monkeypatch, extension(writers=(
        WorkbookWriter("broken", 100, lambda builder, wb: None),
    )))
    with pytest.raises(ValueError, match="Missing component: test_complete"):
        build_training_workbook(financials(), tmp_path / "Broken")
    assert not list(tmp_path.glob("*.xlsx"))


def test_applicable_module_requires_applicable_dependency(monkeypatch):
    register(monkeypatch, extension(applicable=lambda builder: False))
    register(monkeypatch, extension(id="dependent", depends_on=("test_complete",)))
    with pytest.raises(ValueError, match="dependent.*test_complete.*not applicable"):
        ReferenceModelBuilder(financials())

"""Step 9K.1 — PP&E / D&A fixed-asset intensity diagnostics."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    FIXED_ASSET_COMPONENT_CATALOG,
    expand_fixed_asset_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.fixed_asset import (
    compute_fixed_asset_series,
    fixed_asset_applicable,
    fixed_asset_availability,
)
from core.model.historical_expected import fixed_asset_expected_series
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _tiny(*, with_ppe=True, with_da=True, revenue=(1000.0, 1100.0), ppe=(100.0, 120.0), da=(-10.0, -12.0)):
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)

    def vals(a, b):
        return {d1: a, d2: b}

    # Keep BS identity: cash + AR + PPE = AP + bank + equity
    cash, ar, ap, bank = 50.0, 40.0, 30.0, 20.0
    ppe0, ppe1 = ppe
    eq0 = cash + ar + ppe0 - ap - bank
    eq1 = cash + ar + ppe1 - ap - bank
    bs = [
        _li("Cash and cash equivalents", vals(cash, cash)),
        _li("Trade receivables", vals(ar, ar)),
        _li("Trade payables", vals(ap, ap)),
        _li("Bank borrowings", vals(bank, bank)),
        _li("Total equity", vals(eq0, eq1)),
    ]
    if with_ppe:
        bs.insert(2, _li("Property, plant and equipment", vals(*ppe)))
    cf = [
        _li("Net cash from operating activities", vals(80, 90)),
    ]
    if with_da:
        cf.insert(0, _li("Depreciation and amortisation", vals(*da)))
    return StandardizedFinancials(
        ticker="FA",
        company_name="Fixed Asset Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=[
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ],
        income_statement=[
            _li("Revenue", vals(*revenue)),
            _li("Finance costs", vals(-10, -11)),
            _li("Finance income", vals(0, 0)),
            _li("Profit before tax", vals(200, 220)),
            _li("Income tax expense", vals(-30, -33)),
            _li("Profit for the year", vals(170, 187)),
        ],
        balance_sheet=bs,
        cash_flow=cf,
    )


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(FIXED_ASSET_COMPONENT_CATALOG) == 8
    assert [f.order for f in FIXED_ASSET_COMPONENT_CATALOG] == list(range(79, 87))
    specs = expand_fixed_asset_specs(periods, start_order=1000)
    assert len(FIXED_ASSET_COMPONENT_CATALOG) == 8
    assert len(specs) == 35
    assert {s.family_order for s in specs} == set(range(79, 87))
    with pytest.raises(ValueError, match="duplicate"):
        expand_fixed_asset_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_fixed_asset_specs(list(reversed(periods)), start_order=1)


def test_canonical_demo_applicability_and_fy2022_math():
    data = _ingest_demo()
    periods = list(canonical_fiscal_periods(data))
    assert fixed_asset_applicable(data)
    avail = fixed_asset_availability(data)
    assert avail.ppe and avail.depreciation_amortization
    anchor = compute_anchor(data, periods)
    series = compute_fixed_asset_series(data, periods, anchor)
    assert series.average_ppe[0] is None
    assert series.ppe_turnover[0] is None
    assert series.ppe_intensity[0] is None
    assert series.ppe_change[0] is None
    assert series.da_to_average_ppe[0] is None
    assert series.ppe[0] == pytest.approx(3200.0)
    assert series.depreciation_amortization[0] == pytest.approx(340.0)
    assert series.da_to_revenue[0] == pytest.approx(340.0 / 8500.0)
    assert series.average_ppe[1] == pytest.approx((3200.0 + 3400.0) / 2.0)
    assert series.ppe_turnover[1] == pytest.approx(9200.0 / 3300.0)
    assert series.ppe_intensity[1] == pytest.approx(3300.0 / 9200.0)
    assert series.ppe_change[1] == pytest.approx(200.0)
    assert series.da_to_revenue[1] == pytest.approx(368.0 / 9200.0)
    assert series.da_to_average_ppe[1] == pytest.approx(368.0 / 3300.0)
    mapped = fixed_asset_expected_series(series)
    assert set(mapped) == {f.id for f in FIXED_ASSET_COMPONENT_CATALOG}


def test_ppe_only_or_da_only_not_applicable():
    assert not fixed_asset_applicable(_tiny(with_ppe=True, with_da=False))
    assert not fixed_asset_applicable(_tiny(with_ppe=False, with_da=True))
    assert fixed_asset_applicable(_tiny(with_ppe=True, with_da=True))


def test_zero_denominators_and_negative_da_preserved():
    fin = _tiny(revenue=(0.0, 0.0), ppe=(0.0, 0.0), da=(-10.0, -12.0))
    periods = list(canonical_fiscal_periods(fin))
    anchor = compute_anchor(fin, periods)
    series = compute_fixed_asset_series(fin, periods, anchor)
    assert series.da_to_revenue[0] == UNDEFINED_RATIO
    assert series.ppe_turnover[1] == UNDEFINED_RATIO
    assert series.ppe_intensity[1] == UNDEFINED_RATIO
    assert series.da_to_average_ppe[1] == UNDEFINED_RATIO
    assert series.depreciation_amortization[1] == pytest.approx(-12.0)

    fin2 = _tiny(revenue=(1000.0, 1100.0), ppe=(100.0, 120.0), da=(0.0, 0.0))
    periods2 = list(canonical_fiscal_periods(fin2))
    series2 = compute_fixed_asset_series(fin2, periods2, compute_anchor(fin2, periods2))
    assert series2.da_to_revenue[1] == pytest.approx(0.0)
    assert series2.da_to_average_ppe[1] == pytest.approx(0.0)


def test_missing_period_value_fails_closed():
    fin = _tiny()
    periods = list(canonical_fiscal_periods(fin))
    ppe = next(i for i in fin.balance_sheet if i.label.startswith("Property"))
    del ppe.values[periods[1]]
    with pytest.raises(MissingHistoricalValueError):
        compute_fixed_asset_series(fin, periods, compute_anchor(fin, periods))


def test_canonical_demo_surface_and_alt_dupont_section(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "FA_BASE.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 70
    assert len(smap.all_ordered()) == 294
    fa = [c for c in smap.all_ordered() if c.family_id in {f.id for f in FIXED_ASSET_COMPONENT_CATALOG}]
    assert len({c.family_id for c in fa}) == 8
    assert len(fa) == 35
    wb = load_workbook(answer)
    assert "FIXED-ASSET INTENSITY CONTEXT" in [
        wb["ALT DuPont"].cell(row=r, column=1).value
        for r in range(1, (wb["ALT DuPont"].max_row or 1) + 1)
    ]
    assert "Fixed Asset Analysis" not in wb.sheetnames
    wb.close()
    summary = check_workbook(trainer)
    assert (summary.correct, summary.incorrect, summary.blank) == (0, 0, 294)


def test_ppe_only_omits_module(tmp_path):
    fin = _tiny(with_ppe=True, with_da=False)
    builder = ReferenceModelBuilder(fin)
    assert builder.fixed_asset_specs == ()
    trainer, answer = build_training_workbook(fin, tmp_path / "PPE_ONLY.xlsx")
    smap = load_semantic_map(answer)
    assert not any(
        c.family_id in {f.id for f in FIXED_ASSET_COMPONENT_CATALOG}
        for c in smap.all_ordered()
    )


def test_check_accepts_formula_and_rejects_wrong_cache(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "FA_CHK.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "ppe_turnover" and isinstance(c.expected_value, (int, float))
    )
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct >= 1

    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula=f"={float(comp.expected_value)}",
        cached_value=float(comp.expected_value),
    )
    assert check_workbook(trainer).correct >= 1

    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad = check_workbook(trainer)
    assert bad.incorrect >= 1
    wb = load_workbook(trainer, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFC7CE"
    wb.close()


def test_trusted_fixed_asset_check_tamper_fails_before_recolor(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "FA_TAMPER.xlsx")
    smap = load_semantic_map(answer)
    comp = next(c for c in smap.all_ordered() if c.family_id == "ppe_source_link")
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    ws = wb["ALT DuPont"]
    check_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "FIXED-ASSET INTENSITY CHECK"
    )
    ws.cell(check_row, 3).value = '="TAMPERED"'
    wb.save(trainer)
    wb.close()
    with pytest.raises(ValueError, match="Trusted workbook cell was modified"):
        check_workbook(trainer)
    wb = load_workbook(trainer, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

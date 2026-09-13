"""Step 9M.5 — goodwill / intangible-asset intensity & change diagnostics."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    GOODWILL_INTANGIBLES_COMPONENT_CATALOG,
    expand_goodwill_intangibles_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.goodwill_intangibles import (
    compute_goodwill_intangibles_series,
    goodwill_intangibles_applicable,
    goodwill_intangibles_availability,
)
from core.model.historical_expected import goodwill_intangibles_expected_series
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.source_values import MissingHistoricalValueError
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"


def _ingest_demo():
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _tiny(
    *,
    goodwill=(80.0, 80.0),
    intangibles=(100.0, 120.0),
    payments=(-10.0, -12.0),
    revenue=(1000.0, 1100.0),
    with_goodwill: bool = True,
    with_intangibles: bool = True,
    with_payments: bool = True,
    duplicate_goodwill: bool = False,
    missing_period: bool = False,
):
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)

    def vals(a, b):
        return {d1: a, d2: b}

    cash, ar, ap, bank = 50.0, 40.0, 30.0, 20.0
    gw0 = goodwill[0] if with_goodwill else 0.0
    gw1 = goodwill[1] if with_goodwill else 0.0
    ia0 = intangibles[0] if with_intangibles else 0.0
    ia1 = intangibles[1] if with_intangibles else 0.0
    eq0 = cash + ar + gw0 + ia0 - ap - bank
    eq1 = cash + ar + gw1 + ia1 - ap - bank
    bs = [
        _li("Cash and cash equivalents", vals(cash, cash)),
        _li("Trade receivables", vals(ar, ar)),
        _li("Trade payables", vals(ap, ap)),
        _li("Bank borrowings", vals(bank, bank)),
        _li("Total equity", vals(eq0, eq1)),
    ]
    if with_goodwill:
        gw_vals = {d1: goodwill[0]} if missing_period else vals(*goodwill)
        bs.insert(2, _li("Goodwill", gw_vals, concept="goodwill"))
        if duplicate_goodwill:
            bs.insert(3, _li("Goodwill duplicate", vals(*goodwill), concept="goodwill"))
    if with_intangibles:
        bs.insert(
            2 if not with_goodwill else 3,
            _li("Intangible assets", vals(*intangibles), concept="intangible_assets"),
        )
    cf = [
        _li("Net cash from operating activities", vals(80, 90)),
    ]
    if with_payments:
        cf.insert(
            0,
            _li(
                "Payments for intangible assets",
                vals(*payments),
                concept="payments_for_intangible_assets",
            ),
        )
    return StandardizedFinancials(
        ticker="GI",
        company_name="Goodwill Intangibles Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=[
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ],
        income_statement=[
            _li("Revenue", vals(*revenue)),
            _li("Finance costs", vals(-10, -11)),
            _li("Finance income", vals(1, 1)),
            _li("Profit before tax", vals(100, 110)),
            _li("Income tax expense", vals(-16, -18)),
            _li("Profit for the year", vals(84, 92)),
        ],
        balance_sheet=bs,
        cash_flow=cf,
    )


def test_catalog_orders_and_expand_filters():
    assert [f.order for f in GOODWILL_INTANGIBLES_COMPONENT_CATALOG] == list(
        range(98, 112)
    )
    periods = [date(2024, 12, 31), date(2025, 12, 31)]
    from core.model.goodwill_intangibles import GoodwillIntangiblesAvailability

    full = GoodwillIntangiblesAvailability(
        goodwill=True,
        intangible_assets=True,
        goodwill_and_intangibles=True,
        payments_for_intangible_assets=True,
    )
    specs = expand_goodwill_intangibles_specs(periods, start_order=1000, availability=full)
    # 4 balance families × 3 lines × 1 comparable + 2 payment families × 2 periods
    assert len(specs) == 4 * 3 * 1 + 2 * 2

    gw_only = GoodwillIntangiblesAvailability(
        goodwill=True,
        intangible_assets=False,
        goodwill_and_intangibles=False,
        payments_for_intangible_assets=False,
    )
    gw_specs = expand_goodwill_intangibles_specs(
        periods, start_order=1, availability=gw_only
    )
    assert {s.family_id for s in gw_specs} == {
        "goodwill_change",
        "goodwill_growth",
        "average_goodwill",
        "goodwill_to_revenue",
    }
    assert len(gw_specs) == 4

    with pytest.raises(ValueError):
        expand_goodwill_intangibles_specs(
            [periods[0], periods[0]], start_order=1, availability=full
        )
    with pytest.raises(ValueError):
        expand_goodwill_intangibles_specs(
            list(reversed(periods)), start_order=1, availability=full
        )


def test_partial_availability_and_math():
    both = _tiny()
    assert goodwill_intangibles_applicable(both)
    avail = goodwill_intangibles_availability(both)
    assert avail.goodwill and avail.intangible_assets
    assert avail.goodwill_and_intangibles and avail.payments_for_intangible_assets

    periods = canonical_fiscal_periods(both)
    series = compute_goodwill_intangibles_series(
        both, periods, compute_anchor(both, periods)
    )
    assert series.goodwill is not None
    assert series.goodwill.change[0] is None
    assert series.goodwill.change[1] == pytest.approx(0.0)
    assert series.goodwill.growth[1] == pytest.approx(0.0)
    assert series.goodwill.average[1] == pytest.approx(80.0)
    assert series.goodwill.to_revenue[1] == pytest.approx(80.0 / 1100.0)

    assert series.intangible_assets is not None
    assert series.intangible_assets.change[1] == pytest.approx(20.0)
    assert series.intangible_assets.growth[1] == pytest.approx(0.2)
    assert series.intangible_assets.average[1] == pytest.approx(110.0)

    assert series.goodwill_and_intangibles is not None
    assert series.goodwill_and_intangibles.level[1] == pytest.approx(200.0)

    assert series.payments_reported == (-10.0, -12.0)
    assert series.intangible_payments == (10.0, 12.0)
    assert series.intangible_payments_to_revenue[1] == pytest.approx(12.0 / 1100.0)

    mapped = goodwill_intangibles_expected_series(series, avail)
    assert "goodwill_change" in mapped
    assert "intangible_payments" in mapped

    gw_only = _tiny(with_intangibles=False, with_payments=False)
    avail_gw = goodwill_intangibles_availability(gw_only)
    assert avail_gw.goodwill and not avail_gw.intangible_assets
    assert not avail_gw.goodwill_and_intangibles
    assert not avail_gw.payments_for_intangible_assets
    series_gw = compute_goodwill_intangibles_series(
        gw_only, periods, compute_anchor(gw_only, periods)
    )
    assert series_gw.goodwill is not None
    assert series_gw.intangible_assets is None
    assert series_gw.goodwill_and_intangibles is None
    assert series_gw.intangible_payments is None

    ia_only = _tiny(with_goodwill=False, with_payments=True)
    avail_ia = goodwill_intangibles_availability(ia_only)
    assert not avail_ia.goodwill and avail_ia.intangible_assets
    assert not avail_ia.goodwill_and_intangibles
    assert avail_ia.payments_for_intangible_assets

    none = _tiny(with_goodwill=False, with_intangibles=False, with_payments=True)
    assert not goodwill_intangibles_applicable(none)
    avail_none = goodwill_intangibles_availability(none)
    assert not avail_none.goodwill and not avail_none.intangible_assets
    # Payments alone do not activate the module.
    assert avail_none.payments_for_intangible_assets


def test_duplicates_missing_periods_zero_denominators_mixed_signs():
    dup = _tiny(duplicate_goodwill=True)
    avail = goodwill_intangibles_availability(dup)
    assert not avail.goodwill
    assert avail.intangible_assets
    assert not avail.goodwill_and_intangibles

    missing = _tiny(missing_period=True)
    periods = canonical_fiscal_periods(missing)
    with pytest.raises(MissingHistoricalValueError):
        compute_goodwill_intangibles_series(
            missing, periods, compute_anchor(missing, periods)
        )

    zero_rev = _tiny(revenue=(1000.0, 0.0), goodwill=(0.0, 10.0))
    periods = canonical_fiscal_periods(zero_rev)
    series = compute_goodwill_intangibles_series(
        zero_rev, periods, compute_anchor(zero_rev, periods)
    )
    assert series.goodwill is not None
    assert series.goodwill.growth[1] == UNDEFINED_RATIO  # prior goodwill 0
    assert series.goodwill.to_revenue[1] == UNDEFINED_RATIO
    assert series.intangible_payments_to_revenue is not None
    assert series.intangible_payments_to_revenue[1] == UNDEFINED_RATIO

    mixed = _tiny(payments=(-10.0, 5.0))
    periods = canonical_fiscal_periods(mixed)
    series_m = compute_goodwill_intangibles_series(
        mixed, periods, compute_anchor(mixed, periods)
    )
    assert series_m.intangible_payments == (10.0, -5.0)


def test_demo_omits_label_only_goodwill():
    data = _ingest_demo()
    assert not goodwill_intangibles_applicable(data)
    builder = ReferenceModelBuilder(data)
    assert builder.goodwill_intangibles_series is None
    assert builder.goodwill_intangibles_specs == ()


def test_workbook_partial_and_check(tmp_path):
    data = _tiny(with_payments=False)
    trainer, answer = build_training_workbook(data, tmp_path / "GI_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.goodwill_intangibles_series is not None
    assert builder.goodwill_intangibles_availability.payments_for_intangible_assets is False
    families = {c.family_id for c in builder.goodwill_intangibles_specs}
    assert "intangible_payments" not in families
    assert "goodwill_change" in families
    assert "goodwill_and_intangibles_change" in families

    smap = load_semantic_map(answer)
    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    gi_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in GOODWILL_INTANGIBLES_COMPONENT_CATALOG}
    ]
    assert gi_comps
    wb_answer = load_workbook(answer)
    for comp in gi_comps:
        row, col = parse_cell_ref(comp.cell)
        cell = wb_answer[comp.tab].cell(row=row, column=col)
        assert cell.comment is not None
        assert (cell.comment.text or "").strip()
    wb_answer.close()

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in gi_comps if c.family_id == "goodwill_change")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "Change in Goodwill" not in dumped


def _dupont_row_by_label(ws, label: str) -> int:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"ALT DuPont label not found: {label!r}")


@pytest.mark.parametrize(
    "stem,kwargs,balance_labels",
    [
        (
            "BOTH",
            {},
            ("Goodwill", "Intangible Assets", "Goodwill & Intangibles"),
        ),
        (
            "GW",
            {"with_intangibles": False, "with_payments": False},
            ("Goodwill",),
        ),
        (
            "IA",
            {"with_goodwill": False},
            ("Intangible Assets",),
        ),
    ],
)
def test_first_period_balance_diagnostics_blank(tmp_path, stem, kwargs, balance_labels):
    data = _tiny(**kwargs)
    trainer, answer = build_training_workbook(data, tmp_path / f"GI_FP_{stem}.xlsx")
    smap = load_semantic_map(answer)
    practice_cells = {(c.tab, c.cell) for c in smap.all_ordered()}

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        for balance in balance_labels:
            level_row = _dupont_row_by_label(ws, balance)
            diag_rows = (
                _dupont_row_by_label(ws, f"Change in {balance}"),
                _dupont_row_by_label(ws, f"{balance} Growth"),
                _dupont_row_by_label(ws, f"Average {balance}"),
                _dupont_row_by_label(ws, f"Average {balance} / Revenue"),
            )
            assert ws.cell(row=level_row, column=2).value is not None
            for row in diag_rows:
                cell = ws.cell(row=row, column=2)
                assert cell.value is None
                assert ("ALT DuPont", cell.coordinate) not in practice_cells

                later = ws.cell(row=row, column=3)
                if path == answer:
                    assert isinstance(later.value, str) and later.value.startswith("=")
                else:
                    assert later.value is None
        wb.close()


def test_demo_practice_counts_unchanged(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "GI_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312
    assert check_workbook(trainer).blank == 312

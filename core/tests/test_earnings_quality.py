"""Step 9A — historical earnings-quality cash conversion and accruals."""

from __future__ import annotations

import copy
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    QUALITY_COMPONENT_CATALOG,
    expand_quality_specs,
)
from core.engine.reference_model import EARNINGS_QUALITY_SHEET, ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.earnings_quality import (
    UNDEFINED_RATIO,
    compute_earnings_quality_series,
    earnings_quality_availability,
    resolve_sbc_source,
    sbc_diagnostics_applicable,
)
from core.model.financial_math import compute_anchor
from core.model.historical_expected import earnings_quality_expected_series
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.source_values import MissingHistoricalValueError
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    adapter = HKManualDocumentAdapter()
    return adapter.ingest([DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)])


def _li(label, values, *, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _tiny_fin(*, with_cfo=True, with_assets=True, ambiguous_cfo=False):
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
    periods = [
        FinancialPeriod(end_date=d1, label="FY2024"),
        FinancialPeriod(end_date=d2, label="FY2025"),
    ]
    income = [
        _li("Revenue", {d1: 1000, d2: 1100}, concept="revenue"),
        _li("Profit before tax", {d1: 200, d2: 220}),
        _li("Income tax expense", {d1: -30, d2: -33}),
        _li("Finance costs", {d1: -10, d2: -11}),
        _li("Finance income", {d1: 0.0, d2: 0.0}),
        _li("Profit for the year", {d1: 100, d2: 120}),
    ]
    balance = [
        _li("Cash", {d1: 50, d2: 60}),
        _li("Trade receivables", {d1: 40, d2: 45}),
        _li("Property, plant and equipment", {d1: 200, d2: 210}),
    ]
    if with_assets:
        balance.append(_li("Total assets", {d1: 290, d2: 315}))
    balance.extend(
        [
            _li("Trade payables", {d1: 20, d2: 22}),
            _li("Bank borrowings", {d1: 80, d2: 85}),
            _li("Total equity", {d1: 190, d2: 208}),
        ]
    )
    cash_flow = []
    if with_cfo:
        cash_flow.append(
            _li("Net cash from operating activities", {d1: 80, d2: 90})
        )
    if ambiguous_cfo:
        cash_flow = [
            _li("Net cash from operating activities", {d1: 80, d2: 90}),
            _li("Net cash generated from operating activities", {d1: 70, d2: 75}),
        ]
    if not cash_flow:
        cash_flow.append(_li("Net cash from financing activities", {d1: -10, d2: -12}))
    return StandardizedFinancials(
        ticker="EQ",
        company_name="EQ Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=income,
        balance_sheet=balance,
        cash_flow=cash_flow,
    )


def test_operating_cash_flow_concept_and_aliases():
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
    item = resolve_line(
        [_li("Weird label", {d1: 1, d2: 2}, concept="operating_cash_flow")],
        "operating_cash_flow",
    )
    assert item.item is not None
    assert item.item.concept == "operating_cash_flow"

    for label in (
        "Net cash from operating activities",
        "Net cash generated from operating activities",
        "Net cash provided by operating activities",
        "Net cash flow from operating activities",
        "Net cash flows from operating activities",
    ):
        resolved = resolve_line([_li(label, {d1: 1, d2: 2})], "operating_cash_flow")
        assert resolved.item is not None
        assert resolved.item.label == label

    generic = resolve_line(
        [_li("Cash generated from operations", {d1: 1, d2: 2})],
        "operating_cash_flow",
        required=False,
    )
    assert generic.item is None

    with pytest.raises(AmbiguousLineError):
        resolve_line(
            [
                _li("Net cash from operating activities", {d1: 1, d2: 2}),
                _li("Net cash generated from operating activities", {d1: 1, d2: 2}),
            ],
            "operating_cash_flow",
        )


def test_earnings_quality_series_core_cases():
    fin = _tiny_fin()
    periods = canonical_fiscal_periods(fin)
    anchor = compute_anchor(fin, periods)
    series = compute_earnings_quality_series(fin, periods, anchor)
    assert series.operating_cash_flow == (80.0, 90.0)
    assert series.cash_conversion_ratio[0] == pytest.approx(80.0 / 100.0)
    assert series.total_accruals == (20.0, 30.0)
    assert series.average_total_assets[0] is None
    assert series.accrual_ratio[0] is None
    assert series.average_total_assets[1] == pytest.approx((290.0 + 315.0) / 2.0)
    assert series.accrual_ratio[1] == pytest.approx(
        30.0 / series.average_total_assets[1]
    )

    # CFO > NI -> negative accruals
    fin2 = _tiny_fin()
    fin2.income_statement[-1].values[periods[0]] = 50.0
    fin2.cash_flow[0].values[periods[0]] = 80.0
    anchor2 = compute_anchor(fin2, periods)
    series2 = compute_earnings_quality_series(fin2, periods, anchor2)
    assert series2.total_accruals[0] == pytest.approx(-30.0)

    # Zero NI -> undefined cash conversion
    fin3 = _tiny_fin()
    fin3.income_statement[-1].values[periods[0]] = 0.0
    anchor3 = compute_anchor(fin3, periods)
    series3 = compute_earnings_quality_series(fin3, periods, anchor3)
    assert series3.cash_conversion_ratio[0] == UNDEFINED_RATIO
    assert series3.total_accruals[0] == pytest.approx(-80.0)

    # Zero CFO, nonzero NI -> numeric 0.0 conversion
    fin4 = _tiny_fin()
    fin4.cash_flow[0].values[periods[0]] = 0.0
    series4 = compute_earnings_quality_series(
        fin4, periods, compute_anchor(fin4, periods)
    )
    assert series4.cash_conversion_ratio[0] == 0.0
    assert series4.total_accruals[0] == pytest.approx(100.0)

    # Both zero -> still undefined (denominator rules)
    fin5 = _tiny_fin()
    fin5.income_statement[-1].values[periods[0]] = 0.0
    fin5.cash_flow[0].values[periods[0]] = 0.0
    series5 = compute_earnings_quality_series(
        fin5, periods, compute_anchor(fin5, periods)
    )
    assert series5.cash_conversion_ratio[0] == UNDEFINED_RATIO
    assert series5.total_accruals[0] == 0.0


def test_earnings_quality_without_assets_and_missing_cfo():
    fin = _tiny_fin(with_assets=False)
    periods = canonical_fiscal_periods(fin)
    avail = earnings_quality_availability(fin)
    assert avail.operating_cash_flow is True
    assert avail.total_assets is False
    series = compute_earnings_quality_series(fin, periods, compute_anchor(fin, periods))
    assert all(v is None for v in series.average_total_assets)
    assert all(v is None for v in series.accrual_ratio)
    assert series.total_accruals[0] == pytest.approx(20.0)

    fin_no = _tiny_fin(with_cfo=False)
    assert earnings_quality_availability(fin_no).operating_cash_flow is False
    with pytest.raises(MissingLineError):
        compute_earnings_quality_series(
            fin_no, canonical_fiscal_periods(fin_no), compute_anchor(fin_no, canonical_fiscal_periods(fin_no))
        )

    with pytest.raises(AmbiguousLineError):
        earnings_quality_availability(_tiny_fin(ambiguous_cfo=True))


def test_zero_average_assets_accrual_ratio_convention():
    fin = _tiny_fin()
    periods = canonical_fiscal_periods(fin)
    anchor = compute_anchor(fin, periods)
    for item in fin.balance_sheet:
        if item.label == "Total assets":
            item.values[periods[0]] = 0.0
            item.values[periods[1]] = 0.0
    series = compute_earnings_quality_series(fin, periods, anchor)
    assert series.average_total_assets[1] == 0.0
    assert series.accrual_ratio[1] == UNDEFINED_RATIO

    # Zero accruals with nonzero average assets -> numeric 0.0
    fin_num = _tiny_fin()
    periods = canonical_fiscal_periods(fin_num)
    fin_num.income_statement[-1].values[periods[1]] = 90.0
    fin_num.cash_flow[0].values[periods[1]] = 90.0
    series_num = compute_earnings_quality_series(
        fin_num, periods, compute_anchor(fin_num, periods)
    )
    assert series_num.total_accruals[1] == pytest.approx(0.0)
    assert series_num.accrual_ratio[1] == pytest.approx(0.0)


def test_incomplete_or_missing_net_income_fails_without_zero_fabrication():
    fin = _tiny_fin()
    periods = canonical_fiscal_periods(fin)
    anchor = compute_anchor(fin, periods)
    ni_item = fin.income_statement[-1]
    del ni_item.values[periods[1]]
    with pytest.raises(ValueError, match="net_income") as exc_info:
        compute_earnings_quality_series(fin, periods, anchor)
    assert periods[1].isoformat() in str(exc_info.value)

    fin_none = _tiny_fin()
    periods = canonical_fiscal_periods(fin_none)
    anchor_none = compute_anchor(fin_none, periods)
    fin_none.income_statement[-1].values[periods[0]] = None
    with pytest.raises(ValueError, match="net_income") as exc_info:
        compute_earnings_quality_series(fin_none, periods, anchor_none)
    assert periods[0].isoformat() in str(exc_info.value)

    fin_zero = _tiny_fin()
    periods = canonical_fiscal_periods(fin_zero)
    fin_zero.income_statement[-1].values[periods[0]] = 0.0
    series = compute_earnings_quality_series(
        fin_zero, periods, compute_anchor(fin_zero, periods)
    )
    assert series.cash_conversion_ratio[0] == UNDEFINED_RATIO

    fin_mismatch = _tiny_fin()
    periods = canonical_fiscal_periods(fin_mismatch)
    anchor_mismatch = compute_anchor(fin_mismatch, periods)
    fin_mismatch.income_statement[-1].values[periods[0]] = 999.0
    with pytest.raises(ValueError, match="does not match AnchorMetrics"):
        compute_earnings_quality_series(fin_mismatch, periods, anchor_mismatch)

def test_incomplete_or_missing_cfo_period_fails_without_zero_fabrication():
    fin = _tiny_fin()
    periods = canonical_fiscal_periods(fin)
    anchor = compute_anchor(fin, periods)
    missing_period = periods[1]
    del fin.cash_flow[0].values[missing_period]
    with pytest.raises(ValueError, match="operating_cash_flow") as exc_info:
        compute_earnings_quality_series(fin, periods, anchor)
    assert missing_period.isoformat() in str(exc_info.value)

    fin_none = _tiny_fin()
    fin_none.cash_flow[0].values[periods[0]] = None
    with pytest.raises(ValueError, match="operating_cash_flow") as exc_info:
        compute_earnings_quality_series(
            fin_none, periods, compute_anchor(fin_none, periods)
        )
    assert periods[0].isoformat() in str(exc_info.value)

    fin_zero = _tiny_fin()
    fin_zero.cash_flow[0].values[periods[0]] = 0.0
    series = compute_earnings_quality_series(
        fin_zero, periods, compute_anchor(fin_zero, periods)
    )
    assert series.operating_cash_flow[0] == 0.0
    assert series.total_accruals[0] == pytest.approx(100.0)


def test_incomplete_total_assets_fails_without_zero_fabrication():
    fin = _tiny_fin()
    periods = canonical_fiscal_periods(fin)
    # Mutate after anchor so reformulation integrity still passes (same pattern as
    # zero-average-assets). Completeness is enforced inside the quality series.
    anchor = compute_anchor(fin, periods)
    assets = next(item for item in fin.balance_sheet if item.label == "Total assets")
    del assets.values[periods[0]]
    with pytest.raises(ValueError, match="total_assets") as exc_info:
        compute_earnings_quality_series(fin, periods, anchor)
    assert periods[0].isoformat() in str(exc_info.value)

    fin_later = _tiny_fin()
    periods = canonical_fiscal_periods(fin_later)
    anchor_later = compute_anchor(fin_later, periods)
    assets_later = next(
        item for item in fin_later.balance_sheet if item.label == "Total assets"
    )
    assets_later.values[periods[1]] = None
    with pytest.raises(ValueError, match="total_assets") as exc_info:
        compute_earnings_quality_series(fin_later, periods, anchor_later)
    assert periods[1].isoformat() in str(exc_info.value)

    fin_zero = _tiny_fin()
    periods = canonical_fiscal_periods(fin_zero)
    anchor_zero = compute_anchor(fin_zero, periods)
    assets_zero = next(
        item for item in fin_zero.balance_sheet if item.label == "Total assets"
    )
    assets_zero.values[periods[0]] = 0.0
    assets_zero.values[periods[1]] = 0.0
    series = compute_earnings_quality_series(fin_zero, periods, anchor_zero)
    assert series.average_total_assets[1] == 0.0
    assert series.accrual_ratio[1] == UNDEFINED_RATIO


def test_availability_and_incompleteness_are_distinct(tmp_path):
    # CFO absent, Total Assets present -> omit module
    fin_no_cfo = _tiny_fin(with_cfo=False)
    builder = ReferenceModelBuilder(fin_no_cfo)
    assert builder.quality_series is None
    assert builder.quality_specs == ()
    answer = tmp_path / "AbsentCFO_Answer_Key.xlsx"
    builder.build(answer)
    wb = load_workbook(answer)
    assert EARNINGS_QUALITY_SHEET not in wb.sheetnames
    wb.close()

    # CFO complete, Total Assets absent -> core families only
    fin_no_assets = _tiny_fin(with_assets=False)
    builder = ReferenceModelBuilder(fin_no_assets)
    assert builder.quality_series is not None
    assert {s.family_id for s in builder.quality_specs} == {
        "operating_cash_flow_link",
        "cash_conversion_ratio",
        "total_accruals",
    }
    answer = tmp_path / "AbsentAssets_Answer_Key.xlsx"
    builder.build(answer)
    wb = load_workbook(answer)
    assert EARNINGS_QUALITY_SHEET in wb.sheetnames
    wb.close()

    # CFO present but incomplete -> build raises (do not omit module)
    fin_incomplete_cfo = _tiny_fin()
    periods = canonical_fiscal_periods(fin_incomplete_cfo)
    del fin_incomplete_cfo.cash_flow[0].values[periods[1]]
    with pytest.raises(ValueError, match="operating_cash_flow"):
        ReferenceModelBuilder(fin_incomplete_cfo)
    with pytest.raises(ValueError, match="operating_cash_flow"):
        build_training_workbook(
            fin_incomplete_cfo, tmp_path / "IncompleteCFO_Trainer.xlsx"
        )

    # Total Assets present but incomplete -> quality series fails closed (do not
    # omit extension / invent zero). Mutate after a successful builder init so
    # reformulation is not the confounding failure mode.
    fin_incomplete_assets = _tiny_fin()
    periods = canonical_fiscal_periods(fin_incomplete_assets)
    builder = ReferenceModelBuilder(fin_incomplete_assets)
    assert builder.quality_series is not None
    assets = next(
        item
        for item in fin_incomplete_assets.balance_sheet
        if item.label == "Total assets"
    )
    del assets.values[periods[0]]
    with pytest.raises(ValueError, match="total_assets"):
        compute_earnings_quality_series(
            fin_incomplete_assets,
            periods,
            builder.anchor,
        )
    # Fresh build from incomplete Total Assets also fails closed (reformulation
    # integrity rejects the malformed source before/alongside quality).
    fin_build_incomplete = _tiny_fin()
    assets_build = next(
        item
        for item in fin_build_incomplete.balance_sheet
        if item.label == "Total assets"
    )
    del assets_build.values[periods[0]]
    with pytest.raises(ValueError):
        ReferenceModelBuilder(fin_build_incomplete)
    with pytest.raises(ValueError):
        build_training_workbook(
            fin_build_incomplete, tmp_path / "IncompleteAssets_Trainer.xlsx"
        )

    # Both complete -> all five quality families
    fin_complete = _tiny_fin()
    builder = ReferenceModelBuilder(fin_complete)
    assert builder.quality_series is not None
    assert {s.family_id for s in builder.quality_specs} == {
        "operating_cash_flow_link",
        "cash_conversion_ratio",
        "total_accruals",
        "average_total_assets",
        "accrual_ratio",
    }


def test_expand_quality_specs_counts_and_gating():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(QUALITY_COMPONENT_CATALOG) == 8
    assert [f.order for f in QUALITY_COMPONENT_CATALOG] == [
        30,
        31,
        32,
        33,
        34,
        126,
        127,
        128,
    ]
    full = expand_quality_specs(periods, start_order=119, include_asset_scaled=True)
    assert len(full) == 23
    assert full[0].order == 119
    assert full[-1].order == 141
    core = expand_quality_specs(periods, start_order=119, include_asset_scaled=False)
    assert len(core) == 15
    assert {s.family_id for s in core} == {
        "operating_cash_flow_link",
        "cash_conversion_ratio",
        "total_accruals",
    }
    sbc = expand_quality_specs(
        periods, start_order=119, include_asset_scaled=True, include_sbc=True
    )
    assert len(sbc) == 38
    assert {s.family_id for s in sbc} >= {
        "sbc_to_revenue",
        "sbc_to_operating_cash_flow",
        "operating_cash_flow_less_sbc",
    }


def test_demo_quality_surface_base_and_normalization(tmp_path):
    import json

    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "BASE_Trainer.xlsx")
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == 312
    assert len(group_components_by_family(smap)) == 74
    wb = load_workbook(answer)
    assert EARNINGS_QUALITY_SHEET in wb.sheetnames
    wb.close()
    summary = check_workbook(trainer)
    assert summary.total == 312
    assert summary.blank == 312

    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        data, tmp_path / "NORM_Trainer.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(smap_n.all_ordered()) == 332
    assert len(group_components_by_family(smap_n)) == 78
    summary_n = check_workbook(trainer_n)
    assert summary_n.total == 332
    assert summary_n.blank == 332


def test_cfo_unavailable_keeps_prior_surface(tmp_path):
    fin = _tiny_fin(with_cfo=False)
    builder = ReferenceModelBuilder(fin)
    assert builder.quality_series is None
    assert builder.quality_specs == ()
    answer = tmp_path / "NoCFO_Answer_Key.xlsx"
    trainer = tmp_path / "NoCFO_Trainer.xlsx"
    smap = builder.build(answer)
    from core.trainer.workbook import TrainingWorkbookGenerator

    TrainingWorkbookGenerator(answer, smap).generate(trainer)
    loaded = load_semantic_map(answer)
    # 2 periods: historical expandable cells only (no quality)
    assert all(c.category != "earnings_quality" for c in loaded.all_ordered())
    wb = load_workbook(answer)
    assert EARNINGS_QUALITY_SHEET not in wb.sheetnames
    wb.close()


def test_cfo_without_assets_omits_scaled_families(tmp_path):
    fin = _tiny_fin(with_assets=False)
    builder = ReferenceModelBuilder(fin)
    assert builder.quality_availability.operating_cash_flow is True
    assert builder.quality_availability.total_assets is False
    family_ids = {s.family_id for s in builder.quality_specs}
    assert family_ids == {
        "operating_cash_flow_link",
        "cash_conversion_ratio",
        "total_accruals",
    }
    answer = tmp_path / "NoAssets_Answer_Key.xlsx"
    smap = builder.build(answer)
    assert "average_total_assets" not in {c.family_id for c in smap.all_ordered()}


def test_quality_exact_and_equivalent_check(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "Q_Trainer.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "cash_conversion_ratio" and c.period_end == "2025-12-31"
    )
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    summary = check_workbook(trainer)
    assert summary.correct == 1

    # Equivalent cached value
    from core.tests.test_normalization import _inject_formula_and_cached_value

    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula=f"={float(comp.expected_value)}",
        cached_value=float(comp.expected_value),
    )
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer,
        comp.tab,
        comp.cell,
        formula="=0",
        cached_value=0.0,
    )
    assert check_workbook(trainer).incorrect == 1


def test_quality_populated_row_tamper_fails_closed(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "T_Trainer.xlsx")
    smap = load_semantic_map(answer)
    comp = next(c for c in smap.all_ordered() if c.family_id == "total_accruals")
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    # Tamper populated Reported Net Income row (row 6).
    wb[EARNINGS_QUALITY_SHEET].cell(6, 2).value = 1
    wb.save(trainer)
    wb.close()
    with pytest.raises(ValueError, match="Trusted workbook cell was modified: Earnings Quality"):
        check_workbook(trainer)
    wb = load_workbook(trainer, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

    trainer2, _ = build_training_workbook(data, tmp_path / "T2_Trainer.xlsx")
    from core.trainer.semantic_io import answer_key_path_for

    smap2 = load_semantic_map(answer_key_path_for(trainer2))
    comp2 = next(c for c in smap2.all_ordered() if c.family_id == "accrual_ratio")
    wb = load_workbook(trainer2, data_only=False)
    row, col = parse_cell_ref(comp2.cell)
    wb[comp2.tab].cell(row=row, column=col).value = comp2.formula
    wb[EARNINGS_QUALITY_SHEET].cell(10, 2).value = 1  # Total Assets populated link
    wb.save(trainer2)
    wb.close()
    with pytest.raises(ValueError, match="Trusted workbook cell was modified: Earnings Quality"):
        check_workbook(trainer2)


def test_quality_composes_with_normalization_and_classification(tmp_path):
    import json

    data = _ingest_demo()
    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer, answer = build_training_workbook(
        data, tmp_path / "Compose_Trainer.xlsx", assumptions
    )
    smap = load_semantic_map(answer)
    net_debt = max(
        (c for c in smap.all_ordered() if c.family_id == "net_debt"),
        key=lambda c: c.period_index or 0,
    )
    norm = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "pretax_normalization_adjustment"
        and c.period_end == "2023-12-31"
    )
    quality = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "total_accruals" and c.period_end == "2025-12-31"
    )
    wb = load_workbook(trainer, data_only=False)
    wb["Accounting Judgment"].cell(5, 6).value = "Financial Liability"
    wb["Normalization Judgment"].cell(5, 6).value = "Recurring"
    for comp in (net_debt, norm, quality):
        r, c = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=r, column=c).value = comp.formula
    wb.save(trainer)
    wb.close()
    summary = check_workbook(trainer)
    assert summary.correct == 3
    assert summary.incorrect == 0


def test_undefined_ratio_formulas_use_na_and_check_accepts_na(tmp_path):
    from core.tests.test_normalization import _inject_formula_and_cached_value

    # Zero Net Income -> conversion formula uses NA(); expected is #N/A
    fin_ni = _tiny_fin()
    periods = canonical_fiscal_periods(fin_ni)
    fin_ni.income_statement[-1].values[periods[0]] = 0.0
    trainer, answer = build_training_workbook(fin_ni, tmp_path / "ZeroNI_Trainer.xlsx")
    smap = load_semantic_map(answer)
    conv = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "cash_conversion_ratio" and c.period_index == 0
    )
    assert conv.expected_value == UNDEFINED_RATIO
    assert "NA()" in conv.formula
    assert ",0," not in conv.formula.replace("NA()", "")
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(conv.cell)
    assert wb[conv.tab].cell(row=row, column=col).value is None
    assert _fill_rgb(wb[conv.tab].cell(row=row, column=col)) == "FFFF00"
    wb[conv.tab].cell(row=row, column=col).value = conv.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer,
        conv.tab,
        conv.cell,
        formula="=NA()",
        cached_value=UNDEFINED_RATIO,
    )
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer,
        conv.tab,
        conv.cell,
        formula="=0",
        cached_value=0.0,
    )
    assert check_workbook(trainer).incorrect == 1

    # Zero average assets (post-anchor) -> undefined accrual ratio in the series
    fin_assets = _tiny_fin()
    periods = canonical_fiscal_periods(fin_assets)
    anchor = compute_anchor(fin_assets, periods)
    for item in fin_assets.balance_sheet:
        if item.label == "Total assets":
            item.values[periods[0]] = 0.0
            item.values[periods[1]] = 0.0
    series = compute_earnings_quality_series(fin_assets, periods, anchor)
    assert series.accrual_ratio[1] == UNDEFINED_RATIO

    # Demo Answer Key ratio formulas use NA() denominator guards; Trainer stays blank
    data = _ingest_demo()
    demo_trainer, demo_answer = build_training_workbook(
        data, tmp_path / "DemoNA_Trainer.xlsx"
    )
    demo_smap = load_semantic_map(demo_answer)
    accrual = next(c for c in demo_smap.all_ordered() if c.family_id == "accrual_ratio")
    conversion = next(
        c for c in demo_smap.all_ordered() if c.family_id == "cash_conversion_ratio"
    )
    assert "NA()" in accrual.formula
    assert ",0," not in accrual.formula.replace("NA()", "")
    assert "NA()" in conversion.formula
    trainer_wb = load_workbook(demo_trainer, data_only=False)
    r, c = parse_cell_ref(accrual.cell)
    assert trainer_wb[accrual.tab].cell(row=r, column=c).value is None
    trainer_wb.close()


def _eq_row_by_label(ws, label: str) -> int:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row, 1).value == label:
            return row
    raise AssertionError(f"missing Earnings Quality row {label!r}")


def _add_sbc(
    fin: StandardizedFinancials,
    *,
    values=(10.0, 12.0),
    label="Stock-based compensation expense",
    concept="stock_based_compensation",
) -> LineItem:
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
    item = _li(label, {d1: values[0], d2: values[1]}, concept=concept)
    fin.cash_flow.append(item)
    return item


def test_sbc_explicit_concept_excludes_settlement_and_labels():
    fin = _tiny_fin()
    assert not sbc_diagnostics_applicable(fin)
    assert resolve_sbc_source(fin) is None
    avail = earnings_quality_availability(fin)
    assert avail.stock_based_compensation is False
    assert avail.stock_based_compensation_ambiguous is False

    label_only = _tiny_fin()
    _add_sbc(label_only, concept="")
    assert resolve_sbc_source(label_only) is None
    assert not sbc_diagnostics_applicable(label_only)
    series_label = compute_earnings_quality_series(
        label_only,
        list(canonical_fiscal_periods(label_only)),
        compute_anchor(label_only, list(canonical_fiscal_periods(label_only))),
    )
    assert series_label.operating_cash_flow_less_sbc is None
    assert ReferenceModelBuilder(label_only).quality_series is not None
    assert "sbc_to_revenue" not in {
        s.family_id for s in ReferenceModelBuilder(label_only).quality_specs
    }

    misleading = _tiny_fin()
    _add_sbc(
        misleading,
        label="Proceeds from settlement of stock-based compensation",
        concept="proceeds_from_stock_based_compensation",
        values=(5.0, 6.0),
    )
    _add_sbc(
        misleading,
        label="Shares withheld related to net share settlement of stock-based compensation",
        concept="shares_withheld_for_stock_based_compensation",
        values=(-2.0, -3.0),
    )
    _add_sbc(
        misleading,
        label="Repurchase of common stock",
        concept="repurchase_of_common_stock",
        values=(-20.0, -25.0),
    )
    assert resolve_sbc_source(misleading) is None
    assert not sbc_diagnostics_applicable(misleading)

    explicit = _tiny_fin()
    stored = _add_sbc(
        explicit,
        label="Weird SBC add-back label",
        concept="stock_based_compensation",
        values=(10.0, 12.0),
    )
    _add_sbc(
        explicit,
        label="Stock-based compensation expense",
        concept="",
        values=(99.0, 99.0),
    )
    _add_sbc(
        explicit,
        label="Proceeds from settlement of stock-based compensation",
        concept="proceeds_from_stock_based_compensation",
        values=(5.0, 6.0),
    )
    item = resolve_sbc_source(explicit)
    assert item is stored
    assert item.concept == "stock_based_compensation"
    assert item.label == "Weird SBC add-back label"
    resolved = resolve_line(
        explicit.cash_flow, "stock_based_compensation", required=True
    )
    assert resolved.item is stored
    assert sbc_diagnostics_applicable(explicit) is True


def test_sbc_ambiguous_and_absent_cfo_omit_only_extension():
    dup = _tiny_fin()
    _add_sbc(dup, label="SBC A", values=(10.0, 12.0))
    _add_sbc(dup, label="SBC B", values=(11.0, 13.0))
    avail = earnings_quality_availability(dup)
    assert avail.operating_cash_flow is True
    assert avail.stock_based_compensation is False
    assert avail.stock_based_compensation_ambiguous is True
    assert not sbc_diagnostics_applicable(dup)
    assert resolve_sbc_source(dup) is None
    series = compute_earnings_quality_series(
        dup,
        list(canonical_fiscal_periods(dup)),
        compute_anchor(dup, list(canonical_fiscal_periods(dup))),
    )
    assert series.operating_cash_flow == (80.0, 90.0)
    assert series.operating_cash_flow_less_sbc is None
    mapped = earnings_quality_expected_series(series)
    assert "operating_cash_flow_less_sbc" not in mapped
    builder = ReferenceModelBuilder(dup)
    assert {s.family_id for s in builder.quality_specs} == {
        "operating_cash_flow_link",
        "cash_conversion_ratio",
        "total_accruals",
        "average_total_assets",
        "accrual_ratio",
    }

    absent_cfo = _tiny_fin(with_cfo=False)
    _add_sbc(absent_cfo)
    assert not sbc_diagnostics_applicable(absent_cfo)
    assert ReferenceModelBuilder(absent_cfo).quality_series is None

    amb_cfo = _tiny_fin(ambiguous_cfo=True)
    _add_sbc(amb_cfo)
    with pytest.raises(AmbiguousLineError):
        earnings_quality_availability(amb_cfo)
    assert not sbc_diagnostics_applicable(amb_cfo)


def test_sbc_series_signs_zero_undefined_missing_and_immutability():
    fin = _tiny_fin()
    _add_sbc(fin, values=(10.0, 12.0))
    periods = list(canonical_fiscal_periods(fin))
    series = compute_earnings_quality_series(fin, periods, compute_anchor(fin, periods))
    assert series.stock_based_compensation == (10.0, 12.0)
    assert series.sbc_to_revenue[0] == pytest.approx(10.0 / 1000.0)
    assert series.sbc_to_operating_cash_flow[0] == pytest.approx(10.0 / 80.0)
    assert series.operating_cash_flow_less_sbc == (70.0, 78.0)
    mapped = earnings_quality_expected_series(series)
    assert set(mapped) >= {
        "sbc_to_revenue",
        "sbc_to_operating_cash_flow",
        "operating_cash_flow_less_sbc",
    }

    zero_sbc = _tiny_fin()
    _add_sbc(zero_sbc, values=(0.0, 12.0))
    series_z = compute_earnings_quality_series(
        zero_sbc,
        list(canonical_fiscal_periods(zero_sbc)),
        compute_anchor(zero_sbc, list(canonical_fiscal_periods(zero_sbc))),
    )
    assert series_z.stock_based_compensation[0] == 0.0
    assert series_z.sbc_to_revenue[0] == pytest.approx(0.0)
    assert series_z.operating_cash_flow_less_sbc[0] == pytest.approx(80.0)

    neg = _tiny_fin()
    _add_sbc(neg, values=(-10.0, 12.0))
    series_n = compute_earnings_quality_series(
        neg,
        list(canonical_fiscal_periods(neg)),
        compute_anchor(neg, list(canonical_fiscal_periods(neg))),
    )
    assert series_n.stock_based_compensation[0] == -10.0
    assert series_n.operating_cash_flow_less_sbc[0] == pytest.approx(90.0)

    zero_cfo = _tiny_fin()
    zero_cfo.cash_flow[0].values[periods[0]] = 0.0
    _add_sbc(zero_cfo, values=(10.0, 12.0))
    series_c = compute_earnings_quality_series(
        zero_cfo, periods, compute_anchor(zero_cfo, periods)
    )
    assert series_c.sbc_to_operating_cash_flow[0] == UNDEFINED_RATIO
    assert series_c.operating_cash_flow_less_sbc[0] == pytest.approx(-10.0)

    zero_rev = _tiny_fin()
    zero_rev.income_statement[0].values[periods[0]] = 0.0
    _add_sbc(zero_rev, values=(10.0, 12.0))
    series_r = compute_earnings_quality_series(
        zero_rev, periods, compute_anchor(zero_rev, periods)
    )
    assert series_r.sbc_to_revenue[0] == UNDEFINED_RATIO

    missing = _tiny_fin()
    item = _add_sbc(missing, values=(10.0, 12.0))
    del item.values[periods[1]]
    with pytest.raises(MissingHistoricalValueError, match="stock_based_compensation"):
        compute_earnings_quality_series(
            missing, periods, compute_anchor(missing, periods)
        )

    none_period = _tiny_fin()
    none_item = _add_sbc(none_period, values=(10.0, 12.0))
    none_item.values[periods[0]] = None
    with pytest.raises(MissingHistoricalValueError, match="stock_based_compensation"):
        compute_earnings_quality_series(
            none_period, periods, compute_anchor(none_period, periods)
        )

    immutable = _tiny_fin()
    stored = _add_sbc(immutable, values=(10.0, 12.0))
    before = copy.deepcopy(stored.values)
    compute_earnings_quality_series(
        immutable, periods, compute_anchor(immutable, periods)
    )
    assert stored.values == before
    assert stored.concept == "stock_based_compensation"


def test_sbc_workbook_formulas_notes_and_check(tmp_path):
    from core.tests.test_normalization import _inject_formula_and_cached_value

    fin = _tiny_fin()
    stored = _add_sbc(fin, values=(10.0, 12.0))
    builder = ReferenceModelBuilder(fin)
    assert builder.quality_series is not None
    assert builder.quality_series.operating_cash_flow_less_sbc == (70.0, 78.0)
    assert len([s for s in builder.quality_specs if s.family_id in {
        "sbc_to_revenue",
        "sbc_to_operating_cash_flow",
        "operating_cash_flow_less_sbc",
    }]) == 6

    trainer, answer = build_training_workbook(fin, tmp_path / "SBC_Trainer.xlsx")
    smap = load_semantic_map(answer)
    sbc_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id
        in {
            "sbc_to_revenue",
            "sbc_to_operating_cash_flow",
            "operating_cash_flow_less_sbc",
        }
    ]
    assert len(sbc_comps) == 6
    awb = load_workbook(answer, data_only=False)
    twb = load_workbook(trainer, data_only=False)
    ws_a = awb[EARNINGS_QUALITY_SHEET]
    reported = _eq_row_by_label(ws_a, "Stock-based compensation (reported)")
    less_row = _eq_row_by_label(ws_a, "Reported CFO less SBC add-back")
    reported_f = str(ws_a.cell(reported, 2).value).replace(" ", "")
    assert reported_f.startswith("='CashFlowStatement'!")
    less_f = str(ws_a.cell(less_row, 2).value).replace(" ", "")
    assert less_f.startswith("=")
    assert "-" in less_f
    for comp in sbc_comps:
        row, col = parse_cell_ref(comp.cell)
        tcell = twb[comp.tab].cell(row=row, column=col)
        acell = awb[comp.tab].cell(row=row, column=col)
        assert tcell.value is None
        assert tcell.comment is None
        assert _fill_rgb(tcell) == "FFFF00"
        assert acell.value == comp.formula
        note = (acell.comment.text or "") if acell.comment is not None else ""
        assert note.strip()
        assert "does not restate reported CFO" in note
        assert "cash compensation" in note.lower() or "dilution" in note.lower()
        assert "free cash flow" in note.lower()
        assert "tax" in note.lower()
    reported_t = twb[EARNINGS_QUALITY_SHEET].cell(reported, 2)
    assert reported_t.value is not None
    assert reported_t.comment is None
    twb.close()
    awb.close()
    assert stored.concept == "stock_based_compensation"
    assert stored.values[date(2024, 12, 31)] == 10.0

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0

    bad = next(c for c in sbc_comps if c.family_id == "operating_cash_flow_less_sbc")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    dumped = repr(check_workbook(trainer))
    assert "=999" not in dumped
    assert "Reported CFO less SBC add-back" not in dumped

"""Step 9E.1 — historical cash-conversion / accrual trend diagnostics."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    QUALITY_CHANGE_COMPONENT_CATALOG,
    QUALITY_COMPONENT_CATALOG,
    expand_quality_change_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.earnings_quality import EarningsQualitySeries
from core.model.earnings_quality_change import compute_earnings_quality_change_series
from core.model.historical_expected import (
    earnings_quality_change_expected_series,
    expected_value_for_component,
)
from core.model.ratio_values import UNDEFINED_RATIO
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family
from core.tests.test_normalization import _inject_formula_and_cached_value

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    adapter = HKManualDocumentAdapter()
    return adapter.ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values_by_date, concept=""):
    return LineItem(label=label, values=values_by_date, concept=concept)


def _eq_series(
    *,
    cfo,
    conversion,
    accruals,
    accrual_ratio,
    average_total_assets=None,
):
    n = len(cfo)
    avg = average_total_assets
    if avg is None:
        avg = tuple(None for _ in range(n))
    return EarningsQualitySeries(
        operating_cash_flow=tuple(cfo),
        cash_conversion_ratio=tuple(conversion),
        total_accruals=tuple(accruals),
        average_total_assets=tuple(avg),
        accrual_ratio=tuple(accrual_ratio),
    )


def test_ordinary_quality_change_series():
    series = compute_earnings_quality_change_series(
        _eq_series(
            cfo=(80.0, 90.0, 120.0),
            conversion=(0.80, 0.75, 1.00),
            accruals=(20.0, 30.0, 0.0),
            accrual_ratio=(None, 0.10, 0.04),
            average_total_assets=(None, 300.0, 310.0),
        )
    )
    assert series.operating_cash_flow_change[0] is None
    assert series.cash_conversion_ratio_change[0] is None
    assert series.total_accruals_change[0] is None
    assert series.accrual_ratio_change[0] is None
    assert series.accrual_ratio_change[1] is None

    assert series.operating_cash_flow_change[1] == pytest.approx(10.0)
    assert series.operating_cash_flow_change[2] == pytest.approx(30.0)
    assert series.cash_conversion_ratio_change[1] == pytest.approx(-0.05)
    assert series.cash_conversion_ratio_change[2] == pytest.approx(0.25)
    assert series.total_accruals_change[1] == pytest.approx(10.0)
    assert series.total_accruals_change[2] == pytest.approx(-30.0)
    assert series.accrual_ratio_change[2] == pytest.approx(-0.06)


def test_quality_change_edge_cases():
    # Unchanged -> 0.0; negatives retained
    s0 = compute_earnings_quality_change_series(
        _eq_series(
            cfo=(100.0, 100.0, 80.0),
            conversion=(1.0, 1.0, 0.8),
            accruals=(10.0, 10.0, -5.0),
            accrual_ratio=(None, 0.05, 0.05),
        )
    )
    assert s0.operating_cash_flow_change[1] == pytest.approx(0.0)
    assert s0.operating_cash_flow_change[2] == pytest.approx(-20.0)
    assert s0.total_accruals_change[2] == pytest.approx(-15.0)
    assert s0.accrual_ratio_change[2] == pytest.approx(0.0)

    # Conversion #N/A propagation
    sna = compute_earnings_quality_change_series(
        _eq_series(
            cfo=(80.0, 90.0, 100.0),
            conversion=(0.8, UNDEFINED_RATIO, 1.0),
            accruals=(20.0, 10.0, 0.0),
            accrual_ratio=(None, None, None),
        )
    )
    assert sna.cash_conversion_ratio_change[1] == UNDEFINED_RATIO
    assert sna.cash_conversion_ratio_change[2] == UNDEFINED_RATIO
    assert all(v is None for v in sna.accrual_ratio_change)

    # Accrual-ratio #N/A
    sar = compute_earnings_quality_change_series(
        _eq_series(
            cfo=(80.0, 90.0, 100.0),
            conversion=(0.8, 0.9, 1.0),
            accruals=(20.0, 10.0, 0.0),
            accrual_ratio=(None, 0.1, UNDEFINED_RATIO),
        )
    )
    assert sar.accrual_ratio_change[2] == UNDEFINED_RATIO

    # Inconsistent partial accrual-ratio row
    with pytest.raises(ValueError, match="consecutive comparable"):
        compute_earnings_quality_change_series(
            _eq_series(
                cfo=(80.0, 90.0, 100.0),
                conversion=(0.8, 0.9, 1.0),
                accruals=(20.0, 10.0, 0.0),
                accrual_ratio=(None, 0.1, None),
            )
        )

    # Length mismatch
    bad = EarningsQualitySeries(
        operating_cash_flow=(80.0, 90.0),
        cash_conversion_ratio=(0.8, 0.9, 1.0),
        total_accruals=(20.0, 10.0, 0.0),
        average_total_assets=(None, 100.0, 110.0),
        accrual_ratio=(None, 0.1, 0.05),
    )
    with pytest.raises(ValueError, match="length mismatch"):
        compute_earnings_quality_change_series(bad)


def test_catalog_expand_and_expected_keys():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(QUALITY_CHANGE_COMPONENT_CATALOG) == 4
    assert len({f.id for f in QUALITY_CHANGE_COMPONENT_CATALOG}) == 4
    assert [f.order for f in QUALITY_CHANGE_COMPONENT_CATALOG] == [63, 64, 65, 66]
    assert len(QUALITY_COMPONENT_CATALOG) == 5

    full = expand_quality_change_specs(periods, start_order=245, include_asset_scaled=True)
    assert len(full) == 15
    assert any(s.family_id == "accrual_ratio_change" for s in full)

    core = expand_quality_change_specs(
        periods, start_order=245, include_asset_scaled=False
    )
    assert len(core) == 12
    assert all(s.family_id != "accrual_ratio_change" for s in core)

    data = _ingest_demo()
    builder = ReferenceModelBuilder(data)
    assert builder.quality_change_series is not None
    assert len(builder.quality_change_specs) == 15
    series = earnings_quality_change_expected_series(builder.quality_series)
    assert set(series) == {f.id for f in QUALITY_CHANGE_COMPONENT_CATALOG}


def test_quality_change_gating():
    from core.tests.test_earnings_quality import _tiny_fin

    fin_no_cfo = _tiny_fin(with_cfo=False)
    builder = ReferenceModelBuilder(fin_no_cfo)
    assert builder.quality_series is None
    assert builder.quality_change_series is None
    assert builder.quality_change_specs == ()

    fin_no_assets = _tiny_fin(with_assets=False)
    builder2 = ReferenceModelBuilder(fin_no_assets)
    assert builder2.quality_series is not None
    assert builder2.quality_change_series is not None
    fams = {s.family_id for s in builder2.quality_change_specs}
    assert fams == {
        "operating_cash_flow_change",
        "cash_conversion_ratio_change",
        "total_accruals_change",
    }


def test_demo_quality_change_surface(tmp_path):
    import json

    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "BASE_QC.xlsx")
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == 259
    assert len(group_components_by_family(smap)) == 62
    qc = [c for c in smap.all_ordered() if c.category == "earnings_quality_change"]
    assert len(qc) == 15
    assert all(c.tab == "Earnings Quality" for c in qc)

    wb = load_workbook(answer)
    ws = wb["Earnings Quality"]
    assert ws.cell(5, 1).value == "Operating Cash Flow"
    assert ws.cell(12, 1).value == "Accrual Ratio"
    assert ws.cell(14, 1).value == "EARNINGS QUALITY TREND DIAGNOSTICS"
    assert ws.cell(15, 1).value == "Change in Operating Cash Flow"
    assert ws.cell(18, 1).value == "Change in Accrual Ratio"
    assert ws.cell(19, 1).value == "EARNINGS QUALITY CHANGE CHECK"
    assert ws.cell(15, 2).value == "N/A"
    assert ws.cell(18, 2).value == "N/A"
    assert ws.cell(18, 3).value == "N/A"
    practice = {(c.tab, c.cell) for c in smap.all_ordered()}
    assert ("Earnings Quality", "C19") not in practice

    cfo_chg = next(
        c for c in qc if c.family_id == "operating_cash_flow_change" and c.period_index == 1
    )
    assert "IFERROR" not in cfo_chg.formula.upper()
    assert "ABS" not in cfo_chg.formula.upper()

    wb_t = load_workbook(trainer)
    for comp in qc:
        row, col = parse_cell_ref(comp.cell)
        cell = wb_t[comp.tab].cell(row=row, column=col)
        assert cell.value is None
        assert _fill_rgb(cell) == "FFFF00"
    wb.close()
    wb_t.close()

    summary = check_workbook(trainer)
    assert summary.total == 259
    assert summary.blank == 259

    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        data, tmp_path / "NORM_QC.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(smap_n.all_ordered()) == 279
    assert len(group_components_by_family(smap_n)) == 66
    assert check_workbook(trainer_n).blank == 279


def test_no_assets_trend_surface(tmp_path):
    from core.tests.test_earnings_quality import _tiny_fin

    fin = _tiny_fin(with_assets=False)
    # Tiny fixture has 2 periods -> 3 comparable trend cells (no accrual-ratio change)
    trainer, answer = build_training_workbook(fin, tmp_path / "NO_ASSETS_QC.xlsx")
    smap = load_semantic_map(answer)
    qc = [c for c in smap.all_ordered() if c.category == "earnings_quality_change"]
    assert len(qc) == 3
    assert {c.family_id for c in qc} == {
        "operating_cash_flow_change",
        "cash_conversion_ratio_change",
        "total_accruals_change",
    }
    wb = load_workbook(answer)
    ws = wb["Earnings Quality"]
    assert ws.cell(14, 1).value == "EARNINGS QUALITY TREND DIAGNOSTICS"
    for col in range(2, 4):
        assert ws.cell(18, col).value == "N/A"
    wb.close()

    # Five-period expand without assets -> 12 specs
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    specs = expand_quality_change_specs(
        periods, start_order=1, include_asset_scaled=False
    )
    assert len(specs) == 12
    assert all(s.family_id != "accrual_ratio_change" for s in specs)


def test_quality_change_check_and_undefined(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "QC_Check.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "operating_cash_flow_change" and c.period_index == 1
    )
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    # Tamper change check on fresh trainer
    trainer2, _ = build_training_workbook(data, tmp_path / "QC_Tamper.xlsx")
    smap2 = load_semantic_map(answer)
    comp2 = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "operating_cash_flow_change" and c.period_index == 1
    )
    wb = load_workbook(trainer2, data_only=False)
    row, col = parse_cell_ref(comp2.cell)
    wb[comp2.tab].cell(row=row, column=col).value = comp2.formula
    assert _fill_rgb(wb[comp2.tab].cell(row=row, column=col)) == "FFFF00"
    ws = wb["Earnings Quality"]
    check_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "EARNINGS QUALITY CHANGE CHECK"
    )
    ws.cell(check_row, 3).value = '="TAMPERED"'
    wb.save(trainer2)
    wb.close()
    with pytest.raises(
        ValueError, match="Trusted workbook cell was modified: Earnings Quality"
    ):
        check_workbook(trainer2)
    wb = load_workbook(trainer2, data_only=False)
    assert _fill_rgb(wb[comp2.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

    # Zero NI period -> conversion change #N/A
    d1, d2, d3 = date(2023, 12, 31), date(2024, 12, 31), date(2025, 12, 31)
    periods = [
        FinancialPeriod(end_date=d1, label="FY2023"),
        FinancialPeriod(end_date=d2, label="FY2024"),
        FinancialPeriod(end_date=d3, label="FY2025"),
    ]

    def vals(a, b, c):
        return {d1: a, d2: b, d3: c}

    fin = StandardizedFinancials(
        ticker="NI0",
        company_name="Zero NI Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=[
            _li("Revenue", vals(1000, 1100, 1200)),
            _li("Finance costs", vals(-10, -11, -12)),
            _li("Finance income", vals(0, 0, 0)),
            _li("Profit before tax", vals(200, 0, 240)),
            _li("Income tax expense", vals(-30, 0, -36)),
            _li("Profit for the year", vals(170, 0, 204)),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", vals(50, 55, 60)),
            _li("Trade receivables", vals(40, 45, 50)),
            _li("Property, plant and equipment", vals(400, 420, 440)),
            _li("Trade payables", vals(30, 35, 40)),
            _li("Bank borrowings", vals(200, 210, 220)),
            _li("Total equity", vals(260, 275, 290)),
            _li("Total assets", vals(490, 520, 550)),
        ],
        cash_flow=[_li("Net cash from operating activities", vals(80, 90, 100))],
    )
    trainer3, answer3 = build_training_workbook(fin, tmp_path / "QC_NA.xlsx")
    smap3 = load_semantic_map(answer3)
    conv_chg = next(
        c
        for c in smap3.all_ordered()
        if c.family_id == "cash_conversion_ratio_change" and c.period_index == 1
    )
    assert conv_chg.expected_value == UNDEFINED_RATIO
    assert "IFERROR" not in conv_chg.formula.upper()

    wb = load_workbook(trainer3, data_only=False)
    row, col = parse_cell_ref(conv_chg.cell)
    wb[conv_chg.tab].cell(row=row, column=col).value = conv_chg.formula
    wb.save(trainer3)
    wb.close()
    assert check_workbook(trainer3).correct == 1
    _inject_formula_and_cached_value(
        trainer3, conv_chg.tab, conv_chg.cell, formula="=NA()", cached_value=UNDEFINED_RATIO
    )
    assert check_workbook(trainer3).correct == 1
    _inject_formula_and_cached_value(
        trainer3, conv_chg.tab, conv_chg.cell, formula="=0", cached_value=0.0
    )
    assert check_workbook(trainer3).incorrect == 1


def test_quality_change_composes_with_judgment_and_normalization(tmp_path):
    import json

    data = _ingest_demo()
    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer, answer = build_training_workbook(
        data, tmp_path / "QC_COMPOSE.xlsx", assumptions
    )
    smap = load_semantic_map(answer)

    hist = next(c for c in smap.all_ordered() if c.family_id == "noa_agg" and c.period_index == 1)
    norm = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "normalized_nopat" and c.period_index == 1
    )
    qc = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "operating_cash_flow_change" and c.period_index == 1
    )

    wb = load_workbook(trainer, data_only=False)
    # Keep judgment/normalization as supplied reference treatments
    for comp in (hist, norm, qc):
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()

    summary = check_workbook(trainer)
    assert summary.correct == 3
    assert summary.incorrect == 0

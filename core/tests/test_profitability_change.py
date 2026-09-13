"""Step 9C.2 — historical RNOA Margin / Turnover change attribution."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    PROFITABILITY_CHANGE_COMPONENT_CATALOG,
    PROFITABILITY_DRIVER_COMPONENT_CATALOG,
    WORKING_CAPITAL_COMPONENT_CATALOG,
    expand_profitability_change_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.historical_expected import (
    expected_value_for_component,
    profitability_change_expected_series,
)
from core.model.profitability_change import compute_profitability_change_series
from core.model.profitability_drivers import ProfitabilityDriverSeries
from core.model.ratio_values import UNDEFINED_RATIO
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family
from core.tests.test_normalization import _inject_formula_and_cached_value

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    adapter = HKManualDocumentAdapter()
    return adapter.ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values_by_date, concept=""):
    return LineItem(label=label, values=values_by_date, concept=concept)


def _anchor_with_drivers(
    *,
    margin: tuple,
    turnover: tuple,
    rnoa: tuple,
    average_noa: tuple | None = None,
):
    n = len(margin)
    assert len(turnover) == n and len(rnoa) == n
    avg = average_noa or tuple(None if i == 0 else 100.0 + i for i in range(n))
    intensity = tuple(
        None if i == 0 else (UNDEFINED_RATIO if turnover[i] == UNDEFINED_RATIO else 1.0)
        for i in range(n)
    )
    drivers = ProfitabilityDriverSeries(
        average_noa=avg,
        noa_turnover=turnover,
        noa_intensity=intensity,
        rnoa_from_margin_turnover=tuple(
            UNDEFINED_RATIO
            if (m == UNDEFINED_RATIO or t == UNDEFINED_RATIO)
            else (None if m is None or t is None else float(m) * float(t))
            for m, t in zip(margin, turnover)
        ),
    )
    anchor = SimpleNamespace(dupont={"NOPAT Margin": list(margin), "RNOA": list(rnoa)})
    return anchor, drivers


def test_ordinary_midpoint_attribution():
    margin = (0.08, 0.10, 0.12)
    turnover = (1.5, 2.0, 2.5)
    rnoa = (0.12, 0.20, 0.30)
    anchor, drivers = _anchor_with_drivers(margin=margin, turnover=turnover, rnoa=rnoa)
    with patch(
        "core.model.profitability_change.compute_profitability_driver_series",
        return_value=drivers,
    ):
        series = compute_profitability_change_series(anchor)

    for field in (
        series.nopat_margin_change,
        series.noa_turnover_change,
        series.rnoa_change,
        series.margin_effect_on_rnoa,
        series.turnover_effect_on_rnoa,
        series.rnoa_change_from_drivers,
    ):
        assert field[0] is None
        assert field[1] is None

    assert series.nopat_margin_change[2] == pytest.approx(0.02)
    assert series.noa_turnover_change[2] == pytest.approx(0.5)
    assert series.rnoa_change[2] == pytest.approx(0.10)
    assert series.margin_effect_on_rnoa[2] == pytest.approx(0.045)
    assert series.turnover_effect_on_rnoa[2] == pytest.approx(0.055)
    assert series.rnoa_change_from_drivers[2] == pytest.approx(0.10)


def test_profitability_change_edge_cases():
    # Unchanged margin -> margin effect 0; unchanged turnover -> turnover effect 0
    margin = (0.1, 0.1, 0.1)
    turnover = (2.0, 2.0, 2.5)
    rnoa = (0.2, 0.2, 0.25)
    anchor, drivers = _anchor_with_drivers(margin=margin, turnover=turnover, rnoa=rnoa)
    with patch(
        "core.model.profitability_change.compute_profitability_driver_series",
        return_value=drivers,
    ):
        series = compute_profitability_change_series(anchor)
    assert series.margin_effect_on_rnoa[2] == pytest.approx(0.0)
    assert series.turnover_effect_on_rnoa[2] == pytest.approx(0.05)

    margin2 = (0.1, 0.12, 0.12)
    turnover2 = (2.0, 2.0, 2.0)
    rnoa2 = (0.2, 0.24, 0.24)
    anchor2, drivers2 = _anchor_with_drivers(
        margin=margin2, turnover=turnover2, rnoa=rnoa2
    )
    with patch(
        "core.model.profitability_change.compute_profitability_driver_series",
        return_value=drivers2,
    ):
        series2 = compute_profitability_change_series(anchor2)
    assert series2.turnover_effect_on_rnoa[2] == pytest.approx(0.0)
    assert series2.margin_effect_on_rnoa[2] == pytest.approx(0.0)

    # Negative signed effects retained
    margin_n = (0.12, 0.12, 0.10)
    turnover_n = (2.5, 2.5, 2.0)
    rnoa_n = (0.30, 0.30, 0.20)
    anchor_n, drivers_n = _anchor_with_drivers(
        margin=margin_n, turnover=turnover_n, rnoa=rnoa_n
    )
    with patch(
        "core.model.profitability_change.compute_profitability_driver_series",
        return_value=drivers_n,
    ):
        series_n = compute_profitability_change_series(anchor_n)
    assert series_n.margin_effect_on_rnoa[2] == pytest.approx(-0.045)
    assert series_n.turnover_effect_on_rnoa[2] == pytest.approx(-0.055)

    # Undefined margin -> both effects and driver Δ #N/A; numeric direct change allowed
    margin_u = (0.1, 0.1, UNDEFINED_RATIO)
    turnover_u = (2.0, 2.0, 2.5)
    rnoa_u = (0.2, 0.2, 0.25)
    anchor_u, drivers_u = _anchor_with_drivers(
        margin=margin_u, turnover=turnover_u, rnoa=rnoa_u
    )
    with patch(
        "core.model.profitability_change.compute_profitability_driver_series",
        return_value=drivers_u,
    ):
        series_u = compute_profitability_change_series(anchor_u)
    assert series_u.margin_effect_on_rnoa[2] == UNDEFINED_RATIO
    assert series_u.turnover_effect_on_rnoa[2] == UNDEFINED_RATIO
    assert series_u.rnoa_change_from_drivers[2] == UNDEFINED_RATIO
    assert series_u.rnoa_change[2] == pytest.approx(0.05)

    # Undefined turnover
    margin_t = (0.1, 0.1, 0.12)
    turnover_t = (2.0, 2.0, UNDEFINED_RATIO)
    rnoa_t = (0.2, 0.2, 0.25)
    anchor_t, drivers_t = _anchor_with_drivers(
        margin=margin_t, turnover=turnover_t, rnoa=rnoa_t
    )
    with patch(
        "core.model.profitability_change.compute_profitability_driver_series",
        return_value=drivers_t,
    ):
        series_t = compute_profitability_change_series(anchor_t)
    assert series_t.rnoa_change_from_drivers[2] == UNDEFINED_RATIO

    # Numeric driver + mismatched RNOA -> ValueError
    margin_m = (0.1, 0.1, 0.12)
    turnover_m = (2.0, 2.0, 2.5)
    rnoa_m = (0.2, 0.2, 0.99)
    anchor_m, drivers_m = _anchor_with_drivers(
        margin=margin_m, turnover=turnover_m, rnoa=rnoa_m
    )
    with patch(
        "core.model.profitability_change.compute_profitability_driver_series",
        return_value=drivers_m,
    ):
        with pytest.raises(ValueError, match="does not reconcile"):
            compute_profitability_change_series(anchor_m)

    # Length mismatch
    anchor_l = SimpleNamespace(
        dupont={"NOPAT Margin": [0.1, 0.1, 0.12], "RNOA": [0.2, 0.2]}
    )
    drivers_l = ProfitabilityDriverSeries(
        average_noa=(None, 100.0, 110.0),
        noa_turnover=(1.0, 2.0, 2.5),
        noa_intensity=(None, 0.5, 0.4),
        rnoa_from_margin_turnover=(0.1, 0.2, 0.3),
    )
    with patch(
        "core.model.profitability_change.compute_profitability_driver_series",
        return_value=drivers_l,
    ):
        with pytest.raises(ValueError, match="length mismatch"):
            compute_profitability_change_series(anchor_l)


def test_catalog_expand_and_expected_keys():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(PROFITABILITY_CHANGE_COMPONENT_CATALOG) == 6
    assert len({f.id for f in PROFITABILITY_CHANGE_COMPONENT_CATALOG}) == 6
    assert all(
        f.period_scope == "post_comparable" for f in PROFITABILITY_CHANGE_COMPONENT_CATALOG
    )
    assert [f.order for f in PROFITABILITY_CHANGE_COMPONENT_CATALOG] == list(range(50, 56))
    assert len(PROFITABILITY_DRIVER_COMPONENT_CATALOG) == 4
    assert len(WORKING_CAPITAL_COMPONENT_CATALOG) == 11

    specs = expand_profitability_change_specs(periods, start_order=205)
    assert len(specs) == 18

    data = _ingest_demo()
    builder = ReferenceModelBuilder(data)
    assert len(builder.profitability_change_specs) == 18
    series = profitability_change_expected_series(builder.anchor)
    assert set(series) == {f.id for f in PROFITABILITY_CHANGE_COMPONENT_CATALOG}
    for j in range(2, len(builder.periods)):
        driver = series["rnoa_change_from_drivers"][j]
        direct = series["rnoa_change"][j]
        if driver != UNDEFINED_RATIO:
            assert driver == pytest.approx(float(direct))


def test_demo_profitability_change_surface(tmp_path):
    import json

    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "BASE_PC.xlsx")
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == 312
    assert len(group_components_by_family(smap)) == 74
    pc_comps = [c for c in smap.all_ordered() if c.category == "profitability_change"]
    assert len(pc_comps) == 18
    assert all(c.tab == "ALT DuPont" for c in pc_comps)
    assert all(c.period_index >= 2 for c in pc_comps)

    wb = load_workbook(answer)
    ws = wb["ALT DuPont"]
    assert ws.cell(14, 1).value == "RNOA DRIVER DECOMPOSITION"
    assert ws.cell(19, 1).value == "RNOA DRIVER CHECK"
    assert ws.cell(21, 1).value == "RNOA CHANGE ATTRIBUTION"
    assert ws.cell(22, 1).value == "Change in NOPAT Margin"
    assert ws.cell(27, 1).value == "RNOA Change from Drivers"
    assert ws.cell(28, 1).value == "RNOA CHANGE DRIVER CHECK"
    for col in (2, 3):
        for row in range(22, 29):
            assert ws.cell(row, col).value == "N/A"
    practice = {(c.tab, c.cell) for c in smap.all_ordered()}
    assert ("ALT DuPont", "D28") not in practice

    margin_chg = next(
        c for c in pc_comps if c.family_id == "nopat_margin_change" and c.period_index == 2
    )
    assert "ABS" not in margin_chg.formula.upper()
    assert "IFERROR" not in margin_chg.formula.upper()
    effect = next(
        c for c in pc_comps if c.family_id == "rnoa_margin_effect" and c.period_index == 2
    )
    assert "/2)" in effect.formula
    assert "ABS" not in effect.formula.upper()

    wb_t = load_workbook(trainer)
    for comp in pc_comps:
        row, col = parse_cell_ref(comp.cell)
        cell = wb_t[comp.tab].cell(row=row, column=col)
        assert cell.value is None
        assert _fill_rgb(cell) == "FFFF00"
        assert cell.comment is None
    wb.close()
    wb_t.close()

    summary = check_workbook(trainer)
    assert summary.total == 312
    assert summary.blank == 312

    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        data, tmp_path / "NORM_PC.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(smap_n.all_ordered()) == 332
    assert len(group_components_by_family(smap_n)) == 78
    assert check_workbook(trainer_n).blank == 332


def test_live_classification_changes_profitability_change(tmp_path):
    d1, d2, d3 = date(2023, 12, 31), date(2024, 12, 31), date(2025, 12, 31)
    periods = [
        FinancialPeriod(end_date=d1, label="FY2023"),
        FinancialPeriod(end_date=d2, label="FY2024"),
        FinancialPeriod(end_date=d3, label="FY2025"),
    ]

    def vals(a, b, c):
        return {d1: a, d2: b, d3: c}

    fin = StandardizedFinancials(
        ticker="STI3",
        company_name="STI Change Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=[
            _li("Revenue", vals(1000, 1100, 1200)),
            _li("Finance costs", vals(-10, -11, -12)),
            _li("Finance income", vals(0, 0, 0)),
            _li("Profit before tax", vals(200, 220, 240)),
            _li("Income tax expense", vals(-30, -33, -36)),
            _li("Profit for the year", vals(170, 187, 204)),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", vals(50, 55, 60)),
            _li("Trade receivables", vals(40, 45, 50)),
            _li("Short-term investments", vals(100, 120, 140)),
            _li("Property, plant and equipment", vals(400, 420, 440)),
            _li("Trade payables", vals(30, 35, 40)),
            _li("Bank borrowings", vals(200, 210, 220)),
            _li("Total equity", vals(360, 395, 430)),
        ],
        cash_flow=[_li("Net cash from operating activities", vals(50, 60, 70))],
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "STI_PC.xlsx")
    smap = load_semantic_map(answer)
    turn_chg = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "noa_turnover_change" and c.period_index == 2
    )
    turn_effect = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "rnoa_turnover_effect" and c.period_index == 2
    )
    driver_chg = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "rnoa_change_from_drivers" and c.period_index == 2
    )

    ref_turn_chg = turn_chg.expected_value
    alt_anchor = compute_anchor(
        fin,
        [d1, d2, d3],
        classification_overrides={
            "label:Short-term investments": "Operating Working Capital Asset"
        },
    )
    alt_turn_chg = expected_value_for_component(alt_anchor, turn_chg)
    alt_turn_effect = expected_value_for_component(alt_anchor, turn_effect)
    alt_driver_chg = expected_value_for_component(alt_anchor, driver_chg)
    assert alt_turn_chg != pytest.approx(float(ref_turn_chg))
    assert isinstance(alt_turn_effect, float)
    assert isinstance(alt_driver_chg, float)

    wb = load_workbook(trainer, data_only=False)
    wb["Accounting Judgment"].cell(5, 6).value = "Operating Working Capital Asset"
    wb.save(trainer)
    wb.close()

    # Structurally different formula: Check uses live expected vs cached value
    _inject_formula_and_cached_value(
        trainer,
        turn_chg.tab,
        turn_chg.cell,
        formula="=C16-B16",
        cached_value=float(ref_turn_chg),
    )
    assert check_workbook(trainer).incorrect == 1

    _inject_formula_and_cached_value(
        trainer,
        turn_chg.tab,
        turn_chg.cell,
        formula="=C16-B16",
        cached_value=float(alt_turn_chg),
    )
    assert check_workbook(trainer).correct == 1


def test_change_formula_check_and_tamper(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "PC_Check.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "rnoa_margin_effect" and c.period_index == 2
    )
    expected = float(comp.expected_value)

    # Exact formula -> green
    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    # Structurally different formula with correct cached value -> green
    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula="=0.045+0", cached_value=expected
    )
    assert check_workbook(trainer).correct == 1

    # Incorrect cached numeric -> red
    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula="=0.045+0", cached_value=expected + 1.0
    )
    assert check_workbook(trainer).incorrect == 1

    # Tamper RNOA CHANGE DRIVER CHECK on a fresh Trainer (yellow practice cell)
    trainer2, _ = build_training_workbook(data, tmp_path / "PC_Tamper.xlsx")
    smap2 = load_semantic_map(answer)
    comp2 = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "rnoa_margin_effect" and c.period_index == 2
    )
    wb = load_workbook(trainer2, data_only=False)
    row, col = parse_cell_ref(comp2.cell)
    wb[comp2.tab].cell(row=row, column=col).value = comp2.formula
    ws = wb["ALT DuPont"]
    check_row = next(
        r
        for r in range(1, (ws.max_row or 1) + 1)
        if ws.cell(r, 1).value == "RNOA CHANGE DRIVER CHECK"
    )
    ws.cell(check_row, 4).value = '="TAMPERED"'
    wb.save(trainer2)
    wb.close()
    with pytest.raises(ValueError, match="Trusted workbook cell was modified: ALT DuPont"):
        check_workbook(trainer2)
    wb = load_workbook(trainer2, data_only=False)
    assert _fill_rgb(wb[comp2.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


def test_undefined_change_attribution_check(tmp_path):
    d1, d2, d3 = date(2023, 12, 31), date(2024, 12, 31), date(2025, 12, 31)
    periods = [
        FinancialPeriod(end_date=d1, label="FY2023"),
        FinancialPeriod(end_date=d2, label="FY2024"),
        FinancialPeriod(end_date=d3, label="FY2025"),
    ]

    def vals(a, b, c):
        return {d1: a, d2: b, d3: c}

    # Zero Average NOA in period 3 (index 2) via zero NOA endpoints around that period
    # Make ending NOA at d2 and d3 both ~0 so average NOA at index 2 is 0 -> turnover #N/A
    fin = StandardizedFinancials(
        ticker="NA3",
        company_name="NA Change Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=[
            _li("Revenue", vals(1000, 1100, 1200)),
            _li("Finance costs", vals(-10, -11, -12)),
            _li("Finance income", vals(0, 0, 0)),
            _li("Profit before tax", vals(200, 220, 240)),
            _li("Income tax expense", vals(-30, -33, -36)),
            _li("Profit for the year", vals(170, 187, 204)),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", vals(100, 110, 120)),
            _li("Trade receivables", vals(50, 0, 0)),
            _li("Trade payables", vals(50, 0, 0)),
            _li("Bank borrowings", vals(200, 210, 220)),
            _li("Total equity", vals(-100, -100, -100)),
        ],
        cash_flow=[_li("Net cash from operating activities", vals(50, 60, 70))],
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "PC_NA.xlsx")
    smap = load_semantic_map(answer)
    effect = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "rnoa_margin_effect" and c.period_index == 2
    )
    driver = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "rnoa_change_from_drivers" and c.period_index == 2
    )
    assert effect.expected_value == UNDEFINED_RATIO
    assert driver.expected_value == UNDEFINED_RATIO

    wb = load_workbook(answer, data_only=False)
    check_row = next(
        r
        for r in range(1, (wb["ALT DuPont"].max_row or 1) + 1)
        if wb["ALT DuPont"].cell(r, 1).value == "RNOA CHANGE DRIVER CHECK"
    )
    formula = str(wb["ALT DuPont"].cell(check_row, 4).value)
    assert "ISNA(" in formula
    assert '"N/A"' in formula
    wb.close()

    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(effect.cell)
    wb[effect.tab].cell(row=row, column=col).value = effect.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1
    _inject_formula_and_cached_value(
        trainer, effect.tab, effect.cell, formula="=NA()", cached_value=UNDEFINED_RATIO
    )
    assert check_workbook(trainer).correct == 1
    _inject_formula_and_cached_value(
        trainer, effect.tab, effect.cell, formula="=0", cached_value=0.0
    )
    assert check_workbook(trainer).incorrect == 1

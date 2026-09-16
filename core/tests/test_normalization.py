"""Step 8B2 — explicit normalization candidates, series, sheets, and dynamic Check."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    COMPONENT_CATALOG,
    NORMALIZATION_COMPONENT_CATALOG,
    expand_historical_specs,
    expand_normalization_specs,
)
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.historical_expected import expected_value_for_component
from core.model.normalization import (
    NORMALIZATION_TREATMENTS,
    SUPPORTED_NORMALIZATION_SCOPE,
    compute_normalization_series,
    normalization_cases,
    resolve_income_statement_selector,
)
from core.model.period_axis import canonical_fiscal_periods
from core.trainer.checker import check_workbook
from core.trainer.check_context import (
    CHECK_CONTEXT_SCHEMA_VERSION,
    live_normalization_treatment_formula,
    load_check_context,
    normalization_treatments_for_check,
)
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    adapter = HKManualDocumentAdapter()
    return adapter.ingest([DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)])


def _demo_assumptions():
    return json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))


def _build_norm_pair(tmp_path):
    data = _ingest_demo()
    return build_training_workbook(
        data, tmp_path / "DEMO_HK_Trainer.xlsx", _demo_assumptions()
    )


def _li(label, *vals, concept=None):
    periods = [date(2021 + i, 12, 31) for i in range(len(vals))]
    return LineItem(
        label=label,
        values={p: float(v) for p, v in zip(periods, vals)},
        concept=concept,
    )


def _tiny_fin(*is_items):
    d1, d2 = date(2021, 12, 31), date(2022, 12, 31)
    return StandardizedFinancials(
        company_name="Tiny Co",
        ticker="TINY",
        currency="HKD",
        jurisdiction="HK",
        units="HKD millions",
        periods=[
            FinancialPeriod(end_date=d1, label="FY2021"),
            FinancialPeriod(end_date=d2, label="FY2022"),
        ],
        income_statement=list(is_items),
        balance_sheet=[
            LineItem(label="Cash", values={d1: 10, d2: 12}, concept="cash"),
            LineItem(
                label="Total equity",
                values={d1: 10, d2: 12},
                concept="total_equity",
            ),
        ],
        cash_flow=[
            LineItem(
                label="Net cash from operating activities",
                values={d1: 1, d2: 2},
            )
        ],
    )


# --- Task 1: cases / selectors -------------------------------------------------


def test_normalization_candidate_concept_and_label_selectors():
    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("Restructuring expense", -20, 0, concept="restructuring_expense"),
        _li("One-off legal settlement", -5, -1),
    )
    periods = canonical_fiscal_periods(fin)
    by_concept = resolve_income_statement_selector(fin, "concept:restructuring_expense")
    assert by_concept.concept == "restructuring_expense"
    by_label = resolve_income_statement_selector(fin, "label:One-off legal settlement")
    assert by_label.label == "One-off legal settlement"

    cases = normalization_cases(
        fin,
        periods,
        {
            "normalizationCandidates": [
                {
                    "selector": "concept:restructuring_expense",
                    "referenceTreatment": "Non-recurring",
                    "scope": SUPPORTED_NORMALIZATION_SCOPE,
                    "topic": "t",
                    "referenceRationale": "r",
                    "consequenceNote": "c",
                }
            ]
        },
    )
    assert len(cases) == 1
    assert cases[0].reference_treatment == "Non-recurring"
    assert cases[0].alternatives == ("Recurring",)
    assert cases[0].override_selector == "concept:restructuring_expense"


def test_normalization_candidate_unknown_and_ambiguous_and_invalid():
    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("Special charge", -1, 0),
        _li("Special charge", -2, 0),
    )
    periods = canonical_fiscal_periods(fin)
    with pytest.raises(ValueError, match="matched no income-statement line"):
        resolve_income_statement_selector(fin, "concept:missing")
    with pytest.raises(ValueError, match="matched 2"):
        resolve_income_statement_selector(fin, "label:Special charge")
    with pytest.raises(ValueError, match="unsupported scope"):
        normalization_cases(
            fin,
            periods,
            {
                "normalizationCandidates": [
                    {
                        "selector": "concept:revenue",
                        "referenceTreatment": "Recurring",
                        "scope": "special_tax",
                        "topic": "t",
                        "referenceRationale": "r",
                        "consequenceNote": "c",
                    }
                ]
            },
        )
    with pytest.raises(ValueError, match="referenceTreatment must be one of"):
        normalization_cases(
            fin,
            periods,
            {
                "normalizationCandidates": [
                    {
                        "selector": "concept:revenue",
                        "referenceTreatment": "Maybe",
                        "scope": SUPPORTED_NORMALIZATION_SCOPE,
                        "topic": "t",
                        "referenceRationale": "r",
                        "consequenceNote": "c",
                    }
                ]
            },
        )


def test_no_label_heuristic_without_configured_candidate():
    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("One-off restructuring expense", -50, 0),
    )
    periods = canonical_fiscal_periods(fin)
    assert normalization_cases(fin, periods, {}) == ()
    assert normalization_cases(fin, periods, {"normalizationCandidates": []}) == ()


def test_zero_valued_candidate_suppressed_and_order_dense():
    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("Zero charge", 0, 0, concept="zero_charge"),
        _li("Live charge", -10, 0, concept="live_charge"),
    )
    periods = canonical_fiscal_periods(fin)
    cases = normalization_cases(
        fin,
        periods,
        {
            "normalizationCandidates": [
                {
                    "selector": "concept:zero_charge",
                    "referenceTreatment": "Non-recurring",
                    "scope": SUPPORTED_NORMALIZATION_SCOPE,
                    "topic": "zero",
                    "referenceRationale": "r",
                    "consequenceNote": "c",
                },
                {
                    "selector": "concept:live_charge",
                    "referenceTreatment": "Recurring",
                    "scope": SUPPORTED_NORMALIZATION_SCOPE,
                    "topic": "live",
                    "referenceRationale": "r",
                    "consequenceNote": "c",
                },
            ]
        },
    )
    assert len(cases) == 1
    assert cases[0].order == 1
    assert cases[0].override_selector == "concept:live_charge"
    assert cases[0].reference_treatment == "Recurring"
    assert cases[0].alternatives == ("Non-recurring",)


# --- Task 2/3: demo split + series --------------------------------------------


def test_demo_admin_plus_restructuring_preserves_prior_admin_expense():
    data = _ingest_demo()
    by_label = {item.label: item for item in data.income_statement}
    admin = by_label["Administrative expenses"]
    restruct = by_label["Restructuring expense"]
    assert restruct.concept == "restructuring_expense"
    prior = {
        date(2021, 12, 31): -1700.0,
        date(2022, 12, 31): -1840.0,
        date(2023, 12, 31): -2020.0,
        date(2024, 12, 31): -2240.0,
        date(2025, 12, 31): -2500.0,
    }
    for period, expected in prior.items():
        assert float(admin.values[period]) + float(restruct.values[period]) == pytest.approx(
            expected
        )


def test_normalization_series_sign_and_tax_convention():
    data = _ingest_demo()
    assumptions = _demo_assumptions()
    periods = canonical_fiscal_periods(data)
    cases = normalization_cases(data, periods, assumptions)
    assert len(cases) == 1
    anchor = compute_anchor(data, periods)
    series = compute_normalization_series(data, periods, anchor, cases)
    # FY2023 (index 2): pretax add-back = -(-200) = 200
    assert series.pretax_adjustment[2] == pytest.approx(200.0)
    etr = float(anchor.historical.effective_tax_rate[2])
    assert series.after_tax_adjustment[2] == pytest.approx(200.0 * (1.0 - etr))
    assert series.normalized_nopat[2] == pytest.approx(
        float(anchor.historical.nopat[2]) + series.after_tax_adjustment[2]
    )
    recurring = compute_normalization_series(
        data, periods, anchor, cases, {cases[0].id: "Recurring"}
    )
    assert all(v == pytest.approx(0.0) for v in recurring.pretax_adjustment)
    assert recurring.normalized_nopat == tuple(float(x) for x in anchor.historical.nopat)


# --- Task 4/5: sheets + families ----------------------------------------------


def test_normalization_sheets_answer_key_and_trainer_contract(tmp_path):
    from core.engine.reference_model import (
        EARNINGS_NORMALIZATION_SHEET,
        NORMALIZATION_JUDGMENT_SHEET,
        ReferenceModelBuilder,
    )

    data = _ingest_demo()
    assumptions = _demo_assumptions()
    builder = ReferenceModelBuilder(data, assumptions)
    assert len(builder.normalization_cases) == 1
    case = builder.normalization_cases[0]

    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    assert len(smap.all_ordered()) == 332
    assert len(COMPONENT_CATALOG) == 25
    assert len(NORMALIZATION_COMPONENT_CATALOG) == 4
    assert [f.order for f in NORMALIZATION_COMPONENT_CATALOG] == [26, 27, 28, 29]

    groups = group_components_by_family(smap)
    assert len(groups) == 78
    assert groups[-1]["family_order"] == 90

    wb_a = load_workbook(answer_key_path, data_only=False)
    wb_t = load_workbook(trainer_path, data_only=False)
    assert NORMALIZATION_JUDGMENT_SHEET in wb_a.sheetnames
    assert EARNINGS_NORMALIZATION_SHEET in wb_a.sheetnames
    assert NORMALIZATION_JUDGMENT_SHEET in wb_t.sheetnames
    assert EARNINGS_NORMALIZATION_SHEET in wb_t.sheetnames

    ws_a = wb_a[NORMALIZATION_JUDGMENT_SHEET]
    ws_t = wb_t[NORMALIZATION_JUDGMENT_SHEET]
    assert ws_a["A1"].value == "Normalization Judgment"
    assert "reference convention" in str(ws_a["A2"].value).lower()
    assert "blank f uses the supplied reference" in str(ws_a["A3"].value).lower()
    assert ws_a.cell(5, 4).value == case.reference_treatment
    assert ws_a.cell(5, 5).value == ", ".join(case.alternatives)
    assert ws_a.cell(5, 6).value == case.reference_treatment
    assert ws_a.cell(5, 7).value == case.model_rationale
    assert ws_a.cell(5, 8).value == case.model_consequence
    for col in (6, 7, 8):
        assert _fill_rgb(ws_a.cell(5, col)) in {"", "FFFFFF"}
        assert ws_t.cell(5, col).value is None
        assert ws_t.cell(5, col).comment is None
        assert _fill_rgb(ws_t.cell(5, col)) == "FFFF00"

    expected_formula = live_normalization_treatment_formula(5, case.reference_treatment)
    assert "$D$" not in expected_formula
    assert wb_a[EARNINGS_NORMALIZATION_SHEET].cell(5, 2).value == expected_formula
    assert wb_t[EARNINGS_NORMALIZATION_SHEET].cell(5, 2).value == expected_formula
    assert _fill_rgb(wb_a[EARNINGS_NORMALIZATION_SHEET].cell(5, 2)) != "FFFF00"

    # Detail source links resolve the configured line identity (Income Statement row).
    source_row = None
    for row in range(7, (wb_a["Income Statement"].max_row or 7) + 1):
        if wb_a["Income Statement"].cell(row, 1).value == case.label:
            source_row = row
            break
    assert source_row is not None
    assert wb_a[EARNINGS_NORMALIZATION_SHEET].cell(5, 3).value == (
        f"='Income Statement'!B{source_row}"
    )

    # NORMALIZATION CHECK is populated and not a practice cell.
    check_row = None
    for row in range(1, (wb_a[EARNINGS_NORMALIZATION_SHEET].max_row or 1) + 1):
        if wb_a[EARNINGS_NORMALIZATION_SHEET].cell(row, 1).value == "NORMALIZATION CHECK":
            check_row = row
            break
    assert check_row is not None
    assert isinstance(wb_a[EARNINGS_NORMALIZATION_SHEET].cell(check_row, 3).value, str)
    practice_cells = {(c.tab, c.cell) for c in smap.all_ordered()}
    assert (EARNINGS_NORMALIZATION_SHEET, f"C{check_row}") not in practice_cells

    for needle in (case.model_rationale, case.model_consequence):
        for name in wb_t.sheetnames:
            sheet = wb_t[name]
            for row in sheet.iter_rows(
                max_row=sheet.max_row or 1, max_col=sheet.max_column or 1
            ):
                for cell in row:
                    if isinstance(cell.value, str):
                        assert needle not in cell.value
                    if cell.comment is not None:
                        assert needle not in (cell.comment.text or "")

    for wb in (wb_a, wb_t):
        dvs = list(wb[NORMALIZATION_JUDGMENT_SHEET].data_validations.dataValidation)
        assert any("F5" in str(dv.sqref) for dv in dvs)
        assert any(
            "Non-recurring" in str(dv.formula1) and "Recurring" in str(dv.formula1)
            for dv in dvs
        )

    ctx = load_check_context(answer_key_path)
    assert ctx is not None
    assert ctx.schema_version == CHECK_CONTEXT_SCHEMA_VERSION
    assert len(ctx.normalization_bindings) == 1
    blob = json.dumps(
        {
            "bindings": [b.__dict__ for b in ctx.normalization_bindings],
            "source": ctx.source_payload,
        }
    )
    assert case.model_rationale not in blob
    assert case.model_consequence not in blob

    wb_a.close()
    wb_t.close()


def test_no_normalization_assumptions_keeps_quality_surface_and_omits_norm_sheets(tmp_path):
    data = _ingest_demo()
    trainer_path, answer_key_path = build_training_workbook(
        data, tmp_path / "BASE_Trainer.xlsx"
    )
    smap = load_semantic_map(answer_key_path)
    assert len(smap.all_ordered()) == 312
    assert len(expand_historical_specs(canonical_fiscal_periods(data))) == 118
    assert len(group_components_by_family(smap)) == 74
    wb = load_workbook(answer_key_path, data_only=False)
    assert "Normalization Judgment" not in wb.sheetnames
    assert "Earnings Normalization" not in wb.sheetnames
    wb.close()
    summary = check_workbook(trainer_path)
    assert summary.total == 312
    assert summary.blank == 312


def test_expand_normalization_specs_continuous_order():
    data = _ingest_demo()
    periods = canonical_fiscal_periods(data)
    hist = expand_historical_specs(periods)
    specs = expand_normalization_specs(periods, start_order=len(hist) + 1)
    assert len(specs) == 20
    assert [s.order for s in specs] == list(range(119, 139))
    assert specs[0].family_id == "pretax_normalization_adjustment"
    assert specs[-1].family_id == "normalized_net_income"


# --- Task 6: dynamic Check ----------------------------------------------------


def _set_normalization_treatment(trainer_path, treatment):
    wb = load_workbook(trainer_path, data_only=False)
    wb["Normalization Judgment"].cell(5, 6).value = treatment
    wb.save(trainer_path)
    wb.close()


def _inject_formula_and_cached_value(
    workbook_path: Path,
    sheet: str,
    cell: str,
    *,
    formula: str,
    cached_value: float | str,
) -> None:
    import os
    import tempfile
    import zipfile
    import xml.etree.ElementTree as ET

    ssml = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    od_rels = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    formula_body = formula[1:] if formula.startswith("=") else formula
    workbook_path = Path(workbook_path)

    def sheet_path(zf, sheet_name: str) -> str:
        wb = ET.fromstring(zf.read("xl/workbook.xml"))
        rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
        sheets = wb.find(f"{{{ssml}}}sheets")
        if sheets is None:
            raise FileNotFoundError("workbook.xml missing sheets")
        rid = None
        for sheet_el in sheets:
            if sheet_el.attrib.get("name") == sheet_name:
                rid = sheet_el.attrib.get(f"{{{od_rels}}}id")
                break
        if not rid:
            raise FileNotFoundError(f"Sheet not found: {sheet_name}")
        target = None
        for rel in rels:
            if rel.attrib.get("Id") == rid:
                target = rel.attrib.get("Target", "")
                break
        if not target:
            raise FileNotFoundError(f"Relationship missing for sheet {sheet_name}")
        target = target.lstrip("/")
        if not target.startswith("xl/"):
            target = f"xl/{target}"
        return target

    with zipfile.ZipFile(workbook_path, "r") as zf_in:
        path = sheet_path(zf_in, sheet)
        sheet_xml = zf_in.read(path)
        other_members = {
            info.filename: zf_in.read(info.filename)
            for info in zf_in.infolist()
            if info.filename != path
        }

    root = ET.fromstring(sheet_xml)
    cell_el = None
    for c_el in root.iter(f"{{{ssml}}}c"):
        if c_el.attrib.get("r") == cell:
            cell_el = c_el
            break
    if cell_el is None:
        raise FileNotFoundError(f"Cell {cell} not found on sheet {sheet}")

    for child in list(cell_el):
        if child.tag in {f"{{{ssml}}}f", f"{{{ssml}}}v", f"{{{ssml}}}is"}:
            cell_el.remove(child)
    f_el = ET.SubElement(cell_el, f"{{{ssml}}}f")
    f_el.text = formula_body
    v_el = ET.SubElement(cell_el, f"{{{ssml}}}v")
    if isinstance(cached_value, str):
        # Excel error cached results use cell type "e" (e.g. #N/A).
        cell_el.set("t", "e")
        v_el.text = cached_value
    else:
        cell_el.attrib.pop("t", None)
        v_el.text = f"{float(cached_value)}"

    ET.register_namespace("", ssml)
    new_sheet_xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)

    fd, tmp_name = tempfile.mkstemp(
        suffix=".xlsx", prefix=workbook_path.stem + "_inj_", dir=str(workbook_path.parent)
    )
    os.close(fd)
    tmp_path = Path(tmp_name)
    try:
        with zipfile.ZipFile(tmp_path, "w") as zf_out:
            for name, data in other_members.items():
                zf_out.writestr(name, data)
            zf_out.writestr(path, new_sheet_xml)
        tmp_path.replace(workbook_path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise


def _latest(smap, family_id: str):
    return max(
        (c for c in smap.all_ordered() if c.family_id == family_id),
        key=lambda c: c.period_index or 0,
    )


def _fy2023(smap, family_id: str):
    return next(
        c
        for c in smap.all_ordered()
        if c.family_id == family_id and c.period_end == "2023-12-31"
    )


def test_dynamic_normalization_reference_and_recurring(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    assert len(smap.all_ordered()) == 332

    summary = check_workbook(trainer_path)
    assert summary.total == 332
    assert summary.blank == 332

    comp = _fy2023(smap, "after_tax_normalization_adjustment")
    # Exact formula under blank F / Non-recurring reference.
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer_path)
    wb.close()
    summary = check_workbook(trainer_path)
    assert summary.correct == 1
    assert summary.incorrect == 0

    # Equivalent formula with Non-recurring cached value.
    _inject_formula_and_cached_value(
        trainer_path,
        comp.tab,
        comp.cell,
        formula=f"={float(comp.expected_value)}",
        cached_value=float(comp.expected_value),
    )
    summary = check_workbook(trainer_path)
    assert summary.correct == 1

    # Switch to Recurring: exact formula still green; expected becomes 0.
    _set_normalization_treatment(trainer_path, "Recurring")
    wb = load_workbook(trainer_path, data_only=False)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer_path)
    wb.close()
    summary = check_workbook(trainer_path)
    assert summary.correct == 1

    data = _ingest_demo()
    ctx = load_check_context(answer_key_path)
    periods = [date.fromisoformat(p) for p in ctx.modeled_periods]
    cases = normalization_cases(data, periods, _demo_assumptions())
    anchor = compute_anchor(data, periods)
    recurring_series = compute_normalization_series(
        data, periods, anchor, cases, {cases[0].id: "Recurring"}
    )
    recurring_expected = float(
        expected_value_for_component(anchor, comp, normalization=recurring_series)
    )
    assert recurring_expected == pytest.approx(0.0)

    _inject_formula_and_cached_value(
        trainer_path,
        comp.tab,
        comp.cell,
        formula=f"={recurring_expected}",
        cached_value=recurring_expected,
    )
    summary = check_workbook(trainer_path)
    assert summary.correct == 1

    # Stale Non-recurring cache under Recurring -> red.
    _inject_formula_and_cached_value(
        trainer_path,
        comp.tab,
        comp.cell,
        formula="=999",
        cached_value=float(comp.expected_value),
    )
    summary = check_workbook(trainer_path)
    assert summary.correct == 0
    assert summary.incorrect == 1


def test_invalid_normalization_treatment_and_tamper_fail_closed(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    comp = _latest(smap, "normalized_nopat")
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer_path)
    wb.close()

    _set_normalization_treatment(trainer_path, "Not A Treatment")
    with pytest.raises(ValueError, match="Invalid treatment"):
        check_workbook(trainer_path)
    # No partial recolor after failure: cell should still be yellow from Trainer blanking.
    wb = load_workbook(trainer_path, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()

    _set_normalization_treatment(trainer_path, "Recurring")
    wb = load_workbook(trainer_path, data_only=False)
    wb["Normalization Judgment"].cell(5, 4).value = "Tampered"
    wb.save(trainer_path)
    wb.close()
    with pytest.raises(ValueError, match="reference prompt was modified"):
        check_workbook(trainer_path)

    trainer_path, _ = _build_norm_pair(tmp_path / "tamper2")
    wb = load_workbook(trainer_path, data_only=False)
    wb["Earnings Normalization"].cell(5, 2).value = "Recurring"
    wb.save(trainer_path)
    wb.close()
    with pytest.raises(ValueError, match="treatment was modified"):
        check_workbook(trainer_path)


def test_combined_classification_and_normalization_state(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    net_debt = max(
        (c for c in smap.all_ordered() if c.family_id == "net_debt"),
        key=lambda c: c.period_index or 0,
    )
    norm_comp = _fy2023(smap, "normalized_nopat")

    wb = load_workbook(trainer_path, data_only=False)
    wb["Accounting Judgment"].cell(5, 6).value = "Financial Liability"
    wb["Normalization Judgment"].cell(5, 6).value = "Recurring"
    row_nd, col_nd = parse_cell_ref(net_debt.cell)
    row_n, col_n = parse_cell_ref(norm_comp.cell)
    wb[net_debt.tab].cell(row=row_nd, column=col_nd).value = net_debt.formula
    wb[norm_comp.tab].cell(row=row_n, column=col_n).value = norm_comp.formula
    wb.save(trainer_path)
    wb.close()

    summary = check_workbook(trainer_path)
    assert summary.correct == 2
    assert summary.incorrect == 0

    ctx = load_check_context(answer_key_path)
    wb = load_workbook(trainer_path, data_only=False)
    treatments = normalization_treatments_for_check(wb, ctx)
    wb.close()
    assert list(treatments.values()) == ["Recurring"]


def test_legacy_schema_v1_check_context_still_loads(tmp_path):
    from core.data.standardized_io import standardized_to_payload
    from core.trainer.check_context import (
        CHECK_CONTEXT_LEGACY_SCHEMA_VERSION,
        CHECK_CONTEXT_MAGIC,
        CHECK_CONTEXT_SHEET,
        load_check_context,
    )
    from openpyxl import Workbook

    data = _ingest_demo()
    payload = {
        "schema_version": CHECK_CONTEXT_LEGACY_SCHEMA_VERSION,
        "source_payload": standardized_to_payload(data),
        "modeled_periods": [p.isoformat() for p in canonical_fiscal_periods(data)],
        "base_classification_overrides": {},
        "judgment_bindings": [
            {
                "order": 1,
                "worksheet_row": 5,
                "case_id": "j1",
                "line_identity": "concept:lease_liability",
                "override_selector": "concept:lease_liability",
                "reference_treatment": "Operating Long-Term Liability",
                "allowed_treatments": [
                    "Operating Long-Term Liability",
                    "Financial Liability",
                ],
            }
        ],
    }
    wb = Workbook()
    wb.active.title = "Dummy"
    ws = wb.create_sheet(CHECK_CONTEXT_SHEET)
    ws.sheet_state = "hidden"
    ws["A1"] = CHECK_CONTEXT_MAGIC
    ws["A2"] = json.dumps(payload, separators=(",", ":"), sort_keys=True)
    path = tmp_path / "legacy_ak.xlsx"
    wb.save(path)
    wb.close()
    loaded = load_check_context(path)
    assert loaded.schema_version == 1
    assert loaded.normalization_bindings == ()
    assert len(loaded.judgment_bindings) == 1


# --- Step 8B2.1 integrity hardening -------------------------------------------


def _candidate(**overrides):
    base = {
        "selector": "concept:restructuring_expense",
        "referenceTreatment": "Non-recurring",
        "scope": SUPPORTED_NORMALIZATION_SCOPE,
        "topic": "t",
        "referenceRationale": "r",
        "consequenceNote": "c",
    }
    base.update(overrides)
    return base


def _find_earnings_row(ws, label: str) -> int:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"row not found: {label}")


@pytest.mark.parametrize(
    "label,column",
    [
        ("Restructuring expense", 3),  # detail source link (FY1)
        ("Reported NOPAT", 3),
        ("Reported Net Income", 3),
        ("Effective Tax Rate", 3),
        ("NORMALIZATION CHECK", 3),
    ],
)
def test_generated_earnings_normalization_formula_tamper_fails_closed(
    tmp_path, label, column
):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    practice = _fy2023(smap, "pretax_normalization_adjustment")
    wb = load_workbook(trainer_path, data_only=False)
    prow, pcol = parse_cell_ref(practice.cell)
    wb[practice.tab].cell(row=prow, column=pcol).value = practice.formula
    assert _fill_rgb(wb[practice.tab].cell(row=prow, column=pcol)) == "FFFF00"
    sheet = "Earnings Normalization"
    row = _find_earnings_row(wb[sheet], label)
    wb[sheet].cell(row=row, column=column).value = "=1+1"
    wb.save(trainer_path)
    wb.close()

    with pytest.raises(ValueError, match="Trusted workbook cell was modified"):
        check_workbook(trainer_path)

    wb = load_workbook(trainer_path, data_only=False)
    assert _fill_rgb(wb[practice.tab].cell(row=prow, column=pcol)) == "FFFF00"
    wb.close()


def test_learner_practice_formula_not_structurally_rejected(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    comp = _fy2023(smap, "after_tax_normalization_adjustment")
    _inject_formula_and_cached_value(
        trainer_path,
        comp.tab,
        comp.cell,
        formula=f"={float(comp.expected_value)}",
        cached_value=float(comp.expected_value),
    )
    summary = check_workbook(trainer_path)
    assert summary.correct == 1
    assert summary.incorrect == 0


def test_duplicate_normalization_candidates_rejected():
    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("Restructuring expense", -20, 0, concept="restructuring_expense"),
    )
    periods = canonical_fiscal_periods(fin)
    with pytest.raises(ValueError, match="duplicate normalization candidate"):
        normalization_cases(
            fin,
            periods,
            {
                "normalizationCandidates": [
                    _candidate(selector="concept:restructuring_expense"),
                    _candidate(selector="concept:restructuring_expense"),
                ]
            },
        )
    with pytest.raises(ValueError, match="duplicate normalization candidate"):
        normalization_cases(
            fin,
            periods,
            {
                "normalizationCandidates": [
                    _candidate(selector="concept:restructuring_expense"),
                    _candidate(selector="label:Restructuring expense"),
                ]
            },
        )


def test_label_selector_preserves_identity_with_shared_concept():
    from core.data.line_identity import line_identity
    from core.model.normalization import resolve_income_statement_identity

    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("Restructuring charge", -50, 0, concept="special_item"),
        _li("Litigation charge", -10, 0, concept="special_item"),
    )
    periods = canonical_fiscal_periods(fin)
    cases = normalization_cases(
        fin,
        periods,
        {
            "normalizationCandidates": [
                _candidate(selector="label:Restructuring charge")
            ]
        },
    )
    assert len(cases) == 1
    case = cases[0]
    assert case.override_selector == "label:Restructuring charge"
    assert case.line_identity == line_identity(
        resolve_income_statement_selector(fin, "label:Restructuring charge")
    ).key()
    assert "restructuring charge" in case.line_identity
    assert case.override_selector.startswith("label:")

    # Concept selector would be ambiguous; identity resolution stays on the label row.
    with pytest.raises(ValueError, match="matched 2"):
        resolve_income_statement_selector(fin, "concept:special_item")
    item = resolve_income_statement_identity(fin, case.line_identity)
    assert item.label == "Restructuring charge"
    assert float(item.values[periods[0]]) == pytest.approx(-50.0)
    # Adjustment uses only this row's signed amount: -(-50) = +50.
    assert -float(item.values[periods[0]]) == pytest.approx(50.0)


def test_resolve_income_statement_identity_requires_exact_match():
    from core.model.normalization import resolve_income_statement_identity

    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("Restructuring expense", -20, 0, concept="restructuring_expense"),
    )
    with pytest.raises(ValueError, match="matched no income-statement line"):
        resolve_income_statement_identity(fin, "concept:missing")


def test_typographic_apostrophe_label_selector_resolves():
    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("Director\u2019s fee", -8, -9),
    )
    periods = canonical_fiscal_periods(fin)
    item = resolve_income_statement_selector(fin, "label:Director's fee")
    assert item.label == "Director\u2019s fee"
    reverse = resolve_income_statement_selector(fin, "label:Director\u2019s fee")
    assert reverse.label == "Director\u2019s fee"
    cases = normalization_cases(
        fin,
        periods,
        {"normalizationCandidates": [_candidate(selector="label:Director's fee")]},
    )
    assert len(cases) == 1
    assert cases[0].override_selector == "label:Director\u2019s fee"


def test_whitespace_only_normalization_treatment_fails_closed(tmp_path):
    from core.trainer.semantic_io import answer_key_path_for

    trainer_path, _ = _build_norm_pair(tmp_path)
    answer_key_path = answer_key_path_for(trainer_path)
    smap = load_semantic_map(answer_key_path)
    comp = _latest(smap, "normalized_nopat")
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb["Normalization Judgment"].cell(5, 6).value = "   "
    wb.save(trainer_path)
    wb.close()

    with pytest.raises(ValueError, match="surrounding whitespace"):
        check_workbook(trainer_path)
    wb = load_workbook(trainer_path, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


def test_whitespace_only_classification_treatment_fails_closed(tmp_path):
    from core.trainer.workbook import build_training_workbook

    data = _ingest_demo()
    trainer_path, answer_key_path = build_training_workbook(
        data, tmp_path / "CLS_Trainer.xlsx"
    )
    smap = load_semantic_map(answer_key_path)
    comp = max(
        (c for c in smap.all_ordered() if c.family_id == "net_debt"),
        key=lambda c: c.period_index or 0,
    )
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb["Accounting Judgment"].cell(5, 6).value = "   "
    wb.save(trainer_path)
    wb.close()
    with pytest.raises(ValueError, match="surrounding whitespace"):
        check_workbook(trainer_path)
    wb = load_workbook(trainer_path, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


@pytest.mark.parametrize(
    "sheet,cell_value,match",
    [
        ("Normalization Judgment", " Non-recurring ", "surrounding whitespace"),
        ("Normalization Judgment", " Recurring ", "surrounding whitespace"),
        ("Accounting Judgment", " Financial Liability ", "surrounding whitespace"),
    ],
)
def test_padded_treatments_are_rejected(tmp_path, sheet, cell_value, match):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    comp = _fy2023(smap, "pretax_normalization_adjustment")
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb[sheet].cell(5, 6).value = cell_value
    wb.save(trainer_path)
    wb.close()
    with pytest.raises(ValueError, match=match):
        check_workbook(trainer_path)
    wb = load_workbook(trainer_path, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


def test_exact_valid_treatments_still_accepted(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    comp = _fy2023(smap, "pretax_normalization_adjustment")
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb["Normalization Judgment"].cell(5, 6).value = "Recurring"
    wb["Accounting Judgment"].cell(5, 6).value = "Financial Liability"
    wb.save(trainer_path)
    wb.close()

    ctx = load_check_context(answer_key_path)
    wb = load_workbook(trainer_path, data_only=False)
    treatments = normalization_treatments_for_check(wb, ctx)
    from core.trainer.check_context import classification_overrides_for_check

    overrides = classification_overrides_for_check(wb, ctx)
    wb.close()
    assert list(treatments.values()) == ["Recurring"]
    lease_sel = next(
        b.override_selector
        for b in ctx.judgment_bindings
        if "lease" in b.override_selector
    )
    assert overrides[lease_sel] == "Financial Liability"

    summary = check_workbook(trainer_path)
    assert summary.correct == 1
    assert summary.incorrect == 0


def test_blank_normalization_treatment_still_uses_reference(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    comp = _fy2023(smap, "pretax_normalization_adjustment")
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb["Normalization Judgment"].cell(5, 6).value = None
    wb.save(trainer_path)
    wb.close()
    summary = check_workbook(trainer_path)
    assert summary.correct == 1

    wb = load_workbook(trainer_path, data_only=False)
    wb["Normalization Judgment"].cell(5, 6).value = ""
    wb.save(trainer_path)
    wb.close()
    summary = check_workbook(trainer_path)
    assert summary.correct == 1


def test_source_value_tamper_fails_before_exact_formula_green(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    comp = max(
        (c for c in smap.all_ordered() if c.family_id == "revenue_link"),
        key=lambda c: c.period_index or 0,
    )
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    # Tamper an Income Statement source amount for the same period column.
    src_col = col  # Condensed revenue uses IS columns starting at B
    # Find revenue row on Income Statement and alter latest FY value.
    is_ws = wb["Income Statement"]
    rev_row = None
    for r in range(7, (is_ws.max_row or 7) + 1):
        if is_ws.cell(r, 1).value == "Revenue":
            rev_row = r
            break
    assert rev_row is not None
    # Period columns on IS are B..; match practice cell's relative period index.
    period_col = 2 + (comp.period_index or 0)
    is_ws.cell(rev_row, period_col).value = 999999
    wb.save(trainer_path)
    wb.close()

    with pytest.raises(ValueError, match="Trusted workbook cell was modified: Income Statement"):
        check_workbook(trainer_path)
    wb = load_workbook(trainer_path, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


def test_fixed_classification_tamper_fails_before_green(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    comp = max(
        (c for c in smap.all_ordered() if c.family_id == "net_debt"),
        key=lambda c: c.period_index or 0,
    )
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    # Change a non-judgment fixed classification (not Operating lease liabilities).
    cf = wb["Condensed Financials"]
    target_row = None
    for r in range(1, (cf.max_row or 1) + 1):
        label = cf.cell(r, 1).value
        cat = cf.cell(r, 2).value
        if label == "Trade receivables" and isinstance(cat, str) and not cat.startswith("="):
            target_row = r
            break
    assert target_row is not None
    cf.cell(target_row, 2).value = "Financial Asset"
    wb.save(trainer_path)
    wb.close()

    with pytest.raises(ValueError, match="Trusted workbook cell was modified: Condensed Financials"):
        check_workbook(trainer_path)
    wb = load_workbook(trainer_path, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


def test_base_build_source_tamper_fails_closed(tmp_path):
    from core.trainer.workbook import build_training_workbook

    data = _ingest_demo()
    trainer_path, answer_key_path = build_training_workbook(
        data, tmp_path / "BASE_Trainer.xlsx"
    )
    smap = load_semantic_map(answer_key_path)
    assert len(smap.all_ordered()) == 312
    comp = max(
        (c for c in smap.all_ordered() if c.family_id == "revenue_link"),
        key=lambda c: c.period_index or 0,
    )
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    is_ws = wb["Income Statement"]
    rev_row = next(
        r for r in range(7, (is_ws.max_row or 7) + 1) if is_ws.cell(r, 1).value == "Revenue"
    )
    is_ws.cell(rev_row, 2 + (comp.period_index or 0)).value = 1
    wb.save(trainer_path)
    wb.close()
    with pytest.raises(ValueError, match="Trusted workbook cell was modified: Income Statement"):
        check_workbook(trainer_path)
    wb = load_workbook(trainer_path, data_only=False)
    assert _fill_rgb(wb[comp.tab].cell(row=row, column=col)) == "FFFF00"
    wb.close()


def test_judgment_gh_edits_do_not_fail_trusted_validation(tmp_path):
    trainer_path, answer_key_path = _build_norm_pair(tmp_path)
    smap = load_semantic_map(answer_key_path)
    comp = _fy2023(smap, "pretax_normalization_adjustment")
    wb = load_workbook(trainer_path, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb["Normalization Judgment"].cell(5, 6).value = "Non-recurring"
    wb["Normalization Judgment"].cell(5, 7).value = "Learner rationale text"
    wb["Normalization Judgment"].cell(5, 8).value = "Learner consequence text"
    wb["Accounting Judgment"].cell(5, 7).value = "Lease rationale"
    wb["Accounting Judgment"].cell(5, 8).value = "Lease consequence"
    wb.save(trainer_path)
    wb.close()
    summary = check_workbook(trainer_path)
    assert summary.correct == 1
    assert summary.incorrect == 0


def test_normalization_candidate_period_completeness_required():
    from core.model.normalization import NormalizationCase
    from core.model.source_values import MissingHistoricalValueError
    from core.data.line_identity import line_identity

    fin = _tiny_fin(
        _li("Revenue", 100, 110, concept="revenue"),
        _li("One-off charge", -20, -10, concept="one_off"),
    )
    periods = canonical_fiscal_periods(fin)
    candidate = {
        "selector": "concept:one_off",
        "referenceTreatment": "Non-recurring",
        "scope": SUPPORTED_NORMALIZATION_SCOPE,
        "topic": "charge",
        "referenceRationale": "r",
        "consequenceNote": "c",
    }
    charge = next(i for i in fin.income_statement if i.concept == "one_off")
    del charge.values[periods[1]]
    with pytest.raises(MissingHistoricalValueError, match="normalization candidate"):
        normalization_cases(fin, periods, {"normalizationCandidates": [candidate]})

    charge.values[periods[1]] = None
    with pytest.raises(MissingHistoricalValueError, match="normalization candidate"):
        normalization_cases(fin, periods, {"normalizationCandidates": [candidate]})

    charge.values[periods[0]] = 0.0
    charge.values[periods[1]] = 0.0
    assert normalization_cases(fin, periods, {"normalizationCandidates": [candidate]}) == ()

    charge.values[periods[0]] = -20.0
    charge.values[periods[1]] = 0.0
    cases = normalization_cases(fin, periods, {"normalizationCandidates": [candidate]})
    assert len(cases) == 1

    d1, d2 = periods
    full = StandardizedFinancials(
        company_name="Tiny Co",
        ticker="TINY",
        currency="HKD",
        jurisdiction="HK",
        units="HKD millions",
        periods=[
            FinancialPeriod(end_date=d1, label="FY2021"),
            FinancialPeriod(end_date=d2, label="FY2022"),
        ],
        income_statement=[
            LineItem(label="Revenue", values={d1: 100, d2: 110}, concept="revenue"),
            LineItem(label="Finance costs", values={d1: -1, d2: -1}),
            LineItem(label="Finance income", values={d1: 0, d2: 0}),
            LineItem(label="Profit before tax", values={d1: 50, d2: 55}),
            LineItem(label="Income tax expense", values={d1: -5, d2: -6}),
            LineItem(label="Profit for the year", values={d1: 45, d2: 49}),
            LineItem(
                label="One-off charge",
                values={d1: -20.0, d2: None},
                concept="one_off",
            ),
        ],
        balance_sheet=[
            LineItem(label="Cash", values={d1: 10, d2: 12}, concept="cash"),
            LineItem(
                label="Total equity", values={d1: 10, d2: 12}, concept="total_equity"
            ),
        ],
        cash_flow=[
            LineItem(
                label="Net cash from operating activities",
                values={d1: 1, d2: 2},
            )
        ],
    )
    item = next(i for i in full.income_statement if i.concept == "one_off")
    identity = line_identity(item).key()
    manual = NormalizationCase(
        id=f"normalization::{identity}",
        order=1,
        line_identity=identity,
        override_selector="concept:one_off",
        label="One-off charge",
        scope=SUPPORTED_NORMALIZATION_SCOPE,
        topic="charge",
        reference_treatment="Non-recurring",
        alternatives=("Recurring",),
        model_rationale="r",
        consequence_prompt="p",
        model_consequence="c",
    )
    anchor = compute_anchor(full, periods)
    with pytest.raises(MissingHistoricalValueError, match="normalization candidate"):
        compute_normalization_series(full, periods, anchor, (manual,))


def test_normalization_undefined_etr_tax_effect_and_check(tmp_path):
    from core.engine.reference_model import EARNINGS_NORMALIZATION_SHEET
    from core.model.ratio_values import UNDEFINED_RATIO

    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
    fin = StandardizedFinancials(
        ticker="NORMETR",
        company_name="Norm ETR Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=[
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ],
        income_statement=[
            LineItem(label="Revenue", values={d1: 1000, d2: 1100}),
            LineItem(label="Finance costs", values={d1: 0, d2: 0}),
            LineItem(label="Finance income", values={d1: 0, d2: 0}),
            LineItem(label="Profit before tax", values={d1: 0, d2: 0}),
            LineItem(label="Income tax expense", values={d1: 0, d2: 0}),
            LineItem(label="Profit for the year", values={d1: 80, d2: 90}),
            LineItem(
                label="Restructuring expense",
                values={d1: 0.0, d2: -100.0},
                concept="restructuring_expense",
            ),
        ],
        balance_sheet=[
            LineItem(label="Cash and cash equivalents", values={d1: 100, d2: 110}),
            LineItem(label="Trade receivables", values={d1: 80, d2: 90}),
            LineItem(label="Property, plant and equipment", values={d1: 400, d2: 420}),
            LineItem(label="Trade payables", values={d1: 50, d2: 55}),
            LineItem(label="Bank borrowings", values={d1: 200, d2: 210}),
            LineItem(label="Total equity", values={d1: 330, d2: 355}),
        ],
        cash_flow=[
            LineItem(
                label="Net cash from operating activities",
                values={d1: 50, d2: 60},
            )
        ],
    )
    periods = [d1, d2]
    assumptions = {
        "normalizationCandidates": [
            {
                "selector": "concept:restructuring_expense",
                "referenceTreatment": "Non-recurring",
                "scope": SUPPORTED_NORMALIZATION_SCOPE,
                "topic": "restructure",
                "referenceRationale": "one-off",
                "consequenceNote": "normalize",
            }
        ]
    }
    cases = normalization_cases(fin, periods, assumptions)
    assert len(cases) == 1
    anchor = compute_anchor(fin, periods)
    assert anchor.historical.effective_tax_rate == [UNDEFINED_RATIO, UNDEFINED_RATIO]
    series = compute_normalization_series(fin, periods, anchor, cases)
    assert series.pretax_adjustment[0] == pytest.approx(0.0)
    assert series.after_tax_adjustment[0] == pytest.approx(0.0)
    assert series.normalized_net_income[0] == pytest.approx(80.0)
    assert series.normalized_nopat[0] == pytest.approx(float(anchor.historical.nopat[0]))
    assert series.pretax_adjustment[1] == pytest.approx(100.0)
    assert series.after_tax_adjustment[1] == UNDEFINED_RATIO
    assert series.normalized_net_income[1] == UNDEFINED_RATIO
    assert series.normalized_nopat[1] == UNDEFINED_RATIO

    trainer, answer = build_training_workbook(
        fin, tmp_path / "NormEtr_Trainer.xlsx", assumptions
    )
    smap = load_semantic_map(answer)
    after = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "after_tax_normalization_adjustment" and c.period_index == 1
    )
    assert after.expected_value == UNDEFINED_RATIO
    assert "ISNA(" in after.formula
    assert after.formula.startswith("=IF(")
    norm_ni = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "normalized_net_income" and c.period_index == 1
    )
    assert norm_ni.expected_value == UNDEFINED_RATIO

    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(after.cell)
    wb[after.tab].cell(row=row, column=col).value = after.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer,
        after.tab,
        after.cell,
        formula="=NA()",
        cached_value=UNDEFINED_RATIO,
    )
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer,
        after.tab,
        after.cell,
        formula="=0",
        cached_value=0.0,
    )
    assert check_workbook(trainer).incorrect == 1

    _, demo_answer = _build_norm_pair(tmp_path / "demo_norm")
    demo_smap = load_semantic_map(demo_answer)
    demo_after = _fy2023(demo_smap, "after_tax_normalization_adjustment")
    assert "ISNA(" in demo_after.formula
    assert "=IF(" in demo_after.formula
    assert demo_after.tab == EARNINGS_NORMALIZATION_SHEET

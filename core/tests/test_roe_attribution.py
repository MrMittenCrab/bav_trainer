"""Step 9D.1 — historical ROE operating / financing attribution."""

from __future__ import annotations

from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.engine.component_catalog import (
    PROFITABILITY_CHANGE_COMPONENT_CATALOG,
    PROFITABILITY_DRIVER_COMPONENT_CATALOG,
    ROE_ATTRIBUTION_COMPONENT_CATALOG,
    WORKING_CAPITAL_COMPONENT_CATALOG,
    expand_roe_attribution_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.financial_math import compute_anchor
from core.model.historical_expected import (
    expected_value_for_component,
    roe_attribution_expected_series,
)
from core.model.profitability_change import ProfitabilityChangeSeries
from core.model.roe_attribution import compute_roe_attribution_series
from core.model.ratio_values import UNDEFINED_RATIO
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family
from core.tests.test_normalization import _inject_formula_and_cached_value

ROOT = Path(__file__).resolve().parents[2]
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"
DEMO_ASSUMPTIONS = ROOT / "example" / "DEMO_HK_Assumptions.json"


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _ingest_demo():
    adapter = HKManualDocumentAdapter()
    return adapter.ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values_by_date, concept=""):
    return LineItem(label=label, values=values_by_date, concept=concept)


def _pc_series(*, rnoa_change, n):
    none_pad = [None] * n
    def pad(seq):
        out = list(none_pad)
        for i, v in enumerate(seq):
            if i < n:
                out[i] = v
        return tuple(out)

    return ProfitabilityChangeSeries(
        nopat_margin_change=pad([None] * n),
        noa_turnover_change=pad([None] * n),
        rnoa_change=pad(rnoa_change),
        margin_effect_on_rnoa=pad([None] * n),
        turnover_effect_on_rnoa=pad([None] * n),
        rnoa_change_from_drivers=pad([None] * n),
    )


def _anchor_roe(*, rnoa, flev, spread, roe):
    return SimpleNamespace(
        dupont={
            "RNOA": list(rnoa),
            "FLEV": list(flev),
            "Spread": list(spread),
            "ROE (decomposed)": list(roe),
        }
    )


def test_ordinary_midpoint_roe_attribution():
    # Indices 0 unused / N/A; 1 prior comparable; 2 current
    rnoa = (None, 0.12, 0.15)
    flev = (None, 0.40, 0.50)
    spread = (None, 0.04, 0.06)
    roe = (None, 0.136, 0.180)
    rnoa_change = (None, None, 0.03)
    anchor = _anchor_roe(rnoa=rnoa, flev=flev, spread=spread, roe=roe)
    pc = _pc_series(rnoa_change=rnoa_change, n=3)
    with patch(
        "core.model.roe_attribution.compute_profitability_change_series",
        return_value=pc,
    ):
        series = compute_roe_attribution_series(anchor)

    assert series.financing_contribution_to_roe[0] is None
    assert series.financing_contribution_to_roe[1] == pytest.approx(0.016)
    assert series.financing_contribution_to_roe[2] == pytest.approx(0.030)

    for field in (
        series.roe_change,
        series.operating_effect_on_roe_change,
        series.leverage_effect_on_roe_change,
        series.spread_effect_on_roe_change,
        series.financing_effect_on_roe_change,
        series.roe_change_from_drivers,
    ):
        assert field[0] is None
        assert field[1] is None

    assert series.roe_change[2] == pytest.approx(0.044)
    assert series.operating_effect_on_roe_change[2] == pytest.approx(0.030)
    assert series.leverage_effect_on_roe_change[2] == pytest.approx(0.005)
    assert series.spread_effect_on_roe_change[2] == pytest.approx(0.009)
    assert series.financing_effect_on_roe_change[2] == pytest.approx(0.014)
    assert series.roe_change_from_drivers[2] == pytest.approx(0.044)


def test_roe_attribution_edge_cases():
    # Zero FLEV / Spread -> numeric zero contribution; 0 × #N/A stays #N/A
    anchor0 = _anchor_roe(
        rnoa=(None, 0.10, 0.10),
        flev=(None, 0.0, 0.0),
        spread=(None, 0.05, 0.05),
        roe=(None, 0.10, 0.10),
    )
    pc0 = _pc_series(rnoa_change=(None, None, 0.0), n=3)
    with patch(
        "core.model.roe_attribution.compute_profitability_change_series",
        return_value=pc0,
    ):
        s0 = compute_roe_attribution_series(anchor0)
    assert s0.financing_contribution_to_roe[1] == pytest.approx(0.0)
    assert s0.leverage_effect_on_roe_change[2] == pytest.approx(0.0)
    assert s0.spread_effect_on_roe_change[2] == pytest.approx(0.0)

    anchor_s0 = _anchor_roe(
        rnoa=(None, 0.10, 0.10),
        flev=(None, 0.5, 0.5),
        spread=(None, 0.0, 0.0),
        roe=(None, 0.10, 0.10),
    )
    with patch(
        "core.model.roe_attribution.compute_profitability_change_series",
        return_value=pc0,
    ):
        ss0 = compute_roe_attribution_series(anchor_s0)
    assert ss0.financing_contribution_to_roe[1] == pytest.approx(0.0)

    anchor_na = _anchor_roe(
        rnoa=(None, 0.10, 0.12),
        flev=(None, 0.0, UNDEFINED_RATIO),
        spread=(None, 0.05, 0.05),
        roe=(None, 0.10, UNDEFINED_RATIO),
    )
    pc_na = _pc_series(rnoa_change=(None, None, 0.02), n=3)
    with patch(
        "core.model.roe_attribution.compute_profitability_change_series",
        return_value=pc_na,
    ):
        sna = compute_roe_attribution_series(anchor_na)
    assert sna.financing_contribution_to_roe[2] == UNDEFINED_RATIO
    assert sna.financing_effect_on_roe_change[2] == UNDEFINED_RATIO
    assert sna.roe_change_from_drivers[2] == UNDEFINED_RATIO

    # Undefined operating -> driver #N/A
    anchor_op = _anchor_roe(
        rnoa=(None, 0.12, 0.15),
        flev=(None, 0.4, 0.5),
        spread=(None, 0.04, 0.06),
        roe=(None, 0.136, 0.180),
    )
    pc_op = _pc_series(rnoa_change=(None, None, UNDEFINED_RATIO), n=3)
    with patch(
        "core.model.roe_attribution.compute_profitability_change_series",
        return_value=pc_op,
    ):
        sop = compute_roe_attribution_series(anchor_op)
    assert sop.operating_effect_on_roe_change[2] == UNDEFINED_RATIO
    assert sop.roe_change_from_drivers[2] == UNDEFINED_RATIO

    # Negative signed effects
    anchor_n = _anchor_roe(
        rnoa=(None, 0.15, 0.12),
        flev=(None, 0.50, 0.40),
        spread=(None, 0.06, 0.04),
        roe=(None, 0.180, 0.136),
    )
    pc_n = _pc_series(rnoa_change=(None, None, -0.03), n=3)
    with patch(
        "core.model.roe_attribution.compute_profitability_change_series",
        return_value=pc_n,
    ):
        sn = compute_roe_attribution_series(anchor_n)
    assert sn.leverage_effect_on_roe_change[2] == pytest.approx(-0.005)
    assert sn.spread_effect_on_roe_change[2] == pytest.approx(-0.009)
    assert sn.roe_change_from_drivers[2] == pytest.approx(-0.044)

    # Mismatched direct ROE
    bad = _anchor_roe(
        rnoa=(None, 0.12, 0.15),
        flev=(None, 0.40, 0.50),
        spread=(None, 0.04, 0.06),
        roe=(None, 0.136, 0.99),
    )
    with patch(
        "core.model.roe_attribution.compute_profitability_change_series",
        return_value=_pc_series(rnoa_change=(None, None, 0.03), n=3),
    ):
        with pytest.raises(ValueError, match="level bridge does not reconcile"):
            compute_roe_attribution_series(bad)

    # Length mismatch
    bad_len = SimpleNamespace(
        dupont={
            "RNOA": [None, 0.12, 0.15],
            "FLEV": [None, 0.4],
            "Spread": [None, 0.04, 0.06],
            "ROE (decomposed)": [None, 0.136, 0.180],
        }
    )
    with patch(
        "core.model.roe_attribution.compute_profitability_change_series",
        return_value=_pc_series(rnoa_change=(None, None, 0.03), n=3),
    ):
        with pytest.raises(ValueError, match="length mismatch"):
            compute_roe_attribution_series(bad_len)


def test_catalog_expand_and_expected_keys():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(ROE_ATTRIBUTION_COMPONENT_CATALOG) == 7
    assert len({f.id for f in ROE_ATTRIBUTION_COMPONENT_CATALOG}) == 7
    assert [f.order for f in ROE_ATTRIBUTION_COMPONENT_CATALOG] == list(range(56, 63))
    scopes = {f.id: f.period_scope for f in ROE_ATTRIBUTION_COMPONENT_CATALOG}
    assert scopes["financing_contribution_to_roe"] == "comparable"
    assert all(
        scopes[fid] == "post_comparable"
        for fid in scopes
        if fid != "financing_contribution_to_roe"
    )
    assert len(PROFITABILITY_DRIVER_COMPONENT_CATALOG) == 4
    assert len(PROFITABILITY_CHANGE_COMPONENT_CATALOG) == 6
    assert len(WORKING_CAPITAL_COMPONENT_CATALOG) == 11

    specs = expand_roe_attribution_specs(periods, start_order=223)
    assert len(specs) == 22

    data = _ingest_demo()
    builder = ReferenceModelBuilder(data)
    assert len(builder.roe_attribution_specs) == 22
    series = roe_attribution_expected_series(builder.anchor)
    assert set(series) == {f.id for f in ROE_ATTRIBUTION_COMPONENT_CATALOG}
    for j in range(2, len(builder.periods)):
        driver = series["roe_change_from_drivers"][j]
        direct = series["roe_change"][j]
        if driver != UNDEFINED_RATIO:
            assert driver == pytest.approx(float(direct))


def test_demo_roe_attribution_surface(tmp_path):
    import json

    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "BASE_ROE.xlsx")
    smap = load_semantic_map(answer)
    assert len(smap.all_ordered()) == 312
    assert len(group_components_by_family(smap)) == 74
    ra = [c for c in smap.all_ordered() if c.category == "roe_attribution"]
    assert len(ra) == 22
    assert all(c.tab == "ALT DuPont" for c in ra)

    wb = load_workbook(answer)
    ws = wb["ALT DuPont"]
    assert ws.cell(21, 1).value == "RNOA CHANGE ATTRIBUTION"
    assert ws.cell(28, 1).value == "RNOA CHANGE DRIVER CHECK"
    assert ws.cell(30, 1).value == "ROE OPERATING / FINANCING ATTRIBUTION"
    assert ws.cell(31, 1).value == "Financing Contribution to ROE"
    assert ws.cell(32, 1).value == "ROE LEVEL ATTRIBUTION CHECK"
    assert ws.cell(34, 1).value == "ROE CHANGE ATTRIBUTION"
    assert ws.cell(41, 1).value == "ROE CHANGE ATTRIBUTION CHECK"
    assert ws.cell(31, 2).value == "N/A"
    assert ws.cell(35, 2).value == "N/A"
    assert ws.cell(35, 3).value == "N/A"
    practice = {(c.tab, c.cell) for c in smap.all_ordered()}
    assert ("ALT DuPont", "C32") not in practice
    assert ("ALT DuPont", "D41") not in practice

    fin_c = next(
        c
        for c in ra
        if c.family_id == "financing_contribution_to_roe" and c.period_index == 1
    )
    assert "*" in fin_c.formula
    assert "IFERROR" not in fin_c.formula.upper()
    op = next(
        c
        for c in ra
        if c.family_id == "operating_effect_on_roe_change" and c.period_index == 2
    )
    assert "ABS" not in op.formula.upper()
    lev = next(
        c
        for c in ra
        if c.family_id == "leverage_effect_on_roe_change" and c.period_index == 2
    )
    assert "/2)" in lev.formula
    assert "ABS" not in lev.formula.upper()

    wb_t = load_workbook(trainer)
    for comp in ra:
        row, col = parse_cell_ref(comp.cell)
        cell = wb_t[comp.tab].cell(row=row, column=col)
        assert cell.value is None
        assert _fill_rgb(cell) == "FFFF00"
    wb.close()
    wb_t.close()

    summary = check_workbook(trainer)
    assert summary.total == 312
    assert summary.blank == 312

    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer_n, answer_n = build_training_workbook(
        data, tmp_path / "NORM_ROE.xlsx", assumptions
    )
    smap_n = load_semantic_map(answer_n)
    assert len(smap_n.all_ordered()) == 332
    assert len(group_components_by_family(smap_n)) == 78
    assert check_workbook(trainer_n).blank == 332


def test_live_classification_changes_roe_attribution(tmp_path):
    d1, d2, d3 = date(2023, 12, 31), date(2024, 12, 31), date(2025, 12, 31)
    periods = [
        FinancialPeriod(end_date=d1, label="FY2023"),
        FinancialPeriod(end_date=d2, label="FY2024"),
        FinancialPeriod(end_date=d3, label="FY2025"),
    ]

    def vals(a, b, c):
        return {d1: a, d2: b, d3: c}

    fin = StandardizedFinancials(
        ticker="STI3",
        company_name="STI ROE Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=[
            _li("Revenue", vals(1000, 1100, 1200)),
            _li("Finance costs", vals(-10, -11, -12)),
            _li("Finance income", vals(0, 0, 0)),
            _li("Profit before tax", vals(200, 220, 240)),
            _li("Income tax expense", vals(-30, -33, -36)),
            _li("Profit for the year", vals(170, 187, 204)),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", vals(50, 55, 60)),
            _li("Trade receivables", vals(40, 45, 50)),
            _li("Short-term investments", vals(100, 120, 140)),
            _li("Property, plant and equipment", vals(400, 420, 440)),
            _li("Trade payables", vals(30, 35, 40)),
            _li("Bank borrowings", vals(200, 210, 220)),
            _li("Total equity", vals(360, 395, 430)),
        ],
        cash_flow=[_li("Net cash from operating activities", vals(50, 60, 70))],
    )
    trainer, answer = build_training_workbook(fin, tmp_path / "STI_ROE.xlsx")
    smap = load_semantic_map(answer)
    fin_contrib = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "financing_contribution_to_roe" and c.period_index == 2
    )
    lev = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "leverage_effect_on_roe_change" and c.period_index == 2
    )
    fin_eff = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "financing_effect_on_roe_change" and c.period_index == 2
    )
    driver = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "roe_change_from_drivers" and c.period_index == 2
    )

    ref = float(fin_contrib.expected_value)
    alt_anchor = compute_anchor(
        fin,
        [d1, d2, d3],
        classification_overrides={
            "label:Short-term investments": "Operating Working Capital Asset"
        },
    )
    alt_fin = expected_value_for_component(alt_anchor, fin_contrib)
    alt_lev = expected_value_for_component(alt_anchor, lev)
    alt_fin_eff = expected_value_for_component(alt_anchor, fin_eff)
    alt_driver = expected_value_for_component(alt_anchor, driver)
    assert isinstance(alt_fin, float)
    assert alt_fin != pytest.approx(ref)
    assert isinstance(alt_lev, float)
    assert isinstance(alt_fin_eff, float)
    assert isinstance(alt_driver, float)

    wb = load_workbook(trainer, data_only=False)
    wb["Accounting Judgment"].cell(5, 6).value = "Operating Working Capital Asset"
    wb.save(trainer)
    wb.close()

    _inject_formula_and_cached_value(
        trainer,
        fin_contrib.tab,
        fin_contrib.cell,
        formula="=C10*C9",
        cached_value=ref,
    )
    assert check_workbook(trainer).incorrect == 1
    _inject_formula_and_cached_value(
        trainer,
        fin_contrib.tab,
        fin_contrib.cell,
        formula="=C10*C9",
        cached_value=float(alt_fin),
    )
    assert check_workbook(trainer).correct == 1


def test_roe_attribution_check_tamper_and_undefined(tmp_path):
    data = _ingest_demo()
    trainer, answer = build_training_workbook(data, tmp_path / "ROE_Check.xlsx")
    smap = load_semantic_map(answer)
    comp = next(
        c
        for c in smap.all_ordered()
        if c.family_id == "financing_contribution_to_roe" and c.period_index == 1
    )
    expected = float(comp.expected_value)

    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(comp.cell)
    wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    assert check_workbook(trainer).correct == 1

    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula="=0.01+0", cached_value=expected
    )
    assert check_workbook(trainer).correct == 1
    _inject_formula_and_cached_value(
        trainer, comp.tab, comp.cell, formula="=0.01+0", cached_value=expected + 1.0
    )
    assert check_workbook(trainer).incorrect == 1

    for check_label in (
        "ROE LEVEL ATTRIBUTION CHECK",
        "ROE CHANGE ATTRIBUTION CHECK",
    ):
        trainer_t, _ = build_training_workbook(data, tmp_path / f"ROE_Tamper_{check_label}.xlsx")
        smap_t = load_semantic_map(answer)
        comp_t = next(
            c
            for c in smap_t.all_ordered()
            if c.family_id == "financing_contribution_to_roe" and c.period_index == 1
        )
        wb = load_workbook(trainer_t, data_only=False)
        row, col = parse_cell_ref(comp_t.cell)
        wb[comp_t.tab].cell(row=row, column=col).value = comp_t.formula
        ws = wb["ALT DuPont"]
        check_row = next(
            r
            for r in range(1, (ws.max_row or 1) + 1)
            if ws.cell(r, 1).value == check_label
        )
        # Col C for comparable period index 1; col D for change checks
        tamper_col = 3 if "LEVEL" in check_label else 4
        ws.cell(check_row, tamper_col).value = '="TAMPERED"'
        wb.save(trainer_t)
        wb.close()
        with pytest.raises(
            ValueError, match="Trusted workbook cell was modified: ALT DuPont"
        ):
            check_workbook(trainer_t)
        wb = load_workbook(trainer_t, data_only=False)
        assert _fill_rgb(wb[comp_t.tab].cell(row=row, column=col)) == "FFFF00"
        wb.close()

    # Zero average equity -> FLEV #N/A -> financing contribution #N/A
    d1, d2 = date(2024, 12, 31), date(2025, 12, 31)
    periods = [
        FinancialPeriod(end_date=d1, label="FY2024"),
        FinancialPeriod(end_date=d2, label="FY2025"),
    ]

    def v(a, b):
        return {d1: a, d2: b}

    fin = StandardizedFinancials(
        ticker="EQ0",
        company_name="Zero Equity Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=[
            _li("Revenue", v(1000, 1100)),
            _li("Finance costs", v(-10, -11)),
            _li("Finance income", v(0, 0)),
            _li("Profit before tax", v(200, 220)),
            _li("Income tax expense", v(-30, -33)),
            _li("Profit for the year", v(170, 187)),
        ],
        balance_sheet=[
            _li("Cash and cash equivalents", v(100, 110)),
            _li("Trade receivables", v(80, 90)),
            _li("Property, plant and equipment", v(400, 420)),
            _li("Trade payables", v(50, 55)),
            _li("Bank borrowings", v(530, 565)),
            _li("Total equity", v(0, 0)),
        ],
        cash_flow=[_li("Net cash from operating activities", v(50, 60))],
    )
    trainer2, answer2 = build_training_workbook(fin, tmp_path / "ROE_NA.xlsx")
    smap2 = load_semantic_map(answer2)
    fin_c = next(
        c
        for c in smap2.all_ordered()
        if c.family_id == "financing_contribution_to_roe" and c.period_index == 1
    )
    assert fin_c.expected_value == UNDEFINED_RATIO

    wb = load_workbook(trainer2, data_only=False)
    row, col = parse_cell_ref(fin_c.cell)
    wb[fin_c.tab].cell(row=row, column=col).value = fin_c.formula
    wb.save(trainer2)
    wb.close()
    assert check_workbook(trainer2).correct == 1
    _inject_formula_and_cached_value(
        trainer2, fin_c.tab, fin_c.cell, formula="=NA()", cached_value=UNDEFINED_RATIO
    )
    assert check_workbook(trainer2).correct == 1
    _inject_formula_and_cached_value(
        trainer2, fin_c.tab, fin_c.cell, formula="=0", cached_value=0.0
    )
    assert check_workbook(trainer2).incorrect == 1

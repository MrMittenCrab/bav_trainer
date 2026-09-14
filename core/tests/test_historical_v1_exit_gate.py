"""Step 9H.1 / 9N.2 — historical active-catalog freeze (release contract)."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import DocumentManifest, DocumentType
from core.engine.component_catalog import (
    CAPEX_COMPONENT_CATALOG,
    COMPONENT_CATALOG,
    DEFERRED_COMPONENT_SPECS,
    DEFERRED_TAX_COMPONENT_CATALOG,
    FIXED_ASSET_COMPONENT_CATALOG,
    GOODWILL_INTANGIBLES_COMPONENT_CATALOG,
    LEASE_LIABILITY_COMPONENT_CATALOG,
    LEASE_ROU_COMPONENT_CATALOG,
    NORMALIZATION_COMPONENT_CATALOG,
    NORMALIZED_PER_SHARE_COMPONENT_CATALOG,
    OWNERSHIP_ATTRIBUTION_COMPONENT_CATALOG,
    PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG,
    PER_SHARE_COMPONENT_CATALOG,
    PROFITABILITY_CHANGE_COMPONENT_CATALOG,
    PROFITABILITY_DRIVER_COMPONENT_CATALOG,
    QUALITY_CHANGE_COMPONENT_CATALOG,
    QUALITY_COMPONENT_CATALOG,
    ROE_ATTRIBUTION_COMPONENT_CATALOG,
    WORKING_CAPITAL_COMPONENT_CATALOG,
)
from core.engine.reference_model import DEFERRED_PLACEHOLDER, DEFERRED_TAB_NAMES
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.tests.cross_company_fixtures import (
    asset_light_services_case,
    capital_intensive_manufacturer_case,
    inventory_retail_case,
)
from core.tests.test_per_share import DEMO_ASSUMPTIONS, DEMO_JSON
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]

# Step 9 historical active namespace: family orders 1–121 (through capex).
ACTIVE_CATALOGS = (
    COMPONENT_CATALOG,
    NORMALIZATION_COMPONENT_CATALOG,
    QUALITY_COMPONENT_CATALOG,
    WORKING_CAPITAL_COMPONENT_CATALOG,
    PROFITABILITY_DRIVER_COMPONENT_CATALOG,
    PROFITABILITY_CHANGE_COMPONENT_CATALOG,
    ROE_ATTRIBUTION_COMPONENT_CATALOG,
    QUALITY_CHANGE_COMPONENT_CATALOG,
    PER_SHARE_COMPONENT_CATALOG,
    PER_SHARE_ATTRIBUTION_COMPONENT_CATALOG,
    NORMALIZED_PER_SHARE_COMPONENT_CATALOG,
    FIXED_ASSET_COMPONENT_CATALOG,
    LEASE_LIABILITY_COMPONENT_CATALOG,
    OWNERSHIP_ATTRIBUTION_COMPONENT_CATALOG,
    GOODWILL_INTANGIBLES_COMPONENT_CATALOG,
    LEASE_ROU_COMPONENT_CATALOG,
    DEFERRED_TAX_COMPONENT_CATALOG,
    CAPEX_COMPONENT_CATALOG,
)

# Explicit post–9M optional catalog family id → order mappings (98–121).
EXPECTED_POST_V1_FAMILY_ORDERS = {
    # Goodwill / intangibles 98–111
    "goodwill_change": 98,
    "goodwill_growth": 99,
    "average_goodwill": 100,
    "goodwill_to_revenue": 101,
    "intangible_assets_change": 102,
    "intangible_assets_growth": 103,
    "average_intangible_assets": 104,
    "intangible_assets_to_revenue": 105,
    "goodwill_and_intangibles_change": 106,
    "goodwill_and_intangibles_growth": 107,
    "average_goodwill_and_intangibles": 108,
    "goodwill_and_intangibles_to_revenue": 109,
    "intangible_payments": 110,
    "intangible_payments_to_revenue": 111,
    # Lease ROU 112–115
    "rou_assets_change": 112,
    "rou_assets_growth": 113,
    "average_rou_assets": 114,
    "rou_assets_to_revenue": 115,
    # Deferred tax 116–119
    "net_deferred_tax_position": 116,
    "deferred_tax_assets_change": 117,
    "deferred_tax_liabilities_change": 118,
    "net_deferred_tax_position_change": 119,
    # Capex 120–121
    "ppe_capex": 120,
    "ppe_capex_to_revenue": 121,
}


def _fill_rgb(cell) -> str:
    fill = cell.fill
    if not fill or fill.fill_type != "solid":
        return ""
    color = fill.fgColor.rgb or fill.start_color.rgb or ""
    return str(color).upper().lstrip("0")[-6:] if color else ""


def _forecast_must_not_run(*args, **kwargs):
    raise AssertionError("dormant forecast/valuation engine executed during historical build")


def _ingest_demo():
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _assert_deferred_isolation(trainer: Path, answer: Path) -> None:
    for path in (trainer, answer):
        wb = load_workbook(path)
        for name in DEFERRED_TAB_NAMES:
            assert name in wb.sheetnames
            assert wb[name].sheet_state == "hidden"
            assert wb[name]["A1"].value == DEFERRED_PLACEHOLDER
        wb.close()

    smap = load_semantic_map(answer)
    assert not any(
        comp.tab in DEFERRED_TAB_NAMES or comp.category in {"forecasting", "valuation"}
        for comp in smap.all_ordered()
    )


def test_historical_v1_active_catalog_namespace_is_frozen():
    """Freeze the current Step 9 historical active catalog namespace (orders 1–121)."""
    families = [family for catalog in ACTIVE_CATALOGS for family in catalog]
    ids = [family.id for family in families]
    orders = [family.order for family in families]
    by_id = {family.id: family.order for family in families}

    assert len(ids) == len(set(ids))
    assert len(orders) == len(set(orders))
    assert sorted(orders) == list(range(1, 122))

    for family_id, expected_order in EXPECTED_POST_V1_FAMILY_ORDERS.items():
        assert by_id[family_id] == expected_order

    deferred_ids = {spec.id for spec in DEFERRED_COMPONENT_SPECS}
    assert deferred_ids.isdisjoint(ids)
    assert all(family.category not in {"forecasting", "valuation"} for family in families)


@pytest.mark.parametrize(
    "case_key",
    [
        "ordinary_norm",
        "asset_light_services",
        "inventory_retail",
        "capital_intensive_manufacturer",
    ],
)
def test_normal_historical_build_never_executes_dormant_forecast(
    monkeypatch, tmp_path, case_key
):
    import core.engine.reference_model as reference_model

    monkeypatch.setattr(reference_model, "run_scenario", _forecast_must_not_run)
    monkeypatch.setattr(reference_model, "weighted_ivps", _forecast_must_not_run)

    if case_key == "ordinary_norm":
        data = _ingest_demo()
        assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
        out = tmp_path / "DEMO_NORM_Trainer.xlsx"
    elif case_key == "asset_light_services":
        case = asset_light_services_case()
        data, assumptions = case.financials, case.assumptions
        out = tmp_path / "ALS_Trainer.xlsx"
    elif case_key == "inventory_retail":
        case = inventory_retail_case()
        data, assumptions = case.financials, case.assumptions
        out = tmp_path / "IRT_Trainer.xlsx"
    else:
        case = capital_intensive_manufacturer_case()
        data, assumptions = case.financials, case.assumptions
        out = tmp_path / "CIM_Trainer.xlsx"

    trainer, answer = build_training_workbook(data, out, assumptions)
    _assert_deferred_isolation(trainer, answer)


def test_canonical_demo_trainer_answer_key_practice_contract(tmp_path):
    data = _ingest_demo()
    assumptions = json.loads(DEMO_ASSUMPTIONS.read_text(encoding="utf-8"))
    trainer, answer = build_training_workbook(
        data,
        tmp_path / "DEMO_HK_Trainer.xlsx",
        assumptions,
    )
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 78
    assert len(smap.all_ordered()) == 332

    trainer_wb = load_workbook(trainer, data_only=False)
    answer_wb = load_workbook(answer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        trainer_cell = trainer_wb[comp.tab].cell(row=row, column=col)
        answer_cell = answer_wb[comp.tab].cell(row=row, column=col)

        assert trainer_cell.value is None
        assert trainer_cell.comment is None
        assert _fill_rgb(trainer_cell) == "FFFF00"

        assert answer_cell.value == comp.formula
        assert answer_cell.comment is not None
        assert (answer_cell.comment.text or "").strip()
        assert _fill_rgb(answer_cell) == "FFFF00"
    trainer_wb.close()
    answer_wb.close()

    for suffix in (".component_map.json", ".trainer.json", ".assumptions.json"):
        assert not trainer.with_suffix(suffix).exists()

    comps = [
        c
        for c in smap.all_ordered()
        if isinstance(c.expected_value, (int, float))
    ]
    assert len(comps) >= 2
    filled, blank = comps[0], comps[1]

    wb = load_workbook(trainer, data_only=False)
    row, col = parse_cell_ref(filled.cell)
    entered_formula = filled.formula
    wb[filled.tab].cell(row=row, column=col).value = entered_formula
    wb.save(trainer)
    wb.close()

    first = check_workbook(trainer)
    assert first.correct >= 1
    assert first.blank >= 1

    wb = load_workbook(trainer, data_only=False)
    assert wb[filled.tab].cell(row, col).value == entered_formula
    assert _fill_rgb(wb[filled.tab].cell(row, col)) == "C8E6C9"
    blank_row, blank_col = parse_cell_ref(blank.cell)
    assert wb[blank.tab].cell(blank_row, blank_col).value is None
    assert _fill_rgb(wb[blank.tab].cell(blank_row, blank_col)) == "FFFF00"
    wb.close()

    second = check_workbook(trainer)
    assert second.correct >= 1
    wb = load_workbook(trainer, data_only=False)
    assert wb[filled.tab].cell(row, col).value == entered_formula
    assert _fill_rgb(wb[filled.tab].cell(row, col)) == "C8E6C9"
    assert wb[blank.tab].cell(blank_row, blank_col).value is None
    assert _fill_rgb(wb[blank.tab].cell(blank_row, blank_col)) == "FFFF00"
    wb.close()


def test_committed_canonical_demo_pair_matches_current_builder():
    trainer = ROOT / "example" / "DEMO_HK_Trainer.xlsx"
    answer = ROOT / "example" / "DEMO_HK_Answer_Key.xlsx"
    assert trainer.is_file()
    assert answer.is_file()

    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 78
    assert len(smap.all_ordered()) == 332
    summary = check_workbook(trainer)
    assert (summary.correct, summary.incorrect, summary.blank) == (0, 0, 332)

    for suffix in (".component_map.json", ".trainer.json", ".assumptions.json"):
        assert not trainer.with_suffix(suffix).exists()

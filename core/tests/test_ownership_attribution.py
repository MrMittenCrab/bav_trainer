"""Step 9M.3C — parent / NCI ownership attribution."""

from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    FinancialPeriod,
    HistoricalShareData,
    LineItem,
    StandardizedFinancials,
)
from core.data.line_identity import line_identity
from core.engine.component_catalog import (
    OWNERSHIP_ATTRIBUTION_COMPONENT_CATALOG,
    expand_ownership_attribution_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.model.financial_math import compute_anchor
from core.model.line_resolver import AmbiguousLineError, resolve_line
from core.model.ownership_attribution import (
    OWNERSHIP_BRIDGE_TOLERANCE,
    OwnershipAttributionIntegrityError,
    compute_ownership_attribution_series,
    ownership_attribution_applicable,
    ownership_attribution_availability,
    per_share_earnings_numerator,
)
from core.model.per_share import compute_per_share_series
from core.model.period_axis import canonical_fiscal_periods
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook

P1 = date(2024, 12, 31)
P2 = date(2025, 12, 31)


def _li(label: str, v1: float, v2: float, *, concept: str = "") -> LineItem:
    return LineItem(label=label, concept=concept, values={P1: v1, P2: v2})


def _periods() -> list[FinancialPeriod]:
    return [
        FinancialPeriod(end_date=P1, label="FY2024"),
        FinancialPeriod(end_date=P2, label="FY2025"),
    ]


def _ownership_fin(
    *,
    total_profit=(110.0, 110.0),
    parent_profit=(100.0, 100.0),
    nci_profit=(10.0, 10.0),
    total_equity=(200.0, 220.0),
    parent_equity=(180.0, 200.0),
    nci_equity=(20.0, 20.0),
    include_attribution: bool = True,
    omit: str | None = None,
    duplicate_parent_profit: bool = False,
    shares: bool = False,
) -> StandardizedFinancials:
    is_rows = [
        _li("Revenue", 1000, 1100),
        _li("Finance costs", -10, -11),
        _li("Finance income", 0, 0),
        _li("Profit before tax", 200, 220),
        _li("Income tax expense", -30, -33),
        _li("Profit for the year", *total_profit, concept="net_income"),
    ]
    if include_attribution and omit != "parent_profit":
        is_rows.append(
            _li(
                "Owners of the Parent",
                *parent_profit,
                concept="profit_attributable_to_owners",
            )
        )
    if duplicate_parent_profit:
        is_rows.append(
            _li(
                "Profit attributable to owners of the Parent",
                *parent_profit,
                concept="profit_attributable_to_owners",
            )
        )
    if include_attribution and omit != "nci_profit":
        is_rows.append(
            _li(
                "Non-controlling interests",
                *nci_profit,
                concept="profit_attributable_to_nci",
            )
        )

    bs_rows = [
        _li("Trade receivables", total_equity[0] + 30.0, total_equity[1] + 30.0),
        _li("Trade payables", 30.0, 30.0),
    ]
    if include_attribution and omit != "parent_equity":
        bs_rows.append(
            _li(
                "Equity attributable to owners of the Parent",
                *parent_equity,
                concept="equity_attributable_to_owners",
            )
        )
    if include_attribution and omit != "nci_equity":
        bs_rows.append(
            _li(
                "Non-controlling interests",
                *nci_equity,
                concept="noncontrolling_interests",
            )
        )
    bs_rows.append(_li("Total equity", *total_equity, concept="total_equity"))

    return StandardizedFinancials(
        ticker="OWN",
        company_name="Ownership Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=_periods(),
        income_statement=is_rows,
        balance_sheet=bs_rows,
        cash_flow=[_li("Net cash from operating activities", 80, 90)],
        historical_shares=(
            HistoricalShareData(
                scale_basis="financial_statement_units",
                diluted_weighted_average={P1: 10.0, P2: 10.0},
            )
            if shares
            else None
        ),
    )


def test_resolver_exact_ownership_concepts():
    fin = _ownership_fin()
    assert (
        resolve_line(
            fin.income_statement, "profit_attributable_to_owners", required=True
        ).item
        is not None
    )
    assert (
        resolve_line(
            fin.income_statement, "profit_attributable_to_nci", required=True
        ).item
        is not None
    )
    assert (
        resolve_line(
            fin.balance_sheet, "equity_attributable_to_owners", required=True
        ).item
        is not None
    )
    assert (
        resolve_line(
            fin.balance_sheet, "noncontrolling_interests", required=True
        ).item
        is not None
    )


def test_resolver_ownership_label_aliases():
    items = [
        LineItem(
            label="Profit attributable to owners of the Parent",
            values={P1: 1, P2: 2},
        )
    ]
    resolved = resolve_line(items, "profit_attributable_to_owners", required=True)
    assert resolved.item is not None


def test_ownership_availability_states():
    none = _ownership_fin(include_attribution=False)
    avail = ownership_attribution_availability(none)
    assert avail.available is False
    assert avail.partial is False
    assert avail.ambiguous is False

    complete = _ownership_fin()
    avail = ownership_attribution_availability(complete)
    assert avail.available is True
    assert ownership_attribution_applicable(complete) is True

    partial = _ownership_fin(omit="nci_profit")
    avail = ownership_attribution_availability(partial)
    assert avail.available is False
    assert avail.partial is True

    dup = _ownership_fin(duplicate_parent_profit=True)
    avail = ownership_attribution_availability(dup)
    assert avail.ambiguous is True
    assert avail.available is False


def test_ownership_series_and_rounding_envelope():
    fin = _ownership_fin(
        total_profit=(110.0, 111.5),
        parent_profit=(100.0, 100.0),
        nci_profit=(10.0, 10.0),
    )
    periods = [P1, P2]
    series = compute_ownership_attribution_series(fin, periods)
    assert series.profit_attribution_gap[0] == pytest.approx(0.0)
    assert series.profit_attribution_gap[1] == pytest.approx(-1.5)
    assert series.equity_attribution_gap[0] == pytest.approx(0.0)
    assert series.parent_roe[0] is None
    assert series.parent_roe[1] == pytest.approx(100.0 / ((180.0 + 200.0) / 2.0))

    bad = _ownership_fin(
        total_profit=(110.0, 111.6),
        parent_profit=(100.0, 100.0),
        nci_profit=(10.0, 10.0),
    )
    with pytest.raises(OwnershipAttributionIntegrityError, match="profit attribution"):
        compute_ownership_attribution_series(bad, periods)


def test_ownership_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert [f.order for f in OWNERSHIP_ATTRIBUTION_COMPONENT_CATALOG] == list(
        range(91, 98)
    )
    specs = expand_ownership_attribution_specs(periods, start_order=1000)
    assert len(specs) == 34
    assert {s.family_order for s in specs} == set(range(91, 98))


def test_ownership_builder_applicability(tmp_path):
    complete = _ownership_fin()
    builder = ReferenceModelBuilder(complete)
    assert builder.ownership_attribution_series is not None
    # 6 all-period families × 2 periods + 1 comparable family × 1 = 13
    assert len(builder.ownership_attribution_specs) == 13

    none = _ownership_fin(include_attribution=False)
    builder_none = ReferenceModelBuilder(none)
    assert builder_none.ownership_attribution_series is None
    assert builder_none.ownership_attribution_specs == ()

    partial = _ownership_fin(omit="parent_equity")
    builder_partial = ReferenceModelBuilder(partial)
    assert builder_partial.ownership_attribution_series is None
    assert builder_partial.ownership_attribution_specs == ()
    assert builder_partial.expected_specs


def test_ownership_workbook_formulas_and_check(tmp_path):
    fin = _ownership_fin()
    trainer, answer = build_training_workbook(fin, tmp_path / "OWN.xlsx")
    smap = load_semantic_map(answer)
    own = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in OWNERSHIP_ATTRIBUTION_COMPONENT_CATALOG}
    ]
    assert len(own) == 13
    assert "Ownership Attribution" in {c.tab for c in own}

    parent_profit = next(
        c for c in own if c.family_id == "parent_profit_source_link" and c.period_index == 1
    )
    assert "Income Statement" in parent_profit.formula
    gap = next(
        c for c in own if c.family_id == "profit_attribution_gap" and c.period_index == 1
    )
    assert "+" in gap.formula and "-" in gap.formula
    roe = next(c for c in own if c.family_id == "parent_roe")
    assert roe.period_index == 1
    assert "/" in roe.formula

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    summary = check_workbook(trainer)
    assert summary.incorrect == 0
    assert summary.blank == 0


def test_per_share_uses_parent_profit_numerator():
    fin = _ownership_fin(shares=True)
    periods = [P1, P2]
    anchor = compute_anchor(fin, periods)
    series = compute_per_share_series(fin, periods, anchor)
    assert series.earnings_numerator == (100.0, 100.0)
    assert series.reported_diluted_eps == (10.0, 10.0)

    whole = _ownership_fin(include_attribution=False, shares=True)
    # Need balanced BS without attribution — equity only total.
    whole_anchor = compute_anchor(whole, periods)
    whole_series = compute_per_share_series(whole, periods, whole_anchor)
    assert whole_series.earnings_numerator == tuple(whole_anchor.historical.net_income)
    assert whole_series.reported_diluted_eps[0] == pytest.approx(
        whole_anchor.historical.net_income[0] / 10.0
    )


def test_per_share_partial_ownership_fails_closed():
    fin = _ownership_fin(omit="nci_equity", shares=True)
    periods = [P1, P2]
    anchor = compute_anchor(fin, periods)
    with pytest.raises(OwnershipAttributionIntegrityError):
        compute_per_share_series(fin, periods, anchor)


def test_bridge_tolerance_constant():
    assert OWNERSHIP_BRIDGE_TOLERANCE == pytest.approx(0.5 * (2 + 1))

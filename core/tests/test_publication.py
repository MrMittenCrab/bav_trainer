"""Word/PDF publication from canonical research, without rebuilding analysis."""
from __future__ import annotations

from io import BytesIO
import hashlib
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
from xml.etree import ElementTree
import zipfile

import pytest

from core.__main__ import main
from core.research.document import (
    publication_filenames,
    publish_company_documents,
    publish_resolved_company,
)
from core.research.drivers import placeholder_filenames
from core.research.style import LATIN_FACE

ROOT = Path(__file__).resolve().parents[2]
CANONICAL = ROOT / "build" / "output" / "lululemon"
INSPECT = ROOT / ".git" / "autocycle" / "step-7-3-1-word-inspect"

PDF_RENDER_DPI = 150
PDF_RENDER_COLORSPACE = "rgb"
PDF_RENDER_ALPHA = False
PDF_VOLATILE_META_FIELDS = ("creationDate", "modDate", "id")
WORD_METADATA_PARTS = ("docProps/core.xml", "docProps/app.xml")


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _insert_after_title(text: str, extra: str) -> str:
    lines = text.splitlines(keepends=True)
    for index, line in enumerate(lines):
        if line.startswith("# "):
            insert_at = index + 1
            if insert_at < len(lines) and lines[insert_at].strip() == "":
                insert_at += 1
            payload = extra if extra.endswith("\n") else extra + "\n"
            lines.insert(insert_at, payload)
            return "".join(lines)
    raise AssertionError("Drivers Markdown has no title")


def _copy_research(dest: Path) -> Path:
    dest.mkdir(parents=True)
    shutil.copytree(CANONICAL / "research", dest / "research")
    shutil.copytree(CANONICAL / "figures", dest / "figures")
    return dest


def _company(tmp_path: Path, monkeypatch, name: str = "Lululemon"):
    from core import current_build
    monkeypatch.setattr(current_build, "OUTPUT_ROOT", tmp_path)
    return current_build.resolve_company(name)


def _norm_num(value, digits: int = 4):
    return round(float(value), digits)


def _norm_box(box, digits: int = 4):
    return tuple(_norm_num(part, digits) for part in box)


def _norm_color(color):
    if color is None:
        return None
    if isinstance(color, (int, float)):
        return _norm_num(color, 6)
    return tuple(_norm_num(part, 6) for part in color)


def _serialize_geometry(part, digits: int = 4):
    if hasattr(part, "x0") and hasattr(part, "y0"):
        return _norm_box((part.x0, part.y0, part.x1, part.y1), digits)
    if hasattr(part, "x") and hasattr(part, "y"):
        return (_norm_num(part.x, digits), _norm_num(part.y, digits))
    if isinstance(part, (tuple, list)):
        return tuple(_serialize_geometry(item, digits) for item in part)
    if isinstance(part, (int, float)):
        return _norm_num(part, digits)
    return part


def _word_member_payloads(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _parse_core_fields(payload: bytes) -> dict[str, str]:
    if not payload:
        return {}
    root = ElementTree.fromstring(payload)
    return {child.tag.split("}", 1)[-1]: child.text or "" for child in root}


def _parse_app_fields(payload: bytes) -> dict[str, str]:
    if not payload:
        return {}
    root = ElementTree.fromstring(payload)
    fields: dict[str, str] = {}
    for child in root:
        tag = child.tag.split("}", 1)[-1]
        if list(child):
            continue
        fields[tag] = child.text or ""
    return fields


def _xml_field_diffs(first: bytes, second: bytes, parser) -> list[str]:
    left, right = parser(first), parser(second)
    names = sorted(set(left) | set(right))
    return [name for name in names if left.get(name) != right.get(name)]


def _word_payload(data: bytes) -> dict:
    from docx import Document

    members = _word_member_payloads(data)
    document = Document(BytesIO(data))
    text = [paragraph.text for paragraph in document.paragraphs]
    tables = [
        [[cell.text for cell in row.cells] for row in document_table.rows]
        for document_table in document.tables
    ]
    sections = [
        {
            "page_width": int(section.page_width),
            "page_height": int(section.page_height),
            "orientation": str(section.orientation),
            "left_margin": int(section.left_margin),
            "right_margin": int(section.right_margin),
            "top_margin": int(section.top_margin),
            "bottom_margin": int(section.bottom_margin),
        }
        for section in document.sections
    ]
    return {
        "inventory": tuple(sorted(members)),
        "members": {
            name: hashlib.sha256(payload).hexdigest() for name, payload in members.items()
        },
        "text": text,
        "tables": tables,
        "sections": sections,
        "media": {
            name: hashlib.sha256(payload).hexdigest()
            for name, payload in members.items()
            if name.startswith("word/media/")
        },
        "core": hashlib.sha256(members.get("docProps/core.xml", b"")).hexdigest(),
        "core_fields": _parse_core_fields(members.get("docProps/core.xml", b"")),
    }


def _word_documents_equal(first: bytes, second: bytes) -> bool:
    left, right = _word_payload(first), _word_payload(second)
    return left["inventory"] == right["inventory"] and left["members"] == right["members"]


def _word_member_diffs(first: bytes, second: bytes) -> list[str]:
    left, right = _word_member_payloads(first), _word_member_payloads(second)
    names = sorted(set(left) | set(right))
    diffs = []
    for name in names:
        if name not in left:
            diffs.append(f"{name} (only second)")
        elif name not in right:
            diffs.append(f"{name} (only first)")
        elif left[name] != right[name]:
            diffs.append(name)
    return diffs


def _pdf_id(data: bytes) -> str | None:
    match = re.search(
        rb"/ID\s*\[\s*<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>",
        data,
    )
    if match is None:
        return None
    return f"{match.group(1).decode()}/{match.group(2).decode()}"


def _pdf_render_samples(page) -> tuple[bytes, tuple[int, int, int]]:
    import fitz

    pixmap = page.get_pixmap(
        matrix=fitz.Matrix(PDF_RENDER_DPI / 72.0, PDF_RENDER_DPI / 72.0),
        colorspace=fitz.csRGB,
        alpha=PDF_RENDER_ALPHA,
    )
    return bytes(pixmap.samples), (pixmap.width, pixmap.height, pixmap.n)


def _pdf_page_layout(page, doc) -> dict:
    text_layout = []
    for block in page.get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text_layout.append(
                    {
                        "text": span["text"],
                        "origin": _norm_box(span["origin"]),
                        "bbox": _norm_box(span["bbox"]),
                        "font": span["font"],
                        "size": _norm_num(span["size"], 3),
                        "color": span["color"],
                        "direction": _norm_box(line.get("dir", (1, 0))),
                    }
                )
    images = []
    image_layout = []
    for image in page.get_images(full=True):
        xref = image[0]
        digest = hashlib.sha256(doc.extract_image(xref)["image"]).hexdigest()
        images.append(digest)
        image_layout.append(
            {
                "sha": digest,
                "rects": tuple(_norm_box(rect) for rect in page.get_image_rects(xref)),
            }
        )
    drawings = []
    for drawing in page.get_drawings():
        drawings.append(
            {
                "type": drawing.get("type"),
                "color": _norm_color(drawing.get("color")),
                "fill": _norm_color(drawing.get("fill")),
                "width": _norm_num(drawing.get("width") or 0, 4),
                "rect": _norm_box(drawing.get("rect") or (0, 0, 0, 0)),
                "items": tuple(
                    (item[0], *(_serialize_geometry(part) for part in item[1:]))
                    for item in drawing.get("items", ())
                ),
            }
        )
    samples, size = _pdf_render_samples(page)
    return {
        "geometry": {
            "mediabox": _norm_box(page.mediabox),
            "cropbox": _norm_box(page.cropbox),
            "rotation": int(page.rotation),
        },
        "text": page.get_text(),
        "text_layout": text_layout,
        "images": images,
        "image_layout": image_layout,
        "drawings": drawings,
        "render": hashlib.sha256(samples).hexdigest(),
        "render_size": size,
    }


def _pdf_payload(data: bytes) -> dict:
    import fitz

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        pages = [_pdf_page_layout(page, doc) for page in doc]
        meta = {
            key: doc.metadata.get(key)
            for key in ("creationDate", "modDate", "producer", "creator")
        }
        meta["id"] = _pdf_id(data)
        return {
            "pages": doc.page_count,
            "text": [page["text"] for page in pages],
            "images": [image for page in pages for image in page["images"]],
            "geometry": [page["geometry"] for page in pages],
            "text_layout": [page["text_layout"] for page in pages],
            "image_layout": [page["image_layout"] for page in pages],
            "drawings": [page["drawings"] for page in pages],
            "renders": [page["render"] for page in pages],
            "render_size": [page["render_size"] for page in pages],
            "render_spec": {
                "renderer": "pymupdf",
                "dpi": PDF_RENDER_DPI,
                "colorspace": PDF_RENDER_COLORSPACE,
                "alpha": PDF_RENDER_ALPHA,
            },
            "meta": meta,
        }
    finally:
        doc.close()


def _text_layout_features(first: list[dict], second: list[dict]) -> list[str]:
    if len(first) != len(second):
        return ["text_layout"]
    features: list[str] = []
    for left, right in zip(first, second):
        if left["text"] != right["text"]:
            features.append("text_content")
        if left["origin"] != right["origin"] or left["bbox"] != right["bbox"]:
            features.append("text_position")
        if left["size"] != right["size"]:
            features.append("font_size")
        if left["font"] != right["font"] or left["color"] != right["color"]:
            features.append("typography")
    return sorted(set(features))


def _pdf_page_features(first: dict, second: dict) -> list[str]:
    features: list[str] = []
    if first["geometry"] != second["geometry"]:
        features.append("page_geometry")
    features.extend(_text_layout_features(first["text_layout"], second["text_layout"]))
    if first["images"] != second["images"]:
        features.append("image_content")
    if first["image_layout"] != second["image_layout"]:
        features.append("image_position")
    if first["drawings"] != second["drawings"]:
        features.append("drawing_geometry")
    if first["render"] != second["render"] or first["render_size"] != second["render_size"]:
        features.append("render")
    return features


def _pdf_layout_diffs(first: dict, second: dict) -> dict:
    pages = []
    features: list[str] = []
    if first["pages"] != second["pages"]:
        features.append("page_count")
    count = min(first["pages"], second["pages"])
    for index in range(count):
        left = {
            "geometry": first["geometry"][index],
            "text_layout": first["text_layout"][index],
            "images": [
                item["sha"] for item in first["image_layout"][index]
            ],
            "image_layout": first["image_layout"][index],
            "drawings": first["drawings"][index],
            "render": first["renders"][index],
            "render_size": first["render_size"][index],
        }
        right = {
            "geometry": second["geometry"][index],
            "text_layout": second["text_layout"][index],
            "images": [item["sha"] for item in second["image_layout"][index]],
            "image_layout": second["image_layout"][index],
            "drawings": second["drawings"][index],
            "render": second["renders"][index],
            "render_size": second["render_size"][index],
        }
        page_features = _pdf_page_features(left, right)
        if page_features:
            pages.append({"page": index + 1, "features": page_features})
            features.extend(page_features)
    if first["pages"] > second["pages"]:
        for index in range(second["pages"], first["pages"]):
            pages.append({"page": index + 1, "features": ["only_first"]})
    elif second["pages"] > first["pages"]:
        for index in range(first["pages"], second["pages"]):
            pages.append({"page": index + 1, "features": ["only_second"]})
    return {"pages": pages, "features": sorted(set(features))}


def _pdf_documents_equal(first: bytes, second: bytes) -> bool:
    left, right = _pdf_payload(first), _pdf_payload(second)
    return (
        left["pages"] == right["pages"]
        and left["geometry"] == right["geometry"]
        and left["text_layout"] == right["text_layout"]
        and left["image_layout"] == right["image_layout"]
        and left["drawings"] == right["drawings"]
        and left["renders"] == right["renders"]
        and left["text"] == right["text"]
        and left["images"] == right["images"]
    )


def _content_equal(first_word: bytes, first_pdf: bytes, second_word: bytes, second_pdf: bytes) -> bool:
    return _word_documents_equal(first_word, second_word) and _pdf_documents_equal(
        first_pdf, second_pdf
    )


def _metadata_field_diffs(first_word: bytes, first_pdf: bytes, second_word: bytes, second_pdf: bytes) -> dict:
    word_a, word_b = _word_member_payloads(first_word), _word_member_payloads(second_word)
    pdf_a, pdf_b = _pdf_payload(first_pdf), _pdf_payload(second_pdf)
    core_fields = _xml_field_diffs(
        word_a.get("docProps/core.xml", b""),
        word_b.get("docProps/core.xml", b""),
        _parse_core_fields,
    )
    app_fields = _xml_field_diffs(
        word_a.get("docProps/app.xml", b""),
        word_b.get("docProps/app.xml", b""),
        _parse_app_fields,
    )
    pdf_fields = [
        key
        for key in ("creationDate", "modDate", "producer", "creator", "id")
        if pdf_a["meta"].get(key) != pdf_b["meta"].get(key)
    ]
    return {
        "word_core_fields": core_fields,
        "word_app_fields": app_fields,
        "pdf_meta_fields": pdf_fields,
    }


def _publication_diff(first_word: bytes, first_pdf: bytes, second_word: bytes, second_pdf: bytes) -> dict:
    word_members = _word_member_diffs(first_word, second_word)
    pdf_layout = _pdf_layout_diffs(_pdf_payload(first_pdf), _pdf_payload(second_pdf))
    fields = _metadata_field_diffs(first_word, first_pdf, second_word, second_pdf)
    metadata_parts = [name for name in word_members if name.split(" ")[0] in WORD_METADATA_PARTS]
    unexplained = [
        name
        for name in word_members
        if name.split(" ")[0] not in WORD_METADATA_PARTS
    ]
    return {
        "word_members": word_members,
        "pdf_pages": pdf_layout["pages"],
        "pdf_features": pdf_layout["features"],
        "word_metadata_parts": metadata_parts,
        "word_non_metadata_members": unexplained,
        **fields,
        "word_core": bool(fields["word_core_fields"]),
        "pdf_meta": bool(fields["pdf_meta_fields"]),
        "word_sha": hashlib.sha256(first_word).hexdigest() != hashlib.sha256(second_word).hexdigest(),
        "pdf_sha": hashlib.sha256(first_pdf).hexdigest() != hashlib.sha256(second_pdf).hexdigest(),
    }


def _metadata_diff(first_word: bytes, first_pdf: bytes, second_word: bytes, second_pdf: bytes) -> dict:
    return _publication_diff(first_word, first_pdf, second_word, second_pdf)


def _save_docx(document) -> bytes:
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _zip_with_timestamps(data: bytes, date_time: tuple[int, int, int, int, int, int]) -> bytes:
    buffer = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as source, zipfile.ZipFile(buffer, "w") as dest:
        for info in source.infolist():
            copied = zipfile.ZipInfo(filename=info.filename, date_time=date_time)
            copied.compress_type = info.compress_type
            copied.external_attr = info.external_attr
            dest.writestr(copied, source.read(info.filename))
    return buffer.getvalue()


def _pdf_with_volatile_metadata(data: bytes, stamp: bytes) -> bytes:
    updated = re.sub(rb"D:\d{14}\+\d{2}'\d{2}'", stamp, data, count=2)
    if updated == data:
        updated = data.replace(
            _pdf_payload(data)["meta"]["creationDate"].encode(),
            stamp,
        )
    identity = re.search(rb"/ID\s*\[\s*<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", updated)
    if identity is not None:
        replacement = identity.group(0).replace(identity.group(2), b"0" * len(identity.group(2)))
        updated = updated.replace(identity.group(0), replacement, 1)
    return updated


def _rewrite_pdf_stream(data: bytes, page_index: int, replacer) -> bytes:
    import fitz

    doc = fitz.open(stream=data, filetype="pdf")
    try:
        page = doc[page_index]
        xrefs = page.get_contents()
        assert xrefs, f"page {page_index + 1} has no contents"
        stream = doc.xref_stream(xrefs[0])
        rewritten = replacer(stream)
        assert rewritten != stream
        doc.update_stream(xrefs[0], rewritten)
        return doc.tobytes(garbage=0, deflate=False)
    finally:
        doc.close()


def _shift_first_tm(stream: bytes) -> bytes:
    match = re.search(rb"1 0 0 1 ([\d.]+) ([\d.]+) Tm", stream)
    assert match is not None
    new_y = f"{float(match.group(2)) - 18:.3f}".encode()
    return stream[: match.start(2)] + new_y + stream[match.end(2) :]


def _enlarge_first_body_tf(stream: bytes) -> bytes:
    match = re.search(rb"/F2\+0 10 Tf", stream)
    assert match is not None
    return stream[: match.start()] + b"/F2+0 16 Tf" + stream[match.end() :]


def _shift_first_image(stream: bytes) -> bytes:
    match = re.search(
        rb"1 0 0 1 ([\d.]+) ([\d.]+) cm\nq\n[\d.]+ 0 0 [\d.]+ 0 0 cm\n/FormXob",
        stream,
    )
    assert match is not None
    new_y = f"{float(match.group(2)) - 24:.3f}".encode()
    return stream[: match.start(2)] + new_y + stream[match.end(2) :]


def _shift_first_rule(stream: bytes) -> bytes:
    match = re.search(rb"n ([\d.]+) ([\d.]+) m ([\d.]+) ([\d.]+) l S", stream)
    assert match is not None
    new_x = f"{float(match.group(3)) - 40:.3f}".encode()
    return stream[: match.start(3)] + new_x + stream[match.end(3) :]


def _mutate_word(data: bytes, mutator) -> bytes:
    from docx import Document

    document = Document(BytesIO(data))
    mutator(document)
    return _save_docx(document)


def _publish_cli_hiding(module: str, output_root: Path, company: str = "Lululemon"):
    script = f"""
import sys
from pathlib import Path

class _Block:
    def find_spec(self, name, path=None, target=None):
        root = {module!r}
        if name == root or name.startswith(root + "."):
            raise ImportError(root)
        return None

sys.meta_path.insert(0, _Block())
from core import current_build
current_build.OUTPUT_ROOT = Path({str(output_root)!r})
from core.__main__ import main
raise SystemExit(main(["publish", {company!r}]))
"""
    env = os.environ.copy()
    env["PYTHONPATH"] = os.pathsep.join(
        [str(ROOT), env["PYTHONPATH"]] if env.get("PYTHONPATH") else [str(ROOT)]
    )
    return subprocess.run(
        [sys.executable, "-c", script],
        cwd=str(ROOT),
        capture_output=True,
        text=True,
        env=env,
    )


def test_publish_help_is_bav_first():
    from core.__main__ import main as cli
    import io
    from contextlib import redirect_stdout
    buf = io.StringIO()
    with pytest.raises(SystemExit) as raised:
        with redirect_stdout(buf):
            cli(["publish", "--help"])
    assert raised.value.code == 0
    text = buf.getvalue()
    assert "BAV" in text
    assert "canonical Drivers research" in text
    assert "Answer Key" not in text


def test_fast_retailing_has_no_publishable_research(capsys):
    before = {
        path: path.read_bytes()
        for path in (ROOT / "build" / "output" / "fast_retailing").rglob("*")
        if path.is_file()
    }
    assert main(["publish", "FastRetailing"]) != 0
    err = capsys.readouterr().err
    assert "No publishable canonical research" in err
    assert "FastRetailing" in err
    assert "legacy" in err.casefold()
    after = {
        path: path.read_bytes()
        for path in (ROOT / "build" / "output" / "fast_retailing").rglob("*")
        if path.is_file()
    }
    assert after == before
    word, pdf = publication_filenames("FastRetailing")
    assert not (ROOT / "build" / "output" / "fast_retailing" / word).exists()
    assert not (ROOT / "build" / "output" / "fast_retailing" / pdf).exists()


def test_missing_markdown_fails_without_inventing(tmp_path, monkeypatch, capsys):
    company = _company(tmp_path, monkeypatch)
    company.output.mkdir(parents=True)
    assert main(["publish", "Lululemon"]) != 0
    err = capsys.readouterr().err
    assert "No publishable canonical research" in err
    assert not list(company.output.glob("*.docx"))
    assert not list(company.output.glob("*.pdf"))
    assert not (company.output / "research").exists()


def test_missing_figure_fails(tmp_path, monkeypatch, capsys):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    (company.output / "figures" / "drivers" / "growth.png").unlink()
    assert main(["publish", "LULU"]) != 0
    err = capsys.readouterr().err
    assert "growth.png" in err
    assert not list(company.output.glob("*.docx"))


def test_broken_reference_fails(tmp_path, monkeypatch, capsys):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    drivers = company.output / "research" / "Lululemon_Drivers.md"
    text = drivers.read_text(encoding="utf-8")
    drivers.write_text(text.replace("growth.png", "missing.png"), encoding="utf-8")
    assert main(["publish", "lululemon"]) != 0
    err = capsys.readouterr().err
    assert "missing.png" in err or "broken" in err.casefold() or "missing" in err.casefold()
    assert not list(company.output.glob("*.docx"))


def test_valid_document_references_are_preserved(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    note = company.output / "research" / "source-note.md"
    note.write_text("# Source note\n\nItem 7 locator retained.\n", encoding="utf-8")
    drivers = company.output / "research" / "Lululemon_Drivers.md"
    text = drivers.read_text(encoding="utf-8")
    drivers.write_text(
        _insert_after_title(
            text,
            "See the [source note](source-note.md#source-note) and [Appendix](#appendix).\n",
        ),
        encoding="utf-8",
    )
    published = publish_resolved_company(company)
    from docx import Document
    import fitz

    word_text = "\n".join(p.text for p in Document(published.word).paragraphs)
    assert "source note (source-note.md)" in word_text
    assert "Appendix" in word_text
    doc = fitz.open(published.pdf)
    try:
        pdf_text = "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()
    assert "source note (source-note.md)" in pdf_text
    assert "Appendix" in pdf_text


def test_missing_local_document_reference_fails(tmp_path, monkeypatch, capsys):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    drivers = company.output / "research" / "Lululemon_Drivers.md"
    text = drivers.read_text(encoding="utf-8")
    drivers.write_text(
        _insert_after_title(text, "See [missing note](missing-note.md).\n"),
        encoding="utf-8",
    )
    prior = {
        path: path.read_bytes()
        for path in company.output.rglob("*")
        if path.is_file() and path.suffix in {".docx", ".pdf"}
    }
    assert main(["publish", "Lululemon"]) != 0
    err = capsys.readouterr().err
    assert "missing-note.md" in err
    assert str(drivers) in err or "Lululemon_Drivers.md" in err
    assert "broken document reference" in err
    assert not list(company.output.glob("*.docx"))
    for path, data in prior.items():
        assert path.read_bytes() == data


def test_invalid_document_anchor_fails(tmp_path, monkeypatch, capsys):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    drivers = company.output / "research" / "Lululemon_Drivers.md"
    text = drivers.read_text(encoding="utf-8")
    drivers.write_text(
        _insert_after_title(text, "See [missing heading](#no-such-heading).\n"),
        encoding="utf-8",
    )
    assert main(["publish", "Lululemon"]) != 0
    err = capsys.readouterr().err
    assert "no-such-heading" in err
    assert "Lululemon_Drivers.md" in err
    assert "invalid anchor" in err
    assert not list(company.output.glob("*.docx"))


def test_invalid_anchor_on_local_markdown_fails(tmp_path, monkeypatch, capsys):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    (company.output / "research" / "source-note.md").write_text(
        "# Source note\n\nPresent.\n", encoding="utf-8"
    )
    drivers = company.output / "research" / "Lululemon_Drivers.md"
    text = drivers.read_text(encoding="utf-8")
    drivers.write_text(
        _insert_after_title(text, "See [bad target](source-note.md#absent-section).\n"),
        encoding="utf-8",
    )
    assert main(["publish", "Lululemon"]) != 0
    err = capsys.readouterr().err
    assert "source-note.md#absent-section" in err
    assert "invalid anchor" in err
    assert "Lululemon_Drivers.md" in err
    assert not list(company.output.glob("*.docx"))


def test_missing_font_fails(tmp_path, monkeypatch, capsys):
    from core.research import document
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    monkeypatch.setattr(
        document,
        "resolve_required_fonts",
        lambda: (_ for _ in ()).throw(ValueError("required fonts unavailable: Aptos Regular")),
    )
    assert main(["publish", "Lululemon"]) != 0
    assert "required fonts unavailable" in capsys.readouterr().err
    assert not list(company.output.glob("*.docx"))


def test_missing_pandoc_fails(tmp_path, monkeypatch, capsys):
    from core.research import document
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    monkeypatch.setattr(document.shutil, "which", lambda name: None)
    assert main(["publish", "Lululemon"]) != 0
    assert "pandoc" in capsys.readouterr().err
    assert not list(company.output.glob("*.docx"))


def test_converter_failure_preserves_prior_publication(tmp_path, monkeypatch, capsys):
    from core.research import document
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    assert main(["publish", "Lululemon"]) == 0
    word = company.output / "Lululemon_BAV.docx"
    pdf = company.output / "Lululemon_BAV.pdf"
    prior = {word: word.read_bytes(), pdf: pdf.read_bytes()}
    monkeypatch.setattr(
        document,
        "_render_pdf",
        lambda *args, **kwargs: (_ for _ in ()).throw(RuntimeError("forced converter failure")),
    )
    assert main(["publish", "Lululemon"]) != 0
    assert "publication conversion failed" in capsys.readouterr().err
    assert word.read_bytes() == prior[word]
    assert pdf.read_bytes() == prior[pdf]


def test_publish_does_not_require_trainer_or_mutate_inputs(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    workbook = company.output / "Lululemon_BAV.xlsx"
    workbook.write_bytes(b"workbook-bytes")
    trainer = company.output / "Lululemon_BAV_Trainer.xlsx"
    before = {
        path: path.read_bytes()
        for path in company.output.rglob("*")
        if path.is_file()
    }
    assert main(["publish", "Lululemon"]) == 0
    assert not trainer.exists()
    assert workbook.read_bytes() == b"workbook-bytes"
    for name in placeholder_filenames(company.name):
        path = company.output / "research" / name
        assert path.is_file() and path.stat().st_size == 0
    for path in (company.output / "figures" / "drivers").glob("*.png"):
        assert path.read_bytes() == before[path]
    drivers = company.output / "research" / "Lululemon_Drivers.md"
    assert drivers.read_bytes() == before[drivers]


def test_repeat_comparison_detects_content_difference(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    first = publish_resolved_company(company)
    first_word = first.word.read_bytes()
    first_pdf = first.pdf.read_bytes()
    drivers = company.output / "research" / "Lululemon_Drivers.md"
    drivers.write_text(
        _insert_after_title(
            drivers.read_text(encoding="utf-8"),
            "Deliberate publication difference for comparison.\n",
        ),
        encoding="utf-8",
    )
    second = publish_resolved_company(company)
    second_word = second.word.read_bytes()
    second_pdf = second.pdf.read_bytes()
    assert not _content_equal(first_word, first_pdf, second_word, second_pdf)
    assert "Deliberate publication difference for comparison" in "\n".join(
        _word_payload(second_word)["text"]
    )


def test_repeat_publish_equivalent_content(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    first = publish_resolved_company(company)
    first_word = first.word.read_bytes()
    first_pdf = first.pdf.read_bytes()
    second = publish_resolved_company(company)
    second_word = second.word.read_bytes()
    second_pdf = second.pdf.read_bytes()
    assert _content_equal(first_word, first_pdf, second_word, second_pdf)
    metadata = _metadata_diff(first_word, first_pdf, second_word, second_pdf)
    assert metadata["word_members"] == []
    assert metadata["pdf_pages"] == []
    assert metadata["word_core_fields"] == []
    assert metadata["word_non_metadata_members"] == []
    assert set(metadata["pdf_meta_fields"]) <= set(PDF_VOLATILE_META_FIELDS)
    assert "word_core" in metadata
    assert "pdf_meta" in metadata
    assert "word_sha" in metadata
    assert "pdf_sha" in metadata


def test_word_layout_mutations_rejected(tmp_path, monkeypatch):
    from docx.oxml.ns import qn
    from docx.shared import Inches, Mm, Pt

    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    first = publish_resolved_company(company)
    first_word = first.word.read_bytes()
    first_pdf = first.pdf.read_bytes()
    baseline = _word_payload(first_word)
    baseline_text = "\n".join(baseline["text"])
    baseline_tables = baseline["tables"]

    def title_30pt(document):
        for run in document.paragraphs[0].runs:
            run.font.size = Pt(30)

    def paragraph_spacing(document):
        document.paragraphs[0].paragraph_format.space_after = Pt(36)
        document.paragraphs[0].paragraph_format.space_before = Pt(18)

    def style_typography(document):
        normal = document.styles["Normal"]
        normal.font.size = Pt(16)
        normal.font.name = "Times New Roman"

    def table_geometry(document):
        table = document.tables[0]
        width = Inches(2.4)
        for cell in table.columns[0].cells:
            cell.width = width
        grid = table._tbl.find(qn("w:tblGrid"))
        columns = [] if grid is None else grid.findall(qn("w:gridCol"))
        if columns:
            columns[0].set(qn("w:w"), str(int(width)))

    def section_margins(document):
        section = document.sections[0]
        section.left_margin = Mm(40)
        section.right_margin = Mm(40)
        section.top_margin = Mm(36)
        section.bottom_margin = Mm(36)

    cases = (
        ("title_30pt", title_30pt, "word/document.xml"),
        ("paragraph_spacing", paragraph_spacing, "word/document.xml"),
        ("style_typography", style_typography, "word/styles.xml"),
        ("table_geometry", table_geometry, "word/document.xml"),
        ("section_margins", section_margins, "word/document.xml"),
    )
    for name, mutator, expected_member in cases:
        mutated = _mutate_word(first_word, mutator)
        payload = _word_payload(mutated)
        assert "\n".join(payload["text"]) == baseline_text, name
        assert payload["tables"] == baseline_tables, name
        assert not _content_equal(first_word, first_pdf, mutated, first_pdf), name
        diff = _publication_diff(first_word, first_pdf, mutated, first_pdf)
        assert expected_member in diff["word_members"], (name, diff["word_members"])
        assert expected_member in diff["word_non_metadata_members"], name


def test_pdf_layout_mutations_rejected(tmp_path, monkeypatch):
    import fitz

    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    first = publish_resolved_company(company)
    first_word = first.word.read_bytes()
    first_pdf = first.pdf.read_bytes()
    baseline = _pdf_payload(first_pdf)

    def image_page_index() -> int:
        doc = fitz.open(stream=first_pdf, filetype="pdf")
        try:
            for index, page in enumerate(doc):
                if page.get_images():
                    return index
        finally:
            doc.close()
        raise AssertionError("publication PDF has no image page")

    cases = (
        ("text_position", 0, _shift_first_tm, "text_position"),
        ("font_size", 0, _enlarge_first_body_tf, "font_size"),
        ("image_position", image_page_index(), _shift_first_image, "image_position"),
        ("table_rule", 0, _shift_first_rule, "drawing_geometry"),
    )
    for name, page_index, replacer, feature in cases:
        mutated = _rewrite_pdf_stream(first_pdf, page_index, replacer)
        payload = _pdf_payload(mutated)
        assert payload["pages"] == baseline["pages"], name
        assert payload["text"] == baseline["text"], name
        assert payload["images"] == baseline["images"], name
        assert not _content_equal(first_word, first_pdf, first_word, mutated), name
        diff = _publication_diff(first_word, first_pdf, first_word, mutated)
        assert feature in diff["pdf_features"], (name, diff["pdf_features"], diff["pdf_pages"])
        assert any(page["page"] == page_index + 1 for page in diff["pdf_pages"]), (name, diff["pdf_pages"])


def test_allowed_volatile_metadata_remains_equivalent(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    first = publish_resolved_company(company)
    first_word = first.word.read_bytes()
    first_pdf = first.pdf.read_bytes()
    stamped_word = _zip_with_timestamps(first_word, (2020, 1, 2, 3, 4, 5))
    assert hashlib.sha256(stamped_word).digest() != hashlib.sha256(first_word).digest()
    assert _word_member_diffs(first_word, stamped_word) == []
    assert _content_equal(first_word, first_pdf, stamped_word, first_pdf)
    stamped_pdf = _pdf_with_volatile_metadata(first_pdf, b"D:20200101000000+00'00'")
    assert hashlib.sha256(stamped_pdf).digest() != hashlib.sha256(first_pdf).digest()
    assert _content_equal(first_word, first_pdf, first_word, stamped_pdf)
    metadata = _metadata_diff(first_word, first_pdf, stamped_word, stamped_pdf)
    assert metadata["word_members"] == []
    assert metadata["pdf_pages"] == []
    assert metadata["word_core_fields"] == []
    assert set(metadata["pdf_meta_fields"]) <= set(PDF_VOLATILE_META_FIELDS)
    assert metadata["pdf_meta"]
    assert metadata["word_sha"]
    assert metadata["pdf_sha"]


def test_inspect_artifacts_match_fresh_publication(tmp_path, monkeypatch):
    inspect_docx = INSPECT / "Lululemon_BAV.docx"
    inspect_pages = INSPECT / "pdf-pages"
    assert inspect_docx.is_file()
    assert (INSPECT / "manifest.json").is_file()
    assert (INSPECT / "Lululemon_BAV.word.pdf").is_file()
    word_pages = sorted((INSPECT / "word-pages").glob("word-page-*.png"))
    pdf_pages = sorted(inspect_pages.glob("pdf-page-*.png"))
    assert word_pages
    assert pdf_pages
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    published = publish_resolved_company(company)
    from docx import Document
    import fitz

    document = Document(published.word)
    heading_styles = {
        paragraph.style.name
        for paragraph in document.paragraphs
        if paragraph.style is not None and paragraph.style.name.startswith("Heading")
    }
    assert "Heading 1" in heading_styles
    assert "Heading 2" in heading_styles
    word_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert word_text.find("Drivers") < word_text.find("Appendix")
    doc = fitz.open(published.pdf)
    try:
        pdf_text = "\n".join(page.get_text() for page in doc)
        images = sum(len(page.get_images()) for page in doc)
        assert doc.page_count >= 3
        assert pdf_text.find("Drivers") < pdf_text.find("Appendix")
        assert images >= 3
    finally:
        doc.close()


def test_missing_python_docx_reaches_cli(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    assert main(["publish", "Lululemon"]) == 0
    word = company.output / "Lululemon_BAV.docx"
    pdf = company.output / "Lululemon_BAV.pdf"
    prior = {word: word.read_bytes(), pdf: pdf.read_bytes()}
    result = _publish_cli_hiding("docx", tmp_path)
    assert result.returncode != 0
    err = result.stderr
    assert "python-docx" in err
    assert "pip install -r requirements-trainer.txt" in err
    assert "README" in err
    assert word.read_bytes() == prior[word]
    assert pdf.read_bytes() == prior[pdf]


def test_missing_reportlab_reaches_cli(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    assert main(["publish", "Lululemon"]) == 0
    word = company.output / "Lululemon_BAV.docx"
    pdf = company.output / "Lululemon_BAV.pdf"
    prior = {word: word.read_bytes(), pdf: pdf.read_bytes()}
    result = _publish_cli_hiding("reportlab", tmp_path)
    assert result.returncode != 0
    err = result.stderr
    assert "reportlab" in err
    assert "pip install -r requirements-trainer.txt" in err
    assert "README" in err
    assert word.read_bytes() == prior[word]
    assert pdf.read_bytes() == prior[pdf]


def test_lululemon_publication_preserves_analysis(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    published = publish_resolved_company(company)
    from docx import Document
    import fitz

    document = Document(published.word)
    word_text = "\n".join(p.text for p in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    combined = word_text + "\n" + table_text
    required = (
        "Lululemon — Drivers",
        "Appendix",
        "Lululemon BAV",
        "FY2024",
        "2 February 2025",
        "Form 10-K pp. 28–29",
        "approximately $275 million",
        "not store productivity",
        "not measured new-store revenue",
    )
    for item in required:
        assert item in combined, item
    assert "Forecast" not in [p.text for p in document.paragraphs if p.text.startswith("Forecast")]
    assert "Valuation" not in word_text
    assert "# Overview" not in word_text
    doc = fitz.open(published.pdf)
    try:
        pdf_text = "\n".join(page.get_text() for page in doc)
        fonts = {font[3] or font[4] for page in doc for font in page.get_fonts()}
        images = sum(len(page.get_images()) for page in doc)
        assert doc.page_count >= 3
    finally:
        doc.close()
    pdf_norm = " ".join(pdf_text.split())
    for item in required:
        assert item in pdf_norm, item
    assert images >= 3
    assert any(LATIN_FACE.casefold() in str(name).casefold() for name in fonts)
    with zipfile.ZipFile(published.word) as archive:
        assert sum(1 for name in archive.namelist() if name.startswith("word/media/")) >= 3


def test_publish_writes_only_under_canonical_output(tmp_path, monkeypatch):
    company = _company(tmp_path, monkeypatch)
    _copy_research(company.output)
    published = publish_company_documents("LULU")
    assert published.word.parent == company.output
    assert published.pdf.parent == company.output
    assert published.word.name == "Lululemon_BAV.docx"
    assert published.pdf.name == "Lululemon_BAV.pdf"
    assert company.output == tmp_path / "lululemon"

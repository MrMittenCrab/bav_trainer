"""Step 9M.9 — PP&E capex practice, ratios, and workbook Check."""

from __future__ import annotations

import copy
import json
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core.data.interface import (
    DocumentManifest,
    DocumentType,
    FinancialPeriod,
    LineItem,
    StandardizedFinancials,
)
from core.data.standardized_io import standardized_from_payload, standardized_to_payload
from core.engine.component_catalog import (
    CAPEX_COMPONENT_CATALOG,
    expand_capex_specs,
)
from core.engine.reference_model import ReferenceModelBuilder
from core.ingestion.manual_hk import HKManualDocumentAdapter
from core.model.capex import (
    capex_applicable,
    capex_availability,
    cash_after_ppe_capex_applicable,
    compute_capex_series,
    resolve_capex_source,
    resolve_operating_cash_source,
)
from core.model.financial_math import compute_anchor
from core.model.historical_expected import capex_expected_series
from core.model.line_resolver import AmbiguousLineError, MissingLineError, resolve_line
from core.model.period_axis import canonical_fiscal_periods
from core.model.ratio_values import UNDEFINED_RATIO
from core.model.source_values import MissingHistoricalValueError, required_period_value
from core.tests.test_normalization import _inject_formula_and_cached_value
from core.trainer.checker import check_workbook
from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
from core.trainer.workbook import build_training_workbook, group_components_by_family

ROOT = Path(__file__).resolve().parents[2]
STD_JSON = ROOT / "benchmark" / "fast_retailing" / "reconciled" / "standardized.json"
DEMO_JSON = ROOT / "example" / "DEMO_HK_Standardized.json"

P1 = date(2024, 12, 31)
P2 = date(2025, 12, 31)


def _ingest_demo():
    return HKManualDocumentAdapter().ingest(
        [DocumentManifest(path=str(DEMO_JSON), doc_type=DocumentType.OTHER)]
    )


def _li(label, values, concept=""):
    return LineItem(label=label, values=values, concept=concept)


def _tiny(
    *,
    payments=(-100.0, -120.0),
    revenue=(1000.0, 1100.0),
    with_payments: bool = True,
    label_only: bool = False,
    duplicate: bool = False,
    missing_period: bool = False,
    none_period: bool = False,
    missing_revenue: bool = False,
    none_revenue: bool = False,
    payments_first: bool = False,
    on_balance_sheet: bool = False,
    single_period: bool = False,
    payments_concept: str = "payments_for_ppe",
    with_cfo: bool = True,
    cfo=(80.0, 90.0),
    duplicate_cfo: bool = False,
    missing_cfo_period: bool = False,
    none_cfo_period: bool = False,
    cfo_concept: str = "",
):
    if single_period:
        d1 = date(2025, 12, 31)
        periods = [FinancialPeriod(end_date=d1, label="FY2025")]

        def vals(a, b=None):
            return {d1: a}

        pay_vals = (payments[0],)
        rev_vals = (revenue[0],)
    else:
        d1, d2 = P1, P2
        periods = [
            FinancialPeriod(end_date=d1, label="FY2024"),
            FinancialPeriod(end_date=d2, label="FY2025"),
        ]

        def vals(a, b=None):
            return {d1: a, d2: b if b is not None else a}

        pay_vals = payments
        rev_vals = revenue

    concept = "" if label_only else payments_concept
    if missing_period and not single_period:
        pay_values = {P1: pay_vals[0]}
    elif none_period and not single_period:
        pay_values = {P1: pay_vals[0], P2: None}
    elif single_period:
        pay_values = vals(pay_vals[0])
    else:
        pay_values = vals(*pay_vals)

    pay_item = _li(
        "Payments for property, plant and equipment",
        pay_values,
        concept=concept,
    )

    if missing_revenue and not single_period:
        rev_values = {P1: rev_vals[0]}
    elif none_revenue and not single_period:
        rev_values = {P1: rev_vals[0], P2: None}
    elif single_period:
        rev_values = vals(rev_vals[0])
    else:
        rev_values = vals(*rev_vals)

    cash, ar, ap, bank = 50.0, 40.0, 30.0, 20.0
    eq0 = cash + ar - ap - bank
    eq1 = eq0
    bs = [
        _li("Cash and cash equivalents", vals(cash, cash)),
        _li("Trade receivables", vals(ar, ar)),
        _li("Trade payables", vals(ap, ap)),
        _li("Bank borrowings", vals(bank, bank)),
        _li("Total equity", vals(eq0, eq1)),
    ]
    if single_period:
        is_rows = [
            _li("Revenue", rev_values, concept="revenue"),
            _li("Finance costs", vals(-10)),
            _li("Finance income", vals(0)),
            _li("Profit before tax", vals(200)),
            _li("Income tax expense", vals(-30)),
            _li("Profit for the year", vals(170)),
        ]
        cfo_vals = (cfo[0],)
    else:
        is_rows = [
            _li("Revenue", rev_values, concept="revenue"),
            _li("Finance costs", vals(-10, -11)),
            _li("Finance income", vals(0, 0)),
            _li("Profit before tax", vals(200, 220)),
            _li("Income tax expense", vals(-30, -33)),
            _li("Profit for the year", vals(170, 187)),
        ]
        cfo_vals = cfo

    if missing_cfo_period and not single_period:
        cfo_values = {P1: cfo_vals[0]}
    elif none_cfo_period and not single_period:
        cfo_values = {P1: cfo_vals[0], P2: None}
    elif single_period:
        cfo_values = vals(cfo_vals[0])
    else:
        cfo_values = vals(*cfo_vals)

    cf: list[LineItem] = []
    if with_cfo:
        cf.append(
            _li(
                "Net cash from operating activities",
                cfo_values,
                concept=cfo_concept,
            )
        )
        if duplicate_cfo and not on_balance_sheet:
            cf.append(
                _li(
                    "Net cash from operating activities duplicate",
                    vals(*cfo_vals) if not single_period else vals(cfo_vals[0]),
                    concept=cfo_concept or "operating_cash_flow",
                )
            )

    if with_payments:
        target = bs if on_balance_sheet else cf
        if payments_first:
            target.insert(0, pay_item)
        else:
            target.append(pay_item)
        if duplicate and not on_balance_sheet:
            cf.append(
                _li(
                    "Payments for PPE duplicate",
                    vals(*pay_vals) if not single_period else vals(pay_vals[0]),
                    concept=concept,
                )
            )

    return StandardizedFinancials(
        ticker="CAPEX",
        company_name="Capex Co",
        currency="HKD",
        units="HKD mn",
        jurisdiction="HK",
        periods=periods,
        income_statement=is_rows,
        balance_sheet=bs,
        cash_flow=cf,
    )


def _anchor(fin: StandardizedFinancials):
    return compute_anchor(fin, list(canonical_fiscal_periods(fin)))


def test_catalog_expansion_five_periods():
    periods = [date(y, 12, 31) for y in range(2021, 2026)]
    assert len(CAPEX_COMPONENT_CATALOG) == 4
    assert [f.order for f in CAPEX_COMPONENT_CATALOG] == [120, 121, 124, 125]
    specs = expand_capex_specs(periods, start_order=1000)
    assert len(specs) == 10
    assert {s.family_id for s in specs} == {"ppe_capex", "ppe_capex_to_revenue"}
    assert all(s.semantic_key.startswith("capex.") for s in specs)
    cash_specs = expand_capex_specs(
        periods, start_order=1000, include_operating_cash=True
    )
    assert len(cash_specs) == 20
    assert {s.family_id for s in cash_specs} == {
        "ppe_capex",
        "ppe_capex_to_revenue",
        "cash_after_ppe_capex",
        "cash_after_ppe_capex_to_revenue",
    }
    with pytest.raises(ValueError, match="duplicate"):
        expand_capex_specs([periods[0], periods[0]], start_order=1)
    with pytest.raises(ValueError, match="chronological"):
        expand_capex_specs(list(reversed(periods)), start_order=1)


def test_unique_concept_resolution_and_renamed_label():
    fin = _tiny()
    item = resolve_capex_source(fin)
    assert item is not None
    assert item.concept == "payments_for_ppe"
    assert item.label == "Payments for property, plant and equipment"
    assert capex_applicable(fin)
    avail = capex_availability(fin)
    assert avail.payments_for_ppe is True and avail.ambiguous is False
    assert avail.operating_cash_flow is True
    assert avail.operating_cash_flow_ambiguous is False
    assert cash_after_ppe_capex_applicable(fin)

    renamed = _tiny()
    renamed.cash_flow[-1] = _li(
        "Purchase of fixed assets (renamed)",
        {P1: -100.0, P2: -120.0},
        concept="payments_for_ppe",
    )
    assert resolve_capex_source(renamed) is not None
    assert resolve_capex_source(renamed).label == "Purchase of fixed assets (renamed)"


def test_reordered_rows_still_resolve():
    first = _tiny(payments_first=True)
    last = _tiny(payments_first=False)
    assert resolve_capex_source(first) is not None
    assert resolve_capex_source(last) is not None
    assert resolve_capex_source(first).values == resolve_capex_source(last).values


def test_absent_label_only_duplicate_wrong_statement():
    absent = _tiny(with_payments=False)
    assert not capex_applicable(absent)
    assert resolve_capex_source(absent) is None
    assert capex_availability(absent).payments_for_ppe is False
    with pytest.raises(MissingLineError):
        compute_capex_series(absent, list(canonical_fiscal_periods(absent)), _anchor(absent))
    assert ReferenceModelBuilder(absent).capex_specs == ()

    label_only = _tiny(label_only=True)
    assert not capex_applicable(label_only)
    assert resolve_capex_source(label_only) is None
    assert (
        resolve_line(
            label_only.cash_flow, "payments_for_ppe", required=False
        ).item
        is None
    )
    assert ReferenceModelBuilder(label_only).capex_specs == ()

    dup = _tiny(duplicate=True)
    avail = capex_availability(dup)
    assert avail.ambiguous is True
    assert avail.payments_for_ppe is False
    assert not capex_applicable(dup)
    assert resolve_capex_source(dup) is None
    with pytest.raises(AmbiguousLineError):
        resolve_line(dup.cash_flow, "payments_for_ppe", required=False)
    assert ReferenceModelBuilder(dup).capex_specs == ()

    wrong = _tiny(on_balance_sheet=True)
    assert not capex_applicable(wrong)
    assert resolve_capex_source(wrong) is None
    assert (
        resolve_line(wrong.balance_sheet, "payments_for_ppe", required=False).item
        is not None
    )


def test_sign_conversion_ratios_positive_negative_mixed_zero():
    negative = _tiny(payments=(-50.0, -80.0), revenue=(1000.0, 2000.0))
    series_neg = compute_capex_series(
        negative, list(canonical_fiscal_periods(negative)), _anchor(negative)
    )
    assert series_neg.payments_reported == (-50.0, -80.0)
    assert series_neg.ppe_capex == (50.0, 80.0)
    assert series_neg.ppe_capex_to_revenue[0] == pytest.approx(50.0 / 1000.0)
    assert series_neg.ppe_capex_to_revenue[1] == pytest.approx(80.0 / 2000.0)

    positive = _tiny(payments=(50.0, 80.0))
    series_pos = compute_capex_series(
        positive, list(canonical_fiscal_periods(positive)), _anchor(positive)
    )
    assert series_pos.payments_reported == (50.0, 80.0)
    assert series_pos.ppe_capex == (-50.0, -80.0)
    assert series_pos.ppe_capex_to_revenue[0] == pytest.approx(-50.0 / 1000.0)

    mixed = _tiny(payments=(-50.0, 80.0))
    series_mix = compute_capex_series(
        mixed, list(canonical_fiscal_periods(mixed)), _anchor(mixed)
    )
    assert series_mix.payments_reported == (-50.0, 80.0)
    assert series_mix.ppe_capex == (50.0, -80.0)

    zero = _tiny(payments=(0.0, -10.0))
    series_z = compute_capex_series(
        zero, list(canonical_fiscal_periods(zero)), _anchor(zero)
    )
    assert series_z.payments_reported == (0.0, -10.0)
    assert series_z.ppe_capex == (0.0, 10.0)
    assert series_z.ppe_capex_to_revenue[0] == pytest.approx(0.0)

    mapped = capex_expected_series(series_neg)
    assert set(mapped) == {f.id for f in CAPEX_COMPONENT_CATALOG}


def test_zero_revenue_undefined_ratio():
    fin = _tiny(payments=(-50.0, -80.0), revenue=(0.0, 2000.0))
    series = compute_capex_series(
        fin, list(canonical_fiscal_periods(fin)), _anchor(fin)
    )
    assert series.ppe_capex == (50.0, 80.0)
    assert series.ppe_capex_to_revenue[0] == UNDEFINED_RATIO
    assert series.ppe_capex_to_revenue[1] == pytest.approx(80.0 / 2000.0)
    assert series.operating_cash_flow == (80.0, 90.0)
    assert series.cash_after_ppe_capex == (30.0, 10.0)
    assert series.cash_after_ppe_capex_to_revenue[0] == UNDEFINED_RATIO
    assert series.cash_after_ppe_capex_to_revenue[1] == pytest.approx(10.0 / 2000.0)


def test_missing_none_inputs_raise():
    missing = _tiny(missing_period=True)
    periods = list(canonical_fiscal_periods(missing))
    # Revenue is complete; payments missing for P2.
    anchor = compute_anchor(missing, periods)
    with pytest.raises(MissingHistoricalValueError):
        compute_capex_series(missing, periods, anchor)

    none_period = _tiny(none_period=True)
    periods_n = list(canonical_fiscal_periods(none_period))
    anchor_n = compute_anchor(none_period, periods_n)
    with pytest.raises(MissingHistoricalValueError):
        compute_capex_series(none_period, periods_n, anchor_n)

    missing_rev = _tiny(missing_revenue=True)
    periods_r = list(canonical_fiscal_periods(missing_rev))
    with pytest.raises(MissingHistoricalValueError):
        compute_anchor(missing_rev, periods_r)

    none_rev = _tiny(none_revenue=True)
    periods_nr = list(canonical_fiscal_periods(none_rev))
    with pytest.raises(MissingHistoricalValueError):
        compute_anchor(none_rev, periods_nr)

    missing_cfo = _tiny(missing_cfo_period=True)
    periods_c = list(canonical_fiscal_periods(missing_cfo))
    with pytest.raises(MissingHistoricalValueError):
        compute_capex_series(
            missing_cfo, periods_c, compute_anchor(missing_cfo, periods_c)
        )
    none_cfo = _tiny(none_cfo_period=True)
    periods_cn = list(canonical_fiscal_periods(none_cfo))
    with pytest.raises(MissingHistoricalValueError):
        compute_capex_series(
            none_cfo, periods_cn, compute_anchor(none_cfo, periods_cn)
        )


def test_period_ordering_and_unchanged_input():
    fin = _tiny(payments=(-10.0, -20.0))
    reverse = [P2, P1]
    anchor = compute_anchor(fin, reverse)
    series = compute_capex_series(fin, reverse, anchor)
    assert series.payments_reported == (-20.0, -10.0)
    assert series.ppe_capex == (20.0, 10.0)

    before = copy.deepcopy(resolve_capex_source(fin).values)
    compute_capex_series(fin, [P1, P2], compute_anchor(fin, [P1, P2]))
    after = resolve_capex_source(fin).values
    assert after == before


def _dupont_row_by_label(ws, label: str) -> int:
    for row in range(1, (ws.max_row or 1) + 1):
        if ws.cell(row=row, column=1).value == label:
            return row
    raise AssertionError(f"ALT DuPont label not found: {label!r}")


def test_workbook_gating_structure_formulas_notes_and_check(tmp_path):
    data = _tiny()
    trainer, answer = build_training_workbook(data, tmp_path / "CAPEX_BASE.xlsx")
    builder = ReferenceModelBuilder(data)
    assert builder.capex_series is not None
    assert len(builder.capex_specs) == 8
    assert {s.family_id for s in builder.capex_specs} == {
        "ppe_capex",
        "ppe_capex_to_revenue",
        "cash_after_ppe_capex",
        "cash_after_ppe_capex_to_revenue",
    }

    smap = load_semantic_map(answer)
    capex_comps = [
        c
        for c in smap.all_ordered()
        if c.family_id in {f.id for f in CAPEX_COMPONENT_CATALOG}
    ]
    assert len(capex_comps) == 8

    for path in (trainer, answer):
        wb = load_workbook(path, data_only=False)
        ws = wb["ALT DuPont"]
        assert any(
            ws.cell(r, 1).value == "PP&E CAPEX CONTEXT"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        assert any(
            ws.cell(r, 1).value == "OPERATING CASH AFTER PP&E CAPEX"
            for r in range(1, (ws.max_row or 1) + 1)
        )
        reported_row = _dupont_row_by_label(
            ws, "Payments for Property, Plant & Equipment (reported)"
        )
        capex_row = _dupont_row_by_label(ws, "PP&E Capex (−reported)")
        ratio_row = _dupont_row_by_label(ws, "PP&E Capex / Revenue")
        cfo_row = _dupont_row_by_label(ws, "Operating Cash Flow (reported)")
        cash_row = _dupont_row_by_label(ws, "Operating cash after PP&E capex")
        cash_ratio_row = _dupont_row_by_label(
            ws, "Operating cash after PP&E capex / Revenue"
        )
        # Source link populated in both workbooks.
        assert isinstance(ws.cell(reported_row, 2).value, str)
        assert ws.cell(reported_row, 2).value.startswith("=")
        assert isinstance(ws.cell(cfo_row, 2).value, str)
        assert ws.cell(cfo_row, 2).value.startswith("=")
        for col in (2, 3):
            reported = ws.cell(reported_row, column=col)
            capex_cell = ws.cell(capex_row, column=col)
            ratio_cell = ws.cell(ratio_row, column=col)
            cash_cell = ws.cell(cash_row, column=col)
            cash_ratio_cell = ws.cell(cash_ratio_row, column=col)
            if path == answer:
                assert isinstance(capex_cell.value, str) and capex_cell.value.startswith(
                    "="
                )
                assert capex_cell.value.replace(" ", "").startswith("=-")
                assert "IF(" in str(ratio_cell.value) and "NA()" in str(ratio_cell.value)
                assert capex_cell.comment is not None
                assert (capex_cell.comment.text or "").strip()
                assert ratio_cell.comment is not None
                assert (ratio_cell.comment.text or "").strip()
                assert isinstance(cash_cell.value, str) and cash_cell.value.startswith(
                    "="
                )
                assert "free cash flow" in (cash_cell.comment.text or "").lower()
                assert "maintenance" in (cash_cell.comment.text or "").lower()
                assert cash_ratio_cell.comment is not None
                assert (cash_ratio_cell.comment.text or "").strip()
            else:
                assert capex_cell.value is None
                assert ratio_cell.value is None
                assert cash_cell.value is None
                assert cash_ratio_cell.value is None
                assert capex_cell.comment is None
                assert ratio_cell.comment is None
                assert cash_cell.comment is None
                assert cash_ratio_cell.comment is None
                assert reported.value is not None
                assert ws.cell(cfo_row, column=col).value is not None
        wb.close()

    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total
    assert blank.correct == 0

    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    filled = check_workbook(trainer)
    assert filled.correct == filled.total
    assert filled.incorrect == 0
    assert filled.blank == 0

    bad = next(c for c in capex_comps if c.family_id == "ppe_capex")
    _inject_formula_and_cached_value(
        trainer,
        bad.tab,
        bad.cell,
        formula="=999",
        cached_value=999.0,
    )
    bad_summary = check_workbook(trainer)
    assert bad_summary.incorrect >= 1
    dumped = repr(bad_summary)
    assert "=999" not in dumped
    assert "PP&E Capex" not in dumped


def test_workbook_renamed_reordered_source_links(tmp_path):
    for payments_first in (False, True):
        fin = _tiny(payments_first=payments_first)
        fin.cash_flow[
            next(
                i
                for i, item in enumerate(fin.cash_flow)
                if (item.concept or "") == "payments_for_ppe"
            )
        ] = _li(
            "Purchase of fixed assets (renamed)",
            {P1: -100.0, P2: -120.0},
            concept="payments_for_ppe",
        )
        trainer, answer = build_training_workbook(
            fin, tmp_path / f"CAPEX_REORDER_{payments_first}.xlsx"
        )
        idx = next(
            i
            for i, item in enumerate(fin.cash_flow)
            if (item.concept or "") == "payments_for_ppe"
        )
        expected_row = 7 + idx
        wb = load_workbook(answer, data_only=False)
        ws = wb["ALT DuPont"]
        reported_row = _dupont_row_by_label(
            ws, "Payments for Property, Plant & Equipment (reported)"
        )
        level_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
        assert f"'CashFlowStatement'!B{expected_row}" in level_f or (
            f"'Cash Flow Statement'!B{expected_row}" in level_f
        )
        wb.close()


def test_undefined_ratio_check_accepts_na(tmp_path):
    data = _tiny(payments=(-50.0, -80.0), revenue=(0.0, 2000.0))
    trainer, answer = build_training_workbook(data, tmp_path / "CAPEX_NA.xlsx")
    smap = load_semantic_map(answer)
    undef = [
        c
        for c in smap.all_ordered()
        if c.family_id
        in {"ppe_capex_to_revenue", "cash_after_ppe_capex_to_revenue"}
        and c.expected_value == UNDEFINED_RATIO
    ]
    assert undef
    wb = load_workbook(trainer, data_only=False)
    for comp in smap.all_ordered():
        row, col = parse_cell_ref(comp.cell)
        wb[comp.tab].cell(row=row, column=col).value = comp.formula
    wb.save(trainer)
    wb.close()
    # Inject cached #N/A for the undefined-ratio cell so Check can validate.
    for comp in undef:
        _inject_formula_and_cached_value(
            trainer,
            comp.tab,
            comp.cell,
            formula=comp.formula,
            cached_value="#N/A",
        )
    filled = check_workbook(trainer)
    assert filled.incorrect == 0
    assert filled.blank == 0
    assert filled.correct == filled.total


def test_demo_omits_and_practice_counts_unchanged(tmp_path):
    data = _ingest_demo()
    assert not capex_applicable(data)
    builder = ReferenceModelBuilder(data)
    assert builder.capex_series is None
    assert builder.capex_specs == ()
    trainer, answer = build_training_workbook(data, tmp_path / "CAPEX_DEMO.xlsx")
    smap = load_semantic_map(answer)
    assert len(group_components_by_family(smap)) == 74
    assert len(smap.all_ordered()) == 312
    assert check_workbook(trainer).blank == 312


def test_fast_retailing_fy2021_fy2025_capex_anchors_and_ratios():
    payload = json.loads(STD_JSON.read_text(encoding="utf-8"))
    fin = standardized_from_payload(payload)
    assert capex_applicable(fin)
    periods = list(canonical_fiscal_periods(fin))
    assert [p.isoformat() for p in periods] == [
        "2021-08-31",
        "2022-08-31",
        "2023-08-31",
        "2024-08-31",
        "2025-08-31",
    ]
    anchor = compute_anchor(fin, periods)
    series = compute_capex_series(fin, periods, anchor)
    assert series.payments_reported == (
        -56500.0,
        -51271.0,
        -61764.0,
        -73728.0,
        -135535.0,
    )
    assert series.ppe_capex == (56500.0, 51271.0, 61764.0, 73728.0, 135535.0)
    revenues = (2132992.0, 2301122.0, 2766557.0, 3103836.0, 3400539.0)
    assert list(anchor.historical.revenue) == list(revenues)
    for j, rev in enumerate(revenues):
        assert series.ppe_capex_to_revenue[j] == pytest.approx(
            series.ppe_capex[j] / rev
        )

    builder = ReferenceModelBuilder(fin)
    assert len(builder.capex_specs) == 20
    assert len(builder.expected_specs) == 577


def test_capital_expenditures_alias_resolution_sign_zero_and_missing():
    alias = _tiny(payments_concept="capital_expenditures")
    item = resolve_capex_source(alias)
    assert item is not None
    assert item.concept == "capital_expenditures"
    assert item.label == "Payments for property, plant and equipment"
    assert capex_applicable(alias)
    avail = capex_availability(alias)
    assert avail.payments_for_ppe is True and avail.ambiguous is False

    series = compute_capex_series(
        alias, list(canonical_fiscal_periods(alias)), _anchor(alias)
    )
    assert series.payments_reported == (-100.0, -120.0)
    assert series.ppe_capex == (100.0, 120.0)
    assert series.ppe_capex_to_revenue[0] == pytest.approx(100.0 / 1000.0)

    zero = _tiny(payments=(0.0, -10.0), payments_concept="capital_expenditures")
    series_z = compute_capex_series(
        zero, list(canonical_fiscal_periods(zero)), _anchor(zero)
    )
    assert series_z.payments_reported == (0.0, -10.0)
    assert series_z.ppe_capex == (0.0, 10.0)
    assert series_z.ppe_capex_to_revenue[0] == pytest.approx(0.0)

    positive = _tiny(payments=(50.0, 80.0), payments_concept="capital_expenditures")
    series_pos = compute_capex_series(
        positive, list(canonical_fiscal_periods(positive)), _anchor(positive)
    )
    assert series_pos.payments_reported == (50.0, 80.0)
    assert series_pos.ppe_capex == (-50.0, -80.0)

    missing = _tiny(missing_period=True, payments_concept="capital_expenditures")
    periods = list(canonical_fiscal_periods(missing))
    with pytest.raises(MissingHistoricalValueError):
        compute_capex_series(missing, periods, compute_anchor(missing, periods))

    none_period = _tiny(none_period=True, payments_concept="capital_expenditures")
    periods_n = list(canonical_fiscal_periods(none_period))
    with pytest.raises(MissingHistoricalValueError):
        compute_capex_series(
            none_period, periods_n, compute_anchor(none_period, periods_n)
        )


def test_capital_expenditures_alias_ambiguity_statement_boundary_and_label_only():
    both = _tiny(payments_concept="payments_for_ppe")
    both.cash_flow.append(
        _li(
            "Purchase of property and equipment",
            {P1: -100.0, P2: -120.0},
            concept="capital_expenditures",
        )
    )
    avail = capex_availability(both)
    assert avail.ambiguous is True
    assert avail.payments_for_ppe is False
    assert not capex_applicable(both)
    assert resolve_capex_source(both) is None
    with pytest.raises(AmbiguousLineError):
        resolve_line(both.cash_flow, "payments_for_ppe", required=False)
    assert ReferenceModelBuilder(both).capex_specs == ()

    dup_alias = _tiny(duplicate=True, payments_concept="capital_expenditures")
    assert capex_availability(dup_alias).ambiguous is True
    assert not capex_applicable(dup_alias)

    wrong = _tiny(on_balance_sheet=True, payments_concept="capital_expenditures")
    assert not capex_applicable(wrong)
    assert resolve_capex_source(wrong) is None
    assert (
        resolve_line(wrong.balance_sheet, "payments_for_ppe", required=False).item
        is not None
    )

    label_only = _tiny(label_only=True)
    assert not capex_applicable(label_only)
    assert resolve_capex_source(label_only) is None


def test_capital_expenditures_alias_round_trip_preserves_stored_identity():
    fin = _tiny(payments_concept="capital_expenditures")
    original = resolve_capex_source(fin)
    assert original is not None
    restored = standardized_from_payload(standardized_to_payload(fin))
    item = resolve_capex_source(restored)
    assert item is not None
    assert item.concept == "capital_expenditures"
    assert item.label == original.label
    assert item.values == original.values
    assert capex_applicable(restored)
    series = compute_capex_series(
        restored, list(canonical_fiscal_periods(restored)), _anchor(restored)
    )
    assert series.payments_reported == (-100.0, -120.0)
    assert series.ppe_capex == (100.0, 120.0)


def test_capital_expenditures_alias_python_excel_source_identity(tmp_path):
    for payments_first in (False, True):
        fin = _tiny(
            payments_first=payments_first,
            payments_concept="capital_expenditures",
        )
        resolved = resolve_line(fin.cash_flow, "payments_for_ppe", required=True)
        assert resolved.item is not None
        assert resolved.item.concept == "capital_expenditures"
        expected_row = 7 + resolved.index
        python_payments = compute_capex_series(
            fin, list(canonical_fiscal_periods(fin)), _anchor(fin)
        ).payments_reported
        assert python_payments == (-100.0, -120.0)

        trainer, answer = build_training_workbook(
            fin, tmp_path / f"CAPEX_ALIAS_{payments_first}.xlsx"
        )
        builder = ReferenceModelBuilder(fin)
        assert builder.capex_series is not None
        assert builder.capex_series.payments_reported == python_payments
        assert len(builder.capex_specs) == 8

        wb = load_workbook(answer, data_only=False)
        ws = wb["ALT DuPont"]
        reported_row = _dupont_row_by_label(
            ws, "Payments for Property, Plant & Equipment (reported)"
        )
        capex_row = _dupont_row_by_label(ws, "PP&E Capex (−reported)")
        level_f = str(ws.cell(reported_row, 2).value).replace(" ", "")
        assert f"'CashFlowStatement'!B{expected_row}" in level_f or (
            f"'Cash Flow Statement'!B{expected_row}" in level_f
        )
        capex_f = str(ws.cell(capex_row, 2).value).replace(" ", "")
        assert capex_f.startswith("=-")
        wb.close()

        blank = check_workbook(trainer)
        assert blank.incorrect == 0
        assert blank.blank == blank.total
        assert blank.correct == 0

        smap = load_semantic_map(answer)
        wb = load_workbook(trainer, data_only=False)
        for comp in smap.all_ordered():
            row, col = parse_cell_ref(comp.cell)
            wb[comp.tab].cell(row=row, column=col).value = comp.formula
        wb.save(trainer)
        wb.close()
        filled = check_workbook(trainer)
        assert filled.correct == filled.total
        assert filled.incorrect == 0
        assert filled.blank == 0


def test_absent_ambiguous_cfo_preserves_existing_capex(tmp_path):
    absent = _tiny(with_cfo=False)
    assert capex_applicable(absent)
    avail = capex_availability(absent)
    assert avail.operating_cash_flow is False
    assert avail.operating_cash_flow_ambiguous is False
    assert not cash_after_ppe_capex_applicable(absent)
    assert resolve_operating_cash_source(absent) is None
    series = compute_capex_series(
        absent, list(canonical_fiscal_periods(absent)), _anchor(absent)
    )
    assert series.ppe_capex == (100.0, 120.0)
    assert series.operating_cash_flow is None
    assert series.cash_after_ppe_capex is None
    assert series.cash_after_ppe_capex_to_revenue is None
    mapped = capex_expected_series(series)
    assert set(mapped) == {"ppe_capex", "ppe_capex_to_revenue"}
    builder = ReferenceModelBuilder(absent)
    assert len(builder.capex_specs) == 4
    assert {s.family_id for s in builder.capex_specs} == {
        "ppe_capex",
        "ppe_capex_to_revenue",
    }
    trainer, answer = build_training_workbook(absent, tmp_path / "CAPEX_NO_CFO.xlsx")
    smap = load_semantic_map(answer)
    families = {c.family_id for c in smap.all_ordered()}
    assert "ppe_capex" in families
    assert "cash_after_ppe_capex" not in families
    wb = load_workbook(answer, data_only=False)
    ws = wb["ALT DuPont"]
    labels = [
        ws.cell(r, 1).value for r in range(1, (ws.max_row or 1) + 1)
    ]
    assert "PP&E CAPEX CONTEXT" in labels
    assert "OPERATING CASH AFTER PP&E CAPEX" not in labels
    wb.close()
    blank = check_workbook(trainer)
    assert blank.incorrect == 0
    assert blank.blank == blank.total

    amb = _tiny(duplicate_cfo=True, cfo_concept="operating_cash_flow")
    avail_a = capex_availability(amb)
    assert avail_a.payments_for_ppe is True
    assert avail_a.ambiguous is False
    assert avail_a.operating_cash_flow is False
    assert avail_a.operating_cash_flow_ambiguous is True
    assert capex_applicable(amb)
    assert not cash_after_ppe_capex_applicable(amb)
    with pytest.raises(AmbiguousLineError):
        resolve_line(amb.cash_flow, "operating_cash_flow", required=False)
    series_a = compute_capex_series(
        amb, list(canonical_fiscal_periods(amb)), _anchor(amb)
    )
    assert series_a.ppe_capex == (100.0, 120.0)
    assert series_a.cash_after_ppe_capex is None
    try:
        builder_a = ReferenceModelBuilder(amb)
    except AmbiguousLineError:
        builder_a = None
    else:
        assert {s.family_id for s in builder_a.capex_specs} == {
            "ppe_capex",
            "ppe_capex_to_revenue",
        }


def test_cash_after_ppe_capex_signs_zero_reversal_and_immutability():
    negative_cfo = _tiny(cfo=(-40.0, -10.0), payments=(-50.0, -80.0))
    series_n = compute_capex_series(
        negative_cfo,
        list(canonical_fiscal_periods(negative_cfo)),
        _anchor(negative_cfo),
    )
    assert series_n.operating_cash_flow == (-40.0, -10.0)
    assert series_n.ppe_capex == (50.0, 80.0)
    assert series_n.cash_after_ppe_capex == (-90.0, -90.0)

    zero_cfo = _tiny(cfo=(0.0, 90.0), payments=(-50.0, -80.0))
    series_z = compute_capex_series(
        zero_cfo, list(canonical_fiscal_periods(zero_cfo)), _anchor(zero_cfo)
    )
    assert series_z.operating_cash_flow == (0.0, 90.0)
    assert series_z.cash_after_ppe_capex == (-50.0, 10.0)
    assert series_z.cash_after_ppe_capex_to_revenue[0] == pytest.approx(-50.0 / 1000.0)

    reversal = _tiny(payments=(50.0, -80.0), cfo=(100.0, 90.0))
    series_r = compute_capex_series(
        reversal, list(canonical_fiscal_periods(reversal)), _anchor(reversal)
    )
    assert series_r.ppe_capex == (-50.0, 80.0)
    assert series_r.cash_after_ppe_capex == (150.0, 10.0)

    fin = _tiny(cfo=(80.0, 90.0), payments=(-10.0, -20.0))
    before_pay = copy.deepcopy(resolve_capex_source(fin).values)
    before_cfo = copy.deepcopy(resolve_operating_cash_source(fin).values)
    compute_capex_series(fin, [P1, P2], compute_anchor(fin, [P1, P2]))
    assert resolve_capex_source(fin).values == before_pay
    assert resolve_operating_cash_source(fin).values == before_cfo


def test_fast_retailing_cash_after_ppe_capex_independent_arithmetic():
    payload = json.loads(STD_JSON.read_text(encoding="utf-8"))
    fin = standardized_from_payload(payload)
    periods = list(canonical_fiscal_periods(fin))
    cfo_item = resolve_operating_cash_source(fin)
    pay_item = resolve_capex_source(fin)
    assert cfo_item is not None and pay_item is not None
    independent = []
    margins = []
    revenues = (2132992.0, 2301122.0, 2766557.0, 3103836.0, 3400539.0)
    for j, period in enumerate(periods):
        cfo = required_period_value(cfo_item, period, field="operating_cash_flow")
        pay = required_period_value(pay_item, period, field="payments_for_ppe")
        cash_after = cfo - (-pay)
        independent.append(cash_after)
        margins.append(cash_after / revenues[j])
    assert independent == [372468.0, 379546.0, 401452.0, 577793.0, 445083.0]
    series = compute_capex_series(fin, periods, compute_anchor(fin, periods))
    assert series.cash_after_ppe_capex == tuple(independent)
    for j, margin in enumerate(margins):
        assert series.cash_after_ppe_capex_to_revenue[j] == pytest.approx(margin)
    assert len(ReferenceModelBuilder(fin).capex_specs) == 20
    assert len(ReferenceModelBuilder(fin).expected_specs) == 577


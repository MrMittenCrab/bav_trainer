"""Tests for the read-only reference workbook audit utility."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

from scripts.audit_reference_workbook import audit_workbook

ROOT = Path(__file__).resolve().parents[2]
GOOGL = ROOT / "example" / "GOOGL_Demo_Integrated_Financials.xlsx"


def test_audit_workbook_synthetic_inventory(tmp_path):
    path = tmp_path / "synthetic.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Income Statement"
    ws["A1"] = "Revenue"
    ws["B1"] = 100
    ws["C1"] = "=B1*1.1"
    hidden = wb.create_sheet("Model_Base")
    hidden["A1"] = "Deferred example"
    hidden.sheet_state = "hidden"
    wb.save(path)
    wb.close()

    before = path.read_bytes()
    result = audit_workbook(path)
    after = path.read_bytes()
    assert before == after

    assert result["workbook"] == path.name
    assert [s["name"] for s in result["sheets"]] == [
        "Income Statement",
        "Model_Base",
    ]
    assert result["sheets"][0]["state"] == "visible"
    assert result["sheets"][1]["state"] == "hidden"
    assert result["sheets"][0]["formula_cells"] == 1
    assert result["sheets"][0]["nonempty_cells"] == 3
    assert result["sheets"][0]["representative_labels"] == ["Revenue"]
    assert result["sheets"][1]["representative_labels"] == ["Deferred example"]
    assert result["sheets"][1]["formula_cells"] == 0
    assert result["sheets"][1]["nonempty_cells"] == 1


def test_googl_reference_workbook_is_auditable_read_only():
    assert GOOGL.is_file()
    before = GOOGL.read_bytes()
    result = audit_workbook(GOOGL)
    after = GOOGL.read_bytes()
    assert before == after
    assert result["workbook"] == GOOGL.name
    assert len(result["sheets"]) >= 1
    names = [s["name"] for s in result["sheets"]]
    assert len(names) == len(set(names))
    for sheet in result["sheets"]:
        assert set(sheet) >= {
            "name",
            "state",
            "max_row",
            "max_column",
            "nonempty_cells",
            "formula_cells",
            "representative_labels",
        }
        assert isinstance(sheet["representative_labels"], list)
        assert len(sheet["representative_labels"]) <= 40


def test_googl_historical_reference_doc_separates_evidence_from_trainer():
    doc = (ROOT / "docs" / "GOOGL_HISTORICAL_REFERENCE.md").read_text(encoding="utf-8")
    assert "Canonical Answer Key SHA-256:" not in doc
    assert "81faf2882d0df07ecf5def45695431c1935b4f7a94c1e017367596a063063896" in doc

    # Fixed-asset matrix row: GOOGL evidence stays source facts; Trainer names the module.
    assert "BS Property and equipment" in doc
    assert "CF depreciation" in doc or "CF Depreciation" in doc
    assert "Purchases of property and equipment" in doc
    matrix_section = doc.split("## Historical capability gap matrix", 1)[1].split(
        "## Explicitly deferred", 1
    )[0]
    fa_line = next(
        line
        for line in matrix_section.splitlines()
        if line.startswith("| PP&E / D&A / asset intensity |")
    )
    cols = [c.strip() for c in fa_line.strip("|").split("|")]
    googl_evidence, current_trainer = cols[1], cols[2]
    assert "FIXED-ASSET INTENSITY CONTEXT" not in googl_evidence
    assert "FIXED-ASSET INTENSITY CONTEXT" in current_trainer or "fixed-asset" in current_trainer.lower()

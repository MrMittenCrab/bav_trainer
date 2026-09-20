"""Project company resolution and verified, atomic current-snapshot publication.

Company metadata selects evidence locations, never accounting behavior. All
source additions still pass the existing source binding and admission pipeline.
"""
from __future__ import annotations

import ctypes
from dataclasses import dataclass
import json
import os
import re
from pathlib import Path
import sys
import tempfile

from .data.issuer_fiscal import (
    apply_issuer_fiscal_labels,
    issuer_fiscal_years_from_extracted,
)
from .data.standardized_io import standardized_from_payload

ROOT = Path(__file__).resolve().parents[1]
INPUT_ROOT = ROOT / 'build' / 'input'
OUTPUT_ROOT = ROOT / 'build' / 'output'
# These are project routing/admission settings, not analytical special cases.
PROJECTS = tuple(
    (item['name'], item['slug'], tuple(item['aliases']),
     tuple(item['admit_periods']), item['supplemental_facts'],
     item.get('strategy_disclosures'))
    for item in json.loads(Path(__file__).with_name('project_companies.json').read_text())
)


@dataclass(frozen=True)
class Company:
    name: str
    slug: str
    aliases: tuple[str, ...]
    admit_periods: tuple[str, ...]
    supplemental_facts: str | None
    strategy_disclosures: str | None
    input: Path
    output: Path

    @property
    def bav(self):
        return self.output / f'{self.name}_BAV.xlsx'

    @property
    def trainer(self):
        return self.output / f'{self.name}_BAV_Trainer.xlsx'

    @property
    def answer(self):
        """Backward-compatible alias — the BAV replaces the former Answer Key."""
        return self.bav


def resolve_company(query: str) -> Company:
    key = query.strip().casefold()
    matches = [p for p in PROJECTS if key in {s.casefold() for s in (p[0], p[1], *p[2])}]
    if len(matches) != 1:
        label = 'Ambiguous' if matches else 'Unknown'
        candidates = ', '.join(p[0] for p in (matches or PROJECTS))
        raise ValueError(f'{label} company {query!r}; candidates: {candidates}')
    name, slug, aliases, periods, facts, strategy = matches[0]
    return Company(
        name, slug, aliases, periods, facts, strategy,
        INPUT_ROOT / slug, OUTPUT_ROOT / slug,
    )


def check_company_output(query: str) -> int:
    """Resolve canonical output and check it without path arguments."""
    from .research.publish import verify_research_artifacts
    from .trainer.checker import check_workbook
    from .trainer.semantic_io import load_semantic_map, sidecar_paths
    company = resolve_company(query)
    if not company.bav.is_file():
        raise ValueError(
            f'No current build for {company.name}; run python -m bav build {company.name}'
        )
    for path in sidecar_paths(company.bav):
        if not path.is_file():
            raise ValueError(f'Missing required sidecar: {path.name}')
        json.loads(path.read_text(encoding='utf-8'))
    load_semantic_map(company.bav)
    research = company.output / 'research'
    if (research / f'{company.name}_Drivers.md').is_file():
        verify_research_artifacts(company.output, company.name)
    if company.trainer.is_file():
        summary = check_workbook(company.trainer)
        print(
            f'Checked {summary.total} practice cells: '
            f'{summary.correct} correct, {summary.incorrect} incorrect, '
            f'{summary.blank} blank.'
        )
        return 0 if summary.incorrect == 0 else 1
    print(f'Checked {company.name} output: {company.bav}')
    return 0


def current_workbook(query: str, *, answer: bool = False) -> Path:
    company = resolve_company(query)
    if answer:
        path = company.bav
        if not path.is_file():
            raise ValueError(
                f'No current build for {company.name}; run python -m bav build {company.name}'
            )
        return path
    if company.trainer.is_file():
        return company.trainer
    if company.bav.is_file():
        return company.bav
    raise ValueError(
        f'No current build for {company.name}; run python -m bav build {company.name}'
    )


def _write_json(path: Path, payload):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + '\n', encoding='utf-8')


def _publish_company_sidecars(staged: Path, bav_name: str) -> None:
    """Move generated sidecars into supporting/; do not copy persistent inputs."""
    supporting = staged / 'supporting'
    supporting.mkdir(parents=True, exist_ok=True)
    bav = staged / bav_name
    mapping = (
        (bav.with_suffix('.component_map.json'), supporting / 'component_map.json'),
        (bav.with_suffix('.assumptions.json'), supporting / 'assumptions.json'),
        (staged / 'rowmap.json', supporting / 'rowmap.json'),
    )
    for source, destination in mapping:
        if not source.is_file():
            raise ValueError(f'Missing required sidecar: {source.name}')
        source.replace(destination)


def _require_company_input(company: Company) -> Path:
    standardized = company.input / 'reconciled' / 'standardized.json'
    if not company.input.is_dir() or not standardized.is_file():
        raise ValueError(
            f'Canonical input missing for {company.name}; expected {standardized}'
        )
    extracted = company.input / 'extracted'
    source = company.input / 'source'
    if not extracted.is_dir() or not source.is_dir():
        raise ValueError(
            f'Canonical source/extracted input missing for {company.name}'
        )
    return standardized


def prepare_company_input(company: Company, staged: Path | None = None):
    """Load the accepted canonical model; do not republish inputs into output."""
    del staged
    path = _require_company_input(company)
    payload = json.loads(path.read_text(encoding='utf-8'))
    if not isinstance(payload, dict):
        raise ValueError(f'{path} must contain a JSON object')
    payload.pop('historical_strategy', None)
    fin = standardized_from_payload(payload, strict=True)
    mapping = issuer_fiscal_years_from_extracted(company.input / 'extracted')
    apply_issuer_fiscal_labels(fin, mapping, require_complete=True)
    if company.strategy_disclosures:
        from .data.historical_strategy import load_strategy_disclosures
        fin.historical_strategy = load_strategy_disclosures(ROOT / company.strategy_disclosures)
    return fin


def verify_staged(fin, bav: Path, assumptions=None, trainer: Path | None = None):
    """Verify the BAV independently of Trainer existence; optionally verify a Trainer."""
    from openpyxl import load_workbook
    from .engine.reference_model import ReferenceModelBuilder
    from .engine.build_contract import verify_complete_build
    from .engine.component_catalog import is_operating_kpi_source_identity
    from .engine.semantic_map import SemanticMap
    from .trainer.semantic_io import load_semantic_map, sidecar_paths
    from .trainer.checker import check_workbook
    for path in sidecar_paths(bav):
        if not path.is_file():
            raise ValueError(f'Missing required sidecar: {path.name}')
        json.loads(path.read_text())
    smap = load_semantic_map(bav)
    builder = ReferenceModelBuilder(fin, assumptions, current_snapshot=True)
    verify_complete_build(builder.expected_specs, smap)
    embedded = SemanticMap.from_workbook(bav)
    if [c.to_dict() for c in embedded.all_ordered()] != [c.to_dict() for c in smap.all_ordered()]:
        raise ValueError('Embedded and sidecar semantic maps differ')
    paths = [bav] if trainer is None else [bav, trainer]
    for path in paths:
        wb = load_workbook(path)
        try:
            if 'Build Status' not in wb:
                raise ValueError('Missing Build Status')
            if path == bav:
                if 'Overview' not in wb:
                    raise ValueError('Missing professional BAV opening')
                if 'Trainer' in wb:
                    raise ValueError('BAV must not contain a Trainer sheet')
                opening = ' '.join(
                    str(cell.value)
                    for row in wb['Overview'].iter_rows()
                    for cell in row
                    if cell.value is not None
                ).casefold()
                for term in ('trainer', 'answer key', 'exercise', 'practice', 'check'):
                    if term in opening:
                        raise ValueError(f'BAV opening contains exercise framing: {term}')
            for comp in smap.all_ordered():
                cell = wb[comp.tab][comp.cell]
                if is_operating_kpi_source_identity(comp):
                    if cell.value is None:
                        raise ValueError('Missing supplied KPI fact')
                elif path == trainer:
                    if cell.value is not None or cell.comment is not None or cell.fill.fgColor.rgb not in ('00FFFF00', 'FFFFFF00'):
                        raise ValueError('Trainer practice invariant failed')
                elif cell.value != comp.formula or cell.comment is None:
                    raise ValueError('BAV formula/Note invariant failed')
            if path == bav and any(c.fill.fgColor.type == 'rgb' and c.fill.fgColor.rgb in ('FFFFFF00', '00FFFF00') for ws in wb for row in ws for c in row):
                raise ValueError('BAV has yellow fill')
        finally:
            wb.close()
    if trainer is not None:
        result = check_workbook(trainer)
        if result.correct or result.incorrect or result.blank != result.total:
            raise ValueError('Pristine Trainer Check failed')


def _ensure_canonical_dirname(path: Path) -> None:
    """Force the on-disk directory name to the lowercase slug on case-insensitive volumes."""
    if not path.exists():
        return
    parent = path.parent
    desired = path.name
    actual = next(
        (
            item.name
            for item in parent.iterdir()
            if item.is_dir() and os.path.samefile(item, path)
        ),
        desired,
    )
    if actual == desired:
        return
    temporary = parent / f".{desired}.case-fix"
    os.rename(parent / actual, temporary)
    os.rename(temporary, parent / desired)


def _exchange_directories(staged: Path, current: Path):
    """One filesystem transaction: readers see either complete generation.

    Fail closed on platforms without directory exchange; no delete/rename gap.
    After exchange staged contains the old generation, removed by TemporaryDirectory.
    """
    libc = ctypes.CDLL(None, use_errno=True)
    if sys.platform == 'darwin':
        fn = libc.renamex_np
        fn.argtypes = [ctypes.c_char_p, ctypes.c_char_p, ctypes.c_uint]
        result = fn(os.fsencode(staged), os.fsencode(current), 2)  # RENAME_SWAP
    elif sys.platform.startswith('linux') and hasattr(libc, 'renameat2'):
        fn = libc.renameat2
        fn.argtypes = [ctypes.c_int, ctypes.c_char_p, ctypes.c_int, ctypes.c_char_p, ctypes.c_uint]
        result = fn(-100, os.fsencode(staged), -100, os.fsencode(current), 2)
    else:
        raise ValueError('Atomic directory replacement is unsupported on this platform')
    if result:
        number = ctypes.get_errno()
        raise OSError(number, os.strerror(number))


def build_company(company: Company, assumptions=None):
    from .trainer.workbook import build_bav_workbook
    from .build_status import status_rows
    from .trainer.semantic_io import load_semantic_map
    current = company.output
    from .__main__ import _validate_build_output
    _validate_build_output(company.bav, [])
    if current.is_symlink():
        raise ValueError('Current build directory must not be a symlink')
    current.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=f'.{company.name}-staging-', dir=current.parent) as raw:
        staged = Path(raw)
        fin = prepare_company_input(company, staged)
        bav = build_bav_workbook(fin, staged / company.bav.name, assumptions, current_snapshot=True)
        _publish_company_sidecars(staged, company.bav.name)
        verify_staged(fin, bav, assumptions)
        from .research.publish import publish_company_research
        publish_company_research(company.name, fin, staged)
        rows = status_rows(load_semantic_map(bav))
        _write_json(staged / 'supporting' / 'build_status.json', rows)
        if current.exists():
            _exchange_directories(staged, current)
        else:
            os.rename(staged, current)
    _ensure_canonical_dirname(current)
    # Retire only recognizable legacy workbook artifacts for this company.
    # Canonical-directory contents were replaced as a single generation above.
    pattern = re.compile(
        re.escape(company.name)
        + r'(?:_(?:Live|KPI|Preview|[0-9][0-9_-]*))?(?:_BAV_Trainer|_BAV|_Trainer|_Answer_Key)\.(?:xlsx|component_map\.json|assumptions\.json|trainer\.json)$'
    )
    for legacy in current.parent.iterdir():
        if pattern.fullmatch(legacy.name) and legacy.is_file():
            try:
                legacy.unlink()
            except OSError as exc:
                print(f'warning: built successfully but could not retire {legacy}: {exc}', file=sys.stderr)
    return rows

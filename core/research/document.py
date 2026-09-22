"""Word/PDF publication from canonical research Markdown and figures.

Uses pandoc as the Markdown parser, then python-docx and reportlab for
STYLE-controlled Word and PDF. Publication does not recalculate analysis.
"""
from __future__ import annotations

from dataclasses import dataclass
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import zipfile

from ..current_build import Company, resolve_company
from .drivers import drivers_filename, placeholder_filenames
from .publish import verify_research_artifacts
from .style import (
    CJK_FACE,
    LABEL_PT,
    LATIN_FACE,
    LINE_SPACING,
    RELATED_PT,
    SECTION_PT,
    ResolvedFonts,
    resolve_required_fonts,
)

BODY_PT = 10
HEADING_PT = 14
CAPTION_PT = LABEL_PT
PAGE_MARGIN_MM = 18
MM_PT = 72.0 / 25.4
PORTRAIT = (210 * MM_PT, 297 * MM_PT)
LANDSCAPE = (297 * MM_PT, 210 * MM_PT)
_LIBS_LOADED = False
_BAVCanvas = None
DEBUG_TERMS = (
    "fail-closed",
    "fail closed",
    "standardizedfinancials",
    "audit-only",
    "audit only",
    "supported_descriptively",
    "segment_bridge",
    "answer key",
)


@dataclass(frozen=True)
class PublishedDocuments:
    word: Path
    pdf: Path


@dataclass(frozen=True)
class Heading:
    level: int
    text: str


@dataclass(frozen=True)
class Body:
    text: str


@dataclass(frozen=True)
class FigureBlock:
    path: Path
    caption: str


@dataclass(frozen=True)
class TableBlock:
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    layout: str


@dataclass(frozen=True)
class ListBlock:
    items: tuple[str, ...]


@dataclass(frozen=True)
class _ParseCtx:
    source: Path
    origin: Path
    heading_ids: frozenset[str]


Block = Heading | Body | FigureBlock | TableBlock | ListBlock


def publication_filenames(company: str) -> tuple[str, str]:
    return f"{company}_BAV.docx", f"{company}_BAV.pdf"


def require_pandoc() -> Path:
    path = shutil.which("pandoc")
    if not path:
        raise ValueError(
            "required converter unavailable: pandoc. "
            "Install pandoc 3.x and keep it on PATH; see README."
        )
    return Path(path)


def require_publication_libraries() -> None:
    missing: list[str] = []
    try:
        import docx  # noqa: F401
    except ImportError:
        missing.append("python-docx")
    try:
        import reportlab  # noqa: F401
    except ImportError:
        missing.append("reportlab")
    if missing:
        raise ValueError(
            "required converter unavailable: "
            + ", ".join(missing)
            + ". Install with pip install -r requirements-trainer.txt; see README."
        )
    _load_publication_libraries()


def _load_publication_libraries() -> None:
    """Bind python-docx and reportlab names after a successful diagnostic."""
    global Document, WD_ORIENT, WD_LINE_SPACING, OxmlElement, qn, Mm, Pt, RGBColor
    global black, HexColor, TA_LEFT, TA_RIGHT, ParagraphStyle, pdfmetrics, TTFont
    global Canvas, BaseDocTemplate, Frame, RLImage, KeepTogether, NextPageTemplate
    global PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle
    global _BAVCanvas, _LIBS_LOADED
    if _LIBS_LOADED:
        return
    from docx import Document as _Document
    from docx.enum.section import WD_ORIENT as _WD_ORIENT
    from docx.enum.text import WD_LINE_SPACING as _WD_LINE_SPACING
    from docx.oxml import OxmlElement as _OxmlElement
    from docx.oxml.ns import qn as _qn
    from docx.shared import Mm as _Mm, Pt as _Pt, RGBColor as _RGBColor
    from reportlab.lib.colors import HexColor as _HexColor, black as _black
    from reportlab.lib.enums import TA_LEFT as _TA_LEFT, TA_RIGHT as _TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle as _ParagraphStyle
    from reportlab.pdfbase import pdfmetrics as _pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as _TTFont
    from reportlab.pdfgen.canvas import Canvas as _Canvas
    from reportlab.platypus import (
        BaseDocTemplate as _BaseDocTemplate,
        Frame as _Frame,
        Image as _RLImage,
        KeepTogether as _KeepTogether,
        NextPageTemplate as _NextPageTemplate,
        PageBreak as _PageBreak,
        PageTemplate as _PageTemplate,
        Paragraph as _Paragraph,
        Spacer as _Spacer,
        Table as _Table,
        TableStyle as _TableStyle,
    )

    Document = _Document
    WD_ORIENT = _WD_ORIENT
    WD_LINE_SPACING = _WD_LINE_SPACING
    OxmlElement = _OxmlElement
    qn = _qn
    Mm = _Mm
    Pt = _Pt
    RGBColor = _RGBColor
    black = _black
    HexColor = _HexColor
    TA_LEFT = _TA_LEFT
    TA_RIGHT = _TA_RIGHT
    ParagraphStyle = _ParagraphStyle
    pdfmetrics = _pdfmetrics
    TTFont = _TTFont
    Canvas = _Canvas
    BaseDocTemplate = _BaseDocTemplate
    Frame = _Frame
    RLImage = _RLImage
    KeepTogether = _KeepTogether
    NextPageTemplate = _NextPageTemplate
    PageBreak = _PageBreak
    PageTemplate = _PageTemplate
    Paragraph = _Paragraph
    Spacer = _Spacer
    Table = _Table
    TableStyle = _TableStyle

    class BoundCanvas(_Canvas):
        """ReportLab defaults to Helvetica; force the resolved Aptos face."""

        def __init__(self, *args, latin: str = LATIN_FACE, **kwargs):
            super().__init__(*args, **kwargs)
            self._bav_latin = latin
            self.setFont(latin, BODY_PT)

        def setFont(self, psfontname, size, leading=None):
            if psfontname in {"Helvetica", "Helvetica-Bold", "Times-Roman", "Courier"}:
                psfontname = self._bav_latin
            return super().setFont(psfontname, size, leading)

    _BAVCanvas = BoundCanvas
    _LIBS_LOADED = True


def publish_company_documents(query: str) -> PublishedDocuments:
    """Publish Word and PDF under the company's canonical output directory."""
    return publish_resolved_company(resolve_company(query))


def publish_resolved_company(company: Company) -> PublishedDocuments:
    require_publication_libraries()
    pandoc = require_pandoc()
    fonts = resolve_required_fonts()
    drivers = _publishable_drivers(company)
    verify_research_artifacts(company.output, company.name)
    markdown = drivers.read_text(encoding="utf-8")
    _reject_debug_material(markdown, drivers)
    ast = _parse_markdown(pandoc, drivers)
    blocks = _blocks_from_ast(ast, drivers, fonts)
    _assert_no_empty_reserved_sections(blocks)
    word_name, pdf_name = publication_filenames(company.name)
    dest_word = company.output / word_name
    dest_pdf = company.output / pdf_name
    company.output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(
        prefix=f".{company.name}-publish-", dir=company.output.parent
    ) as raw:
        staged = Path(raw)
        stage_word = staged / word_name
        stage_pdf = staged / pdf_name
        try:
            _render_word(stage_word, company.name, fonts, blocks)
            _render_pdf(stage_pdf, company.name, fonts, blocks)
            _validate_staged_publication(stage_word, stage_pdf, company.name, blocks)
        except Exception as exc:
            raise ValueError(f"publication conversion failed: {exc}") from exc
        _install_publications(
            {dest_word: stage_word, dest_pdf: stage_pdf},
        )
    return PublishedDocuments(dest_word, dest_pdf)


def _publishable_drivers(company: Company) -> Path:
    drivers = company.output / "research" / drivers_filename(company.name)
    if not drivers.is_file() or drivers.stat().st_size == 0:
        raise ValueError(
            f"No publishable canonical research for {company.name}; "
            f"expected non-empty {drivers}. "
            "Publication does not invent research or use legacy paths. "
            f"Run python -m bav build {company.name} when Drivers applies."
        )
    for name in placeholder_filenames(company.name):
        path = company.output / "research" / name
        if path.is_file() and path.stat().st_size != 0:
            raise ValueError(f"placeholder must remain a zero-byte file: {path}")
    return drivers


def _reject_debug_material(text: str, path: Path) -> None:
    lowered = text.casefold()
    for term in DEBUG_TERMS:
        if term in lowered:
            raise ValueError(f"internal debug material is not publishable in {path}: {term}")


def _parse_markdown(pandoc: Path, markdown: Path) -> dict:
    result = subprocess.run(
        [str(pandoc), "-f", "gfm+pipe_tables", "-t", "json", "--", str(markdown)],
        capture_output=True,
        text=True,
        check=False,
    )
    if result.returncode != 0:
        detail = (result.stderr or result.stdout).strip() or f"exit {result.returncode}"
        raise ValueError(f"pandoc failed to parse {markdown}: {detail}")
    try:
        return json.loads(result.stdout)
    except json.JSONDecodeError as exc:
        raise ValueError(f"pandoc returned invalid JSON for {markdown}: {exc}") from exc


def _heading_identifier(text: str) -> str:
    ident: list[str] = []
    for ch in text.strip().casefold():
        if ch.isalnum() or ch in "-_":
            ident.append(ch)
        elif ch.isspace():
            ident.append("-")
    return re.sub("-{2,}", "-", "".join(ident)).strip("-")


def _heading_ids_from_ast(ast: dict) -> set[str]:
    ids: set[str] = set()
    for node in ast.get("blocks", ()):
        if not isinstance(node, dict) or node.get("t") != "Header":
            continue
        content = node.get("c")
        if not isinstance(content, list) or len(content) < 3:
            continue
        attr = content[1]
        if isinstance(attr, list) and attr and attr[0]:
            ids.add(str(attr[0]))
        label = _inlines_text(content[2])
        slug = _heading_identifier(label)
        if slug:
            ids.add(slug)
    return ids


def _heading_ids_from_markdown(path: Path) -> set[str]:
    try:
        ast = _parse_markdown(require_pandoc(), path)
    except ValueError:
        ast = {}
    ids = _heading_ids_from_ast(ast)
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("#"):
            slug = _heading_identifier(line.lstrip("#").strip())
            if slug:
                ids.add(slug)
    return ids


def _inlines_text(inlines, ctx: _ParseCtx | None = None) -> str:
    parts: list[str] = []
    for node in inlines or ():
        kind = node.get("t") if isinstance(node, dict) else None
        content = node.get("c") if isinstance(node, dict) else None
        if kind == "Str":
            parts.append(str(content))
        elif kind in {"Space", "SoftBreak"}:
            parts.append(" ")
        elif kind == "LineBreak":
            parts.append("\n")
        elif kind == "Quoted":
            parts.append(_inlines_text(content[1], ctx))
        elif kind in {"Emph", "Strong", "Underline", "Strikeout", "SmallCaps"}:
            parts.append(_inlines_text(content, ctx))
        elif kind == "Math":
            parts.append(content[1] if isinstance(content, list) else str(content))
        elif kind == "RawInline":
            parts.append(content[1] if isinstance(content, list) else "")
        elif kind == "Link":
            target = content[2][0] if isinstance(content, list) and len(content) >= 3 else ""
            if ctx is not None:
                _validate_link_target(ctx.source, ctx.origin, target, ctx.heading_ids)
            label = _inlines_text(content[1], ctx)
            parts.append(_format_link_text(label, target))
        elif kind == "Image":
            parts.append(_inlines_text(content[1], ctx))
        elif kind == "Code":
            parts.append(content[1] if isinstance(content, list) else str(content))
        elif kind == "Span":
            parts.append(_inlines_text(content[1], ctx))
        elif kind == "Note":
            continue
        elif isinstance(content, list):
            parts.append(_inlines_text(content, ctx))
    return "".join(parts)


def _format_link_text(label: str, target: str) -> str:
    label = label.strip()
    if not target:
        return label
    if target.startswith("#"):
        return label or target
    display = Path(target.split("#", 1)[0]).name
    if display and display not in label:
        return f"{label} ({display})" if label else display
    return label or target


def _walk_links(node, acc: list[str]) -> None:
    if isinstance(node, dict):
        if node.get("t") == "Link":
            content = node.get("c")
            if isinstance(content, list) and len(content) >= 3:
                acc.append(str(content[2][0]))
        for value in node.values():
            _walk_links(value, acc)
    elif isinstance(node, list):
        for item in node:
            _walk_links(item, acc)


def _validate_document_links(
    ast: dict, source: Path, origin: Path, heading_ids: set[str]
) -> None:
    targets: list[str] = []
    _walk_links(ast, targets)
    for target in targets:
        _validate_link_target(source, origin, target, heading_ids)


def _validate_link_target(
    source: Path, origin: Path, target: str, heading_ids: set[str]
) -> None:
    if not target or target.startswith(("mailto:", "http://", "https://", "javascript:")):
        raise ValueError(f"unsupported document reference {target!r} in {source}")
    href, frag = target, ""
    if "#" in target:
        href, frag = target.split("#", 1)
    if not href:
        if not frag or frag not in heading_ids:
            raise ValueError(
                f"broken document reference {target!r} in {source}: "
                f"invalid anchor {frag!r}"
            )
        return
    path = _resolve_local_document(origin, href, source)
    if frag:
        if path.suffix.lower() in {".md", ".markdown"}:
            ids = heading_ids if path.resolve() == source.resolve() else _heading_ids_from_markdown(path)
            if frag not in ids:
                raise ValueError(
                    f"broken document reference {target!r} in {source}: "
                    f"invalid anchor {frag!r}"
                )
        else:
            raise ValueError(
                f"unsupported document reference {target!r} in {source}: "
                f"anchor on non-markdown target"
            )


def _blocks_text(blocks, ctx: _ParseCtx | None = None) -> str:
    parts: list[str] = []
    for block in blocks or ():
        kind = block.get("t") if isinstance(block, dict) else None
        content = block.get("c") if isinstance(block, dict) else None
        if kind in {"Para", "Plain"}:
            parts.append(_inlines_text(content, ctx))
        elif kind == "Header":
            parts.append(_inlines_text(content[2], ctx))
        elif isinstance(content, list):
            parts.append(_blocks_text(content, ctx))
    return "\n".join(part for part in parts if part)


def _walk_images(node, acc: list[tuple[str, str]]) -> None:
    if isinstance(node, dict):
        if node.get("t") == "Image":
            inlines, target = node["c"][1], node["c"][2][0]
            acc.append((_inlines_text(inlines), target))
        for value in node.values():
            _walk_images(value, acc)
    elif isinstance(node, list):
        for item in node:
            _walk_images(item, acc)


def _resolve_local_document(origin: Path, target: str, source: Path) -> Path:
    path = (origin / target).resolve()
    research_root = origin.resolve()
    try:
        path.relative_to(research_root.parent.parent.resolve())
    except ValueError as exc:
        raise ValueError(
            f"broken document reference {target!r} in {source}: "
            f"target escapes canonical research"
        ) from exc
    if not path.is_file():
        raise ValueError(
            f"broken document reference {target!r} in {source}: "
            f"missing local file {path}"
        )
    return path


def _resolve_asset(origin: Path, target: str) -> Path:
    if not target or target.startswith(("#", "mailto:", "http://", "https://")):
        raise ValueError(f"unsupported document reference {target!r} in {origin}")
    path = (origin / target).resolve()
    research_root = origin.resolve()
    try:
        path.relative_to(research_root.parent.parent.resolve())
    except ValueError as exc:
        raise ValueError(
            f"broken reference {target!r} escapes canonical research at {origin}"
        ) from exc
    if not path.is_file():
        raise ValueError(f"missing figure or referenced file: {path} (from {target})")
    if path.suffix.lower() == ".png" and path.read_bytes()[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"referenced figure is not a PNG: {path}")
    return path


def _cell_text(cell, ctx: _ParseCtx | None = None) -> str:
    return _blocks_text(cell[4] if isinstance(cell, list) and len(cell) >= 5 else cell, ctx)


def _table_rows(
    header_rows, body, ctx: _ParseCtx | None = None
) -> tuple[tuple[str, ...], tuple[tuple[str, ...], ...]]:
    headers = tuple(_cell_text(cell, ctx) for cell in header_rows[0][1]) if header_rows else ()
    rows = []
    bodies = body if isinstance(body, list) else []
    for tbody in bodies:
        body_rows = tbody[3] if isinstance(tbody, list) and len(tbody) >= 4 else tbody
        for row in body_rows:
            cells = row[1] if isinstance(row, list) and len(row) >= 2 else row
            rows.append(tuple(_cell_text(cell, ctx) for cell in cells))
    if headers:
        width = len(headers)
        normalized = []
        for row in rows:
            values = list(row[:width])
            if len(values) < width:
                values.extend([""] * (width - len(values)))
            normalized.append(tuple(values))
        rows = normalized
    return headers, tuple(rows)


def _table_layout(headers: tuple[str, ...], rows: tuple[tuple[str, ...], ...]) -> str:
    ncols = max(len(headers), max((len(row) for row in rows), default=0))
    body_long = any(len(cell) > 40 for row in rows for cell in row)
    char_w = BODY_PT * 0.5
    min_width = 0.0
    for index in range(ncols):
        samples = [headers[index] if index < len(headers) else ""]
        samples.extend(row[index] if index < len(row) else "" for row in rows)
        longest_word = max(
            (len(word) for sample in samples for word in sample.replace("/", " ").split()),
            default=4,
        )
        min_width += max(longest_word, 6) * char_w
    landscape_width = LANDSCAPE[0] - 2 * PAGE_MARGIN_MM * MM_PT
    portrait_width = PORTRAIT[0] - 2 * PAGE_MARGIN_MM * MM_PT
    if body_long and min_width > landscape_width:
        return "stacked"
    if ncols >= 7 or min_width > portrait_width:
        return "landscape"
    return "portrait"


def _blocks_from_ast(ast: dict, source: Path, fonts: ResolvedFonts) -> list[Block]:
    del fonts
    research_dir = source.parent
    heading_ids = _heading_ids_from_ast(ast)
    ctx = _ParseCtx(source=source, origin=research_dir, heading_ids=frozenset(heading_ids))
    _validate_document_links(ast, source, research_dir, heading_ids)
    blocks: list[Block] = []
    for node in ast.get("blocks", ()):
        kind = node.get("t")
        content = node.get("c")
        if kind == "Header":
            blocks.append(Heading(int(content[0]), _inlines_text(content[2], ctx)))
        elif kind in {"Para", "Plain"}:
            images: list[tuple[str, str]] = []
            _walk_images(node, images)
            if images and not _inlines_text(content, ctx).strip():
                for caption, target in images:
                    blocks.append(FigureBlock(_resolve_asset(research_dir, target), caption))
            elif images:
                text = _inlines_text(content, ctx).strip()
                if text:
                    blocks.append(Body(text))
                for caption, target in images:
                    blocks.append(FigureBlock(_resolve_asset(research_dir, target), caption))
            else:
                text = _inlines_text(content, ctx).strip()
                if text:
                    blocks.append(Body(text))
        elif kind == "Figure":
            caption = _blocks_text(
                content[1][1] if isinstance(content[1], list) else content[1], ctx
            )
            images: list[tuple[str, str]] = []
            _walk_images(content[2], images)
            if not images:
                raise ValueError("figure block has no image reference")
            for alt, target in images:
                blocks.append(
                    FigureBlock(_resolve_asset(research_dir, target), caption or alt)
                )
        elif kind == "Table":
            headers, rows = _table_rows(content[3][1], content[4], ctx)
            if not headers and not rows:
                raise ValueError("table has no readable cells")
            blocks.append(TableBlock(headers, rows, _table_layout(headers, rows)))
        elif kind == "OrderedList":
            items = []
            for item in content[1]:
                text = _blocks_text(item, ctx).strip()
                if text:
                    items.append(text)
            if items:
                blocks.append(ListBlock(tuple(items)))
        elif kind == "BulletList":
            items = []
            for item in content:
                text = _blocks_text(item, ctx).strip()
                if text:
                    items.append(text)
            if items:
                blocks.append(ListBlock(tuple(items)))
        elif kind == "RawBlock":
            continue
        else:
            text = _blocks_text([node], ctx).strip()
            if text:
                blocks.append(Body(text))
    if not blocks:
        raise ValueError("canonical Drivers Markdown produced no publishable blocks")
    return blocks


def _assert_no_empty_reserved_sections(blocks: list[Block]) -> None:
    headings = [block.text.casefold() for block in blocks if isinstance(block, Heading)]
    for reserved in ("forecast", "valuation", "overview"):
        if any(reserved == heading or heading.endswith(f"— {reserved}") for heading in headings):
            raise ValueError(f"empty reserved module {reserved} must not appear in the publication")


def _register_pdf_fonts(fonts: ResolvedFonts) -> tuple[str, str]:
    latin = f"BAV-{LATIN_FACE}"
    cjk = f"BAV-{CJK_FACE}"
    if latin not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(latin, str(fonts.latin_path)))
    if cjk not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(cjk, str(fonts.cjk_path)))
    return latin, cjk


def _set_run_font(run, fonts: ResolvedFonts, size: float, *, italic: bool = False) -> None:
    run.font.name = fonts.latin_name
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.rPr.rFonts.set(qn("w:ascii"), fonts.latin_name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), fonts.latin_name)
    run._element.rPr.rFonts.set(qn("w:cs"), fonts.latin_name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), fonts.cjk_name)


def _apply_paragraph_format(paragraph, *, before: float = 0, after: float = RELATED_PT) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = LINE_SPACING
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def _set_section_page(section, *, wide: bool) -> None:
    section.page_width = Mm(297 if wide else 210)
    section.page_height = Mm(210 if wide else 297)
    section.orientation = WD_ORIENT.LANDSCAPE if wide else WD_ORIENT.PORTRAIT
    for edge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, edge, Mm(PAGE_MARGIN_MM))


def _repeat_header_row(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _set_table_fixed(table, width_mm: float) -> None:
    tbl = table._tbl
    tbl_pr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), str(int(width_mm * 56.7)))
    width.set(qn("w:type"), "dxa")
    existing = tbl_pr.find(qn("w:tblW"))
    if existing is not None:
        tbl_pr.remove(existing)
    tbl_pr.append(width)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    existing_layout = tbl_pr.find(qn("w:tblLayout"))
    if existing_layout is not None:
        tbl_pr.remove(existing_layout)
    tbl_pr.append(layout)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:space"), "0")
        item.set(qn("w:color"), "808080")
        borders.append(item)
    existing_borders = tbl_pr.find(qn("w:tblBorders"))
    if existing_borders is not None:
        tbl_pr.remove(existing_borders)
    tbl_pr.append(borders)


def _soft_wrap_header(text: str) -> str:
    """Break long headers at slashes, spaces or existing hyphens, never mid-word."""
    if "/" in text:
        return text.replace("/", "/\n")
    if " " in text and len(text) > 18:
        parts = text.split(" ")
        lines: list[str] = [parts[0]]
        for part in parts[1:]:
            if len(lines[-1]) + 1 + len(part) <= 18:
                lines[-1] = f"{lines[-1]} {part}"
            else:
                lines.append(part)
        return "\n".join(lines)
    if "-" in text and len(text) > 18:
        return text.replace("-", "-\n")
    return text


def _fill_word_cell(cell, text: str, fonts: ResolvedFonts, size: float) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _apply_paragraph_format(paragraph, before=2, after=2)
    for index, line in enumerate(text.split("\n")):
        if index:
            run = paragraph.add_run()
            run.add_break()
            _set_run_font(run, fonts, size)
        run = paragraph.add_run(line)
        _set_run_font(run, fonts, size)
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    no_wrap = tc_pr.find(qn("w:noWrap"))
    if no_wrap is not None:
        tc_pr.remove(no_wrap)


def _column_widths_mm(block: TableBlock, width_mm: float) -> list[float]:
    ncols = max(len(block.headers), 1)
    weights: list[float] = []
    for index in range(ncols):
        samples = [block.headers[index] if index < len(block.headers) else ""]
        samples.extend(row[index] if index < len(row) else "" for row in block.rows)
        longest_word = max(
            (len(word) for sample in samples for word in re.split(r"[/\s-]+", sample) if word),
            default=4,
        )
        longest = max((len(sample) for sample in samples), default=4)
        weights.append(max(longest_word, 6) + 0.2 * longest)
    total = sum(weights) or ncols
    return [width_mm * weight / total for weight in weights]


def _add_word_grid(document, block: TableBlock, fonts: ResolvedFonts, *, wide: bool) -> None:
    width_mm = (297 if wide else 210) - 2 * PAGE_MARGIN_MM
    table = document.add_table(rows=1 + len(block.rows), cols=len(block.headers))
    _set_table_fixed(table, width_mm)
    for index, header in enumerate(block.headers):
        _fill_word_cell(table.rows[0].cells[index], _soft_wrap_header(header), fonts, BODY_PT)
    _repeat_header_row(table.rows[0])
    for row_index, row in enumerate(block.rows, start=1):
        for col_index, value in enumerate(row):
            _fill_word_cell(table.rows[row_index].cells[col_index], value, fonts, BODY_PT)
    widths = _column_widths_mm(block, width_mm)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Mm(widths[index] if index < len(widths) else widths[-1])


def _add_word_stacked(document, block: TableBlock, fonts: ResolvedFonts) -> None:
    for row in block.rows:
        for header, value in zip(block.headers, row, strict=False):
            label = document.add_paragraph()
            _apply_paragraph_format(label, before=RELATED_PT, after=2)
            run = label.add_run(header)
            _set_run_font(run, fonts, BODY_PT)
            body = document.add_paragraph()
            _apply_paragraph_format(body, before=0, after=RELATED_PT)
            run = body.add_run(value)
            _set_run_font(run, fonts, BODY_PT)
        spacer = document.add_paragraph()
        _apply_paragraph_format(spacer, before=0, after=SECTION_PT)


def _word_header_footer(section, company: str, fonts: ResolvedFonts) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.clear()
    _apply_paragraph_format(paragraph, before=0, after=RELATED_PT)
    run = paragraph.add_run(f"{company} BAV")
    _set_run_font(run, fonts, BODY_PT)
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    _apply_paragraph_format(paragraph, before=RELATED_PT, after=0)
    run = paragraph.add_run("Page ")
    _set_run_font(run, fonts, LABEL_PT)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run = paragraph.add_run()
    run._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run = paragraph.add_run()
    run._r.append(instr)
    _set_run_font(run, fonts, LABEL_PT)
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(fld)


def _render_word(path: Path, company: str, fonts: ResolvedFonts, blocks: list[Block]) -> None:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = fonts.latin_name
    normal.font.size = Pt(BODY_PT)
    normal.font.bold = False
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), fonts.latin_name)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), fonts.latin_name)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), fonts.cjk_name)
    _set_section_page(document.sections[0], wide=False)
    _word_header_footer(document.sections[0], company, fonts)
    core = document.core_properties
    core.title = f"{company} BAV"
    core.author = "BAV"
    title = document.add_paragraph()
    _apply_paragraph_format(title, before=0, after=SECTION_PT)
    run = title.add_run(f"{company} BAV")
    _set_run_font(run, fonts, HEADING_PT)
    current_wide = False
    for index, block in enumerate(blocks):
        following = blocks[index + 1] if index + 1 < len(blocks) else None
        want_wide = isinstance(block, TableBlock) and block.layout == "landscape"
        if isinstance(block, Body) and isinstance(following, TableBlock) and following.layout == "landscape":
            want_wide = True
        if want_wide != current_wide:
            section = document.add_section()
            _set_section_page(section, wide=want_wide)
            _word_header_footer(section, company, fonts)
            current_wide = want_wide
        if isinstance(block, Heading):
            paragraph = document.add_paragraph()
            _apply_paragraph_format(
                paragraph,
                before=SECTION_PT if block.level > 1 else RELATED_PT,
                after=RELATED_PT,
            )
            run = paragraph.add_run(block.text)
            _set_run_font(run, fonts, HEADING_PT)
        elif isinstance(block, Body):
            paragraph = document.add_paragraph()
            _apply_paragraph_format(paragraph, before=0, after=RELATED_PT)
            run = paragraph.add_run(block.text)
            _set_run_font(run, fonts, BODY_PT)
        elif isinstance(block, ListBlock):
            for index, item in enumerate(block.items, start=1):
                paragraph = document.add_paragraph()
                _apply_paragraph_format(paragraph, before=0, after=RELATED_PT)
                run = paragraph.add_run(f"{index}. {item}")
                _set_run_font(run, fonts, BODY_PT)
        elif isinstance(block, FigureBlock):
            width = Mm((297 if current_wide else 210) - 2 * PAGE_MARGIN_MM)
            document.add_picture(str(block.path), width=width)
            caption = document.add_paragraph()
            _apply_paragraph_format(caption, before=RELATED_PT, after=SECTION_PT)
            run = caption.add_run(block.caption)
            _set_run_font(run, fonts, CAPTION_PT)
        elif isinstance(block, TableBlock):
            if block.layout == "stacked":
                _add_word_stacked(document, block, fonts)
            else:
                _add_word_grid(document, block, fonts, wide=block.layout == "landscape")
            spacer = document.add_paragraph()
            _apply_paragraph_format(spacer, before=0, after=SECTION_PT)
    if current_wide:
        section = document.add_section()
        _set_section_page(section, wide=False)
        _word_header_footer(section, company, fonts)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def _escape_xml(text: str) -> str:
    return (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _pdf_styles(latin: str) -> dict[str, ParagraphStyle]:
    leading = BODY_PT * LINE_SPACING
    return {
        "heading": ParagraphStyle(
            "BAVHeading",
            fontName=latin,
            fontSize=HEADING_PT,
            leading=HEADING_PT * LINE_SPACING,
            textColor=black,
            spaceBefore=SECTION_PT,
            spaceAfter=RELATED_PT,
            alignment=TA_LEFT,
        ),
        "body": ParagraphStyle(
            "BAVBody",
            fontName=latin,
            fontSize=BODY_PT,
            leading=leading,
            textColor=black,
            spaceBefore=0,
            spaceAfter=RELATED_PT,
            alignment=TA_LEFT,
        ),
        "caption": ParagraphStyle(
            "BAVCaption",
            fontName=latin,
            fontSize=CAPTION_PT,
            leading=CAPTION_PT * LINE_SPACING,
            textColor=HexColor("#4D4D4D"),
            spaceBefore=RELATED_PT,
            spaceAfter=SECTION_PT,
            alignment=TA_LEFT,
        ),
        "cell": ParagraphStyle(
            "BAVCell",
            fontName=latin,
            fontSize=BODY_PT,
            leading=leading,
            textColor=black,
            alignment=TA_LEFT,
        ),
        "label": ParagraphStyle(
            "BAVLabel",
            fontName=latin,
            fontSize=BODY_PT,
            leading=leading,
            textColor=HexColor("#4D4D4D"),
            spaceBefore=RELATED_PT,
            spaceAfter=2,
            alignment=TA_LEFT,
        ),
        "header": ParagraphStyle(
            "BAVRunning",
            fontName=latin,
            fontSize=BODY_PT,
            leading=leading,
            textColor=black,
            alignment=TA_LEFT,
        ),
        "page": ParagraphStyle(
            "BAVPage",
            fontName=latin,
            fontSize=LABEL_PT,
            leading=LABEL_PT * LINE_SPACING,
            textColor=black,
            alignment=TA_RIGHT,
        ),
    }


def _pdf_header(company: str, latin: str, canvas, doc) -> None:
    canvas.saveState()
    canvas.setFillColor(black)
    canvas.setFont(latin, BODY_PT)
    width, height = canvas._pagesize
    canvas.drawString(PAGE_MARGIN_MM * MM_PT, height - 12 * MM_PT, f"{company} BAV")
    canvas.setFont(latin, LABEL_PT)
    canvas.drawRightString(width - PAGE_MARGIN_MM * MM_PT, 10 * MM_PT, str(doc.page))
    canvas.setStrokeColor(HexColor("#808080"))
    canvas.setLineWidth(0.4)
    canvas.line(
        PAGE_MARGIN_MM * MM_PT,
        height - 14 * MM_PT,
        width - PAGE_MARGIN_MM * MM_PT,
        height - 14 * MM_PT,
    )
    canvas.restoreState()


def _escape_cell(text: str) -> str:
    return _escape_xml(text).replace("\n", "<br/>")


def _pdf_table(block: TableBlock, styles, page_width: float) -> Table:
    data = [
        [Paragraph(_escape_cell(_soft_wrap_header(cell)), styles["cell"]) for cell in block.headers]
    ]
    for row in block.rows:
        data.append([Paragraph(_escape_cell(cell), styles["cell"]) for cell in row])
    ncols = len(block.headers)
    widths = _column_widths_mm(block, page_width / MM_PT)
    col_ws = [width * MM_PT for width in widths]
    table = Table(data, colWidths=col_ws, repeatRows=1, splitByRow=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), styles["cell"].fontName),
                ("FONTSIZE", (0, 0), (-1, -1), BODY_PT),
                ("TEXTCOLOR", (0, 0), (-1, -1), black),
                ("BACKGROUND", (0, 0), (-1, 0), HexColor("#F2F2F2")),
                ("GRID", (0, 0), (-1, -1), 0.4, HexColor("#808080")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    return table


def _figure_flowable(block: FigureBlock, styles, page_width: float):
    image = RLImage(str(block.path))
    image.drawWidth = page_width
    image.drawHeight = page_width * (4.8 / 7.5)
    caption = Paragraph(_escape_xml(block.caption), styles["caption"])
    return KeepTogether([image, caption])


def _render_pdf(path: Path, company: str, fonts: ResolvedFonts, blocks: list[Block]) -> None:
    latin, _cjk = _register_pdf_fonts(fonts)
    styles = _pdf_styles(latin)
    frame_kw = dict(
        x1=PAGE_MARGIN_MM * MM_PT,
        y1=16 * MM_PT,
        width=PORTRAIT[0] - 2 * PAGE_MARGIN_MM * MM_PT,
        height=PORTRAIT[1] - 32 * MM_PT,
    )
    land_kw = dict(
        x1=PAGE_MARGIN_MM * MM_PT,
        y1=16 * MM_PT,
        width=LANDSCAPE[0] - 2 * PAGE_MARGIN_MM * MM_PT,
        height=LANDSCAPE[1] - 32 * MM_PT,
    )
    document = BaseDocTemplate(
        str(path),
        pagesize=PORTRAIT,
        title=f"{company} BAV",
        author="BAV",
        creator="BAV",
    )
    document.addPageTemplates(
        [
            PageTemplate(
                id="portrait",
                frames=[Frame(**frame_kw, id="p")],
                onPage=lambda c, d: _pdf_header(company, latin, c, d),
                pagesize=PORTRAIT,
            ),
            PageTemplate(
                id="landscape",
                frames=[Frame(**land_kw, id="l")],
                onPage=lambda c, d: _pdf_header(company, latin, c, d),
                pagesize=LANDSCAPE,
            ),
        ]
    )
    story: list = [Paragraph(_escape_xml(f"{company} BAV"), styles["heading"])]
    current = "portrait"
    portrait_w = frame_kw["width"]
    landscape_w = land_kw["width"]

    def switch(target: str) -> None:
        nonlocal current
        if current == target:
            return
        story.append(NextPageTemplate(target))
        story.append(PageBreak())
        current = target

    for index, block in enumerate(blocks):
        following = blocks[index + 1] if index + 1 < len(blocks) else None
        if isinstance(block, TableBlock) and block.layout == "landscape":
            switch("landscape")
            story.append(_pdf_table(block, styles, landscape_w))
            story.append(Spacer(1, SECTION_PT))
            continue
        if isinstance(block, Body) and isinstance(following, TableBlock) and following.layout == "landscape":
            switch("landscape")
            story.append(Paragraph(_escape_xml(block.text), styles["body"]))
            continue
        switch("portrait")
        if isinstance(block, Heading):
            story.append(Paragraph(_escape_xml(block.text), styles["heading"]))
        elif isinstance(block, Body):
            story.append(Paragraph(_escape_xml(block.text), styles["body"]))
        elif isinstance(block, ListBlock):
            for index, item in enumerate(block.items, start=1):
                story.append(
                    Paragraph(_escape_xml(f"{index}. {item}"), styles["body"])
                )
        elif isinstance(block, FigureBlock):
            story.append(_figure_flowable(block, styles, portrait_w))
        elif isinstance(block, TableBlock) and block.layout == "stacked":
            for row in block.rows:
                group = []
                for header, value in zip(block.headers, row, strict=False):
                    group.append(Paragraph(_escape_xml(header), styles["label"]))
                    group.append(Paragraph(_escape_xml(value), styles["body"]))
                group.append(Spacer(1, RELATED_PT))
                story.append(KeepTogether(group))
        elif isinstance(block, TableBlock):
            story.append(_pdf_table(block, styles, portrait_w))
            story.append(Spacer(1, SECTION_PT))
    path.parent.mkdir(parents=True, exist_ok=True)
    document.build(
        story,
        canvasmaker=lambda *args, **kwargs: _BAVCanvas(*args, latin=latin, **kwargs),
    )


def _validate_staged_publication(
    word: Path, pdf: Path, company: str, blocks: list[Block]
) -> None:
    if word.stat().st_size == 0 or pdf.stat().st_size == 0:
        raise ValueError("staged publication is empty")
    if word.read_bytes()[:2] != b"PK":
        raise ValueError(f"staged Word file is not a DOCX: {word}")
    if not pdf.read_bytes().startswith(b"%PDF"):
        raise ValueError(f"staged PDF is not a PDF: {pdf}")
    with zipfile.ZipFile(word) as archive:
        names = set(archive.namelist())
        if "word/document.xml" not in names:
            raise ValueError("staged Word file is missing document.xml")
        xml = archive.read("word/document.xml").decode("utf-8")
        media = [name for name in names if name.startswith("word/media/")]
    heading_text = [block.text for block in blocks if isinstance(block, Heading)]
    for heading in heading_text:
        if heading not in xml:
            raise ValueError(f"staged Word is missing heading {heading!r}")
    figures = [block for block in blocks if isinstance(block, FigureBlock)]
    if len(media) < len(figures):
        raise ValueError("staged Word is missing embedded figures")
    try:
        import fitz
    except ImportError as exc:
        raise ValueError("required converter unavailable: pymupdf") from exc
    doc = fitz.open(pdf)
    try:
        if doc.page_count < 1:
            raise ValueError("staged PDF has no pages")
        text = "\n".join(page.get_text() for page in doc)
        images = sum(len(page.get_images()) for page in doc)
        fonts = set()
        for page in doc:
            for block in page.get_text("dict").get("blocks", ()):
                for line in block.get("lines", ()):
                    for span in line.get("spans", ()):
                        if span.get("text", "").strip():
                            fonts.add(span.get("font", ""))
    finally:
        doc.close()
    for heading in heading_text:
        if heading not in text:
            raise ValueError(f"staged PDF is missing heading {heading!r}")
    if images < len(figures):
        raise ValueError("staged PDF is missing embedded figures")
    if not any(LATIN_FACE.casefold() in str(name).casefold() for name in fonts):
        raise ValueError(
            f"staged PDF does not use required {LATIN_FACE} (found {sorted(fonts)})"
        )
    substituted = {
        name
        for name in fonts
        if name
        and LATIN_FACE.casefold() not in name.casefold()
        and CJK_FACE.casefold() not in name.casefold()
    }
    if substituted:
        raise ValueError(
            f"staged PDF rendered text with substituted fonts {sorted(substituted)}; "
            f"required {LATIN_FACE} and {CJK_FACE}"
        )
    if f"{company} BAV" not in xml or f"{company} BAV" not in text:
        raise ValueError("publication is missing BAV title")
    lowered = text.casefold()
    for term in DEBUG_TERMS:
        if term in lowered:
            raise ValueError(f"publication contains internal debug material: {term}")


def _install_publications(mapping: dict[Path, Path]) -> None:
    backups: dict[Path, Path] = {}
    installed: list[Path] = []
    try:
        for dest, source in mapping.items():
            if dest.exists():
                backup = dest.with_name(f".{dest.name}.prev")
                os.replace(dest, backup)
                backups[dest] = backup
            os.replace(source, dest)
            installed.append(dest)
        for backup in backups.values():
            Path(backup).unlink(missing_ok=True)
    except Exception:
        for dest in installed:
            dest.unlink(missing_ok=True)
        for dest, backup in backups.items():
            if Path(backup).exists():
                os.replace(backup, dest)
        raise

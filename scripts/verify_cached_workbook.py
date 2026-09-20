"""Read-only saved-cache acceptance; native Excel execution belongs to AutoCycle.

References are explicit independent cell expectations, not values read from the
copy or calculated by the production model. Check every formula on the selected
sheets and their transitive A1 dependencies. Unsupported references fail closed.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def verify(copy, original, references):
    copy, original, references = map(Path, (copy, original, references))
    report = {'status': 'BLOCKED', 'copy': str(copy.resolve()),
              'source': str(original.resolve()), 'action': ''}
    books = []
    try:
        if copy.resolve() == original.resolve():
            raise ValueError('Verification requires a separate saved copy')
        before = [digest(p) for p in (original, copy, references)]
        refs = json.loads(references.read_text())
        if not refs.get('cells') or not refs.get('sheets'):
            raise ValueError('Nonempty independent references and affected sheets required')
        if refs.get('source_sha256', before[0]) != before[0]:
            raise ValueError('Independent references belong to a different source workbook')
        for path, data_only in ((original, False), (copy, False), (copy, True)):
            books.append(load_workbook(path, data_only=data_only))
        source, formulas, cached = books
        if source.sheetnames != formulas.sheetnames:
            raise ValueError('Workbook sheet topology changed')
        # Includes hidden sheets, newly introduced formulas, and literal inputs.
        for sheet in source:
            saved = formulas[sheet.title]
            if sheet.sheet_state != saved.sheet_state:
                raise ValueError(f'Sheet visibility changed: {sheet.title}')
            for row in sheet.iter_rows(max_row=max(sheet.max_row, saved.max_row),
                                       max_col=max(sheet.max_column, saved.max_column)):
                for cell in row:
                    other = saved[cell.coordinate]
                    if cell.value != other.value or cell.data_type != other.data_type:
                        kind = 'Formula' if 'f' in (cell.data_type, other.data_type) else 'Source cell'
                        raise ValueError(f'{kind} changed: {sheet.title}!{cell.coordinate}')
        pending = set()
        expectations = {}
        for ref in refs['cells']:
            key = (ref['sheet'], ref['cell'])
            if key in expectations or not ref.get('source'):
                raise ValueError('Duplicate reference or missing independent source citation')
            expectations[key] = ref['value']
            pending.add(key)
        for name in refs['sheets']:
            count = 0
            for row in source[name]:
                for cell in row:
                    if cell.data_type == 'f':
                        key = (name, cell.coordinate)
                        if key not in expectations:
                            raise ValueError(f'Independent reference missing for affected formula: {name}!{cell.coordinate}')
                        pending.add(key); count += 1
            if not count:
                raise ValueError(f'No affected formulas on reference sheet: {name}')
        checked = set()
        while pending:
            name, address = pending.pop()
            key = (name, address)
            if key in checked:
                continue
            checked.add(key)
            cell = source[name][address]
            value = cached[name][address]
            if value.data_type == 'e' or value.value is None:
                raise ValueError(f'Cached value unavailable/error: {name}!{address} = {value.value!r}')
            if key in expectations:
                expected = expectations[key]
                actual = value.value
                matches = (isinstance(actual, (int, float)) and not isinstance(actual, bool)
                           and math.isfinite(actual) and math.isclose(actual, expected, rel_tol=0, abs_tol=1e-8)) \
                    if isinstance(expected, (int, float)) and not isinstance(expected, bool) else actual == expected
                if not matches:
                    raise ValueError(f'Independent reference mismatch: {name}!{address}: {actual!r} != {expected!r}')
            if cell.data_type != 'f':
                continue
            for token in Tokenizer(cell.value).items:
                if token.type == 'FUNC' and token.subtype == 'OPEN' and token.value.upper() in ('INDIRECT(', 'OFFSET('):
                    raise ValueError(f'Unsupported dynamic dependency: {name}!{address}')
                if token.type != 'OPERAND' or token.subtype != 'RANGE':
                    continue
                target, separator, area = token.value.rpartition('!')
                if not separator:
                    target, area = name, token.value
                elif target.startswith("'") and target.endswith("'"):
                    target = target[1:-1].replace("''", "'")
                if target not in source.sheetnames or not re.fullmatch(r'\$?[A-Z]+\$?[1-9][0-9]*(?::\$?[A-Z]+\$?[1-9][0-9]*)?', area):
                    raise ValueError(f'Unsupported dependency {token.value!r}: {name}!{address}')
                left, top, right, bottom = range_boundaries(area)
                for row in source[target].iter_rows(min_col=left, max_col=right, min_row=top, max_row=bottom):
                    pending.update((target, c.coordinate) for c in row)
        if before != [digest(p) for p in (original, copy, references)]:
            raise ValueError('Verification inputs changed during inspection')
        report.update(status='VERIFIED', action='NONE', source_sha256=before[0],
                      copy_sha256=before[1], references_sha256=before[2],
                      checked_cells=len(checked), independent_references=len(expectations),
                      affected_sheets=refs['sheets'], formulas_preserved=True)
    except (OSError, ValueError, KeyError, TypeError, BadZipFile) as error:
        report['action'] = str(error)
    finally:
        for book in books:
            book.close()
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('--original', required=True, type=Path)
    parser.add_argument('--references', required=True, type=Path)
    args = parser.parse_args()
    result = verify(args.workbook, args.original, args.references)
    print(json.dumps(result, indent=2))
    return 0 if result['status'] == 'VERIFIED' else 1


if __name__ == '__main__':
    raise SystemExit(main())

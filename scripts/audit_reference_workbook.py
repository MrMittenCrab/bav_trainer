"""Read-only inventory of an Excel workbook for historical reference audits.

Never mutates or saves the inspected workbook.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def audit_workbook(path: Path) -> dict[str, Any]:
    """Return a deterministic sheet-level inventory of ``path``.

    Uses ``data_only=False`` and ``read_only=True``. Never calls ``save()``.
    """
    path = Path(path)
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        sheets: list[dict[str, Any]] = []
        for ws in wb.worksheets:
            nonempty = 0
            formulas = 0
            labels: list[str] = []
            seen: set[str] = set()
            max_row = 0
            max_col = 0
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    nonempty += 1
                    max_row = max(max_row, cell.row)
                    max_col = max(max_col, cell.column)
                    if isinstance(value, str) and value.startswith("="):
                        formulas += 1
                    if (
                        cell.column <= 3
                        and isinstance(value, str)
                        and value
                        and not value.startswith("=")
                        and value not in seen
                        and len(labels) < 40
                    ):
                        seen.add(value)
                        labels.append(value)
            sheets.append(
                {
                    "name": ws.title,
                    "state": ws.sheet_state,
                    "max_row": max_row,
                    "max_column": max_col,
                    "nonempty_cells": nonempty,
                    "formula_cells": formulas,
                    "representative_labels": labels,
                }
            )
        return {"workbook": path.name, "sheets": sheets}
    finally:
        wb.close()


def main(argv: list[str] | None = None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    if len(args) != 1:
        print(
            "Usage: python scripts/audit_reference_workbook.py <workbook.xlsx>",
            file=sys.stderr,
        )
        return 2
    path = Path(args[0])
    if not path.is_file():
        print(f"File not found: {path}", file=sys.stderr)
        return 1
    print(json.dumps(audit_workbook(path), indent=2, default=str))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

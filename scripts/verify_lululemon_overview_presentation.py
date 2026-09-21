"""Inspect Overview fiscal-label presentation and formula preservation."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--original", required=True, type=Path)
    args = parser.parse_args()
    report = {"status": "BLOCKED", "action": "", "copy": str(args.workbook)}
    books = []
    try:
        if args.workbook.resolve() == args.original.resolve():
            raise ValueError("Verification requires a separate saved copy")
        before = [digest(args.original), digest(args.workbook)]
        original = load_workbook(args.original, data_only=False)
        formulas = load_workbook(args.workbook, data_only=False)
        cached = load_workbook(args.workbook, data_only=True)
        books.extend((original, formulas, cached))
        if original.sheetnames != formulas.sheetnames:
            raise ValueError("Workbook sheet topology changed")
        formula_diffs = 0
        for sheet in original:
            saved = formulas[sheet.title]
            if sheet.sheet_state != saved.sheet_state:
                raise ValueError(f"Sheet visibility changed: {sheet.title}")
            for row in sheet.iter_rows(
                max_row=max(sheet.max_row, saved.max_row),
                max_col=max(sheet.max_column, saved.max_column),
            ):
                for cell in row:
                    other = saved[cell.coordinate]
                    if cell.data_type == "f" or other.data_type == "f":
                        if cell.value != other.value or cell.data_type != other.data_type:
                            formula_diffs += 1
        if formula_diffs:
            raise ValueError(f"Formula cells changed: {formula_diffs}")
        overview = cached["Overview"]["A3"].value
        if overview != "Historical coverage: FY2021 – FY2025 (5 periods)":
            raise ValueError(f"Overview A3 unexpected: {overview!r}")
        formula_text = formulas["Overview"]["A3"].value
        if formula_text != overview:
            raise ValueError(f"Overview A3 is not a readable literal: {formula_text!r}")
        if before != [digest(args.original), digest(args.workbook)]:
            raise ValueError("Verification inputs changed during inspection")
        report.update(
            status="VERIFIED",
            action="NONE",
            overview_a3=overview,
            formulas_preserved=True,
            formula_diffs=0,
            source_sha256=before[0],
            copy_sha256=before[1],
        )
    except (OSError, ValueError, KeyError, TypeError) as error:
        report["action"] = str(error)
    finally:
        for book in books:
            book.close()
    print(json.dumps(report, indent=2))
    return 0 if report["status"] == "VERIFIED" else 1


if __name__ == "__main__":
    raise SystemExit(main())

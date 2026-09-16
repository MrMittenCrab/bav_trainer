#!/usr/bin/env python3
"""Stage-by-stage audit of the Fast Retailing benchmark against the current engine.

Writes BASELINE.md. Does not patch production accounting logic.
"""

from __future__ import annotations

import hashlib
import json
import shutil
import sys
import tempfile
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
BENCH = ROOT / "benchmark" / "fast_retailing"
RECONCILED = BENCH / "reconciled"
STD_JSON = RECONCILED / "standardized.json"
PROV_JSON = RECONCILED / "provenance.json"
CONFLICTS_JSON = RECONCILED / "conflicts.json"
MANIFEST = BENCH / "source_manifest.json"
BASELINE = BENCH / "BASELINE.md"
EXPECTED_PRACTICE_TOTAL = 577
EXPECTED_BLANK_CHECK = (0, 0, EXPECTED_PRACTICE_TOTAL, EXPECTED_PRACTICE_TOTAL)
EXPECTED_FILLED_CHECK = (EXPECTED_PRACTICE_TOTAL, 0, 0, EXPECTED_PRACTICE_TOTAL)
PRACTICE_YELLOW_RGB = "FFFF00"
WHITE_RGBS = frozenset({"", "FFFFFF"})
YELLOW_RGBS = frozenset(
    {"FFFF00", "FFF2CC", "FFFF99", "FFEE00", "FFCC00", "FFE599"}
)
# Checkpoint-derived SHA-256 identities for the matched historical pair.
# Do not replace these from candidate files, filenames, or yellow appearance.
FROZEN_COMPATIBILITY_CHECKPOINT = "3f6f5dde023847e3347a4c830d822614a28c81a9"
FROZEN_COMPATIBILITY_TRAINER_SHA256 = (
    "546390001c69b2d05e7f16f730bacfed79a62817ea718e97bef079b4a6d014bd"
)
FROZEN_COMPATIBILITY_ANSWER_KEY_SHA256 = (
    "f01241947745e4c5db4ceff9b445af3811f41eb5e2247a8c82c372de28bdec27"
)
FILL_COMPONENT_KEYS = frozenset(
    {
        "fill_kind",
        "fill_type",
        "pattern_type",
        "fg_color",
        "bg_color",
        "degree",
        "left",
        "right",
        "top",
        "bottom",
        "stops",
    }
)


@dataclass
class StageResult:
    stage: str
    status: str  # pass | fail | skipped
    exception_type: str = ""
    message: str = ""


def _load_payload(path: Path | None = None) -> dict[str, Any]:
    return json.loads((path or STD_JSON).read_text(encoding="utf-8"))


SOURCE_STATEMENT_SHEETS: tuple[tuple[str, str], ...] = (
    ("Income Statement", "income_statement"),
    ("Balance Sheet", "balance_sheet"),
    ("Cash Flow Statement", "cash_flow"),
)
SOURCE_HEADER_ROW = 6
SOURCE_START_ROW = 7
PER_SHARE_SHEET = "Per Share Analysis"
PER_SHARE_HEADER_ROW = 4
PER_SHARE_SHARES_ROW = 7
PER_SHARE_SHARES_LABEL = "Diluted Weighted-Average Shares"
DEFERRED_TAB_NAMES = ("Model_Bear", "Model_Base", "Model_Bull", "Scenario_Summary")
LEASE_INTEREST_LABEL = "Lease interest expense (reported note)"
JUDGMENT_SHEETS = ("Accounting Judgment", "Normalization Judgment")
JUDGMENT_RESPONSE_COLS = (6, 7, 8)
JUDGMENT_FIRST_DATA_ROW = 5
_THEME_SCHEME_ORDER = (
    "lt1",
    "dk1",
    "lt2",
    "dk2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
)
_RGBMAX = 0xFF
_HLSMAX = 240


@dataclass(frozen=True)
class _FillReport:
    rgbs: tuple[str, ...]
    nonsolid: bool
    unresolved: bool
    blank: bool
    solid_fg: str


def _sha256_file(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _authenticate_frozen_compatibility_pair(
    trainer_path: Path,
    answer_key_path: Path,
) -> None:
    """Bind the historical exemption to checkpoint SHA-256 identities only."""
    trainer_hash = _sha256_file(trainer_path)
    answer_hash = _sha256_file(answer_key_path)
    if (
        trainer_hash != FROZEN_COMPATIBILITY_TRAINER_SHA256
        or answer_hash != FROZEN_COMPATIBILITY_ANSWER_KEY_SHA256
    ):
        raise ValueError(
            "frozen compatibility pair is not authenticated against "
            f"checkpoint {FROZEN_COMPATIBILITY_CHECKPOINT}: "
            f"trainer_sha256={trainer_hash} answer_sha256={answer_hash}"
        )


def _workbook_of(cell):
    parent = getattr(cell, "parent", None)
    return getattr(parent, "parent", None) if parent is not None else None


def _color_context(wb) -> tuple[tuple[str, ...], tuple[str, ...]]:
    if wb is None:
        return (), ()
    return _theme_scheme_colors(wb), _indexed_palette(wb)


def _compact_hex(value: str) -> str:
    compact = str(value or "").upper().lstrip("#")
    if len(compact) >= 6:
        return compact[-6:]
    return compact


def _color_hex(
    color,
    *,
    theme_colors: tuple[str, ...],
    palette: tuple[str, ...],
) -> tuple[str, bool]:
    """Return (RRGGBB, unresolved). Missing color is ('', False)."""
    if color is None or getattr(color, "type", None) is None:
        return "", False
    norm = _normalize_color(color, theme_colors=theme_colors, palette=palette)
    if norm is None:
        return "", False
    kind, _value, _tint, effective = norm
    if kind == "auto":
        return "", False
    hex6 = _compact_hex(str(effective or ""))
    if not hex6 or hex6 == "AUTO":
        return "", True
    return hex6, False


def _inspect_fill(
    fill,
    *,
    theme_colors: tuple[str, ...],
    palette: tuple[str, ...],
) -> _FillReport:
    if fill is None:
        return _FillReport((), False, False, True, "")
    fill_type = getattr(fill, "fill_type", None) or getattr(fill, "type", None)
    pattern = getattr(fill, "patternType", None)
    stops = getattr(fill, "stop", None)
    is_gradient = bool(stops) or fill_type in {"linear", "path"}
    if is_gradient and not hasattr(fill, "fgColor") and not hasattr(fill, "patternType"):
        rgbs: list[str] = []
        unresolved = False
        for stop in stops or ():
            rgb, bad = _color_hex(
                getattr(stop, "color", None),
                theme_colors=theme_colors,
                palette=palette,
            )
            if bad:
                unresolved = True
            elif rgb:
                rgbs.append(rgb)
        if not rgbs and not unresolved:
            unresolved = True
        return _FillReport(tuple(rgbs), True, unresolved, False, "")

    if fill_type in (None, "none") and not pattern:
        return _FillReport((), False, False, True, "")

    fg = getattr(fill, "fgColor", None) or getattr(fill, "start_color", None)
    bg = getattr(fill, "bgColor", None) or getattr(fill, "end_color", None)
    rgbs = []
    unresolved = False
    for color in (fg, bg):
        rgb, bad = _color_hex(color, theme_colors=theme_colors, palette=palette)
        if bad:
            unresolved = True
        elif rgb:
            rgbs.append(rgb)
    is_solid = fill_type == "solid" or pattern == "solid"
    if is_solid:
        fg_rgb, fg_bad = _color_hex(fg, theme_colors=theme_colors, palette=palette)
        if fg_bad:
            unresolved = True
        return _FillReport(tuple(rgbs), False, unresolved, False, fg_rgb)
    if not rgbs and not unresolved:
        unresolved = True
    return _FillReport(tuple(rgbs), True, unresolved, False, "")


def _inspect_cell_fill(cell) -> _FillReport:
    theme, palette = _color_context(_workbook_of(cell))
    return _inspect_fill(
        getattr(cell, "fill", None),
        theme_colors=theme,
        palette=palette,
    )


def _fill_has_yellow(report: _FillReport) -> bool:
    return any(_rgb_is_yellow(rgb) for rgb in report.rgbs)


def _fill_rgb(cell) -> str:
    """Resolved solid foreground RGB; non-solid/blank fills are empty."""
    report = _inspect_cell_fill(cell)
    if report.blank or report.nonsolid or report.unresolved:
        return ""
    return report.solid_fg


def _cell_has_yellow(cell) -> bool:
    return _fill_has_yellow(_inspect_cell_fill(cell))


def _cell_is_white_or_none(cell) -> bool:
    report = _inspect_cell_fill(cell)
    if report.nonsolid or report.unresolved or _fill_has_yellow(report):
        return False
    if report.blank:
        return True
    return report.solid_fg in WHITE_RGBS


def _rgb_is_yellow(rgb: str) -> bool:
    if not rgb:
        return False
    compact = _compact_hex(rgb)
    if compact in YELLOW_RGBS:
        return True
    try:
        red = int(compact[0:2], 16)
        green = int(compact[2:4], 16)
        blue = int(compact[4:6], 16)
    except ValueError:
        return False
    return (
        red >= 0xC8
        and green >= 0xC0
        and blue <= 0x80
        and (red - blue) >= 0x40
        and (green - blue) >= 0x40
    )


def _iter_conditional_rule_entries(ws):
    cf = ws.conditional_formatting
    rules_map = getattr(cf, "_cf_rules", None)
    if isinstance(rules_map, dict):
        for cf_obj, rules in rules_map.items():
            sqref = str(getattr(cf_obj, "sqref", None) or cf_obj)
            for rule in rules:
                yield sqref, rule
        return
    try:
        for cf_obj in cf:
            sqref = str(getattr(cf_obj, "sqref", None) or cf_obj)
            rules = (
                getattr(cf_obj, "cfRule", None)
                or getattr(cf_obj, "rules", None)
                or ()
            )
            for rule in rules:
                yield sqref, rule
    except TypeError:
        return


def _lookup_differential_style(wb, dxf_id):
    styles = getattr(wb, "_differential_styles", None)
    if styles is None or dxf_id is None:
        return None
    try:
        return styles[int(dxf_id)]
    except (TypeError, ValueError, IndexError, KeyError):
        return None


def _rule_differential_fill(wb, rule):
    dxf = getattr(rule, "dxf", None)
    fill = getattr(dxf, "fill", None) if dxf is not None else None
    if fill is not None:
        return fill, False
    dxf_id = getattr(rule, "dxfId", None)
    if dxf_id is None:
        return None, False
    dxf = _lookup_differential_style(wb, dxf_id)
    if dxf is None:
        return None, True
    return getattr(dxf, "fill", None), False


def _iter_conditional_color_reports(ws, wb):
    theme, palette = _color_context(wb)
    for sqref, rule in _iter_conditional_rule_entries(ws):
        fill, unresolved_dxf = _rule_differential_fill(wb, rule)
        if unresolved_dxf:
            yield sqref, _FillReport((), False, True, False, ""), "differential"
        elif fill is not None:
            yield (
                sqref,
                _inspect_fill(fill, theme_colors=theme, palette=palette),
                "differential",
            )
        color_scale = getattr(rule, "colorScale", None)
        if color_scale is not None:
            rgbs: list[str] = []
            unresolved = False
            for color in getattr(color_scale, "color", None) or ():
                rgb, bad = _color_hex(
                    color, theme_colors=theme, palette=palette
                )
                if bad:
                    unresolved = True
                elif rgb:
                    rgbs.append(rgb)
            yield (
                sqref,
                _FillReport(tuple(rgbs), True, unresolved, False, ""),
                "colorScale",
            )


def _copy_release_pair_to_temp(
    trainer_path: Path,
    answer_key_path: Path,
    tmp: Path,
) -> tuple[Path, Path]:
    """Copy Trainer/Answer Key and Answer Key sidecars for mutating Check ops."""
    trainer_copy = tmp / trainer_path.name
    answer_copy = tmp / answer_key_path.name
    shutil.copy2(trainer_path, trainer_copy)
    shutil.copy2(answer_key_path, answer_copy)
    for suffix in (".component_map.json", ".assumptions.json", ".trainer.json"):
        sidecar = answer_key_path.with_suffix(suffix)
        if sidecar.is_file():
            shutil.copy2(sidecar, tmp / sidecar.name)
    return trainer_copy, answer_copy


def _cell_addr(row: int, col: int) -> str:
    from openpyxl.utils import get_column_letter

    return f"{get_column_letter(col)}{row}"


def _as_date(value: Any):
    from datetime import date, datetime

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _judgment_case_rows(ws):
    """Yield data rows that look like judgment cases (not the zero-case message)."""
    for row in range(JUDGMENT_FIRST_DATA_ROW, (ws.max_row or 0) + 1):
        order = ws.cell(row=row, column=1).value
        label = ws.cell(row=row, column=2).value
        if isinstance(order, int) and order >= 1 and label not in (None, ""):
            yield row


def _judgment_response_coords(wb) -> set[tuple[str, int, int]]:
    coords: set[tuple[str, int, int]] = set()
    for name in JUDGMENT_SHEETS:
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        for row in _judgment_case_rows(ws):
            for col in JUDGMENT_RESPONSE_COLS:
                coords.add((name, row, col))
    return coords


def _theme_scheme_colors(wb) -> tuple[str, ...]:
    theme_bytes = getattr(wb, "loaded_theme", None)
    if not theme_bytes:
        return ()
    from openpyxl.xml.functions import QName, fromstring

    xlmns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    root = fromstring(theme_bytes)
    theme_el = root.find(QName(xlmns, "themeElements").text)
    if theme_el is None:
        return ()
    schemes = theme_el.findall(QName(xlmns, "clrScheme").text)
    if not schemes:
        return ()
    first = schemes[0]
    colors: list[str] = []
    for name in _THEME_SCHEME_ORDER:
        node = first.find(QName(xlmns, name).text)
        if node is None:
            colors.append("")
            continue
        child = next(iter(node), None)
        if child is None:
            colors.append("")
            continue
        val = child.attrib.get("val", "")
        if "window" in val:
            colors.append(str(child.attrib.get("lastClr", "")).upper())
        else:
            colors.append(str(val).upper())
    return tuple(colors)


def _indexed_palette(wb) -> tuple[str, ...]:
    colors = getattr(wb, "_colors", None)
    if colors:
        return tuple(str(c).upper() for c in colors)
    from openpyxl.styles.colors import COLOR_INDEX

    return tuple(str(c).upper() for c in COLOR_INDEX)


def _rgb_to_ms_hls(rgb_hex: str) -> tuple[int, int, int]:
    from colorsys import rgb_to_hls

    hex6 = rgb_hex.upper().lstrip("#")[-6:]
    red = int(hex6[0:2], 16) / _RGBMAX
    green = int(hex6[2:4], 16) / _RGBMAX
    blue = int(hex6[4:6], 16) / _RGBMAX
    h, l, s = rgb_to_hls(red, green, blue)
    return (int(round(h * _HLSMAX)), int(round(l * _HLSMAX)), int(round(s * _HLSMAX)))


def _ms_hls_to_rgb_hex(h: int, l: int, s: int) -> str:
    from colorsys import hls_to_rgb

    red, green, blue = hls_to_rgb(h / _HLSMAX, l / _HLSMAX, s / _HLSMAX)
    return (
        f"{int(round(red * _RGBMAX)):02X}"
        f"{int(round(green * _RGBMAX)):02X}"
        f"{int(round(blue * _RGBMAX)):02X}"
    )


def _tint_luminance(tint: float, lum: int) -> int:
    if tint < 0:
        return int(round(lum * (1.0 + tint)))
    return int(round(lum * (1.0 - tint) + (_HLSMAX - _HLSMAX * (1.0 - tint))))


def _apply_tint(rgb_hex: str, tint: float) -> str:
    hex6 = rgb_hex.upper().lstrip("#")[-6:]
    if not hex6 or abs(float(tint or 0.0)) < 1e-12:
        return hex6
    h, l, s = _rgb_to_ms_hls(hex6)
    return _ms_hls_to_rgb_hex(h, _tint_luminance(float(tint), l), s)


def _normalize_color(
    color,
    *,
    theme_colors: tuple[str, ...],
    palette: tuple[str, ...],
) -> tuple | None:
    """Normalize openpyxl Color by active type only; include resolved effective RGB."""
    if color is None:
        return None
    ctype = color.type
    tint = float(getattr(color, "tint", 0.0) or 0.0)
    if ctype == "rgb":
        raw = color.value
        rgb = str(raw).upper() if raw is not None else ""
        effective = _apply_tint(rgb[-6:], tint) if len(rgb) >= 6 else rgb
        return ("rgb", rgb, tint, effective)
    if ctype == "theme":
        theme = color.value
        base = ""
        if isinstance(theme, int) and 0 <= theme < len(theme_colors):
            base = theme_colors[theme]
        effective = _apply_tint(base, tint) if base else ""
        return ("theme", theme, tint, effective)
    if ctype == "indexed":
        indexed = color.value
        base = ""
        if isinstance(indexed, int) and 0 <= indexed < len(palette):
            base = palette[indexed]
        effective = _apply_tint(base[-6:], tint) if len(base) >= 6 else str(base).upper()
        return ("indexed", indexed, tint, effective)
    if ctype == "auto":
        return ("auto", bool(color.value), tint, "AUTO")
    return (ctype, None, tint, "")


def _side_token(side, *, theme_colors, palette) -> tuple | None:
    if side is None:
        return None
    return (
        side.style,
        _normalize_color(side.color, theme_colors=theme_colors, palette=palette),
    )


def _border_components(border, *, theme_colors, palette) -> dict[str, Any]:
    if border is None:
        return {
            "border_flags": None,
            "border_left": None,
            "border_right": None,
            "border_top": None,
            "border_bottom": None,
            "border_diagonal": None,
            "border_vertical": None,
            "border_horizontal": None,
            "border_start": None,
            "border_end": None,
        }
    return {
        "border_flags": (
            bool(border.outline),
            bool(border.diagonalUp),
            bool(border.diagonalDown),
        ),
        "border_left": _side_token(
            border.left, theme_colors=theme_colors, palette=palette
        ),
        "border_right": _side_token(
            border.right, theme_colors=theme_colors, palette=palette
        ),
        "border_top": _side_token(
            border.top, theme_colors=theme_colors, palette=palette
        ),
        "border_bottom": _side_token(
            border.bottom, theme_colors=theme_colors, palette=palette
        ),
        "border_diagonal": _side_token(
            border.diagonal, theme_colors=theme_colors, palette=palette
        ),
        "border_vertical": _side_token(
            getattr(border, "vertical", None),
            theme_colors=theme_colors,
            palette=palette,
        ),
        "border_horizontal": _side_token(
            getattr(border, "horizontal", None),
            theme_colors=theme_colors,
            palette=palette,
        ),
        "border_start": _side_token(
            getattr(border, "start", None),
            theme_colors=theme_colors,
            palette=palette,
        ),
        "border_end": _side_token(
            getattr(border, "end", None),
            theme_colors=theme_colors,
            palette=palette,
        ),
    }


def _fill_components(fill, *, theme_colors, palette) -> dict[str, Any]:
    if fill is None:
        return {"fill_kind": None}
    fill_type = getattr(fill, "fill_type", None) or getattr(fill, "type", None)
    pattern = getattr(fill, "patternType", None) or getattr(fill, "fill_type", None)
    # PatternFill
    if hasattr(fill, "fgColor") or hasattr(fill, "patternType"):
        fg = getattr(fill, "fgColor", None) or getattr(fill, "start_color", None)
        bg = getattr(fill, "bgColor", None) or getattr(fill, "end_color", None)
        return {
            "fill_kind": "pattern",
            "fill_type": fill_type,
            "pattern_type": pattern,
            "fg_color": _normalize_color(
                fg, theme_colors=theme_colors, palette=palette
            ),
            "bg_color": _normalize_color(
                bg, theme_colors=theme_colors, palette=palette
            ),
        }
    # GradientFill
    stops = []
    for stop in getattr(fill, "stop", None) or ():
        stops.append(
            (
                getattr(stop, "position", None),
                _normalize_color(
                    getattr(stop, "color", None),
                    theme_colors=theme_colors,
                    palette=palette,
                ),
            )
        )
    return {
        "fill_kind": "gradient",
        "fill_type": fill_type,
        "degree": getattr(fill, "degree", None),
        "left": getattr(fill, "left", None),
        "right": getattr(fill, "right", None),
        "top": getattr(fill, "top", None),
        "bottom": getattr(fill, "bottom", None),
        "stops": tuple(stops),
    }


def _format_components(
    cell,
    *,
    theme_colors: tuple[str, ...],
    palette: tuple[str, ...],
) -> dict[str, Any]:
    font = cell.font
    alignment = cell.alignment
    protection = cell.protection
    comps: dict[str, Any] = {
        "number_format": cell.number_format,
        "font_name": font.name if font else None,
        "font_size": font.size if font else None,
        "font_bold": bool(font.bold) if font else False,
        "font_italic": bool(font.italic) if font else False,
        "font_underline": font.underline if font else None,
        "font_strikethrough": bool(font.strike) if font else False,
        "font_vert_align": font.vertAlign if font else None,
        "font_outline": bool(font.outline) if font else False,
        "font_shadow": bool(font.shadow) if font else False,
        "font_condense": bool(font.condense) if font else False,
        "font_extend": bool(font.extend) if font else False,
        "font_family": font.family if font else None,
        "font_charset": font.charset if font else None,
        "font_scheme": font.scheme if font else None,
        "font_color": _normalize_color(
            font.color if font else None,
            theme_colors=theme_colors,
            palette=palette,
        ),
        "align_horizontal": alignment.horizontal if alignment else None,
        "align_vertical": alignment.vertical if alignment else None,
        "align_wrap_text": bool(alignment.wrap_text) if alignment else False,
        "align_shrink_to_fit": bool(alignment.shrink_to_fit) if alignment else False,
        "align_text_rotation": alignment.textRotation if alignment else 0,
        "align_indent": alignment.indent if alignment else 0,
        "align_relative_indent": (
            alignment.relativeIndent if alignment else 0
        ),
        "align_justify_last_line": (
            bool(alignment.justifyLastLine) if alignment else False
        ),
        "align_reading_order": alignment.readingOrder if alignment else 0,
        "protection_locked": bool(protection.locked) if protection else True,
        "protection_hidden": bool(protection.hidden) if protection else False,
    }
    comps.update(
        _fill_components(cell.fill, theme_colors=theme_colors, palette=palette)
    )
    comps.update(
        _border_components(cell.border, theme_colors=theme_colors, palette=palette)
    )
    return comps


def _assert_format_parity(
    ct,
    ca,
    *,
    sheet: str,
    row: int,
    col: int,
    theme_t: tuple[str, ...],
    theme_a: tuple[str, ...],
    palette_t: tuple[str, ...],
    palette_a: tuple[str, ...],
    skip_fill: bool = False,
) -> None:
    comps_t = _format_components(ct, theme_colors=theme_t, palette=palette_t)
    comps_a = _format_components(ca, theme_colors=theme_a, palette=palette_a)
    keys = list(dict.fromkeys([*comps_t.keys(), *comps_a.keys()]))
    for key in keys:
        if skip_fill and key in FILL_COMPONENT_KEYS:
            continue
        vt = comps_t.get(key)
        va = comps_a.get(key)
        if vt != va:
            raise ValueError(
                f"effective formatting mismatch: sheet={sheet!r} "
                f"cell={_cell_addr(row, col)} component={key} "
                f"trainer={vt!r} answer={va!r}"
            )


def _comment_text(cell) -> str:
    if cell.comment is None:
        return ""
    return str(cell.comment.text or "")


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _source_failure(
    workbook: str,
    sheet: str,
    row: int,
    col: int,
    source_identity: str,
    period: Any,
    expected: Any,
    actual: Any,
    reason: str = "",
) -> ValueError:
    period_s = "" if period is None else str(period)
    detail = reason or "value mismatch"
    return ValueError(
        f"{detail}: workbook={workbook} sheet={sheet!r} cell={_cell_addr(row, col)} "
        f"source_identity={source_identity!r} period={period_s!r} "
        f"expected={expected!r} actual={actual!r}"
    )


def _canonical_periods(fin) -> list:
    from core.model.period_axis import canonical_fiscal_periods

    return canonical_fiscal_periods(fin)


def _verify_period_headers(
    ws,
    *,
    workbook: str,
    sheet: str,
    header_row: int,
    periods: list,
    start_col: int = 2,
) -> None:
    for j, expected in enumerate(periods):
        col = start_col + j
        actual = ws.cell(row=header_row, column=col).value
        actual_date = _as_date(actual)
        if actual_date != expected:
            raise _source_failure(
                workbook,
                sheet,
                header_row,
                col,
                "period_header",
                expected,
                expected,
                actual,
                reason="period header mismatch",
            )


def _verify_source_value(
    cell,
    *,
    workbook: str,
    sheet: str,
    row: int,
    col: int,
    source_identity: str,
    period: Any,
    expected: Any,
) -> None:
    actual = cell.value
    if _is_formula(actual):
        raise _source_failure(
            workbook,
            sheet,
            row,
            col,
            source_identity,
            period,
            expected,
            actual,
            reason="source fact replaced by formula",
        )
    if expected is None:
        if actual is not None:
            raise _source_failure(
                workbook,
                sheet,
                row,
                col,
                source_identity,
                period,
                None,
                actual,
                reason="missing source fact blanked incorrectly",
            )
        return
    if actual is None:
        raise _source_failure(
            workbook,
            sheet,
            row,
            col,
            source_identity,
            period,
            expected,
            actual,
            reason="required source fact blanked",
        )
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        if float(actual) != float(expected):
            raise _source_failure(
                workbook,
                sheet,
                row,
                col,
                source_identity,
                period,
                expected,
                actual,
            )
        return
    if actual != expected:
        raise _source_failure(
            workbook,
            sheet,
            row,
            col,
            source_identity,
            period,
            expected,
            actual,
        )


def _verify_statement_sheet(ws, items, periods, *, workbook: str, sheet: str) -> None:
    from core.data.line_identity import line_identity

    _verify_period_headers(
        ws,
        workbook=workbook,
        sheet=sheet,
        header_row=SOURCE_HEADER_ROW,
        periods=periods,
    )
    for idx, item in enumerate(items):
        row = SOURCE_START_ROW + idx
        identity = line_identity(item).key()
        label_cell = ws.cell(row=row, column=1)
        if label_cell.value != item.label:
            raise _source_failure(
                workbook,
                sheet,
                row,
                1,
                identity,
                None,
                item.label,
                label_cell.value,
                reason="source row label mismatch",
            )
        for j, period in enumerate(periods):
            col = 2 + j
            expected = item.values.get(period)
            _verify_source_value(
                ws.cell(row=row, column=col),
                workbook=workbook,
                sheet=sheet,
                row=row,
                col=col,
                source_identity=identity,
                period=period,
                expected=expected,
            )


def _verify_workbook_source_fidelity(wb, fin, *, workbook: str) -> None:
    periods = _canonical_periods(fin)
    statement_attrs = {
        "Income Statement": fin.income_statement,
        "Balance Sheet": fin.balance_sheet,
        "Cash Flow Statement": fin.cash_flow,
    }
    for sheet, _attr in SOURCE_STATEMENT_SHEETS:
        if sheet not in wb.sheetnames:
            raise ValueError(
                f"missing required source sheet: workbook={workbook} sheet={sheet!r}"
            )
        ws = wb[sheet]
        expected_units = f"Units: {fin.units}"
        actual_units = ws.cell(row=3, column=1).value
        if actual_units != expected_units:
            raise _source_failure(
                workbook,
                sheet,
                3,
                1,
                "units",
                None,
                expected_units,
                actual_units,
                reason="units mismatch",
            )
        _verify_statement_sheet(
            ws,
            statement_attrs[sheet],
            periods,
            workbook=workbook,
            sheet=sheet,
        )
        if sheet == "Income Statement" and fin.historical_lease is not None:
            lease_row = None
            for row in range(SOURCE_START_ROW, (ws.max_row or SOURCE_START_ROW) + 1):
                if ws.cell(row=row, column=1).value == LEASE_INTEREST_LABEL:
                    lease_row = row
                    break
            if lease_row is None:
                raise ValueError(
                    f"missing lease interest source row: workbook={workbook} "
                    f"sheet={sheet!r} source_identity={LEASE_INTEREST_LABEL!r}"
                )
            for j, period in enumerate(periods):
                col = 2 + j
                expected = fin.historical_lease.lease_interest_expense.get(period)
                expected = None if expected is None else float(expected)
                _verify_source_value(
                    ws.cell(row=lease_row, column=col),
                    workbook=workbook,
                    sheet=sheet,
                    row=lease_row,
                    col=col,
                    source_identity="historical_lease.lease_interest_expense",
                    period=period,
                    expected=expected,
                )

    if fin.historical_shares is None:
        return
    if PER_SHARE_SHEET not in wb.sheetnames:
        raise ValueError(
            f"missing required source sheet: workbook={workbook} "
            f"sheet={PER_SHARE_SHEET!r}"
        )
    ws = wb[PER_SHARE_SHEET]
    _verify_period_headers(
        ws,
        workbook=workbook,
        sheet=PER_SHARE_SHEET,
        header_row=PER_SHARE_HEADER_ROW,
        periods=periods,
    )
    label = ws.cell(row=PER_SHARE_SHARES_ROW, column=1).value
    if label != PER_SHARE_SHARES_LABEL:
        raise _source_failure(
            workbook,
            PER_SHARE_SHEET,
            PER_SHARE_SHARES_ROW,
            1,
            "historical_shares.diluted_weighted_average",
            None,
            PER_SHARE_SHARES_LABEL,
            label,
            reason="historical shares label mismatch",
        )
    for j, period in enumerate(periods):
        col = 2 + j
        expected = fin.historical_shares.diluted_weighted_average.get(period)
        _verify_source_value(
            ws.cell(row=PER_SHARE_SHARES_ROW, column=col),
            workbook=workbook,
            sheet=PER_SHARE_SHEET,
            row=PER_SHARE_SHARES_ROW,
            col=col,
            source_identity="historical_shares.diluted_weighted_average",
            period=period,
            expected=expected,
        )


def _comparable_sheet_names(wb) -> list[str]:
    """Non-metadata sheets in workbook order (includes deferred placeholders)."""
    return [name for name in wb.sheetnames if not name.startswith("_")]


def _requires_visible(name: str) -> bool:
    return not name.startswith("_") and name not in DEFERRED_TAB_NAMES


def _verify_required_visibility(wb, *, workbook: str) -> None:
    for sheet, _attr in SOURCE_STATEMENT_SHEETS:
        if sheet not in wb.sheetnames:
            raise ValueError(
                f"missing required source sheet: workbook={workbook} sheet={sheet!r}"
            )
    for name in wb.sheetnames:
        if not _requires_visible(name):
            continue
        state = wb[name].sheet_state
        if state != "visible":
            raise ValueError(
                f"required historical/practice sheet hidden: workbook={workbook} "
                f"sheet={name!r} sheet_state={state!r}"
            )


def _dim_hidden(dimension) -> bool:
    return bool(getattr(dimension, "hidden", False))


def _verify_visible_layout_parity(
    wb_t,
    wb_a,
    practice_coords: set[tuple[str, int, int]],
    *,
    fill_exempt: set[tuple[str, int, int]] | None = None,
) -> None:
    from openpyxl.utils import get_column_letter

    comparable_t = _comparable_sheet_names(wb_t)
    comparable_a = _comparable_sheet_names(wb_a)
    if comparable_t != comparable_a:
        raise ValueError(
            f"sheet order/state mismatch: trainer={comparable_t} answer={comparable_a}"
        )
    for name in comparable_t:
        if wb_t[name].sheet_state != wb_a[name].sheet_state:
            raise ValueError(
                f"sheet_state mismatch for {name!r}: "
                f"trainer={wb_t[name].sheet_state!r} answer={wb_a[name].sheet_state!r}"
            )

    judgment_t = _judgment_response_coords(wb_t)
    judgment_a = _judgment_response_coords(wb_a)
    if judgment_t != judgment_a:
        raise ValueError(
            f"judgment response coordinate mismatch: "
            f"trainer={sorted(judgment_t)} answer={sorted(judgment_a)}"
        )
    content_exempt = set(practice_coords) | judgment_t
    fill_exempt = set(fill_exempt or ())
    theme_t = _theme_scheme_colors(wb_t)
    theme_a = _theme_scheme_colors(wb_a)
    palette_t = _indexed_palette(wb_t)
    palette_a = _indexed_palette(wb_a)

    visible = [
        name
        for name in comparable_t
        if wb_t[name].sheet_state == "visible" and wb_a[name].sheet_state == "visible"
    ]
    for name in visible:
        ws_t = wb_t[name]
        ws_a = wb_a[name]
        if (ws_t.max_row, ws_t.max_column) != (ws_a.max_row, ws_a.max_column):
            raise ValueError(
                f"visible dimensions mismatch on {name!r}: "
                f"trainer={(ws_t.max_row, ws_t.max_column)} "
                f"answer={(ws_a.max_row, ws_a.max_column)}"
            )
        if ws_t.freeze_panes != ws_a.freeze_panes:
            raise ValueError(
                f"freeze_panes mismatch on {name!r}: "
                f"trainer={ws_t.freeze_panes!r} answer={ws_a.freeze_panes!r}"
            )
        if list(ws_t.merged_cells.ranges) != list(ws_a.merged_cells.ranges):
            raise ValueError(f"merged ranges mismatch on {name!r}")
        max_row = max(ws_t.max_row or 1, ws_a.max_row or 1)
        max_col = max(ws_t.max_column or 1, ws_a.max_column or 1)
        for row in range(1, max_row + 1):
            if _dim_hidden(ws_t.row_dimensions[row]) != _dim_hidden(
                ws_a.row_dimensions[row]
            ):
                raise ValueError(f"row hidden mismatch on {name!r} row={row}")
            if ws_t.row_dimensions[row].height != ws_a.row_dimensions[row].height:
                raise ValueError(f"row height mismatch on {name!r} row={row}")
        for col in range(1, max_col + 1):
            letter = get_column_letter(col)
            if _dim_hidden(ws_t.column_dimensions[letter]) != _dim_hidden(
                ws_a.column_dimensions[letter]
            ):
                raise ValueError(f"column hidden mismatch on {name!r} col={letter}")
            if ws_t.column_dimensions[letter].width != ws_a.column_dimensions[letter].width:
                raise ValueError(f"column width mismatch on {name!r} col={letter}")
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                ct = ws_t.cell(row=row, column=col)
                ca = ws_a.cell(row=row, column=col)
                if (name, row, col) not in content_exempt:
                    if ct.value != ca.value:
                        raise ValueError(
                            f"non-practice value mismatch: sheet={name!r} "
                            f"cell={_cell_addr(row, col)} "
                            f"trainer={ct.value!r} answer={ca.value!r}"
                        )
                    if _comment_text(ct) != _comment_text(ca):
                        raise ValueError(
                            f"non-practice Note mismatch: sheet={name!r} "
                            f"cell={_cell_addr(row, col)}"
                        )
                _assert_format_parity(
                    ct,
                    ca,
                    sheet=name,
                    row=row,
                    col=col,
                    theme_t=theme_t,
                    theme_a=theme_a,
                    palette_t=palette_t,
                    palette_a=palette_a,
                    skip_fill=(name, row, col) in fill_exempt,
                )


def _verify_answer_key_no_yellow(wb_a) -> None:
    """Reject yellow fill or yellow conditional highlighting anywhere on the Answer Key."""
    theme, palette = _color_context(wb_a)
    for ws in wb_a.worksheets:
        for sqref, report, kind in _iter_conditional_color_reports(ws, wb_a):
            if report.unresolved:
                raise ValueError(
                    f"Answer Key unresolved conditional formatting: sheet={ws.title!r} "
                    f"range={sqref} kind={kind}"
                )
            if _fill_has_yellow(report):
                raise ValueError(
                    f"Answer Key yellow conditional formatting: sheet={ws.title!r} "
                    f"range={sqref} kind={kind} fill={report.rgbs!r}"
                )
        max_row = ws.max_row or 1
        max_col = ws.max_column or 1
        for row in range(1, max_row + 1):
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row, column=col)
                report = _inspect_fill(
                    cell.fill, theme_colors=theme, palette=palette
                )
                if _fill_has_yellow(report):
                    raise ValueError(
                        f"Answer Key yellow fill: sheet={ws.title!r} "
                        f"cell={_cell_addr(row, col)}"
                    )


def _verify_practice_contract(
    wb_t,
    wb_a,
    comps,
    *,
    allow_frozen_yellow_answer_key: bool = False,
) -> None:
    from core.trainer.semantic_io import parse_cell_ref

    for comp in comps:
        row, col = parse_cell_ref(comp.cell)
        tc = wb_t[comp.tab].cell(row=row, column=col)
        ac = wb_a[comp.tab].cell(row=row, column=col)
        if tc.value is not None:
            raise ValueError(f"Trainer practice cell {comp.tab}!{comp.cell} not blank")
        if tc.comment is not None:
            raise ValueError(f"Trainer practice cell {comp.tab}!{comp.cell} has Note")
        if _fill_rgb(tc) != PRACTICE_YELLOW_RGB:
            raise ValueError(f"Trainer practice cell {comp.tab}!{comp.cell} not yellow")
        if not (isinstance(ac.value, str) and ac.value.startswith("=")):
            raise ValueError(f"Answer Key {comp.tab}!{comp.cell} missing formula")
        if ac.value != comp.formula:
            raise ValueError(
                f"Answer Key {comp.tab}!{comp.cell} formula != semantic map"
            )
        if ac.comment is None or not str(ac.comment.text or "").strip():
            raise ValueError(
                f"Answer Key {comp.tab}!{comp.cell} missing non-empty Note"
            )
        if allow_frozen_yellow_answer_key:
            if _fill_rgb(ac) != PRACTICE_YELLOW_RGB:
                raise ValueError(
                    f"Answer Key practice cell {comp.tab}!{comp.cell} not yellow"
                )
        else:
            if _cell_has_yellow(ac) or not _cell_is_white_or_none(ac):
                raise ValueError(
                    f"Answer Key practice cell {comp.tab}!{comp.cell} not white/no-fill"
                )

    judgment = _judgment_response_coords(wb_a)
    if judgment != _judgment_response_coords(wb_t):
        raise ValueError(
            f"judgment response coordinate mismatch: "
            f"trainer={sorted(_judgment_response_coords(wb_t))} "
            f"answer={sorted(judgment)}"
        )
    if allow_frozen_yellow_answer_key:
        return
    for sheet, row, col in sorted(judgment):
        tc = wb_t[sheet].cell(row=row, column=col)
        ac = wb_a[sheet].cell(row=row, column=col)
        addr = _cell_addr(row, col)
        if tc.value is not None:
            raise ValueError(f"Trainer judgment cell {sheet}!{addr} not blank")
        if tc.comment is not None:
            raise ValueError(f"Trainer judgment cell {sheet}!{addr} has Note")
        if _fill_rgb(tc) != PRACTICE_YELLOW_RGB:
            raise ValueError(f"Trainer judgment cell {sheet}!{addr} not yellow")
        if ac.value in (None, ""):
            raise ValueError(f"Answer Key judgment cell {sheet}!{addr} missing response")
        if _cell_has_yellow(ac) or not _cell_is_white_or_none(ac):
            raise ValueError(
                f"Answer Key judgment cell {sheet}!{addr} not white/no-fill"
            )


def _assert_frozen_yellow_answer_key_signature(wb_a, comps) -> None:
    """Refuse the frozen-yellow exception unless Answer-Key practice cells are yellow."""
    from core.trainer.semantic_io import parse_cell_ref

    for comp in comps:
        row, col = parse_cell_ref(comp.cell)
        ac = wb_a[comp.tab].cell(row=row, column=col)
        if _fill_rgb(ac) != PRACTICE_YELLOW_RGB:
            raise ValueError(
                "frozen yellow Answer Key exception does not apply to "
                "current-style pairs"
            )


def _verify_release_pair_contents(
    trainer_path: Path,
    answer_key_path: Path,
    fin,
    *,
    skip_answer_key_yellow: bool = False,
) -> str:
    """Source fidelity, practice contract, visibility, and visible structural parity."""
    from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
    from openpyxl import load_workbook

    smap = load_semantic_map(answer_key_path)
    comps = smap.all_ordered()
    if len(comps) != EXPECTED_PRACTICE_TOTAL:
        raise ValueError(
            f"semantic practice cells={len(comps)} != {EXPECTED_PRACTICE_TOTAL}"
        )
    practice_coords = {
        (comp.tab, *parse_cell_ref(comp.cell)) for comp in comps
    }

    wb_t = load_workbook(trainer_path, data_only=False)
    wb_a = load_workbook(answer_key_path, data_only=False)
    try:
        _verify_workbook_source_fidelity(wb_t, fin, workbook="Trainer")
        _verify_workbook_source_fidelity(wb_a, fin, workbook="Answer Key")
        _verify_required_visibility(wb_t, workbook="Trainer")
        _verify_required_visibility(wb_a, workbook="Answer Key")
        if skip_answer_key_yellow:
            _assert_frozen_yellow_answer_key_signature(wb_a, comps)
        else:
            _verify_answer_key_no_yellow(wb_a)
        judgment_coords = _judgment_response_coords(wb_a)
        fill_exempt = (
            set()
            if skip_answer_key_yellow
            else set(practice_coords) | judgment_coords
        )
        _verify_visible_layout_parity(
            wb_t, wb_a, practice_coords, fill_exempt=fill_exempt
        )
        _verify_practice_contract(
            wb_t,
            wb_a,
            comps,
            allow_frozen_yellow_answer_key=skip_answer_key_yellow,
        )
    finally:
        wb_t.close()
        wb_a.close()
    return (
        f"practice_cells={len(comps)} source_fidelity=ok "
        f"visibility=ok layout_parity=ok"
    )


def _verify_release_pair_contract(
    trainer_path: Path,
    answer_key_path: Path,
    fin,
    *,
    allow_frozen_yellow_answer_key: bool = False,
) -> str:
    """Public pair contract. Frozen yellow exemption requires authenticated hashes."""
    if allow_frozen_yellow_answer_key:
        _authenticate_frozen_compatibility_pair(trainer_path, answer_key_path)
    return _verify_release_pair_contents(
        trainer_path,
        answer_key_path,
        fin,
        skip_answer_key_yellow=allow_frozen_yellow_answer_key,
    )


def _module_applicability(fin, anchor=None) -> dict[str, Any]:
    from core.model.earnings_quality import earnings_quality_availability
    from core.model.deferred_tax import (
        deferred_tax_applicable,
        deferred_tax_availability,
    )
    from core.model.fixed_asset import fixed_asset_applicable
    from core.model.goodwill_intangibles import (
        goodwill_intangibles_applicable,
        goodwill_intangibles_availability,
    )
    from core.model.lease_liability import (
        lease_liability_applicable,
        lease_liability_availability,
    )
    from core.model.lease_rou import (
        lease_rou_applicable,
        lease_rou_availability,
    )
    from core.model.ownership_attribution import (
        ownership_attribution_applicable,
        ownership_attribution_availability,
    )
    from core.model.per_share import per_share_available
    from core.model.working_capital import working_capital_applicable

    lease_avail = lease_liability_availability(fin)
    lease_rou_avail = lease_rou_availability(fin)
    ownership_avail = ownership_attribution_availability(fin)
    gi_avail = goodwill_intangibles_availability(fin)
    deferred_tax_avail = deferred_tax_availability(fin)
    eq = earnings_quality_availability(fin)
    wc_applicable: bool | None
    if anchor is None:
        wc_applicable = None
    else:
        wc_applicable = working_capital_applicable(anchor)
    return {
        "earnings_quality": {
            "applicable": bool(eq.operating_cash_flow),
            "detail": asdict(eq),
        },
        "working_capital": {
            "applicable": wc_applicable,
            "detail": "requires AnchorMetrics; filled after successful builder"
            if anchor is None
            else "",
        },
        "fixed_asset": {"applicable": fixed_asset_applicable(fin)},
        "lease_liability": {
            "applicable": lease_liability_applicable(fin),
            "availability": asdict(lease_avail),
        },
        "lease_rou": {
            "applicable": lease_rou_applicable(fin),
            "availability": asdict(lease_rou_avail),
        },
        "ownership_attribution": {
            "applicable": ownership_attribution_applicable(fin),
            "availability": asdict(ownership_avail),
        },
        "goodwill_intangibles": {
            "applicable": goodwill_intangibles_applicable(fin),
            "availability": asdict(gi_avail),
        },
        "deferred_tax": {
            "applicable": deferred_tax_applicable(fin),
            "availability": asdict(deferred_tax_avail),
        },
        "per_share": {"applicable": per_share_available(fin)},
        "normalization": {
            "applicable": False,
            "detail": "no normalization assumptions supplied for benchmark audit",
        },
    }


def run_audit(
    *,
    standardized_json: Path | None = None,
    provenance_json: Path | None = None,
    conflicts_json: Path | None = None,
    trainer_path: Path | None = None,
    answer_key_path: Path | None = None,
    require_check_counts: bool | None = None,
    verify_release_pair: bool = False,
    allow_frozen_yellow_answer_key: bool = False,
) -> dict[str, Any]:
    """Run the Fast Retailing stage audit.

    Default call with no arguments preserves the historical temporary-workbook
    interface used by existing tests. Pass explicit release workbook paths and
    generated standardized JSON to verify a persisted pair without regenerating
    substitute workbooks.
    """
    std_path = Path(standardized_json) if standardized_json else STD_JSON
    explicit_pair = trainer_path is not None or answer_key_path is not None
    if explicit_pair:
        if trainer_path is None or answer_key_path is None:
            raise ValueError("trainer_path and answer_key_path must be provided together")
        trainer_path = Path(trainer_path)
        answer_key_path = Path(answer_key_path)
        if require_check_counts is None:
            require_check_counts = True
    elif require_check_counts is None:
        require_check_counts = False

    stages: list[StageResult] = []
    context: dict[str, Any] = {
        "modules": {},
        "first_failure": None,
        "workbook_paths": {},
        "release_fingerprints_before": {},
        "release_fingerprints_after": {},
        "standardized_json": str(std_path),
        "provenance_json": str(provenance_json or PROV_JSON),
        "conflicts_json": str(conflicts_json or CONFLICTS_JSON),
    }

    def _fingerprint(path: Path) -> str:
        return hashlib.sha256(path.read_bytes()).hexdigest()

    release_files: list[Path] = []
    if explicit_pair:
        release_files = [
            trainer_path,
            answer_key_path,
            answer_key_path.with_suffix(".component_map.json"),
        ]
        for path in (trainer_path, answer_key_path):
            if not path.is_file():
                stages.append(
                    StageResult(
                        "1_source_fixture_load",
                        "fail",
                        "FileNotFoundError",
                        f"missing release artifact: {path}",
                    )
                )
                context["first_failure"] = stages[-1]
                for name in (
                    "2_identity_validation",
                    "3_reconciliation",
                    "4_reference_model_builder",
                    "5_workbook_generation",
                    "6_blank_check",
                    "7_filled_check",
                ):
                    stages.append(
                        StageResult(name, "skipped", message="prior stage failed")
                    )
                return {"stages": stages, "context": context}
        context["release_fingerprints_before"] = {
            str(p): _fingerprint(p) for p in release_files if p.is_file()
        }

    # Stage 1
    try:
        if not std_path.is_file():
            raise FileNotFoundError(f"missing standardized input: {std_path}")
        payload = _load_payload(std_path)
        assert payload["ticker"] == "6288.HK"
        assert payload.get("company_name") == "FAST RETAILING CO., LTD."
        assert [p["end_date"] for p in payload["periods"]] == [
            "2021-08-31",
            "2022-08-31",
            "2023-08-31",
            "2024-08-31",
            "2025-08-31",
        ]
        stages.append(StageResult("1_source_fixture_load", "pass", message="loaded"))
    except Exception as exc:  # noqa: BLE001 — audit must capture all failures
        stages.append(
            StageResult(
                "1_source_fixture_load",
                "fail",
                type(exc).__name__,
                str(exc),
            )
        )
        context["first_failure"] = stages[-1]
        for name in (
            "2_identity_validation",
            "3_reconciliation",
            "4_reference_model_builder",
            "5_workbook_generation",
            "6_blank_check",
            "7_filled_check",
        ):
            stages.append(StageResult(name, "skipped", message="prior stage failed"))
        return {"stages": stages, "context": context}

    from core.data.standardized_io import standardized_from_payload

    fin = standardized_from_payload(payload)

    # Stage 2
    try:
        from core.data.line_identity import validate_financials_identities

        validate_financials_identities(fin)
        stages.append(StageResult("2_identity_validation", "pass"))
    except Exception as exc:  # noqa: BLE001
        stages.append(
            StageResult("2_identity_validation", "fail", type(exc).__name__, str(exc))
        )
        context["first_failure"] = stages[-1]
        for name in (
            "3_reconciliation",
            "4_reference_model_builder",
            "5_workbook_generation",
            "6_blank_check",
            "7_filled_check",
        ):
            stages.append(StageResult(name, "skipped", message="prior stage failed"))
        context["modules"] = _module_applicability(fin)
        return {"stages": stages, "context": context}

    # Stage 3
    try:
        from core.ingestion.reconciler import reconcile_financials

        report = reconcile_financials(fin)
        if not all(report.checksums.values()):
            raise ValueError(
                f"reconciliation checksums={report.checksums} warnings={report.warnings}"
            )
        stages.append(StageResult("3_reconciliation", "pass"))
    except Exception as exc:  # noqa: BLE001
        stages.append(
            StageResult("3_reconciliation", "fail", type(exc).__name__, str(exc))
        )
        if context["first_failure"] is None:
            context["first_failure"] = stages[-1]
        # Continue measuring later stages where possible — identity already passed.
        # But workbook generation may still be attempted; mark dependent skips only
        # if builder cannot run. We still try builder for gap visibility.

    context["modules"] = _module_applicability(fin)

    # Stage 4
    builder = None
    try:
        from core.engine.reference_model import ReferenceModelBuilder

        builder = ReferenceModelBuilder(fin)
        stages.append(
            StageResult(
                "4_reference_model_builder",
                "pass",
                message=(
                    f"expected_specs={len(builder.expected_specs)} "
                    f"lease_specs={len(builder.lease_liability_specs)} "
                    f"lease_rou_specs={len(builder.lease_rou_specs)} "
                    f"ownership_specs={len(builder.ownership_attribution_specs)} "
                    f"per_share_specs={len(builder.per_share_specs)} "
                    f"per_share_attribution_specs={len(builder.per_share_attribution_specs)} "
                    f"fixed_asset_specs={len(builder.fixed_asset_specs)} "
                    f"goodwill_intangibles_specs={len(builder.goodwill_intangibles_specs)} "
                    f"deferred_tax_specs={len(builder.deferred_tax_specs)} "
                    f"capex_specs={len(builder.capex_specs)}"
                ),
            )
        )
        context["modules"] = _module_applicability(fin, builder.anchor)
    except Exception as exc:  # noqa: BLE001
        stages.append(
            StageResult(
                "4_reference_model_builder",
                "fail",
                type(exc).__name__,
                str(exc),
            )
        )
        if context["first_failure"] is None:
            context["first_failure"] = stages[-1]
        for name in ("5_workbook_generation", "6_blank_check", "7_filled_check"):
            stages.append(StageResult(name, "skipped", message="prior stage failed"))
        return {"stages": stages, "context": context}

    # Stage 5–7: either use persisted release pair or generate temporary workbooks.
    try:
        from core.trainer.checker import check_workbook
        from core.trainer.semantic_io import load_semantic_map, parse_cell_ref
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory(prefix="fr_bench_") as tmp:
            tmp_path = Path(tmp)
            if explicit_pair:
                assert trainer_path is not None and answer_key_path is not None
                if verify_release_pair:
                    contract_msg = _verify_release_pair_contract(
                        trainer_path,
                        answer_key_path,
                        fin,
                        allow_frozen_yellow_answer_key=allow_frozen_yellow_answer_key,
                    )
                else:
                    contract_msg = "release pair supplied"
                # Mutating Check / fill on copies only.
                trainer, answer = _copy_release_pair_to_temp(
                    trainer_path, answer_key_path, tmp_path
                )
                context["workbook_paths"] = {
                    "trainer": str(trainer_path),
                    "answer": str(answer_key_path),
                    "check_copies": str(tmp_path),
                    "note": "persisted release pair; Check used temporary copies",
                    "contract": contract_msg,
                }
                stages.append(
                    StageResult(
                        "5_workbook_generation",
                        "pass",
                        message="persisted release pair (not regenerated)",
                    )
                )
            else:
                from core.trainer.workbook import build_training_workbook

                trainer, answer = build_training_workbook(
                    fin, tmp_path / "FastRetailing_Trainer.xlsx"
                )
                context["workbook_paths"] = {
                    "trainer": str(trainer),
                    "answer": str(answer),
                    "note": "temporary only; not committed",
                }
                stages.append(StageResult("5_workbook_generation", "pass"))

            # Stage 6
            try:
                summary = check_workbook(trainer)
                blank_tuple = (
                    summary.correct,
                    summary.incorrect,
                    summary.blank,
                    summary.total,
                )
                if require_check_counts and blank_tuple != EXPECTED_BLANK_CHECK:
                    raise ValueError(
                        f"pristine Check counts {blank_tuple} != {EXPECTED_BLANK_CHECK}"
                    )
                stages.append(
                    StageResult(
                        "6_blank_check",
                        "pass",
                        message=(
                            f"correct={summary.correct} incorrect={summary.incorrect} "
                            f"blank={summary.blank} total={summary.total}"
                        ),
                    )
                )
            except Exception as exc:  # noqa: BLE001
                stages.append(
                    StageResult("6_blank_check", "fail", type(exc).__name__, str(exc))
                )
                if context["first_failure"] is None:
                    context["first_failure"] = stages[-1]
                stages.append(
                    StageResult("7_filled_check", "skipped", message="prior stage failed")
                )
                if explicit_pair:
                    context["release_fingerprints_after"] = {
                        str(p): _fingerprint(p)
                        for p in release_files
                        if p.is_file()
                    }
                return {"stages": stages, "context": context}

            # Stage 7
            try:
                smap = load_semantic_map(answer)
                wb = load_workbook(trainer, data_only=False)
                for comp in smap.all_ordered():
                    row, col = parse_cell_ref(comp.cell)
                    wb[comp.tab].cell(row=row, column=col).value = comp.formula
                wb.save(trainer)
                wb.close()
                filled = check_workbook(trainer)
                filled_tuple = (
                    filled.correct,
                    filled.incorrect,
                    filled.blank,
                    filled.total,
                )
                if require_check_counts:
                    if filled_tuple != EXPECTED_FILLED_CHECK:
                        raise ValueError(
                            f"filled Check counts {filled_tuple} != {EXPECTED_FILLED_CHECK}"
                        )
                elif filled.incorrect or filled.blank:
                    raise ValueError(
                        f"filled check incomplete: correct={filled.correct} "
                        f"incorrect={filled.incorrect} blank={filled.blank}"
                    )
                stages.append(
                    StageResult(
                        "7_filled_check",
                        "pass",
                        message=f"correct={filled.correct} total={filled.total}",
                    )
                )
            except Exception as exc:  # noqa: BLE001
                stages.append(
                    StageResult("7_filled_check", "fail", type(exc).__name__, str(exc))
                )
                if context["first_failure"] is None:
                    context["first_failure"] = stages[-1]
    except Exception as exc:  # noqa: BLE001
        stages.append(
            StageResult(
                "5_workbook_generation",
                "fail",
                type(exc).__name__,
                str(exc),
            )
        )
        if context["first_failure"] is None:
            context["first_failure"] = stages[-1]
        for name in ("6_blank_check", "7_filled_check"):
            stages.append(StageResult(name, "skipped", message="prior stage failed"))

    if explicit_pair:
        context["release_fingerprints_after"] = {
            str(p): _fingerprint(p) for p in release_files if p.is_file()
        }
        before = context["release_fingerprints_before"]
        after = context["release_fingerprints_after"]
        if before and after and before != after:
            stages.append(
                StageResult(
                    "8_release_pristine",
                    "fail",
                    "AssertionError",
                    "release artifacts mutated during verification",
                )
            )
            if context["first_failure"] is None:
                context["first_failure"] = stages[-1]
        elif before:
            stages.append(
                StageResult(
                    "8_release_pristine",
                    "pass",
                    message="release fingerprints unchanged",
                )
            )

    return {"stages": stages, "context": context}


def write_baseline(result: dict[str, Any]) -> None:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    provenance = json.loads(PROV_JSON.read_text(encoding="utf-8"))
    conflicts = (
        json.loads(CONFLICTS_JSON.read_text(encoding="utf-8"))
        if CONFLICTS_JSON.is_file()
        else {}
    )
    stages: list[StageResult] = result["stages"]
    ctx = result["context"]
    overlap = conflicts.get(
        "overlap_conflict_count", provenance.get("overlap_conflict_count", 0)
    )
    supplemental = conflicts.get(
        "supplemental_conflict_count",
        provenance.get("supplemental_conflict_count", 0),
    )
    lines = [
        "# Fast Retailing Benchmark Baseline (Step 9M.7)",
        "",
        "- Accounting engine phase: Step 9M.7 (deferred-tax balance diagnostics on 9M.6 base)",
        "- Input path: generic `extracted/` → `validate-source` → `reconcile` → `reconciled/`",
        "- Benchmark phase: measurement only — G1/G1B/G2/G2B/G2C/G3/G4/G5/G6/G7 closed",
        "- Five fiscal periods: 2021-08-31 … 2025-08-31",
        "",
        "## Source hashes",
        "",
    ]
    for src in manifest["sources"]:
        lines.append(
            f"- FY{src['fiscal_year']}: `{src['sha256']}` ({src['bytes']} bytes)"
        )
    lines.extend(
        [
            "",
            f"- Overlap conflicts recorded in conflicts.json: **{overlap}**",
            f"- Supplemental conflicts recorded in conflicts.json: **{supplemental}**",
            f"- Standardized payload: `{STD_JSON.relative_to(ROOT)}`",
            "- Supplemental provenance source-bound: yes",
            "- Portable source-path validation: yes",
            "- Silent repeated-share overwrite removed: yes",
            "- G1/G1B/G2/G2B/G2C/G3/G4/G5/G6/G7 closed; G7 disagreements retained with verified deterministic selection and provenance.",
            "",
            "## Stage results",
            "",
            "| Stage | Status | Detail |",
            "|---|---|---|",
        ]
    )
    for stage in stages:
        detail = stage.message or stage.exception_type
        detail = detail.replace("|", "\\|").replace("\n", " ")
        lines.append(f"| {stage.stage} | {stage.status} | {detail} |")

    first = ctx.get("first_failure")
    lines.extend(["", "## First failure", ""])
    if first is None:
        lines.append("None — all stages passed.")
    else:
        lines.append(f"- Stage: `{first.stage}`")
        lines.append(f"- Status: `{first.status}`")
        lines.append(f"- Exception: `{first.exception_type}`")
        lines.append(f"- Message: {first.message}")

    lines.extend(["", "## Module applicability (source fixture)", ""])
    for name, info in (ctx.get("modules") or {}).items():
        applicable = info.get("applicable")
        lines.append(f"- **{name}**: {'applicable' if applicable else 'omitted/not applicable'}")
        if name == "lease_liability":
            avail = info.get("availability") or {}
            lines.append(
                f"  - availability: lease_liability={avail.get('lease_liability')} "
                f"ambiguous={avail.get('ambiguous')}"
            )
        if name == "lease_rou":
            avail = info.get("availability") or {}
            lines.append(
                f"  - availability: right_of_use_assets={avail.get('right_of_use_assets')} "
                f"ambiguous={avail.get('ambiguous')}"
            )
        if name == "ownership_attribution":
            avail = info.get("availability") or {}
            lines.append(
                f"  - availability: available={avail.get('available')} "
                f"partial={avail.get('partial')} ambiguous={avail.get('ambiguous')}"
            )
        if name == "goodwill_intangibles":
            avail = info.get("availability") or {}
            lines.append(
                "  - availability: "
                f"goodwill={avail.get('goodwill')} "
                f"intangible_assets={avail.get('intangible_assets')} "
                f"goodwill_and_intangibles={avail.get('goodwill_and_intangibles')} "
                f"payments={avail.get('payments_for_intangible_assets')}"
            )
        if name == "deferred_tax":
            avail = info.get("availability") or {}
            lines.append(
                "  - availability: "
                f"deferred_tax_assets={avail.get('deferred_tax_assets')} "
                f"deferred_tax_liabilities={avail.get('deferred_tax_liabilities')} "
                f"ambiguous={avail.get('ambiguous')}"
            )

    lines.extend(
        [
            "",
            "## Notes",
            "",
            "- Temporary Trainer/Answer Key artifacts are not committed.",
            "- Synthetic DEMO / cross-company surfaces remain unchanged.",
            "- Forecasting / valuation remain deferred.",
            "- Comparable diluted-WAS axis (financial-statement units): "
            "306.871785, 306.969624, 307.13887, 307.231804, 307.247804 "
            "(basis=split_adjusted; FY2021 factor=3; FY2022–FY2025 factor=1).",
            "- Raw extracted share facts and G7 conflicts preserved unchanged.",
            "",
        ]
    )
    BASELINE.write_text("\n".join(lines), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    import argparse

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--standardized-json",
        type=Path,
        help="Generated standardized.json (default: benchmark reconciled fixture)",
    )
    parser.add_argument("--provenance-json", type=Path)
    parser.add_argument("--conflicts-json", type=Path)
    parser.add_argument("--trainer", type=Path, help="Persisted Trainer workbook")
    parser.add_argument("--answer-key", type=Path, help="Persisted Answer Key workbook")
    parser.add_argument(
        "--require-check-counts",
        action="store_true",
        help="Require pristine (0,0,577,577) and filled (577,0,0,577) Check counts",
    )
    parser.add_argument(
        "--verify-release-pair",
        action="store_true",
        help="Verify practice contract and structural parity on the persisted pair",
    )
    parser.add_argument(
        "--allow-frozen-yellow-answer-key",
        action="store_true",
        help=(
            "Narrow frozen-pair compatibility: permit historical yellow Answer-Key "
            "practice/judgment fills only when both workbook SHA-256 identities match "
            f"checkpoint {FROZEN_COMPATIBILITY_CHECKPOINT}. "
            "Current-generation, altered, mixed, and freshly generated pairs must not "
            "use this flag."
        ),
    )
    parser.add_argument(
        "--no-baseline",
        action="store_true",
        help="Skip writing benchmark/fast_retailing/BASELINE.md",
    )
    args = parser.parse_args(argv)

    result = run_audit(
        standardized_json=args.standardized_json,
        provenance_json=args.provenance_json,
        conflicts_json=args.conflicts_json,
        trainer_path=args.trainer,
        answer_key_path=args.answer_key,
        require_check_counts=True if args.require_check_counts else None,
        verify_release_pair=args.verify_release_pair,
        allow_frozen_yellow_answer_key=args.allow_frozen_yellow_answer_key,
    )
    if not args.no_baseline:
        write_baseline(result)
        print(f"wrote {BASELINE.relative_to(ROOT)}")
    for stage in result["stages"]:
        print(f"{stage.stage}: {stage.status}")
        if stage.message:
            print(f"  {stage.message[:300]}")
        if stage.exception_type and stage.status == "fail":
            print(f"  {stage.exception_type}: {stage.message[:300]}")
    failed = any(s.status == "fail" for s in result["stages"])
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
